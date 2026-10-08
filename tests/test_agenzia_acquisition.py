"""Behavioral regressions for the single local Agenzia acquisition workflow."""

from __future__ import annotations

import importlib
import importlib.util
import json
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "plugins/browser-automation/scripts"
PACKAGE = "test_vera_ade_acquisition"
spec = importlib.util.spec_from_file_location(
    PACKAGE,
    SCRIPTS / "ade_acquisition/__init__.py",
    submodule_search_locations=[str(SCRIPTS / "ade_acquisition")],
)
assert spec and spec.loader
package = importlib.util.module_from_spec(spec)
sys.modules[PACKAGE] = package
spec.loader.exec_module(package)
contracts = importlib.import_module(f"{PACKAGE}.contracts")
artifacts = importlib.import_module(f"{PACKAGE}.artifacts")
storage = importlib.import_module(f"{PACKAGE}.storage")
engine = importlib.import_module(f"{PACKAGE}.engine")
reports = importlib.import_module(f"{PACKAGE}.reports")
CLIENT = {
    "name": "Ditta esempio",
    "tax_code": "01234567890",
    "vat_number": "01234567890",
}


def plan_data(operations=None):
    return {
        "schema_version": "vera-agenzia-plan/v1",
        "clients": [CLIENT.copy()],
        "date_from": "2026-01-01",
        "date_to": "2026-01-31",
        "operations": operations or ["ricevute"],
    }


def invoice(number="1", customer="01234567890", supplier="09876543210"):
    return f"""<FatturaElettronica><FatturaElettronicaHeader><CedentePrestatore><DatiAnagrafici><IdFiscaleIVA><IdCodice>{supplier}</IdCodice></IdFiscaleIVA></DatiAnagrafici></CedentePrestatore><CessionarioCommittente><DatiAnagrafici><IdFiscaleIVA><IdCodice>{customer}</IdCodice></IdFiscaleIVA></DatiAnagrafici></CessionarioCommittente></FatturaElettronicaHeader><FatturaElettronicaBody><DatiGenerali><DatiGeneraliDocumento><TipoDocumento>TD01</TipoDocumento><Data>2026-01-02</Data><Numero>{number}</Numero></DatiGeneraliDocumento></DatiGenerali></FatturaElettronicaBody></FatturaElettronica>""".encode()


def der(tag, data):
    size = len(data)
    length = (
        bytes([size])
        if size < 128
        else bytes([128 + (size.bit_length() + 7) // 8])
        + size.to_bytes((size.bit_length() + 7) // 8, "big")
    )
    return bytes([tag]) + length + data


def cms(content):
    encapsulated = der(
        0x30, der(6, bytes.fromhex("2a864886f70d010701")) + der(0xA0, der(4, content))
    )
    signed = der(0x30, der(2, b"\x01") + der(0x31, b"") + encapsulated + der(0x31, b""))
    return der(0x30, der(6, bytes.fromhex("2a864886f70d010702")) + der(0xA0, signed))


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("1.234,56 €", Decimal("1234.56")),
        ("0,00", Decimal("0.00")),
        ("-2,10", Decimal("-2.10")),
        ("—", None),
        (None, None),
    ],
)
def test_money_preserves_exact_amount_and_absence(value, expected):
    assert contracts.money(value) == expected


@pytest.mark.parametrize("value", ["1,2,3", "EUR due", "1.20", "NaN", "2e4"])
def test_money_rejects_unknown_formats(value):
    with pytest.raises(contracts.AcquisitionError, match="monetary-value-invalid"):
        contracts.money(value)


@pytest.mark.parametrize(
    "mutation",
    [
        {"password": "never-store"},
        {"operations": ["ricevute", "ricevute"]},
        {"date_to": "2025-01-01"},
        {"clients": [CLIENT, CLIENT]},
        {"access_mode": "incaricato"},
        {"max_documents": True},
    ],
)
def test_plan_rejects_unsafe_or_ambiguous_scope(mutation):
    with pytest.raises(contracts.AcquisitionError):
        contracts.Plan.parse({**plan_data(), **mutation})


def test_quarters_partition_exact_inclusive_dates():
    assert contracts.quarters(date(2025, 12, 31), date(2026, 4, 1)) == [
        (2025, 4, date(2025, 12, 31), date(2025, 12, 31)),
        (2026, 1, date(2026, 1, 1), date(2026, 3, 31)),
        (2026, 2, date(2026, 4, 1), date(2026, 4, 1)),
    ]


def test_client_csv_preserves_leading_zero_and_same_name_clients(tmp_path):
    path = tmp_path / "clients.csv"
    path.write_text(
        "NOMINATIVO CLIENTE;CODICE FISCALE;PARTITA IVA\nUguale;01234567890;01234567890\nUguale;09876543210;09876543210\n"
    )
    result = contracts.read_clients(path)
    assert [c["vat_number"] for c in result] == ["01234567890", "09876543210"]


def test_client_xlsx_restores_numeric_vat_leading_zero(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["CLIENTE", "CODICE FISCALE", "PARTITA IVA"])
    sheet.append(["Esempio", 1234567890, 1234567890])
    path = tmp_path / "clients.xlsx"
    workbook.save(path)
    result = contracts.read_clients(path)
    assert result[0]["tax_code"] == result[0]["vat_number"] == "01234567890"


def test_client_file_rejects_incomplete_row_instead_of_skipping(tmp_path):
    path = tmp_path / "clients.csv"
    path.write_text("CLIENTE;CODICE FISCALE;PARTITA IVA\nEsempio;01234567890;\n")
    with pytest.raises(contracts.AcquisitionError, match="client-row-2"):
        contracts.read_clients(path)


def test_signed_original_retained_and_xml_extracted_without_signature_claim():
    original = cms(invoice())
    metadata, extracted = artifacts.inspect_invoice(
        original, "fattura.xml.p7m", contracts.Client.parse(CLIENT), "ricevute"
    )
    assert extracted == invoice()
    assert metadata["signature_validation"] == "not_performed"
    assert metadata["byte_length"] == len(original)


@pytest.mark.parametrize(
    ("data", "name", "code"),
    [
        (b"<html>login</html>", "file.xml", "download-not-fatturapa"),
        (invoice(customer="11111111111"), "file.xml", "invoice-client-mismatch"),
        (b"metadata", "file.xml", "download-not-invoice-xml"),
        (b"\x30\x80", "file.p7m", "p7m-invalid-structure"),
        (
            b'<!DOCTYPE a [<!ENTITY xx SYSTEM "file:///etc/passwd">]><FatturaElettronica>&xx;</FatturaElettronica>',
            "file.xml",
            "download-not-invoice-xml",
        ),
    ],
)
def test_download_acceptance_rejects_wrong_or_unsafe_artifact(data, name, code):
    with pytest.raises(contracts.AcquisitionError, match=code):
        artifacts.inspect_invoice(
            data, name, contracts.Client.parse(CLIENT), "ricevute"
        )


class PortalFixture:
    """An observable source population supporting exact date splits and page restoration."""

    def __init__(self, count=2, *, mismatch=False, unavailable=False, cash_count=2):
        self.count, self.mismatch, self.unavailable = count, mismatch, unavailable
        self.cash_count = cash_count
        self.downloads = 0
        self.page = 1
        self.selected = []
        self.active = []

    def select_client(self, plan, client):
        self.selected.append(client.vat_number)

    def search_invoices(self, client, category, start, end):
        self.active = [
            i
            for i in range(self.count)
            if start <= date(2026, 1, 2 if i % 2 == 0 else 20) <= end
        ]
        self.page = 1
        return len(self.active)

    def invoice_rows(self):
        items = self.active[(self.page - 1) * 50 : self.page * 50]
        return [{"key": str(item), "index": index} for index, item in enumerate(items)]

    def download_invoice(self, row):
        self.downloads += 1
        if self.unavailable:
            raise contracts.AcquisitionError("invoice-original-unavailable")
        return "same-name.xml", invoice(
            row["key"],
            customer="11111111111" if self.mismatch else CLIENT["vat_number"],
        )

    def restore_invoices(self, *args):
        return None

    def page_fingerprint(self):
        return str(self.page)

    def next_page(self, number, fingerprint):
        self.page = number

    def cash_devices(self, client, start, end):
        return [
            {"name": "RT", "count": self.cash_count, "index": 0},
            {"name": "Altro", "count": 0, "index": 1},
        ]

    def cash_rows(self, client, start, end, device):
        return [
            {
                "Id invio": str(index),
                "Matricola dispositivo": "RT001",
                "Data e ora rilevazione": "02/01/2026 12:00:00",
                "Data e ora trasmissione": "02/01/2026 12:01:00",
                "Ammontare vendite": "12,20",
                "Imponibile vendite": "10,00",
                "Imposta vendite": "2,20",
            }
            for index in range(device["count"])
        ]

    def stamp_duty(self, client, year, quarter):
        return {
            "status": "acquired",
            "amount": "2.00",
            "payment_status": "pagato",
            "list_a_count": 1,
            "list_b_count": 0,
            "document_count": 1,
        }


def execute(tmp_path, portal, operations=None):
    plan = contracts.Plan.parse(plan_data(operations))
    store = storage.Store(tmp_path.resolve() / "acquisition")
    run = store.new_run(plan.payload())
    outcome = engine.Acquisition(store, run, plan, portal).execute()
    return store, run, plan, outcome


def test_repeat_run_verifies_existing_files_and_repairs_deleted_original(tmp_path):
    portal = PortalFixture()
    store, run, plan, first = execute(tmp_path, portal)
    original = store.safe_path(
        engine.load_events(run)[0]["record"]["artifacts"][0]["path"]
    )
    original.unlink()
    resumed = store.new_run(plan.payload())
    result = engine.Acquisition(store, resumed, plan, portal).execute()
    assert first["downloaded"] == 2
    assert result["state"] == "complete"
    assert result["downloaded"] == 1
    assert result["verified_existing"] == 1
    assert portal.downloads == 3


def test_date_splitting_and_same_day_pagination_reconcile_all_originals(tmp_path):
    portal = PortalFixture(count=102)
    store, run, plan, result = execute(tmp_path, portal)
    assert result["state"] == "complete"
    assert result["downloaded"] == 102
    paths = [
        event["record"]["artifacts"][0]["path"]
        for event in engine.load_events(run)
        if event["kind"] == "invoice"
    ]
    assert len(set(paths)) == 102
    assert all(store.safe_path(path).name == "0_same-name.xml" for path in paths)


@pytest.mark.parametrize("options", [{"mismatch": True}, {"unavailable": True}])
def test_unusable_originals_never_mark_scope_complete(tmp_path, options):
    store, run, plan, result = execute(tmp_path, PortalFixture(**options))
    assert result["state"] == "partial"
    assert result["downloaded"] == 0
    assert result["failed_documents"] == 2


def test_cancel_retains_run_and_reports_unattempted_scopes(tmp_path):
    plan = contracts.Plan.parse(plan_data(["ricevute", "bolli"]))
    store = storage.Store(tmp_path.resolve() / "private")
    run = store.new_run(plan.payload())
    storage.atomic_json(run / "cancel.json", {"cancel": True})
    result = engine.Acquisition(store, run, plan, PortalFixture()).execute()
    assert result["state"] == "cancelled"
    assert result["scopes_complete"] == 0


def test_duplicate_portal_page_is_incomplete_not_skipped(tmp_path):
    class Repeating(PortalFixture):
        def invoice_rows(self):
            return [{"key": "same", "index": 0}]

    store, run, plan, result = execute(tmp_path, Repeating())
    assert result["state"] == "partial"
    assert result["scopes"][0]["error"] == "portal-invoice-rows-repeated"


def test_corrupted_file_is_not_an_incremental_success(tmp_path):
    store, run, plan, result = execute(tmp_path, PortalFixture(count=1))
    event = engine.load_events(run)[0]
    artifact = store.safe_path(event["record"]["artifacts"][0]["path"])
    artifact.write_bytes(b"corruption")
    assert store.lookup(plan.clients[0], "ricevute", event["portal_key"]) is None


def test_symlink_or_repository_output_is_rejected(tmp_path):
    root = tmp_path.resolve()
    (root / "repository").mkdir()
    (root / "repository/.git").mkdir()
    with pytest.raises(contracts.AcquisitionError, match="outside-git"):
        storage.Store(root / "repository/output")
    (root / "link").symlink_to(root / "repository")
    with pytest.raises(contracts.AcquisitionError, match="outside-git"):
        storage.Store(root / "link/output")


def test_worker_lock_rejects_concurrent_access(tmp_path):
    store = storage.Store(tmp_path.resolve() / "private")
    with store.lock():
        with pytest.raises(contracts.AcquisitionError, match="already-running"):
            with store.lock():
                pytest.fail("Second worker acquired the same archive")


def test_full_acquisition_writes_reviewable_reports(tmp_path):
    from openpyxl import load_workbook

    store, run, plan, result = execute(
        tmp_path, PortalFixture(), ["ricevute", "corrispettivi", "bolli"]
    )
    output = reports.write_reports(run, plan, result)
    workbook = load_workbook(run / "riepilogo.xlsx", data_only=False)
    assert result["scopes_complete"] == 3
    assert output["cash_review_issue_count"] == 0
    assert workbook["Corrispettivi giornalieri"]["E2"].value == 24.4
    assert workbook["Bolli e prospetto F24"]["J2"].value == "pagato"
    assert (run / "report.html").is_file()
    assert (run / "model_data_report.md").is_file()
    assert not any(
        cell.data_type == "f" for sheet in workbook for row in sheet for cell in row
    )


def test_missing_cash_value_and_duplicate_id_are_visible_not_zero(tmp_path):
    class IncompleteCash(PortalFixture):
        def cash_rows(self, *args):
            rows = super().cash_rows(*args)
            for row in rows:
                row["Imposta vendite"] = "—"
                row["Id invio"] = "same"
            return rows

    store, run, plan, result = execute(tmp_path, IncompleteCash(), ["corrispettivi"])
    records, issues = reports.cash_records(engine.load_events(run))
    assert records[0]["vat"] is None
    assert "vat-missing" in records[0]["issues"]
    assert "submission-id-repeated" in records[1]["issues"]


def test_f24_uses_only_reviewed_amount_and_supports_fourth_quarter(tmp_path):
    from pypdf import PdfReader

    payload = {
        **plan_data(["bolli"]),
        "date_from": "2026-10-01",
        "date_to": "2026-12-31",
    }
    plan = contracts.Plan.parse(payload)
    store = storage.Store(tmp_path.resolve() / "private")
    run = store.new_run(plan.payload())
    engine.Acquisition(store, run, plan, PortalFixture()).execute()
    reviewed = {
        "schema_version": "vera-agenzia-f24/v1",
        "reviewed": True,
        "items": [
            {
                "client": CLIENT["vat_number"],
                "year": 2026,
                "quarter": 4,
                "amount": "1.00",
                "due_date": "2027-03-01",
                "note": "Residuo verificato dal professionista",
            }
        ],
    }
    pdf = reports.prepare_f24(run, plan, reviewed)
    text = " ".join(page.extract_text() for page in PdfReader(pdf).pages)
    assert "2524" in text
    assert "1.00" in text
    assert "BOZZA" in text
    assert (
        json.loads(pdf.with_suffix(".json").read_text())["items"][0]["source_amount"]
        == "2.00"
    )


def test_f24_rejects_unreviewed_or_missing_source(tmp_path):
    store, run, plan, result = execute(tmp_path, PortalFixture(), ["ricevute"])
    with pytest.raises(contracts.AcquisitionError, match="explicit-review"):
        reports.prepare_f24(run, plan, {"reviewed": False})
