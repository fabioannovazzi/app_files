from __future__ import annotations

import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from src.slides.semantic_pptx import (
    SLIDES_PPTX_SPEC_FILENAME,
    render_slides_pptx_from_template,
)


@pytest.mark.parametrize("position", ["right", "bottom"])
def test_text_and_visual_remain_separate_and_preserve_financial_caveat(
    tmp_path: Path, position: str
) -> None:
    body = "Margin fell 2 percentage points. The forecast excludes financing costs."
    Image.new("RGB", (640, 360), "navy").save(tmp_path / "evidence.png")
    payload = {
        "templateKey": "uniform",
        "promptStyle": "uniform",
        "slides": [
            {
                "slideId": "slide0.html",
                "kind": "text_visual",
                "layoutVariant": f"text_visual_{position}",
                "title": "Forecast limitations",
                "body": body,
                "visualPath": "evidence.png",
                "visualType": "figure",
            }
        ],
    }
    (tmp_path / SLIDES_PPTX_SPEC_FILENAME).write_text(json.dumps(payload))

    result = render_slides_pptx_from_template(tmp_path)

    presentation = Presentation(result)
    slide = presentation.slides[0]
    text_boxes = [
        shape for shape in slide.shapes if shape.has_text_frame and shape.text == body
    ]
    pictures = [shape for shape in slide.shapes if int(shape.shape_type) == 13]
    assert len(text_boxes) == 1
    assert len(pictures) == 1
    text, picture = text_boxes[0], pictures[0]
    assert (
        text.left + text.width <= picture.left or text.top + text.height <= picture.top
    )
    assert picture.left >= 0
    assert picture.top >= 0
    assert picture.left + picture.width <= presentation.slide_width
    assert picture.top + picture.height <= presentation.slide_height
