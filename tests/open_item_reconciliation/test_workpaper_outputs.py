from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from scripts.validate_plugin_review_contract import validate_contract

SCRIPTS = (
    Path(__file__).resolve().parents[2]
    / "plugins"
    / "open-item-reconciliation"
    / "scripts"
)
WORKFLOW = SCRIPTS / "reconciliation_workflow.py"


def load_outputs():
    spec = importlib.util.spec_from_file_location(
        "audit_workpaper_outputs_entrypoint",
        WORKFLOW,
    )
    assert spec and spec.loader
    entrypoint = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = entrypoint
    spec.loader.exec_module(entrypoint)
    return sys.modules["workpaper_outputs"]


def load_review_session():
    spec = importlib.util.spec_from_file_location(
        "audit_workpaper_review_session_entrypoint",
        WORKFLOW,
    )
    assert spec and spec.loader
    entrypoint = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = entrypoint
    spec.loader.exec_module(entrypoint)
    return sys.modules["review_session"]


def document_text(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


def test_partial_workpaper_writes_without_inventing_missing_aggregate_origins(tmp_path):
    outputs = load_outputs()
    path = tmp_path / "aggregate-only.xlsx"
    cells = []

    outputs.write_excel_workpaper(
        path,
        {
            "External evidence aggregate": [
                {"external_category": "TOTAL", "amount_total": "123.00"}
            ]
        },
        projection_cells=cells,
    )

    workbook = load_workbook(path, read_only=True)
    try:
        assert workbook["External evidence aggregate"]["A2"].value == "123.00"
    finally:
        workbook.close()
    assert cells[1][0] is None


@pytest.mark.parametrize("field", ["review_notes", "unregistered_custom_field"])
def test_canonical_report_keeps_supplied_row_text_unverified(field):
    outputs = load_outputs()
    source = {"record_id": "invoice-1", "amount": "1200.00"}
    payload = {
        "schema_version": "vera.reconciliation_results.v1",
        "client_engagement": None,
        "source_processing": {},
        "reconciliation_rows": [source],
        "allocation_ledgers": [],
        "analyses": {},
        "checks": [],
        "review_rows": [{**source, field: "PRIVATE_SUPPLIED_TEXT"}],
        "source_qualifications": [],
    }
    context = {"sheets": {"Normalized records": [source]}, "allocation_projections": {}}

    origins = outputs.canonical_report_origins(payload, context)

    assert origins is None


def test_build_audit_workbook_sheets_has_standard_tabs():
    outputs = load_outputs()
    sheets = outputs.build_audit_workbook_sheets(
        assumptions={"cutoff_date": "2023-12-31"},
        source_inventory=[{"source_file": "ledger.pdf"}],
        normalized_records=[{"record_id": "n1"}],
        reconciliation_rows=[
            {
                "record_id": "open-1",
                "amount": "100.00",
                "reconciliation_status": "closed",
                "rule_applied": "direct_external_or_documented",
            }
        ],
        checks=[{"check": "row_count", "status": "PASS"}],
        bank_allocation_candidates=[
            {
                "candidate_id": "candidate-1",
                "candidate_type": "unallocated_counterparty_bank_pool",
            }
        ],
        external_evidence_summary=[
            {
                "external_category": "TOTAL",
                "settlement_effect_signed_net_debit_minus_credit": "-180.00",
            }
        ],
        external_evidence_detail=[
            {
                "record_id": "bank-direct",
                "external_category": "direct_counterparty_bank_receipt",
            }
        ],
        ledger_balance_rows=[
            {"account": "TOTAL", "closing_balance_signed_debit_minus_credit": "550.00"}
        ],
        account_rollforward_check=[{"account": "TOTAL", "status": "PASS"}],
        aggregate_rollforward_summary=[
            {"account": "TOTAL", "closing_net_debit_minus_credit": "550.00"}
        ],
        aggregate_rollforward_rows=[{"record_id": "journal_rollforward:1"}],
        post_cutoff_candidates=[{"candidate_id": "post-cutoff-1"}],
        aging_summary=[{"aging_bucket": "181-365", "rows": 1}],
        evidence_concentration=[{"support_bucket": "bank", "rows": 1}],
        review_signals=[{"record_id": "open-1", "review_signal_rank": 1}],
        document_source_map=[{"document_key": "INV1|2023", "open_item_rows": 1}],
        reversal_candidates=[{"candidate_id": "reversal-1"}],
        cutoff_window_movements=[{"record_id": "cutoff-1"}],
        review_rows=[{"record_id": "open-1", "review_result": "PASS"}],
    )

    assert list(sheets) == outputs.STANDARD_SHEET_ORDER
    assert sheets["Summary"][0]["rows"] == 1
    assert sheets["Bank allocation candidates"][0]["candidate_id"] == "candidate-1"
    assert (
        sheets["External evidence aggregate"][0][
            "settlement_effect_signed_net_debit_minus_credit"
        ]
        == "-180.00"
    )
    assert sheets["External evidence detail"][0]["record_id"] == "bank-direct"
    assert (
        sheets["Ledger balance check"][0]["closing_balance_signed_debit_minus_credit"]
        == "550.00"
    )
    assert sheets["Account rollforward check"][0]["status"] == "PASS"
    assert (
        sheets["Journal rollforward"][0]["closing_net_debit_minus_credit"] == "550.00"
    )
    assert sheets["Journal detail"][0]["record_id"] == "journal_rollforward:1"
    assert sheets["Post-cutoff candidates"][0]["candidate_id"] == "post-cutoff-1"
    assert sheets["Open item aging"][0]["aging_bucket"] == "181-365"
    assert sheets["Evidence concentration"][0]["support_bucket"] == "bank"
    assert sheets["Review signals"][0]["review_signal_rank"] == 1
    assert sheets["Document source map"][0]["document_key"] == "INV1|2023"
    assert sheets["Reversal candidates"][0]["candidate_id"] == "reversal-1"
    assert sheets["Cutoff window movements"][0]["record_id"] == "cutoff-1"
    assert sheets["Index"][0]["Sheet"] == "Assumptions"


def test_write_excel_workpaper_creates_auditable_tabs(tmp_path):
    outputs = load_outputs()
    path = tmp_path / "audit.xlsx"
    sheets = outputs.build_audit_workbook_sheets(
        assumptions={"cutoff_date": "2023-12-31"},
        source_inventory=[],
        normalized_records=[],
        reconciliation_rows=[],
        checks=[],
        review_rows=[],
    )

    outputs.write_excel_workpaper(path, sheets)

    workbook = load_workbook(path)
    assert workbook.sheetnames == outputs.STANDARD_SHEET_ORDER
    assert workbook["Index"].freeze_panes == "A2"


def test_write_excel_workpaper_adds_italian_presentation_tabs_without_removing_raw(
    tmp_path,
):
    outputs = load_outputs()
    path = tmp_path / "audit.xlsx"
    sheets = outputs.build_audit_workbook_sheets(
        assumptions={"factoring_pro_soluto_closes_item": True},
        source_inventory=[{"source_role": "bank_statement"}],
        normalized_records=[],
        reconciliation_rows=[
            {
                "record_id": "open-1",
                "amount": "100.00",
                "reconciliation_status": "needs_evidence",
                "rule_applied": "internal_closure_without_external",
                "matched_evidence_type": "internal_closure",
            }
        ],
        checks=[],
        review_rows=[],
    )

    outputs.write_excel_workpaper(path, sheets, language="it")

    workbook = load_workbook(path)
    assert "Dettaglio riconciliazione" in workbook.sheetnames
    assert "Reconciliation detail" in workbook.sheetnames
    assert workbook["Reconciliation detail"].sheet_state == "hidden"
    localized_sheet = workbook["Dettaglio riconciliazione"]
    headers = [cell.value for cell in localized_sheet[1]]
    values = [cell.value for cell in localized_sheet[2]]
    assert "Esito riconciliazione" in headers
    assert "Regola applicata" in headers
    assert "Serve evidenza aggiuntiva" in values
    assert "Chiusura interna senza evidenza esterna" in values
    assert "needs_evidence" not in {str(value) for value in values}
    assert "internal_closure_without_external" not in {str(value) for value in values}
    assumption_values = [
        cell.value
        for row in workbook["Assunzioni"].iter_rows(values_only=False)
        for cell in row
    ]
    assert (
        "Factoring pro-soluto chiude con evidenza esterna allocata" in assumption_values
    )
    assert "factoring_pro_soluto_closes_item" not in {
        str(value) for value in assumption_values
    }


def test_spanish_workpapers_and_review_contract_are_fully_localized(tmp_path):
    outputs = load_outputs()
    excel_path = tmp_path / "conciliacion_auditoria.xlsx"
    word_path = tmp_path / "informe_conciliacion.docx"
    assumptions = {
        "scope_year": "2024",
        "cutoff_date": "2024-12-31",
        "currency": "EUR",
        "factoring_pro_soluto_closes_item": True,
        "post_cutoff_events_excluded": True,
    }
    reconciliation_rows = [
        {
            "record_id": "open-1",
            "document_no": "FAC-001",
            "amount": "100.00",
            "reconciliation_status": "needs_evidence",
            "rule_applied": "internal_closure_without_external",
            "matched_evidence_type": "internal_closure",
        }
    ]
    checks = [{"check": "review_packet", "status": "WARN"}]
    review_rows = [
        {
            "record_id": "open-1",
            "review_status": "PENDING",
            "source_file": "diario.xlsx",
            "source_page": 2,
            "source_row": 14,
        }
    ]
    sheets = outputs.build_audit_workbook_sheets(
        assumptions=assumptions,
        source_inventory=[{"source_role": "bank_statement"}],
        normalized_records=[],
        reconciliation_rows=reconciliation_rows,
        checks=checks,
        review_rows=review_rows,
        language="es",
    )

    outputs.write_excel_workpaper(excel_path, sheets, language="es-ES")
    outputs.write_word_report(
        word_path,
        title="Informe de conciliación de auditoría",
        metadata={"scope_year": "2024", "cutoff_date": "2024-12-31"},
        summary_rows=sheets["Summary"],
        assumptions=assumptions,
        next_steps=["Obtener la evidencia pendiente."],
        source_inventory=[{"source_role": "bank_statement"}],
        checks=checks,
        review_rows=review_rows,
        language="spa",
    )

    workbook = load_workbook(excel_path)
    assert "Índice" in workbook.sheetnames
    assert "Supuestos" in workbook.sheetnames
    assert "Detalle de conciliación" in workbook.sheetnames
    assert "Revisión Codex" in workbook.sheetnames
    assert workbook["Reconciliation detail"].sheet_state == "hidden"
    detail_values = {
        str(cell.value)
        for row in workbook["Detalle de conciliación"].iter_rows()
        for cell in row
        if cell.value is not None
    }
    assert "Estado de conciliación" in detail_values
    assert "Regla aplicada" in detail_values
    assert "Requiere evidencia adicional" in detail_values
    assert "Cierre interno sin evidencia externa" in detail_values
    assert "Asiento de cierre" in detail_values
    assert "needs_evidence" not in detail_values
    assert "internal_closure_without_external" not in detail_values
    assumption_values = {
        str(cell.value)
        for row in workbook["Supuestos"].iter_rows()
        for cell in row
        if cell.value is not None
    }
    assert "Año del alcance" in assumption_values
    assert (
        "El factoring sin recurso cierra con evidencia externa asignada"
        in assumption_values
    )
    assert "Sí" in assumption_values

    report = document_text(word_path)
    for expected in (
        "Resumen ejecutivo",
        "Papel de trabajo de conciliación determinista",
        "Alcance y método",
        "Cómo interpretar los resultados",
        "Requiere evidencia adicional",
        "Cierre interno sin evidencia externa",
        "Extracto bancario",
        "Controles automáticos",
        "Aviso",
        "Revisión manual de Codex",
        "Pendiente",
        "Limitaciones del procedimiento",
        "Referencia al archivo Excel",
    ):
        assert expected in report
    for internal_or_english in (
        "needs_evidence",
        "internal_closure_without_external",
        "bank_statement",
        "Deterministic reconciliation workpaper",
        "Scope and Method",
        "Procedure Limits",
    ):
        assert internal_or_english not in report

    review_session = load_review_session()
    run_intake = review_session.write_run_intake(
        tmp_path,
        assumptions=assumptions,
        language="spa",
        source_hint="auditoria-es",
    )
    result = {
        "excel_path": excel_path,
        "word_path": word_path,
        "assumptions": assumptions,
        "reconciliation_rows": reconciliation_rows,
        "checks": checks,
        "review_rows": review_rows,
        "checks_pass": False,
    }
    review_result = review_session.write_review_session_artifacts(
        tmp_path,
        run_id=run_intake.run_id,
        run_intake_path=run_intake.path,
        result=result,
        language="spa",
    )
    run_intake_payload = json.loads(run_intake.path.read_text(encoding="utf-8"))
    review_payload = json.loads(
        review_result.review_payload_path.read_text(encoding="utf-8")
    )
    final_artifacts = json.loads(
        review_result.final_artifacts_path.read_text(encoding="utf-8")
    )
    assert review_payload["columns"] == [
        {"field": "item_type", "label": "Tipo"},
        {"field": "title", "label": "Línea o artefacto"},
        {"field": "recommended_action", "label": "Acción sugerida"},
        {"field": "source_path", "label": "Fuente"},
        {"field": "output_path", "label": "Salida"},
        {"field": "status", "label": "Estado"},
    ]
    review_row_item = next(
        item for item in review_payload["items"] if item["id"] == "open-1"
    )
    assert review_row_item["source_path"] == "diario.xlsx; página 2; fila 14"
    assert review_row_item["data"]["edit_hint"].startswith(
        "Editar esta línea de revisión"
    )
    artifact_titles = {
        item["title"]
        for item in review_payload["items"]
        if item["item_type"] in {"workpaper_artifact", "report_artifact"}
    }
    assert artifact_titles == {
        "Libro de conciliación de auditoría",
        "Informe narrativo de conciliación",
    }
    assert run_intake_payload["data_posture"]["notes"] == [
        "Los scripts de conciliación leen las rutas locales de evidencias contables registradas en input_paths.",
        "Los datos de revisión muestran líneas de conciliación y referencias de evidencias acotadas para su revisión en la interfaz.",
        "De forma predeterminada no se utiliza ningún conector externo, ruta de carga, SQL remoto ni cuaderno alojado.",
    ]
    spanish_review_contract = json.dumps(
        {"run_intake": run_intake_payload, "review_payload": review_payload},
        ensure_ascii=False,
    )
    for english_copy in (
        "Row or artifact",
        "Suggested action",
        "Editing this review row",
        "Audit reconciliation workbook",
        "Narrative reconciliation report",
        "Review payloads expose bounded",
    ):
        assert english_copy not in spanish_review_contract
    artifact_card = review_result.artifact_card_path.read_text(encoding="utf-8")
    assert "# Ficha de artefactos de la conciliación de auditoría" in artifact_card
    assert "## Siguiente acción" in artifact_card
    assert "Open-item Reconciliation Artifact Card" not in artifact_card
    assert "Los datos mostrados en el navegador están acotados" in " ".join(
        final_artifacts["caveats"]
    )
    assert "Abra el servidor de revisión" in final_artifacts["next_actions"][0]
    assert "Open the browser review server" not in json.dumps(
        final_artifacts, ensure_ascii=False
    )
    audit_output = next(
        output
        for output in final_artifacts["outputs"]
        if output.get("artifact_role") == "audit_workpaper"
    )
    word_output = next(
        output
        for output in final_artifacts["outputs"]
        if output.get("artifact_role") == "word_report"
    )
    assert audit_output["required_sheets"] == [
        "Índice",
        "Supuestos",
        "Problemas de fuentes",
        "Detalle de conciliación",
        "Resumen",
        "Controles",
        "Revisión Codex",
    ]
    assert audit_output["required_sheet_headers"] == {
        "Índice": ["Hoja", "Líneas"],
        "Supuestos": ["Campo", "Valor"],
    }
    assert audit_output["required_cells"]["Supuestos"]["A2"] == "Año del alcance"
    assert word_output["required_text"][0] == "Resumen ejecutivo"
    assert word_output["required_text"][-1] == "Referencia al archivo Excel"
    contract_report = validate_contract(tmp_path, strict_output_content=True)
    artifact_errors = [
        error
        for error in contract_report.errors
        if excel_path.name in error or word_path.name in error
    ]
    assert artifact_errors == []


def test_write_word_report_defaults_to_italian_labels(tmp_path):
    outputs = load_outputs()
    path = tmp_path / "relazione.docx"

    outputs.write_word_report(
        path,
        title="Relazione di riconciliazione",
        metadata={"Periodo": "2023"},
        summary_rows=[{"Esito": "Chiuso", "Righe": 1}],
        assumptions={"cutoff_date": "2023-12-31"},
        next_steps=["Acquisire documentazione mancante."],
        narrative="Sintesi tecnica.",
    )

    text = document_text(path)
    assert "Conclusioni" in text
    assert "Assunzioni" in text
    assert "Prossimi passi" in text


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_report_keeps_review_diagnostics_in_canonical_evidence(tmp_path, language):
    import copy

    outputs = load_outputs()
    assumptions = json.loads(
        (
            Path(__file__).parent
            / "fixtures"
            / "reviewed_settlement_case/reviewed-assumptions.json"
        ).read_text()
    )
    assumptions["professional_note"] = "Valutare separatamente la nota di credito."
    before = copy.deepcopy(assumptions)
    path = tmp_path / "report.docx"
    outputs.write_word_report(
        path,
        title="Report",
        metadata={},
        summary_rows=[],
        assumptions=assumptions,
        next_steps=[],
        language=language,
    )
    text = document_text(path)
    assert assumptions == before
    assert "Valutare separatamente la nota di credito." in text
    assert "source_path_keys" not in text
    assert "/studio-private/" not in text
    assert "unsigned synthetic agent test" not in text
    assert "source_review_notes" not in outputs._report_assumptions(assumptions)
    assert outputs._report_assumptions(assumptions)["professional_approval"] is False
    if language == "it":
        assert "Tolleranza sugli importi: 0" in text
        assert "Conferma umana: No" in text
        assert "Approvazione professionale: No" in text
        assert "Prova sintetica automatizzata" in text
        for key in [
            "Jurisdiction",
            "Unit:",
            "Entity ref",
            "Party ref",
            "Amount tolerance",
            "Human confirmation",
            "Professional approval",
            "Source review notes",
            "agent_test",
        ]:
            assert key not in text


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_report_renders_owned_settlement_labels_without_rewriting_source_text(
    tmp_path, language
):
    outputs = load_outputs()
    categories = [
        "other_external",
        "factor_or_operator_debit_fee",
        "bank_advance_repayment",
        "bank_advance_credit",
        "factor_operator_cash_inflow",
        "direct_counterparty_bank_receipt",
        "direct_counterparty_bank_payment",
        "counterparty_bank_other",
        "factor_operator_other",
    ]
    path = tmp_path / "report.docx"
    original_note = (
        "Source note: other_external; invoice exact_document_partial_bank_payment"
    )
    outputs.write_word_report(
        path,
        title="Report",
        metadata={},
        summary_rows=[
            {
                "reconciliation_status": "partially_paid",
                "rule_applied": "exact_document_partial_bank_payment",
                "rows": 1,
                "amount": "600.00",
            }
        ],
        assumptions={},
        next_steps=[],
        narrative=original_note,
        external_evidence_summary=[
            {
                "external_category": category,
                "rows": 1,
                "amount_total": "100.00",
            }
            for category in categories
        ],
        language=language,
    )
    text = document_text(path)
    assert original_note in text
    table_cells = [
        cell.text
        for table in Document(path).tables
        for row in table.rows
        for cell in row.cells
    ]
    assert all(category not in table_cells for category in categories)
    assert "exact_document_partial_bank_payment" not in table_cells
    assert "Exact document partial bank payment" not in table_cells
    if language == "it":
        assert "Altra evidenza esterna" in table_cells
        assert (
            "Pagamento bancario parziale con riferimento esatto al documento"
            in table_cells
        )


def test_write_word_report_localizes_italian_internal_codes(tmp_path):
    outputs = load_outputs()
    path = tmp_path / "relazione.docx"

    outputs.write_word_report(
        path,
        title="Relazione di riconciliazione",
        metadata={},
        summary_rows=[
            {
                "reconciliation_status": "needs_evidence",
                "rule_applied": "internal_closure_without_external",
                "rows": 1,
                "amount": "100.00",
            }
        ],
        assumptions={
            "factoring_pro_soluto_closes_item": True,
            "post_cutoff_events_excluded": True,
        },
        next_steps=["Acquisire evidenze aggiuntive."],
        source_inventory=[{"source_role": "bank_statement"}],
        checks=[{"check": "review_packet", "status": "WARN"}],
        review_rows=[{"review_status": "PENDING"}],
    )

    text = document_text(path)
    assert "Serve evidenza aggiuntiva" in text
    assert "Chiusura interna senza evidenza esterna" in text
    assert "Factoring pro-soluto chiude con evidenza esterna allocata: Sì" in text
    assert "Eventi post cut-off esclusi dalla chiusura: Sì" in text
    assert "Estratto conto bancario" in text
    assert "needs_evidence" not in text
    assert "internal_closure_without_external" not in text
    assert "factoring_pro_soluto_closes_item" not in text
    assert "post_cutoff_events_excluded" not in text


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_report_localizes_emitted_review_signals(tmp_path, language):
    outputs = load_outputs()
    helpers = sys.modules["reconciliation_helpers"]
    signals = helpers.review_signal_rows(
        [
            {
                "record_id": f"open-{index}",
                "document_no": f"INV-{index}",
                "document_date": "2025-01-01",
                "amount": "1000.00",
                "reconciliation_status": status,
                "evidence_level": level,
                "matched_evidence_type": "payment_order",
            }
            for index, (status, level) in enumerate(
                [
                    ("needs_evidence", "weak_internal"),
                    ("unresolved", "none"),
                    ("probable_payment", "bridge_only"),
                ]
            )
        ],
        {"cutoff_date": "2026-09-30", "review_high_value_threshold": "100"},
    )
    emitted = {
        token.strip() for row in signals for token in row["review_signals"].split(";")
    }
    assert {
        "old_open_item",
        "evidence_weak_internal",
        "evidence_bridge_only",
        "evidence_none",
        "payment_order_bridge",
    } <= emitted
    path = tmp_path / "report.docx"
    outputs.write_word_report(
        path,
        title="Report",
        metadata={},
        summary_rows=[],
        assumptions={},
        next_steps=[],
        review_signals=signals,
        language=language,
    )
    text = document_text(path)
    assert "INV-0" in text and "INV-1" in text and "INV-2" in text
    for code in emitted:
        assert code not in text
    assert "Evidence weak internal" not in text
    assert "Payment order bridge" not in text
    if language == "it":
        assert "Evidenza interna debole" in text
        assert "Solo evidenza ponte" in text
        assert "Partita datata" in text


def test_write_word_report_includes_deterministic_check_summaries(tmp_path):
    outputs = load_outputs()
    path = tmp_path / "relazione.docx"

    outputs.write_word_report(
        path,
        title="Relazione di riconciliazione",
        metadata={},
        summary_rows=[],
        assumptions={"cutoff_date": "2023-12-31"},
        next_steps=[],
        account_rollforward_check=[
            {
                "account": "TOTAL",
                "account_name": "Conti confrontati",
                "opening_difference_journal_minus_ledger": "0.00",
                "closing_difference_journal_minus_ledger": "8941820.33",
                "status": "DIFFERENCE",
                "review_note": "Il saldo ricostruito dal giornale non coincide con il saldo finale del mastro.",
            }
        ],
        post_cutoff_candidates=[
            {
                "document_no": "23FE01/000425",
                "document_date": "2023-01-31",
                "open_amount": "109484.65",
                "evidence_date": "2024-01-12",
                "evidence_type": "external_bank",
                "exact_amount_match": "YES",
                "evidence_source_file": "banca.pdf",
                "evidence_amount": "109484.65",
            }
        ],
        aging_summary=[
            {"aging_bucket": "181-365", "rows": 1, "amount_abs_total": "109484.65"}
        ],
        evidence_concentration=[
            {
                "support_bucket": "payment_order",
                "reconciliation_status": "needs_evidence",
                "rows": 1,
                "amount_abs_total": "109484.65",
                "share_of_abs_amount_percent": "100.00",
            }
        ],
        review_signals=[
            {
                "review_signal_rank": 1,
                "document_no": "23FE01/000425",
                "amount": "109484.65",
                "age_days_at_reference": 334,
                "reconciliation_status": "needs_evidence",
                "review_signals": "high_value; needs_evidence",
            }
        ],
        document_source_map=[
            {
                "document_no_examples": "23FE01/000425",
                "open_amount_total": "109484.65",
                "open_item_rows": 1,
                "ledger_rows": 1,
                "journal_rows": 1,
                "bank_rows": 0,
                "payment_order_rows": 1,
                "factoring_rows": 0,
                "review_note": "Documento presente in distinta/compensazione: verificare supporto di chiusura.",
            }
        ],
        reversal_candidates=[
            {
                "document_no": "23FE01/000425",
                "open_amount": "109484.65",
                "evidence_amount": "-109484.65",
                "evidence_type": "compensation",
                "evidence_date": "2023-12-15",
                "candidate_reasons": "opposite_sign_amount",
            }
        ],
        cutoff_window_movements=[
            {
                "cutoff_window_timing": "before_cutoff",
                "evidence_type": "compensation",
                "amount": "-109484.65",
            }
        ],
    )

    text = document_text(path)
    assert "Controllo saldi da mastro e giornale" in text
    assert "Differenza" in text
    assert "8,941,820.33" in text
    assert "Evidenze successive al cut-off" in text
    assert "Principali candidati" in text
    assert "23FE01/000425" in text
    assert "Analisi deterministiche aggiuntive" in text
    assert "Aging partite aperte" in text
    assert "Mappa documento-fonti" in text
    assert "Possibili storni, giroconti o compensazioni" in text


def test_word_report_keeps_source_control_receipts_out_of_financial_prose(tmp_path):
    outputs = load_outputs()
    assumptions = {
        "cutoff_date": "2026-09-30",
        "reviewed_source_decisions": {
            "bank.pdf": {"money": {"reported_increment": "0.01"}}
        },
        "_reviewed_source_decision_receipts": {"private_control_marker": "receipt"},
    }
    path = tmp_path / "report.docx"

    outputs.write_word_report(
        path,
        title="Riconciliazione",
        metadata={},
        summary_rows=[],
        assumptions=assumptions,
        next_steps=[],
        language="it",
    )

    text = document_text(path)
    assert "2026-09-30" in text
    assert "reported_increment" not in text
    assert "private_control_marker" not in text
    assert (
        assumptions["reviewed_source_decisions"]["bank.pdf"]["money"][
            "reported_increment"
        ]
        == "0.01"
    )


def test_word_report_leads_with_results_and_keeps_wide_evidence_readable(tmp_path):
    outputs = load_outputs()
    path = tmp_path / "report.docx"
    outputs.write_word_report(
        path,
        title="Riconciliazione — esercitazione",
        metadata={"input_dir": "/local/workspace/fictional-inputs"},
        summary_rows=[
            {"reconciliation_status": "closed", "rows": 2, "amount": "1952.00"},
            {"reconciliation_status": "unresolved", "rows": 1, "amount": "488.00"},
        ],
        assumptions={"cutoff_date": "2026-03-31"},
        next_steps=["Verificare la fattura 26FF01/000003."],
        document_source_map=[
            {
                "document_no_examples": "26FF01/000003",
                "record_refs": "open-source-record",
                "open_amount_total": "488.00",
                "open_item_rows": 1,
                "ledger_rows": 2,
                "journal_rows": 3,
                "bank_rows": 4,
                "payment_order_rows": 5,
                "factoring_rows": 6,
                "review_note": "Nota completa da conservare nel rapporto.",
            }
        ],
        cutoff_window_movements=[
            {
                "record_id": "open:imports/fictional-inputs/open-items.pdf:p1:l17",
                "cutoff_window_timing": "before_cutoff",
                "evidence_type": "open_item",
                "amount": "488.00",
            }
        ],
    )

    document = Document(path)
    first_table = document.tables[0]
    assert [cell.text for cell in first_table.rows[1].cells] == [
        "Chiusa da evidenza",
        "2",
        "1.952,00",
    ]
    assert [cell.text for cell in first_table.rows[2].cells] == [
        "Non risolta",
        "1",
        "488,00",
    ]
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert paragraphs.index("Prossimi passi") < paragraphs.index("Perimetro e metodo")
    assert paragraphs.index("Metadati run") > paragraphs.index("Rinvio al file Excel")
    evidence_tables = [
        table for table in document.tables if table.rows[0].cells[0].text == "Documento"
    ]
    assert len(evidence_tables) == 2
    evidence = {}
    for table in evidence_tables:
        assert table.rows[1].cells[0].text == "26FF01/000003"
        evidence.update(
            zip(
                [cell.text for cell in table.rows[0].cells],
                [cell.text for cell in table.rows[1].cells],
            )
        )
        assert len(table.columns) <= 6
        assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
        assert [column.width for column in table.columns] == [
            cell.width for cell in table.rows[0].cells
        ]
    assert evidence == {
        "Documento": "26FF01/000003",
        "Riferimenti alle righe": "open-source-record",
        "Importo aperto": "488,00",
        "Righe aperte": "1",
        "Mastro": "2",
        "Giornale": "3",
        "Banca": "4",
        "Distinte": "5",
        "Factor": "6",
        "Nota": "Nota completa da conservare nel rapporto.",
    }

    references = next(
        table
        for table in document.tables
        if table.rows[0].cells[-1].text == "Riferimenti alle righe"
    )
    assert references.columns[-1].width.inches >= 2.0
    assert references.rows[1].cells[-1].text == (
        "open:imports/fictional-inputs/open-items.pdf:p1:l17"
    )
