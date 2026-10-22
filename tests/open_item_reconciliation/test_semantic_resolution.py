"""Component regressions; mocked launches are not model-quality acceptance."""

from __future__ import annotations

import hashlib
import importlib.util
import json
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest

SCRIPTS = (
    Path(__file__).resolve().parents[2] / "plugins/open-item-reconciliation/scripts"
)


@pytest.fixture
def runtime():
    sys.path.insert(0, str(SCRIPTS))
    spec = importlib.util.spec_from_file_location(
        "semantic_test_workflow", SCRIPTS / "reconciliation_workflow.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    import reconciliation_helpers as helpers
    import semantic_resolution as semantic

    return module, helpers, semantic


def population(amounts=("366", "244"), payments=("610",), payer="Zeta Cumulo Srl"):
    items = [
        {
            "record_id": f"invoice-{index}",
            "document_no": f"FT-2026-{141 + index:04}",
            "document_date": "2026-10-01",
            "amount": amount,
            "currency": "EUR",
            "counterparty": "Zeta Cumulo Srl",
        }
        for index, amount in enumerate(amounts)
    ]
    bank = [
        {
            "record_id": f"bank-{index}",
            "description": "saldo fatture FT-2026-0141 e FT-2026-0142",
            "posting_date": "2026-10-10",
            "evidence_type": "external_bank",
            "amount": amount,
            "currency": "EUR",
            "counterparty": payer,
        }
        for index, amount in enumerate(payments)
    ]
    return items, bank


def decision(target, allocations, outcome="supported"):
    return {
        "target_id": target,
        "outcome": outcome,
        "reason": "Explicit invoice references in the bank description; same payer and currency.",
        "allocations": [
            {
                "source_id": bank,
                "amount": amount,
                "citations": [
                    {
                        "record_id": target,
                        "field": "document_no",
                        "quote": f"FT-2026-{141 + int(target[-1]):04}",
                    },
                    {
                        "record_id": bank,
                        "field": "description",
                        "quote": "saldo fatture FT-2026-0141 e FT-2026-0142",
                    },
                ],
            }
            for bank, amount in allocations
        ],
    }


def prepare(
    runtime,
    monkeypatch,
    tmp_path,
    items,
    bank,
    decisions,
    classified_rows=None,
    scope=None,
):
    _, h, semantic = runtime
    response = {"decisions": decisions}

    def worker(**kwargs):
        launch = {
            "workflow_id": "open-item-reconciliation",
            "packet_sha256": kwargs["packet_sha256"],
            "requested_worker_configuration": {
                "model": "test-model",
                "reasoning_effort": "high",
                "direct_model_api": False,
                "ephemeral": True,
            },
            "packet": {
                "prompt_sha256": hashlib.sha256(kwargs["prompt"].encode()).hexdigest(),
                "output_schema_sha256": hashlib.sha256(
                    (
                        json.dumps(
                            kwargs["output_schema"], ensure_ascii=False, indent=2
                        )
                        + "\n"
                    ).encode()
                ).hexdigest(),
            },
            "boundary": {
                "canaries": {
                    "exact_schema_read_succeeded": True,
                    "outside_capsule_read_denied": True,
                    "required_cli_capabilities_available": True,
                }
            },
            "process": {
                "return_code": 0,
                "response_sha256": hashlib.sha256(
                    (json.dumps(response, ensure_ascii=False, indent=2) + "\n").encode()
                ).hexdigest(),
            },
        }
        launch["content_sha256"] = semantic.canonical_json_sha256(launch)
        path = kwargs["output_dir"] / "luna_launch_receipt.json"
        path.write_text(json.dumps(launch))
        return {"response_payload": response, "launch_receipt": path}

    monkeypatch.setattr(
        semantic,
        "_capsule",
        lambda: SimpleNamespace(
            discover_worker_selection=lambda *_: {},
            run_isolated_native_worker=worker,
        ),
    )
    assumptions = {"currency": "EUR", **(scope or {})}
    rows = (
        classified_rows
        if classified_rows is not None
        else h.reconcile_open_items(items, bank, assumptions)
    )
    assumptions[semantic.KEY] = semantic.prepare_resolution(
        rows, bank, assumptions, tmp_path / "results" / "semantic-review"
    )
    return assumptions


def test_cumulative_receipt_creates_two_native_closures_and_conserved_ledger(
    runtime, monkeypatch, tmp_path
):
    _, h, semantic = runtime
    items, bank = population()
    assumptions = prepare(
        runtime,
        monkeypatch,
        tmp_path,
        items,
        bank,
        [
            decision("invoice-0", [("bank-0", "366")]),
            decision("invoice-1", [("bank-0", "244")]),
        ],
    )
    rows = h.reconcile_open_items(items, bank, assumptions)
    ledgers, failures = h.closed_bank_allocation_controls(rows, bank, assumptions)
    assert [row["reconciliation_status"] for row in rows] == ["closed", "closed"]
    assert [row["allocated_amount"] for row in rows] == ["366", "244"]
    assert failures == {}
    assert ledgers[0]["balanced"] is True
    assert assumptions[semantic.KEY]["status"] == "completed"


@pytest.mark.parametrize(
    ("amounts", "payment", "allocated", "bank_residual"),
    [(("366", "244"), "610", ("366", "244"), "0"), (("610",), "700", ("610",), "90")],
    ids=["cumulative", "overpayment"],
)
def test_included_post_cutoff_receipt_reaches_worker_native_ledger_and_workpapers(
    runtime, monkeypatch, tmp_path, amounts, payment, allocated, bank_residual
):
    workflow, h, semantic = runtime
    items, bank = population(amounts=amounts, payments=(payment,))
    for item in items:
        item["document_date"] = "2026-07-01"
    bank[0]["posting_date"] = "2026-07-24"
    if len(items) == 1:
        bank[0]["description"] = "saldo fattura FT-2026-0141"
    proposals = [
        decision(f"invoice-{index}", [("bank-0", amount)])
        for index, amount in enumerate(allocated)
    ]
    for proposal in proposals:
        proposal["allocations"][0]["citations"][1]["quote"] = bank[0]["description"]

    assumptions = prepare(
        runtime,
        monkeypatch,
        tmp_path,
        items,
        bank,
        proposals,
        scope={"cutoff_date": "2026-07-01", "post_cutoff_events_excluded": False},
    )

    resolution = assumptions[semantic.KEY]
    assert resolution["status"] == "completed"
    assert len(resolution["batches"]) == 1
    assert (
        resolution["batches"][0]["packet"]["sources"][0]["posting_date"] == "2026-07-24"
    )
    rows = h.reconcile_open_items(items, bank, assumptions)
    assert [row["reconciliation_status"] for row in rows] == ["closed"] * len(items)
    assert [row["allocated_amount"] for row in rows] == list(allocated)
    ledgers, failures = h.closed_bank_allocation_controls(rows, bank, assumptions)
    assert failures == {}
    assert ledgers[0]["balanced"] is (bank_residual == "0")
    assert ledgers[0]["source_residuals"][0]["residual"] == bank_residual
    assert all(row["residual"] == "0" for row in ledgers[0]["target_residuals"])
    result = workflow.build_reconciliation_artifacts(
        output_dir=tmp_path / "results",
        open_items=items,
        evidence_rows=bank,
        assumptions=assumptions,
        require_completed_review=False,
    )
    from audit_assurance import validate_assurance_run
    from openpyxl import load_workbook

    assert len(result["relationship_allocation_ledgers"]) == 1
    assert Path(result["excel_path"]).is_file()
    assert Path(result["word_path"]).is_file()
    workbook = load_workbook(result["accountant_report_path"], read_only=True)
    operating = list(workbook["Scheda operativa"].values)
    allocated_column = operating[0].index("importo attribuito")
    residual_column = operating[0].index("residuo aperto")
    assert [row[allocated_column] for row in operating[1:]] == list(allocated)
    assert all(row[residual_column] == "0" for row in operating[1:])
    if bank_residual != "0":
        from decimal import Decimal

        detail = list(workbook["Dettaglio riscontri"].values)
        difference_column = detail[0].index("differenza")
        assert Decimal(detail[1][difference_column]) == Decimal(bank_residual)
    workbook.close()
    assert validate_assurance_run(tmp_path / "results")["content_sha256"]


@pytest.mark.parametrize("policy", [True, None], ids=["excluded", "default-excluded"])
def test_excluded_post_cutoff_receipt_never_launches_worker(
    runtime, monkeypatch, tmp_path, policy
):
    _, h, semantic = runtime
    items, bank = population()
    for item in items:
        item["document_date"] = "2026-07-01"
    bank[0]["posting_date"] = "2026-07-24"
    assumptions = {"cutoff_date": "2026-07-01"}
    if policy is not None:
        assumptions["post_cutoff_events_excluded"] = policy
    rows = h.reconcile_open_items(items, bank, assumptions)
    monkeypatch.setattr(
        semantic, "_capsule", lambda: pytest.fail("Excluded receipt sent")
    )

    result = semantic.prepare_resolution(
        rows, bank, assumptions, tmp_path / "semantic-review"
    )

    assert result["status"] == "missing_evidence"
    assert result["batches"] == []
    assert all(row["reconciliation_status"] == "unresolved" for row in rows)


@pytest.mark.parametrize("tamper", ["hide_excess", "underallocate", "inflate_invoice"])
def test_overpayment_assurance_rejects_ledger_changes(
    runtime, monkeypatch, tmp_path, tamper
):
    workflow, h, _ = runtime
    items, bank = population(amounts=("610",), payments=("700",))
    assumptions = prepare(
        runtime,
        monkeypatch,
        tmp_path,
        items,
        bank,
        [decision("invoice-0", [("bank-0", "610")])],
    )
    original = workflow.closed_bank_allocation_controls

    def altered_ledger(*args, **kwargs):
        ledgers, failures = original(*args, **kwargs)
        ledger = json.loads(json.dumps(ledgers[0]))
        if tamper == "hide_excess":
            ledger["source_records"][0]["amount"] = "610"
        elif tamper == "underallocate":
            ledger["allocations"][0]["amount"] = "609"
        else:
            ledger["target_records"][0]["amount"] = "700"
        rebuilt = h.build_allocation_ledger(
            **{
                key: ledger[key]
                for key in (
                    "ledger_id",
                    "policy",
                    "source_records",
                    "target_records",
                    "allocations",
                )
            }
        )
        return [rebuilt], failures

    monkeypatch.setattr(workflow, "closed_bank_allocation_controls", altered_ledger)
    from audit_assurance import AssuranceRunError

    with pytest.raises(AssuranceRunError, match="exact evidence allocation"):
        workflow.build_reconciliation_artifacts(
            output_dir=tmp_path / "results",
            open_items=items,
            evidence_rows=bank,
            assumptions=assumptions,
            require_completed_review=False,
        )


def test_sorted_json_replay_retains_worker_receipt_and_rejects_response_edit(
    runtime, monkeypatch, tmp_path
):
    _, h, semantic = runtime
    items, bank = population()
    assumptions = prepare(
        runtime,
        monkeypatch,
        tmp_path,
        items,
        bank,
        [
            decision("invoice-0", [("bank-0", "366")]),
            decision("invoice-1", [("bank-0", "244")]),
        ],
    )
    persisted = json.loads(json.dumps(assumptions, sort_keys=True))

    rows = h.reconcile_open_items(items, bank, persisted)

    assert [row["reconciliation_status"] for row in rows] == ["closed", "closed"]
    persisted["semantic_resolution"]["batches"][0]["response"]["decisions"][0][
        "reason"
    ] = "edited"
    receipt = persisted[semantic.KEY]
    receipt["content_sha256"] = semantic.canonical_json_sha256(
        {key: value for key, value in receipt.items() if key != "content_sha256"}
    )
    with pytest.raises(ValueError, match="retained serialization"):
        h.reconcile_open_items(items, bank, persisted)


@pytest.mark.parametrize(
    ("payments", "allocations", "status", "target_residual", "source_residual"),
    [
        (("100",), [("bank-0", "100")], "partially_paid", "266", "0"),
        (("100", "266"), [("bank-0", "100"), ("bank-1", "266")], "closed", "0", "0"),
        (("400",), [("bank-0", "366")], "closed", "0", "34"),
    ],
)
def test_partial_multiple_and_excess_payments_retain_exact_residuals(
    runtime,
    monkeypatch,
    tmp_path,
    payments,
    allocations,
    status,
    target_residual,
    source_residual,
):
    _, h, _ = runtime
    items, bank = population(amounts=("366",), payments=payments)
    assumptions = prepare(
        runtime,
        monkeypatch,
        tmp_path,
        items,
        bank,
        [decision("invoice-0", allocations)],
    )
    rows = h.reconcile_open_items(items, bank, assumptions)
    ledger = h.closed_bank_allocation_controls(rows, bank, assumptions)[0][0]
    assert rows[0]["reconciliation_status"] == status
    assert rows[0]["residual_amount"] == target_residual
    assert ledger["source_residuals"][0]["residual"] == source_residual
    assert ledger["target_residuals"][0]["residual"] == target_residual
    import accountant_report

    main, details, _ = accountant_report.build_accountant_report_rows(
        rows, normalized_records=bank
    )
    assert main[0]["non capita"] == "NO"
    assert main[0]["differenza importo"] == target_residual
    assert len(details) == len(payments)


@pytest.mark.parametrize("second_date_field", ["posting_date", "value_date"])
def test_multiple_semantic_payments_keep_each_actual_bank_date(
    runtime, monkeypatch, tmp_path, second_date_field
):
    _, h, _ = runtime
    items, bank = population(amounts=("366",), payments=("100", "266"))
    del bank[1]["posting_date"]
    bank[1][second_date_field] = "2026-10-12"
    assumptions = prepare(
        runtime,
        monkeypatch,
        tmp_path,
        items,
        bank,
        [decision("invoice-0", [("bank-0", "100"), ("bank-1", "266")])],
    )

    rows = h.reconcile_open_items(items, bank, assumptions)

    assert rows[0]["supporting_bank_record_id"] == "bank-0; bank-1"
    assert rows[0]["supporting_bank_date"] == "2026-10-10; 2026-10-12"
    import accountant_report

    main, details, _ = accountant_report.build_accountant_report_rows(
        rows, normalized_records=bank
    )
    assert "2026-10-10" in str(main) and "2026-10-12" in str(main)
    assert main[0]["_cells"]["data pagamento"] == [
        {"record_id": "bank-0", "fields": ["posting_date"]},
        {"record_id": "bank-1", "fields": [second_date_field]},
    ]
    assert len(details) == 2


def test_mismatched_payer_cannot_be_closed_by_equal_amount(
    runtime, monkeypatch, tmp_path
):
    _, h, semantic = runtime
    items, bank = population(amounts=("610",), payer="Different Payer Srl")
    assumptions = prepare(
        runtime,
        monkeypatch,
        tmp_path,
        items,
        bank,
        [decision("invoice-0", [("bank-0", "610")])],
    )
    rows = h.reconcile_open_items(items, bank, assumptions)
    assert rows[0]["reconciliation_status"] != "closed"
    assert "party mismatch" in assumptions[semantic.KEY]["blocks"][0]["reason"]


@pytest.mark.parametrize("outcome", ["ambiguous", "conflict", "missing_evidence"])
def test_uncertainty_never_becomes_a_paid_item(runtime, monkeypatch, tmp_path, outcome):
    _, h, _ = runtime
    items, bank = population(amounts=("610",))
    assumptions = prepare(
        runtime,
        monkeypatch,
        tmp_path,
        items,
        bank,
        [decision("invoice-0", [], outcome)],
    )
    rows = h.reconcile_open_items(items, bank, assumptions)
    assert rows[0]["reconciliation_status"] != "closed"
    assert rows[0]["semantic_outcome"] == outcome
    assert h.closed_bank_allocation_controls(rows, bank, assumptions)[0] == []


def test_allocations_cannot_reuse_cash_or_exceed_invoice_amount(
    runtime, monkeypatch, tmp_path
):
    _, h, semantic = runtime
    items, bank = population(amounts=("610", "610"))
    assumptions = prepare(
        runtime,
        monkeypatch,
        tmp_path,
        items,
        bank,
        [
            decision("invoice-0", [("bank-0", "610")]),
            decision("invoice-1", [("bank-0", "610")]),
        ],
    )
    assert assumptions[semantic.KEY]["status"] == "incomplete"
    assert assumptions[semantic.KEY]["batches"] == []
    assert all(
        row["reconciliation_status"] != "closed"
        for row in h.reconcile_open_items(items, bank, assumptions)
    )


def test_source_change_invalidates_semantic_replay(runtime, monkeypatch, tmp_path):
    _, h, _ = runtime
    items, bank = population(amounts=("610",))
    assumptions = prepare(
        runtime,
        monkeypatch,
        tmp_path,
        items,
        bank,
        [decision("invoice-0", [("bank-0", "610")])],
    )
    bank[0]["description"] = "unrelated payment"
    with pytest.raises(ValueError, match="stale"):
        h.reconcile_open_items(items, bank, assumptions)


@pytest.mark.parametrize(
    ("input_tolerance", "sealed_tolerance"), [("0.00", "0"), ("0.0100", "0.01")]
)
def test_equivalent_tolerance_spelling_preserves_semantic_replay(
    runtime, tmp_path, input_tolerance, sealed_tolerance
):
    _, h, semantic = runtime
    items, bank = population(payments=())
    assumptions = {"amount_tolerance": input_tolerance}
    preliminary = h.reconcile_open_items(items, bank, assumptions)
    assumptions[semantic.KEY] = semantic.prepare_resolution(
        preliminary, bank, assumptions, tmp_path / "semantic-review"
    )
    assumptions["amount_tolerance"] = sealed_tolerance

    rows = h.reconcile_open_items(items, bank, assumptions)

    assert rows == preliminary


def test_changed_tolerance_value_invalidates_semantic_replay(runtime, tmp_path):
    _, h, semantic = runtime
    items, bank = population(payments=())
    assumptions = {"amount_tolerance": "0.00"}
    preliminary = h.reconcile_open_items(items, bank, assumptions)
    assumptions[semantic.KEY] = semantic.prepare_resolution(
        preliminary, bank, assumptions, tmp_path / "semantic-review"
    )
    assumptions["amount_tolerance"] = "0.01"

    with pytest.raises(ValueError, match="stale"):
        h.reconcile_open_items(items, bank, assumptions)


def test_correct_semantics_reach_native_workbooks_and_assurance(
    runtime, monkeypatch, tmp_path
):
    workflow, _, _ = runtime
    items, bank = population()
    assumptions = prepare(
        runtime,
        monkeypatch,
        tmp_path,
        items,
        bank,
        [
            decision("invoice-0", [("bank-0", "366")]),
            decision("invoice-1", [("bank-0", "244")]),
        ],
    )
    result = workflow.build_reconciliation_artifacts(
        output_dir=tmp_path / "results",
        open_items=items,
        evidence_rows=bank,
        assumptions=assumptions,
    )
    from audit_assurance import validate_assurance_run
    from openpyxl import load_workbook

    workbook = load_workbook(result["accountant_report_path"])
    assert "366" in {
        str(cell.value) for row in workbook["Scheda operativa"] for cell in row
    }
    assert len(result["relationship_allocation_ledgers"]) == 1
    assert validate_assurance_run(tmp_path / "results")["content_sha256"]


def test_fabricated_quote_never_becomes_a_native_closure(
    runtime, monkeypatch, tmp_path
):
    _, h, semantic = runtime
    items, bank = population(amounts=("610",))
    proposal = decision("invoice-0", [("bank-0", "610")])
    proposal["allocations"][0]["citations"][1]["quote"] = "fabricated source phrase"
    assumptions = prepare(runtime, monkeypatch, tmp_path, items, bank, [proposal])
    assert "exact quote" in assumptions[semantic.KEY]["blocks"][0]["reason"]
    assert (
        h.reconcile_open_items(items, bank, assumptions)[0]["reconciliation_status"]
        != "closed"
    )


@pytest.mark.parametrize(
    "case", ["no_bank", "oversized", "foreign_currency", "foreign_entity"]
)
def test_incompatible_or_oversized_population_is_not_sent(
    runtime, monkeypatch, tmp_path, case
):
    _, h, semantic = runtime
    items, bank = population(amounts=("610",))
    if case == "no_bank":
        bank = []
    elif case == "oversized":
        bank[0]["description"] = "x" * semantic.MAX_PACKET_BYTES
    elif case == "foreign_currency":
        bank[0]["currency"] = "USD"
    else:
        items[0]["entity_ref"] = "one-entity"
        bank[0]["entity_ref"] = "other-entity"
    monkeypatch.setattr(
        semantic,
        "_capsule",
        lambda: pytest.fail("No compatible bounded packet may be sent"),
    )
    rows = h.reconcile_open_items(items, bank, {})
    result = semantic.prepare_resolution(rows, bank, {}, tmp_path / "semantic-review")
    assert result["batches"] == []
    assert result["blocks"]


def test_retained_worker_artifact_tampering_is_rejected(runtime, monkeypatch, tmp_path):
    _, _, semantic = runtime
    items, bank = population(amounts=("610",))
    assumptions = prepare(
        runtime,
        monkeypatch,
        tmp_path,
        items,
        bank,
        [decision("invoice-0", [("bank-0", "610")])],
    )
    output = tmp_path / "results"
    (output / "prepared_records.json").write_text(
        json.dumps({"assumptions": assumptions})
    )
    (output / "semantic-review/batch-0000/request.json").write_text("{}")
    with pytest.raises(ValueError, match="artifact changed"):
        semantic.semantic_artifact_paths(output)


def test_established_partial_payment_cannot_be_taken_by_another_target(
    runtime, monkeypatch, tmp_path
):
    _, h, semantic = runtime
    items, bank = population()
    rows = h.reconcile_open_items(items, bank, {"currency": "EUR"})
    rows[0].update(
        reconciliation_status="partially_paid",
        matched_evidence_id="bank-0",
        allocated_amount="100",
    )
    assumptions = prepare(
        runtime,
        monkeypatch,
        tmp_path,
        items,
        bank,
        [
            decision("invoice-0", [], "missing_evidence"),
            decision("invoice-1", [("bank-0", "244")]),
        ],
        classified_rows=rows,
    )
    assert (
        "partial payment cannot be reused"
        in assumptions[semantic.KEY]["blocks"][0]["reason"]
    )
    assert (
        semantic.apply_resolution(rows, bank, assumptions)[0]["allocated_amount"]
        == "100"
    )


@pytest.mark.parametrize("host", ["cowork", "antigravity"])
def test_native_handoff_resumes_same_population_and_closes_cumulative_payment(
    runtime, monkeypatch, tmp_path, host
):
    from tests.plugins.test_native_host_workers import completed_record

    workflow, helpers, semantic = runtime
    monkeypatch.syspath_prepend(
        str(SCRIPTS.parents[1] / "journal-bank-reconciliation/scripts")
    )
    import native_host_worker
    import worker_capabilities

    monkeypatch.setattr(
        semantic,
        "_capsule",
        lambda: SimpleNamespace(
            discover_worker_selection=lambda workflow: worker_capabilities.select_worker(
                workflow, host=host
            ),
            run_isolated_native_worker=lambda **kwargs: native_host_worker.run_worker(
                **{
                    key: value
                    for key, value in kwargs.items()
                    if key != "source_records"
                }
            ),
        ),
    )
    items, bank = population()
    assumptions = {"currency": "EUR"}
    rows = helpers.reconcile_open_items(items, bank, assumptions)
    output = tmp_path / "results/semantic-review"
    pending = semantic.prepare_resolution(rows, bank, assumptions, output)
    request_path = output / "batch-0000/native_host_request.json"
    original_request = request_path.read_bytes()
    response = {
        "decisions": [
            decision("invoice-0", [("bank-0", "366")]),
            decision("invoice-1", [("bank-0", "244")]),
        ]
    }
    (request_path.parent / "native_host_record.json").write_text(
        json.dumps(completed_record(json.loads(original_request), response))
    )

    assumptions[semantic.KEY] = semantic.prepare_resolution(
        rows, bank, assumptions, output
    )
    result = helpers.reconcile_open_items(items, bank, assumptions)

    assert pending["status"] == "incomplete"
    assert "NATIVE_WORKER_PENDING" in pending["blocks"][0]["reason"]
    assert request_path.read_bytes() == original_request
    assert assumptions[semantic.KEY]["status"] == "completed"
    assert [row["reconciliation_status"] for row in result] == ["closed", "closed"]
    delivered = workflow.build_reconciliation_artifacts(
        output_dir=output.parent,
        open_items=items,
        evidence_rows=bank,
        assumptions=assumptions,
        require_completed_review=False,
    )
    assert delivered["assurance"]["content_sha256"]
    assert Path(delivered["excel_path"]).is_file()
    assert Path(delivered["word_path"]).is_file()
