from __future__ import annotations

import importlib.util
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

SCRIPTS = (
    Path(__file__).resolve().parents[2]
    / "plugins"
    / "open-item-reconciliation"
    / "scripts"
)
WORKFLOW = SCRIPTS / "reconciliation_workflow.py"


def load_workflow():
    sys.path.insert(0, str(SCRIPTS))
    spec = importlib.util.spec_from_file_location(
        "open_item_reconciliation_workflow", WORKFLOW
    )
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def document_text(path: str | Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


def test_build_reconciliation_artifacts_writes_excel_and_word(tmp_path):
    workflow = load_workflow()
    result = workflow.build_reconciliation_artifacts(
        output_dir=tmp_path,
        open_items=[
            {
                "record_id": "open-1",
                "document_key": "INV1|2023",
                "document_date": "2023-03-01",
                "amount": "100.00",
            }
        ],
        evidence_rows=[
            {
                "record_id": "bank-1",
                "evidence_type": "external_bank",
                "document_key": "INV1|2023",
                "posting_date": "2023-12-15",
                "amount": "100.00",
                "source_file": "bank.pdf",
                "source_page": "1",
            }
        ],
        assumptions={"scope_year": "2023", "cutoff_date": "2023-12-31"},
        source_inventory=[{"source_file": "bank.pdf", "source_role": "bank_statement"}],
        ledger_balance_rows=[
            {"account": "TOTAL", "closing_balance_signed_debit_minus_credit": "100.00"}
        ],
        account_rollforward_check=[{"account": "TOTAL", "status": "PASS"}],
        aggregate_rollforward_summary=[
            {"account": "TOTAL", "closing_net_debit_minus_credit": "100.00"}
        ],
        aggregate_rollforward_rows=[{"record_id": "journal_rollforward:1"}],
        metadata={"Periodo": "2023"},
        narrative="Riconciliazione completata sui dati normalizzati.",
    )

    assert result["checks_pass"] is True
    assert result["reconciliation_rows"][0]["reconciliation_status"] == "closed"
    assert result["reconciliation_rows"][0]["relationship_control_status"] == "passed"
    assert len(result["relationship_allocation_ledgers"]) == 1
    assert Path(result["excel_path"]).exists()
    assert Path(result["accountant_report_path"]).exists()
    assert Path(result["word_path"]).exists()
    assert Path(result["review_session"]["review_html_path"]).exists()
    assert Path(result["review_session"]["review_html_path"]).name == "review_ui.html"

    workbook = load_workbook(result["excel_path"])
    assert "Reconciliation detail" in workbook.sheetnames
    assert "Bank allocation candidates" in workbook.sheetnames
    assert "External evidence aggregate" in workbook.sheetnames
    assert "External evidence detail" in workbook.sheetnames
    assert "Ledger balance check" in workbook.sheetnames
    assert "Account rollforward check" in workbook.sheetnames
    assert "Journal rollforward" in workbook.sheetnames
    assert "Journal detail" in workbook.sheetnames
    assert "Post-cutoff candidates" in workbook.sheetnames
    assert "Open item aging" in workbook.sheetnames
    assert "Evidence concentration" in workbook.sheetnames
    assert "Review signals" in workbook.sheetnames
    assert "Document source map" in workbook.sheetnames
    assert "Reversal candidates" in workbook.sheetnames
    assert "Cutoff window movements" in workbook.sheetnames
    assert "Review" in workbook.sheetnames
    assert "bank_allocation_candidates" in result
    accountant_workbook = load_workbook(result["accountant_report_path"])
    assert "Scheda operativa" in accountant_workbook.sheetnames
    assert "Dettaglio riscontri" in accountant_workbook.sheetnames
    assert "external_evidence_summary" in result
    assert "external_evidence_detail" in result
    assert result["ledger_balance_rows"][0]["account"] == "TOTAL"
    assert result["account_rollforward_check"][0]["status"] == "PASS"
    assert result["aggregate_rollforward_summary"][0]["account"] == "TOTAL"
    assert result["aging_summary"][0]["aging_bucket"] == "181-365"
    assert result["evidence_concentration"][0]["support_bucket"] == "bank"
    assert result["document_source_map"][0]["bank_rows"] == 1
    assert result["cutoff_window_movements"][0]["record_id"] == "bank-1"
    assert any(
        check["check"] == "codex_review_packet_present" for check in result["checks"]
    )
    assert result["review_rows"][0]["review_status"] == "PENDING"
    text = document_text(result["word_path"])
    assert "Conclusioni" in text
    assert "Controllo saldi da mastro e giornale" in text
    assert "Analisi deterministiche aggiuntive" in text


def test_workflow_can_be_loaded_without_preconfigured_python_path(tmp_path):
    spec = importlib.util.spec_from_file_location("audit_workflow_direct", WORKFLOW)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    result = module.build_reconciliation_artifacts(
        output_dir=tmp_path,
        open_items=[],
        evidence_rows=[],
        assumptions={"scope_year": "2023"},
    )

    assert result["checks_pass"] is True
    assert Path(result["excel_path"]).exists()
    assert Path(result["word_path"]).exists()


def test_default_next_steps_mentions_unresolved_rows():
    workflow = load_workflow()
    steps = workflow.default_next_steps(
        [{"reconciliation_status": "unresolved"}],
        language="it",
    )

    assert any("non risolte" in step for step in steps)


def test_default_next_steps_uses_selected_language():
    workflow = load_workflow()
    steps = workflow.default_next_steps(
        [{"reconciliation_status": "needs_evidence"}],
        language="en_US",
    )

    assert any("Obtain the evidence" in step for step in steps)


def test_default_report_title_uses_spanish_language(tmp_path):
    workflow = load_workflow()

    result = workflow.build_reconciliation_artifacts(
        output_dir=tmp_path,
        open_items=[],
        evidence_rows=[],
        assumptions={"scope_year": "2026"},
        language="es_ES",
    )

    assert "Informe de conciliación contable" in document_text(result["word_path"])


def test_workflow_can_require_completed_review(tmp_path):
    workflow = load_workflow()

    try:
        workflow.build_reconciliation_artifacts(
            output_dir=tmp_path,
            open_items=[
                {
                    "record_id": "open-1",
                    "document_key": "INV1|2023",
                    "document_date": "2023-03-01",
                    "amount": "100.00",
                }
            ],
            evidence_rows=[],
            assumptions={"scope_year": "2023"},
            require_completed_review=True,
        )
    except ValueError as exc:
        assert "codex_review_completed" in str(exc)
    else:
        raise AssertionError(
            "expected pending Codex review to fail when completion is required"
        )


def test_office_outputs_are_byte_replayable_and_have_fixed_package_metadata(
    tmp_path,
):
    workflow = load_workflow()
    open_items = [
        {
            "record_id": "open-1",
            "document_key": "INV1|2025",
            "document_date": "2025-01-01",
            "amount": "10.00",
            "currency": "EUR",
        }
    ]
    evidence_rows = [
        {
            "record_id": "bank-1",
            "source_role": "bank_statement",
            "evidence_type": "external_bank",
            "document_key": "INV1|2025",
            "posting_date": "2025-01-02",
            "amount": "10.00",
            "currency": "EUR",
            "source_file": "bank.csv",
            "source_row": "2",
        }
    ]

    first = workflow.build_reconciliation_artifacts(
        output_dir=tmp_path / "first",
        open_items=open_items,
        evidence_rows=evidence_rows,
        assumptions={"scope_year": "2025", "amount_tolerance": "0"},
        require_completed_review=False,
    )
    second = workflow.build_reconciliation_artifacts(
        output_dir=tmp_path / "second",
        open_items=open_items,
        evidence_rows=evidence_rows,
        assumptions={"scope_year": "2025", "amount_tolerance": "0"},
        require_completed_review=False,
    )

    for key in ("excel_path", "accountant_report_path", "word_path"):
        first_path = Path(first[key])
        second_path = Path(second[key])
        assert first_path.read_bytes() == second_path.read_bytes()
        with zipfile.ZipFile(first_path) as package:
            assert all(
                info.date_time == (1980, 1, 1, 0, 0, 0) for info in package.infolist()
            )
            core_properties = package.read("docProps/core.xml")
        assert b"2000-01-01T00:00:00Z" in core_properties


def test_workflow_seals_thousands_formatted_amount_without_losing_digits(tmp_path):
    workflow = load_workflow()
    result = workflow.build_reconciliation_artifacts(
        output_dir=tmp_path,
        open_items=[
            {
                "record_id": "invoice-thousands",
                "document_key": "INV-THOUSANDS",
                "document_date": "2026-09-01",
                "amount": "1220.00",
                "currency": "EUR",
            }
        ],
        evidence_rows=[],
        assumptions={"scope_year": "2026", "cutoff_date": "2026-09-30"},
        language="it",
    )

    assert result["checks_pass"] is True
    assert result["reconciliation_rows"][0]["amount"] == "1220.00"
    assert Path(result["word_path"]).is_file()


def test_workflow_seals_partial_payment_summary_with_source_record_reference(tmp_path):
    workflow = load_workflow()
    result = workflow.build_reconciliation_artifacts(
        output_dir=tmp_path,
        open_items=[
            {
                "record_id": "invoice-partial",
                "document_key": "INV-PARTIAL",
                "document_date": "2026-09-01",
                "amount": "1220.00",
                "currency": "EUR",
            }
        ],
        evidence_rows=[
            {
                "record_id": "payment-partial",
                "document_key": "INV-PARTIAL",
                "posting_date": "2026-09-20",
                "amount": "488.00",
                "currency": "EUR",
                "evidence_type": "external_bank",
            }
        ],
        assumptions={"scope_year": "2026", "cutoff_date": "2026-09-30"},
        language="it",
    )

    assert result["checks_pass"] is True
    assert "payment-partial" in document_text(result["word_path"])
    row = result["reconciliation_rows"][0]
    assert row["reconciliation_status"] == "partially_paid"
    assert row["amount"] == "1220.00"
    assert row["allocated_amount"] == "488.00"
    assert row["residual_amount"] == "732.00"
    assert row["relationship_control_status"] == "passed"
    assert "732.00" in document_text(result["word_path"])


@pytest.mark.parametrize(
    "close_exact,expected_status", [(False, "unresolved"), (True, "closed")]
)
def test_grouped_bank_payment_uses_one_source_capacity_and_seals_report(
    tmp_path, close_exact, expected_status
):
    workflow = load_workflow()
    invoices = [
        {
            "record_id": "I1",
            "document_key": "1FE|2026",
            "document_no": "1-FE",
            "document_date": "2026-09-01",
            "amount": "1220.00",
            "currency": "EUR",
        },
        {
            "record_id": "I2",
            "document_key": "2FE|2026",
            "document_no": "2-FE",
            "document_date": "2026-09-01",
            "amount": "610.00",
            "currency": "EUR",
        },
    ]
    payment = {
        "record_id": "P1",
        "source_role": "bank_statement",
        "document_key": "1FE|2026",
        "document_keys": "1FE|2026;2FE|2026",
        "description": "Bonifico fatture 1-FE e 2-FE",
        "posting_date": "2026-09-20",
        "amount": "1830.00",
        "currency": "EUR",
        "evidence_type": "external_bank",
    }

    result = workflow.build_reconciliation_artifacts(
        output_dir=tmp_path,
        open_items=invoices,
        evidence_rows=[payment],
        assumptions={
            "scope_year": "2026",
            "cutoff_date": "2026-09-30",
            "promote_probable_bank_payments": close_exact,
            "probable_bank_exact_matches_close": close_exact,
        },
        language="it",
    )

    assert result["checks_pass"] is True
    assert result["reconciliation_rows"][0]["reconciliation_status"] == expected_status
    assert result["reconciliation_rows"][1]["reconciliation_status"] == expected_status
    assert Path(result["word_path"]).is_file()


def test_post_cutoff_payment_seals_reports_without_closing_invoice(tmp_path):
    workflow = load_workflow()

    result = workflow.build_reconciliation_artifacts(
        output_dir=tmp_path,
        open_items=[
            {
                "record_id": "I1",
                "source_role": "open_items",
                "evidence_type": "open_item",
                "document_key": "1FE|2026",
                "document_date": "2026-09-01",
                "amount": "1000.00",
                "currency": "EUR",
            }
        ],
        evidence_rows=[
            {
                "record_id": "P1",
                "document_key": "1FE|2026",
                "posting_date": "2026-10-02",
                "amount": "1000.00",
                "currency": "EUR",
                "source_role": "bank_statement",
                "evidence_type": "external_bank",
            }
        ],
        assumptions={"scope_year": "2026", "cutoff_date": "2026-09-30"},
        language="it",
    )

    assert result["checks_pass"] is True
    assert result["reconciliation_rows"][0]["reconciliation_status"] == "unresolved"
    assert "allocated_amount" not in result["reconciliation_rows"][0]
    assert result["post_cutoff_candidates"]
    assert Path(result["word_path"]).is_file()


def test_public_report_counts_distinct_rows_once_after_document_alias_merge(tmp_path):
    row_count = 1
    workflow = load_workflow()
    open_items = [
        {
            "record_id": f"open-{index}",
            "document_key": "3FE|2026",
            "document_no": "26FE01/000003",
            "document_date": "2026-09-01",
            "evidence_type": "open_item",
            "amount": "1220.00",
        }
        for index in range(row_count)
    ]

    result = workflow.build_reconciliation_artifacts(
        output_dir=tmp_path,
        open_items=open_items,
        evidence_rows=[],
        assumptions={"scope_year": "2026", "cutoff_date": "2026-09-30"},
    )

    workbook = load_workbook(result["excel_path"], read_only=True)
    sheet = workbook["Document source map"]
    headers, values = list(sheet.values)
    saved_row = dict(zip(headers, values))
    workbook.close()
    assert saved_row["document_aliases"] == "3FE|2026; 3|2026"
    assert saved_row["open_item_rows"] == row_count
    assert saved_row["reconciliation_status_counts"] == f"unresolved:{row_count}"
    assert len(result["reconciliation_rows"]) == row_count
