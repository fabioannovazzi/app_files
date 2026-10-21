from __future__ import annotations

import copy
import io
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime
from pathlib import Path

import pytest
from defusedxml.common import EntitiesForbidden
from openpyxl import Workbook

from tests.open_item_reconciliation.test_audit_assurance import (
    SCRIPTS,
    load_assurance,
    load_script_module,
    reviewed_source_decision,
)


def spec(path: Path, role="open_items"):
    value = reviewed_source_decision()
    value.update(role=role, adapter_family="reviewed_tabular_accounting_v1")
    columns = {"document_no": "InvoiceX", "document_date": "DateX", "amount": "ValueX"}
    if role != "open_items":
        columns["posting_date"] = "PostedX"
    value["tabular"] = {
        "sheet": "Export" if path.suffix == ".xlsx" else None,
        "header_row": 1,
        "first_data_row": 2,
        "last_data_row": 3,
        "delimiter": ";" if path.suffix == ".csv" else None,
        "columns": columns,
        "evidence_type_map": {},
        "default_evidence_type": (
            "open_item"
            if role == "open_items"
            else "external_bank" if role == "bank_statement" else "internal_accounting"
        ),
    }
    return value


def write_source(path: Path, *, bad=None):
    rows = [
        ["InvoiceX", "DateX", "ValueX", "PostedX"],
        ["INV-A", datetime(2026, 9, 3), 100.25, datetime(2026, 9, 8)],
        ["INV-B", datetime(2026, 9, 4), 200.50, datetime(2026, 9, 9)],
    ]
    if bad == "formula":
        rows[2][2] = "=200+0.5"
    if bad == "date":
        rows[2][1] = "31/02/2026"
    if bad == "precision":
        rows[2][2] = 200.501
    if bad == "header":
        rows[0][3] = "ValueX"
    if path.suffix == ".xlsx":
        wb = Workbook()
        sheet = wb.active
        sheet.title = "Export"
        for row in rows:
            sheet.append(row)
        wb.save(path)
    else:
        import csv

        with path.open("w", newline="") as handle:
            writer = csv.writer(handle, delimiter=";")
            writer.writerows(
                [
                    [
                        v.strftime("%d/%m/%Y") if isinstance(v, datetime) else str(v)
                        for v in row
                    ]
                    for row in rows
                ]
            )


def reviewed(assurance, path, value):
    receipts = assurance.build_source_receipts(path.parent, [path])
    decisions, errors = assurance.build_reviewed_source_decisions(
        input_root=path.parent,
        source_receipts=receipts,
        adapter_families={path.name: "reviewed_tabular_accounting_v1"},
        assumptions={
            "assurance_run_date": "2026-10-18",
            "reviewed_source_decisions": {path.name: value},
        },
    )
    assert errors == {}
    return receipts, decisions[path.name]


@pytest.mark.parametrize(
    "key_style", ["relative", "absolute", "absolute_native", "both_identical"]
)
def test_nested_source_review_keeps_exact_receipt_identity(tmp_path, key_style):
    a = load_assurance()
    source = tmp_path / "inputs"
    path = source / "imports/input_a/population.csv"
    path.parent.mkdir(parents=True)
    write_source(path)
    relative = path.relative_to(source).as_posix()
    choice = spec(path)
    keys = {
        "relative": [relative],
        "absolute": [path.as_posix()],
        "absolute_native": [str(path)],
        "both_identical": [relative, path.as_posix(), str(path)],
    }[key_style]
    receipts = a.build_source_receipts(source, [path])

    decisions, errors = a.build_reviewed_source_decisions(
        input_root=source,
        source_receipts=receipts,
        adapter_families={relative: "reviewed_tabular_accounting_v1"},
        assumptions={
            "assurance_run_date": "2026-10-21",
            "reviewed_source_decisions": {key: choice for key in keys},
        },
    )

    assert errors == {}
    assert list(decisions) == [relative]
    assert decisions[relative]["content"]["source_path"] == relative
    assert decisions[relative]["source_artifact_refs"] == [receipts[0]["artifact_id"]]


@pytest.mark.parametrize(
    "key_style", ["foreign_absolute", "basename", "conflicting_aliases"]
)
def test_nested_source_review_does_not_substitute_other_mapping(tmp_path, key_style):
    a = load_assurance()
    source = tmp_path / "inputs"
    path = source / "imports/input_a/population.csv"
    path.parent.mkdir(parents=True)
    write_source(path)
    relative = path.relative_to(source).as_posix()
    choice = spec(path)
    conflicting = copy.deepcopy(choice)
    conflicting["perimeter"]["party_ref"] = "party.other"
    selections = {
        "foreign_absolute": {str(tmp_path / "other_run/inputs" / relative): choice},
        "basename": {path.name: choice},
        "conflicting_aliases": {relative: choice, path.as_posix(): conflicting},
    }[key_style]
    receipts = a.build_source_receipts(source, [path])

    decisions, errors = a.build_reviewed_source_decisions(
        input_root=source,
        source_receipts=receipts,
        adapter_families={relative: "reviewed_tabular_accounting_v1"},
        assumptions={
            "assurance_run_date": "2026-10-21",
            "reviewed_source_decisions": selections,
        },
    )

    assert decisions == {}
    assert list(errors) == [relative]


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
@pytest.mark.parametrize("role", ["open_items", "bank_statement", "journal"])
def test_reviewed_columns_drive_source_rows_without_header_guessing(
    tmp_path, suffix, role
):
    a = load_assurance()
    path = tmp_path / ("ordinary" + suffix)
    write_source(path)
    receipts, decision = reviewed(a, path, spec(path, role))
    rows = a.read_reviewed_tabular_source(path, decision)
    assert [r["amount"] for r in rows] == ["100.25", "200.5"]
    assert [r["document_date"] for r in rows] == ["2026-09-03", "2026-09-04"]
    assert {r["source_role"] for r in rows} == {role}
    assert {r["evidence_type"] for r in rows} == {
        (
            "open_item"
            if role == "open_items"
            else "external_bank" if role == "bank_statement" else "internal_accounting"
        )
    }
    locator, amount = a._source_value_locator(
        row=rows[0], source_root=tmp_path, receipt=receipts[0], decision=decision
    )
    assert amount == "100.25" and "reviewed_tabular_record_value" in locator
    wrong = {**rows[0], "document_no": "WRONG-INVOICE"}
    with pytest.raises(a.AssuranceRunError, match="stale"):
        a._source_value_locator(
            row=wrong, source_root=tmp_path, receipt=receipts[0], decision=decision
        )


@pytest.mark.parametrize("bad", ["formula", "date", "precision", "header"])
def test_invalid_mapped_cell_rejects_entire_source(tmp_path, bad):
    a = load_assurance()
    path = tmp_path / "source.xlsx"
    write_source(path, bad=bad)
    _, decision = reviewed(a, path, spec(path))
    with pytest.raises(a.AssuranceRunError):
        a.read_reviewed_tabular_source(path, decision)


def test_workbook_xml_entities_cannot_supply_source_values(tmp_path):
    a = load_assurance()
    path = tmp_path / "source.xlsx"
    write_source(path)
    rewritten = io.BytesIO()
    with (
        zipfile.ZipFile(path) as source_zip,
        zipfile.ZipFile(rewritten, "w") as target_zip,
    ):
        for item in source_zip.infolist():
            data = source_zip.read(item.filename)
            if item.filename == "xl/workbook.xml":
                data = (
                    b'<!DOCTYPE workbook [<!ENTITY sheetname "Export">]>'
                    + data.replace(b'name="Export"', b'name="&sheetname;"')
                )
            target_zip.writestr(item, data)
    path.write_bytes(rewritten.getvalue())
    _, decision = reviewed(a, path, spec(path))

    with pytest.raises(EntitiesForbidden):
        a.read_reviewed_tabular_source(path, decision)


def test_source_bytes_and_perimeter_are_not_replaced_by_review_defaults(tmp_path):
    a = load_assurance()
    path = tmp_path / "source.csv"
    write_source(path)
    _, old = reviewed(a, path, spec(path))
    path.write_text(path.read_text().replace("INV-B", "OTHER"))
    with pytest.raises(a.AssuranceRunError, match="changed"):
        a.read_reviewed_tabular_source(path, old)
    path.write_text(
        path.read_text()
        .replace("PostedX", "CurrencyX")
        .replace("08/09/2026", "USD")
        .replace("09/09/2026", "EUR")
    )
    value = spec(path)
    value["tabular"]["columns"]["currency"] = "CurrencyX"
    _, decision = reviewed(a, path, value)
    with pytest.raises(a.AssuranceRunError, match="perimeter"):
        a.read_reviewed_tabular_source(path, decision)


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
@pytest.mark.parametrize("role", ["open_items", "bank_statement", "journal"])
def test_reviewed_party_column_preserves_each_source_counterparty(
    tmp_path, suffix, role
):
    a = load_assurance()
    path = tmp_path / ("parties" + suffix)
    source_rows = [
        ["InvoiceX", "DateX", "ValueX", "PostedX", "PartyX"],
        ["F001", "01/09/2026", "100", "05/09/2026", "Alfa"],
        ["F002", "02/09/2026", "200", "06/09/2026", "Beta"],
    ]
    if suffix == ".xlsx":
        workbook = Workbook()
        workbook.active.title = "Export"
        for row in source_rows:
            workbook.active.append(row)
        workbook.save(path)
        workbook.close()
    else:
        path.write_text("\n".join(";".join(row) for row in source_rows))
    value = spec(path, role)
    # The source has no common party: this explicit mapping is row authority.
    value["perimeter"]["party_ref"] = None
    value["tabular"]["columns"]["party_ref"] = "PartyX"
    _, decision = reviewed(a, path, value)

    rows = a.read_reviewed_tabular_source(path, decision)

    assert [r["party_ref"] for r in rows] == ["Alfa", "Beta"]
    assert {r["entity_ref"] for r in rows} == {"entity.test"}


@pytest.mark.parametrize("party,fixed_party", [("", None), ("Beta", "Alfa")])
def test_party_column_cannot_be_blank_or_contradict_a_fixed_party(
    tmp_path, party, fixed_party
):
    a = load_assurance()
    path = tmp_path / "source.csv"
    path.write_text(
        "InvoiceX;DateX;ValueX;PartyX\n"
        "F001;01/09/2026;100;Alfa\n"
        f"F002;02/09/2026;200;{party}\n"
    )
    value = spec(path)
    value["perimeter"]["party_ref"] = fixed_party
    value["tabular"]["columns"]["party_ref"] = "PartyX"
    _, decision = reviewed(a, path, value)

    with pytest.raises(a.AssuranceRunError, match="party_ref"):
        a.read_reviewed_tabular_source(path, decision)


@pytest.mark.parametrize(
    "language, decision_key_style",
    [
        ("it", "relative"),
        ("en", "relative"),
        ("es", "relative"),
        ("de", "relative"),
        ("fr", "relative"),
        ("it", "absolute"),
        ("it", "absolute_posix"),
    ],
)
def test_multi_party_exports_reconcile_later_payments_and_replay_source_parties(
    tmp_path, language, decision_key_style
):
    a = load_assurance()
    raw = load_script_module("multi_party_raw", SCRIPTS / "raw_input_runner.py")
    workflow = load_script_module(
        "multi_party_workflow", SCRIPTS / "reconciliation_workflow.py"
    )
    source = tmp_path / "inputs"
    source.mkdir()
    population = source / "population.csv"
    population.write_text(
        "InvoiceX;DateX;ValueX;DueX;PartyX\n"
        "F001;01/09/2026;1200.00;30/09/2026;Cliente Alfa Test\n"
        "F002;05/09/2026;800.00;30/09/2026;Cliente Beta Test\n"
        "F003;15/09/2026;500.00;15/10/2026;Cliente Gamma Test\n"
    )
    bank = source / "bank.csv"
    bank.write_text(
        "InvoiceX;PostedX;ValueX;PartyX\n"
        "F001;05/10/2026;1200.00;Cliente Alfa Test\n"
        "F002;10/10/2026;300.00;Cliente Beta Test\n"
    )
    population_spec = spec(population)
    population_spec["tabular"]["last_data_row"] = 4
    population_spec["tabular"]["columns"]["due_date"] = "DueX"
    bank_spec = spec(bank, "bank_statement")
    del bank_spec["tabular"]["columns"]["document_date"]
    for decision in (population_spec, bank_spec):
        decision["perimeter"]["party_ref"] = None
        decision["tabular"]["columns"]["party_ref"] = "PartyX"
    assumptions = {
        "assurance_run_date": "2026-10-20",
        "scope_year": 2026,
        "cutoff_date": "2026-09-30",
        "post_cutoff_events_excluded": False,
        "amount_tolerance": "0",
        "reviewed_source_decisions": {
            population.name: population_spec,
            bank.name: bank_spec,
        },
    }
    if decision_key_style.startswith("absolute"):
        assumptions["reviewed_source_decisions"] = {
            (
                (source / name).as_posix()
                if decision_key_style == "absolute_posix"
                else str(source / name)
            ): decision
            for name, decision in assumptions["reviewed_source_decisions"].items()
        }
    extracted = raw.extract_normalized_records(
        source, assumptions, output_dir=tmp_path / "extracted"
    )
    assert extracted["extraction_errors"] == []
    assert [r["party_ref"] for r in extracted["open_items"]] == [
        "Cliente Alfa Test",
        "Cliente Beta Test",
        "Cliente Gamma Test",
    ]
    output = tmp_path / "workpapers"

    result = workflow.build_reconciliation_artifacts(
        output_dir=output,
        open_items=extracted["open_items"],
        evidence_rows=extracted["evidence_rows"],
        assumptions=assumptions,
        source_inventory=extracted["source_inventory"],
        source_qualifications=extracted["source_qualifications"],
        source_artifact_root=source,
        source_artifact_receipts=extracted["source_artifact_receipts"],
        reviewed_source_decision_receipts=extracted[
            "reviewed_source_decision_receipts"
        ],
        language=language,
        require_completed_review=False,
    )

    rows = result["reconciliation_rows"]
    assert [r["reconciliation_status"] for r in rows] == [
        "closed",
        "partially_paid",
        "unresolved",
    ]
    assert rows[1]["allocated_amount"] == "300.00"
    assert rows[1]["residual_amount"] == "500.00"
    assert rows[2]["amount"] == "500"
    assert (
        result["assurance"]["gate_register"]["gates"]["reconciliation"]["status"]
        == "passed"
    )
    assert Path(result["excel_path"]).is_file()
    assert Path(result["word_path"]).is_file()
    assert (
        a.validate_assurance_run(output)["content_sha256"]
        == result["assurance"]["content_sha256"]
    )

    forged = copy.deepcopy(extracted["evidence_rows"])
    forged[1]["party_ref"] = "Cliente Alfa Test"
    with pytest.raises(a.AssuranceRunError, match="source cells"):
        a.prepare_assurance_run(
            output_dir=tmp_path / "forged",
            open_items=extracted["open_items"],
            evidence_rows=forged,
            assumptions=assumptions,
            source_root=source,
            source_receipts=extracted["source_artifact_receipts"],
            reviewed_source_decisions=extracted["reviewed_source_decision_receipts"],
            source_qualifications=extracted["source_qualifications"],
        )


def test_journal_evidence_meaning_is_explicit_review_not_description_keyword(tmp_path):
    a = load_assurance()
    path = tmp_path / "source.csv"
    write_source(path)
    path.write_text(
        path.read_text()
        .replace("PostedX", "KindX;PostedX")
        .replace("08/09/2026", "credit-note;08/09/2026")
        .replace("09/09/2026", "unproved-note;09/09/2026")
    )
    value = spec(path, "journal")
    value["tabular"]["columns"]["evidence_type"] = "KindX"
    value["tabular"].update(
        default_evidence_type=None,
        evidence_type_map={
            "credit-note": "compensation",
            "unproved-note": "internal_accounting",
        },
    )
    _, decision = reviewed(a, path, value)
    rows = a.read_reviewed_tabular_source(path, decision)
    assert [r["evidence_type"] for r in rows] == ["compensation", "internal_accounting"]
    value["tabular"]["evidence_type_map"].pop("unproved-note")
    _, incomplete = reviewed(a, path, value)
    with pytest.raises(a.AssuranceRunError, match="explicit review"):
        a.read_reviewed_tabular_source(path, incomplete)


def test_raw_extraction_qualifies_tabular_sources_and_preparation_replays_all_rows(
    tmp_path,
):
    a = load_assurance()
    runner = load_script_module(
        "reviewed_tabular_raw_runner", SCRIPTS / "raw_input_runner.py"
    )
    source = tmp_path / "inputs"
    source.mkdir()
    population = source / "population.xlsx"
    bank = source / "bank.csv"
    write_source(population)
    write_source(bank)
    assumptions = {
        "assurance_run_date": "2026-10-18",
        "scope_year": 2026,
        "reviewed_source_decisions": {
            population.name: spec(population),
            bank.name: spec(bank, "bank_statement"),
        },
    }
    extracted = runner.extract_normalized_records(
        source, assumptions, output_dir=tmp_path / "extract"
    )
    assert len(extracted["open_items"]) == len(extracted["evidence_rows"]) == 2
    assert all(x["status"] == "qualified" for x in extracted["source_qualifications"])
    args = dict(
        output_dir=tmp_path / "prepared",
        open_items=extracted["open_items"],
        evidence_rows=extracted["evidence_rows"],
        assumptions=assumptions,
        source_root=source,
        source_receipts=extracted["source_artifact_receipts"],
        reviewed_source_decisions=extracted["reviewed_source_decision_receipts"],
        source_qualifications=extracted["source_qualifications"],
    )
    a.prepare_assurance_run(**args)
    forged = copy.deepcopy(extracted["evidence_rows"])
    forged[0]["amount"] = "999.00"
    with pytest.raises(a.AssuranceRunError, match="source cells"):
        a.prepare_assurance_run(
            **{**args, "output_dir": tmp_path / "forged", "evidence_rows": forged}
        )


def test_source_replay_cache_is_bound_to_the_reviewed_mapping(tmp_path):
    a = load_assurance()
    path = tmp_path / "bank.csv"
    write_source(path)
    receipts, first = reviewed(a, path, spec(path, "bank_statement"))
    rows = a.read_reviewed_tabular_source(path, first)
    cache = {}
    a._source_value_locator(
        row=rows[0],
        source_root=tmp_path,
        receipt=receipts[0],
        decision=first,
        tabular_cache=cache,
    )
    _, changed = reviewed(a, path, spec(path, "journal"))
    with pytest.raises(a.AssuranceRunError, match="stale"):
        a._source_value_locator(
            row=rows[0],
            source_root=tmp_path,
            receipt=receipts[0],
            decision=changed,
            tabular_cache=cache,
        )


@pytest.mark.parametrize(
    "field,value",
    [
        ("delimiter", []),
        ("default_evidence_type", []),
        ("evidence_type_map", {"x": []}),
    ],
)
def test_malformed_review_layout_is_an_explicit_validation_error(
    tmp_path, field, value
):
    a = load_assurance()
    path = tmp_path / "source.csv"
    write_source(path)
    proposed = spec(path)
    proposed["tabular"][field] = value
    receipts = a.build_source_receipts(tmp_path, [path])
    decisions, errors = a.build_reviewed_source_decisions(
        input_root=tmp_path,
        source_receipts=receipts,
        adapter_families={path.name: "reviewed_tabular_accounting_v1"},
        assumptions={
            "assurance_run_date": "2026-10-18",
            "reviewed_source_decisions": {path.name: proposed},
        },
    )
    assert not decisions and path.name in errors


@pytest.mark.parametrize(
    "dimensions", ["normal", "missing", "stale", "other_invoice_year"]
)
def test_tabular_sources_generate_workpapers_and_replay_final_assurance(
    tmp_path, dimensions
):
    a = load_assurance()
    runner = load_script_module("tabular_full_raw", SCRIPTS / "raw_input_runner.py")
    workflow = load_script_module(
        "tabular_full_workflow", SCRIPTS / "reconciliation_workflow.py"
    )
    source = tmp_path / "inputs"
    source.mkdir()
    population = source / "population.xlsx"
    bank = source / "bank.csv"
    write_source(population)
    write_source(bank)
    if dimensions == "other_invoice_year":
        # Same printed invoice numbers refer to different source years.
        bank.write_text(bank.read_text().replace("/09/2026", "/09/2025"))
    if dimensions in {"missing", "stale"}:
        original = population.read_bytes()
        rewritten = io.BytesIO()
        with (
            zipfile.ZipFile(io.BytesIO(original)) as source_zip,
            zipfile.ZipFile(rewritten, "w") as target_zip,
        ):
            for info in source_zip.infolist():
                data = source_zip.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    tree = ET.fromstring(data)
                    dim = tree.find(
                        "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}dimension"
                    )
                    assert dim is not None
                    if dimensions == "missing":
                        tree.remove(dim)
                    else:
                        dim.set("ref", "A1:A1")
                    data = ET.tostring(tree)
                target_zip.writestr(info, data)
        population.write_bytes(rewritten.getvalue())
    assumptions = {
        "assurance_run_date": "2026-10-18",
        "scope_year": 2026,
        "amount_tolerance": "0",
        "reviewed_source_decisions": {
            population.name: spec(population),
            bank.name: spec(bank, "bank_statement"),
        },
    }
    extracted = runner.extract_normalized_records(
        source, assumptions, output_dir=tmp_path / "extracted"
    )
    output = tmp_path / "workpapers"
    result = workflow.build_reconciliation_artifacts(
        output_dir=output,
        open_items=extracted["open_items"],
        evidence_rows=extracted["evidence_rows"],
        assumptions=assumptions,
        source_inventory=extracted["source_inventory"],
        source_qualifications=extracted["source_qualifications"],
        source_artifact_root=source,
        source_artifact_receipts=extracted["source_artifact_receipts"],
        reviewed_source_decision_receipts=extracted[
            "reviewed_source_decision_receipts"
        ],
        require_completed_review=False,
    )
    expected_status = "unresolved" if dimensions == "other_invoice_year" else "closed"
    assert [r["reconciliation_status"] for r in result["reconciliation_rows"]] == [
        expected_status,
        expected_status,
    ]
    assert Path(result["excel_path"]).is_file()
    assert Path(result["word_path"]).is_file()
    replay = a.validate_assurance_run(output)
    assert replay["content_sha256"] == result["assurance"]["content_sha256"]
    # A final sealed output must remain bound to the original source bytes.
    bank.write_text(bank.read_text().replace("100.25", "99.25"))
    with pytest.raises(
        ValueError, match="artifact receipt does not match current bytes"
    ):
        a.validate_assurance_run(output)


def test_representative_reviewed_exports_replay_the_full_settlement_case(tmp_path):
    import json
    import shutil
    from decimal import Decimal

    a = load_assurance()
    raw = load_script_module(
        "representative_tabular_raw", SCRIPTS / "raw_input_runner.py"
    )
    workflow = load_script_module(
        "representative_tabular_workflow", SCRIPTS / "reconciliation_workflow.py"
    )
    fixtures = Path(__file__).parent / "fixtures" / "reviewed_settlement_case"
    source = tmp_path / "inputs"
    source.mkdir()
    assumptions = json.loads((fixtures / "reviewed-assumptions.json").read_text())
    reviewed = assumptions.pop("reviewed_source_decisions")
    nested_reviewed = {}
    for index, filename in enumerate(
        ("partite_aperte.xlsx", "movimenti_contabili.xlsx", "movimenti_banca.csv")
    ):
        relative = Path("imports") / f"input_{index:024x}" / filename
        target = source / relative
        target.parent.mkdir(parents=True)
        shutil.copyfile(fixtures / filename, target)
        nested_reviewed[relative.as_posix()] = reviewed[filename]
    assumptions["reviewed_source_decisions"] = nested_reviewed
    extracted = raw.extract_normalized_records(
        source, assumptions, output_dir=tmp_path / "extraction"
    )
    assert len(extracted["open_items"]) == 6
    assert len(extracted["evidence_rows"]) == 7
    output = tmp_path / "workpapers"
    result = workflow.build_reconciliation_artifacts(
        output_dir=output,
        open_items=extracted["open_items"],
        evidence_rows=extracted["evidence_rows"],
        assumptions=assumptions,
        source_inventory=extracted["source_inventory"],
        source_qualifications=extracted["source_qualifications"],
        source_artifact_root=source,
        source_artifact_receipts=extracted["source_artifact_receipts"],
        reviewed_source_decision_receipts=extracted[
            "reviewed_source_decision_receipts"
        ],
        extraction_errors=extracted["extraction_errors"],
        normalized_records=extracted["normalized_records"],
        language="it",
        require_completed_review=False,
    )
    rows = {row["source_record_id"]: row for row in result["reconciliation_rows"]}
    assert rows["P001"]["allocated_amount"] == "600.00"
    assert rows["P001"]["residual_amount"] == "400.00"
    assert rows["P002"]["reconciliation_status"] == "closed"
    assert rows["P003"]["rule_applied"] == "exact_document_partial_compensation"
    assert rows["P003"]["residual_amount"] == "600.00"
    assert rows["P004"]["reconciliation_status"] == "unresolved"
    assert rows["P005"]["reconciliation_status"] == "unresolved"
    assert rows["P006"]["reconciliation_status"] == "needs_evidence"
    assert sum(
        Decimal(
            row.get(
                "residual_amount",
                "0" if row["reconciliation_status"] == "closed" else row["amount"],
            )
        )
        for row in rows.values()
    ) == Decimal("1900")
    assert (
        a.validate_assurance_run(output)["content_sha256"]
        == result["assurance"]["content_sha256"]
    )


@pytest.mark.parametrize(
    "alter_request_amount, decision_key_style",
    [(False, "relative"), (True, "relative"), (False, "absolute")],
)
def test_representative_managed_raw_producer_replays_followup_workbook(
    tmp_path, monkeypatch, alter_request_amount, decision_key_style
):
    import json
    from decimal import Decimal

    a = load_assurance()
    raw = load_script_module(
        "representative_managed_raw", SCRIPTS / "raw_input_runner.py"
    )
    ledger = load_script_module(
        "representative_managed_ledger",
        SCRIPTS.parents[1] / "studio-archive" / "scripts" / "client_ledger.py",
    )
    fixtures = Path(__file__).parent / "fixtures" / "reviewed_settlement_case"
    client_root = tmp_path / "Studio customer"
    client_root.mkdir()
    client_id = "client_333333333333333333333333"
    ledger.create_client_manifest(client_root, client_id)
    engagement = ledger.create_engagement(
        client_root, client_id, "Reviewed source case"
    )
    input_ids = []
    for filename in (
        "partite_aperte.xlsx",
        "movimenti_contabili.xlsx",
        "movimenti_banca.csv",
    ):
        imported = ledger.import_document(
            client_root,
            client_id,
            engagement["engagement_id"],
            fixtures / filename,
            "source",
        )
        input_ids.append(imported["receipt"]["input_id"])
    version = json.loads(
        (SCRIPTS.parent / ".codex-plugin" / "plugin.json").read_text()
    )["version"]
    prepared = ledger.prepare_run(
        client_root,
        client_id,
        engagement["engagement_id"],
        "open-item-reconciliation",
        version,
        input_ids=input_ids,
    )
    started = ledger.start_run(
        client_root, engagement["engagement_id"], prepared["run"]["run_id"]
    )
    context = started["context"]
    source = Path(context["input_dir"])
    assumptions = json.loads((fixtures / "reviewed-assumptions.json").read_text())
    reviewed = assumptions.pop("reviewed_source_decisions")
    assumptions["reviewed_source_decisions"] = {
        Path(b["path"]).relative_to(source).as_posix(): reviewed[Path(b["path"]).name]
        for b in context["input_bindings"]
    }
    if decision_key_style == "absolute":
        assumptions["reviewed_source_decisions"] = {
            str(source / name): decision
            for name, decision in assumptions["reviewed_source_decisions"].items()
        }
    output = Path(context["output_dir"]) / "reconciliation"
    if alter_request_amount:
        from openpyxl import load_workbook

        original = raw.write_missing_evidence_workbook

        def altered_request(path, pack, **kwargs):
            written = original(path, pack, **kwargs)
            workbook = load_workbook(written)
            changed = False
            for sheet in workbook:
                headers = [cell.value for cell in sheet[1]]
                if "id_riga_workpaper" not in headers or "importo" not in headers:
                    continue
                id_col = headers.index("id_riga_workpaper") + 1
                amount_col = headers.index("importo") + 1
                for row in range(2, sheet.max_row + 1):
                    value = str(sheet.cell(row, id_col).value or "")
                    if value.endswith("partite_aperte.xlsx:Partite:r2"):
                        sheet.cell(row, amount_col).value = "1000.00"
                        changed = True
            assert changed
            workbook.save(written)
            workbook.close()
            return written

        monkeypatch.setattr(raw, "write_missing_evidence_workbook", altered_request)
        with pytest.raises(ValueError, match="request row amount does not match"):
            raw.run_raw_input_reconciliation(
                input_dir=source,
                prepared_client_engagement=context,
                assumptions=assumptions,
                language="it",
                output_subdirectory="reconciliation",
            )
        assert not (output / "assurance_receipts.json").exists()
        return
    result = raw.run_raw_input_reconciliation(
        input_dir=source,
        prepared_client_engagement=context,
        assumptions=assumptions,
        language="it",
        output_subdirectory="reconciliation",
    )
    output = Path(context["output_dir"]) / "reconciliation"
    assert (output / "richieste_mirate_evidenze.xlsx").is_file()
    from docx import Document

    report_path = next(output.glob("*.docx"))
    report = Document(report_path)
    report_text = "\n".join(
        [paragraph.text for paragraph in report.paragraphs]
        + [
            cell.text
            for table in report.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    assert (
        "Pagamento bancario parziale con riferimento esatto al documento" in report_text
    )
    assert "Altra evidenza esterna" in report_text
    assert "Exact document partial bank payment" not in report_text
    assert "other_external" not in report_text
    assert "Evidenza interna debole" in report_text
    assert "Evidence weak internal" not in report_text
    aging = result["aging_summary"]
    assert sum(Decimal(row["partially_paid_amount"]) for row in aging) == Decimal(
        "1750.00"
    )
    for row in aging:
        assert sum(
            Decimal(row[field])
            for field in (
                "closed_amount",
                "partially_paid_amount",
                "probable_payment_amount",
                "open_supported_amount",
                "needs_evidence_amount",
                "unresolved_amount",
            )
        ) == Decimal(row["amount_total"])
    assert "Importo delle partite regolate parzialmente" in report_text
    assert "Tolleranza sugli importi: 0" in report_text
    assert "Conferma umana: No" in report_text
    assert "Approvazione professionale: No" in report_text
    assert "Prova sintetica automatizzata" in report_text
    assert "Amount tolerance" not in report_text
    assert "Source review notes" not in report_text
    assert "source_path_keys" not in report_text
    assert "/studio-private/" not in report_text
    prepared = json.loads((output / "prepared_records.json").read_text())
    assert (
        prepared["assumptions"]["source_review_notes"]
        == assumptions["source_review_notes"]
    )
    rows = {r["source_record_id"]: r for r in result["reconciliation_rows"]}
    assert rows["P001"]["residual_amount"] == "400.00"
    assert rows["P003"]["residual_amount"] == "600.00"
    assert (
        a.validate_assurance_run(output)["content_sha256"]
        == result["assurance"]["content_sha256"]
    )
