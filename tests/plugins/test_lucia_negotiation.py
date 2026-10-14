from __future__ import annotations

import hashlib
import json
import socket
import sys
from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from lxml import etree as ET
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "plugins/lucia/scripts"))
from legal_documents import prepare
from legal_docx import W
from legal_negotiation import (
    inventory_changes,
    main,
    render_positions,
    scaffold,
    validate_positions,
    verify_delivery,
)


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def refuse(*args, **kwargs):
        raise AssertionError("Negotiation helpers must not contact a service")

    monkeypatch.setattr(socket.socket, "connect", refuse)
    monkeypatch.setattr(socket, "create_connection", refuse)


def tracked_document(tmp_path: Path) -> Path:
    source = tmp_path / "tracked.docx"
    document = Document()
    document.add_paragraph("1. Limite di responsabilità.")
    document.add_paragraph("2. Foro di Milano.")
    document.add_paragraph("3. Definizione di cliente.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "4. Allegato A"
    document.save(source)
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = ET.fromstring(parts["word/document.xml"])
    paragraphs = list(root.iter(W + "p"))
    # Duplicate Word IDs cannot collapse distinct revision records.
    for paragraph, kind, text in (
        (paragraphs[0], "del", "EUR 100.000"),
        (paragraphs[0], "ins", "EUR 50.000"),
        (paragraphs[1], "del", "Milano"),
        (paragraphs[1], "ins", "Roma"),
        (paragraphs[2], "moveFrom", "Affiliati"),
        (paragraphs[2], "moveTo", "Affiliati e controllate"),
    ):
        revision = ET.SubElement(
            paragraph,
            W + kind,
            {
                W + "id": "7",
                W + "author": "Controparte",
                W + "date": "2026-10-14T08:00:00Z",
            },
        )
        run = ET.SubElement(revision, W + "r")
        ET.SubElement(run, W + ("delText" if kind == "del" else "t")).text = text
    properties = ET.SubElement(paragraphs[2], W + "pPr")
    change = ET.SubElement(
        properties, W + "pPrChange", {W + "id": "8", W + "author": "Controparte"}
    )
    ET.SubElement(ET.SubElement(change, W + "pPr"), W + "numPr")
    row = next(root.iter(W + "tr"))
    row_properties = ET.SubElement(row, W + "trPr")
    ET.SubElement(
        row_properties, W + "ins", {W + "id": "9", W + "author": "Controparte"}
    )
    ET.SubElement(paragraphs[0], W + "commentRangeStart", {W + "id": "1"})
    ET.SubElement(paragraphs[1], W + "commentRangeEnd", {W + "id": "1"})
    parts["word/document.xml"] = ET.tostring(root)
    comments = ET.Element(W + "comments", nsmap={"w": W[1:-1]})
    for cid, text in (
        ("1", "Il commento dichiara che nulla cambia."),
        ("2", "Commento senza ancora."),
    ):
        node = ET.SubElement(
            comments, W + "comment", {W + "id": cid, W + "author": "Controparte"}
        )
        ET.SubElement(
            ET.SubElement(ET.SubElement(node, W + "p"), W + "r"), W + "t"
        ).text = text
    parts["word/comments.xml"] = ET.tostring(comments)
    footer = ET.Element(W + "ftr", nsmap={"w": W[1:-1]})
    paragraph = ET.SubElement(footer, W + "p")
    insert = ET.SubElement(
        paragraph, W + "ins", {W + "id": "10", W + "author": "Controparte"}
    )
    ET.SubElement(ET.SubElement(insert, W + "r"), W + "t").text = "Rinvio modificato."
    parts["word/footer1.xml"] = ET.tostring(footer)
    parts["word/people.xml"] = (
        b'<w15:people xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml">'
        b'<w15:person w15:author="Controparte"/></w15:people>'
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return source


def review_case(tmp_path: Path):
    source = tracked_document(tmp_path)
    run = tmp_path / "review"
    pack = prepare(run, [source], "revisione-documentale", ["Posizioni negoziali"])
    review = json.loads(scaffold(run, "D001").read_text())
    inventory = inventory_changes(run, "D001")
    review["context"].update(
        {
            key: "Non fornito; confronto testuale sintetico"
            for key in review["context"]
            if key != "assumptions"
        }
    )
    review["context"]["language"] = "it"
    review["positions"] = [
        {
            "id": "P1",
            "title": "Modifiche negoziali da valutare",
            "effect": "Modifica il limite e il foro.",
            "grouping_reason": "Raggruppamento dichiarato solo per la prova tecnica.",
            "change_ids": [c["id"] for c in inventory["changes"]],
            "citations": [
                {
                    "source_id": "D001",
                    "anchor": "word/document.xml#p1",
                    "quote": "Limite di responsabilità.",
                }
            ],
            "response": "",
        }
    ]
    review["comment_reviews"] = [
        {
            "comment_id": c["id"],
            "analysis": "Annotazione distinta dall'effetto del testo.",
            "position_ids": ["P1"],
        }
        for c in inventory["comments"]
    ]
    return run, source, review, pack


def save_review(run: Path, review: dict) -> Path:
    path = run / "negotiation.json"
    path.write_text(json.dumps(review, ensure_ascii=False), encoding="utf-8")
    return path


def test_inventory_retains_text_moves_formatting_rows_footer_and_orphan_comments(
    tmp_path,
):
    run, source, _, _ = review_case(tmp_path)
    original = source.read_bytes()

    inventory = inventory_changes(run, "D001")

    assert len(inventory["changes"]) == 9
    assert len({c["id"] for c in inventory["changes"]}) == 9
    assert {c["kind"] for c in inventory["changes"]} == {
        "ins",
        "del",
        "moveFrom",
        "moveTo",
        "pPrChange",
    }
    assert any(
        c["locator"].startswith("word/footer1.xml#") for c in inventory["changes"]
    )
    assert any(
        c["kind"] == "ins" and not c["paragraph_anchor"] for c in inventory["changes"]
    )
    assert inventory["comments"][0]["anchors"] == [
        "word/document.xml#p1",
        "word/document.xml#p2",
    ]
    assert inventory["comments"][1]["anchors"] == []
    assert source.read_bytes() == original


@pytest.mark.parametrize(
    "problem",
    [
        "missing",
        "duplicate",
        "unknown",
        "missing-comment",
        "unknown-comment-position",
        "wrong-quote",
        "unsolicited-response",
        "wrong-inventory",
    ],
)
def test_validation_rejects_incomplete_or_inconsistent_accounting(tmp_path, problem):
    run, _, review, _ = review_case(tmp_path)
    position = review["positions"][0]
    if problem == "missing":
        position["change_ids"].pop()
    elif problem == "duplicate":
        review["other_changes"] = [
            {"change_id": position["change_ids"][0], "reason": "Duplicated assignment."}
        ]
    elif problem == "unknown":
        position["change_ids"].append("C99999")
    elif problem == "missing-comment":
        review["comment_reviews"].pop()
    elif problem == "unknown-comment-position":
        review["comment_reviews"][0]["position_ids"] = ["P99"]
    elif problem == "wrong-quote":
        position["citations"][0]["quote"] = "Una clausola inesistente"
    elif problem == "unsolicited-response":
        position["response"] = "Accettiamo"
    else:
        review["inventory_sha256"] = "0" * 64

    result = validate_positions(run, review)

    assert not result["valid"]
    assert not result["accounted"]
    assert result["errors"]


def test_visual_pending_does_not_become_professional_approval_from_complete_counts(
    tmp_path,
):
    run, _, review, _ = review_case(tmp_path)

    result = validate_positions(run, review)

    assert result["valid"]
    assert result["counts"]["unaccounted"] == 0
    assert "Visual review is pending" in " ".join(result["limitations"])
    assert "not proof of legal correctness" in result["meaning"]


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
@pytest.mark.parametrize("product", ["Vera", "Lucia"])
def test_render_delivers_complete_ledger_and_safe_text_in_both_products(
    tmp_path, language, product
):
    run, source, review, _ = review_case(tmp_path)
    original = source.read_bytes()
    review["context"]["language"] = language
    review["positions"][0][
        "title"
    ] = '=HYPERLINK("https://example.invalid") <script>alert(1)</script>'
    path = save_review(run, review)

    target = render_positions(run, path, product)

    page = target.read_text()
    assert f"lang='{language}'" in page
    assert product in page
    assert "<script>" not in page
    assert "&lt;script&gt;" in page
    workbook = load_workbook(run / "negotiation.xlsx", data_only=False)
    assert workbook.worksheets[0]["A2"].data_type == "s"
    assert workbook.worksheets[1].max_row == 10
    assert workbook.worksheets[2].max_row == 3
    assert (
        (run / "negotiation.csv")
        .read_text(encoding="utf-8-sig")
        .splitlines()[1]
        .startswith("\"'=")
    )
    receipt = json.loads((run / "negotiation-receipt.json").read_text())
    assert (
        receipt["outputs"]["negotiation.html"]
        == hashlib.sha256(target.read_bytes()).hexdigest()
    )
    assert source.read_bytes() == original


def test_failed_rerender_withdraws_all_previous_outputs(tmp_path):
    run, _, review, _ = review_case(tmp_path)
    path = save_review(run, review)
    render_positions(run, path)
    review["positions"][0]["change_ids"].pop()
    save_review(run, review)

    with pytest.raises(ValueError, match="accounting failed"):
        render_positions(run, path)

    assert not (run / "negotiation.html").exists()
    assert not (run / "negotiation.xlsx").exists()
    assert not (run / "negotiation.csv").exists()
    assert not (run / "negotiation-receipt.json").exists()
    assert (run / "previous-1-negotiation.html").is_file()
    assert not json.loads((run / "negotiation-validation.json").read_text())["valid"]


def test_inventory_rehash_cannot_hide_a_fabricated_change(tmp_path):
    run, _, review, _ = review_case(tmp_path)
    path = run / "negotiation-inventory.json"
    inventory = json.loads(path.read_text())
    inventory["changes"][0]["text"] = "Invented source text"
    path.write_text(json.dumps(inventory))
    (run / "negotiation-inventory.sha256").write_text(
        hashlib.sha256(path.read_bytes()).hexdigest()
    )

    with pytest.raises(ValueError, match="does not match"):
        validate_positions(run, review)


def test_resume_refuses_reinitialisation_and_detects_source_drift(tmp_path):
    run, _, review, _ = review_case(tmp_path)
    with pytest.raises(ValueError, match="resume"):
        scaffold(run, "D001")
    (run / "originals/D001.docx").write_bytes(b"Changed source")

    with pytest.raises(ValueError, match="snapshot changed"):
        validate_positions(run, review)


def test_second_round_preserves_earlier_snapshot_and_requires_full_position_mapping(
    tmp_path,
):
    previous_dir = tmp_path / "previous"
    previous_dir.mkdir()
    prior_run, _, prior_review, _ = review_case(previous_dir)
    prior_path = save_review(prior_run, prior_review)
    before = prior_path.read_bytes()
    current_dir = tmp_path / "current"
    current_dir.mkdir()
    source = tracked_document(current_dir)
    run = current_dir / "review"
    prepare(run, [source], "revisione-documentale", ["Nuovo round"])
    review = json.loads(scaffold(run, "D001", prior_path).read_text())
    for key in ("context", "positions", "comment_reviews"):
        review[key] = deepcopy(prior_review[key])
    review["round_comparison"] = [
        {
            "previous_position_ids": ["P1"],
            "current_position_ids": ["P1"],
            "status": "unchanged",
            "analysis": "Posizioni rilette nella prova sintetica.",
        }
    ]

    result = validate_positions(run, review)

    assert result["valid"]
    assert prior_path.read_bytes() == before
    assert (run / "negotiation-prior.json").is_file()
    assert "not automatically established" in " ".join(result["limitations"])
    review["round_comparison"] = []
    assert not validate_positions(run, review)["valid"]


def test_no_tracked_changes_does_not_fabricate_a_negotiation_inventory(tmp_path):
    source = tmp_path / "clean.docx"
    document = Document()
    document.add_paragraph("Un contratto senza revisioni.")
    document.save(source)
    run = tmp_path / "review"
    prepare(run, [source], "revisione-documentale", ["Posizioni"])

    with pytest.raises(ValueError, match="No tracked revisions"):
        scaffold(run, "D001")

    assert not (run / "negotiation.json").exists()


def test_cli_renders_real_saved_artifacts(tmp_path):
    run, _, review, _ = review_case(tmp_path)
    path = save_review(run, review)

    result = main(
        ["render", "--run-dir", str(run), "--review", str(path), "--product", "Vera"]
    )

    assert result == 0
    assert (run / "negotiation.html").is_file()


def test_delivery_verifies_retained_review_then_rejects_changed_html(tmp_path):
    run, _, review, _ = review_case(tmp_path)
    path = save_review(run, review)
    render_positions(run, path)
    assert verify_delivery(run)["accounted"]
    assert (run / "negotiation-reviewed.json").is_file()
    (run / "negotiation.html").write_text("Changed report")

    with pytest.raises(ValueError, match="delivery changed"):
        verify_delivery(run)


def test_symlink_output_is_rejected_without_writing_foreign_file(tmp_path):
    run, _, review, _ = review_case(tmp_path)
    path = save_review(run, review)
    foreign = tmp_path / "foreign.txt"
    foreign.write_text("Do not change")
    try:
        (run / "negotiation.xlsx").symlink_to(foreign)
    except OSError:
        pytest.skip("This host does not permit symlinks")

    with pytest.raises(ValueError, match="ordinary local files"):
        render_positions(run, path)

    assert foreign.read_text() == "Do not change"


def test_source_drift_after_delivery_withdraws_outputs_and_clears_valid_status(
    tmp_path,
):
    run, _, review, _ = review_case(tmp_path)
    path = save_review(run, review)
    render_positions(run, path)
    (run / "originals/D001.docx").write_bytes(b"Changed original")

    with pytest.raises(ValueError, match="snapshot changed"):
        render_positions(run, path)

    assert not (run / "negotiation.html").exists()
    assert not (run / "negotiation-reviewed.json").exists()
    assert not json.loads((run / "negotiation-validation.json").read_text())["valid"]
