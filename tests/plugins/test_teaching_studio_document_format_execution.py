"""Exercise actual course sources with explicitly fictional review decisions."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from tests.plugins._teaching_release import record_native_check
from tests.plugins.test_studio_document_format import (
    ROOT,
    _command,
    _helpers,
    _json,
    _native,
)

__all__ = []
REVIEWER = "Synthetic course reviewer, not the learner"


def read(path: Path) -> dict:
    """Read one actual local regression artifact."""
    return json.loads(path.read_text(encoding="utf-8"))


def approve(review: Path) -> None:
    """Approve only this fictional regression proposal after checking previews."""
    manifest = read(review / "preview_manifest.json")
    assert len(manifest["outputs"]) == 2
    assert Document(review / "preview-short.docx").paragraphs
    assert Document(review / "preview-long.docx").tables
    _native(
        "approve",
        "--review-dir",
        str(review),
        "--review-digest",
        read(review / "format_review.json")["review_digest"],
        "--reviewer",
        REVIEWER,
        "--confirmed-by-user",
    )


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
@pytest.mark.parametrize(
    "phase,expected_total", [("demo", "14500"), ("practice", "25100")]
)
def test_course_format_preview_adoption_report_and_revision(
    tmp_path: Path, language: str, phase: str, expected_total: str, record_property
) -> None:
    """Run the selected raw exercise through the real format and report pipelines."""
    spec_path = (
        ROOT
        / f"scripts/course_materials/inputs/studio-document-format/{phase}-{language}.json"
    )
    spec = read(spec_path)
    outputs = tmp_path / "outputs"
    outputs.mkdir()
    material = outputs / "sources"
    _native(
        "practice-materials", "--spec", str(spec_path), "--output-dir", str(material)
    )
    evidence = outputs / "format-evidence.json"
    _native(
        "inspect",
        "--sample",
        str(material / "studio-example.docx"),
        "--output",
        str(evidence),
    )
    assert read(evidence)["samples"][0]["sections"][0]["margins_mm"][
        "left"
    ] == pytest.approx(spec["margins_mm"]["left"], abs=0.02)
    assert spec["body"] not in evidence.read_text()
    studio_id = spec["studio_name"].replace(" ", "-")
    workspace = outputs / studio_id
    _command(
        "initialize_workspace.py",
        "--workspace",
        str(workspace),
        "--workspace-id",
        studio_id,
        "--owner",
        REVIEWER,
        "--retention-owner",
        REVIEWER,
        "--confirmed-by-user",
    )
    brand = _helpers("test_comunicazione_professionale")._brand()
    brand["studio_name"] = spec["studio_name"]
    settings = {
        "font_family": spec["font"],
        "body_font_size_pt": spec["size_pt"],
        "heading_1_size_pt": spec["heading_pt"],
        "heading_color": spec["color"],
        "margins_mm": spec["margins_mm"],
        "header_text": spec["studio_name"],
        "footer_text": spec["footer"],
        "page_numbers": True,
    }
    _native(
        "prepare",
        "--workspace",
        str(workspace),
        "--review-id",
        "initial",
        "--settings",
        str(_json(outputs / "proposed-settings.json", settings)),
        "--sample",
        str(material / "studio-example.docx"),
        "--brand",
        str(_json(outputs / "brand.json", brand)),
        "--language",
        language,
    )
    initial = workspace / "runs/initial"
    denied = _native(
        "approve",
        "--review-dir",
        str(initial),
        "--review-digest",
        read(initial / "format_review.json")["review_digest"],
        "--reviewer",
        REVIEWER,
        check=False,
    )
    assert denied.returncode != 0
    assert not (workspace / "studio_profile.json").exists()
    approve(initial)
    adopted = read(workspace / "studio_profile.json")
    descriptor = _json(
        tmp_path / "tutorial_case.json",
        {
            "tutorial": True,
            "local_only": True,
            "workflow_id": "studio-document-format",
            "directory": str(tmp_path),
            "output_dir": str(outputs),
        },
    )
    _native(
        "report-case",
        "--tutorial-case",
        str(descriptor),
        "--source",
        str(material / "financial-source.csv"),
        "--studio-workspace",
        str(workspace),
    )
    reuse = read(tmp_path / "report_reuse_case.json")
    helper = _helpers("test_report_builder_plugin")
    core = helper.load_core()
    context = read(Path(reuse["client_engagement_path"]))
    source = next(Path(reuse["input_dir"]).rglob("*.csv"))
    report_root = Path(reuse["output_dir"])
    inspected = core.inspect_inputs(
        source,
        report_root / "inspection",
        language=language,
        document_language=language,
    )
    table_id = inspected.inspection["tables"][0]["table_id"]
    recipe = inspected.suggested_recipe
    recipe.update(entity="Fictional course company", period=spec["report_period"])
    recipe["sections"]["income_statement"]["assigned_table"] = table_id
    recipe = core.review_numeric_measure_columns(
        inspected.inspection,
        recipe,
        section_key="income_statement",
        **helper._numeric_review_args(inspected.inspection, table_id, ["amount"]),
        reviewer_ref="synthetic.studio-course",
        reviewed_on="2026-10-01",
        numeric_locale="en",
        currency="EUR",
        unit="currency",
        scale="1",
        parse_policy="strict_all_nonblank_v1",
    )
    recipe_path = report_root / "recipe.json"
    core.write_json(recipe_path, recipe)
    report = report_root / "initial-report"
    core.build_report(
        source,
        report,
        recipe_path=recipe_path,
        run_id=context["run_id"],
        client_engagement=context,
        studio_workspace=workspace,
        studio_id=studio_id,
        studio_name=spec["studio_name"],
    )
    analysis = read(report / "report_analysis.json")
    assert analysis["sections"][1]["numeric_columns"][0]["sum"] == expected_total
    assert Document(report / "report.docx").styles["Normal"].font.name == spec["font"]
    assert read(report / "used_recipe.json")["studio_format"]["studio_id"] == studio_id
    _native(
        "prepare",
        "--workspace",
        str(workspace),
        "--review-id",
        "revision",
        "--settings",
        str(_json(outputs / "revised-settings.json", {"paragraph_after_pt": 12})),
        "--language",
        language,
    )
    revision = workspace / "runs/revision"
    approve(revision)
    revised = read(workspace / "studio_profile.json")
    assert revised["version"] == 2
    assert read(workspace / "profiles/studio_profile-v001.json") == adopted
    assert revised["profile"]["email"] == adopted["profile"]["email"]
    assert revised["profile"]["document"]["docx"]["paragraph_after_pt"] == 12
    revised_report = report_root / "revised-report"
    core.build_report(
        source,
        revised_report,
        recipe_path=recipe_path,
        run_id=context["run_id"],
        client_engagement=context,
        studio_workspace=workspace,
        studio_id=studio_id,
        studio_name=spec["studio_name"],
    )
    assert read(revised_report / "report_analysis.json") == analysis
    assert (
        Document(revised_report / "report.docx").styles["Normal"].font.name
        == spec["font"]
    )
    assert (
        read(revised_report / "used_recipe.json")["studio_format"]["settings"][
            "paragraph_after_pt"
        ]
        == 12
    )
    assert (
        read(report / "used_recipe.json")["studio_format"]["settings"][
            "paragraph_after_pt"
        ]
        != 12
    )
    _json(
        outputs / "reviewed-artifacts.json",
        {
            "phase": phase,
            "language": language,
            "artifacts": [
                str(initial / "preview-short.docx"),
                str(initial / "preview-long.docx"),
                str(workspace / "studio_profile.json"),
                str(revised_report / "report.docx"),
                str(revised_report / "report_analysis.json"),
            ],
            "approval": "Explicitly simulated regression decisions; no learner acceptance",
        },
    )
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="studio-document-format",
        language=language,
        phase=phase,
    )
