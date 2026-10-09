from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

from tests.plugins import test_open_item_reconciliation_plugin as engine
from tests.plugins import test_vera_native_workspace as native


def isolated_engine(code: str, *arguments: object) -> dict:
    """Avoid reloading PyMuPDF native extensions between pytest cases."""
    completed = subprocess.run(
        [
            sys.executable,
            "-I",
            "-B",
            "-c",
            "import sys,json;sys.path.insert(0,sys.argv[1]);import raw_input_runner as runner;import audit_assurance as assurance;import review_server as review_server;"
            + code,
            str(engine.SCRIPT_DIR),
            *[json.dumps(value) for value in arguments],
        ],
        capture_output=True,
        text=True,
        check=True,
    )
    return json.loads(completed.stdout)


def prepared_raw_case(
    tmp_path: Path, monkeypatch, nested: bool = True, complete_review: bool = True
) -> dict:
    """Use real fictional teaching PDFs and the unchanged accounting adapters."""
    ledger = engine._load_customer_ledger()
    client = tmp_path / "Fictional open-item customer"
    client.mkdir()
    client_id = "client_" + "a7" * 12
    ledger.create_client_manifest(client, client_id)
    engagement = ledger.create_engagement(client, client_id, "Fictional reconciliation")
    sources = (
        native.ROOT / "plugins/vera/assets/courses/open-item-reconciliation/files/input"
    )
    receipts = [
        ledger.import_document(
            client, client_id, engagement["engagement_id"], sources / name, "source"
        )["receipt"]
        for name in ("open-items-it.pdf", "bank-march-it.pdf")
    ]
    prepared = ledger.prepare_run(
        client,
        client_id,
        engagement["engagement_id"],
        "open-item-reconciliation",
        "synthetic",
        input_ids=[row["input_id"] for row in receipts],
    )
    run = ledger.start_run(
        client, engagement["engagement_id"], prepared["run"]["run_id"]
    )
    context = run["context"]
    decisions = {}
    for row in context["input_bindings"]:
        path = Path(row["path"])
        population = path.name == "open-items-it.pdf"
        decisions[path.relative_to(Path(context["input_dir"])).as_posix()] = {
            "role": "open_items" if population else "bank_statement",
            "adapter_family": (
                "open_items_text_v1" if population else "bank_statement_text_v1"
            ),
            "reviewer_ref": "reviewer.synthetic",
            "reviewed_on": "2026-09-14",
            "perimeter": {
                "entity_ref": "entity.arco",
                "party_ref": "party.servizi_esempio",
                "currency": "EUR",
                "unit": "currency_amount",
                "direction_policy": "supplier",
                "allocation_policy": "one_to_one",
            },
            "money": {
                "decimal_separator": ",",
                "thousands_separator": ".",
                "reported_unit": "EUR",
                "reported_increment": "0.01",
            },
            "date": {"order": "day_first"},
        }
    parameters = dict(
        input_dir=context["input_dir"],
        prepared_client_engagement=context,
        output_subdirectory="reconciliation" if nested else None,
        title="Titolo della revisione fittizia",
        narrative="Testo professionale fittizio da conservare.",
        assumptions={
            "scope_year": "2026",
            "cutoff_date": "2026-03-31",
            "assurance_run_date": "2026-09-14",
            "currency": "EUR",
            "post_cutoff_events_excluded": True,
            "reviewed_source_decisions": decisions,
            "counterparty_keywords": ["servizi esempio"],
            "ocr_scanned": False,
        },
    )
    initial = isolated_engine(
        "\nresult=runner.run_raw_input_reconciliation(**json.loads(sys.argv[2]));output=runner.Path(result['run_output_dir']);checkpoint=json.loads((output/'assurance_receipts.json').read_bytes())['content_sha256'];review=json.loads((output/'review_payload.json').read_bytes());items=review['items'] if json.loads(sys.argv[3]) else review['items'][:1];review_server.apply_decisions(output,{'expected_predecessor_checkpoint':checkpoint,'reviewer':'fictional_reviewer','decisions':[{'item_id':row['id'],'action':'accept'} for row in items]});sys.stdout.write(json.dumps({'output':str(output),'checkpoint':checkpoint}))",
        parameters,
        complete_review,
    )
    output = Path(initial["output"])
    checkpoint = initial["checkpoint"]
    binding, _ = native.archived_review_binding(output, "open-item-reconciliation")
    native.configure(monkeypatch, tmp_path, [binding])
    service = native.workspace_module()
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    return {
        "output": output,
        "context": context,
        "service": service,
        "view": view,
        "request": {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "expected_predecessor_checkpoint": checkpoint,
            "human_reviewed": True,
            "idempotency_key": "fictional-regeneration",
        },
    }


@pytest.mark.parametrize("nested", [False, True])
def test_native_raw_regeneration_replays_full_package_preserves_sources_and_report(
    tmp_path, monkeypatch, nested
):
    case = prepared_raw_case(tmp_path, monkeypatch, nested)
    output = case["output"]
    source_bytes = {
        row["path"]: Path(row["path"]).read_bytes()
        for row in case["context"]["input_bindings"]
    }
    history = (
        output
        / "assurance_transition_history"
        / case["request"]["expected_predecessor_checkpoint"]
    )
    predecessor = engine._audit_tree_image(history)

    result = case["service"].dispatch(
        "vera_workspace_open_items_regenerate", case["request"]
    )

    assert result["regenerated"] is True
    assert result["revision"] != case["view"]["revision"]
    assert engine._audit_tree_image(history) == predecessor
    assert source_bytes == {p: Path(p).read_bytes() for p in source_bytes}
    assert (
        case["service"].dispatch(
            "vera_workspace_open_items_regenerate", case["request"]
        )
        == result
    )
    replay = isolated_engine(
        "sys.stdout.write(json.dumps(assurance.validate_assurance_run(runner.Path(json.loads(sys.argv[2])),expected_predecessor_checkpoint=json.loads(sys.argv[3]))))",
        str(output),
        case["request"]["expected_predecessor_checkpoint"],
    )
    assert replay["content_sha256"] == result["assurance_sha256"]
    assert replay["gate_register"]["gates"]["publication"]["status"] == "withheld"
    manifest = json.loads((output / "run_manifest.json").read_bytes())
    assert ".audit-review-transaction-" not in json.dumps(manifest)
    assert manifest["report_options"]["title"] == "Titolo della revisione fittizia"
    document = Document(output / "relazione_riconciliazione_audit.docx")
    paragraphs = "\n".join(
        [p.text for p in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    assert "Titolo della revisione fittizia" in paragraphs
    assert "Testo professionale fittizio da conservare." in paragraphs
    assert (output / "richieste_mirate_evidenze.xlsx").is_file()
    assert (output / "scheda_operativa_commercialista.xlsx").is_file()
    canonical = json.loads(
        (output / "assurance_final_outputs/reconciliation_results.json").read_bytes()
    )
    assert [
        row["reconciliation_status"] for row in canonical["reconciliation_rows"]
    ] == ["closed", "closed", "unresolved"]
    assert canonical["reconciliation_rows"][2]["amount"] == "488"
    assert all(row["review_status"] == "PASS" for row in canonical["review_rows"])
    reopened = case["service"].dispatch(
        "vera_workspace_view", {"work_ref": case["request"]["work_ref"]}
    )
    assert reopened["revision"] == result["revision"]
    assert reopened["data"]["ui_decisions"]["decisions"]


@pytest.mark.parametrize(
    "fault",
    [
        "checkpoint",
        "human",
        "viewer",
        "stale",
        "settings",
        "history",
        "source",
        "saved",
        "missing_settings",
        "cache",
    ],
)
def test_native_regeneration_rejects_missing_authority_without_mutation(
    tmp_path, monkeypatch, fault
):
    case = prepared_raw_case(tmp_path, monkeypatch)
    request = dict(case["request"])
    if fault == "checkpoint":
        request["expected_predecessor_checkpoint"] = "0" * 64
    if fault == "human":
        request["human_reviewed"] = False
    if fault == "viewer":
        monkeypatch.setenv("VERA_WORKSPACE_ROLES", "VIEWER")
    if fault == "stale":
        request["revision"] = "0" * 64
    if fault == "settings":
        path = case["output"] / "run_manifest.json"
        value = json.loads(path.read_bytes())
        value["report_options"]["narrative"] = "Substituted"
        path.write_text(json.dumps(value))
    if fault in {"missing_settings", "cache"}:
        path = case["output"] / "run_manifest.json"
        value = json.loads(path.read_bytes())
        if fault == "missing_settings":
            value.pop("report_options")
        else:
            value["assumptions"]["cache_dir"] = str(tmp_path / "outside-cache")
        path.write_text(json.dumps(value))
    if fault == "history":
        path = (
            case["output"]
            / "assurance_transition_history"
            / request["expected_predecessor_checkpoint"]
            / "predecessor_run/run_manifest.json"
        )
        path.write_text("{}")
    if fault == "source":
        Path(case["context"]["input_bindings"][0]["path"]).write_bytes(
            b"changed source"
        )
    if fault == "saved":
        path = case["output"] / "ui_decisions.json"
        value = json.loads(path.read_bytes())
        value["decisions"] = []
        value["decision_count"] = 0
        path.write_text(json.dumps(value))
    if fault in {"settings", "saved", "missing_settings", "cache"}:
        request["revision"] = case["service"].output_revision(
            Path(case["context"]["output_dir"]), "open-item-reconciliation"
        )
    before = engine._audit_tree_image(case["output"])

    with pytest.raises(
        (ValueError, PermissionError, sys.modules["client_ledger"].LedgerError)
    ):
        case["service"].dispatch("vera_workspace_open_items_regenerate", request)

    assert engine._audit_tree_image(case["output"]) == before


def test_native_regeneration_mcp_refuses_forged_ticket_before_engine(
    tmp_path, monkeypatch
):
    case = prepared_raw_case(tmp_path, monkeypatch)
    before = engine._audit_tree_image(case["output"])
    request = {
        "jsonrpc": "2.0",
        "id": 1,
        "method": "tools/call",
        "params": {
            "name": "vera_workspace_open_items_regenerate",
            "arguments": {**case["request"], "review_ticket": "forged.invalid"},
        },
    }

    completed = subprocess.run(
        [native.NODE, str(native.SERVER)],
        input=json.dumps(request) + "\n",
        capture_output=True,
        text=True,
        check=True,
    )

    assert json.loads(completed.stdout)["result"]["isError"] is True
    assert "Invalid review ticket" in completed.stdout
    assert engine._audit_tree_image(case["output"]) == before


def test_native_regeneration_keeps_incomplete_professional_review_blocked(
    tmp_path, monkeypatch
):
    case = prepared_raw_case(tmp_path, monkeypatch, complete_review=False)

    result = case["service"].dispatch(
        "vera_workspace_open_items_regenerate", case["request"]
    )

    assert result["regenerated"] is True
    assert result["report_ready"] is False
    canonical = json.loads(
        (
            case["output"] / "assurance_final_outputs/reconciliation_results.json"
        ).read_bytes()
    )
    assert [row["review_status"] for row in canonical["review_rows"]] == [
        "PASS",
        "UNRESOLVED",
        "UNRESOLVED",
    ]
    assert (
        json.loads((case["output"] / "final_artifacts.json").read_bytes())["status"]
        != "final_ready"
    )


def test_raw_regeneration_downstream_failure_rolls_back_entire_package(
    tmp_path, monkeypatch
):
    case = prepared_raw_case(tmp_path, monkeypatch)
    before = engine._audit_tree_image(case["output"])

    result = isolated_engine(
        """
def refuse_report(*args, **kwargs):
    raise ValueError('synthetic downstream report failure')
runner.write_missing_evidence_workbook=refuse_report
try:
    runner.regenerate_raw_input_reconciliation(runner.Path(json.loads(sys.argv[2])),json.loads(sys.argv[3]),expected_predecessor_checkpoint=json.loads(sys.argv[4]))
except ValueError as error:
    sys.stdout.write(json.dumps({'error':str(error)}))
else:
    sys.stdout.write('{}')
""",
        str(case["output"]),
        case["context"],
        case["request"]["expected_predecessor_checkpoint"],
    )

    assert result["error"] == "synthetic downstream report failure"
    assert engine._audit_tree_image(case["output"]) == before
