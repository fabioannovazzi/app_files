"""Real registered six-table preparation and retained public Treasury review."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from tests.plugins.test_treasury_delivery import csv_sources
from tests.plugins.test_treasury_forecast import accept, first
from tests.plugins.test_vera_client_workflow_filesystem import _load_customer_ledger
from tests.plugins.test_vera_native_archive_lifecycle import rpc_program
from tests.plugins.test_vera_native_workspace import configure, workspace_module


@pytest.fixture
def initial_treasury(tmp_path, monkeypatch, request):
    ledger = _load_customer_ledger()
    client = tmp_path / "Fictional treasury client"
    client.mkdir()
    client_id = "client_111111111111111111111111"
    ledger.create_client_manifest(client, client_id)
    engagement = ledger.create_engagement(client, client_id, "Synthetic treasury")
    staging = tmp_path / "staging"
    staging.mkdir()
    ids = []
    variant = getattr(request, "param", "csv")
    data = first()
    if variant in {"many", "invalid-last-row"}:
        data["accounts"] += [
            {"account_id": f"bank-extra-{index}", "balance": "0.00"}
            for index in range(1, 35)
        ]
        if variant == "invalid-last-row":
            data["accounts"][-1]["balance"] = "0.001"
    for name, content in csv_sources(data).items():
        source = staging / name
        source.write_text(content)
        ids.append(
            ledger.import_document(
                client, client_id, engagement["engagement_id"], source, "source"
            )["receipt"]["input_id"]
        )
    if variant == "xlsx":
        source = staging / "declared-balances.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Balances exact"
        sheet.append(["account_id", "balance"])
        sheet.append(["bank-a", "40000.00"])
        book.save(source)
        book.close()
        ids.append(
            ledger.import_document(
                client, client_id, engagement["engagement_id"], source, "source"
            )["receipt"]["input_id"]
        )
    if variant == "previous":
        data = first()
        data.update(client_id=client_id, engagement_id=engagement["engagement_id"])
        source = staging / "accepted-previous.json"
        source.write_text(json.dumps(accept(data)))
        ids.append(
            ledger.import_document(
                client, client_id, engagement["engagement_id"], source, "source"
            )["receipt"]["input_id"]
        )
    if variant == "xml":
        source = staging / "invoice.xml"
        source.write_text(
            """<FatturaElettronica><FatturaElettronicaHeader>
<CedentePrestatore><DatiAnagrafici><IdFiscaleIVA><IdPaese>IT</IdPaese><IdCodice>01234567890</IdCodice></IdFiscaleIVA><Anagrafica><Denominazione>Synthetic supplier</Denominazione></Anagrafica></DatiAnagrafici></CedentePrestatore>
<CessionarioCommittente><DatiAnagrafici><CodiceFiscale>99999999999</CodiceFiscale><Anagrafica><Denominazione>Synthetic customer</Denominazione></Anagrafica></DatiAnagrafici></CessionarioCommittente>
</FatturaElettronicaHeader><FatturaElettronicaBody><DatiGenerali><DatiGeneraliDocumento><TipoDocumento>TD01</TipoDocumento><Divisa>EUR</Divisa><Data>2026-09-01</Data><Numero>XML-1</Numero><ImportoTotaleDocumento>122.00</ImportoTotaleDocumento></DatiGeneraliDocumento></DatiGenerali>
<DatiPagamento><CondizioniPagamento>TP02</CondizioniPagamento><DettaglioPagamento><ModalitaPagamento>MP05</ModalitaPagamento><DataScadenzaPagamento>2026-09-30</DataScadenzaPagamento><ImportoPagamento>122.00</ImportoPagamento></DettaglioPagamento></DatiPagamento>
</FatturaElettronicaBody></FatturaElettronica>"""
        )
        ids.append(
            ledger.import_document(
                client, client_id, engagement["engagement_id"], source, "source"
            )["receipt"]["input_id"]
        )
    prepared = ledger.prepare_run(
        client,
        client_id,
        engagement["engagement_id"],
        "treasury-forecast",
        "0.1.0",
        input_ids=ids,
    )
    run = ledger.start_run(
        client, engagement["engagement_id"], prepared["run"]["run_id"]
    )
    context = run["context"]
    binding = {
        "work_ref": "initial-treasury",
        "client_root": str(client),
        "client_id": client_id,
        "engagement_id": engagement["engagement_id"],
        "run_id": context["run_id"],
        "workflow_id": "treasury-forecast",
    }
    return (
        configure(monkeypatch, tmp_path, [binding]),
        Path(context["output_dir"]),
        binding,
    )


SETUP = """
const setup=payload(call('vera_workspace_treasury_setup',{work_ref:'initial-treasury'}));
const authority={work_ref:setup.work_ref,revision:setup.revision,review_ticket:setup.review_ticket,expected_draft_revision:setup.draft.draft_revision};
"""


def save_program(fields: dict | None = None) -> str:
    data = first()
    values = (
        {
            k: data[k]
            for k in (
                "company_id",
                "company_name",
                "currency",
                "as_of",
                "horizon_end",
                "coverage",
            )
        }
        if fields is None
        else fields
    )
    return (
        SETUP
        + f"const fields={json.dumps(values)};"
        + (
            "fields.tables=Object.fromEntries(Object.keys(setup.headers).map(role=>[role,{input_id:setup.items.find(item=>item.title===role+'.csv').id}]));"
            if fields is None
            else ""
        )
        + """
const saved=payload(call('vera_workspace_treasury_draft_save',{...authority,fields}));
const exact={...authority,expected_draft_revision:saved.draft_revision};
const args={...exact,human_reviewed:true,idempotency_key:'fictional-treasury-preparation'};
"""
    )


def test_native_treasury_prepares_without_a_registered_manifest(initial_treasury):
    env, output, binding = initial_treasury
    result = rpc_program(
        env,
        save_program()
        + """
const inspected=payload(call('vera_workspace_treasury_inspect',{work_ref:setup.work_ref,revision:setup.revision,expected_draft_revision:saved.draft_revision,table:'open_items'}));
const prepared=payload(call('vera_workspace_treasury_prepare',args));
const reopened=payload(call('vera_workspace_treasury_setup',{work_ref:setup.work_ref}));
const view=payload(call('vera_workspace_view',{work_ref:setup.work_ref}));
const retry=payload(call('vera_workspace_treasury_prepare',args));
const result={setup,inspected,prepared,reopened,view,retry};
""",
    )
    assert len(result["setup"]["items"]) == 6
    assert result["inspected"]["opening_cash"] == "40000.00"
    assert result["inspected"]["total"] == 3
    assert result["prepared"] == result["retry"]
    assert result["prepared"]["professional_approval"] is False
    assert result["prepared"]["run_completed"] is False
    assert result["view"]["kind"] == "treasury"
    assert result["reopened"]["status"] == "prepared"
    assert result["reopened"]["can_prepare"] is False
    record = json.loads(
        (
            output / "versions" / result["prepared"]["record_sha256"] / "forecast.json"
        ).read_bytes()
    )
    assert record["status"] == "draft_for_review"
    assert (output / "native_treasury_intake.json").is_file()
    assert "human_reviewed" not in result["reopened"]["draft"]["fields"]
    assert workspace_module().load_binding(binding)["run"]["status"] == "running"


def test_native_treasury_catalogue_exposes_initial_preparation(initial_treasury):
    env, _, _ = initial_treasury
    result = rpc_program(env, "const result=payload(call('vera_workspace_open',{}));")
    assert result["works"][0]["setup_available"] is True


def test_native_treasury_partial_choices_reopen_without_outputs(initial_treasury):
    env, output, _ = initial_treasury
    result = rpc_program(
        env,
        save_program({"company_name": "Partial fictional company"})
        + "const recovered=payload(call('vera_workspace_treasury_setup',{work_ref:setup.work_ref}));const refused=call('vera_workspace_treasury_prepare',args);const result={recovered,refused};",
    )
    assert result["recovered"]["draft"]["fields"] == {
        "company_name": "Partial fictional company"
    }
    assert result["refused"]["isError"] is True
    assert "Declare company" in result["refused"]["content"][0]["text"]
    assert list(output.iterdir()) == []


@pytest.mark.parametrize(
    "fields",
    [
        {"human_reviewed": True},
        {"tables": {"guessed": {}}},
        {"invoice_input_ids": [{}]},
        {"tables": {"accounts": {"path": "/foreign.csv"}}},
    ],
)
def test_native_treasury_refuses_unsafe_draft_fields(initial_treasury, fields):
    env, output, _ = initial_treasury
    result = rpc_program(
        env,
        SETUP
        + "const result=call('vera_workspace_treasury_draft_save',{...authority,fields:"
        + json.dumps(fields)
        + "});",
    )
    assert result["isError"] is True
    assert list(output.iterdir()) == []


@pytest.mark.parametrize(
    "key,value,error",
    [
        ("expected_draft_revision", "0" * 64, "draft changed"),
        ("human_reviewed", False, "Invalid human_reviewed"),
        ("review_ticket", "invalid.signature", "Invalid review ticket"),
        ("revision", "0" * 64, "mismatched review ticket"),
    ],
)
def test_native_treasury_refuses_changed_or_unconfirmed_scope(
    initial_treasury, key, value, error
):
    env, output, _ = initial_treasury
    result = rpc_program(
        env,
        save_program()
        + f"args[{json.dumps(key)}]={json.dumps(value)};const result=call('vera_workspace_treasury_prepare',args);",
    )
    assert result["isError"] is True
    assert error in result["content"][0]["text"]
    assert list(output.iterdir()) == []


def test_native_treasury_viewer_cannot_write_a_draft(initial_treasury):
    env, output, _ = initial_treasury
    env["VERA_WORKSPACE_ROLES"] = "VIEWER"
    result = rpc_program(
        env,
        SETUP
        + "const refused=call('vera_workspace_treasury_draft_save',{...authority,fields:{company_name:'attempt'}});const result={setup,refused};",
    )
    assert result["setup"]["can_prepare"] is False
    assert result["refused"]["isError"] is True
    assert list(output.iterdir()) == []


def test_native_treasury_preparation_retains_first_version_through_date_review(
    initial_treasury,
):
    env, output, binding = initial_treasury
    result = rpc_program(
        env,
        save_program()
        + """
const prepared=payload(call('vera_workspace_treasury_prepare',args));
const view=payload(call('vera_workspace_view',{work_ref:setup.work_ref}));
const selected=view.items[0];
const exactView=payload(call('vera_workspace_view',{work_ref:setup.work_ref,item_id:selected.event_id}));
const date=payload(call('vera_workspace_save',{work_ref:setup.work_ref,revision:exactView.revision,review_ticket:exactView.review_ticket,item_id:selected.event_id,human_reviewed:true,idempotency_key:'fictional-date-review',decisions:{[selected.event_id]:{expected_date:'2026-10-01',basis:'Explicit fictional date review'}}}));
const reopened=payload(call('vera_workspace_view',{work_ref:setup.work_ref}));
const retry=payload(call('vera_workspace_treasury_prepare',args));
const result={prepared,date,reopened,retry};
""",
    )
    assert result["retry"] == result["prepared"]
    assert result["reopened"]["revision"] != result["prepared"]["record_sha256"]
    assert (
        output / "versions" / result["prepared"]["record_sha256"] / "forecast.json"
    ).is_file()
    assert workspace_module().load_binding(binding)["run"]["status"] == "running"


def test_native_treasury_wrong_headers_refuse_before_output_writes(initial_treasury):
    env, output, _ = initial_treasury
    result = rpc_program(
        env,
        save_program()
        + """
fields.tables.accounts=fields.tables.open_items;
const changed=payload(call('vera_workspace_treasury_draft_save',{...exact,fields}));
const refused=call('vera_workspace_treasury_prepare',{...args,expected_draft_revision:changed.draft_revision});
const result=refused;
""",
    )
    assert result["isError"] is True
    assert "headers" in result["content"][0]["text"]
    assert list(output.iterdir()) == []


def test_native_treasury_foreign_receipt_refuses_before_source_read(initial_treasury):
    env, output, _ = initial_treasury
    result = rpc_program(
        env,
        save_program()
        + """
fields.tables.accounts.input_id='foreign-receipt';
const changed=payload(call('vera_workspace_treasury_draft_save',{...exact,fields}));
const result=call('vera_workspace_treasury_prepare',{...args,expected_draft_revision:changed.draft_revision});
""",
    )
    assert result["isError"] is True
    assert "outside this run" in result["content"][0]["text"]
    assert list(output.iterdir()) == []


@pytest.mark.parametrize("target", ["native_treasury_manifest.json", "initial-version"])
def test_native_treasury_changed_retained_evidence_blocks_review(
    initial_treasury, target
):
    env, output, binding = initial_treasury
    prepared = rpc_program(
        env,
        save_program()
        + "const result=payload(call('vera_workspace_treasury_prepare',args));",
    )
    path = (
        output / "native_treasury_manifest.json"
        if target == "native_treasury_manifest.json"
        else output / "versions" / prepared["record_sha256"] / "undeclared.txt"
    )
    path.write_text("changed retained evidence")
    service = workspace_module()
    with pytest.raises(ValueError, match="changed"):
        service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    setup = service.dispatch(
        "vera_workspace_treasury_setup", {"work_ref": binding["work_ref"]}
    )
    assert setup["status"] == "recovery_required"
    assert setup["can_prepare"] is False


def test_native_treasury_interrupted_receipt_never_adopts_found_session(
    initial_treasury,
):
    env, output, binding = initial_treasury
    rpc_program(
        env,
        save_program()
        + "const result=payload(call('vera_workspace_treasury_prepare',args));",
    )
    service = workspace_module()
    intent = next(
        service.ui_state_directory(output, create=False).glob(
            "treasury-prepare-request-*.json"
        )
    )
    retained = json.loads(intent.read_bytes())
    retained.pop("result")
    intent.write_text(json.dumps(retained))
    with pytest.raises(ValueError, match="Interrupted"):
        service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    setup = service.dispatch(
        "vera_workspace_treasury_setup", {"work_ref": binding["work_ref"]}
    )
    assert setup["status"] == "recovery_required"
    assert setup["can_prepare"] is False


def test_native_treasury_saved_draft_cas_prevents_overwrite(initial_treasury):
    env, output, _ = initial_treasury
    result = rpc_program(
        env,
        save_program()
        + "const refused=call('vera_workspace_treasury_draft_save',{...authority,fields:{company_name:'replacement'}});const recovered=payload(call('vera_workspace_treasury_setup',{work_ref:setup.work_ref}));const result={refused,recovered,saved};",
    )
    assert result["refused"]["isError"] is True
    assert (
        result["recovered"]["draft"]["draft_revision"]
        == result["saved"]["draft_revision"]
    )
    assert list(output.iterdir()) == []


def test_native_treasury_successful_key_cannot_be_repurposed(initial_treasury):
    env, output, _ = initial_treasury
    result = rpc_program(
        env,
        save_program()
        + "const prepared=payload(call('vera_workspace_treasury_prepare',args));const refused=call('vera_workspace_treasury_prepare',{...args,expected_draft_revision:'0'.repeat(64)});const result={prepared,refused};",
    )
    assert result["refused"]["isError"] is True
    assert "different request" in result["refused"]["content"][0]["text"]
    assert len(list((output / "versions").iterdir())) == 1


@pytest.mark.parametrize("initial_treasury", ["xlsx"], indirect=True)
@pytest.mark.parametrize(
    "sheet,success", [("Balances exact", True), ("Missing sheet", False), ("", False)]
)
def test_native_treasury_xlsx_requires_the_declared_exact_sheet(
    initial_treasury, sheet, success
):
    env, output, _ = initial_treasury
    result = rpc_program(
        env,
        save_program()
        + f"fields.tables.accounts={{input_id:setup.items.find(row=>row.kind==='.xlsx').id,sheet:{json.dumps(sheet)}}};"
        + """
const changed=payload(call('vera_workspace_treasury_draft_save',{...exact,fields}));
const prepared=call('vera_workspace_treasury_prepare',{...args,expected_draft_revision:changed.draft_revision});
const result=prepared;
""",
    )
    assert bool(result.get("isError")) is not success
    if success:
        assert result["_meta"]["workspace"]["calculation_complete"] is True
        assert (output / "native_treasury_manifest.json").is_file()
    else:
        assert list(output.iterdir()) == []


@pytest.mark.parametrize("initial_treasury", ["previous"], indirect=True)
def test_native_treasury_accepts_exact_registered_predecessor_without_guessed_digest(
    initial_treasury,
):
    env, output, _ = initial_treasury
    result = rpc_program(
        env,
        save_program()
        + """
fields.as_of='2026-09-11';fields.previous_input_id=setup.items.find(row=>row.kind==='.json').id;
const changed=payload(call('vera_workspace_treasury_draft_save',{...exact,fields}));
const inspected=payload(call('vera_workspace_treasury_inspect',{work_ref:setup.work_ref,revision:setup.revision,expected_draft_revision:changed.draft_revision}));
const prepared=payload(call('vera_workspace_treasury_prepare',{...args,expected_draft_revision:changed.draft_revision}));
const result={inspected,prepared};
""",
    )
    assert result["inspected"]["previous"]["status"] == "accepted"
    manifest = json.loads((output / "native_treasury_manifest.json").read_bytes())
    assert (
        manifest["previous"]["record_sha256"]
        == result["inspected"]["previous"]["record_sha256"]
    )
    assert result["prepared"]["professional_approval"] is False


def test_native_treasury_private_source_rows_do_not_enter_model_facing_result(
    initial_treasury,
):
    env, output, _ = initial_treasury
    result = rpc_program(
        env,
        save_program()
        + "const result=call('vera_workspace_treasury_inspect',{work_ref:setup.work_ref,revision:setup.revision,expected_draft_revision:saved.draft_revision,table:'open_items'});",
    )
    assert result["_meta"]["workspace"]["rows"][0]["party_name"] == "Synthetic C101"
    assert "Synthetic C101" not in json.dumps(result["content"])
    assert "50000.00" not in json.dumps(result["content"])
    assert "Synthetic C101" not in json.dumps(result["structuredContent"])
    assert "50000.00" not in json.dumps(result["structuredContent"])
    assert list(output.iterdir()) == []


def test_native_treasury_new_draft_owner_does_not_adopt_other_actor_choices(
    initial_treasury,
):
    env, output, _ = initial_treasury
    rpc_program(env, save_program() + "const result=saved;")
    binding_path = Path(env["VERA_WORKSPACE_BINDINGS"])
    config = json.loads(binding_path.read_bytes())
    config["actor_id"] = "other-fictional-reviewer"
    binding_path.write_text(json.dumps(config))
    env["VERA_WORKSPACE_ACTOR_ID"] = config["actor_id"]
    result = rpc_program(
        env,
        "const result=payload(call('vera_workspace_treasury_setup',{work_ref:'initial-treasury'}));",
    )
    assert result["draft"]["fields"] == {}
    assert result["draft"]["draft_revision"] == ""
    assert list(output.iterdir()) == []


@pytest.mark.parametrize("initial_treasury", ["xml"], indirect=True)
def test_native_treasury_optional_invoice_is_evidence_not_outstanding_balance(
    initial_treasury,
):
    env, output, _ = initial_treasury
    result = rpc_program(
        env,
        save_program()
        + """
fields.invoice_input_ids=[setup.items.find(row=>row.kind==='.xml').id];
const changed=payload(call('vera_workspace_treasury_draft_save',{...exact,fields}));
const inspected=payload(call('vera_workspace_treasury_inspect',{work_ref:setup.work_ref,revision:setup.revision,expected_draft_revision:changed.draft_revision,table:'open_items'}));
const prepared=payload(call('vera_workspace_treasury_prepare',{...args,expected_draft_revision:changed.draft_revision}));
const result={inspected,prepared};
""",
    )
    assert result["inspected"]["invoice_evidence_count"] == 1
    assert result["inspected"]["rows"][0]["amount"] == "50000.00"
    retained = json.loads((output / "prepared_inputs.json").read_bytes())["inputs"]
    assert retained["invoice_evidence"][0]["payments"][0]["amount"] == "122.00"
    assert retained["open_items"][0]["amount"] == "50000.00"


@pytest.mark.parametrize("initial_treasury", ["many"], indirect=True)
def test_native_treasury_source_pages_count_and_reach_every_registered_row(
    initial_treasury,
):
    env, output, _ = initial_treasury
    result = rpc_program(
        env,
        save_program()
        + """
const pageScope={work_ref:setup.work_ref,revision:setup.revision,expected_draft_revision:saved.draft_revision,table:'accounts'};
const first=payload(call('vera_workspace_treasury_inspect',pageScope));
const next=payload(call('vera_workspace_treasury_inspect',{...pageScope,offset:20}));
const result={first,next};
""",
    )
    assert result["first"]["total"] == result["next"]["total"] == 35
    assert len(result["first"]["rows"]) == 20
    assert result["first"]["has_more"] is True
    assert len(result["next"]["rows"]) == 15
    assert result["next"]["has_more"] is False
    assert result["next"]["rows"][-1]["account_id"] == "bank-extra-34"
    assert list(output.iterdir()) == []


@pytest.mark.parametrize("initial_treasury", ["invalid-last-row"], indirect=True)
def test_native_treasury_invalid_row_after_preview_is_not_hidden_by_pagination(
    initial_treasury,
):
    env, output, _ = initial_treasury
    result = rpc_program(
        env,
        save_program()
        + "const inspected=payload(call('vera_workspace_treasury_inspect',{work_ref:setup.work_ref,revision:setup.revision,expected_draft_revision:saved.draft_revision,table:'accounts'}));const prepared=call('vera_workspace_treasury_prepare',args);const result={inspected,prepared};",
    )
    assert result["inspected"]["ok"] is False
    assert result["prepared"]["isError"] is True
    assert list(output.iterdir()) == []
