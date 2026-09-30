"""Operating-capital roll-forward, economic profit and independent FCFF checks."""

from __future__ import annotations

import hashlib
import shutil
from decimal import Decimal
from pathlib import Path

import pytest
from test_business_valuation import (
    FIXTURE,
    ValuationError,
    build_valuation,
    case_data,
    claimed_case,
    compile_html,
    normalized_case,
    planning_case,
    prepare_archive_run,
    read_json,
    restore_imports,
    review,
    run_valuation,
    write_package,
)


def economic_case(base: dict | None = None) -> dict:
    """Add a fictional operating model with an explicit second-period release."""
    case = case_data() if base is None else base
    case["sources"].append(
        {
            **case["sources"][0],
            "id": "economic-basis",
            "path": "economic-profit.txt",
            "sha256": hashlib.sha256(
                (FIXTURE / "economic-profit.txt").read_bytes()
            ).hexdigest(),
            "description": "Fictional operating capital, NOPAT and reinvestment",
        }
    )
    template = case["inputs"][0]
    case["inputs"].extend(
        [
            {
                **template,
                "id": key,
                "value": value,
                "unit": unit,
                "source_ids": ["economic-basis"],
                "description": f"Fictional {key}",
                "locator": "economic-profit.txt: declared amounts",
            }
            for key, value, unit in (
                ("capital0", "1000", "EUR"),
                ("capital1", "1100", "EUR"),
                ("capital2", "1050", "EUR"),
                ("nopat1", "180", "EUR"),
                ("nopat2", "160", "EUR"),
                ("reinvestment1", "100", "EUR"),
                ("reinvestment2", "-50", "EUR"),
                ("terminal-operating", "1400", "EUR"),
                ("economic-rate2", "0.2", "ratio"),
            )
        ]
    )
    case["mandate"].update(valuation_date="2024-12-31", information_cutoff="2024-12-31")
    case["methods"].append(
        {
            "id": "economic",
            "kind": "ECONOMIC_PROFIT",
            "selected": True,
            "rationale": "Synthetic economic-profit comparison with an explicit capital roll-forward.",
            "inputs": {
                "operating_capital": ["capital0", "capital1", "capital2"],
                "nopat": ["nopat1", "nopat2"],
                "net_reinvestment": ["reinvestment1", "reinvestment2"],
                "terminal_enterprise_value": "terminal-operating",
            },
            "limitations": [
                "Synthetic accounting, tax, WACC and terminal assumptions require professional review."
            ],
            "timing": {
                "valuation_date": "2024-12-31",
                "period_end_dates": ["2025-12-31", "2026-12-31"],
                "cash_flow_timing": "end_period",
                "day_count": "ACT/365F",
                "rate_compounding": "effective_annual",
                "rate_model": "flat",
                "rate_ids": ["rate"],
                "rationale": "NOPAT and net reinvestment occur at the specified period ends.",
            },
            "economic_basis": {
                "operating_perimeter": "Only fictional operating capital; no cash surplus or financing balances.",
                "accounting_adjustments": "One consistent assumed basis; no omitted revaluations or currency effects.",
                "nopat_tax_basis": "Explicit NOPAT after assumed unlevered operating tax; tax amounts are not estimated.",
                "reinvestment_basis": "Signed net investment of 100 then release of 50 captures every capital change.",
                "terminal_enterprise_basis": "Independent 1400 EUR operating value at the final date.",
                "capital_cost_basis": "Explicit operating capital cost; interval charges share discount factors.",
                "source_ids": ["economic-basis"],
                "locator": "economic-profit.txt: assumptions",
                "status": "confirmed",
            },
        }
    )
    return case


@pytest.mark.parametrize(
    ("timing", "expected"),
    [
        ({}, 1403.305785123967),
        (
            {
                "rate_model": "spot_curve",
                "rate_ids": ["rate", "economic-rate2"],
                "terminal_discount_rate": "economic-rate2",
            },
            1190.7828282828284,
        ),
        (
            {"rate_model": "forward_curve", "rate_ids": ["rate", "economic-rate2"]},
            1292.4242424242425,
        ),
        ({"rate_compounding": "continuous"}, 1390.5435058984274),
        ({"period_end_dates": ["2025-06-30", "2025-12-31"]}, 1539.943253238575),
    ],
)
def test_economic_profit_matches_independently_discounted_operating_flows(
    timing, expected
):
    case = economic_case()
    case["methods"][-1]["timing"].update(timing)
    result = build_valuation(case, FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert float(values["economic/value"]) == pytest.approx(expected, rel=1e-12)
    assert Decimal(
        values["economic/operating_cashflow_difference"]
    ).copy_abs() < Decimal("1e-28")
    assert result["methods"][-1]["value_type"] == "operating_enterprise"
    assert result["methods"][-1]["equity_id"] is None


def test_capital_charge_uses_opening_capital_and_deducts_terminal_capital_once():
    result = build_valuation(economic_case(), FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert Decimal(values["economic/economic/1/capital_charge"]) == 100
    assert Decimal(values["economic/economic/2/capital_charge"]) == 110
    assert Decimal(values["economic/economic/2/fcff"]) == 210
    assert Decimal(values["economic/continuing_economic_profit"]) == 350


def test_optional_complete_bridge_deducts_financing_once():
    case = economic_case()
    case["methods"][-1]["bridge"] = case["methods"][0]["bridge"]
    result = build_valuation(case, FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert result["methods"][-1]["equity_id"] == "economic/equity"
    assert Decimal(values["economic/value"]) - Decimal(values["economic/equity"]) == 250


@pytest.mark.parametrize(
    ("key", "value"),
    [
        ("capital0", "1001"),
        ("capital1", "1101"),
        ("capital2", "1051"),
        ("reinvestment1", "0"),
        ("reinvestment2", "50"),
    ],
)
def test_unreconciled_capital_blocks_only_affected_method(key, value):
    case = economic_case()
    next(row for row in case["inputs"] if row["id"] == key)["value"] = value
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"
    assert "does not reconcile" in result["methods"][-1]["reason"]
    assert result["methods"][0]["status"] == "ready_for_professional_review"


@pytest.mark.parametrize(
    ("key", "value"),
    [
        ("operating_capital", ["capital0", "capital1"]),
        ("nopat", ["nopat1"]),
        ("net_reinvestment", []),
        ("net_reinvestment", ["reinvestment1"]),
        ("terminal_enterprise_value", "rate"),
        ("terminal_enterprise_value", "nopat2"),
        ("terminal_enterprise_value", "reinvestment2"),
    ],
)
def test_incomplete_schedule_or_reused_flow_terminal_blocks(key, value):
    case = economic_case()
    case["methods"][-1]["inputs"][key] = value
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"


@pytest.mark.parametrize(
    "key", ["capital1", "nopat1", "reinvestment1", "terminal-operating"]
)
def test_unknown_economic_amount_is_not_filled_or_assumed_zero(key):
    case = economic_case()
    next(row for row in case["inputs"] if row["id"] == key)["value"] = None
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"


@pytest.mark.parametrize("field", ["timing", "economic_basis"])
def test_explicit_timing_and_professional_basis_are_required(field):
    case = economic_case()
    del case["methods"][-1][field]
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"


@pytest.mark.parametrize(
    "timing",
    [
        {"cash_flow_timing": "mid_period"},
        {
            "rate_model": "spot_curve",
            "rate_ids": ["rate", "economic-rate2"],
            "terminal_discount_rate": "rate",
        },
    ],
)
def test_inconsistent_capital_and_discount_timing_blocks(timing):
    case = economic_case()
    case["methods"][-1]["timing"].update(timing)
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"


@pytest.mark.parametrize(
    "field",
    [
        "operating_perimeter",
        "accounting_adjustments",
        "nopat_tax_basis",
        "reinvestment_basis",
        "terminal_enterprise_basis",
        "capital_cost_basis",
        "locator",
    ],
)
def test_changed_basis_expires_affected_review_and_preserves_unrelated_method(field):
    case = economic_case()
    before = build_valuation(case, FIXTURE)
    case["methods"][-1]["review"] = review(before["methods"][-1]["dependency_sha256"])
    case["methods"][0]["review"] = review(before["methods"][0]["dependency_sha256"])
    case["methods"][-1]["economic_basis"][field] += " Revised."
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["stale_review"] is True
    assert result["methods"][0]["status"] == "accepted_workpaper"


@pytest.mark.parametrize(
    ("field", "value", "status"),
    [
        ("status", "proposed", "partial"),
        ("source_ids", ["missing"], "blocked"),
        ("nopat_tax_basis", "", "blocked"),
    ],
)
def test_missing_or_unconfirmed_basis_does_not_approve_economic_method(
    field, value, status
):
    case = economic_case()
    case["methods"][-1]["economic_basis"][field] = value
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == status


def test_unreviewed_basis_source_keeps_method_partial():
    case = economic_case()
    case["sources"][-1]["status"] = "unverified"
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "partial"
    assert "economic-basis" in result["methods"][-1]["source_ids"]


def test_signed_operating_capital_and_losses_are_not_floored():
    case = economic_case()
    changed = {
        "capital0": "-1000",
        "capital1": "-900",
        "capital2": "-950",
        "nopat1": "-180",
        "nopat2": "-160",
        "terminal-operating": "-1400",
    }
    for row in case["inputs"]:
        if row["id"] in changed:
            row["value"] = changed[row["id"]]
    result = build_valuation(case, FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert result["methods"][-1]["status"] == "ready_for_professional_review"
    assert float(values["economic/value"]) == pytest.approx(
        -1502.4793388429753, rel=1e-12
    )


def test_explicit_zero_cost_retains_signed_reinvestment_cashflows():
    case = economic_case()
    next(row for row in case["inputs"] if row["id"] == "rate")["value"] = "0"
    result = build_valuation(case, FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert Decimal(values["economic/value"]) == 1690
    assert Decimal(values["economic/operating_cashflow_difference"]) == 0


def test_explicit_terminal_equal_to_capital_has_zero_continuing_profit():
    case = economic_case()
    case["methods"][-1]["inputs"]["terminal_enterprise_value"] = "capital2"
    result = build_valuation(case, FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert Decimal(values["economic/continuing_economic_profit"]) == 0
    assert float(values["economic/value"]) == pytest.approx(
        1114.0495867768595, rel=1e-12
    )


def test_plan_fcff_cannot_be_relabelled_as_nopat(tmp_path):
    case = economic_case(planning_case(tmp_path))
    shutil.copyfile(FIXTURE / "economic-profit.txt", tmp_path / "economic-profit.txt")
    case["mandate"].update(valuation_date="2026-12-31", information_cutoff="2026-12-31")
    case["methods"][-1]["timing"].update(
        valuation_date="2026-12-31", period_end_dates=["2027-12-31", "2028-12-31"]
    )
    case["methods"][-1]["inputs"]["nopat"] = ["flow", "flow"]
    result = build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")
    assert result["methods"][-1]["status"] == "blocked"
    assert "cannot be relabelled" in result["methods"][-1]["reason"]


def test_economic_export_links_capital_checks_and_escapes_basis(tmp_path):
    from openpyxl import load_workbook

    case = economic_case()
    case["methods"][-1]["economic_basis"][
        "operating_perimeter"
    ] = '=HYPERLINK("https://example.invalid")'
    case["methods"][-1]["economic_basis"][
        "accounting_adjustments"
    ] = '<script>alert("source")</script>'
    result = build_valuation(case, FIXTURE)
    output = tmp_path / "report"
    write_package(result, FIXTURE, output)
    workbook = load_workbook(output / "valuation_workbook.xlsx", data_only=False)
    assert workbook["Base profitto economico"]["C2"].data_type == "s"
    assert workbook["Capitale operativo"]["C2"].data_type == "f"
    assert workbook["Capitale operativo"].max_row == 3
    html = compile_html(result)
    assert "<script>alert" not in html
    assert "&lt;script&gt;alert" in html
    assert "perpetuità annuale" not in html.split("Profitto economico operativo", 1)[1]
    assert read_json(output / "valuation.json") == result


def test_economic_method_runs_with_exact_archive_receipts(tmp_path, record_property):
    case_path, context = prepare_archive_run(
        tmp_path,
        supplied_case=economic_case(),
        source_files=[FIXTURE / "evidence.txt", FIXTURE / "economic-profit.txt"],
    )
    output = run_valuation.run_case(case_path, context)
    report = read_json(Path(output["output_dir"]) / "valuation.json")
    assert report["methods"][-1]["status"] == "ready_for_professional_review"
    assert report["methods"][-1]["economic_basis"]["status"] == "confirmed"
    record_property("economic_profit_output", output["output_dir"])


def test_economic_basis_is_rejected_on_an_unrelated_method():
    case = economic_case()
    case["methods"][0]["economic_basis"] = case["methods"][-1]["economic_basis"]
    result = build_valuation(case, FIXTURE)
    assert result["methods"][0]["status"] == "blocked"
    assert result["methods"][-1]["status"] == "ready_for_professional_review"


def test_changed_tax_basis_expires_claim_even_with_identical_cashflow_amount():
    case = economic_case(claimed_case())
    claim = case["claims"][0]
    claim.update(
        calculation_ids=["economic/economic/1/fcff"],
        text="Il primo FCFF sintetico è 80 EUR alle ipotesi dichiarate.",
        values=[
            {"calculation_id": "economic/economic/1/fcff", "value": "80", "unit": "EUR"}
        ],
    )
    first = build_valuation(case, FIXTURE)
    case["methods"][-1]["review"] = review(first["methods"][-1]["dependency_sha256"])
    case["mandate_details"]["review"] = review(
        first["mandate_assessment"]["dependency_sha256"]
    )
    second = build_valuation(case, FIXTURE)
    claim["review"] = review(second["claims"][0]["dependency_sha256"])
    case["methods"][-1]["economic_basis"][
        "nopat_tax_basis"
    ] += " Revised tax interpretation."
    result = build_valuation(case, FIXTURE)
    assert result["claims"][0]["stale_review"] is True
    assert result["claims"][0]["status"] != "accepted_workpaper"


def test_normalized_nopat_retains_adjustment_dependencies():
    case = economic_case(normalized_case())
    case["methods"][-1]["inputs"]["nopat"] = ["income", "nopat2"]
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "partial"
    assert result["methods"][-1]["normalization_ids"] == ["income-2026"]
    assert {"reported-income", "signed-addition", "signed-deduction"} <= set(
        result["methods"][-1]["input_ids"]
    )


def test_normalization_cannot_hide_relabelled_plan_fcff(tmp_path):
    case = economic_case(planning_case(tmp_path))
    shutil.copyfile(FIXTURE / "economic-profit.txt", tmp_path / "economic-profit.txt")
    case["mandate"].update(valuation_date="2026-12-31", information_cutoff="2026-12-31")
    case["methods"][-1]["timing"].update(
        valuation_date="2026-12-31", period_end_dates=["2027-12-31", "2028-12-31"]
    )
    template = case["inputs"][0]
    case["inputs"].append(
        {**template, "id": "invented-nopat-adjustment", "value": "1380"}
    )
    normalization = normalized_case()["normalizations"][0]
    normalization.update(
        id="plan-to-nopat", reported_input="flow", adjusted_input="nopat1"
    )
    normalization["adjustments"] = [
        {**normalization["adjustments"][0], "amount_input": "invented-nopat-adjustment"}
    ]
    case["normalizations"] = [normalization]
    with pytest.raises(ValuationError, match="Revise and replay the plan"):
        build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")
