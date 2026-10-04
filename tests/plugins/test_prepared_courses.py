"""Course isolation, stale-source handling, real corpus rendering and privacy."""

from __future__ import annotations

import hashlib
import importlib.util
import io
import json
import shutil
import sys
from pathlib import Path

import pymupdf
import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "plugins/_shared/vendor/modules"))
from courseware.library import CourseError, CourseLibrary, main

spec = importlib.util.spec_from_file_location(
    "course_build", ROOT / "scripts/course_materials/build_catalog.py"
)
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)
COURSES = [
    (product, workflow)
    for product in ("vera", "clara", "lucia")
    for workflow in sorted(builder._eligible(product))
]
RELEASE = json.loads((ROOT / "scripts/course_materials/release_plan.json").read_text())
PREPARED = [(p, w) for p, w in COURSES if f"{p}/{w}" in RELEASE["prepared"]]
LOCALIZED = [
    (product, workflow, language)
    for product, workflow in PREPARED
    for language in builder.LIMITED_LANGUAGES.get(
        f"{product}/{workflow}", builder.LANGUAGES
    )
]


@pytest.fixture(autouse=True)
def course_runtime_imports(monkeypatch):
    """Restore lazy course imports after the repository's isolation hook."""
    monkeypatch.syspath_prepend(str(ROOT / "plugins/_shared/vendor/modules"))


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture
def isolated(tmp_path):
    """An installed product with one authored course and a local method source."""
    root = tmp_path / "vera"
    original = ROOT / "plugins/vera/assets/courses/fatture-xml-check/course.json"
    course = json.loads(original.read_text(encoding="utf-8"))
    skill = root / "skills/variance-analysis/SKILL.md"
    skill.parent.mkdir(parents=True)
    skill.write_text("# Own installed variance method\n", encoding="utf-8")
    course["sources"] = [
        {"path": "skills/variance-analysis/SKILL.md", "sha256": digest(skill)}
    ]
    course["workflow"] = "variance-analysis"
    manifest = root / "assets/courses/variance-analysis/course.json"
    for asset in course["files"]:
        target = manifest.parent / asset["path"]
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(original.parent / asset["path"], target)
    write_json(manifest, course)
    write_json(root / ".codex-plugin/plugin.json", {"name": "vera"})
    index = {
        "product": "vera",
        "courses": {
            "variance-analysis": {
                "path": "variance-analysis/course.json",
                "sha256": digest(manifest),
                "languages": list(course["locales"]),
            }
        },
    }
    write_json(root / "assets/courses/index.json", index)
    return root, manifest, skill


def alter_course(isolated, change):
    root, path, _ = isolated
    course = json.loads(path.read_text(encoding="utf-8"))
    change(course)
    write_json(path, course)
    index_path = root / "assets/courses/index.json"
    index = json.loads(index_path.read_text(encoding="utf-8"))
    index["courses"]["variance-analysis"]["sha256"] = digest(path)
    write_json(index_path, index)


@pytest.mark.parametrize("product,workflow", COURSES)
def test_every_installed_workflow_has_a_source_current_course(product, workflow):
    library = CourseLibrary(ROOT / "plugins" / product, builder._eligible(product))

    course = library.load(workflow, "it")

    assert course["product"] == product
    assert course["workflow"] == workflow
    assert sum(course["seconds"]) == 390
    assert course["sources"]
    expected_schema = (
        "mparanza.teaching_kit.v2"
        if (product, workflow) in PREPARED
        else "mparanza.course.v1"
    )
    assert course["schema"] == expected_schema


@pytest.mark.parametrize("product,workflow", PREPARED)
def test_each_kit_explains_first_use_and_supplies_actual_inputs(product, workflow):
    course = CourseLibrary(ROOT / "plugins" / product, builder._eligible(product)).load(
        workflow, "it"
    )
    content = course["locales"]["it"]
    assert content["request"]
    assert content["steps"]
    assert content["deliverables"]
    assert content["review"]
    assert content["practice"] != content["request"]
    assert content["success"]
    assert content["repeat"]
    assert not {"output_rows", "conclusion", "answer", "result_title"} & set(content)


@pytest.mark.parametrize("product", ["vera", "clara", "lucia"])
def test_catalog_exactly_covers_current_product_and_no_retired_ids(product):
    expected = builder._eligible(product)
    catalog = CourseLibrary(ROOT / "plugins" / product, expected).catalog()

    assert {row["workflow"] for row in catalog} == expected
    assert (
        not {"report-builder", "check-entries", "bilancio-xbrl-it", "interview"}
        & expected
    )


@pytest.mark.parametrize(
    "workflow", ["brand-fit", "hosted-interview", "research-video"]
)
def test_hosted_clara_workflows_remain_installed_but_have_no_local_lesson(
    workflow, monkeypatch
):
    monkeypatch.syspath_prepend(str(ROOT / "plugins/_shared/vendor/modules"))
    from desktop_teaching.onboarding import (
        OnboardingError,
        eligible_workflows,
        teaching_contract,
    )

    root = ROOT / "plugins/clara"
    assert (root / "skills" / workflow / "SKILL.md").is_file()
    assert workflow not in builder._eligible("clara")
    assert workflow not in eligible_workflows(root)
    library = CourseLibrary(root, {workflow})
    assert library.catalog() == []
    with pytest.raises(CourseError, match="requires hosted services"):
        library.load(workflow, "it")
    with pytest.raises(OnboardingError, match="requires hosted services"):
        teaching_contract(root, workflow)


@pytest.mark.parametrize(
    "foreign", ["reporting-engine", "hosted-interview", "../clara/reporting-engine"]
)
def test_foreign_workflow_never_resolves_from_course_catalog(isolated, foreign):
    root, _, _ = isolated
    with pytest.raises(CourseError, match="this product"):
        CourseLibrary(root, {"variance-analysis"}).load(foreign, "it")


def test_removed_workflow_cannot_reuse_its_old_course(isolated):
    root, _, _ = isolated
    library = CourseLibrary(root, set())
    assert library.catalog() == []
    with pytest.raises(CourseError, match="this product"):
        library.load("variance-analysis", "it")


def test_unsupported_language_is_not_silently_translated(isolated):
    root, _, _ = isolated
    with pytest.raises(CourseError, match="select one"):
        CourseLibrary(root, {"variance-analysis"}).load("variance-analysis", "ja")


def test_source_change_requires_editorial_refresh(isolated):
    root, _, skill = isolated
    skill.write_text("# Changed method\n", encoding="utf-8")
    with pytest.raises(CourseError, match="editorial refresh"):
        CourseLibrary(root, {"variance-analysis"}).load("variance-analysis", "it")


@pytest.mark.parametrize(
    "product,workflow,packaged,repository",
    [
        (
            "clara",
            "attribute-reporting",
            "modules/attribute-reporting/vendor/modules/pdp/attribute_table_templates.py",
            "modules/pdp/attribute_table_templates.py",
        ),
        (
            "vera",
            "bilancio-oic",
            "modules/bilancio-xbrl-it/rulepacks/it/disclosures-2026.1.json",
            "plugins/bilancio-xbrl-it/rulepacks/it/disclosures-2026.1.json",
        ),
        (
            "vera",
            "browser-automation",
            "modules/browser-automation/scripts/discovery_runtime.mjs",
            "plugins/browser-automation/scripts/discovery_runtime.mjs",
        ),
    ],
)
def test_kit_fingerprints_include_actual_bundled_engine_dependencies(
    product, workflow, packaged, repository
):
    records = builder._source_records(product, workflow)
    matches = [record for record in records if record["path"] == packaged]
    assert matches == [
        {
            "path": packaged,
            "repository_path": repository,
            "sha256": digest(ROOT / repository),
        }
    ]


def test_missing_source_cannot_fall_back_to_another_installation(isolated):
    root, _, skill = isolated
    skill.unlink()
    with pytest.raises(CourseError, match="Missing or foreign"):
        CourseLibrary(root, {"variance-analysis"}).load("variance-analysis", "it")


def test_foreign_symlink_source_is_rejected(isolated, tmp_path):
    root, _, skill = isolated
    foreign = tmp_path / "foreign.md"
    foreign.write_bytes(skill.read_bytes())
    skill.unlink()
    skill.symlink_to(foreign)
    with pytest.raises(CourseError, match="Missing or foreign"):
        CourseLibrary(root, {"variance-analysis"}).load("variance-analysis", "it")


def test_altered_material_is_detected_before_rendering(isolated, tmp_path):
    root, manifest, _ = isolated
    manifest.write_text(manifest.read_text(encoding="utf-8") + " ", encoding="utf-8")
    destination = tmp_path / "untouched"
    with pytest.raises(CourseError, match="content changed"):
        CourseLibrary(root, {"variance-analysis"}).render(
            "variance-analysis", "it", destination
        )
    assert not destination.exists()


@pytest.mark.parametrize(
    "field,value",
    [
        ("product", "clara"),
        ("workflow", "reporting-engine"),
        ("seconds", [20] * 6),
        ("schema", "unknown"),
    ],
)
def test_invalid_course_contract_is_rejected(isolated, field, value):
    root, _, _ = isolated
    alter_course(isolated, lambda course: course.update({field: value}))
    with pytest.raises(CourseError, match="Invalid course"):
        CourseLibrary(root, {"variance-analysis"}).load("variance-analysis", "it")


@pytest.mark.parametrize("path", ["../foreign.md", "/tmp/foreign.md", ""])
def test_course_file_paths_cannot_escape_their_root(isolated, path):
    root, _, _ = isolated
    alter_course(
        isolated,
        lambda course: course.update(files=[{"path": path, "sha256": "0" * 64}]),
    )
    # Windows resolves a POSIX root path against the current drive; both hosts
    # must reject it, whether at syntax or resolved-containment checks.
    with pytest.raises(
        CourseError, match="relative and contained|Missing or foreign course file"
    ):
        CourseLibrary(root, {"variance-analysis"}).load("variance-analysis", "it")


def test_changed_attached_example_is_rejected(isolated):
    root, manifest, _ = isolated
    attachment = manifest.parent / "fixture.csv"
    attachment.write_text("amount\n10\n", encoding="utf-8")
    alter_course(
        isolated,
        lambda course: course.update(
            files=[{"path": "fixture.csv", "sha256": digest(attachment)}]
        ),
    )
    attachment.write_text("amount\n11\n", encoding="utf-8")
    with pytest.raises(CourseError, match="Teaching input changed"):
        CourseLibrary(root, {"variance-analysis"}).load("variance-analysis", "it")


@pytest.mark.parametrize(
    "product,workflow",
    [
        ("vera", "fatture-xml-check"),
        ("clara", "html-deck"),
        ("lucia", "apertura-pratica"),
    ],
)
def test_render_supplies_inputs_and_worker_request_without_manufacturing_a_result(
    tmp_path, product, workflow
):
    library = CourseLibrary(ROOT / "plugins" / product, builder._eligible(product))
    output = tmp_path / product

    receipt = library.render(workflow, "it", output)

    assert receipt["prepared_material_only"] is True
    assert receipt["execution_receipt"] is False
    assert receipt["understanding_confirmed"] is False
    assert receipt["source_files"]
    assert receipt["practice_files"]
    request = json.loads(
        (output / "execution-request.json").read_text(encoding="utf-8")
    )
    assert request["product"] == product
    assert request["workflow"] == workflow
    assert request["execution_receipt"] is False
    assert request["source_files"] == receipt["source_files"]
    assert not (output / "example.html").exists()
    page = (output / "course.html").read_text(encoding="utf-8")
    assert "Due chat, una lezione" in page
    assert "execution-request.json" not in page
    assert "<html lang='it'>" in page
    assert "<script" not in page
    assert "https://" not in page
    assert not list(output.rglob("profile.json"))
    assert not list(output.rglob("session.json"))


def test_render_preserves_existing_files(isolated, tmp_path):
    root, _, _ = isolated
    output = tmp_path / "lesson"
    output.mkdir()
    sentinel = output / "notes.md"
    sentinel.write_text("My notes", encoding="utf-8")
    with pytest.raises(CourseError, match="preserve existing"):
        CourseLibrary(root, {"variance-analysis"}).render(
            "variance-analysis", "it", output
        )
    assert sentinel.read_text(encoding="utf-8") == "My notes"


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_xml_input_opens_reading_view_but_execution_keeps_original(
    isolated, tmp_path, language
):
    from html.parser import HTMLParser

    class Links(HTMLParser):
        def __init__(self):
            super().__init__()
            self.hrefs = []

        def handle_starttag(self, tag, attrs):
            if tag == "a":
                self.hrefs.append(dict(attrs).get("href"))

    root, _, _ = isolated
    output = tmp_path / "reading"
    receipt = CourseLibrary(root, {"variance-analysis"}).render(
        "variance-analysis", language, output
    )
    links = Links()
    links.feed((output / "course.html").read_text())
    first_view = output / next(
        href for href in links.hrefs if href.startswith("input-")
    )
    view = first_view.read_text()
    original = Path(receipt["source_files"][0])
    assert original.suffix == ".xml"
    assert (
        original.read_bytes()
        == (
            root / "assets/courses/variance-analysis/files/input/invoice-01.xml"
        ).read_bytes()
    )
    assert "&lt;p:FatturaElettronica" in view
    assert "DEMO-001" in view
    assert "<table>" in view
    assert "<details><summary>" in view
    assert "1220.00" in view
    assert "course.html#step-1" in view
    assert f"<html lang='{language}'>" in view
    assert "<script" not in view
    assert receipt["execution_receipt"] is False


def test_invoice_source_reader_preserves_bodies_and_refuses_entity_expansion():
    from courseware.library import _xml_input

    ui = json.loads(
        (
            ROOT / "plugins/_shared/vendor/modules/courseware/assets/languages.json"
        ).read_text()
    )["it"]
    xml = """<FatturaElettronica><FatturaElettronicaBody><DatiGenerali>
    <DatiGeneraliDocumento><Numero>A&lt;script&gt;</Numero><ImportoTotaleDocumento>001.20</ImportoTotaleDocumento>
    </DatiGeneraliDocumento></DatiGenerali></FatturaElettronicaBody>
    <FatturaElettronicaBody><DatiGenerali><DatiGeneraliDocumento><Numero>B</Numero>
    </DatiGeneraliDocumento></DatiGenerali></FatturaElettronicaBody></FatturaElettronica>"""
    page = _xml_input(xml, ui)
    assert page.count("<table>") == 2
    assert "<td>A&lt;script&gt;</td>" in page and "<td>001.20</td>" in page
    assert "<script>" not in page
    for unsupported in (
        "<broken>",
        '<!DOCTYPE x [<!ENTITY e "EXPANDED">]><FatturaElettronica>&e;</FatturaElettronica>',
        "<Other><Numero>A</Numero></Other>",
    ):
        assert "<table>" not in _xml_input(unsupported, ui)


def test_render_refuses_symlink_destination(isolated, tmp_path):
    root, _, _ = isolated
    target = tmp_path / "real"
    target.mkdir()
    link = tmp_path / "linked"
    link.symlink_to(target, target_is_directory=True)
    with pytest.raises(CourseError, match="not a symlink"):
        CourseLibrary(root, {"variance-analysis"}).render(
            "variance-analysis", "it", link / "course"
        )


@pytest.fixture
def result_view_case(isolated, tmp_path):
    root, _, skill = isolated
    lesson = tmp_path / "live-lesson"
    lesson.mkdir()
    (lesson / "source.txt").write_text("Explicit synthetic source for unit test")
    (lesson / "run.json").write_text('{"synthetic_test_fixture": true}')
    (lesson / "fatture_summary.csv").write_text(
        'invoice_number,total_amount,detail\nDEMO-004,976.00,"<script>bad</script>"\n'
    )
    (lesson / "summary.md").write_text("# Findings\n\n- <script>literal</script>\n")

    def record(name):
        return {"path": name, "sha256": digest(lesson / name)}

    receipt = {
        "schema": "mparanza.teaching_execution.v1",
        "evidence_kind": "host_attested_local_execution",
        "product": "vera",
        "workflow_id": "variance-analysis",
        "phase": "practice",
        "worker_thread_id": "unit-test-worker",
        "outcome": "completed",
        "skill_sha256": digest(skill),
        "inputs": [record("source.txt")],
        "native_records": [record("run.json")],
        "outputs": [record("fatture_summary.csv"), record("summary.md")],
    }
    write_json(lesson / "execution.json", receipt)
    return CourseLibrary(root, {"variance-analysis"}), lesson


@pytest.mark.parametrize(
    "language,pending",
    [
        ("it", "Da rivedere"),
        ("en", "Pending review"),
        ("fr", "À revoir"),
        ("de", "Zu prüfen"),
        ("es", "Por revisar"),
    ],
)
def test_inps_evidence_view_keeps_review_and_source_with_complete_download(
    result_view_case, tmp_path, language, pending
):
    _, lesson = result_view_case
    library = CourseLibrary(ROOT / "plugins/vera", {"previdenza-inps"})
    original = lesson / "evidence_matrix.csv"
    payload = (
        "fact_id,statement,review_status,document_id,locator_kind,locator_value,quote,visual_confirmation_by_id\n"
        'F-1,Document observation,pending,DOC-2,document,1,"Original <text>",test-only-reviewer\n'
    ).encode()
    original.write_bytes(payload)
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt.update(
        workflow_id="previdenza-inps",
        skill_sha256=digest(ROOT / "plugins/vera/skills/previdenza-inps/SKILL.md"),
        outputs=[{"path": original.name, "sha256": digest(original)}],
    )
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "result-view"

    library.results("previdenza-inps", language, lesson, "execution.json", destination)

    view = (destination / "course.html").read_text()
    assert f"<td>{pending}</td>" in view
    assert "<td>DOC-2</td>" in view
    assert "Original &lt;text&gt;" in view
    assert "<th scope='col'>visual_confirmation_by_id</th>" not in view
    assert "test-only-reviewer" in view  # retained in the complete-file disclosure
    assert (destination / "outputs/01-evidence_matrix.csv").read_bytes() == payload
    assert original.read_bytes() == payload


def test_result_view_uses_verified_output_and_keeps_all_fields(
    result_view_case, tmp_path
):
    library, lesson = result_view_case
    before = (lesson / "fatture_summary.csv").read_bytes()
    destination = tmp_path / "result-view"
    result = library.results(
        "variance-analysis",
        "it",
        lesson,
        "execution.json",
        destination,
    )
    page = Path(result["course"]).read_text()
    assert "<th scope='col'>Numero</th>" in page
    assert "<td>976.00</td>" in page
    assert "&lt;script&gt;bad&lt;/script&gt;" in page
    assert "<script>" not in page
    assert "<h2>Findings</h2>" in page
    assert "<h3>Findings</h3>" not in page
    assert "summary.md" in page
    assert "<li>&lt;script&gt;literal&lt;/script&gt;</li>" in page
    assert "<details>" in page
    assert result["presentation_only"] is True
    assert (lesson / "fatture_summary.csv").read_bytes() == before
    assert not (lesson / "session.json").exists()


@pytest.mark.parametrize("extension", ["xlsx", "docx", "pdf"])
@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_result_view_download_preserves_binary_output(
    result_view_case, tmp_path, extension, language
):
    library, lesson = result_view_case
    original = lesson / ("reviewed-result." + extension)
    payload = b"\x00\xffBinary result fixture; must not be decoded or rewritten"
    if extension == "xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Summary"
        sheet.append(["Invoice", "Amount", "Calculated"])
        sheet.append(["<script>literal</script>", 488, "=1+1"])
        hidden = workbook.create_sheet("Hidden evidence")
        hidden.sheet_state = "hidden"
        hidden.append(["Not a visible sheet"])
        workbook.save(original)
    elif extension == "docx":
        from docx import Document

        document = Document()
        document.add_heading("Actual result", level=1)
        document.add_paragraph("<script>literal</script>")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Invoice"
        table.cell(0, 1).text = "Amount"
        table.cell(1, 0).text = "3-FF"
        table.cell(1, 1).text = "488"
        document.add_paragraph("After the table")
        document.save(original)
    else:
        original.write_bytes(payload)
    payload = original.read_bytes()
    receipt_path = lesson / "execution.json"
    receipt = json.loads(receipt_path.read_text())
    receipt["outputs"].append({"path": original.name, "sha256": digest(original)})
    write_json(receipt_path, receipt)
    destination = tmp_path / "result-view"
    result = library.results(
        "variance-analysis", language, lesson, "execution.json", destination
    )
    page = Path(result["course"]).read_text()
    if extension in {"xlsx", "docx"}:
        assert "href='result-03.html'" in page
        page = (destination / "result-03.html").read_text()
        assert "href='course.html'" in page
        assert "&lt;script&gt;literal&lt;/script&gt;" in page
        assert "<script>" not in page
        assert "488" in page
        if extension == "xlsx":
            assert "Hidden evidence" not in page
            assert "Not a visible sheet" not in page
            assert "=1+1" not in page
            assert "<summary>Summary</summary>" in page
        else:
            assert (
                page.index("Actual result")
                < page.index("<table>")
                < page.index("After the table")
            )
    assert f"href='outputs/03-{original.name}'" in page
    assert f"download='{original.name}'" in page
    assert (destination / "outputs" / ("03-" + original.name)).read_bytes() == payload
    assert original.read_bytes() == payload
    assert "Binary result fixture" not in page
    assert (destination / "outputs/01-fatture_summary.csv").read_bytes() == (
        lesson / "fatture_summary.csv"
    ).read_bytes()


def test_result_view_rejects_active_document(result_view_case, tmp_path):
    library, lesson = result_view_case
    original = lesson / "active.html"
    original.write_text("<script>alert(1)</script>")
    receipt_path = lesson / "execution.json"
    receipt = json.loads(receipt_path.read_text())
    receipt["outputs"] = [{"path": original.name, "sha256": digest(original)}]
    write_json(receipt_path, receipt)
    destination = tmp_path / "result-view"
    with pytest.raises(ValueError, match="Unsupported result format"):
        library.results(
            "variance-analysis", "it", lesson, "execution.json", destination
        )
    assert not destination.exists()


def test_attribute_report_reading_keeps_evidence_without_active_controls(
    result_view_case, tmp_path
):
    _, lesson = result_view_case
    library = CourseLibrary(ROOT / "plugins/clara", {"attribute-reporting"})
    original = lesson / "report.html"
    payload = b"""<!doctype html><html><head><meta charset="utf-8"></head><body>
    <main data-report-id="native"><aside data-correctness-verdict="correct_with_caveats">Correct with caveats</aside>
    <nav><a href="#comparison">Compare</a></nav><section id="comparison" data-section-id="winning_now">
    <figure data-table-sha256="audit"><figcaption>Evidence</figcaption><table><tr><td>75.0%</td></tr></table></figure>
    <a href="https://example.invalid/product" target="_blank" rel="noreferrer">Fictional product</a></section>
    <button type="button" onclick="window.print()">Print report</button></main></body></html>"""
    original.write_bytes(payload)
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt.update(
        product="clara",
        workflow_id="attribute-reporting",
        skill_sha256=digest(ROOT / "plugins/clara/skills/attribute-reporting/SKILL.md"),
        outputs=[{"path": original.name, "sha256": digest(original)}],
    )
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "attribute-result"
    library.results("attribute-reporting", "it", lesson, "execution.json", destination)
    view = (destination / "result-01.html").read_text()
    assert "Correct with caveats" in view and "<td>75.0%</td>" in view
    assert "<figcaption>Evidence</figcaption>" in view
    assert "href='#html-result-comparison'" in view
    assert "<span>Fictional product</span>" in view
    assert "onclick" not in view and "<button" not in view
    assert "https://example.invalid" not in view and "data-report-id" not in view
    assert (destination / "outputs/01-report.html").read_bytes() == payload
    assert original.read_bytes() == payload


@pytest.mark.parametrize(
    "active",
    [
        '<p onclick="x()">Click</p>',
        '<a href="javascript:x()">Click</a>',
        '<iframe src="https://example.com"></iframe>',
    ],
)
def test_attribute_report_reader_still_rejects_active_content(active):
    from courseware.html_view import passive_html_body

    with pytest.raises(ValueError, match="Unsupported result format"):
        passive_html_body(active, omit_controls=True, report_metadata=True)


@pytest.mark.parametrize("product", ["vera", "clara"])
def test_business_plan_reading_omits_controls_and_keeps_report_and_original(
    result_view_case, tmp_path, product
):
    _, lesson = result_view_case
    library = CourseLibrary(ROOT / f"plugins/{product}", {"business-planning"})
    original = lesson / "business_plan_review.html"
    payload = b"""<!doctype html><html><head><meta charset="utf-8"></head>
    <body><main><button data-report-print hidden>Print</button>
    <nav><a href="#economics">Economics</a></nav><h1>Test before launch</h1>
    <section id="economics"><table><tr><td>January</td><td>-260.00</td></tr></table></section>
    <p><a href="#economics" data-calculation-id="pilot/2027-01/commercial_operating_result">-260.00 EUR</a></p>
    <details><summary>Sources</summary><p>pilot-economics.csv</p></details>
    </main><script type="application/json" id="validated-plan">{"private_payload":true}</script>
    <script>fetch('https://example.com/leak')</script></body></html>"""
    original.write_bytes(payload)
    receipt = json.loads((lesson / "execution.json").read_text())
    source_csv = lesson / "pilot-economics.csv"
    source_csv.write_text("period,units,net_price_eur\n2027-01,20,45\n")
    receipt["inputs"].append({"path": source_csv.name, "sha256": digest(source_csv)})
    receipt.update(
        product=product,
        workflow_id="business-planning",
        skill_sha256=digest(
            ROOT / f"plugins/{product}/skills/business-planning/SKILL.md"
        ),
        outputs=[{"path": original.name, "sha256": digest(original)}],
    )
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "result-view"

    library.results("business-planning", "it", lesson, "execution.json", destination)

    view = (destination / "result-01.html").read_text()
    assert "<h1>Test before launch</h1>" in view
    assert "<td>-260.00</td>" in view
    assert "href='#html-result-economics'" in view
    assert "id='html-result-economics'" in view
    assert "<summary>Sources</summary>" in view
    assert "data-calculation-id" not in view
    assert ">-260.00 EUR</a>" in view
    assert "<script" not in view and "<button" not in view
    assert "private_payload" not in view and "example.com/leak" not in view
    assert "href='source-02.html'>pilot-economics.csv</a>" in view
    source_view = (destination / "source-02.html").read_text()
    assert "<td>20</td>" in source_view and "<td>45</td>" in source_view
    assert (
        destination / "outputs/source-02-pilot-economics.csv"
    ).read_bytes() == source_csv.read_bytes()
    assert (
        destination / "outputs/01-business_plan_review.html"
    ).read_bytes() == payload
    assert original.read_bytes() == payload


@pytest.mark.parametrize(
    "active", ['<form action="/send">Send</form>', '<p onclick="x()">Click</p>']
)
def test_business_reading_still_rejects_other_active_content(active):
    from courseware.html_view import passive_html_body

    with pytest.raises(ValueError, match="Unsupported result format"):
        passive_html_body(active, omit_controls=True, report_metadata=True)


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_result_view_opens_passive_html_without_rewriting(
    result_view_case, tmp_path, language
):
    library, lesson = result_view_case
    original = lesson / "preview.html"
    payload = "<!doctype html><html><meta charset='utf-8'><style>body{color:#123}</style><main><h1>Invoice DID-001</h1><table><tr><td>122.00 EUR</td></tr></table><details><summary>Sources</summary><p>invoice-data.txt</p></details></main></html>".encode()
    original.write_bytes(payload)
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"] = [{"path": original.name, "sha256": digest(original)}]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "result-view"
    library.results(
        "variance-analysis", language, lesson, "execution.json", destination
    )
    assert "href='result-01.html'" in (destination / "course.html").read_text()
    view = (destination / "result-01.html").read_text()
    assert "<h1>Invoice DID-001</h1>" in view
    assert "<td>122.00 EUR</td>" in view
    assert "<summary>Sources</summary>" in view
    assert "<iframe" not in view and "<style>" not in view
    assert "href='outputs/01-preview.html'" in view
    assert "href='course.html'" in view
    assert (destination / "outputs/01-preview.html").read_bytes() == payload
    assert original.read_bytes() == payload


@pytest.mark.parametrize(
    "html",
    [
        '<div onclick="alert(1)">Click</div>',
        '<meta http-equiv="refresh" content="0;url=https://example.com">',
        '<iframe src="https://example.com"></iframe>',
        '<a href="https://example.com">External</a>',
        '<form action="https://example.com">Send</form>',
        '<svg onload="alert(1)"></svg>',
    ],
)
def test_result_view_rejects_active_html_before_copying(
    result_view_case, tmp_path, html
):
    library, lesson = result_view_case
    original = lesson / "preview.html"
    original.write_text(html)
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"] = [{"path": original.name, "sha256": digest(original)}]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "result-view"
    with pytest.raises(ValueError, match="Unsupported result format"):
        library.results(
            "variance-analysis", "it", lesson, "execution.json", destination
        )
    assert not destination.exists()


def test_result_view_rejects_changed_execution_before_writing(
    result_view_case, tmp_path
):
    library, lesson = result_view_case
    (lesson / "fatture_summary.csv").write_text("changed output")
    destination = tmp_path / "result-view"
    with pytest.raises(ValueError):
        library.results(
            "variance-analysis", "it", lesson, "execution.json", destination
        )
    assert not destination.exists()


def test_html_escapes_case_text_and_keeps_it_as_evidence(isolated, tmp_path):
    root, _, _ = isolated
    alter_course(
        isolated,
        lambda course: course["locales"]["it"].update(
            title="<script>alert('x')</script>"
        ),
    )
    output = tmp_path / "escaped"
    CourseLibrary(root, {"variance-analysis"}).render("variance-analysis", "it", output)
    page = (output / "course.html").read_text(encoding="utf-8")
    assert "<script>" not in page
    assert "&lt;script&gt;" in page


def test_cli_lists_own_courses(isolated, capsys):
    root, _, _ = isolated
    assert main(root, {"variance-analysis"}, ["list"]) == 0
    assert json.loads(capsys.readouterr().out)[0]["workflow"] == "variance-analysis"


def test_cli_shows_prepared_source_checked_content(isolated, capsys):
    root, _, _ = isolated
    assert (
        main(
            root,
            {"variance-analysis"},
            ["show", "--workflow", "variance-analysis", "--language", "it"],
        )
        == 0
    )
    result = json.loads(capsys.readouterr().out)
    assert result["workflow"] == "variance-analysis"
    assert result["language"] == "it"
    assert result["sources_current"] is True
    assert "locales" not in result
    assert result["content"]["checkpoints"]


def test_cli_preserves_localized_symbols_on_legacy_windows_console(
    monkeypatch, isolated
):
    root, _, _ = isolated
    alter_course(
        isolated, lambda course: course["locales"]["en"].update(title="Δ 预算")
    )
    buffer = io.BytesIO()
    console = io.TextIOWrapper(buffer, encoding="cp1252")
    monkeypatch.setattr(sys, "stdout", console)

    assert (
        main(
            root,
            {"variance-analysis"},
            ["show", "--workflow", "variance-analysis", "--language", "en"],
        )
        == 0
    )

    console.flush()
    result = json.loads(buffer.getvalue().decode("cp1252"))
    assert result["content"]["title"] == "Δ 预算"


@pytest.mark.parametrize("default_encoding", ["utf-8", "cp1252"])
def test_complete_compiled_catalog_matches_authoring_and_current_sources(
    monkeypatch, default_encoding
):
    original_read_text = Path.read_text

    def read_with_host_default(path, encoding=None, errors=None):
        return original_read_text(
            path, encoding=encoding or default_encoding, errors=errors
        )

    monkeypatch.setattr(Path, "read_text", read_with_host_default)
    assert builder.build(check=True) == {
        "workflow_kits": len(COURSES),
        "localized_kits": len(LOCALIZED)
        + sum(
            len(
                json.loads(
                    (
                        ROOT
                        / "scripts/course_materials/published"
                        / key
                        / "course.json"
                    ).read_text()
                )["locales"]
            )
            for key in RELEASE["retained"]
        ),
        "missing_translations": [],
    }


def test_compiler_removes_retired_specimens_and_rejects_stale_extra_files(tmp_path):
    kit = tmp_path / "generated-kit"
    current = kit / "files/input/invoice.xml"
    current.parent.mkdir(parents=True)
    current.write_text("current fictional source", encoding="utf-8")
    old = kit / "files/executed-starter/old-result.html"
    old.parent.mkdir()
    old.write_text("retired output specimen", encoding="utf-8")
    expected = {"course.json", "files/input/invoice.xml"}
    with pytest.raises(ValueError, match="Unreferenced generated teaching files"):
        builder._generated_files(kit, expected, check=True)
    assert old.exists()
    builder._generated_files(kit, expected)
    assert not old.parent.exists()
    assert current.read_text(encoding="utf-8") == "current fictional source"
    builder._generated_files(kit, expected, check=True)


def test_compiler_refuses_linked_generated_assets_without_touching_target(tmp_path):
    kit = tmp_path / "generated-kit"
    kit.mkdir()
    outside = tmp_path / "outside"
    outside.write_text("preserve", encoding="utf-8")
    (kit / "linked-input").symlink_to(outside)
    with pytest.raises(ValueError, match="symlinks"):
        builder._generated_files(kit, set())
    assert outside.read_text(encoding="utf-8") == "preserve"


def test_compiler_rejects_cross_language_file_collision_before_changing_kit(
    tmp_path, monkeypatch
):
    authoring = tmp_path / "authoring"
    write_json(
        authoring / "kits.json",
        {
            "vera/previdenza-inps": {
                "files": [
                    {
                        "source": "italian.csv",
                        "path": "files/input/periods.csv",
                        "role": "source",
                        "languages": ["it"],
                    },
                    {
                        "source": "english.csv",
                        "path": "files/input/periods.csv",
                        "role": "source",
                        "languages": ["en"],
                    },
                ]
            }
        },
    )
    write_json(
        authoring / "release_plan.json",
        {"prepared": ["vera/previdenza-inps"], "retained": {}},
    )
    kit = tmp_path / "plugins/vera/assets/courses/previdenza-inps"
    kit.mkdir(parents=True)
    existing = kit / "existing-artifact.txt"
    existing.write_text("Preserve the existing kit on invalid authoring input")
    monkeypatch.setattr(builder, "ROOT", tmp_path)
    monkeypatch.setattr(builder, "AUTHORING", authoring)
    monkeypatch.setattr(
        builder,
        "_eligible",
        lambda product: {"previdenza-inps"} if product == "vera" else set(),
    )
    monkeypatch.setattr(
        builder,
        "_locales",
        lambda: {"vera/previdenza-inps": {lang: {} for lang in builder.LANGUAGES}},
    )
    with pytest.raises(ValueError, match="unique destinations across all languages"):
        builder.build(require_complete=False)
    assert list(kit.iterdir()) == [existing]
    assert (
        existing.read_text() == "Preserve the existing kit on invalid authoring input"
    )


@pytest.mark.parametrize("product,workflow,language", LOCALIZED)
def test_every_supported_locale_has_complete_renderable_material(
    tmp_path, product, workflow, language
):
    library = CourseLibrary(ROOT / "plugins" / product, builder._eligible(product))
    content = library.load(workflow, language)["locales"][language]
    assert set(content) == {
        "title",
        "goal",
        "scenario",
        "request",
        "scope",
        "inputs",
        "steps",
        "deliverables",
        "review",
        "checkpoints",
        "practice",
        "success",
        "repeat",
    }
    assert all(content.values())
    assert all(content["steps"])
    assert all(content["checkpoints"])
    destination = tmp_path / "lesson"

    library.render(workflow, language, destination)

    page = (destination / "course.html").read_text(encoding="utf-8")
    assert f"<html lang='{language}'>" in page
    assert page.count("class='lesson-step'") == 6
    assert "execution-request.json" not in page
    assert "<script" not in page
    assert "https://" not in page
    assert not (destination / "example.html").exists()
    assert (destination / "execution-request.json").is_file()


@pytest.mark.parametrize("language", builder.LANGUAGES)
def test_lucia_website_kit_remains_owned_by_lucia(language):
    library = CourseLibrary(ROOT / "plugins/lucia", builder._eligible("lucia"))
    course = library.load("presenza-digitale-studio", language)
    content = course["locales"][language]
    assert "Lucia" in content["request"]
    assert "Vera" not in json.dumps(content)
    assert course["product"] == "lucia"


def test_precomputed_output_cannot_be_packaged_as_a_teaching_input(isolated):
    root, _, _ = isolated
    alter_course(isolated, lambda course: course["files"][0].update(role="output"))
    with pytest.raises(CourseError, match="source and practice files"):
        CourseLibrary(root, {"variance-analysis"}).load("variance-analysis", "it")


def test_kit_cannot_omit_the_learners_practice_inputs(isolated):
    root, _, _ = isolated
    alter_course(
        isolated,
        lambda course: course.update(
            files=[f for f in course["files"] if f["role"] == "source"]
        ),
    )
    with pytest.raises(CourseError, match="needs source and practice"):
        CourseLibrary(root, {"variance-analysis"}).load("variance-analysis", "it")


def test_cli_refuses_foreign_workflow(isolated, capsys):
    root, _, _ = isolated
    with pytest.raises(SystemExit) as error:
        main(
            root,
            {"variance-analysis"},
            ["show", "--workflow", "reporting-engine", "--language", "it"],
        )
    assert error.value.code == 2
    assert "Course unavailable" in capsys.readouterr().err


@pytest.mark.parametrize(
    "argv",
    [["show"], ["render", "--workflow", "variance-analysis", "--language", "it"]],
)
def test_cli_requires_explicit_identity_and_destination(isolated, argv):
    root, _, _ = isolated
    with pytest.raises(SystemExit) as error:
        main(root, {"variance-analysis"}, argv)
    assert error.value.code == 2


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_pdf_inputs_have_visible_page_previews_and_keep_originals(tmp_path, language):
    from html.parser import HTMLParser

    class Links(HTMLParser):
        def __init__(self):
            super().__init__()
            self.links = []
            self.images = []

        def handle_starttag(self, tag, attrs):
            values = dict(attrs)
            if tag == "a":
                self.links.append(values)
            if tag == "img":
                self.images.append(values)

    library = CourseLibrary(ROOT / "plugins/vera", builder._eligible("vera"))
    destination = tmp_path / "pdf-lesson"
    result = library.render("open-item-reconciliation", language, destination)
    guide = Links()
    guide.feed((destination / "course.html").read_text())
    input_links = [
        link for link in guide.links if link.get("href", "").startswith("input-")
    ]
    assert len(input_links) == 4
    for link in input_links:
        view = Links()
        view.feed((destination / link["href"]).read_text())
        assert len(view.images) == 1  # Every prepared PDF in this kit has one page.
        page = destination / view.images[0]["src"]
        assert page.read_bytes().startswith(b"\x89PNG\r\n\x1a\n")
        assert view.images[0]["alt"]
        original = next(item for item in view.links if "download" in item)
        supplied = destination / original["href"]
        source = (
            ROOT
            / "plugins/vera/assets/courses/open-item-reconciliation"
            / original["href"]
        )
        with pymupdf.open(source) as source_document:
            assert len(view.images) == source_document.page_count
        assert supplied.read_bytes() == source.read_bytes()
        expected_step = "step-5" if "/practice/" in original["href"] else "step-1"
        assert any(
            item.get("href") == "course.html#" + expected_step for item in view.links
        )
    assert all(Path(path).suffix == ".pdf" for path in result["source_files"])
    assert all(Path(path).suffix == ".pdf" for path in result["practice_files"])


def test_result_view_keeps_technical_card_after_learner_outputs(
    result_view_case, tmp_path
):
    library, lesson = result_view_case
    card = lesson / "artifact_card.md"
    card.write_text("# Technical card\n\npython internal.py --run test")
    receipt_path = lesson / "execution.json"
    receipt = json.loads(receipt_path.read_text())
    receipt["outputs"].insert(0, {"path": card.name, "sha256": digest(card)})
    write_json(receipt_path, receipt)
    result = library.results(
        "variance-analysis", "it", lesson, "execution.json", tmp_path / "view"
    )
    page = Path(result["course"]).read_text()
    assert page.index("<td>976.00</td>") < page.index("Dati tecnici del risultato")
    assert "<h3>Technical card</h3>" not in page
    assert "<pre class='source-text'># Technical card" in page


def test_result_view_shows_summary_when_it_is_the_delivered_result(
    result_view_case, tmp_path
):
    library, lesson = result_view_case
    card = lesson / "artifact_card.md"
    card.write_text(
        "# Risultato del riordino\n\nDue documenti archiviati.\n\n| Prima | Dopo |\n|---|---|\n| Download/verbale.md | Riunioni/verbale.md |\n"
    )
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"] = [{"path": card.name, "sha256": digest(card)}]
    write_json(lesson / "execution.json", receipt)
    result = library.results(
        "variance-analysis", "it", lesson, "execution.json", tmp_path / "view"
    )
    page = Path(result["course"]).read_text()
    assert "<h2>Risultato del riordino</h2>" in page
    assert "<p>Due documenti archiviati.</p>" in page
    assert "<td>Riunioni/verbale.md</td>" in page
    assert "Dati tecnici del risultato" not in page
    assert (
        tmp_path / "view/outputs/01-artifact_card.md"
    ).read_bytes() == card.read_bytes()


def test_clara_assignment_plan_is_readable_first_with_bound_output_links(tmp_path):
    root = ROOT / "plugins/clara"
    workflow = "advisory-brief-planner"
    lesson = tmp_path / "lesson"
    lesson.mkdir()
    sources = {
        "assignment.md": "Fictional assignment for reader regression.",
        "validation.json": '{"synthetic_test_fixture": true}',
        "advisory_contract.json": '{"decision": "Synthetic decision"}',
        "generation_handoff.md": "# Prossimo lavoro\n\nConfrontare le opzioni.",
        "codex_run_review.md": (
            "# Incarico da rivedere\n\nIl piano conserva le lacune.\n\n"
            "Leggi [le consegne](generation_handoff.md) e "
            "[la fonte](assignment.md). "
            "Non attivare [il file estraneo](unrecorded.md) o "
            "[la rete](https://example.invalid).\n"
        ),
    }
    for name, content in sources.items():
        (lesson / name).write_text(content)

    def record(name):
        return {"path": name, "sha256": digest(lesson / name)}

    write_json(
        lesson / "execution.json",
        {
            "schema": "mparanza.teaching_execution.v1",
            "evidence_kind": "host_attested_local_execution",
            "product": "clara",
            "workflow_id": workflow,
            "phase": "demo",
            "worker_thread_id": "synthetic-reader-test",
            "outcome": "completed",
            "skill_sha256": digest(root / f"skills/{workflow}/SKILL.md"),
            "inputs": [record("assignment.md")],
            "native_records": [record("validation.json")],
            "outputs": [
                record("advisory_contract.json"),
                record("generation_handoff.md"),
                record("codex_run_review.md"),
            ],
        },
    )

    result = CourseLibrary(root, {workflow}).results(
        workflow, "it", lesson, "execution.json", tmp_path / "view"
    )

    page = Path(result["course"]).read_text()
    assert page.index("<h2>Incarico da rivedere</h2>") < page.index(
        "<h2>advisory_contract.json</h2>"
    )
    assert "<p>Il piano conserva le lacune.</p>" in page
    assert "href='#result-03'" in page
    assert "id='result-03'" in page
    assert "href='source-01.html'" in page
    assert "href='unrecorded.md'" not in page
    assert "href='https://example.invalid'" not in page
    assert (tmp_path / "view/outputs/01-codex_run_review.md").read_bytes() == (
        lesson / "codex_run_review.md"
    ).read_bytes()


@pytest.mark.parametrize(
    "language,previous",
    [
        ("it", "Versione precedente"),
        ("en", "Previous version"),
        ("fr", "Version précédente"),
        ("de", "Vorherige Version"),
        ("es", "Versión anterior"),
    ],
)
@pytest.mark.parametrize("history_folder", ["history", "previous-practice"])
def test_director_history_is_named_and_reachable_before_long_evidence_map(
    tmp_path, language, previous, history_folder
):
    root = ROOT / "plugins/clara"
    workflow = "advisory-case-director"
    lesson = tmp_path / "lesson"
    lesson.mkdir()
    history = history_folder + "/advisory_workpaper.20260927T1958310000.12345678.md"
    files = {
        "assignment.md": "Fictional test input.",
        "validation.json": '{"synthetic_test_fixture": true}',
        "advisory_workpaper.md": "# Current answer\n\nThe new source changes the answer.",
        "advisory_evidence_map.md": "# Evidence map\n\n" + "Source details.\n\n" * 80,
        history: "# Original answer\n\nThe original wording must remain unchanged.",
    }
    for name, content in files.items():
        target = lesson / name
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(content)
    record = lambda name: {"path": name, "sha256": digest(lesson / name)}
    write_json(
        lesson / "execution.json",
        {
            "schema": "mparanza.teaching_execution.v1",
            "evidence_kind": "host_attested_local_execution",
            "product": "clara",
            "workflow_id": workflow,
            "phase": "practice",
            "worker_thread_id": "synthetic-reader-test",
            "outcome": "review_required",
            "skill_sha256": digest(root / f"skills/{workflow}/SKILL.md"),
            "inputs": [record("assignment.md")],
            "native_records": [record("validation.json")],
            "outputs": [
                record(name)
                for name in [
                    "advisory_workpaper.md",
                    "advisory_evidence_map.md",
                    history,
                ]
            ],
        },
    )
    result = CourseLibrary(root, {workflow}).results(
        workflow, language, lesson, "execution.json", tmp_path / "view"
    )
    page = Path(result["course"]).read_text()
    label = previous + " — Original answer"
    assert f"<a href='#result-03'>{label}</a>" in page
    assert page.index("<nav ") < page.index("<section id='result-01'")
    assert f"<section id='result-03'><h2>{label}</h2>" in page
    assert previous + " — Current answer" not in page
    assert (tmp_path / "view/outputs" / ("03-" + Path(history).name)).read_bytes() == (
        lesson / history
    ).read_bytes()


def test_result_view_displays_xml_as_escaped_data_and_preserves_original(
    result_view_case, tmp_path
):
    library, lesson = result_view_case
    original = lesson / "invoice.xml"
    payload = b'<?xml version="1.0"?><Invoice><Total>244.00</Total><script>literal only</script></Invoice>'
    original.write_bytes(payload)
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"] = [{"path": original.name, "sha256": digest(original)}]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "xml-results"
    library.results("variance-analysis", "it", lesson, "execution.json", destination)
    page = (destination / "course.html").read_text()
    assert "File XML prodotto" in page and "&lt;Total&gt;244.00&lt;/Total&gt;" in page
    assert "<script>" not in page
    assert "download='invoice.xml'" in page
    assert (destination / "outputs/01-invoice.xml").read_bytes() == payload
    assert original.read_bytes() == payload


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_journal_inputs_open_as_tables_preserving_originals(tmp_path, language):
    import hashlib

    library = CourseLibrary(ROOT / "plugins/vera", builder._eligible("vera"))
    destination = tmp_path / "journal-lesson"
    library.render("journal-sampling", language, destination)
    guide = (destination / "course.html").read_text()
    for phase, name in [
        ("input", "journal-march.csv"),
        ("practice", "journal-april.csv"),
    ]:
        relative = f"files/{phase}/{name}"
        view = "input-" + hashlib.sha256(relative.encode()).hexdigest()[:16] + ".html"
        assert f"href='{view}'" in guide
        page = (destination / view).read_text()
        assert "<table>" in page and "<th scope='col'>Entry</th>" in page
        assert "href='course.html#step-" in page
        assert f"href='{relative}' download" in page
        original = ROOT / "plugins/vera/assets/courses/journal-sampling" / relative
        assert (destination / relative).read_bytes() == original.read_bytes()


def test_source_csv_reading_preserves_quoted_cells_and_escapes_markup():
    from courseware.library import _csv_input

    page = _csv_input('conto;importo;nota\n001;"1.234,00";"<script>bad</script>"\n')
    assert "<td>001</td>" in page and "<td>1.234,00</td>" in page
    assert "&lt;script&gt;bad&lt;/script&gt;" in page and "<script>" not in page


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_quotation_reading_labels_status_without_changing_values(language):
    from courseware.library import _csv_input

    ui = json.loads(
        (
            ROOT / "plugins/_shared/vendor/modules/courseware/assets/languages.json"
        ).read_text()
    )[language]
    page = _csv_input(
        "item_id,description,amount_eur,status\n"
        "Q-001,Software,6000.00,fictional_unaccepted_quotation\n"
        "Q-002,fictional_unaccepted_quotation,001.00,unknown_status\n",
        ui,
    )
    assert f"<th scope='col'>{ui['csv_input_fields']['amount_eur']}</th>" in page
    label = ui["csv_input_values"]["status"]["fictional_unaccepted_quotation"]
    assert f"<td>{label}</td>" in page
    assert "<td>6000.00</td>" in page and "<td>001.00</td>" in page
    assert "<td>unknown_status</td>" in page
    assert "<td>fictional_unaccepted_quotation</td>" in page


def test_invoice_result_starts_with_exceptions_and_keeps_technical_details(
    result_view_case, tmp_path
):
    from openpyxl import Workbook

    library, lesson = result_view_case
    original = lesson / "exception_workpaper.xlsx"
    workbook = Workbook()
    workbook.active.title = "Summary"
    workbook.active.append(["run_fingerprint", "technical-only"])
    exceptions = workbook.create_sheet("Exceptions")
    exceptions.append(
        [
            "invoice_number",
            "invoice_date",
            "gross_amount",
            "booked_accounts",
            "luna_status",
            "luna_reason",
            "invoice_evidence",
            "booked_account_evidence",
            "professional_should_inspect",
        ]
    )
    exceptions.append(
        [
            "DEMO-003",
            "2026-03-22",
            488,
            "6050 — Carburante",
            "review_required",
            "<unsafe>source</unsafe>",
            "Consulting",
            "Fuel",
            "Review classification",
        ]
    )
    workbook.save(original)
    summary = lesson / "run_summary.md"
    summary.write_text("# Technical summary\n\n- run_fingerprint: technical-only\n")
    receipt_path = lesson / "execution.json"
    receipt = json.loads(receipt_path.read_text())
    receipt["outputs"] = [
        {"path": p.name, "sha256": digest(p)} for p in [original, summary]
    ]
    write_json(receipt_path, receipt)
    before = original.read_bytes()
    destination = tmp_path / "result-view"
    library.results("variance-analysis", "it", lesson, "execution.json", destination)
    page = (destination / "result-01.html").read_text()
    assert "class='sheet-view' open><summary>Exceptions</summary>" in page
    assert page.index("<summary>Exceptions") < page.index("<summary>Summary")
    assert "class='sheet-view'><summary>Summary</summary>" in page
    assert "&lt;unsafe&gt;source&lt;/unsafe&gt;" in page
    assert "<th scope='row'>Motivo della segnalazione</th>" in page
    assert "<td>Review classification</td>" in page
    landing = (destination / "course.html").read_text()
    assert "Eccezioni da rivedere" in landing
    assert "<details><summary>Leggi il file completo</summary><pre" in landing
    assert original.read_bytes() == before
    assert (destination / "outputs/01-exception_workpaper.xlsx").read_bytes() == before


def test_result_view_renders_markdown_tables_without_executing_cell_content(
    result_view_case, tmp_path
):
    library, lesson = result_view_case
    note = lesson / "codex_run_review.md"
    note.write_text(
        "# Net debt\n\nBefore.\n\n| Item | EUR |\n| :--- | ---: |\n| Loan | 65000 |\n| Cash | 27000 |\n| Net debt | 38000 |\n| <script>bad</script> | a\\|b |\n\nAfter.\n\n| not | a table |\n| invalid | separator |\n"
    )
    original = note.read_bytes()
    receipt_path = lesson / "execution.json"
    receipt = json.loads(receipt_path.read_text())
    receipt["outputs"] = [{"path": note.name, "sha256": digest(note)}]
    write_json(receipt_path, receipt)
    destination = tmp_path / "markdown-results"
    library.results("variance-analysis", "it", lesson, "execution.json", destination)
    page = (destination / "course.html").read_text()
    assert "<th scope='col'>Item</th>" in page
    assert "<td>Net debt</td><td>38000</td>" in page
    assert "<td>&lt;script&gt;bad&lt;/script&gt;</td><td>a|b</td>" in page
    assert "<script>" not in page
    assert "<p>After.</p>" in page
    assert "| not | a table | | invalid | separator |" in page
    assert note.read_bytes() == original
    assert (destination / "outputs/01-codex_run_review.md").read_bytes() == original


def test_result_procedure_keeps_numbered_steps_separate_and_original_unchanged(
    result_view_case, tmp_path
):
    library, lesson = result_view_case
    note = lesson / "PROCEDURA.md"
    note.write_text(
        "# Procedura\n\n1. Apri la pagina.\n2. Seleziona **Invoice**.\n"
        "3. Controlla `<script>bad</script>`.\n- Conserva il risultato.\n\n"
        "4) Riprendi la procedura.\n5) Verifica la nota di credito.\n\n"
        "| Esito | Stato |\n| --- | --- |\n| Pacchetto | Pronto |\n"
    )
    original = note.read_bytes()
    receipt_path = lesson / "execution.json"
    receipt = json.loads(receipt_path.read_text())
    receipt["outputs"] = [{"path": note.name, "sha256": digest(note)}]
    write_json(receipt_path, receipt)
    destination = tmp_path / "procedure-results"
    library.results("variance-analysis", "it", lesson, "execution.json", destination)
    page = (destination / "course.html").read_text()
    assert "<ol><li>Apri la pagina.</li><li>Seleziona <strong>Invoice</strong>." in page
    assert "<code>&lt;script&gt;bad&lt;/script&gt;</code>.</li></ol><ul>" in page
    assert "</ul><ol start='4'><li>Riprendi la procedura.</li>" in page
    assert "</li></ol><div class='table-scroll markdown-table'" in page
    assert "<script>" not in page
    assert note.read_bytes() == original
    assert (destination / "outputs/01-PROCEDURA.md").read_bytes() == original


@pytest.mark.parametrize("extension", ["xlsx", "docx"])
@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_prepared_office_inputs_open_readable_views_without_changing_sources(
    isolated, tmp_path, extension, language
):
    root, manifest, _ = isolated
    relative = f"files/input/source.{extension}"
    source = manifest.parent / relative
    if extension == "xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.title = "Creditori"
        workbook.active.append([None, None, None])
        workbook.active.append(["Input di prova", None, None])
        workbook.active.append(["Creditore", "Importo"])
        workbook.active.append(["<script>literal</script>", 120000])
        workbook.active.cell(row=20, column=6).number_format = "0.00"
        workbook.save(source)
    else:
        from docx import Document

        document = Document()
        document.add_heading("Creditori", level=1)
        document.add_paragraph("<script>literal</script> 120000")
        document.save(source)
    original = source.read_bytes()

    def replace_source(course):
        course["files"] = [x for x in course["files"] if x["role"] != "source"]
        course["files"].append(
            {
                "path": relative,
                "role": "source",
                "languages": ["it", "en", "fr", "de", "es"],
                "sha256": digest(source),
            }
        )

    alter_course(isolated, replace_source)
    destination = tmp_path / "office-guide"
    CourseLibrary(root, {"variance-analysis"}).render(
        "variance-analysis", language, destination
    )
    preview = "input-" + hashlib.sha256(relative.encode()).hexdigest()[:16] + ".html"
    page = (destination / preview).read_text()
    assert f"href='{preview}'" in (destination / "course.html").read_text()
    assert "Creditori" in page and "120000" in page
    assert "&lt;script&gt;literal&lt;/script&gt;" in page
    assert "<script>" not in page
    if extension == "xlsx":
        assert "colspan='2'>Input di prova" in page
        assert "<th scope='col'></th>" not in page
        assert "<tr><td></td><td></td></tr>" not in page
    assert f"href='{relative}' download" in page
    assert "href='course.html#step-1'" in page
    assert (destination / relative).read_bytes() == original
    assert source.read_bytes() == original


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_macro_workbook_input_link_explains_download_and_preserves_workflow_inspection(
    isolated, tmp_path, language
):
    root, manifest, _ = isolated
    relative = "files/input/balances.xlsm"
    source = manifest.parent / relative
    original = (
        b"\x00\xffOpaque Office data: the course must not parse this as model context."
    )
    source.write_bytes(original)

    def replace_source(course):
        course["files"] = [x for x in course["files"] if x["role"] != "source"]
        course["files"].append(
            {
                "path": relative,
                "role": "source",
                "languages": ["it", "en", "fr", "de", "es"],
                "sha256": digest(source),
            }
        )

    alter_course(isolated, replace_source)
    destination = tmp_path / "office-guide"
    CourseLibrary(root, {"variance-analysis"}).render(
        "variance-analysis", language, destination
    )
    preview = "input-" + hashlib.sha256(relative.encode()).hexdigest()[:16] + ".html"
    assert f"href='{preview}'" in (destination / "course.html").read_text()
    page = (destination / preview).read_text()
    assert f"href='{relative}' download" in page
    assert "href='course.html#step-1'" in page
    assert "Vera" in page and "{product}" not in page
    assert "Opaque Office data" not in page
    assert (destination / relative).read_bytes() == original


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_result_view_displays_original_chart_bytes(
    result_view_case, tmp_path, language
):
    from PIL import Image

    library, lesson = result_view_case
    chart = lesson / "waterfall.png"
    Image.new("RGB", (160, 90), "white").save(chart)
    original = chart.read_bytes()
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"].append({"path": chart.name, "sha256": digest(chart)})
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "charts"
    result = library.results(
        "variance-analysis", language, lesson, "execution.json", destination
    )
    page = Path(result["course"]).read_text()
    assert "<img src='outputs/03-waterfall.png'" in page
    assert "alt='waterfall.png'" in page
    assert (destination / "outputs/03-waterfall.png").read_bytes() == original
    assert chart.read_bytes() == original


def test_result_view_rejects_disguised_chart_before_writing(result_view_case, tmp_path):
    library, lesson = result_view_case
    chart = lesson / "waterfall.png"
    chart.write_text("<html><script>unsafe()</script></html>")
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"].append({"path": chart.name, "sha256": digest(chart)})
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "charts"
    with pytest.raises(CourseError, match="image is unreadable"):
        library.results(
            "variance-analysis", "it", lesson, "execution.json", destination
        )
    assert not destination.exists()


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_result_markdown_embeds_only_evidenced_sibling_images(
    result_view_case, tmp_path, language
):
    from PIL import Image

    library, lesson = result_view_case
    chart = lesson / "waterfall.png"
    Image.new("RGB", (160, 90), "white").save(chart)
    report = lesson / "report.md"
    report.write_text(
        "**Reading**: <script>unsafe()</script>\n\n"
        "![Actual chart](waterfall.png)\n\n"
        "![Remote](https://example.invalid/tracker.png)\n\n"
        "![Outside](../private.png)\n"
    )
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"].extend(
        {"path": p.name, "sha256": digest(p)} for p in (report, chart)
    )
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "markdown-images"
    library.results(
        "variance-analysis", language, lesson, "execution.json", destination
    )
    page = (destination / "course.html").read_text()
    assert "<strong>Reading</strong>: &lt;script&gt;" in page
    assert "<img src='outputs/04-waterfall.png' alt='Actual chart'" in page
    assert "<img src='https:" not in page
    assert "<img src='../" not in page
    assert "<script>unsafe()" not in page
    assert (destination / "outputs/04-waterfall.png").read_bytes() == chart.read_bytes()


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_sales_plan_view_keeps_totals_and_assumptions_readable(
    result_view_case, tmp_path, language
):
    _, lesson = result_view_case
    library = CourseLibrary(ROOT / "plugins/vera", {"sales-plan"})
    summary = lesson / "scenario_summary.csv"
    summary.write_text(
        "summary_level,dimension_name,dimension_value,metric,unit,actual,plan,delta,delta_pct_rounded_4dp\n"
        "dimension,product,<unsafe>,units,count,25,27.5,2.5,10\n"
        "total,,,gross_sales_reporting,EUR,91000,100100,9100,10\n"
    )
    ledger = lesson / "assumption_application_ledger.csv"
    ledger.write_text(
        "assumption_id,source_row_id,target_period,driver,before_value,after_value,change_pct,status\n"
        "growth,trekking-feb,2027-02,units_pct,25,27.5,10,applied\n"
    )
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["workflow_id"] = "sales-plan"
    receipt["skill_sha256"] = digest(ROOT / "plugins/vera/skills/sales-plan/SKILL.md")
    receipt["outputs"] = [
        {"path": p.name, "sha256": digest(p)} for p in (summary, ledger)
    ]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "sales-results"
    library.results("sales-plan", language, lesson, "execution.json", destination)
    page = (destination / "course.html").read_text()
    labels = json.loads(
        (
            ROOT / "plugins/_shared/vendor/modules/courseware/assets/languages.json"
        ).read_text()
    )[language]["sales_plan"]
    assert page.index("<h3>" + labels["total"]) < page.index("&lt;unsafe&gt;")
    assert labels["metrics"]["gross_sales_reporting"] in page
    assert "<th scope='col'>summary_level" not in page
    assert ("27.5" if language == "en" else "27,5") in page
    assert "<unsafe>" not in page
    assert (
        destination / "outputs/01-scenario_summary.csv"
    ).read_bytes() == summary.read_bytes()
    assert (
        destination / "outputs/02-assumption_application_ledger.csv"
    ).read_bytes() == ledger.read_bytes()


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_result_view_respects_hidden_excel_rows_and_column_ranges(
    result_view_case, tmp_path, language
):
    from openpyxl import Workbook

    library, lesson = result_view_case
    original = lesson / "variance_results.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "Category",
            "Amount",
            "Internal price",
            "Internal volume",
            "Internal mix",
            "Visible control",
        ]
    )
    sheet.append(
        [
            "Revenue",
            20000,
            "HIDDEN_PRICE",
            "HIDDEN_VOLUME",
            "HIDDEN_MIX",
            "VISIBLE_CONTROL",
        ]
    )
    sheet.append(["HIDDEN_ROW", 999, None, None, None, "HIDDEN_CONTROL"])
    sheet.column_dimensions.group("C", "E", hidden=True)
    sheet.row_dimensions[3].hidden = True
    workbook.save(original)
    before = original.read_bytes()
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"] = [{"path": original.name, "sha256": digest(original)}]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "visible"
    library.results(
        "variance-analysis", language, lesson, "execution.json", destination
    )
    page = (destination / "result-01.html").read_text()
    assert "Revenue" in page and "20000" in page and "VISIBLE_CONTROL" in page
    assert "HIDDEN_" not in page and "Internal volume" not in page
    assert (destination / "outputs/01-variance_results.xlsx").read_bytes() == before
    assert original.read_bytes() == before


def test_result_view_preserves_percentage_units_without_changing_workbook(
    result_view_case, tmp_path
):
    from openpyxl import Workbook

    library, lesson = result_view_case
    source = lesson / "percentages.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Scostamento %", "Rapporto grezzo", "Testo", "Formula"])
    sheet.append([0.12765957, 0.12765957, "<script>bad</script>", "=1/8"])
    sheet["A2"].number_format = "0.00%"
    sheet["D2"].number_format = "0.0%"
    sheet.append([-0.20689655, 0, "unchanged", None])
    sheet["A3"].number_format = "0.0%"
    workbook.save(source)
    before = source.read_bytes()
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"] = [{"path": source.name, "sha256": digest(source)}]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "percentage-results"
    library.results("variance-analysis", "it", lesson, "execution.json", destination)
    page = (destination / "result-01.html").read_text()
    assert "<td>12.77%</td>" in page
    assert "<td>-20.7%</td>" in page
    assert "<td>0.12765957</td>" in page
    assert "&lt;script&gt;bad&lt;/script&gt;" in page
    assert "<td>=1/8</td>" not in page
    assert source.read_bytes() == before
    assert (destination / "outputs/01-percentages.xlsx").read_bytes() == before


def test_answer_results_keep_full_audit_collapsed_after_readable_note(
    result_view_case, tmp_path
):
    library, lesson = result_view_case
    note = lesson / "validated_document.md"
    note.write_text("# Nota per il professionista\n\nConclusione da rivedere.")
    audit = lesson / "validation_package.md"
    audit.write_text('# Registro completo\n\n{"claim_count": 13, "status": "pending"}')
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"] = [
        {"path": audit.name, "sha256": digest(audit)},
        {"path": note.name, "sha256": digest(note)},
    ]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "answer-results"

    library.results("variance-analysis", "it", lesson, "execution.json", destination)

    page = (destination / "course.html").read_text()
    assert page.index("Nota per il professionista") < page.index(
        "Registro completo della verifica"
    )
    audit_section = page.split("Registro completo della verifica</h2>")[1].split(
        "</section>"
    )[0]
    assert "<details><summary>" in audit_section
    assert "claim_count" in audit_section.split("<details>")[1]
    assert "<details open" not in audit_section
    assert (
        destination / "outputs/02-validation_package.md"
    ).read_bytes() == audit.read_bytes()


def test_result_native_disclosures_render_without_activating_html(
    result_view_case, tmp_path
):
    library, lesson = result_view_case
    note = lesson / "studio_checklist.md"
    note.write_text(
        "# Piano della pratica\n\nIl nuovo indirizzo resta da verificare.\n\n"
        "<details><summary>Riferimenti tecnici</summary>\n\n"
        "- Fonte: CASE-OPERATION\n\n</details>\n\n"
        "<details open ontoggle=alert(1)><summary>Unsafe</summary>\n"
        "<script>alert(1)</script>\n\n"
        "<!-- Review Handoff -->\n\nProssimo controllo.\n"
    )
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"] = [{"path": note.name, "sha256": digest(note)}]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "disclosure-result"
    library.results("variance-analysis", "it", lesson, "execution.json", destination)
    page = (destination / "course.html").read_text()
    assert "<details><summary>Riferimenti tecnici</summary><ul>" in page
    assert "<li>Fonte: CASE-OPERATION</li></ul></details>" in page
    assert "<details open ontoggle" not in page
    assert "<script>alert(1)</script>" not in page
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in page
    assert "<p>&lt;!-- Review Handoff --&gt;</p>" not in page
    assert (
        destination / "outputs/01-studio_checklist.md"
    ).read_bytes() == note.read_bytes()


def test_result_citations_open_only_recorded_inputs(result_view_case, tmp_path):
    library, lesson = result_view_case
    source = lesson / "source.txt"
    previous = lesson / "previous.json"
    previous.write_text('{"assessment": "<script>literal</script>"}')
    answer = lesson / "outputs/summary.md"
    answer.parent.mkdir()
    answer.write_text(
        f"[Read source](<{source}>)\n\n"
        "[Relative source](../source.txt)\n\n"
        "[Previous assessment](<../previous.json>)\n\n"
        "[Unrecorded relative](../private.txt)\n\n"
        "[Remote](https://example.invalid/private)\n\n"
        "[Unrecorded](/outside/secret.txt)\n\n"
        "[Script](javascript:alert)\n"
    )
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["inputs"].append({"path": "previous.json", "sha256": digest(previous)})
    receipt["outputs"] = [
        {"path": answer.relative_to(lesson).as_posix(), "sha256": digest(answer)}
    ]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "citation-view"

    library.results("variance-analysis", "it", lesson, "execution.json", destination)

    page = (destination / "course.html").read_text()
    assert "<a href='source-01.html'>Read source</a>" in page
    assert "<a href='source-01.html'>Relative source</a>" in page
    assert "<a href='source-02.html'>Previous assessment</a>" in page
    previous_view = (destination / "source-02.html").read_text()
    assert "&lt;script&gt;literal&lt;/script&gt;" in previous_view
    assert "<script>" not in previous_view
    assert (
        destination / "outputs/source-02-previous.json"
    ).read_bytes() == previous.read_bytes()
    assert "href='../private.txt'" not in page
    assert "href='https:" not in page
    assert "href='/outside/" not in page
    assert "href='javascript:" not in page
    assert "Remote (https://example.invalid/private)" in page
    assert "<p>[Remote](https://example.invalid/private)</p>" not in page
    view = (destination / "source-01.html").read_text()
    assert source.read_text() in view
    assert "href='course.html'" in view
    assert (
        destination / "outputs/source-01-source.txt"
    ).read_bytes() == source.read_bytes()
    assert (
        json.loads((destination / "result-view.json").read_text())["outputs"]
        == receipt["outputs"]
    )


def test_result_sources_remain_reachable_without_inline_file_citations(
    result_view_case, tmp_path
):
    library, lesson = result_view_case
    source = lesson / "source.txt"
    answer = lesson / "faq.md"
    answer.write_text(
        "# A short FAQ\n\n## Where do I start?\n\nRead the selected source.\n"
    )
    before = answer.read_bytes()
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"] = [{"path": answer.name, "sha256": digest(answer)}]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "faq-results"

    library.results("variance-analysis", "it", lesson, "execution.json", destination)

    page = (destination / "course.html").read_text()
    assert "<a href='source-01.html'>source.txt</a>" in page
    assert source.read_text() in (destination / "source-01.html").read_text()
    assert "href='course.html'" in (destination / "source-01.html").read_text()
    assert (destination / "outputs/01-faq.md").read_bytes() == before
    assert answer.read_bytes() == before


@pytest.mark.parametrize("extension", ["xlsx", "docx", "pdf"])
def test_result_document_sources_open_from_verified_input_index(
    result_view_case, tmp_path, extension
):
    library, lesson = result_view_case
    source = lesson / f"evidence.{extension}"
    if extension == "xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["Finanza", 50000])
        workbook.save(source)
    elif extension == "docx":
        from docx import Document

        document = Document()
        document.add_paragraph("Finanza 50000")
        document.save(source)
    else:
        import pymupdf

        with pymupdf.open() as document:
            page = document.new_page()
            page.insert_text((72, 72), "Finanza 50000")
            document.save(source)
    before = source.read_bytes()
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["inputs"] = [{"path": source.name, "sha256": digest(source)}]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "document-sources"

    library.results("variance-analysis", "it", lesson, "execution.json", destination)

    assert (
        f"href='source-01.html'>{source.name}</a>"
        in (destination / "course.html").read_text()
    )
    page = (destination / "source-01.html").read_text()
    assert "href='course.html'" in page
    if extension == "pdf":
        assert "source-01-page-1.png" in page
        assert (destination / "source-01-page-1.png").is_file()
    else:
        assert "Finanza" in page and "50000" in page
    assert (destination / f"outputs/source-01-{source.name}").read_bytes() == before
    assert source.read_bytes() == before


def test_client_email_is_readable_and_keeps_exact_plain_text(
    result_view_case, tmp_path
):
    library, lesson = result_view_case
    email = lesson / "client_email.txt"
    email.write_text(
        "Oggetto: Come usare la FAQ\n\nGentile cliente,\n\n"
        "Leggi la FAQ prima della revisione. <script>literal</script>\n\nCordiali saluti\n"
    )
    before = email.read_bytes()
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"] = [{"path": email.name, "sha256": digest(email)}]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "email-results"

    library.results("variance-analysis", "it", lesson, "execution.json", destination)

    page = (destination / "course.html").read_text()
    assert "<h2>Email al cliente</h2>" in page
    assert "<p>Oggetto: Come usare la FAQ</p>" in page
    assert "<p>Gentile cliente,</p>" in page
    assert "&lt;script&gt;literal&lt;/script&gt;" in page and "<script>" not in page
    assert (destination / "outputs/01-client_email.txt").read_bytes() == before
    assert email.read_bytes() == before


def test_native_dossier_audit_links_do_not_block_reading_or_lose_navigation(
    result_view_case, tmp_path
):
    library, lesson = result_view_case
    original = lesson / "review_dossier.html"
    payload = b"""<!doctype html><html><head><meta charset="utf-8"></head><body>
    <main><nav><a href="#sources">Fonti</a></nav>
    <details id="condition"><summary>Condizione</summary><p>Software</p></details>
    <section id="sources"><h2>Fonti</h2><a href="#condition">Torna alla condizione</a></section>
    <footer><a href="review_dossier.md">Registro</a><a href="case_intake.json">Dati</a></footer>
    </main></body></html>"""
    original.write_bytes(payload)
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"] = [{"path": original.name, "sha256": digest(original)}]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "result-view"
    library.results("variance-analysis", "it", lesson, "execution.json", destination)
    view = (destination / "result-01.html").read_text()
    assert "href='#html-result-sources'" in view
    assert "id='html-result-sources'" in view
    assert "href='#html-result-condition'" in view
    assert "id='html-result-condition'" in view
    assert "<summary>Condizione</summary>" in view
    assert "<span>Registro</span><span>Dati</span>" in view
    assert 'href="case_intake.json"' not in view
    assert (destination / "outputs/01-review_dossier.html").read_bytes() == payload
    assert original.read_bytes() == payload


@pytest.mark.parametrize(
    "href",
    [
        "../secret.json",
        "/private/file.json",
        "//example.com/data.json",
        "javascript:alert(1)",
        "data:text/html,x",
    ],
)
def test_passive_html_never_activates_nonlocal_audit_links(href):
    from courseware.html_view import passive_html_body

    with pytest.raises(ValueError, match="Unsupported result format"):
        passive_html_body(f'<a href="{href}">Dati</a>')


def test_accounts_preview_metadata_and_keyboard_regions_remain_passive():
    from courseware.html_view import passive_html_body

    rendered = passive_html_body(
        '<html lang="it"><head><meta name="color-scheme" content="light"></head>'
        '<body><a href="#main-content">Vai al bilancio</a>'
        '<main id="main-content" tabindex="-1" data-output-language="it">'
        '<h2 id="statements">Prospetti</h2>'
        '<div role="region" aria-labelledby="statements" tabindex="0">'
        "<table><tr><td>Attivo</td><td>27.000,00</td></tr></table>"
        "</div></main></body></html>"
    )

    assert "#html-result-main-content" in rendered
    assert "tabindex='0'" in rendered and "tabindex='-1'" in rendered
    assert "aria-labelledby='html-result-statements'" in rendered
    assert "27.000,00" in rendered
    assert "<meta" not in rendered and "data-output-language" not in rendered


def test_html_reading_preserves_safe_semantic_spacing_only():
    from courseware.html_view import passive_html_body

    rendered = passive_html_body(
        '<div class="meta overlay"><span>Da verificare</span><span>Proposta</span></div><summary><span class="code">REQ-1</span>Condizione</summary><p class="narrative" style="position:fixed">Primo.\n\nSecondo.</p>'
    )
    assert "class='html-meta'" in rendered
    assert "class='html-code'" in rendered
    assert "class='html-narrative'" in rendered
    assert "Primo.\n\nSecondo." in rendered
    assert "overlay" not in rendered and "position" not in rendered


def test_grant_html_is_primary_and_technical_markdown_stays_available(
    result_view_case, tmp_path
):
    library, lesson = result_view_case
    html = lesson / "review_dossier.html"
    html.write_text("<h1>Dossier per un bando</h1><p>Risultato leggibile.</p>")
    technical = lesson / "review_dossier.md"
    technical.write_text(
        "# Registro tecnico\n\n| Hash | Stato |\n| --- | --- |\n| abc | proposed |\n"
    )
    receipt = json.loads((lesson / "execution.json").read_text())
    receipt["outputs"] = [
        {"path": p.name, "sha256": digest(p)} for p in (html, technical)
    ]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "result-view"
    library.results("variance-analysis", "it", lesson, "execution.json", destination)
    page = (destination / "course.html").read_text()
    assert "<h2>Dossier da rivedere</h2>" in page
    assert "<h2>Registro dettagliato del dossier</h2>" in page
    assert "<table>" not in page
    assert "<details><summary>Leggi il file completo</summary>" in page
    assert (
        destination / "outputs/02-review_dossier.md"
    ).read_bytes() == technical.read_bytes()


def test_result_sources_distinguish_demo_and_practice_versions(
    result_view_case, tmp_path
):
    library, lesson = result_view_case
    receipt = json.loads((lesson / "execution.json").read_text())
    demo = lesson / "demo-original/package/products.csv"
    practice = lesson / "practice-updated/package/products.csv"
    demo.parent.mkdir(parents=True)
    practice.parent.mkdir(parents=True)
    demo.write_text("product,group\nCardigan,New\n")
    practice.write_text("product,group\nSweater,New\n")
    report = lesson / "comparison.html"
    report.write_text("<h1>Updated assortment</h1><p>New arrivals changed.</p>")
    receipt["inputs"] = [
        {"path": str(demo.relative_to(lesson)), "sha256": digest(demo)},
        {"path": str(practice.relative_to(lesson)), "sha256": digest(practice)},
    ]
    receipt["outputs"] = [{"path": report.name, "sha256": digest(report)}]
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "versioned-sources"
    library.results("variance-analysis", "it", lesson, "execution.json", destination)
    index = (destination / "course.html").read_text()
    reader = (destination / "result-01.html").read_text()
    assert "href='source-01.html'>Dimostrazione — products.csv</a>" in index
    assert "href='source-02.html'>Esercizio — products.csv</a>" in index
    assert "href='source-01.html'>Dimostrazione — products.csv</a>" in reader
    assert "href='source-02.html'>Esercizio — products.csv</a>" in reader
    assert (
        "<h1>Dimostrazione — products.csv</h1>"
        in (destination / "source-01.html").read_text()
    )
    assert "<td>Cardigan</td>" in (destination / "source-01.html").read_text()
    assert (
        "<h1>Esercizio — products.csv</h1>"
        in (destination / "source-02.html").read_text()
    )
    assert "<td>Sweater</td>" in (destination / "source-02.html").read_text()
    assert (
        destination / "outputs/source-01-products.csv"
    ).read_bytes() == demo.read_bytes()
    assert (
        destination / "outputs/source-02-products.csv"
    ).read_bytes() == practice.read_bytes()


@pytest.mark.parametrize("workflow", ["deck-correction", "html-deck"])
def test_native_deck_result_keeps_navigation_original_and_download(
    result_view_case, tmp_path, workflow
):
    _, lesson = result_view_case
    source = (
        ROOT
        / "plugins/clara/assets/courses/deck-correction/files/input/original-deck-it.html"
    )
    original = lesson / "original.html"
    output = lesson / "index.html"
    original.write_bytes(source.read_bytes())
    corrected = source.read_bytes().replace(b"Sara", b"Elena")
    output.write_bytes(corrected)
    package = lesson / "presentation.zip"
    package.write_bytes(b"synthetic download; never extracted")
    validation = lesson / "static-validation.json"
    write_json(
        validation,
        {
            "schema_version": "clara.html_deck_validation.v1",
            "result": "pass",
            "profile": "stage",
            "input": {"sha256": digest(output)},
        },
    )

    def record(path):
        return {"path": path.name, "sha256": digest(path)}

    receipt = json.loads((lesson / "execution.json").read_text())
    receipt.update(
        product="clara",
        workflow_id=workflow,
        skill_sha256=digest(ROOT / f"plugins/clara/skills/{workflow}/SKILL.md"),
        inputs=[record(original)],
        native_records=[record(validation)],
        outputs=[record(output), record(package)],
    )
    write_json(lesson / "execution.json", receipt)
    destination = tmp_path / "deck-result"
    library = CourseLibrary(ROOT / "plugins/clara", {workflow})
    library.results(workflow, "it", lesson, "execution.json", destination)
    index = (destination / "course.html").read_text()
    assert "href='decks/result-01.html'" in index
    assert "href='decks/source-01.html'" in index
    assert (destination / "decks/result-01.html").read_bytes() == corrected
    assert (destination / "decks/source-01.html").read_bytes() == source.read_bytes()
    assert (
        destination / "outputs/02-presentation.zip"
    ).read_bytes() == package.read_bytes()
    from courseware.deck_view import deck_csp

    policy = deck_csp(output.read_text())
    assert "sandbox allow-scripts;" in policy and "allow-same-origin" not in policy
    assert "default-src 'none'" in policy and "script-src 'sha256-" in policy
    assert "connect-src *" not in policy and "form-action 'none'" in policy
    output.write_text(
        output.read_text().replace(
            "</body>", "<script>alert('unexpected')</script></body>"
        )
    )
    receipt["outputs"][0]["sha256"] = digest(output)
    write_json(
        validation,
        {
            "schema_version": "clara.html_deck_validation.v1",
            "result": "pass",
            "profile": "stage",
            "input": {"sha256": digest(output)},
        },
    )
    receipt["native_records"][0]["sha256"] = digest(validation)
    write_json(lesson / "execution.json", receipt)
    with pytest.raises(ValueError, match="scripts differ"):
        library.results(
            workflow, "it", lesson, "execution.json", tmp_path / "unsafe-deck"
        )
    assert not (tmp_path / "unsafe-deck").exists()


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_public_transformation_course_discloses_cowork_unavailability(
    language: str,
) -> None:
    page = (
        ROOT / "static/shared/courses/vera/trasformazione" / language / "course.html"
    ).read_text(encoding="utf-8")
    startup_copy = json.loads(
        (ROOT / "scripts/cowork_teaching/public-start.json").read_text(encoding="utf-8")
    )[language]

    assert 'id="course-start-request"' in page
    assert startup_copy["unavailable"] in page
    assert 'id="cowork-start-request"' not in page
    assert 'data-copy-target="cowork-start-request"' not in page
