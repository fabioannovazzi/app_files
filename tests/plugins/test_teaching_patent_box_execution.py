"""Execute packaged fictional Patent Box kits; attest no learner or real approval."""

from __future__ import annotations

import copy
import csv
import hashlib
import json
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from tests.plugins._teaching_release import record_native_check
from tests.plugins.test_patent_box_workflow import load_module
from tests.plugins.test_teaching_kit_execution import (
    ROOT,
    _bound_case,
    _complete_teaching_case,
    _read,
)

__all__ = []
WORKFLOW = "patent-box-review"
MODULE = ROOT / "plugins" / WORKFLOW


def _proposal(run: dict, workflow, language: str, phase: str) -> tuple[dict, dict]:
    """Interpret only this fixed authored case in a test-only semantic fixture."""
    context = Path(run["context_path"])
    session = workflow.initialize(context, as_of="2026-09-23", demo=True)
    inputs = {row["description"]: row for row in session["inputs"]}
    ledger = inputs["ledger-demo.csv" if phase == "demo" else "ledger-practice.csv"]
    note = inputs[
        (
            f"evidence-{language}.md"
            if phase == "demo"
            else f"evidence-practice-{language}.md"
        )
    ]
    source = inputs["synthetic-rules.txt"]
    table = workflow.inspect_ledger(
        context,
        evidence_id=ledger["evidence_id"],
        options={
            "format": "CSV",
            "sheet": None,
            "header_row": 1,
            "first_row": 2,
            "last_row": 3 if phase == "demo" else 4,
            "delimiter": ",",
            "encoding": "utf-8",
            "pdf_extraction": None,
        },
    )
    with Path(ledger["selected_path"]).open() as stream:
        rows = list(csv.DictReader(stream))
    plan = {
        "schema_version": "1.0",
        "purpose": "Test-only interpretation of fictional teaching ledger",
        "mappings": [
            {
                "table_id": table["table_id"],
                "amount_column": "F",
                "currency_column": None,
                "currency_constant": "EUR",
                "row_key_column": "B",
                "decimal_separator": ".",
                "thousands_separator": None,
                "parentheses_negative": False,
                "economic_key_columns": ["B"],
                "mapping_rationale": "Note section 3; exact ledger header",
            }
        ],
        "control_totals": [
            {
                "table_id": table["table_id"],
                "currency": "EUR",
                "value": "110000.00" if phase == "demo" else "130000.00",
                "evidence_id": note["evidence_id"],
                "locator": "section 3" if phase == "demo" else "section 5",
                "source_cell": None,
            }
        ],
        "fx_rates": [],
        "costs": [
            {
                **{
                    key: row[key]
                    for key in (
                        "cost_id",
                        "period_id",
                        "account",
                        "category",
                        "income_max",
                        "irap_max",
                    )
                },
                "components": [{"row_ref": native["row_ref"], "fx_rate_id": None}],
                "rationale": "Candidate bases from note and source rows; synthetic only",
                "evidence_ids": [ledger["evidence_id"], note["evidence_id"]],
            }
            for row, native in zip(rows, table["rows"], strict=True)
        ],
        "excluded_rows": [],
        "non_data_rows": [],
        "duplicate_reviews": [],
        "population_duplicate_review": {
            "assessment": "Separate fictional economic costs GL.1/GL.2/GL.3; no duplicate declared",
            "evidence_ids": [ledger["evidence_id"], note["evidence_id"]],
        },
        "rounding_policy": "HALF_UP_AT_NORMALIZED_COST_TOTAL",
    }
    normalized = workflow.normalize_ledger(context, plan)
    case = _read(MODULE / "examples/case.ordinary.json")
    case.update(
        case_id=session["run_id"],
        demo=True,
        ledger_control_total=normalized["ledger_control_total"],
        costs=normalized["costs"],
        evidence=[
            {k: row[k] for k in ("evidence_id", "path", "sha256", "description")}
            for row in session["inputs"]
        ],
        controls=[],
    )
    case["ips"][0].update(name="Faro", controls=[])
    base = case["allocations"][0]
    case["allocations"] = []
    for index, row in enumerate(rows, start=1):
        allocation = copy.deepcopy(base)
        allocation.update(
            allocation_id=f"A{index}",
            cost_id=row["cost_id"],
            income_amount=row["income_max"],
            irap_amount=row["irap_max"],
            controls=[],
            allocation_method="Note sections 2/3; synthetic candidate allocation",
        )
        case["allocations"].append(allocation)
    rules = _read(MODULE / "examples/rules.demo.json")
    rules["sources"][0].update(
        snapshot_evidence_id=source["evidence_id"], snapshot_sha256=source["sha256"]
    )
    controls = []
    groups = {
        "case": [
            "SUBJECT",
            "EXCLUSIONS",
            "OPTION",
            "TRANSITION",
            "DUPLICATES",
            "EXTRAORDINARY",
            "DECLARATION",
            "ADVERSARIAL",
            "SOURCES",
        ],
        "ip:IP.SOFT": ["IP.SOFTWARE", "RIGHTS", "USE", "ACTIVITY"],
    }
    for allocation in case["allocations"]:
        groups[f"allocation:{allocation['allocation_id']}"] = [
            "COST",
            "PERSONNEL",
            "LINKAGE",
            "ALLOCATION",
            "INCENTIVES",
            "RD.CREDIT",
        ]
    for scope, names in groups.items():
        for name in names:
            status = "PASS"
            conclusion = "Test-only interpretation of declared fictional facts, note sections 1-4"
            if scope == "allocation:A2" and name == "LINKAGE":
                status, conclusion = (
                    "FAIL",
                    "Note section 2: ordinary maintenance, not original development",
                )
            if scope == "allocation:A3" and name == "LINKAGE":
                status, conclusion = (
                    "NOT_TESTED",
                    "Note section 5: missing activity/project linkage; request records",
                )
            controls.append(
                {
                    "key": f"{scope}/PB.{name}",
                    "status": status,
                    "conclusion": conclusion,
                    "evidence_ids": [note["evidence_id"], ledger["evidence_id"]],
                    "source_ids": ["DEMO.SOURCE"],
                }
            )
    proposal = {
        "case": case,
        "rules": rules,
        "controls": controls,
        "normalization_digest": normalized["normalization_digest"],
        "narratives": [
            {
                "section": "A",
                "text": Path(note["selected_path"]).read_text(),
                "evidence_ids": [note["evidence_id"]],
                "locator": "sections 1-5",
            },
            {
                "section": "B",
                "text": "GL.1 / GL.2" + (" / GL.3" if phase == "practice" else ""),
                "evidence_ids": [ledger["evidence_id"]],
                "locator": "CSV rows 2-4",
            },
        ],
    }
    return proposal, session


def execute_kit(tmp_path: Path, monkeypatch, language: str, phase: str):
    """Run all real mechanics with clearly synthetic semantic review inputs."""
    run = _bound_case(
        tmp_path, monkeypatch, WORKFLOW, WORKFLOW, phase, language=language
    )
    workflow = load_module(
        "patent_box_teaching", MODULE / "scripts/patent_box_workflow.py"
    )
    context = Path(run["context_path"])
    proposal, session = _proposal(run, workflow, language, phase)
    review = workflow.propose(context, proposal)
    digest = review["proposal_digest"]
    readable = Path(review["review_path"])
    assert "GL.1" in readable.read_text()
    with pytest.raises((ValueError, FileNotFoundError)):
        workflow.calculate_draft(context, digest=digest)
    workflow.review(
        context,
        digest=digest,
        reviewer="SYNTHETIC_TEACHING_TEST",
        confirmation_ref=f"automated-kit-{language}-{phase}",
        confirmed=True,
        synthetic=True,
    )
    actual = workflow.calculate_draft(context, digest=digest)
    folder = Path(actual["output_dir"])
    result = actual["result"]
    assert result["bases"]["INCLUDED"] == {"income": "100000.00", "irap": "80000.00"}
    assert result["bases"]["EXCLUDED"] == {"income": "10000.00", "irap": "10000.00"}
    assert result["bases"]["SUSPENDED"] == (
        {"income": "20000.00", "irap": "16000.00"}
        if phase == "practice"
        else {"income": "0.00", "irap": "0.00"}
    )
    assert result["additional_deduction"] == {"income": "110000.00", "irap": "88000.00"}
    assert result["tax_saving"] is None
    assert _read(folder / "dossier_document.json")["status"] == "DRAFT_UNSIGNED"
    document = Document(folder / "fascicolo_A_B.docx")
    word_text = "\n".join(
        [p.text for p in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    pdf_text = "\n".join(
        page.extract_text() for page in PdfReader(folder / "fascicolo_A_B.pdf").pages
    )
    for text in (word_text, pdf_text):
        assert "BOZZA SINTETICA DA RIVEDERE" in text
        assert "Faro" in text and "110000.00" in text
        assert "GL.2" in text
        if phase == "practice":
            assert "GL.3" in text and "NOT_TESTED" in text
    assert (
        proposal["case"]["costs"][1]["ledger_row_key"]
        in (folder / "cost_reconciliation.csv").read_text()
    )
    before = hashlib.sha256((folder / "result.json").read_bytes()).hexdigest()
    with pytest.raises(FileExistsError):
        workflow.calculate_draft(context, digest=digest)
    assert hashlib.sha256((folder / "result.json").read_bytes()).hexdigest() == before
    _complete_teaching_case(run, tmp_path / "case")
    assert (Path(run["output_dir"]) / "model_data_report.md").is_file()
    return run, folder, session


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
@pytest.mark.parametrize("phase", ["demo", "practice"])
def test_patent_box_kit_executes_native_dossier(
    tmp_path, monkeypatch, language, phase, record_property
):
    execute_kit(tmp_path, monkeypatch, language, phase)
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow=WORKFLOW,
        language=language,
        phase=phase,
    )


def test_practice_preserves_demo_and_binds_distinct_inputs(tmp_path, monkeypatch):
    _, demo, demo_session = execute_kit(tmp_path / "demo", monkeypatch, "it", "demo")
    original = {
        p.name: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in demo.iterdir()
        if p.is_file()
    }
    _, practice, practice_session = execute_kit(
        tmp_path / "practice", monkeypatch, "it", "practice"
    )
    assert demo_session["run_id"] != practice_session["run_id"]
    assert demo != practice
    assert original == {
        p.name: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in demo.iterdir()
        if p.is_file()
    }


def test_learning_session_registers_patent_box_without_manufacturing_progress(
    tmp_path, monkeypatch
):
    monkeypatch.syspath_prepend(str(ROOT / "plugins/_shared/vendor/modules"))
    scripts = ROOT / "plugins/vera/scripts"
    monkeypatch.syspath_prepend(str(scripts))
    from courseware.library import CourseLibrary
    from local_teaching import TeachingStore

    store = TeachingStore(tmp_path / "profile")
    state = store.begin(
        {
            "workflow_id": WORKFLOW,
            "title": "Patent Box",
            "goal": "Review a synthetic dossier",
            "mode": "together",
            "pair": {
                "teacher_thread_id": "synthetic-teacher",
                "worker_thread_id": "synthetic-worker",
            },
        }
    )["session"]
    handoff = store.worker("synthetic-worker", WORKFLOW, state["worker_token"])
    assert handoff["assignment"] == "tutorial" and handoff["local_only"] is True
    assert (
        Path(handoff["workflow_contract"]["skill_path"])
        == ROOT / "plugins/vera/skills/patent-box-review/SKILL.md"
    )
    CourseLibrary(ROOT / "plugins/vera", {WORKFLOW}).render(
        WORKFLOW, "it", Path(state["directory"]) / "kit"
    )
    assert "demo" not in store.status()["session"]
    assert "practice" not in store.status()["session"]
    with pytest.raises(ValueError):
        store.change(
            "finish",
            state["revision"],
            {
                "confirmed_by_user": True,
                "understanding": "Synthetic fixture cannot complete an unexecuted lesson",
            },
        )


def test_course_binds_engine_controls_and_historical_demo_rules():
    course = _read(ROOT / "plugins/vera/assets/courses/patent-box-review/course.json")
    sources = {row["repository_path"]: row for row in course["sources"]}
    for name in [
        "patent_box/engine.py",
        "config/control_catalog.json",
        "examples/rules.demo.json",
    ]:
        source = "plugins/patent-box-review/" + name
        assert (
            sources[source]["sha256"]
            == hashlib.sha256((ROOT / source).read_bytes()).hexdigest()
        )
