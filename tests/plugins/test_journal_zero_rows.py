from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl
import pytest

from scripts.validate_plugin_review_contract import validate_contract
from tests.plugins.test_journal_sampling_plugin import (
    _approve_suggested_recipe,
    _save_workbook,
    _write_registration_journal_pdf,
    load_core,
)


def _zero_row_source(root: Path, source_kind: str, *, only_zero: bool = False) -> Path:
    """Write actual source cells, including explicit zero and blank metadata."""
    suffix = "csv" if source_kind in {"csv", "amount"} else "xlsx"
    source = root / f"journal.{suffix}"
    postings = (
        []
        if only_zero
        else [
            ["2025-01-01", "1000", "Debit", "100.00", None],
            ["2025-01-01", "2000", "Credit", None, "100.00"],
        ]
    )
    postings.extend(
        [
            ["2025-01-01", "3000", "Zero debit", "0.0", None],
            ["2025-01-01", "4000", "Zero credit", None, "0.00"],
            ["2025-01-01", "5000", "Zero both", "0.00", "0.00"],
            ["2025-01-01", "META", "Account-only metadata", None, None],
        ]
    )
    rows = [["Date", "Account", "Description", "Debit", "Credit"], *postings]
    if source_kind == "amount":
        rows = [
            ["Date", "Account", "Description", "Amount"],
            *[
                [
                    *row[:3],
                    (
                        row[3]
                        if row[3] is not None
                        else ("-" + row[4] if row[4] is not None else None)
                    ),
                ]
                for row in postings
            ],
        ]
    if source_kind == "print":
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for col, label in {
            2: "Data Reg.",
            5: "Descrizione",
            6: "Conto",
            7: "Descrizione Conto",
            8: "Dare (EUR)",
            11: "Avere (EUR)",
            14: "Nr. Reg",
        }.items():
            sheet.cell(7, col, label)
        sheet.cell(8, 2, "01/01/2025")
        sheet.cell(8, 5, "Registration")
        sheet.cell(8, 14, 1)
        for index, row in enumerate(postings, start=9):
            for col, value in ((6, row[1]), (7, row[2]), (9, row[3]), (12, row[4])):
                sheet.cell(index, col, value)
        workbook.save(source)
    elif suffix == "csv":
        with source.open("w", newline="", encoding="utf-8") as handle:
            csv.writer(handle).writerows(rows)
    else:
        _save_workbook(source, rows)
    return source


@pytest.mark.parametrize("source_kind", ["csv", "xlsx", "amount", "print"])
def test_zero_source_rows_are_retained_but_excluded_from_sample(
    tmp_path: Path, source_kind: str
) -> None:
    core = load_core()
    source = _zero_row_source(tmp_path, source_kind)
    normalized_dir = tmp_path / "normalized"
    core.inspect_path(source, normalized_dir)
    recipe = normalized_dir / "suggested_recipe.json"
    _approve_suggested_recipe(recipe)
    normalized = core.normalize_path(source, normalized_dir, recipe)

    result = core.run_sample(
        normalized_dir / "normalized_journal.csv", tmp_path / "sample", size=25
    )

    assert normalized.diagnostics["population_status"] == "complete"
    assert normalized.frame.height == 5
    assert normalized.frame.get_column("amount_abs").to_list() == [
        "100",
        "100",
        "0",
        "0",
        "0",
    ]
    assert normalized.frame.get_column("reported_increment")[2] == "0.1"
    assert normalized.diagnostics["files"][0]["rejected_row_count"] == 0
    assert normalized.diagnostics["files"][0]["excluded_non_monetary_row_count"] == 1
    assert sorted(result.frame.get_column("amount_signed").to_list()) == ["-100", "100"]
    assert result.audit["population_size_before_filters"] == 5
    assert result.audit["zero_value_rows_excluded"] == 3
    assert result.audit["population_size_after_filters"] == 2
    assert result.audit["sampling_status"] == "sampled"
    assert (
        core.validate_sample_assurance(tmp_path / "sample")["reproducibility"][
            "zero_value_rows_excluded"
        ]
        == 3
    )


@pytest.mark.parametrize("method", ["random", "systematic", "stratified", "mus"])
@pytest.mark.parametrize("min_abs", [None, "0.01"])
def test_complete_pdf_with_zero_row_samples_nonzero_rows_and_replays(
    tmp_path: Path, method: str, min_abs: str | None
) -> None:
    core = load_core()
    source = tmp_path / "journal.pdf"
    layout = _write_registration_journal_pdf(source, zero_value="0,00")
    recipe = tmp_path / "layout.json"
    recipe.write_text(
        json.dumps(
            {"files": {source.name: {"parser": "registration_pdf", "layout": layout}}}
        )
    )
    output = tmp_path / "normalized"
    core.inspect_path(source, output, recipe)
    _approve_suggested_recipe(output / "suggested_recipe.json")
    normalized = core.normalize_path(source, output, output / "suggested_recipe.json")

    result = core.run_sample(
        output / "normalized_journal.csv",
        tmp_path / "sample",
        method=method,
        size=25,
        min_abs=min_abs,
        language="it",
    )

    assert normalized.diagnostics["population_status"] == "complete"
    assert normalized.frame.get_column("amount_signed").to_list() == [
        "100",
        "-100",
        "0",
    ]
    assert result.audit["zero_value_rows_excluded"] == 1
    assert result.frame.get_column("amount_abs").to_list() == ["100", "100"]
    review = json.loads((tmp_path / "sample/review_payload.json").read_text())
    assert review["summary"]["zero_value_rows_excluded"] == 1
    assert (
        "1 righe a importo zero" in (tmp_path / "sample/review_handoff.md").read_text()
    )
    assert (
        core.validate_sample_assurance(tmp_path / "sample")["reproducibility"][
            "zero_value_rows_excluded"
        ]
        == 1
    )


@pytest.mark.parametrize("method", ["random", "systematic", "stratified", "mus"])
def test_all_zero_population_returns_explicit_empty_sample(
    tmp_path: Path, method: str
) -> None:
    core = load_core()
    source = _zero_row_source(tmp_path, "csv", only_zero=True)
    output = tmp_path / "normalized"
    core.inspect_path(source, output)
    _approve_suggested_recipe(output / "suggested_recipe.json")
    normalized = core.normalize_path(source, output, output / "suggested_recipe.json")

    result = core.run_sample(
        output / "normalized_journal.csv", tmp_path / "sample", method=method
    )

    assert normalized.diagnostics["population_status"] == "complete"
    assert normalized.frame.height == 3
    assert result.frame.is_empty()
    assert result.audit["sampling_status"] == "no_eligible_rows"
    assert result.audit["zero_value_rows_excluded"] == 3
    assert (tmp_path / "sample/journal_sample.xlsx").is_file()
    assert "No eligible rows" in (tmp_path / "sample/review_handoff.md").read_text()
    assert (
        core.validate_sample_assurance(tmp_path / "sample")["reproducibility"][
            "selected_prepared_row_numbers"
        ]
        == []
    )
    contract = validate_contract(
        tmp_path / "sample",
        strict_output_paths=True,
        strict_output_content=True,
    )
    assert contract.ok, contract.as_dict()


@pytest.mark.parametrize("bad_value", ["unreadable", "1.2.3"])
def test_explicit_zero_does_not_hide_unreadable_opposite_amount(
    tmp_path: Path, bad_value: str
) -> None:
    core = load_core()
    source = tmp_path / "journal.xlsx"
    _save_workbook(
        source,
        [
            ["Date", "Account", "Description", "Debit", "Credit"],
            ["2025-01-01", "1000", "Bad amount", "0.00", bad_value],
        ],
    )

    result = core.inspect_path(source, tmp_path / "out")

    assert result.files[0]["qualification_status"] == "unsupported_source_layout"


def test_numeric_excel_zero_retains_its_displayed_precision(tmp_path: Path) -> None:
    core = load_core()
    source = tmp_path / "journal.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Date", "Account", "Description", "Debit", "Credit"])
    sheet.append(["2025-01-01", "1000", "Printed zero", 0, None])
    sheet["D2"].number_format = "0.00"
    workbook.save(source)
    output = tmp_path / "normalized"
    core.inspect_path(source, output)
    _approve_suggested_recipe(output / "suggested_recipe.json")

    result = core.normalize_path(source, output, output / "suggested_recipe.json")

    assert result.diagnostics["population_status"] == "complete"
    assert result.frame.get_column("amount_signed").to_list() == ["0"]
    assert result.frame.get_column("reported_increment").to_list() == ["0.01"]
