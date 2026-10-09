"""Residual income, changing capital, clean surplus and review dependencies."""

from __future__ import annotations

import hashlib
import shutil
from decimal import Decimal
from pathlib import Path

import pytest

from tests.plugins.test_business_valuation import (
    FIXTURE,
    build_valuation,
    case_data,
    compile_html,
    planning_case,
    prepare_archive_run,
    read_json,
    restore_imports,
    review,
    run_valuation,
    write_package,
)


def residual_case(base: dict | None = None) -> dict:
    """Append an independent synthetic common-equity model to the existing case."""
    case = case_data() if base is None else base
    case["sources"].append(
        {
            **case["sources"][0],
            "id": "residual-basis",
            "path": "residual-income.txt",
            "sha256": hashlib.sha256(
                (FIXTURE / "residual-income.txt").read_bytes()
            ).hexdigest(),
            "description": "Fictional book equity, clean-surplus income and owner transactions",
        }
    )
    template = case["inputs"][0]
    case["inputs"].extend(
        [
            {
                **template,
                "id": key,
                "value": value,
                "unit": unit,
                "source_ids": ["residual-basis"],
                "description": f"Fictional {key}",
                "locator": "residual-income.txt: declared amounts",
            }
            for key, value, unit in (
                ("book0", "1000", "EUR"),
                ("book1", "1120", "EUR"),
                ("book2", "1230", "EUR"),
                ("ri-income1", "150", "EUR"),
                ("ri-income2", "180", "EUR"),
                ("dividend1", "50", "EUR"),
                ("dividend2", "70", "EUR"),
                ("contribution1", "20", "EUR"),
                ("contribution2", "0", "EUR"),
                ("terminal-equity", "1500", "EUR"),
                ("second-rate", "0.2", "ratio"),
            )
        ]
    )
    case["mandate"].update(valuation_date="2024-12-31", information_cutoff="2024-12-31")
    case["methods"].append(
        {
            "id": "residual",
            "kind": "RESIDUAL_INCOME_EQUITY",
            "selected": True,
            "rationale": "Synthetic clean-surplus model with changing common-equity capital.",
            "inputs": {
                "book_equity": ["book0", "book1", "book2"],
                "incomes": ["ri-income1", "ri-income2"],
                "distributions": ["dividend1", "dividend2"],
                "contributions": ["contribution1", "contribution2"],
                "terminal_equity_value": "terminal-equity",
            },
            "limitations": [
                "Synthetic accounting assumptions and terminal value require professional review."
            ],
            "timing": {
                "valuation_date": "2024-12-31",
                "period_end_dates": ["2025-12-31", "2026-12-31"],
                "cash_flow_timing": "end_period",
                "day_count": "ACT/365F",
                "rate_compounding": "effective_annual",
                "rate_model": "flat",
                "rate_ids": ["rate"],
                "rationale": "All owner transactions and period results occur at explicit period ends.",
            },
            "residual_basis": {
                "accounting_basis": "Consistent synthetic common-equity perimeter and accounting basis.",
                "clean_surplus_adjustments": "No OCI or other non-owner movements assumed; income captures all changes.",
                "owner_transactions": "Distributions and contributions are independent nonnegative amounts at period end.",
                "terminal_equity_basis": "Independent synthetic 1500 EUR total equity, not continuing residual income.",
                "capital_cost_basis": "Explicit equity cost; same factors determine interval charges and discounting.",
                "source_ids": ["residual-basis"],
                "locator": "residual-income.txt: assumptions",
                "status": "confirmed",
            },
        }
    )
    return case


@pytest.mark.parametrize(
    ("timing", "expected"),
    [
        ({}, 1324.793388429752),
        (
            {
                "rate_model": "spot_curve",
                "rate_ids": ["rate", "second-rate"],
                "terminal_discount_rate": "second-rate",
            },
            1117.550505050505,
        ),
        (
            {"rate_model": "forward_curve", "rate_ids": ["rate", "second-rate"]},
            1216.666666666667,
        ),
        ({"rate_compounding": "continuous"}, 1312.5524048735103),
        ({"period_end_dates": ["2025-06-30", "2025-12-31"]}, 1455.8878108735564),
    ],
)
def test_residual_income_matches_independent_owner_cashflows(timing, expected):
    case = residual_case()
    case["methods"][-1]["timing"].update(timing)
    result = build_valuation(case, FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert float(values["residual/value"]) == pytest.approx(expected, rel=1e-12)
    assert Decimal(values["residual/owner_cashflow_difference"]).copy_abs() < Decimal(
        "0.000000000000000000000000000001"
    )
    assert result["methods"][-1]["value_type"] == "equity"


def test_equity_charge_uses_beginning_book_and_subtracts_terminal_book_once():
    result = build_valuation(residual_case(), FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert Decimal(values["residual/residual/1/equity_charge"]) == 100
    assert Decimal(values["residual/residual/2/equity_charge"]) == 112
    assert Decimal(values["residual/continuing_residual"]) == 270
    assert result["methods"][-1]["equity_id"] == "residual/value"


@pytest.mark.parametrize(
    ("key", "value"),
    [
        ("book1", "1121"),
        ("ri-income2", "181"),
        ("contribution1", "0"),
        ("dividend2", "71"),
    ],
)
def test_unreconciled_clean_surplus_blocks_only_its_method(key, value):
    case = residual_case()
    next(row for row in case["inputs"] if row["id"] == key)["value"] = value
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"
    assert "does not reconcile" in result["methods"][-1]["reason"]
    assert result["methods"][0]["status"] == "ready_for_professional_review"


@pytest.mark.parametrize(
    ("key", "value"), [("book1", None), ("dividend1", "-50"), ("contribution1", "-20")]
)
def test_missing_or_signed_owner_flows_are_not_silently_reinterpreted(key, value):
    case = residual_case()
    next(row for row in case["inputs"] if row["id"] == key)["value"] = value
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"


@pytest.mark.parametrize(
    ("key", "value"),
    [
        ("book_equity", ["book0", "book1"]),
        ("distributions", ["dividend1"]),
        ("contributions", []),
        ("incomes", ["ri-income1"]),
        ("terminal_equity_value", "rate"),
        ("terminal_equity_value", "ri-income2"),
    ],
)
def test_residual_method_rejects_wrong_schedule_or_terminal_basis(key, value):
    case = residual_case()
    case["methods"][-1]["inputs"][key] = value
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"


@pytest.mark.parametrize("field", ["timing", "residual_basis"])
def test_residual_method_requires_explicit_dated_professional_basis(field):
    case = residual_case()
    del case["methods"][-1][field]
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"


@pytest.mark.parametrize(
    "timing",
    [
        {"cash_flow_timing": "mid_period"},
        {
            "rate_model": "spot_curve",
            "rate_ids": ["rate", "second-rate"],
            "terminal_discount_rate": "rate",
        },
    ],
)
def test_residual_income_rejects_inconsistent_discount_conventions(timing):
    case = residual_case()
    case["methods"][-1]["timing"].update(timing)
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"


def test_residual_method_never_deducts_debt_again():
    case = residual_case()
    case["methods"][-1]["bridge"] = case["methods"][0]["bridge"]
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"


@pytest.mark.parametrize(
    "field",
    [
        "accounting_basis",
        "clean_surplus_adjustments",
        "owner_transactions",
        "terminal_equity_basis",
        "capital_cost_basis",
        "locator",
    ],
)
def test_changed_residual_basis_invalidates_its_review_and_preserves_others(field):
    case = residual_case()
    before = build_valuation(case, FIXTURE)
    case["methods"][-1]["review"] = review(before["methods"][-1]["dependency_sha256"])
    case["methods"][0]["review"] = review(before["methods"][0]["dependency_sha256"])
    case["methods"][-1]["residual_basis"][field] += " Revised."
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["stale_review"] is True
    assert result["methods"][0]["status"] == "accepted_workpaper"


@pytest.mark.parametrize(
    ("field", "value", "status"),
    [
        ("status", "proposed", "partial"),
        ("source_ids", ["missing"], "blocked"),
        ("clean_surplus_adjustments", "", "blocked"),
    ],
)
def test_unconfirmed_or_missing_residual_basis_cannot_approve_a_method(
    field, value, status
):
    case = residual_case()
    case["methods"][-1]["residual_basis"][field] = value
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == status


def test_negative_income_equity_and_terminal_values_remain_signed():
    case = residual_case()
    changed = {
        "book0": "-1000",
        "book1": "-1200",
        "book2": "-1450",
        "ri-income1": "-170",
        "ri-income2": "-180",
        "terminal-equity": "-1400",
    }
    for row in case["inputs"]:
        if row["id"] in changed:
            row["value"] = changed[row["id"]]
    result = build_valuation(case, FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert result["methods"][-1]["status"] == "ready_for_professional_review"
    assert float(values["residual/value"]) == pytest.approx(
        -1071.900826446281, rel=1e-12
    )


def test_plan_fcff_cannot_be_relabelled_as_clean_surplus_income(tmp_path):
    case = residual_case(planning_case(tmp_path))
    shutil.copyfile(FIXTURE / "residual-income.txt", tmp_path / "residual-income.txt")
    case["mandate"].update(valuation_date="2026-12-31", information_cutoff="2026-12-31")
    case["methods"][-1]["timing"].update(
        valuation_date="2026-12-31", period_end_dates=["2027-12-31", "2028-12-31"]
    )
    case["methods"][-1]["inputs"]["incomes"] = ["flow", "flow"]
    result = build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")
    assert result["methods"][-1]["status"] == "blocked"
    assert "cannot be relabelled" in result["methods"][-1]["reason"]


def test_residual_export_links_all_clean_surplus_cells_and_escapes_basis(tmp_path):
    from openpyxl import load_workbook

    case = residual_case()
    case["methods"][-1]["residual_basis"][
        "accounting_basis"
    ] = '=HYPERLINK("https://example.invalid")'
    case["methods"][-1]["residual_basis"][
        "owner_transactions"
    ] = '<script>alert("source")</script>'
    result = build_valuation(case, FIXTURE)
    output = tmp_path / "report"
    write_package(result, FIXTURE, output)
    workbook = load_workbook(output / "valuation_workbook.xlsx", data_only=False)
    assert workbook["Base residuale"]["C2"].data_type == "s"
    assert workbook["Clean surplus"]["C2"].data_type == "f"
    assert workbook["Clean surplus"].max_row == 3
    html = compile_html(result)
    assert "<script>alert" not in html
    assert "&lt;script&gt;alert" in html
    assert (
        "perpetuità annuale"
        not in html.split("Reddito residuale del capitale proprio", 1)[1]
    )
    assert read_json(output / "valuation.json") == result


def test_residual_method_runs_with_exact_archive_receipts(tmp_path, record_property):
    case_path, context = prepare_archive_run(
        tmp_path,
        supplied_case=residual_case(),
        source_files=[FIXTURE / "evidence.txt", FIXTURE / "residual-income.txt"],
    )
    output = run_valuation.run_case(case_path, context)
    report = read_json(Path(output["output_dir"]) / "valuation.json")
    assert report["methods"][-1]["status"] == "ready_for_professional_review"
    assert report["methods"][-1]["residual_basis"]["status"] == "confirmed"
    record_property("residual_income_output", output["output_dir"])


def test_explicit_zero_cost_preserves_all_net_owner_payments():
    case = residual_case()
    next(row for row in case["inputs"] if row["id"] == "rate")["value"] = "0"
    result = build_valuation(case, FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert Decimal(values["residual/value"]) == 1600
    assert Decimal(values["residual/owner_cashflow_difference"]) == 0


def test_terminal_book_value_is_an_explicit_zero_continuing_excess():
    case = residual_case()
    case["methods"][-1]["inputs"]["terminal_equity_value"] = "book2"
    result = build_valuation(case, FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert Decimal(values["residual/continuing_residual"]) == 0
    assert float(values["residual/value"]) == pytest.approx(
        1101.652892561983, rel=1e-12
    )


def test_unreviewed_residual_basis_source_prevents_method_acceptance():
    case = residual_case()
    case["sources"][-1]["status"] = "unverified"
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "partial"
    assert "residual-basis" in result["methods"][-1]["source_ids"]
