"""Actual Word/PDF export content and immutable archive-artifact acceptance."""

from __future__ import annotations

import io
import json
import sys
from pathlib import Path
from types import ModuleType
from typing import Any

import pytest
from docx import Document
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_patent_box_casebook import proposal_with_book
from test_patent_box_workflow import model_proposal, reviewed, running_case, workflow


def prepared(tmp_path: Path, workflow: ModuleType) -> tuple[dict[str, Any], str]:
    run, proposal = proposal_with_book(tmp_path, workflow)
    proposal["casebook"]["template"] = {
        "version": "Synthetic outline version 1",
        "source_id": "DEMO.SOURCE",
        "evidence_id": "E0001",
        "required_paragraphs": [
            {"paragraph_id": "A.1", "section": "A", "title": "Fatti della pratica"},
            {"paragraph_id": "B.1", "section": "B", "title": "Spese da documentare"},
        ],
        "simplifications": [],
    }
    return run, reviewed(run, workflow, proposal)


def all_docx_text(path: Path) -> str:
    doc = Document(path)
    return "\n".join(
        [p.text for p in doc.paragraphs]
        + [c.text for t in doc.tables for r in t.rows for c in r.cells]
    )


def test_calculation_exports_editable_docx_pdf_and_manifest_with_open_paragraphs(
    tmp_path: Path, workflow: ModuleType
) -> None:
    run, digest = prepared(tmp_path, workflow)

    output = workflow.calculate_draft(run["context"], digest=digest)

    folder = Path(output["output_dir"])
    docx_text = all_docx_text(folder / "fascicolo_A_B.docx")
    pdf_text = "\n".join(
        p.extract_text() for p in PdfReader(folder / "fascicolo_A_B.pdf").pages
    )
    assert "DA COMPLETARE: paragrafo B.1" in docx_text
    assert "DA COMPLETARE: paragrafo B.1" in pdf_text
    assert "F1" in docx_text and "synthetic.txt" in docx_text
    assert "110000.00" in docx_text and "110000.00" in pdf_text
    assert "BOZZA SINTETICA DA RIVEDERE" in docx_text
    assert "BOZZA SINTETICA DA RIVEDERE" in pdf_text
    manifest = json.loads((folder / "manifest.json").read_text())
    assert manifest["artifacts"]["fascicolo_A_B.docx"] == workflow.file_hash(
        folder / "fascicolo_A_B.docx"
    )
    assert manifest["artifacts"]["fascicolo_A_B.pdf"] == workflow.file_hash(
        folder / "fascicolo_A_B.pdf"
    )
    document = json.loads((folder / "dossier_document.json").read_text())
    assert document["status"] == "DRAFT_UNSIGNED"
    assert document["proposal_digest"] == digest


def test_pdf_treats_markup_in_source_as_literal_text(
    tmp_path: Path, workflow: ModuleType
) -> None:
    document = {
        "title": "Fascicolo Patent Box",
        "blocks": [
            {"kind": "paragraph", "text": "Proposta <b>non verificata</b> & citazione"}
        ],
    }

    raw = workflow.render_pdf(document)

    text = PdfReader(io.BytesIO(raw)).pages[0].extract_text()
    assert "<b>non verificata</b>" in text
    assert "& citazione" in text


def test_pdf_unsupported_script_fails_instead_of_losing_glyphs(
    workflow: ModuleType,
) -> None:
    document = {
        "title": "Fascicolo Patent Box",
        "blocks": [{"kind": "paragraph", "text": "漢字"}],
    }

    with pytest.raises(workflow.ContractError, match="unsupported"):
        workflow.render_pdf(document)


def test_document_model_rejects_mismatched_review_digest(
    tmp_path: Path, workflow: ModuleType
) -> None:
    run, digest = prepared(tmp_path, workflow)
    output = workflow.calculate_draft(run["context"], digest=digest)
    proposal = json.loads((run["output"] / f"proposal_{digest}.json").read_text())
    decision = json.loads((run["output"] / f"decision_{digest}.json").read_text())
    decision["proposal_digest"] = "0" * 64

    with pytest.raises(workflow.ContractError, match="different reviewed versions"):
        workflow.compose_dossier(proposal, output["result"], decision)


def test_pdf_multipage_long_source_text_keeps_final_evidence_and_page_numbers(
    workflow: ModuleType,
) -> None:
    document = {
        "title": "Fascicolo Patent Box",
        "blocks": [
            {"kind": "heading", "text": "Riferimenti", "level": 1},
            {
                "kind": "paragraph",
                "text": (
                    "Dettaglio contabile con prova selezionata e riserva esplicita. "
                    * 650
                ),
            },
            {"kind": "paragraph", "text": "EVIDENZA FINALE 7654321"},
        ],
    }

    raw = workflow.render_pdf(document)

    pdf = PdfReader(io.BytesIO(raw))
    assert len(pdf.pages) > 2
    assert "EVIDENZA FINALE 7654321" in pdf.pages[-1].extract_text()
    assert f"Pagina {len(pdf.pages)}" in pdf.pages[-1].extract_text()
