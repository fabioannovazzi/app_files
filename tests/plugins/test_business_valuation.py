"""Numerical, evidence and portable-run acceptance for PMI valuation workpapers."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import sys
from calendar import monthrange
from copy import deepcopy
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "plugins/business-valuation/scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(ROOT / "plugins/_shared/vendor/modules"))

import run_valuation
import valuation_case
import valuation_claims
import valuation_engine
import valuation_normalization
import valuation_report
import valuation_schema
from valuation_case import build_valuation, digest, read_json
from valuation_engine import ValuationError, decimal, evaluate
from valuation_report import compile_html, write_package, write_workbook

FIXTURE = ROOT / "tests/fixtures/business_valuation"


@pytest.fixture(autouse=True)
def restore_imports(monkeypatch: pytest.MonkeyPatch) -> None:
    """Keep isolated component imports available after repository test cleanup."""
    monkeypatch.syspath_prepend(str(SCRIPTS))
    monkeypatch.syspath_prepend(str(ROOT / "plugins/_shared/vendor/modules"))
    for module in (
        run_valuation,
        valuation_case,
        valuation_claims,
        valuation_engine,
        valuation_normalization,
        valuation_report,
        valuation_schema,
    ):
        monkeypatch.setitem(sys.modules, module.__name__, module)


def case_data() -> dict:
    return read_json(FIXTURE / "case.json")


@pytest.mark.parametrize(
    ("path", "value"),
    [
        (("currency",), "eur"),
        (("synthetic",), "true"),
        (("mandate", "valuation_date"), "2026-02-30"),
        (("mandate", "information_cutoff"), "20261231"),
        (("sources", 0, "allowed_audiences"), "internal"),
        (("sources", 0, "sha256"), "a" * 63),
        (("inputs", 0, "value"), "private-value-not-a-number"),
        (("inputs", 0, "value"), "100\n"),
        (("inputs", 0, "value"), 100),
        (("inputs", 0, "source_ids"), [["evidence"]]),
        (("limitations",), [" "]),
        (("undeclared-field",), "private-value-not-a-number"),
    ],
)
def test_schema_rejects_bad_case_before_nested_source_access(
    path: tuple, value: object, tmp_path: Path
) -> None:
    case = case_data()
    parent = case
    for key in path[:-1]:
        parent = parent[key]
    parent[path[-1]] = value

    with pytest.raises(ValuationError, match="schema") as error:
        build_valuation(case, tmp_path / "source-root-does-not-exist")

    assert "private-value-not-a-number" not in str(error.value)


def test_source_path_discovery_applies_schema_before_receipt_expansion(
    tmp_path: Path,
) -> None:
    case = case_data()
    del case["sources"][0]["path"]
    with pytest.raises(ValuationError, match="at /sources/0: schema required"):
        valuation_case.source_paths(case, tmp_path / "missing")


def test_schema_bounds_case_record_arrays() -> None:
    case = case_data()
    case["sources"] = case["sources"] * 2001
    with pytest.raises(ValuationError, match="schema maxItems"):
        valuation_schema.validate_case(case)


@pytest.mark.parametrize("payload", [None, [], {"flows": ["flow"]}])
def test_schema_preserves_malformed_method_and_sensitivity_as_blocked(
    payload: object,
) -> None:
    case = case_data()
    case["methods"][0]["inputs"] = payload
    case["sensitivity"] = [
        dict(
            id="scenario",
            method_id="fcff",
            discount_rate="rate",
            terminal_rate="rate",
            terminal_growth="growth",
        )
    ]

    result = build_valuation(case, FIXTURE)

    assert result["methods"][0]["status"] == "blocked"
    assert "schema" in result["methods"][0]["reason"]
    assert result["sensitivity"][0]["status"] == "blocked"
    assert result["methods"][1]["status"] == "ready_for_professional_review"


def test_schema_keeps_excluded_specialist_method_without_a_numeric_payload() -> None:
    case = case_data()
    case["methods"][0].update(kind="SPECIALIST_METHOD", selected=False, inputs=None)
    result = build_valuation(case, FIXTURE)
    assert result["methods"][0]["status"] == "excluded"
    assert result["methods"][1]["status"] == "ready_for_professional_review"


def test_schema_does_not_grant_acceptance_to_an_incomplete_attestation() -> None:
    case = case_data()
    case["methods"][0]["review"] = {"decision": "accepted"}
    result = build_valuation(case, FIXTURE)
    assert result["methods"][0]["status"] == "ready_for_professional_review"
    assert result["methods"][0]["stale_review"] is True


def test_published_schema_uses_only_internal_references() -> None:
    from jsonschema import Draft202012Validator

    schema = read_json(SCRIPTS.parent / "references/valuation-case.schema.json")
    Draft202012Validator.check_schema(schema)
    references = re.findall(r'"\$ref": "([^"]+)"', json.dumps(schema))
    assert references
    assert set(ref.split("/")[0] for ref in references) == {"#"}


def review(dependency: str) -> dict:
    return {
        "dependency_sha256": dependency,
        "decision": "accepted",
        "reviewer": "Synthetic reviewer",
        "reviewed_at": "2026-09-29T17:00:00+02:00",
    }


def rights_case() -> dict:
    """Supply a fictional 40% interest without treating it as a valuation rule."""
    case = case_data()
    case["mandate"].update(
        subject="Partecipazione ordinaria sintetica",
        rights="40% di una classe inventata; nessuna rettifica automatica del valore",
    )
    case["mandate_details"]["subject_type"]["value"] = "equity_interest"
    case["mandate_details"]["subject_type"].update(
        source_ids=["rights"], locator="rights.txt: separate interest variant"
    )
    case["sources"].append(
        {
            **case["sources"][0],
            "id": "rights",
            "path": "rights.txt",
            "description": "Fictional rights variant",
            "sha256": hashlib.sha256((FIXTURE / "rights.txt").read_bytes()).hexdigest(),
        }
    )
    case["inputs"].append(
        {
            **case["inputs"][2],
            "id": "ownership",
            "value": "0.4",
            "description": "Invented 40% ordinary share interest",
            "source_ids": ["rights"],
            "locator": "rights.txt: percentage and denominator",
        }
    )
    case["mandate_details"]["interests"] = [
        {
            "id": "ordinary",
            "description": "Classe ordinaria sintetica",
            "ownership_input_id": "ownership",
            "ownership_basis": "40% della classe ordinaria; denominatore dichiarato nella prova",
            "economic_rights": "Diritti economici inventati, nessuna preferenza stimata",
            "administrative_rights": "Voto e poteri da giudicare separatamente",
            "statutes": "Riferimento sintetico all'articolo 1",
            "agreements": "Nessun patto nella prova inventata",
            "restrictions": "Vincolo sintetico di trasferimento; nessuno sconto applicato",
            "thresholds": "Nessuna soglia quantificata nella prova",
            "source_ids": ["rights"],
            "locator": "rights.txt: fictional instrument terms",
            "status": "confirmed",
        }
    ]
    return case


def test_missing_structured_mandate_keeps_calculations_but_marks_case_partial() -> None:
    case = case_data()
    del case["mandate_details"]

    result = build_valuation(case, FIXTURE)

    assert result["status"] == "partial"
    assert result["mandate_assessment"]["details"] is None
    assert result["mandate_assessment"]["issues"]
    assert Decimal(
        next(x["value"] for x in result["calculations"] if x["id"] == "fcff/equity")
    ) == Decimal("750")


@pytest.mark.parametrize(
    ("field", "value"),
    [("value", None), ("status", "proposed"), ("source_ids", []), ("locator", None)],
)
def test_mandate_unknown_or_unconfirmed_field_prevents_complete_case(
    field: str, value: object
) -> None:
    case = case_data()
    case["mandate_details"]["commissioning_party"][field] = value

    result = build_valuation(case, FIXTURE)

    assert result["mandate_assessment"]["status"] == "partial"
    assert result["status"] == "partial"
    assert result["methods"][0]["status"] == "ready_for_professional_review"


@pytest.mark.parametrize("field", ["engagement_date", "report_date"])
def test_mandate_date_schema_rejects_invalid_date_before_source_reads(
    field: str, tmp_path: Path
) -> None:
    case = case_data()
    case["mandate_details"][field]["value"] = "2026-02-30"
    with pytest.raises(ValuationError, match="schema"):
        build_valuation(case, tmp_path / "missing")


def test_mandate_unresolved_evidence_is_rejected() -> None:
    case = case_data()
    case["mandate_details"]["conflicts"]["source_ids"] = ["absent"]
    with pytest.raises(ValuationError, match="Unresolved mandate evidence"):
        build_valuation(case, FIXTURE)


def test_unreviewed_mandate_source_prevents_method_acceptance() -> None:
    case = case_data()
    case["sources"][0]["status"] = "unverified"
    result = build_valuation(case, FIXTURE)
    assert "Mandate source review pending" in result["mandate_assessment"]["issues"]
    assert result["mandate_assessment"]["status"] == "partial"
    assert result["methods"][0]["status"] == "partial"


def test_rights_and_percentage_do_not_multiply_equity_or_infer_a_discount() -> None:
    result = build_valuation(rights_case(), FIXTURE)
    assert result["mandate_assessment"]["input_ids"] == ["ownership"]
    assert result["mandate_assessment"]["status"] == "ready_for_professional_review"
    assert Decimal(
        next(x["value"] for x in result["calculations"] if x["id"] == "fcff/equity")
    ) == Decimal("750")
    assert "ownership" not in result["methods"][0]["input_ids"]


@pytest.mark.parametrize("amount", [None, "-0.1", "1.1"])
def test_unknown_or_out_of_range_ownership_remains_partial_without_changing_values(
    amount: str | None,
) -> None:
    case = rights_case()
    case["inputs"][-1]["value"] = amount
    result = build_valuation(case, FIXTURE)
    assert result["mandate_assessment"]["status"] == "partial"
    assert result["methods"][0]["status"] == "ready_for_professional_review"


@pytest.mark.parametrize(
    "missing", ["interests", "ownership_input_id", "economic_rights"]
)
def test_equity_interest_requires_evidenced_rights_and_ownership(missing: str) -> None:
    case = rights_case()
    if missing == "interests":
        case["mandate_details"]["interests"] = []
    else:
        case["mandate_details"]["interests"][0][missing] = None
    result = build_valuation(case, FIXTURE)
    assert result["mandate_assessment"]["status"] == "partial"


def test_specific_right_can_have_explicit_nonpercentage_basis() -> None:
    case = rights_case()
    case["mandate_details"]["subject_type"]["value"] = "specific_right"
    case["mandate_details"]["interests"][0].update(
        ownership_input_id=None,
        ownership_basis="Diritto contrattuale sintetico non espresso come quota di capitale",
    )
    result = build_valuation(case, FIXTURE)
    assert result["mandate_assessment"]["status"] == "ready_for_professional_review"
    assert result["mandate_assessment"]["input_ids"] == []


@pytest.mark.parametrize("invalid", ["duplicate", "source", "input", "unit"])
def test_rights_references_and_units_are_explicit(invalid: str) -> None:
    case = rights_case()
    item = case["mandate_details"]["interests"][0]
    if invalid == "duplicate":
        case["mandate_details"]["interests"].append(deepcopy(item))
    elif invalid == "source":
        item["source_ids"] = ["absent"]
    elif invalid == "input":
        item["ownership_input_id"] = "absent"
    else:
        case["inputs"][-1]["unit"] = "EUR"
    with pytest.raises(
        ValuationError,
        match="Duplicate mandate interest|Unresolved rights evidence|Unresolved ownership input|Ownership requires",
    ):
        build_valuation(case, FIXTURE)


@pytest.mark.parametrize("field", ["commissioning_party", "conflicts", "report_date"])
def test_changed_mandate_revokes_review_and_dependent_method(field: str) -> None:
    case = case_data()
    initial = build_valuation(case, FIXTURE)
    case["mandate_details"]["review"] = review(
        initial["mandate_assessment"]["dependency_sha256"]
    )
    case["methods"][0]["review"] = review(initial["methods"][0]["dependency_sha256"])
    case["mandate_details"][field]["value"] = (
        "2027-02-01" if field == "report_date" else "Changed evidenced choice"
    )
    result = build_valuation(case, FIXTURE)
    assert result["mandate_assessment"]["stale_review"] is True
    assert result["methods"][0]["stale_review"] is True
    assert result["methods"][0]["status"] == "ready_for_professional_review"


def test_review_only_change_preserves_arithmetic_review_but_expires_conclusion() -> (
    None
):
    case = reviewed_claimed_case()
    case["conclusion"] = {
        "text": "Synthetic conclusion",
        "method_ids": ["fcff"],
        "review": None,
    }
    before = build_valuation(case, FIXTURE)
    case["conclusion"]["review"] = review(before["conclusion"]["dependency_sha256"])
    case["mandate_details"]["review"]["reviewer"] = "A different synthetic reviewer"

    result = build_valuation(case, FIXTURE)

    assert result["mandate_assessment"]["status"] == "accepted_workpaper"
    assert result["methods"][0]["status"] == "accepted_workpaper"
    assert result["conclusion"]["status"] == "draft"


def test_mandate_change_invalidates_direct_adjustment_and_claim_attestations() -> None:
    case = reviewed_normalized_case()
    case["claims"] = claimed_case()["claims"]
    prepared = build_valuation(case, FIXTURE)
    case["claims"][0]["review"] = review(prepared["claims"][0]["dependency_sha256"])
    case["mandate_details"]["participant_perspective"][
        "value"
    ] = "Different explicit perspective"
    result = build_valuation(case, FIXTURE)
    assert result["normalizations"][0]["adjustments"][0]["stale_review"] is True
    assert result["claims"][0]["stale_review"] is True


def test_unrelated_numeric_input_does_not_invalidate_mandate_review() -> None:
    case = case_data()
    initial = build_valuation(case, FIXTURE)
    case["mandate_details"]["review"] = review(
        initial["mandate_assessment"]["dependency_sha256"]
    )
    case["inputs"][0]["value"] = "110"
    result = build_valuation(case, FIXTURE)
    assert result["mandate_assessment"]["status"] == "accepted_workpaper"


def test_complete_case_requires_mandate_review_in_addition_to_method_and_conclusion() -> (
    None
):
    case = case_data()
    initial = build_valuation(case, FIXTURE)
    case["methods"][0]["review"] = review(initial["methods"][0]["dependency_sha256"])
    case["conclusion"] = {
        "text": "Synthetic conclusion",
        "method_ids": ["fcff"],
        "review": None,
    }
    draft = build_valuation(case, FIXTURE)
    case["conclusion"]["review"] = review(draft["conclusion"]["dependency_sha256"])
    result = build_valuation(case, FIXTURE)
    assert result["conclusion"]["status"] == "draft"
    assert result["status"] == "ready_for_professional_review"


def test_mandate_exports_preserve_rights_and_escape_untrusted_text(
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    case = rights_case()
    case["mandate_details"]["interests"][0][
        "agreements"
    ] = '=1+1 <script>alert("x")</script>'
    result = build_valuation(case, FIXTURE)
    output = tmp_path / "exports"
    write_package(result, FIXTURE, output)
    workbook = load_workbook(output / "valuation_workbook.xlsx")
    values = [
        cell
        for row in workbook["Incarico"]
        for cell in row
        if cell.value == '=1+1 <script>alert("x")</script>'
    ]
    assert len(values) == 1
    assert values[0].data_type == "s"
    assert "&lt;script&gt;" in (output / "valuation_report.html").read_text()
    assert (
        read_json(output / "mandate.json")["data"]["assessment"]
        == result["mandate_assessment"]
    )
    assert (
        read_json(output / "professional_review.json")["data"]["mandate"]
        == result["mandate_assessment"]
    )


def normalized_case() -> dict:
    """Reconcile 90 reported plus 20 minus 10 to an unchanged 100-unit income."""
    case = case_data()
    income = next(row for row in case["inputs"] if row["id"] == "income")
    case["inputs"].extend(
        [
            {**income, "id": "reported-income", "value": "90"},
            {**income, "id": "signed-addition", "value": "20"},
            {**income, "id": "signed-deduction", "value": "-10"},
        ]
    )
    entry = {
        "id": "one-off",
        "amount_input": "signed-addition",
        "reason": "Synthetic positive adjustment for arithmetic testing.",
        "accounting_check": "90 plus 20 minus 10 equals 100 in the fictional schedule.",
        "economic_rationale": "Invented test case; no economic appropriateness asserted.",
        "tax_treatment": "Amounts are supplied after tax for this income line; no automatic tax factor.",
        "reversibility": "Synthetic nonrecurring event; professional review pending.",
        "source_ids": ["evidence"],
        "locator": "evidence.txt, synthetic schedule",
    }
    case["normalizations"] = [
        {
            "id": "income-2026",
            "year": 2026,
            "line": "Reddito netto sostenibile",
            "reported_input": "reported-income",
            "adjusted_input": "income",
            "adjustments": [
                entry,
                {
                    **entry,
                    "id": "recurring-cost",
                    "amount_input": "signed-deduction",
                    "reason": "Synthetic negative adjustment; do not choose only increases.",
                },
            ],
        }
    ]
    return case


def reviewed_normalized_case() -> dict:
    """Prepare explicit synthetic journal attestations, then method attestations."""
    case = normalized_case()
    initial = build_valuation(case, FIXTURE)
    for entry, row in zip(
        case["normalizations"][0]["adjustments"],
        initial["normalizations"][0]["adjustments"],
    ):
        entry["review"] = review(row["dependency_sha256"])
    prepared = build_valuation(case, FIXTURE)
    for method, row in zip(case["methods"], prepared["methods"]):
        method["review"] = review(row["dependency_sha256"])
    return case


def test_signed_normalization_reconciles_without_approving_adjustments() -> None:
    result = build_valuation(normalized_case(), FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert values["normalization/income-2026/adjustments"] == "10"
    assert values["normalization/income-2026/adjusted"] == "100"
    assert values["normalization/income-2026/difference"] == "0"
    assert values["income-method/value"] == "1000"
    assert result["normalizations"][0]["status"] == "partial"
    assert result["methods"][2]["status"] == "partial"
    assert result["methods"][3]["status"] == "ready_for_professional_review"


def test_reviewed_adjustments_allow_separately_reviewed_method() -> None:
    result = build_valuation(reviewed_normalized_case(), FIXTURE)
    assert result["normalizations"][0]["status"] == "accepted_workpaper"
    assert result["methods"][2]["status"] == "accepted_workpaper"
    assert result["methods"][2]["normalization_ids"] == ["income-2026"]
    assert set(result["methods"][2]["input_ids"]) >= {
        "reported-income",
        "signed-addition",
        "signed-deduction",
    }
    assert result["piv_conformity"] == "not_assessed"


@pytest.mark.parametrize(
    "field",
    [
        "reason",
        "tax_treatment",
        "reversibility",
        "accounting_check",
        "economic_rationale",
    ],
)
def test_changed_adjustment_invalidates_only_dependent_reviews(field: str) -> None:
    case = reviewed_normalized_case()
    case["normalizations"][0]["adjustments"][0][field] = "Revised synthetic explanation"
    result = build_valuation(case, FIXTURE)
    assert result["normalizations"][0]["adjustments"][0]["stale_review"] is True
    assert result["methods"][2]["stale_review"] is True
    assert result["methods"][2]["status"] == "partial"
    assert result["methods"][3]["status"] == "accepted_workpaper"


@pytest.mark.parametrize(
    ("ref", "value"),
    [("income", "999"), ("signed-addition", None), ("reported-income", None)],
)
def test_unreconciled_normalization_blocks_dependent_methods_only(
    ref: str, value: str | None
) -> None:
    case = normalized_case()
    next(row for row in case["inputs"] if row["id"] == ref)["value"] = value
    result = build_valuation(case, FIXTURE)
    assert result["normalizations"][0]["status"] == "blocked"
    assert result["methods"][2]["status"] == "blocked"
    assert result["methods"][4]["status"] == "blocked"
    assert result["methods"][3]["status"] == "ready_for_professional_review"


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("reason", ""),
        ("tax_treatment", ""),
        ("reversibility", ""),
        ("source_ids", ["unknown"]),
        ("amount_input", "income"),
        ("amount_input", "unknown"),
    ],
)
def test_incomplete_adjustment_contract_is_rejected(field: str, value) -> None:
    case = normalized_case()
    case["normalizations"][0]["adjustments"][0][field] = value
    with pytest.raises(ValuationError):
        build_valuation(case, FIXTURE)


def test_duplicate_adjustment_amount_is_not_summed_twice() -> None:
    case = normalized_case()
    case["normalizations"][0]["adjustments"][1]["amount_input"] = "signed-addition"
    with pytest.raises(ValuationError, match="counted twice"):
        build_valuation(case, FIXTURE)


def test_normalization_does_not_reuse_adjusted_output_as_reported_input() -> None:
    case = normalized_case()
    case["normalizations"][0]["reported_input"] = "income"
    with pytest.raises(ValuationError, match="independent"):
        build_valuation(case, FIXTURE)


def test_normalization_calculation_and_workbook_preserve_full_formula_lineage(
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    result = build_valuation(normalized_case(), FIXTURE)
    output = tmp_path / "normalized"
    write_package(result, FIXTURE, output)
    rows = {row["id"]: index for index, row in enumerate(result["calculations"], 2)}
    workbook = load_workbook(output / "valuation_workbook.xlsx")
    normalized = f"B{rows['normalization/income-2026/adjusted']}"
    dependent = f"B{rows['income-method/input/income']}"
    assert workbook["Calcoli"][dependent].value == f"=SUM({normalized})"
    assert workbook["Rettifiche"]["E2"].value == f"='Calcoli'!{normalized}"
    assert "Reversibilità:" in workbook["Rettifiche"]["F3"].value
    assert (
        "Rettificato calcolato: 100,00 EUR"
        in (output / "valuation_report.md").read_text()
    )
    assert "Trattamento fiscale:" in (output / "valuation_report.md").read_text()


def test_missing_normalization_amount_exports_a_visible_block_without_zero(
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    case = normalized_case()
    next(row for row in case["inputs"] if row["id"] == "signed-addition")[
        "value"
    ] = None
    result = build_valuation(case, FIXTURE)
    output = tmp_path / "missing-adjustment"
    write_package(result, FIXTURE, output)
    workbook = load_workbook(output / "valuation_workbook.xlsx")
    assert workbook["Rettifiche"]["E3"].value is None
    assert (
        "Rettifica one-off: non disponibile"
        in (output / "valuation_report.md").read_text()
    )
    assert result["methods"][2]["status"] == "blocked"
    assert result["methods"][3]["status"] == "ready_for_professional_review"


def test_untrusted_adjustment_prose_remains_literal_in_exports(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    case = normalized_case()
    case["normalizations"][0]["adjustments"][0][
        "reason"
    ] = '=HYPERLINK("https://example.org")<script>x</script>'
    result = build_valuation(case, FIXTURE)
    path = tmp_path / "literal-journal.xlsx"
    write_workbook(path, result)
    assert load_workbook(path)["Rettifiche"]["F3"].data_type == "s"
    assert "&lt;script&gt;x&lt;/script&gt;" in compile_html(result)


def claimed_case() -> dict:
    """Author one conditional claim; the numeric binding never validates its prose."""
    case = case_data()
    case["claims"] = [
        {
            "id": "fcff-equity",
            "kind": "hypothesis",
            "text": "Alle ipotesi sintetiche dichiarate il DCF FCFF indica un equity di 750 EUR.",
            "location": "Confronto dei metodi",
            "basis": "Raccordo esplicito dal valore operativo, debito e cassa della prova.",
            "source_ids": [],
            "input_ids": [],
            "calculation_ids": ["fcff/equity"],
            "method_ids": [],
            "limitations": ["Dati inventati; non è una valutazione professionale."],
            "values": [
                {"calculation_id": "fcff/equity", "value": "750", "unit": "EUR"}
            ],
        }
    ]
    return case


def reviewed_claimed_case() -> dict:
    """Prepare explicit synthetic method and claim attestations separately."""
    case = claimed_case()
    initial = build_valuation(case, FIXTURE)
    case["mandate_details"]["review"] = review(
        initial["mandate_assessment"]["dependency_sha256"]
    )
    case["methods"][0]["review"] = review(initial["methods"][0]["dependency_sha256"])
    prepared = build_valuation(case, FIXTURE)
    case["claims"][0]["review"] = review(prepared["claims"][0]["dependency_sha256"])
    return case


def test_claim_resolves_full_numeric_evidence_without_approving_prose() -> None:
    result = build_valuation(claimed_case(), FIXTURE)
    claim = result["claims"][0]
    assert Decimal(claim["resolved_values"][0]["value"]) == Decimal("750")
    assert claim["resolved_values"][0]["calculation_id"] == "fcff/equity"
    assert claim["resolved_values"][0]["unit"] == "EUR"
    assert claim["resolved_source_ids"] == ["evidence"]
    assert claim["resolved_method_ids"] == ["fcff"]
    assert "fcff/pv_terminal" in claim["resolved_calculation_ids"]
    assert claim["status"] == "ready_for_professional_review"
    assert claim["review_dependencies_ready"] is False
    assert claim["semantic_support"] == "requires_professional_review"


def test_method_and_claim_need_separate_local_review() -> None:
    result = build_valuation(reviewed_claimed_case(), FIXTURE)
    assert result["claims"][0]["status"] == "accepted_workpaper"
    assert (
        result["claims"][0]["semantic_support"]
        == "locally_attested_not_independently_verified"
    )
    assert result["piv_conformity"] == "not_assessed"


def test_claim_acceptance_cannot_precede_its_method_review() -> None:
    case = claimed_case()
    prepared = build_valuation(case, FIXTURE)
    case["claims"][0]["review"] = review(prepared["claims"][0]["dependency_sha256"])
    result = build_valuation(case, FIXTURE)
    assert result["claims"][0]["status"] == "ready_for_professional_review"
    assert result["claims"][0]["review_dependencies_ready"] is False


@pytest.mark.parametrize(("field", "value"), [("value", "751"), ("unit", "USD")])
def test_mismatched_claim_number_blocks_claim_without_changing_calculation(
    field: str, value: str
) -> None:
    case = claimed_case()
    case["claims"][0]["values"][0][field] = value
    result = build_valuation(case, FIXTURE)
    assert result["claims"][0]["status"] == "blocked"
    assert "differs" in result["claims"][0]["issues"][0]
    assert result["methods"][0]["status"] == "ready_for_professional_review"
    assert Decimal(
        next(row for row in result["calculations"] if row["id"] == "fcff/equity")[
            "value"
        ]
    ) == Decimal("750")


@pytest.mark.parametrize(
    ("field", "refs"),
    [
        ("calculation_ids", ["fcff/missing"]),
        ("source_ids", ["missing"]),
        ("input_ids", ["missing"]),
        ("method_ids", ["missing"]),
    ],
)
def test_unavailable_claim_reference_stays_visible(field: str, refs: list[str]) -> None:
    case = claimed_case()
    case["claims"][0][field] = refs
    result = build_valuation(case, FIXTURE)
    assert result["claims"][0]["status"] == "blocked"
    assert result["claims"][0]["issues"]
    assert result["methods"][0]["status"] == "ready_for_professional_review"


def test_changed_claim_text_revokes_claim_and_conclusion_but_preserves_method() -> None:
    case = reviewed_claimed_case()
    case["conclusion"] = {
        "text": "Conclusione sintetica condizionata.",
        "method_ids": ["fcff"],
        "claim_ids": ["fcff-equity"],
        "review": None,
    }
    prepared = build_valuation(case, FIXTURE)
    case["conclusion"]["review"] = review(prepared["conclusion"]["dependency_sha256"])
    case["claims"][0]["text"] = "Un testo interpretativo diverso deve essere rivisto."
    result = build_valuation(case, FIXTURE)
    assert result["claims"][0]["stale_review"] is True
    assert result["conclusion"]["status"] == "draft"
    assert result["methods"][0]["status"] == "accepted_workpaper"


def test_claim_based_conclusion_needs_its_own_review() -> None:
    case = reviewed_claimed_case()
    case["conclusion"] = {
        "text": "Conclusione sintetica condizionata.",
        "method_ids": ["fcff"],
        "claim_ids": ["fcff-equity"],
        "review": None,
    }
    prepared = build_valuation(case, FIXTURE)
    case["conclusion"]["review"] = review(prepared["conclusion"]["dependency_sha256"])
    result = build_valuation(case, FIXTURE)
    assert result["conclusion"]["status"] == "accepted_workpaper"
    assert result["conclusion"]["narrative_traceability"] == "explicit_claim_bindings"


def test_direct_normalization_claim_keeps_tax_review_dependency() -> None:
    case = reviewed_normalized_case()
    claim = claimed_case()["claims"][0]
    claim.update(
        {
            "id": "adjusted-income",
            "text": "Reddito rettificato sintetico.",
            "calculation_ids": ["normalization/income-2026/adjusted"],
            "values": [
                {
                    "calculation_id": "normalization/income-2026/adjusted",
                    "value": "100",
                    "unit": "EUR",
                }
            ],
        }
    )
    case["claims"] = [claim]
    prepared = build_valuation(case, FIXTURE)
    claim["review"] = review(prepared["claims"][0]["dependency_sha256"])
    case["normalizations"][0]["adjustments"][0][
        "tax_treatment"
    ] = "Trattamento fiscale modificato."
    result = build_valuation(case, FIXTURE)
    assert result["claims"][0]["stale_review"] is True
    assert result["claims"][0]["status"] == "partial"


def test_source_only_claim_has_no_invented_calculation() -> None:
    case = claimed_case()
    case["claims"][0].update(
        {
            "kind": "fact",
            "text": "Il documento è materiale di prova sintetico.",
            "source_ids": ["evidence"],
            "calculation_ids": [],
            "values": [],
        }
    )
    result = build_valuation(case, FIXTURE)
    assert result["claims"][0]["resolved_calculation_ids"] == []
    assert result["claims"][0]["resolved_values"] == []
    assert result["claims"][0]["review_dependencies_ready"] is True
    assert result["claims"][0]["status"] == "ready_for_professional_review"


def test_sensitivity_claim_preserves_conditional_scenario_and_base_method() -> None:
    case = claimed_case()
    case["sensitivity"] = [
        {
            "id": "stress",
            "method_id": "fcff",
            "discount_rate": "rate",
            "terminal_rate": "rate",
            "terminal_growth": "growth",
        }
    ]
    case["claims"][0]["calculation_ids"] = ["stress/equity"]
    case["claims"][0]["values"][0]["calculation_id"] = "stress/equity"
    result = build_valuation(case, FIXTURE)
    assert result["claims"][0]["scenario_ids"] == ["stress"]
    assert result["claims"][0]["resolved_method_ids"] == ["fcff"]
    assert "Risultati condizionati degli scenari: stress" in compile_html(result)


def test_unreconciled_adjusted_value_cannot_support_an_accepted_claim() -> None:
    case = normalized_case()
    next(row for row in case["inputs"] if row["id"] == "income")["value"] = "999"
    claim = claimed_case()["claims"][0]
    claim.update(
        {
            "calculation_ids": ["normalization/income-2026/adjusted"],
            "values": [
                {
                    "calculation_id": "normalization/income-2026/adjusted",
                    "value": "100",
                    "unit": "EUR",
                }
            ],
        }
    )
    case["claims"] = [claim]
    result = build_valuation(case, FIXTURE)
    assert result["claims"][0]["status"] == "blocked"
    assert result["claims"][0]["issues"] == ["Unreconciled normalization: income-2026"]


def test_claim_and_named_workpapers_export_from_one_canonical_register(
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    report = build_valuation(reviewed_claimed_case(), FIXTURE)
    output = tmp_path / "claims"
    write_package(report, FIXTURE, output)
    registry = read_json(output / "claim_registry.json")
    calculations = read_json(output / "calculations.json")
    assert registry["data"] == report["claims"]
    assert registry["report_sha256"] == report["report_sha256"]
    assert calculations["data"] == report["calculations"]
    assert read_json(output / "forecast_binding.json")["data"] is None
    assert not (output / "model_data_report.json").exists()
    workbook = load_workbook(output / "valuation_workbook.xlsx")
    assert workbook["Affermazioni"]["E3"].value.startswith("='Calcoli'!B")
    assert "Affermazioni e riscontri" in (output / "valuation_report.md").read_text()


def dated_case(**timing_changes) -> dict:
    """Use a synthetic 100-unit flow and independently supplied annual terminal flow."""
    case = case_data()
    timing = {
        "valuation_date": "2027-01-01",
        "period_end_dates": ["2028-01-01"],
        "cash_flow_timing": "end_period",
        "day_count": "ACT/365F",
        "rate_compounding": "effective_annual",
        "rate_model": "flat",
        "rate_ids": ["rate"],
        "rationale": "Synthetic timing choice; no claim of economic suitability.",
        **timing_changes,
    }
    case["mandate"]["valuation_date"] = timing["valuation_date"]
    method = case["methods"][0]
    method["timing"] = timing
    method["inputs"].pop("discount_rate")
    method["inputs"]["terminal_rate"] = "rate"
    method["inputs"]["flows"] = ["flow"] * len(timing["period_end_dates"])
    rate_input = next(row for row in case["inputs"] if row["id"] == "rate")
    case["inputs"].extend(
        [
            {**rate_input, "id": "rate20", "value": "0.2"},
            {**rate_input, "id": "rate25", "value": "0.25"},
        ]
    )
    return case


@pytest.mark.parametrize(
    ("timing", "expected"),
    [
        ({}, 1000.0),
        ({"cash_flow_timing": "mid_period"}, 1004.4371680154683),
        (
            {"valuation_date": "2026-06-30", "period_end_dates": ["2026-12-31"]},
            1048.3981252157032,
        ),
        (
            {
                "valuation_date": "2026-12-31",
                "period_end_dates": ["2027-01-31", "2027-02-28"],
            },
            1182.3767274027036,
        ),
        ({"rate_compounding": "continuous"}, 995.3211598395554),
        (
            {
                "period_end_dates": ["2028-01-01", "2029-01-01"],
                "day_count": "ACT/ACT_ISDA",
                "rate_model": "spot_curve",
                "rate_ids": ["rate", "rate20"],
                "terminal_discount_rate": "rate25",
            },
            800.3535353535353,
        ),
        (
            {
                "period_end_dates": ["2028-01-01", "2029-01-01"],
                "day_count": "ACT/ACT_ISDA",
                "rate_model": "forward_curve",
                "rate_ids": ["rate", "rate20"],
            },
            924.2424242424242,
        ),
        (
            {
                "period_end_dates": ["2028-01-01", "2029-01-01"],
                "day_count": "ACT/ACT_ISDA",
                "cash_flow_timing": "mid_period",
                "rate_model": "forward_curve",
                "rate_ids": ["rate", "rate20"],
            },
            935.9102827889783,
        ),
    ],
)
def test_dated_dcf_matches_independent_present_values(
    timing: dict, expected: float
) -> None:
    result = build_valuation(dated_case(**timing), FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert float(values["fcff/value"]) == pytest.approx(expected, rel=0, abs=1e-9)
    assert result["methods"][0]["convention"] == "explicit_dated_discount"


@pytest.mark.parametrize(
    ("basis", "expected"),
    [("ACT/ACT_ISDA", "1"), ("ACT/365F", "1.002739726027397260273972602739726027397")],
)
def test_dated_leap_year_preserves_declared_day_count(
    basis: str, expected: str
) -> None:
    case = dated_case(
        valuation_date="2024-01-01", period_end_dates=["2025-01-01"], day_count=basis
    )
    result = build_valuation(case, FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert Decimal(values["fcff/time/1/end"]) == Decimal(expected)


@pytest.mark.parametrize(
    ("field", "invalid"),
    [
        ("valuation_date", "2026-12-31"),
        ("period_end_dates", ["2027-01-01"]),
        ("period_end_dates", ["20280101"]),
        ("period_end_dates", ["2028-02-30"]),
        ("period_end_dates", ["2028-01-01", "2029-01-01"]),
        ("period_end_dates", ["2200-01-01"]),
        ("day_count", "ACT/ACT_ICMA"),
        ("day_count", []),
        ("cash_flow_timing", "automatic"),
        ("rate_compounding", "monthly_nominal"),
        ("rate_model", "automatic"),
        ("rate_model", "spot_curve"),
        ("rate_ids", ["missing_rate"]),
        ("rate_ids", []),
        ("rationale", ""),
    ],
)
def test_invalid_timing_blocks_only_its_method(field: str, invalid: object) -> None:
    case = dated_case()
    case["methods"][0]["timing"][field] = invalid
    result = build_valuation(case, FIXTURE)
    assert result["methods"][0]["status"] == "blocked"
    assert result["methods"][1]["status"] == "ready_for_professional_review"


def test_midperiod_cash_does_not_shift_terminal_value_to_midperiod() -> None:
    result = build_valuation(dated_case(cash_flow_timing="mid_period"), FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert Decimal(values["fcff/time/1/cash"]) == Decimal("0.5")
    assert Decimal(values["fcff/time/1/end"]) == Decimal("1")
    assert Decimal(values["fcff/discount/terminal-horizon"]) == Decimal("1.1")


def test_changed_timing_invalidates_only_its_method_review() -> None:
    case = dated_case()
    prior = build_valuation(case, FIXTURE)
    case["methods"][0]["review"] = review(prior["methods"][0]["dependency_sha256"])
    case["methods"][1]["review"] = review(prior["methods"][1]["dependency_sha256"])
    case["methods"][0]["timing"]["cash_flow_timing"] = "mid_period"
    result = build_valuation(case, FIXTURE)
    assert result["methods"][0]["stale_review"] is True
    assert result["methods"][1]["status"] == "accepted_workpaper"


def test_dated_continuous_report_exports_formula_and_timing(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    result = build_valuation(dated_case(rate_compounding="continuous"), FIXTURE)
    output = tmp_path / "dated-report"
    write_package(result, FIXTURE, output)
    workbook = load_workbook(output / "valuation_workbook.xlsx", data_only=False)
    formulas = [
        cell.value for cell in workbook["Calcoli"]["B"] if cell.data_type == "f"
    ]
    assert any(formula.startswith("=EXP(") for formula in formulas)
    assert "ACT/365F" in (output / "valuation_report.md").read_text()
    assert "annuale e distinto" in (output / "valuation_report.md").read_text()


@pytest.mark.parametrize(
    ("calculation_id", "expected"),
    [
        ("fcff/value", "1000"),
        ("fcff/equity", "750"),
        ("fcfe/value", "1000"),
        ("income-method/value", "1000"),
        ("nav/value", "650"),
        ("mixed/value", "681.8181818181818181818181818181818181818"),
        ("multiple-method/equity", "250"),
        ("apv/equity", "930"),
    ],
)
def test_supported_methods_have_independent_expected_results(
    calculation_id: str, expected: str
) -> None:
    result = build_valuation(case_data(), FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert Decimal(values[calculation_id]) == Decimal(expected)
    assert result["piv_conformity"] == "not_assessed"
    assert result["status"] == "ready_for_professional_review"


@pytest.mark.parametrize(
    "invalid",
    [None, 1, 0.1, True, "NaN", "Infinity", "1e5", "1,000", "+1", "01", "9" * 81],
)
def test_noncanonical_or_missing_number_is_never_zero(invalid: object) -> None:
    with pytest.raises(ValuationError):
        decimal(invalid)


def test_missing_flow_blocks_only_dependent_methods() -> None:
    case = case_data()
    case["inputs"][0]["value"] = None
    result = build_valuation(case, FIXTURE)
    states = {row["method_id"]: row["status"] for row in result["methods"]}
    assert result["status"] == "partial"
    assert states["fcff"] == "blocked"
    assert states["nav"] == "ready_for_professional_review"


def test_unsupported_method_remains_visible_in_partial_exports(tmp_path: Path) -> None:
    case = case_data()
    case["methods"][0]["kind"] = "SPECIALIST_METHOD"
    result = build_valuation(case, FIXTURE)
    output = tmp_path / "partial"
    write_package(result, FIXTURE, output)
    assert result["methods"][0]["status"] == "blocked"
    assert (
        "Metodo non supportato: SPECIALIST_METHOD"
        in (output / "valuation_report.md").read_text()
    )


def test_equity_method_rejects_second_debt_deduction() -> None:
    case = case_data()
    case["methods"][1]["bridge"] = case["methods"][0]["bridge"]
    result = build_valuation(case, FIXTURE)
    assert result["methods"][1]["status"] == "blocked"
    assert "again" in result["methods"][1]["reason"]


def test_invalid_terminal_growth_is_blocked() -> None:
    case = case_data()
    case["inputs"][3]["value"] = "0.1"
    result = build_valuation(case, FIXTURE)
    assert result["methods"][0]["status"] == "blocked"


def test_unit_mismatch_is_not_implicitly_converted() -> None:
    case = case_data()
    case["inputs"][0]["unit"] = "ratio"
    result = build_valuation(case, FIXTURE)
    assert "unit" in result["methods"][0]["reason"]


def test_source_tampering_prevents_any_export(tmp_path: Path) -> None:
    shutil.copy(FIXTURE / "evidence.txt", tmp_path / "evidence.txt")
    (tmp_path / "evidence.txt").write_text("changed source")
    with pytest.raises(ValuationError, match="Source bytes changed"):
        build_valuation(case_data(), tmp_path)


def test_audience_change_requires_permitted_sources() -> None:
    case = case_data()
    case["audience"] = "bank"
    with pytest.raises(ValuationError, match="audience"):
        build_valuation(case, FIXTURE)


def test_source_traversal_is_rejected() -> None:
    case = case_data()
    case["sources"][0]["path"] = "../business_valuation/evidence.txt"
    with pytest.raises(ValuationError, match="contained"):
        build_valuation(case, FIXTURE)


def test_changed_input_invalidates_only_dependent_reviews() -> None:
    case = case_data()
    initial = build_valuation(case, FIXTURE)
    case["methods"][0]["review"] = review(initial["methods"][0]["dependency_sha256"])
    case["methods"][3]["review"] = review(initial["methods"][3]["dependency_sha256"])
    case["inputs"][0]["value"] = "120"
    result = build_valuation(case, FIXTURE)
    assert result["methods"][0]["stale_review"] is True
    assert result["methods"][0]["status"] == "ready_for_professional_review"
    assert result["methods"][3]["status"] == "accepted_workpaper"


def test_proposed_input_cannot_be_accepted_with_a_hash() -> None:
    case = case_data()
    case["inputs"][0]["status"] = "proposed"
    initial = build_valuation(case, FIXTURE)
    case["methods"][0]["review"] = review(initial["methods"][0]["dependency_sha256"])
    result = build_valuation(case, FIXTURE)
    assert result["methods"][0]["status"] == "partial"


@pytest.mark.parametrize(
    "profile_id",
    [
        "sale",
        "contribution",
        "transformation",
        "merger",
        "demerger",
        "capital_increase",
        "withdrawal",
        "exclusion",
        "inheritance",
        "family",
        "tax",
        "accounting",
        "ppa",
        "litigation",
        "distress",
        "liquidation",
        "collateral",
        "strategy",
        "fairness",
        "review",
        "custom",
    ],
)
def test_explicit_purpose_intake_does_not_enable_professional_use(
    profile_id: str,
) -> None:
    case = case_data()
    case["purpose_profile"] = {
        "id": profile_id,
        "selection_reason": "Explicit choice for synthetic intake only",
        "source_ids": [case["sources"][0]["id"]],
        "locator": "Synthetic mandate paragraph 1",
    }

    result = build_valuation(case, FIXTURE)

    assert result["purpose_coverage"]["profile"]["id"] == profile_id
    assert result["purpose_coverage"]["coverage"]["professional_use_enabled"] is False
    assert result["purpose_coverage"]["coverage"]["professional_review"] == "pending"
    assert result["piv_conformity"] == "not_assessed"


def test_purpose_change_invalidates_previous_method_acceptance() -> None:
    case = case_data()
    prior = build_valuation(case, FIXTURE)
    case["methods"][0]["review"] = review(prior["methods"][0]["dependency_sha256"])
    case["purpose_profile"] = {
        "id": "withdrawal",
        "selection_reason": "The clarified mandate concerns withdrawal rights",
        "source_ids": [case["sources"][0]["id"]],
        "locator": "Synthetic mandate paragraph 1",
    }

    result = build_valuation(case, FIXTURE)

    assert result["methods"][0]["stale_review"] is True
    assert result["methods"][0]["status"] == "ready_for_professional_review"


def test_unknown_purpose_profile_is_not_guessed_from_its_name() -> None:
    case = case_data()
    case["purpose_profile"] = {
        "id": "sale_and_merger",
        "selection_reason": "Ambiguous synthetic choice",
        "source_ids": [case["sources"][0]["id"]],
        "locator": "Synthetic mandate paragraph 1",
    }

    with pytest.raises(ValuationError, match="Unknown purpose profile"):
        build_valuation(case, FIXTURE)


def test_benchmark_publication_after_cutoff_is_rejected() -> None:
    case = case_data()
    case["inputs"][2]["benchmark"] = dict(
        source_url="https://example.org/source",
        observed_on="2026-12-30",
        published_on="2027-01-03",
        retrieved_on="2027-01-04",
        vintage="synthetic",
        definition="Annual rate",
        geography="Synthetic",
        max_age_days=10,
        selection_reason="Synthetic test",
    )
    with pytest.raises(ValuationError, match="look-ahead"):
        build_valuation(case, FIXTURE)


def test_invalid_sensitivity_is_retained_without_corrupting_base() -> None:
    case = case_data()
    case["sensitivity"] = [
        dict(
            id="bad-rate",
            method_id="fcff",
            discount_rate="rate",
            terminal_rate="rate",
            terminal_growth="rate",
        )
    ]
    result = build_valuation(case, FIXTURE)
    assert result["sensitivity"][0]["status"] == "blocked"
    assert result["methods"][0]["status"] == "ready_for_professional_review"


def test_sensitivity_cannot_revive_an_excluded_method() -> None:
    case = case_data()
    case["methods"][0]["selected"] = False
    case["sensitivity"] = [
        dict(
            id="scenario",
            method_id="fcff",
            discount_rate="rate",
            terminal_rate="rate",
            terminal_growth="growth",
        )
    ]
    with pytest.raises(ValuationError, match="selected DCF"):
        build_valuation(case, FIXTURE)


def test_plan_lineage_without_a_replayed_plan_is_rejected() -> None:
    case = case_data()
    case["inputs"][0]["plan_calculation_ids"] = ["invented"]
    with pytest.raises(ValuationError, match="plan binding"):
        build_valuation(case, FIXTURE)


@pytest.mark.parametrize("change_conclusion", [False, True])
def test_conclusion_acceptance_is_bound_to_its_exact_text(
    change_conclusion: bool,
) -> None:
    case = case_data()
    initial = build_valuation(case, FIXTURE)
    case["mandate_details"]["review"] = review(
        initial["mandate_assessment"]["dependency_sha256"]
    )
    case["methods"][0]["review"] = review(initial["methods"][0]["dependency_sha256"])
    case["conclusion"] = {
        "text": "Synthetic conclusion",
        "method_ids": ["fcff"],
        "review": None,
    }
    draft = build_valuation(case, FIXTURE)
    case["conclusion"]["review"] = review(draft["conclusion"]["dependency_sha256"])
    if change_conclusion:
        case["conclusion"]["text"] = "Changed conclusion"
    result = build_valuation(case, FIXTURE)
    assert result["conclusion"]["status"] == (
        "draft" if change_conclusion else "accepted_workpaper"
    )
    assert result["status"] == (
        "ready_for_professional_review" if change_conclusion else "accepted_workpaper"
    )


def test_report_replay_rejects_tampered_derived_values(tmp_path: Path) -> None:
    result = build_valuation(case_data(), FIXTURE)
    result["calculations"][-1]["value"] = "999999"
    result["report_sha256"] = digest(
        {key: value for key, value in result.items() if key != "report_sha256"}
    )
    with pytest.raises(ValuationError, match="canonical replay"):
        write_package(result, FIXTURE, tmp_path / "export")
    assert not (tmp_path / "export").exists()


def test_html_escapes_authored_content() -> None:
    case = case_data()
    case["entity_name"] = "<script>alert('x')</script>"
    result = compile_html(build_valuation(case, FIXTURE))
    assert "<script>" not in result
    assert "&lt;script&gt;" in result


def test_workbook_formulas_replay_to_engine_values(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    report = build_valuation(case_data(), FIXTURE)
    path = tmp_path / "valuation.xlsx"
    write_workbook(path, report)
    workbook = load_workbook(path, data_only=False)
    cells = {}
    for index, record in enumerate(report["calculations"], 2):
        value = workbook["Calcoli"].cell(index, 2).value
        if isinstance(value, (int, float)):
            calculated = Decimal(str(value))
        elif value.startswith("='Dati'!"):
            calculated = Decimal(str(workbook["Dati"][value.split("!")[1]].value))
        elif value == "=0":
            calculated = Decimal(0)
        elif value.startswith("=SUM("):
            calculated = evaluate("sum", [cells[key] for key in value[5:-1].split(",")])
        else:
            match = re.fullmatch(r"=(B\d+)([-*/^])(B\d+)", value)
            assert match is not None
            op = {"-": "subtract", "*": "multiply", "/": "divide", "^": "power"}[
                match[2]
            ]
            calculated = evaluate(op, [cells[match[1]], cells[match[3]]])
        cells[f"B{index}"] = calculated
        assert abs(calculated - Decimal(record["value"])) < Decimal("0.00000001")
    assert workbook["Sintesi"]["C5"].value.startswith("='Calcoli'!")


def test_workbook_untrusted_labels_are_literal(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    case = case_data()
    case["entity_name"] = '=HYPERLINK("https://example.org")'
    case["inputs"][0]["description"] = "=1+1"
    path = tmp_path / "literal.xlsx"
    write_workbook(path, build_valuation(case, FIXTURE))
    workbook = load_workbook(path)
    assert workbook["Sintesi"]["B1"].data_type == "s"
    assert workbook["Dati"]["B2"].data_type == "s"


def test_export_formats_are_readable_and_preserve_status(tmp_path: Path) -> None:
    from docx import Document
    from pypdf import PdfReader

    report = build_valuation(case_data(), FIXTURE)
    output = tmp_path / "report"
    artifacts = write_package(report, FIXTURE, output)
    assert {item["path"] for item in artifacts} == {
        "valuation.json",
        "valuation_report.html",
        "valuation_report.md",
        "valuation_report.docx",
        "valuation_report.pdf",
        "valuation_workbook.xlsx",
        "calculations.csv",
        "mandate.json",
        "evidence.json",
        "normalizations.json",
        "forecast_binding.json",
        "method_decisions.json",
        "benchmark_observations.json",
        "calculations.json",
        "sensitivity.json",
        "valuation_conclusion.json",
        "claim_registry.json",
        "professional_review.json",
    }
    assert (
        Document(output / "valuation_report.docx").paragraphs[0].text
        == "Valutazione d’impresa"
    )
    assert (
        "Carte di lavoro"
        in PdfReader(output / "valuation_report.pdf").pages[0].extract_text()
    )
    assert (
        read_json(output / "valuation.json")["status"]
        == "ready_for_professional_review"
    )


def test_duplicate_json_keys_are_rejected(tmp_path: Path) -> None:
    path = tmp_path / "duplicate.json"
    path.write_text('{"value":"1","value":"2"}')
    with pytest.raises(ValuationError, match="Duplicate JSON"):
        read_json(path)


def prepare_archive_run(
    tmp_path: Path,
    workflow: str = "business-valuation",
    *,
    supplied_case: dict | None = None,
    source_files: list[Path] | None = None,
) -> tuple[Path, Path]:
    from tests.plugins.test_vera_client_workflow_filesystem import _load_customer_ledger

    ledger = _load_customer_ledger()
    client = tmp_path / "Client"
    client.mkdir()
    identity = "client_111111111111111111111111"
    ledger.create_client_manifest(client, identity)
    engagement = ledger.create_engagement(client, identity, "Synthetic valuation")
    case = deepcopy(supplied_case) if supplied_case is not None else case_data()
    input_ids = []
    for index, source_file in enumerate(source_files or [FIXTURE / "evidence.txt"]):
        receipt = ledger.import_document(
            client, identity, engagement["engagement_id"], source_file, "source"
        )["receipt"]
        input_ids.append(receipt["input_id"])
        case["sources"][index][
            "path"
        ] = f"imports/{receipt['input_id']}/{Path(receipt['relative_path']).name}"
    original = tmp_path / "case.json"
    original.write_text(json.dumps(case))
    imported = ledger.import_document(
        client, identity, engagement["engagement_id"], original, "source"
    )
    prepared = ledger.prepare_run(
        client,
        identity,
        engagement["engagement_id"],
        workflow,
        "development",
        input_ids=[*input_ids, imported["receipt"]["input_id"]],
    )
    running = ledger.start_run(
        client, engagement["engagement_id"], prepared["run"]["run_id"]
    )
    case_path = next(
        Path(row["path"])
        for row in running["context"]["input_bindings"]
        if row["binding_id"] == imported["receipt"]["input_id"]
    )
    return case_path, Path(running["context_path"])


@pytest.mark.parametrize("phase", ["demo", "practice"])
def test_valuation_teaching_sources_run_the_bound_workflow(
    tmp_path: Path, phase: str, record_property
) -> None:
    """Interpret the supplied fictional note explicitly; never ship a preapproved case."""
    from tests.plugins._teaching_release import record_native_check

    folder = ROOT / "plugins/vera/assets/courses/business-valuation/files"
    sources = [folder / "input/caso-it.md"]
    case = case_data()
    case["entity_name"] = "Officina Arco — esercizio sintetico"
    # The teaching note is not a complete engagement letter: preserve unknowns.
    case["mandate_details"] = {
        key: {"value": None, "status": "proposed", "source_ids": [], "locator": None}
        for key in case["mandate_details"]
        if key not in {"interests", "standards", "review"}
    }
    case["mandate_details"]["interests"] = []
    case["mandate_details"]["standards"] = []
    for key, value in (
        ("subject_type", "enterprise"),
        ("recipients", "Soli destinatari interni della lezione"),
    ):
        case["mandate_details"][key] = {
            "value": value,
            "status": "confirmed",
            "source_ids": ["evidence"],
            "locator": "caso-it.md: carte interne e intera attività operativa",
        }
    case["purpose_profile"] = {
        "id": "strategy",
        "selection_reason": "La nota richiede un confronto interno su ipotesi didattiche, senza finalità legale.",
        "source_ids": [case["sources"][0]["id"]],
        "locator": "caso-it.md: incarico e base ipotetica per confronto interno",
    }
    case["methods"] = [case["methods"][0]]
    case["methods"][0][
        "rationale"
    ] = "La nota propone flussi FCFF annuali e un raccordo esplicito; le altre basi non sono disponibili."
    case["methods"][0]["limitations"] = [
        "Tassi didattici, sostenibilità non dimostrata, nessuna finalità legale qualificata."
    ]
    needed = {"flow", "terminal", "rate", "growth", "debt", "cash", "zero"}
    case["inputs"] = [item for item in case["inputs"] if item["id"] in needed]
    for item in case["inputs"]:
        if item["unit"] == "EUR":
            item["value"] = str(Decimal(item["value"]) * 1000)
        item["description"] = f"Ipotesi didattica: {item['id']}"
        item["locator"] = "caso-it.md: flussi, tassi e raccordo"
    assert "100.000" in sources[0].read_text()
    case["sources"][0].update(
        path="caso-it.md",
        description="Nota didattica sintetica",
        sha256=hashlib.sha256(sources[0].read_bytes()).hexdigest(),
    )
    if phase == "practice":
        sources.append(folder / "practice/pratica-it.md")
        assert "12%" in sources[1].read_text()
        case["sources"].append(
            {
                **case["sources"][0],
                "id": "update",
                "path": "pratica-it.md",
                "sha256": hashlib.sha256(sources[1].read_bytes()).hexdigest(),
            }
        )
        rate = next(item for item in case["inputs"] if item["id"] == "rate")
        rate.update(
            value="0.12",
            source_ids=["update"],
            locator="pratica-it.md: tasso di sconto e terminale",
        )
    case_path, context = prepare_archive_run(
        tmp_path, supplied_case=case, source_files=sources
    )
    output = run_valuation.run_case(case_path, context)
    result_path = Path(output["output_dir"]) / "valuation.json"
    report = read_json(result_path)
    amounts = {row["id"]: Decimal(row["value"]) for row in report["calculations"]}
    expected = (
        Decimal("750000")
        if phase == "demo"
        else Decimal("583333.3333333333333333333333333333333333")
    )
    assert abs(amounts["fcff/equity"] - expected) < Decimal("0.000001")
    assert report["status"] == "partial"
    assert report["mandate_assessment"]["status"] == "partial"
    assert report["conclusion"] is None
    assert report["purpose_coverage"]["profile"]["id"] == "strategy"
    assert report["purpose_coverage"]["coverage"]["professional_use_enabled"] is False
    record_property("teaching_output", output["output_dir"])
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="business-valuation",
        language="it",
        phase=phase,
    )


def test_portable_archive_run_exports_and_replays_idempotently(tmp_path: Path) -> None:
    case_path, context = prepare_archive_run(tmp_path)
    first = run_valuation.run_case(case_path, context)
    second = run_valuation.run_case(case_path, context)
    assert first == second
    assert first["status"] == "ready_for_professional_review"
    assert Path(first["output_dir"]).is_relative_to(tmp_path / "Client")


def test_other_workflow_context_is_rejected(tmp_path: Path) -> None:
    case_path, context = prepare_archive_run(tmp_path, "financial-analysis")
    with pytest.raises(ValueError, match="different Vera workflow"):
        run_valuation.run_case(case_path, context)


def planning_case(tmp_path: Path, months: int = 12) -> dict:
    """Reuse the real v3 synthetic case and compiler for an explicit horizon."""
    scripts = ROOT / "plugins/business-planning/scripts"
    sys.path.insert(0, str(scripts))
    from planning_workflow import build_plan

    planning_fixture = ROOT / "tests/fixtures/business_planning"
    upstream = read_json(planning_fixture / "case.json")
    periods = [f"{2027 + index // 12}-{index % 12 + 1:02}" for index in range(months)]
    upstream["periods"] = periods
    upstream["assumptions"][0]["effective_periods"] = periods
    for scenario in upstream["financial"]["scenarios"]:
        original = deepcopy(scenario["schedule"][-1])
        scenario["schedule"].extend(
            {**deepcopy(original), "period": period} for period in periods[3:]
        )
    plan = build_plan(upstream, source_root=planning_fixture)
    assert plan["status"] == "ready_for_professional_review"
    case = case_data()
    shutil.copy(FIXTURE / "evidence.txt", tmp_path / "evidence.txt")
    (tmp_path / "plan.json").write_text(json.dumps(plan))
    case["sources"].append(
        dict(
            id="plan",
            path="plan.json",
            sha256=hashlib.sha256((tmp_path / "plan.json").read_bytes()).hexdigest(),
            description="Replayed synthetic v3 plan",
            allowed_audiences=["internal"],
            status="reviewed",
        )
    )
    mapping = {}
    for original in upstream["sources"]:
        name = f"plan-{original['id']}.txt"
        shutil.copy(planning_fixture / original["path"], tmp_path / name)
        identifier = f"plan-{original['id']}"
        mapping[original["id"]] = identifier
        case["sources"].append(
            dict(
                id=identifier,
                path=name,
                sha256=original["sha256"],
                description="Original plan source",
                allowed_audiences=["internal"],
                status="reviewed",
            )
        )
    case["inputs"][0].update(
        value=str(-100 * months),
        source_ids=["plan"],
        plan_calculation_ids=[
            f"base/{period}/{metric}"
            for period in periods
            for metric in ("ebit", "working_capital")
        ],
    )
    case["plan_binding"] = dict(
        source_id="plan",
        source_map=mapping,
        scenario_id="base",
        cash_operating_taxes={period: "zero" for period in periods},
        opening_operating_nwc="zero",
        annual_input_ids=["flow"],
        operating_classification="All supplied current positions are operating; synthetic zero opening NWC",
        tax_refund_basis="",
    )
    return case


def monthly_planning_case(
    tmp_path: Path, periods: list[str] | None = None, *, months: int = 3
) -> dict:
    """Bind explicit synthetic monthly FCFF, keeping annual terminal evidence separate."""
    case = planning_case(tmp_path, months)
    selected = periods or ["2027-01", "2027-02", "2027-03"]
    case["inputs"][0] = deepcopy(case_data()["inputs"][0])
    template = case["inputs"][0]
    refs = [f"monthly-{period}" for period in selected]
    for period, ref in zip(selected, refs):
        case["inputs"].append(
            {
                **template,
                "id": ref,
                "value": "-100",
                "source_ids": ["plan"],
                "plan_calculation_ids": [
                    f"base/{period}/ebit",
                    f"base/{period}/working_capital",
                ],
            }
        )
    case["inputs"].append({**template, "id": "operating-tax", "value": "0"})
    binding = case["plan_binding"]
    binding.pop("annual_input_ids")
    binding.update(
        flow_frequency="monthly",
        selected_periods=selected,
        monthly_input_ids=refs,
        cash_operating_taxes={period: "operating-tax" for period in selected},
        operating_classification="Synthetic operating current balances; the declared opening stock refers to the selected valuation date.",
    )
    valuation_date = (
        date.fromisoformat(selected[0] + "-01") - timedelta(days=1)
    ).isoformat()
    case["mandate"]["valuation_date"] = valuation_date
    case["mandate"]["information_cutoff"] = valuation_date
    method = case["methods"][0]
    method["inputs"].pop("discount_rate")
    method["inputs"].update(flows=refs, terminal_rate="rate")
    method["timing"] = {
        **dated_case()["methods"][0]["timing"],
        "valuation_date": valuation_date,
        "period_end_dates": [
            f"{period}-{monthrange(int(period[:4]), int(period[5:]))[1]:02}"
            for period in selected
        ],
    }
    return case


def test_monthly_partial_year_plan_preserves_actual_flows_and_dated_value(
    tmp_path: Path,
) -> None:
    case = monthly_planning_case(tmp_path)

    result = build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")

    assert [row["fcff"] for row in result["plan_bridge"]["monthly"]] == [
        "-100",
        "-100",
        "-100",
    ]
    assert result["plan_bridge"]["annual"] == []
    assert result["plan_bridge"]["period_end_dates"] == [
        "2027-01-31",
        "2027-02-28",
        "2027-03-31",
    ]
    values = {row["id"]: float(row["value"]) for row in result["calculations"]}
    assert values["fcff/value"] == pytest.approx(681.4306143628385)
    assert values["fcff/equity"] == pytest.approx(431.43061436283847)
    assert result["methods"][0]["status"] == "ready_for_professional_review"


def test_selected_plan_months_use_declared_opening_stock_and_cash_taxes(
    tmp_path: Path,
) -> None:
    case = monthly_planning_case(tmp_path, ["2027-04", "2027-05", "2027-06"], months=12)
    template = case["inputs"][0]
    case["inputs"].append({**template, "id": "opening-nwc", "value": "50"})
    case["plan_binding"]["opening_operating_nwc"] = "opening-nwc"
    next(row for row in case["inputs"] if row["id"] == "operating-tax")["value"] = "5"
    next(row for row in case["inputs"] if row["id"] == "monthly-2027-04")[
        "value"
    ] = "-55"
    next(row for row in case["inputs"] if row["id"] == "monthly-2027-05")[
        "value"
    ] = "-105"
    next(row for row in case["inputs"] if row["id"] == "monthly-2027-06")[
        "value"
    ] = "-105"

    result = build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")

    assert [row["fcff"] for row in result["plan_bridge"]["monthly"]] == [
        "-55",
        "-105",
        "-105",
    ]
    assert [row["opening_nwc"] for row in result["plan_bridge"]["monthly"]] == [
        "50",
        "0",
        "0",
    ]
    assert result["plan_bridge"]["annual"] == []
    assert result["plan_bridge"]["monthly"][0]["plan_calculation_ids"] == [
        "base/2027-04/ebit",
        "base/2027-04/working_capital",
    ]


def test_monthly_plan_across_year_end_retains_leap_day_without_annualizing(
    tmp_path: Path,
) -> None:
    case = monthly_planning_case(
        tmp_path, ["2027-11", "2027-12", "2028-01", "2028-02"], months=24
    )

    result = build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")

    assert result["plan_bridge"]["period_end_dates"] == [
        "2027-11-30",
        "2027-12-31",
        "2028-01-31",
        "2028-02-29",
    ]
    assert result["plan_bridge"]["annual"] == []


@pytest.mark.parametrize(
    "periods",
    [
        ["2027-01", "2027-03"],
        ["2027-03", "2027-02", "2027-01"],
        ["2027-04"],
    ],
)
def test_monthly_plan_rejects_gaps_reordering_and_unavailable_periods(
    tmp_path: Path, periods: list[str]
) -> None:
    case = monthly_planning_case(tmp_path)
    case["plan_binding"]["selected_periods"] = periods

    with pytest.raises(
        ValuationError, match="existing plan periods|contiguous and ordered"
    ):
        build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")


def test_monthly_plan_rejects_midmonth_valuation_without_invented_proration(
    tmp_path: Path,
) -> None:
    case = monthly_planning_case(tmp_path)
    case["mandate"]["valuation_date"] = "2027-01-15"

    with pytest.raises(ValuationError, match="month end before"):
        build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")


def test_monthly_plan_requires_every_selected_month_tax(tmp_path: Path) -> None:
    case = monthly_planning_case(tmp_path)
    del case["plan_binding"]["cash_operating_taxes"]["2027-02"]

    with pytest.raises(ValuationError, match="every month"):
        build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")


def test_monthly_plan_rejects_substituted_lineage(tmp_path: Path) -> None:
    case = monthly_planning_case(tmp_path)
    next(row for row in case["inputs"] if row["id"] == "monthly-2027-02")[
        "plan_calculation_ids"
    ] = ["base/2027-01/ebit"]

    with pytest.raises(ValuationError, match="exact upstream calculation IDs"):
        build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")


@pytest.mark.parametrize(
    "change", ["undated", "reordered", "wrong-date", "terminal", "fcfe"]
)
def test_monthly_plan_cannot_change_flow_clock_or_reuse_fcff_as_annual_income(
    tmp_path: Path, change: str
) -> None:
    case = monthly_planning_case(tmp_path)
    method = case["methods"][0]
    if change == "undated":
        method.pop("timing")
        method["inputs"]["discount_rate"] = method["inputs"].pop("terminal_rate")
    elif change == "reordered":
        method["inputs"]["flows"] = list(reversed(method["inputs"]["flows"]))
    elif change == "wrong-date":
        method["timing"]["period_end_dates"][0] = "2027-01-30"
    elif change == "terminal":
        method["inputs"]["terminal_next_flow"] = "monthly-2027-03"
    else:
        method["kind"] = "DCF_FCFE"
        method.pop("bridge")

    result = build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")

    assert result["methods"][0]["status"] == "blocked"
    assert "Monthly plan FCFF requires dated DCF_FCFF" in result["methods"][0]["reason"]
    assert result["methods"][3]["status"] == "ready_for_professional_review"


def test_monthly_plan_tax_evidence_change_invalidates_only_dependent_review(
    tmp_path: Path,
) -> None:
    case = monthly_planning_case(tmp_path)
    prepared = build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")
    for method, result in zip(case["methods"], prepared["methods"]):
        method["review"] = review(result["dependency_sha256"])
    next(row for row in case["inputs"] if row["id"] == "operating-tax")[
        "description"
    ] = "Revised synthetic cash tax basis"

    result = build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")

    assert result["methods"][0]["status"] == "ready_for_professional_review"
    assert result["methods"][3]["status"] == "accepted_workpaper"


def test_monthly_plan_workbook_links_reconciliation_to_valuation_inputs(
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    case = monthly_planning_case(tmp_path)
    report = build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")
    output = tmp_path / "monthly.xlsx"

    write_workbook(output, report)

    workbook = load_workbook(output)
    rows = {row[0].value: row[0].row for row in workbook["Dati"].iter_rows(min_row=2)}
    assert workbook["Dati"].cell(rows["monthly-2027-01"], 3).value == "='Piano FCFF'!I2"
    assert workbook["Piano FCFF"]["I2"].value == "=B2-C2+D2-E2-H2"
    assert workbook["Piano FCFF"]["F3"].value == "=G2"
    assert workbook["Piano FCFF"]["J2"].value == "-100"


@pytest.mark.parametrize(
    ("change", "reason"),
    [
        ("amount", "differs from replayed"),
        ("unit", "incompatible bridge input"),
        ("count", "Bind one FCFF"),
        ("cycle", "must differ from tax"),
    ],
)
def test_monthly_plan_requires_exact_amounts_units_and_independent_inputs(
    tmp_path: Path, change: str, reason: str
) -> None:
    case = monthly_planning_case(tmp_path)
    first = next(row for row in case["inputs"] if row["id"] == "monthly-2027-01")
    if change == "amount":
        first["value"] = "-90"
    elif change == "unit":
        first["unit"] = "ratio"
    elif change == "count":
        case["plan_binding"]["monthly_input_ids"] = case["plan_binding"][
            "monthly_input_ids"
        ][:-1]
    else:
        case["plan_binding"]["cash_operating_taxes"]["2027-01"] = first["id"]

    with pytest.raises(ValuationError, match=reason):
        build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")


def test_annual_bridge_can_select_a_later_complete_year_from_same_plan(
    tmp_path: Path,
) -> None:
    case = planning_case(tmp_path, 24)
    periods = [f"2028-{month:02}" for month in range(1, 13)]
    case["mandate"]["valuation_date"] = "2027-12-31"
    case["plan_binding"]["selected_periods"] = periods
    case["plan_binding"]["cash_operating_taxes"] = {
        period: "zero" for period in periods
    }
    case["inputs"][0].update(
        value="-1200",
        plan_calculation_ids=[
            f"base/{period}/{metric}"
            for period in periods
            for metric in ("ebit", "working_capital")
        ],
    )

    result = build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")

    assert result["plan_bridge"]["annual"][0]["year"] == "2028"
    assert result["plan_bridge"]["annual"][0]["fcff"] == "-1200"
    assert result["plan_bridge"]["period_end_dates"] == ["2028-12-31"]


def test_existing_plan_bridge_replays_and_reconciles_annual_fcff(
    tmp_path: Path,
) -> None:
    case = planning_case(tmp_path)
    result = build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")
    assert result["plan_bridge"]["annual"][0]["fcff"] == "-1200"
    assert result["plan_bridge"]["annual"][0]["ending_nwc"] == "0"
    assert len(result["plan_bridge"]["monthly"]) == 12
    assert (
        result["plan_bridge"]["monthly"][0]["plan_calculation_ids"][0]
        == "base/2027-01/ebit"
    )


def test_plan_bridge_requires_every_month_cash_tax(tmp_path: Path) -> None:
    case = planning_case(tmp_path)
    del case["plan_binding"]["cash_operating_taxes"]["2027-01"]
    with pytest.raises(ValuationError, match="every month"):
        build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")


def test_plan_bridge_rejects_arbitrary_annual_flow(tmp_path: Path) -> None:
    case = planning_case(tmp_path)
    case["inputs"][0]["value"] = "-1190"
    with pytest.raises(ValuationError, match="differs from replayed"):
        build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")


def test_plan_bridge_pending_cash_tax_prevents_acceptance(tmp_path: Path) -> None:
    case = planning_case(tmp_path)
    next(item for item in case["inputs"] if item["id"] == "zero")["status"] = "proposed"
    result = build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")
    assert result["methods"][0]["status"] == "partial"


def test_dated_plan_cannot_relabel_annual_amount_as_monthly(tmp_path: Path) -> None:
    case = planning_case(tmp_path)
    method = case["methods"][0]
    method["inputs"].pop("discount_rate")
    method["inputs"]["terminal_rate"] = "rate"
    method["timing"] = dated_case()["methods"][0]["timing"]
    method["timing"]["valuation_date"] = case["mandate"]["valuation_date"]
    method["timing"]["period_end_dates"] = ["2027-01-31"]

    result = build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")

    assert result["methods"][0]["status"] == "blocked"
    assert "no monthly relabelling" in result["methods"][0]["reason"]


@pytest.mark.parametrize("model", ["flat", "forward_curve"])
def test_single_rate_sensitivity_keeps_dated_curve_semantics(model: str) -> None:
    case = dated_case(rate_model=model)
    case["sensitivity"] = [
        {
            "id": "higher-rate",
            "method_id": "fcff",
            "discount_rate": "rate20",
            "terminal_rate": "rate20",
            "terminal_growth": "growth",
        }
    ]
    expected = {"flat": "illustrative_sensitivity", "forward_curve": "blocked"}

    result = build_valuation(case, FIXTURE)

    assert result["sensitivity"][0]["status"] == expected[model]
    assert result["methods"][0]["status"] == "ready_for_professional_review"


def statement_case() -> dict:
    """Two independent fictional balance sheets with separately supplied movements."""
    case = case_data()
    case["sources"].append(
        {
            "id": "statement-source",
            "path": "statements.txt",
            "sha256": hashlib.sha256(
                (FIXTURE / "statements.txt").read_bytes()
            ).hexdigest(),
            "description": "Fictional statements and opening/movement records",
            "allowed_audiences": ["internal"],
            "status": "reviewed",
        }
    )
    for identifier, value in {
        "assets-prior": "900",
        "liabilities-prior": "250",
        "equity-prior": "650",
        "assets-opening": "900",
        "liabilities-opening": "250",
        "equity-opening": "650",
        "assets-closing": "1000",
        "liabilities-closing": "300",
        "equity-closing": "700",
        "assets-movement": "100",
        "liabilities-movement": "50",
        "equity-movement": "50",
    }.items():
        case["inputs"].append(
            {
                **case["inputs"][0],
                "id": identifier,
                "value": value,
                "description": identifier,
                "source_ids": ["statement-source"],
                "locator": "statements.txt, fictional balance and movement schedule",
            }
        )
    base = {
        "title": "Prospetto sintetico",
        "perimeter_id": "standalone",
        "perimeter_description": "Società sintetica, individuale",
        "basis": "reported",
        "source_ids": ["statement-source"],
        "locator": "statements.txt",
        "limitations": ["Totali sintetici; nessuna attestazione contabile."],
    }
    case["statements"] = [
        {
            **base,
            "id": "prior",
            "period_start": "2025-01-01",
            "period_end": "2025-12-31",
            "coverage": "balance_only",
            "assets": ["assets-prior"],
            "liabilities": ["liabilities-prior"],
            "equity": ["equity-prior"],
            "bound_input_ids": [],
            "rollforwards": [],
        },
        {
            **base,
            "id": "current",
            "period_start": "2026-01-01",
            "period_end": "2026-12-31",
            "coverage": "balance_and_movements",
            "assets": ["assets-closing"],
            "liabilities": ["liabilities-closing"],
            "equity": ["equity-closing"],
            "bound_input_ids": ["income"],
            "rollforwards": [
                {
                    "id": key,
                    "description": label,
                    "opening_input": f"{key}-opening",
                    "closing_input": f"{key}-closing",
                    "movement_inputs": [f"{key}-movement"],
                    "prior_statement_id": "prior",
                    "prior_closing_input": f"{key}-prior",
                    "comparison_basis": "Saldi dichiarati distintamente, stesso perimetro e base.",
                }
                for key, label in [
                    ("assets", "Attivo"),
                    ("liabilities", "Passività"),
                    ("equity", "Patrimonio netto"),
                ]
            ],
        },
    ]
    return case


def reviewed_statement_case() -> dict:
    """Record synthetic decisions in dependency order, without real approval."""
    case = statement_case()
    for index in range(2):
        prepared = build_valuation(case, FIXTURE)
        case["statements"][index]["review"] = review(
            prepared["statements"][index]["dependency_sha256"]
        )
    prepared = build_valuation(case, FIXTURE)
    for method, output in zip(case["methods"], prepared["methods"]):
        method["review"] = review(output["dependency_sha256"])
    return case


def test_statements_reconcile_independent_balances_and_movements() -> None:
    result = build_valuation(statement_case(), FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert values["statement/current/assets"] == "1000"
    assert values["statement/current/liabilities-equity"] == "1000"
    assert values["statement/current/balance-difference"] == "0"
    assert values["statement/current/equity/expected-close"] == "700"
    assert values["statement/current/equity/difference"] == "0"
    assert values["statement/current/equity/continuity-difference"] == "0"
    assert result["statements"][0]["status"] == "ready_for_professional_review"
    assert result["statements"][1]["status"] == "partial"
    assert result["methods"][2]["status"] == "partial"
    assert result["methods"][0]["status"] == "ready_for_professional_review"


@pytest.mark.parametrize(
    ("identifier", "value", "reason"),
    [
        ("equity-closing", "699", "Assets do not equal"),
        ("assets-movement", "99", "opening plus movements"),
        ("assets-opening", "899", "opening differs from prior closing"),
        ("assets-movement", None, "Missing input"),
        ("assets-prior", "899", "Prior statement is unreconciled"),
    ],
)
def test_statement_failure_blocks_bound_method_only(
    identifier: str, value: str | None, reason: str
) -> None:
    case = statement_case()
    next(row for row in case["inputs"] if row["id"] == identifier)["value"] = value
    result = build_valuation(case, FIXTURE)
    assert result["statements"][1]["status"] == "blocked"
    assert reason in "; ".join(result["statements"][1]["issues"])
    assert result["methods"][2]["status"] == "blocked"
    assert result["methods"][0]["status"] == "ready_for_professional_review"


@pytest.mark.parametrize(
    ("key", "value", "reason"),
    [
        ("period_start", "2026-02-01", "immediately before"),
        ("period_start", "2027-01-01", "period is reversed"),
        ("period_end", "2027-12-31", "information cutoff"),
        ("perimeter_id", "different-branch", "same explicit perimeter"),
        ("basis", "adjusted", "accounting basis"),
    ],
)
def test_statement_period_and_perimeter_mismatch_stay_visible(
    key: str, value: str, reason: str
) -> None:
    case = statement_case()
    case["statements"][1][key] = value
    result = build_valuation(case, FIXTURE)
    assert result["statements"][1]["status"] == "blocked"
    assert reason in "; ".join(result["statements"][1]["issues"])


@pytest.mark.parametrize(
    ("key", "value", "reason"),
    [
        ("prior_statement_id", "missing", "Unresolved prior statement"),
        ("prior_statement_id", None, "supplied together"),
        ("prior_closing_input", "flow", "belong to that statement"),
        ("opening_input", "assets-prior", "independently supplied"),
        ("opening_input", "assets-closing", "independent inputs"),
        ("movement_inputs", ["missing"], "Unresolved statement input"),
        (
            "movement_inputs",
            ["assets-movement", "assets-movement"],
            "schema uniqueItems",
        ),
    ],
)
def test_statement_rejects_tautologies_and_invalid_references(
    key: str, value: object, reason: str
) -> None:
    case = statement_case()
    case["statements"][1]["rollforwards"][0][key] = value
    with pytest.raises(ValuationError, match=reason):
        build_valuation(case, FIXTURE)


def test_statement_movement_coverage_cannot_omit_a_closing_line() -> None:
    case = statement_case()
    case["statements"][1]["rollforwards"].pop()
    with pytest.raises(ValuationError, match="every declared closing line"):
        build_valuation(case, FIXTURE)


def test_statement_duplicate_category_input_rejects_double_counting() -> None:
    case = statement_case()
    case["statements"][0]["equity"] = ["assets-prior"]
    with pytest.raises(ValuationError, match="multiple statement lines"):
        build_valuation(case, FIXTURE)


def test_statement_review_change_expires_only_bound_method() -> None:
    case = reviewed_statement_case()
    case["statements"][1]["perimeter_description"] = "Perimetro da riesaminare"
    result = build_valuation(case, FIXTURE)
    assert result["statements"][1]["stale_review"] is True
    assert result["methods"][2]["stale_review"] is True
    assert result["methods"][2]["status"] == "partial"
    assert result["methods"][0]["status"] == "accepted_workpaper"


def test_prior_statement_review_change_expires_current_statement() -> None:
    case = reviewed_statement_case()
    case["statements"][0]["review"]["reviewer"] = "Another synthetic reviewer"
    result = build_valuation(case, FIXTURE)
    assert result["statements"][0]["status"] == "accepted_workpaper"
    assert result["statements"][1]["stale_review"] is True
    assert result["methods"][0]["status"] == "accepted_workpaper"


def test_statement_claim_cannot_accept_an_unreviewed_balance() -> None:
    case = statement_case()
    calculation = "statement/current/balance-difference"
    case["claims"] = [
        {
            "id": "balance",
            "kind": "fact",
            "text": "Differenza sintetica zero",
            "location": "Quadrature",
            "basis": "Somma delle sole voci dichiarate",
            "source_ids": ["statement-source"],
            "input_ids": [],
            "calculation_ids": [calculation],
            "method_ids": [],
            "limitations": [],
            "values": [{"calculation_id": calculation, "value": "0", "unit": "EUR"}],
        }
    ]
    initial = build_valuation(case, FIXTURE)
    case["claims"][0]["review"] = review(initial["claims"][0]["dependency_sha256"])
    result = build_valuation(case, FIXTURE)
    assert result["claims"][0]["status"] == "partial"
    assert result["claims"][0]["review_dependencies_ready"] is False


def test_statement_workbook_and_evidence_export_share_the_register(
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    case = statement_case()
    case["statements"][1]["perimeter_description"] = '=1+1 <script>alert("x")</script>'
    result = build_valuation(case, FIXTURE)
    output = tmp_path / "exports"
    artifacts = write_package(result, FIXTURE, output)
    workbook = load_workbook(output / "valuation_workbook.xlsx")
    assert len(artifacts) == 18
    assert workbook["Quadrature"].max_row == 9
    assert workbook["Quadrature"]["C3"].data_type == "s"
    assert workbook["Quadrature"]["C3"].value == '=1+1 <script>alert("x")</script>'
    assert workbook["Quadrature"]["E3"].data_type == "f"
    assert "&lt;script&gt;" in (output / "valuation_report.html").read_text()
    assert (
        read_json(output / "evidence.json")["data"]["statements"]
        == result["statements"]
    )


def test_statement_signed_movements_and_negative_equity_remain_explicit() -> None:
    case = statement_case()
    replacements = {
        "assets-prior": "200",
        "assets-opening": "200",
        "assets-closing": "150",
        "assets-movement": "-50",
        "equity-prior": "-50",
        "equity-opening": "-50",
        "equity-closing": "-150",
        "equity-movement": "-100",
    }
    for row in case["inputs"]:
        row["value"] = replacements.get(row["id"], row["value"])
    result = build_valuation(case, FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert values["statement/current/equity"] == "-150"
    assert values["statement/current/equity/expected-close"] == "-150"
    assert result["statements"][1]["issues"] == []


def test_statement_failure_blocks_sensitivity_using_bound_flow() -> None:
    case = statement_case()
    case["statements"][1]["bound_input_ids"] = ["flow"]
    next(row for row in case["inputs"] if row["id"] == "equity-closing")[
        "value"
    ] = "699"
    case["sensitivity"] = [
        {
            "id": "higher",
            "method_id": "fcff",
            "discount_rate": "rate20",
            "terminal_rate": "rate20",
            "terminal_growth": "growth",
        }
    ]
    result = build_valuation(case, FIXTURE)
    assert result["sensitivity"][0]["status"] == "blocked"
    assert result["methods"][2]["status"] == "ready_for_professional_review"


def test_statement_adjusted_balance_links_normalization_formula_and_review() -> None:
    case = statement_case()
    case["statements"][1]["basis"] = "adjusted"
    for row in case["statements"][1]["rollforwards"]:
        row.update(
            {
                "prior_statement_id": None,
                "prior_closing_input": None,
                "comparison_basis": "Apertura autonoma rettificata; nessuna conversione del comparativo riportato.",
            }
        )
    template = next(row for row in case["inputs"] if row["id"] == "assets-closing")
    case["inputs"].extend(
        [
            {**template, "id": "reported-assets", "value": "950"},
            {**template, "id": "asset-adjustment", "value": "50"},
        ]
    )
    entry = normalized_case()["normalizations"][0]["adjustments"][0]
    entry.update(
        {
            "amount_input": "asset-adjustment",
            "source_ids": ["statement-source"],
            "reason": "Variante sintetica 950 + 50 = 1000",
            "accounting_check": "Rettifica attivo +50",
            "economic_rationale": "Variante aritmetica, nessuna stima economica",
            "tax_treatment": "Nessun effetto fiscale implicito",
            "reversibility": "Da revisione professionale",
            "locator": "statements.txt",
        }
    )
    case["normalizations"] = [
        {
            "id": "assets-adjusted",
            "year": 2026,
            "line": "Attivo sintetico",
            "reported_input": "reported-assets",
            "adjusted_input": "assets-closing",
            "adjustments": [entry],
        }
    ]
    result = build_valuation(case, FIXTURE)
    formula = next(
        row
        for row in result["calculations"]
        if row["id"] == "statement/current/input/assets-closing"
    )
    assert formula["op"] == "sum"
    assert formula["arguments"] == ["normalization/assets-adjusted/adjusted"]
    assert result["statements"][1]["status"] == "partial"
    assert "asset-adjustment" in result["methods"][2]["input_ids"]


def test_statement_unknown_unit_blocks_amount_without_substitution() -> None:
    case = statement_case()
    next(row for row in case["inputs"] if row["id"] == "assets-movement")[
        "unit"
    ] = "ratio"
    result = build_valuation(case, FIXTURE)
    assert result["statements"][1]["status"] == "blocked"
    assert "Incompatible unit" in result["statements"][1]["issues"][0]
