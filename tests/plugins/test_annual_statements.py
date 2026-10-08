"""Annual statement regressions use invented amounts, never private filings."""

from __future__ import annotations

import copy
import hashlib
import importlib
import importlib.util
import io
import json
import sys
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import openpyxl
import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "plugins" / "financial-analysis" / "scripts"


@pytest.fixture
def annual(monkeypatch: pytest.MonkeyPatch) -> Any:
    """Isolate script module names from other plugin loaders."""
    monkeypatch.syspath_prepend(str(SCRIPTS))
    names = (
        "annual_statements",
        "annual_statements_render",
        "run_annual_statements",
        "check_dependencies",
        "preparation_contract_kernel",
    )
    original = {name: sys.modules.get(name) for name in names}
    for name in names:
        sys.modules.pop(name, None)
    try:
        yield SimpleNamespace(
            core=importlib.import_module("annual_statements"),
            render=importlib.import_module("annual_statements_render"),
            runner=importlib.import_module("run_annual_statements"),
        )
    finally:
        for name in names:
            sys.modules.pop(name, None)
            if original[name] is not None:
                sys.modules[name] = original[name]


def _case(tmp_path: Path) -> dict[str, Any]:
    """An internally balanced model whose reported subtotal is inconsistent."""
    amounts = {
        "revenue": (1000, 800),
        "production_value": (1100, 900),
        "production_costs": (900, 750),
        "ebit": (200, 150),
        "depreciation": (20, 30),
        "amortisation": (5, 10),
        "credit_impairment": (2, 7),
        "other_impairment": (0, 0),
        "b10": (27, 47),
        "pretax": (190, 140),
        "tax": (50, 40),
        "net_profit": (140, 100),
        "fixed_assets": (100, 120),
        "inventory": (20, 30),
        "receivables": (40, 50),
        "current_financial_assets": (0, 0),
        "cash": (60, 80),
        "current_assets": (120, 167),
        "prepayments": (10, 10),
        "total_assets": (230, 290),
        "total_liabilities_equity": (230, 290),
        "equity": (100, 150),
        "total_debt": (110, 110),
        "financial_debt": (20, 25),
        "current_liabilities": (100, 100),
        "personnel": (300, 350),
        "third_party_assets": (20, 30),
    }
    path = tmp_path / "statements.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "SP e CE"
    sheet.append(["Voce", "2021", "2022"])
    fields = {}
    for row, (key, values) in enumerate(amounts.items(), 2):
        sheet.append([key, *values])
        fields[key] = {
            year: {
                "source_id": "accounts",
                "sheet": sheet.title,
                "cell": f"{column}{row}",
                "scale": "1",
            }
            for year, column in (("2021", "B"), ("2022", "C"))
        }
    book.save(path)
    return {
        "schema_version": "vera.annual_statements_case.v1",
        "entity": "Impresa sintetica",
        "accounting_basis": "italian_oic_positive_expenses",
        "currency": "EUR",
        "years": ["2021", "2022"],
        "tolerance": "0",
        "mapping_basis": "Explicit synthetic source-cell mapping, not professional approval",
        "sources": {
            "accounts": {
                "path": path.name,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            }
        },
        "fields": fields,
        "coverage": {
            "accounts": {"status": "read", "basis": "All synthetic data cells read"}
        },
        "commentary": [
            {
                "kind": "observed",
                "text": "Il subtotale dichiarato differisce dai componenti.",
                "references": [
                    {"source_id": "accounts", "locator": "SP e CE, attivo circolante"}
                ],
            }
        ],
        "limitations": ["Synthetic test; no professional approval"],
    }


def _item(
    result: dict[str, Any], population: str, key: str, year: str = "2022"
) -> dict[str, Any]:
    return next(
        item
        for item in result[population]
        if item["key"] == key and item["year"] == year
    )


def test_source_subtotal_exception_survives_balanced_reconstruction(
    annual: Any, tmp_path: Path
) -> None:
    case = _case(tmp_path)

    result = annual.core.analyse(case, tmp_path)

    assert _item(result, "checks", "current_assets_components")["difference"] == "7"
    assert _item(result, "checks", "assets_subtotals")["difference"] == "-7"
    assert _item(result, "checks", "balance_sheet")["status"] == "passed"
    assert _item(result, "checks", "assets_components")["status"] == "passed"
    assert result["status"] == "qualified"
    assert result["report_ready"] is False


def test_ebitda_definitions_reconcile_to_impairment(
    annual: Any, tmp_path: Path
) -> None:
    result = annual.core.analyse(_case(tmp_path), tmp_path)

    assert _item(result, "metrics", "ebitda_da")["value"] == "190"
    assert _item(result, "metrics", "ebitda_b10")["value"] == "197"
    assert _item(result, "metrics", "ebitda_definition_gap")["value"] == "7"
    assert _item(result, "checks", "b10_components")["status"] == "passed"
    assert _item(result, "metrics", "revenue_change_percent")["value"] == "-20.000000"


def test_missing_detail_does_not_become_zero_or_total_debt(
    annual: Any, tmp_path: Path
) -> None:
    case = _case(tmp_path)
    case["fields"]["financial_debt"]["2022"] = None
    case["fields"]["current_financial_assets"]["2022"] = None
    case["fields"].pop("current_liabilities")

    result = annual.core.analyse(case, tmp_path)

    assert _item(result, "metrics", "mapped_debt_less_cash")["value"] is None
    assert _item(result, "metrics", "current_ratio_reported")["value"] is None
    assert (
        _item(result, "checks", "current_assets_components")["status"]
        == "not_available"
    )
    assert _item(result, "facts", "total_debt")["value"] == "110"


@pytest.mark.parametrize("cell_value", [None, "=SUM(B2:B4)", "1.234,56", True, "#N/A"])
def test_ambiguous_blank_formula_and_error_cells_fail(
    annual: Any, tmp_path: Path, cell_value: Any
) -> None:
    case = _case(tmp_path)
    path = tmp_path / "statements.xlsx"
    book = openpyxl.load_workbook(path)
    book["SP e CE"]["C2"] = cell_value
    book.save(path)
    case["sources"]["accounts"]["sha256"] = hashlib.sha256(
        path.read_bytes()
    ).hexdigest()

    with pytest.raises(ValueError):
        annual.core.analyse(case, tmp_path)


def test_zero_denominator_returns_unavailable(annual: Any, tmp_path: Path) -> None:
    case = _case(tmp_path)
    path = tmp_path / "statements.xlsx"
    book = openpyxl.load_workbook(path)
    book["SP e CE"]["C2"] = 0
    book.save(path)
    case["sources"]["accounts"]["sha256"] = hashlib.sha256(
        path.read_bytes()
    ).hexdigest()

    result = annual.core.analyse(case, tmp_path)

    assert _item(result, "metrics", "ebitda_margin_percent")["value"] is None


def test_duplicate_source_cell_cannot_supply_different_years(
    annual: Any, tmp_path: Path
) -> None:
    case = _case(tmp_path)
    case["fields"]["revenue"]["2022"] = copy.deepcopy(case["fields"]["revenue"]["2021"])

    with pytest.raises(ValueError, match="multiple years"):
        annual.core.analyse(case, tmp_path)


def test_changed_source_is_rejected(annual: Any, tmp_path: Path) -> None:
    case = _case(tmp_path)
    with (tmp_path / "statements.xlsx").open("ab") as stream:
        stream.write(b"changed")

    with pytest.raises(ValueError, match="hash changed"):
        annual.core.analyse(case, tmp_path)


def test_unread_supporting_source_remains_visible(annual: Any, tmp_path: Path) -> None:
    case = _case(tmp_path)
    note = tmp_path / "notes.txt"
    note.write_text("Industry code disclosed here", encoding="utf-8")
    case["sources"]["notes"] = {
        "path": note.name,
        "sha256": hashlib.sha256(note.read_bytes()).hexdigest(),
    }
    case["coverage"]["notes"] = {
        "status": "unread",
        "basis": "Notes supplied but not read",
    }

    result = annual.core.analyse(case, tmp_path)
    outputs = annual.render.render_outputs(result)

    assert result["coverage"]["notes"]["status"] == "unread"
    assert "Notes supplied but not read" in outputs["annual_report.md"].decode()


def test_unknown_commentary_source_and_missing_coverage_fail(
    annual: Any, tmp_path: Path
) -> None:
    case = _case(tmp_path)
    case["commentary"][0]["references"][0]["source_id"] = "nonexistent"

    with pytest.raises(ValueError, match="Unknown commentary source"):
        annual.core.analyse(case, tmp_path)


def test_word_and_excel_preserve_exceptions_definitions_and_editable_formulas(
    annual: Any, tmp_path: Path
) -> None:
    result = annual.core.analyse(_case(tmp_path), tmp_path)

    outputs = annual.render.render_outputs(result)
    word = Document(io.BytesIO(outputs["annual_report.docx"]))
    book = openpyxl.load_workbook(
        io.BytesIO(outputs["annual_analysis.xlsx"]), data_only=False
    )

    assert "Attivo circolante dichiarato meno componenti: 7 EUR" in "\n".join(
        p.text for p in word.paragraphs
    )
    assert len(word.tables) == 3
    assert book["Calcoli"]["C2"].data_type == "f"
    assert book["Controlli"]["C9"].data_type == "f"
    assert "Le svalutazioni".lower() in outputs["annual_report.md"].decode().lower()
    assert book["Risultati Decimal"].max_row == 27
    assert book.calculation.fullCalcOnLoad


def _managed_case(annual: Any, tmp_path: Path) -> tuple[Path, Path, Path]:
    """Use the real portable ledger, imports and running context for CLI tests."""
    case = _case(tmp_path)
    spec = importlib.util.spec_from_file_location(
        "annual_test_ledger", ROOT / "plugins/studio-archive/scripts/client_ledger.py"
    )
    assert spec and spec.loader
    ledger = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = ledger
    spec.loader.exec_module(ledger)
    client = tmp_path / "client"
    client.mkdir()
    client_id = "client_777777777777777777777777"
    ledger.create_client_manifest(client, client_id)
    engagement = ledger.create_engagement(client, client_id, "Annual review")
    imported = ledger.import_document(
        client,
        client_id,
        engagement["engagement_id"],
        tmp_path / "statements.xlsx",
        "source",
    )
    prepared = ledger.prepare_run(
        client,
        client_id,
        engagement["engagement_id"],
        "financial-analysis",
        "test",
        input_ids=[imported["receipt"]["input_id"]],
    )
    running = ledger.start_run(
        client, engagement["engagement_id"], prepared["run"]["run_id"]
    )
    output = Path(running["output_dir"])
    case_dir = output / "case"
    case_dir.mkdir()
    binding = running["context"]["input_bindings"][0]
    (case_dir / "statements.xlsx").write_bytes(Path(binding["path"]).read_bytes())
    case["sources"]["accounts"]["input_binding_id"] = binding["binding_id"]
    case_path = case_dir / "case.json"
    case_path.write_text(json.dumps(case), encoding="utf-8")
    return case_path, output / "annual", Path(running["context_path"])


def test_managed_runner_delivers_qualified_package_and_exact_receipts(
    annual: Any, tmp_path: Path
) -> None:
    case_path, output, context = _managed_case(annual, tmp_path)

    receipt = annual.runner.run_annual_statements(case_path, output, context)

    assert receipt["status"] == "qualified"
    assert receipt["report_ready"] is False
    assert {p.name for p in output.iterdir()} == {
        "annual_report.docx",
        "annual_analysis.xlsx",
        "annual_report.md",
        "annual_result.json",
        "used_case.json",
        "reconciliation.json",
        "annual_execution_receipt.json",
    }
    assert (
        receipt["output_artifacts"][0]["sha256"]
        == hashlib.sha256(
            (output / receipt["output_artifacts"][0]["path"]).read_bytes()
        ).hexdigest()
    )
    with pytest.raises(FileExistsError):
        annual.runner.run_annual_statements(case_path, output, context)


def test_managed_runner_rejects_foreign_output(annual: Any, tmp_path: Path) -> None:
    case_path, _output, context = _managed_case(annual, tmp_path)

    with pytest.raises(ValueError, match="outside"):
        annual.runner.run_annual_statements(case_path, tmp_path / "unmanaged", context)


def test_managed_runner_rejects_source_not_bound_to_import(
    annual: Any, tmp_path: Path
) -> None:
    case_path, output, context = _managed_case(annual, tmp_path)
    case = json.loads(case_path.read_text())
    case["sources"]["accounts"]["input_binding_id"] = "not-a-real-binding"
    case_path.write_text(json.dumps(case), encoding="utf-8")

    with pytest.raises(ValueError, match="exact imported input"):
        annual.runner.run_annual_statements(case_path, output, context)
    assert not output.exists()
