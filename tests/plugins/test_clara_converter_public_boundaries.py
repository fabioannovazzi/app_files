from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import pytest

SCRIPTS = Path(__file__).resolve().parents[2] / "plugins/clara/scripts"


def load_script(name: str, path: Path):
    sys.path.insert(0, str(SCRIPTS))
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def cancelled(*args, **kwargs):
    raise InterruptedError("operator cancelled conversion")


def test_frame_timeline_propagates_operator_cancellation(tmp_path, monkeypatch):
    module = load_script(
        "frame_cancel_regression",
        SCRIPTS / "build_voice_feedback_timeline.py",
    )
    video = tmp_path / "input.mp4"
    video.write_bytes(b"synthetic process-boundary input")
    monkeypatch.setattr(module, "run_process", cancelled)

    with pytest.raises(InterruptedError, match="operator cancelled"):
        module.build_feedback_timeline_payload(
            clean_transcript="Change the title.",
            timed_transcript_segments=[
                {"text": "Change the title.", "start_ms": 0, "end_ms": 3000}
            ],
            video_path=video,
            output_path=tmp_path / "timeline.json",
            ffmpeg_path="/usr/bin/true",
        )


def test_slide_matcher_propagates_operator_cancellation(tmp_path, monkeypatch):
    module = load_script(
        "matcher_cancel_regression",
        SCRIPTS / "match_feedback_frames_to_deck_slides.py",
    )
    source = make_deck(tmp_path)
    monkeypatch.setattr(module, "run_process", cancelled)

    with pytest.raises(InterruptedError, match="operator cancelled"):
        module.match_feedback_timeline_to_deck_payload(
            feedback_timeline={"entries": []},
            deck_path=source / "current.pptx",
            deck_snapshot_path=source / "snapshot.json",
            slide_render_dir=tmp_path / "renders",
            soffice_path="/usr/bin/true",
        )


def make_deck(tmp_path):
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = "Current source deck"
    deck.save(tmp_path / "current.pptx")
    (tmp_path / "snapshot.json").write_text(
        json.dumps({"slides": [{"slide_number": 1}]})
    )
    return tmp_path


def test_slide_matcher_rejects_xml_entities_before_conversion(tmp_path, monkeypatch):
    from zipfile import ZipFile

    from defusedxml.common import EntitiesForbidden

    module = load_script(
        "matcher_xml_entity_regression",
        SCRIPTS / "match_feedback_frames_to_deck_slides.py",
    )
    source = make_deck(tmp_path)
    deck = source / "current.pptx"
    with ZipFile(deck) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    original = entries["ppt/presentation.xml"]
    declaration, body = original.split(b"?>", 1)
    entries["ppt/presentation.xml"] = (
        declaration
        + b'?>\n<!DOCTYPE p:presentation [<!ENTITY payload "untrusted">]>'
        + body
    )
    with ZipFile(deck, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)

    def unexpected_conversion(*args, **kwargs):
        pytest.fail("Unsafe XML reached the converter")

    monkeypatch.setattr(module, "run_process", unexpected_conversion)

    with pytest.raises(EntitiesForbidden):
        module.match_feedback_timeline_to_deck_payload(
            feedback_timeline={"entries": []},
            deck_path=deck,
            deck_snapshot_path=source / "snapshot.json",
            slide_render_dir=tmp_path / "renders",
            soffice_path="/usr/bin/true",
        )


def test_noop_slide_converter_does_not_reuse_old_pdf(tmp_path, monkeypatch):
    import fitz

    module = load_script(
        "matcher_stale_regression",
        SCRIPTS / "match_feedback_frames_to_deck_slides.py",
    )
    source = make_deck(tmp_path)
    renders = tmp_path / "renders"
    old_pdf = renders / "_pdf" / "old.pdf"
    old_pdf.parent.mkdir(parents=True)
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((30, 30), "Unrelated previous deck")
        document.save(old_pdf)
    original = old_pdf.read_bytes()
    monkeypatch.setattr(
        module,
        "run_process",
        lambda command, **kwargs: subprocess.CompletedProcess(command, 0),
    )

    result = module.match_feedback_timeline_to_deck_payload(
        feedback_timeline={"entries": []},
        deck_path=source / "current.pptx",
        deck_snapshot_path=source / "snapshot.json",
        slide_render_dir=renders,
        soffice_path="/usr/bin/true",
    )

    assert result["slide_matching"]["status"] == "skipped"
    assert old_pdf.read_bytes() == original
    assert not list(renders.glob("*.png"))
    assert not (renders / "render_identity.json").exists()


def test_frame_failure_retains_full_converter_log(tmp_path, monkeypatch):
    module = load_script(
        "frame_log_regression",
        SCRIPTS / "build_voice_feedback_timeline.py",
    )
    video = tmp_path / "input.mp4"
    video.write_bytes(b"synthetic process-boundary input")
    diagnostic = "conversion failed " * 5000

    def fail_conversion(command, **kwargs):
        kwargs["stderr"].write(diagnostic)
        return subprocess.CompletedProcess(command, 1)

    monkeypatch.setattr(module, "run_process", fail_conversion)

    result = module.build_feedback_timeline_payload(
        clean_transcript="Change the title.",
        timed_transcript_segments=[
            {"text": "Change the title.", "start_ms": 0, "end_ms": 3000}
        ],
        video_path=video,
        output_path=tmp_path / "timeline.json",
        ffmpeg_path="/usr/bin/true",
    )

    assert result["entries"][0]["frame_extraction_status"] == "partial_or_failed"
    logs = list((tmp_path / "frames").rglob("*.stderr.log"))
    assert logs
    assert logs[0].read_text() == diagnostic
    assert len(result["entries"][0]["frame_extraction_note"]) <= 800


def test_voice_import_cancellation_does_not_register_partial_session(
    tmp_path, monkeypatch
):
    from zipfile import ZipFile

    sys.path.insert(0, str(SCRIPTS))
    import advisor_case_core

    case = tmp_path / "case"
    advisor_case_core.initialize_case(
        case,
        client="Synthetic Client",
        project="Cancellation regression",
        objective="Review deck",
        audience="Advisor",
        output_language="en",
    )
    module = load_script(
        "import_cancel_regression",
        SCRIPTS / "import_hosted_voice_bundle.py",
    )
    payload = {
        "schema_version": 1,
        "source": "case_notes_hosted_voice",
        "captured_at": "2026-01-02T10:30:00+00:00",
        "user_transcript": "Change the title.",
        "assistant_transcript": "",
        "extraction_json": {
            "cleaned_notes_markdown": "",
            "entries": [],
            "open_questions": [],
        },
        "extraction_text": "",
        "video_file_name": "screen.webm",
        "video_content_type": "video/webm",
        "timed_transcript_segments": [
            {"text": "Change the title.", "start_ms": 0, "end_ms": 3000}
        ],
    }
    bundle = tmp_path / "voice.zip"
    with ZipFile(bundle, "w") as archive:
        archive.writestr("case-notes-voice-20260102.json", json.dumps(payload))
        archive.writestr("screen.webm", b"synthetic converter input")
    before = (case / "material_registry.json").read_bytes()
    monkeypatch.setattr(module, "build_feedback_timeline", cancelled)

    with pytest.raises(InterruptedError, match="operator cancelled"):
        module.import_hosted_voice_bundle(case, bundle)

    assert (case / "material_registry.json").read_bytes() == before
    assert not list(case.rglob("feedback_timeline.json"))


def test_normalization_timeout_retains_logs_without_publishing_output(
    tmp_path, monkeypatch
):
    module = load_script(
        "normalization_timeout_regression",
        SCRIPTS / "advisor_case_core.py",
    )
    source = make_deck(tmp_path)
    output = tmp_path / "normalized.pptx"

    def timeout_conversion(command, **kwargs):
        assert kwargs["timeout"] == 90
        kwargs["stderr"].write("converter timed out after starting")
        raise subprocess.TimeoutExpired(command, kwargs["timeout"])

    monkeypatch.setattr(module, "run_process", timeout_conversion)

    with pytest.raises(module.CaseWorkspaceError, match="logs retained"):
        module.normalize_legacy_pptx_for_editable_merge(
            source / "current.pptx",
            output_path=output,
            force=True,
            soffice_binary=Path("/usr/bin/true"),
        )

    assert not output.exists()
    logs = list((tmp_path / ".normalization-attempts").rglob("soffice.stderr.log"))
    assert len(logs) == 1
    assert logs[0].read_text() == "converter timed out after starting"
