from __future__ import annotations

import hashlib
import importlib.util
import json
import os
import shutil
import stat
import subprocess
import sys
from email.message import EmailMessage
from pathlib import Path
from types import ModuleType, SimpleNamespace
from typing import Any

import pytest

try:
    import fitz as _fitz
except ImportError:
    _fitz = None

ROOT = Path(__file__).resolve().parents[2]
COMPONENT_ROOT = ROOT / "plugins" / "studio-archive"
ARCHIVE_CORE_PATH = COMPONENT_ROOT / "scripts" / "archive_core.py"
MCP_SERVER_PATH = COMPONENT_ROOT / "mcp" / "server.cjs"


@pytest.fixture(scope="module")
def archive_core() -> ModuleType:
    """Load the component core without changing production import paths."""

    module_name = "test_vera_studio_archive_core"
    spec = importlib.util.spec_from_file_location(module_name, ARCHIVE_CORE_PATH)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


@pytest.fixture
def indexed_archive(
    tmp_path: Path,
    archive_core: ModuleType,
) -> SimpleNamespace:
    """Create a small indexed archive with two client scopes and root material."""

    archive_root = tmp_path / "Studio"
    rossi = archive_root / "Rossi"
    bianchi = archive_root / "Bianchi"
    rossi.mkdir(parents=True)
    bianchi.mkdir()
    source = rossi / "precedente.md"
    source.write_text(
        "Verbale cessione quote\nIl socio approva la cessione delle quote.",
        encoding="utf-8",
    )
    (bianchi / "nota.txt").write_text(
        "Promemoria interno sul ravvedimento operoso.",
        encoding="utf-8",
    )
    (archive_root / "procedura.txt").write_text(
        "Procedura generale dello studio per il controllo documentale.",
        encoding="utf-8",
    )
    state_dir = tmp_path / "private-state"
    archive_core.configure_archive(archive_root, state_dir=state_dir)
    archive_core.refresh_archive(state_dir=state_dir)
    status = archive_core.studio_archive_status(state_dir=state_dir)
    scopes = {item["display_name"]: item["scope_id"] for item in status["scopes"]}
    return SimpleNamespace(
        root=archive_root,
        state=state_dir,
        source=source,
        scopes=scopes,
    )


def _node_executable() -> str:
    node = shutil.which("node")
    if node is not None:
        return node
    candidates = sorted(
        (Path.home() / ".cache" / "codex-runtimes").glob("*/dependencies/node/bin/node")
    )
    if not candidates:
        pytest.skip("The Codex-bundled Node.js runtime is required.")
    return candidates[-1].as_posix()


def _fitz_or_skip() -> ModuleType:
    """Return the collection-stable PyMuPDF module used by PDF fixtures."""

    if _fitz is None:
        pytest.skip("PyMuPDF creates the test PDF.")
    return _fitz


def test_client_folder_binding_uses_exact_scope_without_mutating_sources(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    registration = archive_core.set_studio_client_identity(
        indexed_archive.scopes["Rossi"],
        legal_names=["Rossi SRL"],
        state_dir=indexed_archive.state,
    )
    client_id = registration["client"]["client_id"]
    before = {
        path.relative_to(indexed_archive.root).as_posix(): path.read_bytes()
        for path in indexed_archive.root.rglob("*")
        if path.is_file()
    }

    result = archive_core.get_studio_client_folder(
        client_id,
        state_dir=indexed_archive.state,
    )

    binding = result["client_folder"]
    after = {
        path.relative_to(indexed_archive.root).as_posix(): path.read_bytes()
        for path in indexed_archive.root.rglob("*")
        if path.is_file()
    }
    assert binding["schema_version"] == "vera.studio_client_folder.v2"
    assert binding["studio_client_id"] == client_id
    assert binding["scope_id"] == indexed_archive.scopes["Rossi"]
    assert binding["client_root"] == str(indexed_archive.root / "Rossi")
    assert binding["scope_relative_dir"] == "Rossi"
    assert len(binding["content_sha256"]) == 64
    assert result["source_archive_mutated"] is False
    assert after == before


def test_archive_root_scope_cannot_be_exported_as_client_folder(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    with pytest.raises(archive_core.ArchiveError, match="not a client folder"):
        archive_core.set_studio_client_identity(
            indexed_archive.scopes["Studio"],
            legal_names=["Studio"],
            state_dir=indexed_archive.state,
        )


def test_client_folder_binding_excludes_private_identity_values(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    scope_id = indexed_archive.scopes["Rossi"]
    registration = archive_core.set_studio_client_identity(
        scope_id,
        email_addresses=["amministrazione@example.com"],
        legal_names=["Example Legal Name"],
        tax_identifiers=["01234567890"],
        state_dir=indexed_archive.state,
    )

    result = archive_core.get_studio_client_folder(
        registration["client"]["client_id"],
        state_dir=indexed_archive.state,
    )

    serialized = json.dumps(result["client_folder"])
    assert "amministrazione@example.com" not in serialized
    assert "Example Legal Name" not in serialized
    assert "01234567890" not in serialized


def test_existing_client_journal_and_support_share_one_engagement(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    # Arrange
    registration = archive_core.set_studio_client_identity(
        indexed_archive.scopes["Rossi"],
        legal_names=["Rossi SRL"],
        state_dir=indexed_archive.state,
    )
    client_id = registration["client"]["client_id"]
    journal = indexed_archive.root.parent / "giornale.xlsx"
    journal_bytes = b"journal source bytes"
    journal.write_bytes(journal_bytes)

    # Act
    journal_import = archive_core.import_studio_client_document(
        client_id,
        journal,
        "journal",
        engagement_label="2025 journal sample",
        state_dir=indexed_archive.state,
    )
    engagement_id = journal_import["engagement"]["engagement_id"]
    normalization_dir = (
        Path(journal_import["client_engagement"]["output_dir"]) / "normalization"
    )
    normalization_dir.mkdir(parents=True)
    (normalization_dir / "normalized_journal.csv").write_text(
        "movement_number,amount_signed\nM-1,10\n",
        encoding="utf-8",
    )
    (normalization_dir / "normalization_diagnostics.json").write_text(
        "{}\n",
        encoding="utf-8",
    )
    support = indexed_archive.root.parent / "fatture.zip"
    support_bytes = b"support source bytes"
    support.write_bytes(support_bytes)
    support_import = archive_core.import_studio_client_document(
        client_id,
        support,
        "support",
        engagement_id=engagement_id,
        state_dir=indexed_archive.state,
    )
    listed = archive_core.list_studio_client_engagements(
        client_id,
        state_dir=indexed_archive.state,
    )

    # Assert
    assert journal.read_bytes() == journal_bytes
    assert support.read_bytes() == support_bytes
    assert Path(journal_import["imported_path"]).read_bytes() == journal_bytes
    assert Path(support_import["imported_path"]).read_bytes() == support_bytes
    assert journal_import["original_preserved"] is True
    assert support_import["original_preserved"] is True
    assert journal_import["client_engagement"]["workflow_id"] == "journal-sampling"
    assert support_import["client_engagement"]["workflow_id"] == "check-entries"
    assert (
        journal_import["client_engagement"]["engagement_id"]
        == support_import["client_engagement"]["engagement_id"]
        == engagement_id
    )
    assert (
        journal_import["client_engagement"]["studio_client_folder"]["studio_client_id"]
        == support_import["client_engagement"]["studio_client_folder"][
            "studio_client_id"
        ]
        == client_id
    )
    assert listed["engagement_count"] == 1
    listed_engagement = listed["engagements"][0]
    assert [item["role"] for item in listed_engagement["imports"]] == [
        "journal",
        "support",
    ]
    assert listed_engagement["workflow_run_count"] == 2
    journal_runs = [
        run
        for run in listed_engagement["workflow_runs"]
        if run["workflow_id"] == "journal-sampling"
    ]
    assert len(journal_runs) == 1
    assert journal_runs[0]["normalization_available"] is True
    assert journal_runs[0]["normalized_journal_path"] == str(
        normalization_dir / "normalized_journal.csv"
    )
    assert Path(journal_runs[0]["client_engagement_path"]).is_file()


def test_new_client_creation_derives_folder_and_defers_relationship_setup(
    tmp_path: Path,
    archive_core: ModuleType,
) -> None:
    # Arrange
    archive_root = tmp_path / "Studio"
    archive_root.mkdir()
    state_dir = tmp_path / "private-state"
    archive_core.configure_archive(archive_root, state_dir=state_dir)

    # Act
    result = archive_core.create_studio_client(
        "Zecca SPA",
        tax_identifiers=["01234567890"],
        state_dir=state_dir,
    )

    # Assert
    assert result["status"] == "created"
    assert result["client"]["client_id"].startswith("client_")
    assert result["client"]["registration_status"] == "registered"
    assert result["client_folder"]["scope_relative_dir"] == "Zecca SPA"
    assert (archive_root / "Zecca SPA").is_dir()
    assert result["relationship_setup_status"] == "new_client_workflow_pending"
    assert result["next_workflow"] == "new-client"


def test_support_import_rejects_another_clients_engagement(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    # Arrange
    rossi = archive_core.set_studio_client_identity(
        indexed_archive.scopes["Rossi"],
        legal_names=["Rossi SRL"],
        state_dir=indexed_archive.state,
    )["client"]["client_id"]
    bianchi = archive_core.set_studio_client_identity(
        indexed_archive.scopes["Bianchi"],
        legal_names=["Bianchi SRL"],
        state_dir=indexed_archive.state,
    )["client"]["client_id"]
    journal = indexed_archive.root.parent / "journal.csv"
    journal.write_text("entry\n", encoding="utf-8")
    engagement_id = archive_core.import_studio_client_document(
        rossi,
        journal,
        "journal",
        state_dir=indexed_archive.state,
    )["engagement"]["engagement_id"]
    support = indexed_archive.root.parent / "invoice.pdf"
    support.write_bytes(b"%PDF support")

    # Act / Assert
    with pytest.raises(archive_core.ArchiveError, match="another client"):
        archive_core.import_studio_client_document(
            bianchi,
            support,
            "support",
            engagement_id=engagement_id,
            state_dir=indexed_archive.state,
        )
    assert not any(
        path.name == support.name for path in indexed_archive.root.rglob("*")
    )


def test_journal_import_rejects_source_from_another_client_scope(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    # Arrange
    rossi = archive_core.set_studio_client_identity(
        indexed_archive.scopes["Rossi"],
        legal_names=["Rossi SRL"],
        state_dir=indexed_archive.state,
    )["client"]["client_id"]
    bianchi_source = indexed_archive.root / "Bianchi" / "journal.xlsx"
    source_bytes = b"Bianchi journal"
    bianchi_source.write_bytes(source_bytes)

    # Act / Assert
    with pytest.raises(archive_core.ArchiveError, match="another Studio Archive scope"):
        archive_core.import_studio_client_document(
            rossi,
            bianchi_source,
            "journal",
            state_dir=indexed_archive.state,
        )
    assert bianchi_source.read_bytes() == source_bytes
    assert not (indexed_archive.root / "Rossi" / "Vera engagements").exists()


def _mcp_request(
    payload: dict[str, Any],
    *,
    state_dir: Path,
) -> dict[str, Any]:
    environment = os.environ.copy()
    environment["VERA_STUDIO_ARCHIVE_PYTHON"] = sys.executable
    environment["VERA_STUDIO_ARCHIVE_STATE_DIR"] = str(state_dir)
    completed = subprocess.run(
        [_node_executable(), str(MCP_SERVER_PATH), "--stdio"],
        cwd=COMPONENT_ROOT,
        env=environment,
        input=json.dumps(payload) + "\n",
        text=True,
        capture_output=True,
        check=True,
        timeout=30,
    )
    responses = [
        json.loads(line) for line in completed.stdout.splitlines() if line.strip()
    ]
    return next(
        response for response in responses if response.get("id") == payload["id"]
    )


def _mcp_tool(
    name: str,
    arguments: dict[str, Any],
    *,
    state_dir: Path,
    request_id: int = 1,
) -> dict[str, Any]:
    response = _mcp_request(
        {
            "jsonrpc": "2.0",
            "id": request_id,
            "method": "tools/call",
            "params": {"name": name, "arguments": arguments},
        },
        state_dir=state_dir,
    )
    return response["result"]


def test_status_without_configuration_does_not_create_state(
    tmp_path: Path,
    archive_core: ModuleType,
) -> None:
    state_dir = tmp_path / "absent-state"

    result = archive_core.studio_archive_status(state_dir=state_dir)

    assert result == {
        "configured": False,
        "document_count": 0,
        "chunk_count": 0,
        "last_refresh_at": None,
        "scopes": [],
        "needs_ocr_document_count": 0,
        "partial_document_count": 0,
        "failed_document_count": 0,
        "scan_issue_count": 0,
        "scan_issues": [],
        "scan_issues_truncated": False,
        "document_issue_count": 0,
        "document_issues": [],
        "document_issues_truncated": False,
    }
    assert not state_dir.exists()


def test_configure_rejects_state_inside_source_without_writing_it(
    tmp_path: Path,
    archive_core: ModuleType,
) -> None:
    archive_root = tmp_path / "Studio"
    archive_root.mkdir()
    (archive_root / "Client").mkdir()
    nested_state = archive_root / ".private-index"

    with pytest.raises(
        archive_core.ArchiveError,
        match="must not contain one another",
    ):
        archive_core.configure_archive(archive_root, state_dir=nested_state)

    assert not nested_state.exists()


def test_two_professionals_build_separate_indexes_from_same_archive(
    tmp_path: Path,
    archive_core: ModuleType,
) -> None:
    archive_root = tmp_path / "Shared Studio"
    client_root = archive_root / "Rossi"
    client_root.mkdir(parents=True)
    (client_root / "memo.txt").write_text(
        "Precedente dello studio sulla trasformazione societaria.",
        encoding="utf-8",
    )
    fabio_state = tmp_path / "Fabio private"
    paolo_state = tmp_path / "Paolo private"

    for state_dir in (fabio_state, paolo_state):
        archive_core.configure_archive(archive_root, state_dir=state_dir)
        archive_core.refresh_archive(state_dir=state_dir)

    fabio_status = archive_core.studio_archive_status(state_dir=fabio_state)
    paolo_status = archive_core.studio_archive_status(state_dir=paolo_state)
    assert fabio_status["document_count"] == paolo_status["document_count"] == 1
    assert fabio_status["scopes"] == paolo_status["scopes"]
    assert (fabio_state / "archive.sqlite3").is_file()
    assert (paolo_state / "archive.sqlite3").is_file()
    assert (fabio_state / "archive.sqlite3") != (paolo_state / "archive.sqlite3")
    assert stat.S_IMODE(fabio_state.stat().st_mode) == 0o700
    assert stat.S_IMODE((fabio_state / "archive.sqlite3").stat().st_mode) == 0o600
    assert stat.S_IMODE((fabio_state / "config.json").stat().st_mode) == 0o600


def test_incremental_refresh_skips_unchanged_files(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    result = archive_core.refresh_archive(state_dir=indexed_archive.state)

    assert result["discovered_files"] == 3
    assert result["unchanged_files"] == 3
    assert result["indexed_files"] == 0
    assert result["removed_files"] == 0


def test_repeating_same_configuration_is_idempotent(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    config_path = indexed_archive.state / "config.json"
    before = config_path.read_bytes()

    result = archive_core.configure_archive(
        indexed_archive.root,
        state_dir=indexed_archive.state,
    )

    assert result["index_requires_refresh"] is False
    assert config_path.read_bytes() == before


def test_reconfigure_recovers_after_top_level_scope_is_renamed(
    tmp_path: Path,
    archive_core: ModuleType,
) -> None:
    archive_root = tmp_path / "Studio"
    original_scope = archive_root / "Rossi"
    original_scope.mkdir(parents=True)
    state_dir = tmp_path / "state"
    archive_core.configure_archive(archive_root, state_dir=state_dir)
    original_scope.rename(archive_root / "Rossi-Srl")

    result = archive_core.configure_archive(archive_root, state_dir=state_dir)

    assert [scope["display_name"] for scope in result["scopes"]] == ["Rossi-Srl"]
    assert result["index_requires_refresh"] is True


def test_status_detects_new_top_level_scope_before_refresh(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    new_scope = indexed_archive.root / "Verdi"
    new_scope.mkdir()
    (new_scope / "nota.txt").write_text(
        "Nuova pratica aggiunta dopo la configurazione.",
        encoding="utf-8",
    )

    result = archive_core.studio_archive_status(state_dir=indexed_archive.state)

    assert result["scope_configuration_changed"] is True
    assert result["index_requires_refresh"] is True
    assert {scope["display_name"] for scope in result["scopes"]} == {
        "Bianchi",
        "Rossi",
        "Studio",
        "Verdi",
    }


def test_refresh_adopts_and_indexes_new_top_level_scope(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    new_scope = indexed_archive.root / "Verdi"
    new_scope.mkdir()
    (new_scope / "nota.txt").write_text(
        "Nuova pratica sulla liquidazione societaria.",
        encoding="utf-8",
    )

    result = archive_core.refresh_archive(state_dir=indexed_archive.state)

    assert result["scope_configuration_changed"] is True
    assert result["document_count"] == 4
    assert {scope["display_name"] for scope in result["scopes"]} == {
        "Bianchi",
        "Rossi",
        "Studio",
        "Verdi",
    }


def test_incremental_refresh_reindexes_changed_file(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    indexed_archive.source.write_text(
        "Verbale cessione quote aggiornato\nNuova clausola di prelazione.",
        encoding="utf-8",
    )

    result = archive_core.refresh_archive(state_dir=indexed_archive.state)

    assert result["indexed_files"] == 1
    assert result["unchanged_files"] == 2
    assert result["removed_files"] == 0


def test_refresh_hashes_same_size_same_timestamp_content_changes(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    original = "alpha evidence"
    replacement = "omega evidence"
    assert len(original) == len(replacement)
    indexed_archive.source.write_text(original, encoding="utf-8")
    archive_core.refresh_archive(state_dir=indexed_archive.state)
    metadata = indexed_archive.source.stat()
    indexed_archive.source.write_text(replacement, encoding="utf-8")
    os.utime(
        indexed_archive.source,
        ns=(metadata.st_atime_ns, metadata.st_mtime_ns),
    )

    refresh = archive_core.refresh_archive(state_dir=indexed_archive.state)
    candidates = archive_core.search_archive(
        "omega",
        scope_id=indexed_archive.scopes["Rossi"],
        state_dir=indexed_archive.state,
    )

    assert refresh["indexed_files"] == 1
    assert refresh["unchanged_files"] == 2
    assert candidates["result_count"] == 1


def test_incremental_refresh_removes_deleted_file(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    indexed_archive.source.unlink()

    result = archive_core.refresh_archive(state_dir=indexed_archive.state)

    assert result["removed_files"] == 1
    assert result["document_count"] == 2


def test_refresh_securely_removes_deleted_fts_text(
    tmp_path: Path,
    archive_core: ModuleType,
) -> None:
    archive_root = tmp_path / "Studio"
    client_root = archive_root / "Rossi"
    client_root.mkdir(parents=True)
    source = client_root / "secret.txt"
    unique_text = "supersecretuniqueterm"
    source.write_text(unique_text, encoding="utf-8")
    state_dir = tmp_path / "state"
    archive_core.configure_archive(archive_root, state_dir=state_dir)
    archive_core.refresh_archive(state_dir=state_dir)
    source.unlink()

    result = archive_core.refresh_archive(state_dir=state_dir)

    assert result["removed_files"] == 1
    assert (
        unique_text.encode("utf-8") not in (state_dir / "archive.sqlite3").read_bytes()
    )


def test_search_is_confined_to_exact_scope(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    result = archive_core.search_archive(
        "cessione quote",
        scope_id=indexed_archive.scopes["Bianchi"],
        state_dir=indexed_archive.state,
    )

    assert result["result_count"] == 0
    assert result["scope_id"] == indexed_archive.scopes["Bianchi"]


def test_search_and_open_return_verified_citation(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    candidates = archive_core.search_archive(
        "cessione quote",
        scope_id=indexed_archive.scopes["Rossi"],
        state_dir=indexed_archive.state,
    )
    source_id = candidates["results"][0]["source_id"]

    result = archive_core.open_archive_source(
        source_id,
        state_dir=indexed_archive.state,
    )

    assert result["source_verified"] is True
    assert result["relative_path"] == "Rossi/precedente.md"
    assert result["citation"] == "Rossi/precedente.md, lines 1-2"
    assert "cessione delle quote" in result["fragments"][0]["text"]


def test_open_rejects_source_changed_after_search(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    candidates = archive_core.search_archive(
        "cessione quote",
        scope_id=indexed_archive.scopes["Rossi"],
        state_dir=indexed_archive.state,
    )
    source_id = candidates["results"][0]["source_id"]
    indexed_archive.source.write_text(
        "The source changed after the search result was returned.",
        encoding="utf-8",
    )

    with pytest.raises(
        archive_core.SourceChangedError,
        match="changed after indexing",
    ):
        archive_core.open_archive_source(
            source_id,
            state_dir=indexed_archive.state,
        )


def test_status_search_and_open_leave_database_bytes_unchanged(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    database_path = indexed_archive.state / "archive.sqlite3"
    before = hashlib.sha256(database_path.read_bytes()).hexdigest()
    candidates = archive_core.search_archive(
        "cessione quote",
        scope_id=indexed_archive.scopes["Rossi"],
        state_dir=indexed_archive.state,
    )
    archive_core.open_archive_source(
        candidates["results"][0]["source_id"],
        state_dir=indexed_archive.state,
    )
    status_result = archive_core.studio_archive_status(state_dir=indexed_archive.state)

    after = hashlib.sha256(database_path.read_bytes()).hexdigest()

    assert status_result["configured"]
    assert after == before


def test_root_files_have_a_non_overlapping_scope(
    indexed_archive: SimpleNamespace,
    archive_core: ModuleType,
) -> None:
    root_scope = indexed_archive.scopes["Studio"]

    result = archive_core.search_archive(
        "controllo documentale",
        scope_id=root_scope,
        state_dir=indexed_archive.state,
    )

    assert result["result_count"] == 1
    assert result["results"][0]["relative_path"] == "procedura.txt"


def test_refresh_reports_skipped_paths_and_reasons(
    tmp_path: Path,
    archive_core: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    archive_root = tmp_path / "Studio"
    client_root = archive_root / "Rossi"
    client_root.mkdir(parents=True)
    (client_root / "legacy.msg").write_text("unsupported", encoding="utf-8")
    (client_root / "large.txt").write_text("too large", encoding="utf-8")
    monkeypatch.setattr(archive_core, "MAX_FILE_BYTES", 4)
    state_dir = tmp_path / "state"
    archive_core.configure_archive(archive_root, state_dir=state_dir)

    result = archive_core.refresh_archive(state_dir=state_dir)

    assert result["scan_issue_count"] == 2
    assert result["scan_issues_truncated"] is False
    assert result["scan_issues"] == [
        {
            "scope_id": result["scopes"][0]["scope_id"],
            "relative_path": "Rossi/large.txt",
            "reason": "file_size_limit_exceeded",
            "size_bytes": 9,
        },
        {
            "scope_id": result["scopes"][0]["scope_id"],
            "relative_path": "Rossi/legacy.msg",
            "reason": "unsupported_extension",
            "size_bytes": 11,
        },
    ]


def test_pdf_search_preserves_physical_page_locator(
    tmp_path: Path,
    archive_core: ModuleType,
) -> None:
    fitz = _fitz_or_skip()
    archive_root = tmp_path / "Studio"
    client_root = archive_root / "Rossi"
    client_root.mkdir(parents=True)
    pdf_path = client_root / "verbale.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "Verbale assemblea: deliberata trasformazione societaria con voto unanime.",
    )
    document.save(pdf_path)
    document.close()
    state_dir = tmp_path / "state"
    archive_core.configure_archive(archive_root, state_dir=state_dir)
    archive_core.refresh_archive(state_dir=state_dir)
    scope_id = archive_core.studio_archive_status(state_dir=state_dir)["scopes"][0][
        "scope_id"
    ]
    candidates = archive_core.search_archive(
        "trasformazione societaria",
        scope_id=scope_id,
        state_dir=state_dir,
    )

    result = archive_core.open_archive_source(
        candidates["results"][0]["source_id"],
        state_dir=state_dir,
    )

    assert result["locator_kind"] == "page"
    assert result["locator_value"] == "1"
    assert result["citation"] == "Rossi/verbale.pdf, p. 1"


def test_short_pdf_exposes_partial_status_and_ocr_limitation(
    tmp_path: Path,
    archive_core: ModuleType,
) -> None:
    fitz = _fitz_or_skip()
    archive_root = tmp_path / "Studio"
    client_root = archive_root / "Rossi"
    client_root.mkdir(parents=True)
    pdf_path = client_root / "codice.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "VAT123")
    document.save(pdf_path)
    document.close()
    state_dir = tmp_path / "state"
    archive_core.configure_archive(archive_root, state_dir=state_dir)
    archive_core.refresh_archive(state_dir=state_dir)
    scope_id = archive_core.studio_archive_status(state_dir=state_dir)["scopes"][0][
        "scope_id"
    ]

    candidates = archive_core.search_archive(
        "VAT123",
        scope_id=scope_id,
        state_dir=state_dir,
    )
    opened = archive_core.open_archive_source(
        candidates["results"][0]["source_id"],
        state_dir=state_dir,
    )

    assert candidates["results"][0]["document_status"] == "partial"
    assert candidates["results"][0]["needs_ocr"] is True
    assert candidates["results"][0]["limitations"] == ["page_1_no_extractable_text"]
    assert opened["document_status"] == "partial"
    assert opened["needs_ocr"] is True
    assert opened["limitations"] == ["page_1_no_extractable_text"]


def test_empty_ocr_result_does_not_erase_short_native_pdf_text(
    tmp_path: Path,
    archive_core: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fitz = _fitz_or_skip()
    archive_root = tmp_path / "Studio"
    client_root = archive_root / "Rossi"
    client_root.mkdir(parents=True)
    pdf_path = client_root / "codice.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "VAT123")
    document.save(pdf_path)
    document.close()
    fake_ocr = ModuleType("vera_ocr")
    fake_ocr.extract_text_from_image_bytes = lambda *args, **kwargs: SimpleNamespace(  # type: ignore[attr-defined]
        network_used=False,
        warnings=("no_text_detected",),
        status="ok",
        text="",
    )
    monkeypatch.setitem(sys.modules, "vera_ocr", fake_ocr)
    state_dir = tmp_path / "state"
    archive_core.configure_archive(archive_root, state_dir=state_dir)
    archive_core.refresh_archive(enable_ocr=True, state_dir=state_dir)
    scope_id = archive_core.studio_archive_status(state_dir=state_dir)["scopes"][0][
        "scope_id"
    ]

    result = archive_core.search_archive(
        "VAT123",
        scope_id=scope_id,
        state_dir=state_dir,
    )

    assert result["result_count"] == 1
    assert result["results"][0]["needs_ocr"] is True
    assert "page_1_no_text_detected" in result["results"][0]["limitations"]


def test_context_fragments_each_include_their_own_citation(
    tmp_path: Path,
    archive_core: ModuleType,
) -> None:
    archive_root = tmp_path / "Studio"
    client_root = archive_root / "Rossi"
    client_root.mkdir(parents=True)
    lines = [f"ordinary evidence line {number}" for number in range(1, 121)]
    lines.append("unique closing evidence")
    (client_root / "long.txt").write_text("\n".join(lines), encoding="utf-8")
    state_dir = tmp_path / "state"
    archive_core.configure_archive(archive_root, state_dir=state_dir)
    archive_core.refresh_archive(state_dir=state_dir)
    scope_id = archive_core.studio_archive_status(state_dir=state_dir)["scopes"][0][
        "scope_id"
    ]
    candidates = archive_core.search_archive(
        "unique closing",
        scope_id=scope_id,
        state_dir=state_dir,
    )

    result = archive_core.open_archive_source(
        candidates["results"][0]["source_id"],
        context_chunks=1,
        state_dir=state_dir,
    )

    assert [fragment["citation"] for fragment in result["fragments"]] == [
        "Rossi/long.txt, lines 1-120",
        "Rossi/long.txt, lines 121",
    ]


def test_office_and_email_sources_keep_structural_locators(
    tmp_path: Path,
    archive_core: ModuleType,
) -> None:
    from docx import Document
    from openpyxl import Workbook

    archive_root = tmp_path / "Studio"
    client_root = archive_root / "Rossi"
    client_root.mkdir(parents=True)

    document = Document()
    document.add_paragraph("Mandato professionale per assistenza societaria.")
    document.save(client_root / "mandato.docx")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Dati"
    worksheet.append(["Voce", "Importo"])
    worksheet.append(["Acconto imposte", 1250])
    workbook.save(client_root / "acconti.xlsx")
    workbook.close()

    message = EmailMessage()
    message["Subject"] = "Conferma assemblea"
    message["From"] = "cliente@example.invalid"
    message["To"] = "studio@example.invalid"
    message.set_content("Confermo la data della delibera assembleare.")
    (client_root / "conferma.eml").write_bytes(message.as_bytes())

    state_dir = tmp_path / "state"
    archive_core.configure_archive(archive_root, state_dir=state_dir)
    archive_core.refresh_archive(state_dir=state_dir)
    scope_id = archive_core.studio_archive_status(state_dir=state_dir)["scopes"][0][
        "scope_id"
    ]

    docx_candidates = archive_core.search_archive(
        "mandato professionale",
        scope_id=scope_id,
        state_dir=state_dir,
    )
    xlsx_candidates = archive_core.search_archive(
        "acconto imposte",
        scope_id=scope_id,
        state_dir=state_dir,
    )
    eml_candidates = archive_core.search_archive(
        "delibera assembleare",
        scope_id=scope_id,
        state_dir=state_dir,
    )
    docx_source = archive_core.open_archive_source(
        docx_candidates["results"][0]["source_id"],
        state_dir=state_dir,
    )
    xlsx_source = archive_core.open_archive_source(
        xlsx_candidates["results"][0]["source_id"],
        state_dir=state_dir,
    )
    eml_source = archive_core.open_archive_source(
        eml_candidates["results"][0]["source_id"],
        state_dir=state_dir,
    )

    assert docx_source["locator_kind"] == "paragraphs"
    assert docx_source["citation"] == "Rossi/mandato.docx, paragraphs 1"
    assert xlsx_source["locator_kind"] == "sheet"
    assert xlsx_source["locator_value"] == "Dati!rows 1-2"
    assert eml_source["locator_kind"] == "message lines"
    assert eml_source["relative_path"] == "Rossi/conferma.eml"


def test_enabling_ocr_retries_unchanged_scan_without_model_download(
    tmp_path: Path,
    archive_core: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    image_module = pytest.importorskip(
        "PIL.Image",
        reason="Pillow creates the local scan fixture.",
    )
    archive_root = tmp_path / "Studio"
    client_root = archive_root / "Rossi"
    client_root.mkdir(parents=True)
    image_module.new("RGB", (120, 60), color="white").save(
        client_root / "scansione.png"
    )
    state_dir = tmp_path / "state"
    archive_core.configure_archive(archive_root, state_dir=state_dir)
    initial = archive_core.refresh_archive(
        enable_ocr=False,
        state_dir=state_dir,
    )
    calls: list[tuple[str, bool]] = []
    fake_ocr = ModuleType("vera_ocr")

    def extract_text_from_image_bytes(
        image_bytes: bytes,
        *,
        language: str,
        allow_model_download: bool,
    ) -> SimpleNamespace:
        assert image_bytes
        calls.append((language, allow_model_download))
        return SimpleNamespace(
            network_used=False,
            warnings=(),
            status="ok",
            text="Scansione locale della dichiarazione fiscale verificabile.",
        )

    fake_ocr.extract_text_from_image_bytes = extract_text_from_image_bytes  # type: ignore[attr-defined]
    monkeypatch.setitem(sys.modules, "vera_ocr", fake_ocr)

    result = archive_core.refresh_archive(
        enable_ocr=True,
        state_dir=state_dir,
    )

    assert initial["needs_ocr_files"] == 1
    assert result["indexed_files"] == 1
    assert result["unchanged_files"] == 0
    assert result["needs_ocr_files"] == 0
    assert calls == [("it", False)]


def test_zero_chunk_document_is_named_in_all_evidence_responses(
    tmp_path: Path,
    archive_core: ModuleType,
) -> None:
    image_module = pytest.importorskip(
        "PIL.Image",
        reason="Pillow creates the local scan fixture.",
    )
    archive_root = tmp_path / "Studio"
    client_root = archive_root / "Rossi"
    client_root.mkdir(parents=True)
    image_module.new("RGB", (120, 60), color="white").save(
        client_root / "scansione.png"
    )
    (client_root / "nota.txt").write_text(
        "Precedente verificabile sul ravvedimento.",
        encoding="utf-8",
    )
    state_dir = tmp_path / "state"
    configured = archive_core.configure_archive(archive_root, state_dir=state_dir)
    scope_id = configured["scopes"][0]["scope_id"]

    refreshed = archive_core.refresh_archive(state_dir=state_dir)
    status = archive_core.studio_archive_status(state_dir=state_dir)
    candidates = archive_core.search_archive(
        "ravvedimento",
        scope_id=scope_id,
        state_dir=state_dir,
    )
    opened = archive_core.open_archive_source(
        candidates["results"][0]["source_id"],
        state_dir=state_dir,
    )

    expected_issue = {
        "scope_id": scope_id,
        "relative_path": "Rossi/scansione.png",
        "document_status": "partial",
        "needs_ocr": True,
        "limitations": ["ocr_disabled"],
        "chunk_count": 0,
    }
    assert refreshed["document_issue_count"] == 1
    assert refreshed["document_issues"] == [expected_issue]
    assert refreshed["document_issues_truncated"] is False
    assert status["document_issue_count"] == 1
    assert status["document_issues"] == [expected_issue]
    assert status["document_issues_truncated"] is False
    assert candidates["document_issue_count"] == 1
    assert candidates["document_issues"] == [expected_issue]
    assert candidates["document_issues_truncated"] is False
    assert opened["document_issue_count"] == 1
    assert opened["document_issues"] == [expected_issue]
    assert opened["document_issues_truncated"] is False


def test_scoped_search_and_open_hide_other_scope_issues(
    tmp_path: Path,
    archive_core: ModuleType,
) -> None:
    archive_root = tmp_path / "Studio"
    bianchi_root = archive_root / "Bianchi"
    rossi_root = archive_root / "Rossi"
    bianchi_root.mkdir(parents=True)
    rossi_root.mkdir()
    (bianchi_root / "nota.txt").write_text(
        "Evidenza Bianchi verificabile.",
        encoding="utf-8",
    )
    (rossi_root / "scansione.png").write_bytes(b"not-decoded-with-ocr-disabled")
    (rossi_root / "legacy.msg").write_text("unsupported", encoding="utf-8")
    state_dir = tmp_path / "state"
    configured = archive_core.configure_archive(archive_root, state_dir=state_dir)
    scopes = {
        scope["display_name"]: scope["scope_id"] for scope in configured["scopes"]
    }
    archive_core.refresh_archive(state_dir=state_dir)

    scoped = archive_core.search_archive(
        "Evidenza Bianchi",
        scope_id=scopes["Bianchi"],
        state_dir=state_dir,
    )
    opened = archive_core.open_archive_source(
        scoped["results"][0]["source_id"],
        state_dir=state_dir,
    )
    studio_wide = archive_core.search_archive(
        "Evidenza Bianchi",
        scope_id="all",
        state_dir=state_dir,
    )

    assert scoped["results"][0]["scope_id"] == scopes["Bianchi"]
    assert scoped["scan_issue_count"] == 0
    assert scoped["scan_issues"] == []
    assert scoped["document_issue_count"] == 0
    assert scoped["document_issues"] == []
    assert opened["scope_id"] == scopes["Bianchi"]
    assert opened["scan_issue_count"] == 0
    assert opened["scan_issues"] == []
    assert opened["document_issue_count"] == 0
    assert opened["document_issues"] == []
    assert studio_wide["scan_issue_count"] == 1
    assert studio_wide["scan_issues"][0]["relative_path"] == "Rossi/legacy.msg"
    assert studio_wide["document_issue_count"] == 1
    assert studio_wide["document_issues"][0]["relative_path"] == "Rossi/scansione.png"


@pytest.mark.skipif(os.name != "posix", reason="POSIX mode bits are required.")
def test_status_rejects_index_configuration_with_broad_permissions(
    tmp_path: Path,
    archive_core: ModuleType,
) -> None:
    archive_root = tmp_path / "Studio"
    (archive_root / "Rossi").mkdir(parents=True)
    state_dir = tmp_path / "state"
    archive_core.configure_archive(archive_root, state_dir=state_dir)
    (state_dir / "config.json").chmod(0o644)

    with pytest.raises(
        archive_core.ArchiveError,
        match="configuration must not be accessible",
    ):
        archive_core.studio_archive_status(state_dir=state_dir)


def test_symlinked_source_is_not_indexed(
    tmp_path: Path,
    archive_core: ModuleType,
) -> None:
    archive_root = tmp_path / "Studio"
    client_root = archive_root / "Rossi"
    client_root.mkdir(parents=True)
    outside = tmp_path / "outside.txt"
    outside.write_text("This file must not enter the archive index.", encoding="utf-8")
    linked = client_root / "linked.txt"
    try:
        linked.symlink_to(outside)
    except OSError:
        pytest.skip("Symbolic links are unavailable in this environment.")
    state_dir = tmp_path / "state"
    archive_core.configure_archive(archive_root, state_dir=state_dir)

    result = archive_core.refresh_archive(state_dir=state_dir)

    assert result["document_count"] == 0
    assert result["chunk_count"] == 0
    assert result["scan_issue_count"] == 1
    assert result["scan_issues"] == [
        {
            "scope_id": result["scopes"][0]["scope_id"],
            "relative_path": "Rossi/linked.txt",
            "reason": "symbolic_link_not_followed",
            "size_bytes": None,
        }
    ]


def test_mcp_lists_fourteen_strict_local_tools(tmp_path: Path) -> None:
    response = _mcp_request(
        {
            "jsonrpc": "2.0",
            "id": 1,
            "method": "tools/list",
            "params": {},
        },
        state_dir=tmp_path / "state",
    )

    tools = response["result"]["tools"]
    assert {tool["name"] for tool in tools} == {
        "studio_archive_status",
        "list_studio_archive_clients",
        "get_studio_client_folder",
        "create_studio_archive_client",
        "import_studio_client_document",
        "list_studio_client_engagements",
        "prepare_studio_client_workflow",
        "configure_studio_archive",
        "refresh_studio_archive",
        "search_studio_archive",
        "open_studio_archive_source",
        "configure_studio_archive_client",
        "plan_studio_archive_gmail_search",
        "match_studio_archive_email",
    }
    assert all(tool["inputSchema"]["additionalProperties"] is False for tool in tools)
    assert all(tool["annotations"]["openWorldHint"] is False for tool in tools)
    tool_by_name = {tool["name"]: tool for tool in tools}
    assert (
        tool_by_name["configure_studio_archive_client"]["annotations"]["idempotentHint"]
        is False
    )


def test_mcp_rejects_non_object_json_request(tmp_path: Path) -> None:
    environment = os.environ.copy()
    environment["VERA_STUDIO_ARCHIVE_PYTHON"] = sys.executable
    environment["VERA_STUDIO_ARCHIVE_STATE_DIR"] = str(tmp_path / "state")

    completed = subprocess.run(
        [_node_executable(), str(MCP_SERVER_PATH), "--stdio"],
        cwd=COMPONENT_ROOT,
        env=environment,
        input="null\n",
        text=True,
        capture_output=True,
        check=True,
        timeout=30,
    )

    response = json.loads(completed.stdout)
    assert response == {
        "jsonrpc": "2.0",
        "id": None,
        "error": {"code": -32600, "message": "invalid request"},
    }


def test_mcp_configure_refresh_search_and_open(tmp_path: Path) -> None:
    archive_root = tmp_path / "Shared Studio"
    client_root = archive_root / "Rossi"
    client_root.mkdir(parents=True)
    (client_root / "memo.txt").write_text(
        "Precedente verificato sulla cessione di partecipazioni.",
        encoding="utf-8",
    )
    state_dir = tmp_path / "Fabio private"
    configure = _mcp_tool(
        "configure_studio_archive",
        {"archive_root": str(archive_root)},
        state_dir=state_dir,
    )
    scope_id = configure["structuredContent"]["scopes"][0]["scope_id"]
    registration = _mcp_tool(
        "configure_studio_archive_client",
        {"scope_id": scope_id, "legal_names": ["Rossi SRL"]},
        state_dir=state_dir,
    )
    client_id = registration["structuredContent"]["client"]["client_id"]
    refresh = _mcp_tool(
        "refresh_studio_archive",
        {},
        state_dir=state_dir,
    )
    client_folder = _mcp_tool(
        "get_studio_client_folder",
        {"client_id": client_id},
        state_dir=state_dir,
    )
    search = _mcp_tool(
        "search_studio_archive",
        {"query": "cessione partecipazioni", "scope_id": scope_id},
        state_dir=state_dir,
    )
    source_id = search["structuredContent"]["results"][0]["source_id"]

    result = _mcp_tool(
        "open_studio_archive_source",
        {"source_id": source_id},
        state_dir=state_dir,
    )

    assert configure["isError"] is False
    assert refresh["structuredContent"]["document_count"] == 1
    assert client_folder["structuredContent"]["client_folder"]["client_root"] == str(
        client_root
    )
    assert client_folder["structuredContent"]["source_archive_mutated"] is False
    assert search["structuredContent"]["result_count"] == 1
    assert result["isError"] is False
    assert result["structuredContent"]["source_verified"] is True
    assert result["structuredContent"]["citation"] == "Rossi/memo.txt, lines 1"


def test_mcp_configures_plans_and_matches_client_scoped_gmail(tmp_path: Path) -> None:
    archive_root = tmp_path / "Shared Studio"
    (archive_root / "Rossi").mkdir(parents=True)
    state_dir = tmp_path / "Fabio private"
    configured = _mcp_tool(
        "configure_studio_archive",
        {"archive_root": str(archive_root)},
        state_dir=state_dir,
    )
    scope_id = configured["structuredContent"]["scopes"][0]["scope_id"]
    client = _mcp_tool(
        "configure_studio_archive_client",
        {
            "scope_id": scope_id,
            "email_addresses": ["amministrazione@rossi.it"],
            "legal_names": ["Rossi SRL"],
            "tax_identifiers": ["01234567890"],
        },
        state_dir=state_dir,
    )
    plan = _mcp_tool(
        "plan_studio_archive_gmail_search",
        {"scope_id": scope_id, "topic": "rateazione INPS"},
        state_dir=state_dir,
    )

    result = _mcp_tool(
        "match_studio_archive_email",
        {
            "header_addresses": ["Rossi <amministrazione@rossi.it>"],
            "headers_complete": True,
            "expected_scope_id": scope_id,
        },
        state_dir=state_dir,
    )

    assert client["isError"] is False
    assert client["structuredContent"]["gmail_credentials_stored"] is False
    assert plan["structuredContent"]["connector"] == "gmail"
    assert plan["structuredContent"]["gmail_connector_called"] is False
    assert result["isError"] is False
    assert result["structuredContent"]["routing_status"] == "exact"
    assert result["structuredContent"]["may_use_in_scoped_answer"] is True


def test_mcp_lists_and_rebinds_orphaned_client_profile(tmp_path: Path) -> None:
    archive_root = tmp_path / "Shared Studio"
    original_folder = archive_root / "Rossi"
    original_folder.mkdir(parents=True)
    state_dir = tmp_path / "Fabio private"
    configured = _mcp_tool(
        "configure_studio_archive",
        {"archive_root": str(archive_root)},
        state_dir=state_dir,
    )
    original_scope_id = configured["structuredContent"]["scopes"][0]["scope_id"]
    _mcp_tool(
        "configure_studio_archive_client",
        {
            "scope_id": original_scope_id,
            "email_addresses": ["amministrazione@rossi.it"],
        },
        state_dir=state_dir,
    )
    original_folder.rename(archive_root / "Rossi Nuovo")
    refreshed = _mcp_tool(
        "refresh_studio_archive",
        {},
        state_dir=state_dir,
    )
    renamed_scope_id = refreshed["structuredContent"]["scopes"][0]["scope_id"]

    profiles = _mcp_tool(
        "list_studio_archive_clients",
        {},
        state_dir=state_dir,
    )
    rebound = _mcp_tool(
        "configure_studio_archive_client",
        {
            "scope_id": renamed_scope_id,
            "replace_orphaned_scope_id": original_scope_id,
        },
        state_dir=state_dir,
    )
    plan = _mcp_tool(
        "plan_studio_archive_gmail_search",
        {"scope_id": renamed_scope_id},
        state_dir=state_dir,
    )

    assert profiles["structuredContent"]["orphaned_profile_count"] == 1
    assert (
        profiles["structuredContent"]["orphaned_profiles"][0]["scope_id"]
        == original_scope_id
    )
    assert rebound["structuredContent"]["status"] == "rebound"
    assert plan["structuredContent"]["profile_status"] == "configured"


def test_mcp_rejects_gmail_credentials_without_writing_them(tmp_path: Path) -> None:
    archive_root = tmp_path / "Shared Studio"
    (archive_root / "Rossi").mkdir(parents=True)
    state_dir = tmp_path / "Fabio private"
    configured = _mcp_tool(
        "configure_studio_archive",
        {"archive_root": str(archive_root)},
        state_dir=state_dir,
    )
    scope_id = configured["structuredContent"]["scopes"][0]["scope_id"]

    result = _mcp_tool(
        "configure_studio_archive_client",
        {
            "scope_id": scope_id,
            "email_addresses": ["amministrazione@rossi.it"],
            "password": "must-not-be-stored",
        },
        state_dir=state_dir,
    )

    assert result["isError"] is True
    assert (
        "Unknown tool argument: password"
        in result["structuredContent"]["error"]["message"]
    )
    assert not (state_dir / "client-identities.json").exists()


def test_vera_registers_archive_as_embedded_workflow() -> None:
    components = json.loads(
        (ROOT / "plugins" / "vera" / "components.json").read_text(encoding="utf-8")
    )
    vera_mcp = json.loads(
        (ROOT / "plugins" / "vera" / ".mcp.json").read_text(encoding="utf-8")
    )
    manifest = json.loads(
        (ROOT / "plugins" / "vera" / ".codex-plugin" / "plugin.json").read_text(
            encoding="utf-8"
        )
    )

    assert "studio-archive" in components["plugins"]
    assert "studio-archive" not in components["workflow_roles"]
    assert vera_mcp["mcpServers"]["veraStudioArchive"]["args"][-1] == "studio-archive"
    assert (
        ROOT / "plugins" / "vera" / "skills" / "studio-archive" / "SKILL.md"
    ).is_file()
