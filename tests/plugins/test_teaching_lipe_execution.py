"""Run the authored LIPE sources; confirmations are test-only interpretations.

This verifies the current local engine, not learner or professional acceptance.
No reviewed case, calculated answer or simulated approval ships in the lesson.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from tests.plugins._teaching_release import record_native_check
from tests.plugins.test_lipe import (  # noqa: F401
    case_data,
    lipe_script_imports,
    main,
    read_json,
)

__all__: list[str] = []
ROOT = Path(__file__).resolve().parents[2]


def _interpret(source: Path, *, updated: bool = False) -> dict:
    """Bind test-only interpretation to the exact authored physical source."""
    case = case_data()
    case["case_id"] = "fictional-bottega-aurora"
    case["sources"][0].update(
        path=source.name, sha256=hashlib.sha256(source.read_bytes()).hexdigest()
    )
    if updated:
        register = next(
            r for r in case["registers"] if r["register_id"] == "PURCHASES-6"
        )
        register.update(printed_base="750.00", printed_tax="165.00")
        register["rows"][0].update(base="750.00", tax="165.00", deductible_tax="165.00")
        quote = "PURCHASES-6: base 750.00; tax 165.00."
        register["evidence"]["quote"] = quote
        register["rows"][0]["evidence"]["quote"] = quote
        case["modules"][2]["defer_small_debit"] = False
    return case


def _run(
    case: dict, source: Path, output: Path, expected_code: int
) -> tuple[Path, dict]:
    """Persist and execute a synthetic interpretation through the actual CLI."""
    case_path = output.parent / "interpretation.json"
    case_path.write_text(json.dumps(case), encoding="utf-8")
    previous = set(output.glob("lipe-*"))
    assert (
        main(
            [
                "calculate",
                "--case",
                str(case_path),
                "--source-root",
                str(source.parent),
                "--output",
                str(output),
            ]
        )
        == expected_code
    )
    (folder,) = set(output.glob("lipe-*")) - previous
    return folder, read_json(folder / "result.json")


@pytest.mark.parametrize("phase", ["demo", "practice"])
def test_lipe_kit_runs_reviewed_sources_and_preserves_revision(
    tmp_path, monkeypatch, record_property, phase
):
    monkeypatch.syspath_prepend(str(ROOT / "plugins/_shared/vendor/modules"))
    from courseware.library import CourseLibrary

    kit = CourseLibrary(ROOT / "plugins/vera", {"lipe"}).render(
        "lipe", "it", tmp_path / "kit"
    )
    source = next(
        Path(p) for p in kit["source_files"] if Path(p).name == "registri.txt"
    )
    case = _interpret(source)
    case["mappings"][0]["review"]["status"] = "PROPOSED"
    output = tmp_path / "results"
    _, blocked = _run(case, source, output, 2)
    assert blocked["status"] == "BLOCKED"
    assert blocked["modules"] == []

    original_folder, result = _run(_interpret(source), source, output, 0)
    assert result["status"] == "DRAFT_FOR_REVIEW"
    assert [m["rows"]["vp14_debit"] for m in result["modules"]] == ["110.00"] * 3
    original = {p: p.read_bytes() for p in original_folder.iterdir() if p.is_file()}
    folder = original_folder

    if phase == "practice":
        changed = next(
            Path(p)
            for p in kit["practice_files"]
            if Path(p).name == "registri-rettificati.txt"
        )
        proposal = _interpret(changed, updated=True)
        proposal["registers"][5]["rows"][0]["review"]["status"] = "PROPOSED"
        _, reopened = _run(proposal, changed, output, 2)
        assert reopened["modules"] == []
        folder, result = _run(_interpret(changed, updated=True), changed, output, 0)
        assert folder != original_folder
        assert {p: p.read_bytes() for p in original} == original
        assert [m["rows"]["vp14_debit"] for m in result["modules"]] == [
            "110.00",
            "110.00",
            "55.00",
        ]
        assert result["modules"][2]["rows"]["vp3"] == "750.00"
        assert result["modules"][2]["rows"]["vp5"] == "165.00"
        assert result["modules"][2]["payment_status"] == "DIFFERENCE_TO_REVIEW"
        assert {f["code"] for f in result["findings"]} >= {
            "PAYMENT_DIFFERENCE",
            "REGISTER_LIQUIDATION_DIFFERENCE",
        }

    assert result["export_status"] == "NOT_AUTHORIZED"
    workbook = load_workbook(folder / "workpaper.xlsx", read_only=True)
    assert {"Riconciliazione", "VP", "F24"}.issubset(workbook.sheetnames)
    workbook.close()
    assert "LIPE" in PdfReader(folder / "summary.pdf").pages[0].extract_text()
    assert (folder / "review-request.md").stat().st_size > 0
    artifacts = [
        {
            "name": name,
            "sha256": hashlib.sha256((folder / name).read_bytes()).hexdigest(),
        }
        for name in (
            "result.json",
            "workpaper.xlsx",
            "summary.pdf",
            "vp.csv",
            "review-request.md",
        )
    ]
    (tmp_path / "reviewed-artifacts.json").write_text(
        json.dumps(artifacts), encoding="utf-8"
    )
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="lipe",
        language="it",
        phase=phase,
    )
