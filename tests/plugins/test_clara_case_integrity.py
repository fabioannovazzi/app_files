from __future__ import annotations

import importlib.util
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import pytest

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture
def core(monkeypatch: pytest.MonkeyPatch) -> Any:
    scripts = ROOT / "plugins/clara/scripts"
    monkeypatch.syspath_prepend(str(scripts))
    spec = importlib.util.spec_from_file_location(
        "clara_integrity_core", scripts / "advisor_case_core.py"
    )
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    monkeypatch.setitem(sys.modules, spec.name, module)
    spec.loader.exec_module(module)
    return module


def case(core: Any, path: Path) -> Path:
    core.initialize_case(
        path,
        client="Synthetic",
        project=path.name,
        objective="Integrity",
        audience="Reviewer",
    )
    return path


def exchange(core: Any, root: Path) -> tuple[Path, dict[str, bytes]]:
    source = case(core, root / "source")
    target = case(core, root / "target")
    core.ingest_note_text(source, title="Note", text="Original evidence")
    archive_path = core.export_case_update(source).package_path
    with ZipFile(archive_path) as archive:
        return target, {name: archive.read(name) for name in archive.namelist()}


def write_exchange(
    path: Path, members: dict[str, bytes], payload: dict[str, Any]
) -> Path:
    with ZipFile(path, "w") as archive:
        for name, content in members.items():
            archive.writestr(
                name, json.dumps(payload) if name == "case_update.json" else content
            )
    return path


@pytest.mark.parametrize(
    "identifier",
    [
        "../../escaped",
        "/tmp/escaped",
        r"C:\escaped",
        r"..\escaped",
        ".",
        "..",
        "a/b",
        "CON",
        "aux",
        "trailing.",
        "",
    ],
)
def test_exchange_rejects_unsafe_root_before_mutation(
    core: Any, tmp_path: Path, identifier: str
) -> None:
    target, members = exchange(core, tmp_path)
    payload = json.loads(members["case_update.json"])
    payload["exchange_id"] = identifier
    archive = write_exchange(tmp_path / "bad.zip", members, payload)
    before = {
        p.relative_to(target): p.read_bytes() for p in target.rglob("*") if p.is_file()
    }

    with pytest.raises(core.CaseWorkspaceError, match="unsafe exchange_id"):
        core.import_case_update(target, archive)

    assert {
        p.relative_to(target): p.read_bytes() for p in target.rglob("*") if p.is_file()
    } == before
    assert not (tmp_path / "escaped").exists()


def test_exchange_rejects_symlink_destination_and_preserves_external_file(
    core: Any, tmp_path: Path
) -> None:
    target, members = exchange(core, tmp_path)
    payload = json.loads(members["case_update.json"])
    outside = tmp_path / "outside"
    outside.mkdir()
    sentinel = outside / "keep.txt"
    sentinel.write_text("Do not touch")
    (target / "exchange_imports").symlink_to(outside, target_is_directory=True)
    archive = write_exchange(tmp_path / "bad.zip", members, payload)

    with pytest.raises(core.CaseWorkspaceError, match="symbolic-link"):
        core.import_case_update(target, archive)

    assert sentinel.read_text() == "Do not touch"
    assert list(outside.iterdir()) == [sentinel]


def test_exchange_rejects_symbolic_link_member(core: Any, tmp_path: Path) -> None:
    target, members = exchange(core, tmp_path)
    payload = json.loads(members["case_update.json"])
    archive = write_exchange(tmp_path / "bad.zip", members, payload)
    with ZipFile(archive, "a") as package:
        info = ZipInfo("link")
        info.create_system = 3
        info.external_attr = 0o120777 << 16
        package.writestr(info, "../../outside")

    with pytest.raises(core.CaseWorkspaceError, match="symbolic-link"):
        core.import_case_update(target, archive)

    assert not (target / "exchange_imports").exists()


def test_exchange_rejects_duplicate_members(core: Any, tmp_path: Path) -> None:
    target, members = exchange(core, tmp_path)
    archive = write_exchange(
        tmp_path / "bad.zip", members, json.loads(members["case_update.json"])
    )
    with (
        ZipFile(archive, "a") as package,
        pytest.warns(UserWarning, match="Duplicate name"),
    ):
        package.writestr("case_update.json", members["case_update.json"])

    with pytest.raises(core.CaseWorkspaceError, match="duplicate"):
        core.import_case_update(target, archive)

    assert not (target / "exchange_imports").exists()


@pytest.mark.parametrize("second_title", ["Same title", "Same-title"])
def test_same_second_note_captures_preserve_both_sources(
    core: Any, tmp_path: Path, second_title: str
) -> None:
    target = case(core, tmp_path / "case")
    now = datetime(2026, 9, 6, 12, tzinfo=timezone.utc)
    first = core.ingest_note_text(
        target, title="Same title", text="First evidence", now=now
    )
    original = Path(first["path"]).read_bytes()

    second = core.ingest_note_text(
        target, title=second_title, text="Second evidence", now=now
    )

    assert second["id"] != first["id"]
    assert second["path"] != first["path"]
    assert Path(first["path"]).read_bytes() == original
    assert "Second evidence" in Path(second["path"]).read_text()


def test_interrupted_capture_recovers_previous_case_revision(
    core: Any, tmp_path: Path
) -> None:
    import subprocess

    target = case(core, tmp_path / "case")
    program = """
import os, sys
from pathlib import Path
sys.path.insert(0, sys.argv[1])
import advisor_case_core as core
core._touch_manifest = lambda *args: os._exit(23)
core.ingest_note_text(Path(sys.argv[2]), title='Interrupted', text='Must not leave partial state')
"""
    child = subprocess.run(
        [
            sys.executable,
            "-c",
            program,
            str(ROOT / "plugins/clara/scripts"),
            str(target),
        ],
        check=False,
        timeout=20,
    )
    assert child.returncode == 23
    assert (target / ".clara-transaction").exists()

    core.initialize_case(
        target,
        client="Synthetic",
        project=target.name,
        objective="Integrity",
        audience="Reviewer",
    )

    assert (
        json.loads((target / "material_registry.json").read_text())["materials"] == []
    )
    assert not list((target / "notes").glob("*.md"))
    assert not (target / ".clara-transaction").exists()
    assert core.validate_case_workspace(target) == []


def test_concurrent_captures_keep_distinct_registry_entries(
    core: Any, tmp_path: Path
) -> None:
    import subprocess

    target = case(core, tmp_path / "case")
    program = """
import sys
from pathlib import Path
sys.path.insert(0, sys.argv[1])
import advisor_case_core as core
core.ingest_note_text(Path(sys.argv[2]), title='Concurrent', text=sys.argv[3])
"""
    children = [
        subprocess.Popen(
            [
                sys.executable,
                "-c",
                program,
                str(ROOT / "plugins/clara/scripts"),
                str(target),
                f"Evidence {i}",
            ]
        )
        for i in range(4)
    ]

    returncodes = [child.wait(timeout=20) for child in children]

    assert returncodes == [0, 0, 0, 0]
    materials = json.loads((target / "material_registry.json").read_text())["materials"]
    assert len(materials) == 4
    assert len({item["id"] for item in materials}) == 4
    assert len({item["path"] for item in materials}) == 4
    assert core.validate_case_workspace(target) == []


def test_page_reference_in_substantive_sentence_is_not_scaffolding(core: Any) -> None:
    assert (
        core.audit_human_visible_document_text("See page 12 for the signed assumption.")
        == []
    )


def test_docx_index_records_partial_coverage_and_preserves_omitted_source(
    core: Any,
    tmp_path: Path,
) -> None:
    from docx import Document

    case_dir = case(core, tmp_path / "case")
    source = tmp_path / "long.docx"
    document = Document()
    for index in range(12):
        document.add_paragraph(f"Opening paragraph {index}")
    document.add_paragraph("Critical contradiction after the preview")
    document.save(source)
    original_bytes = source.read_bytes()

    material = core.register_material(case_dir, source)

    assert (
        material["source_metadata"]["preview_coverage"]["semantic_review_performed"]
        is False
    )
    assert material["source_metadata"]["preview_coverage"]["scope"].startswith(
        "first 12 body paragraphs"
    )
    assert "Critical contradiction" not in material["summary"]
    assert source.read_bytes() == original_bytes


@pytest.mark.parametrize(
    "archive_path", ["../../outside", "/tmp/outside", r"C:\outside", r"..\outside"]
)
def test_exchange_rejects_malicious_lineage_path_before_mutation(
    core: Any, tmp_path: Path, archive_path: str
) -> None:
    target, members = exchange(core, tmp_path)
    payload = json.loads(members["case_update.json"])
    payload["included_lineage_files"] = [{"archive_path": archive_path}]
    archive = write_exchange(tmp_path / "bad-lineage.zip", members, payload)
    sentinel = tmp_path / "outside"
    sentinel.write_bytes(b"unrelated evidence")
    before = {
        p.relative_to(target): p.read_bytes() for p in target.rglob("*") if p.is_file()
    }

    with pytest.raises(core.CaseWorkspaceError, match="unsafe exchange path"):
        core.import_case_update(target, archive)

    assert sentinel.read_bytes() == b"unrelated evidence"
    assert {
        p.relative_to(target): p.read_bytes() for p in target.rglob("*") if p.is_file()
    } == before


def test_identical_note_replay_is_a_distinct_capture_with_original_timestamp(
    core: Any, tmp_path: Path
) -> None:
    target = case(core, tmp_path / "case")
    captured_at = datetime(2026, 9, 6, 12, 30, 45, tzinfo=timezone.utc)
    first = core.ingest_note_text(
        target, title="Replay", text="Same evidence", now=captured_at
    )

    repeated = core.ingest_note_text(
        target, title="Replay", text="Same evidence", now=captured_at
    )

    assert first["id"] != repeated["id"]
    assert first["path"] != repeated["path"]
    assert Path(first["path"]).read_bytes() == Path(repeated["path"]).read_bytes()
    assert "Captured: 2026-09-06T12:30:45+00:00" in Path(repeated["path"]).read_text()


@pytest.mark.parametrize("limit_kind", ["archive", "expanded", "members", "manifest"])
def test_exchange_size_limits_reject_before_case_mutation(
    core: Any, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, limit_kind: str
) -> None:
    import case_exchange_safety as safety

    target, members = exchange(core, tmp_path)
    archive_path = tmp_path / "bounded.zip"
    with ZipFile(archive_path, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
        archive.writestr("large-synthetic.txt", b"A" * 65536)
    with ZipFile(archive_path) as archive:
        expanded_bytes = sum(info.file_size for info in archive.infolist())
        member_count = len(archive.infolist())
        manifest_bytes = archive.getinfo("case_update.json").file_size
    limits = {
        "archive": ("MAX_ARCHIVE_BYTES", archive_path.stat().st_size - 1),
        "expanded": ("MAX_ARCHIVE_BYTES", expanded_bytes - 1),
        "members": ("MAX_MEMBERS", member_count - 1),
        "manifest": ("MAX_MANIFEST_BYTES", manifest_bytes - 1),
    }
    name, value = limits[limit_kind]
    monkeypatch.setattr(safety, name, value)
    before = {
        p.relative_to(target): p.read_bytes() for p in target.rglob("*") if p.is_file()
    }

    with pytest.raises(core.CaseWorkspaceError, match="exceeds.*limit"):
        core.import_case_update(target, archive_path)

    assert {
        p.relative_to(target): p.read_bytes() for p in target.rglob("*") if p.is_file()
    } == before
    assert not (target / "exchange_imports").exists()
