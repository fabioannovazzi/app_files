"""Keep native answer exports readable without an installed Pandoc binary."""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE

ROOT = Path(__file__).resolve().parents[2]


def _package_module():
    name = "answer_docx_fallback_under_test"
    spec = importlib.util.spec_from_file_location(
        name,
        ROOT / "plugins/deep-research-validator/scripts/package_validation.py",
    )
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def test_native_fallback_exports_clickable_sources_and_readable_hierarchy(
    tmp_path: Path, monkeypatch
) -> None:
    module = _package_module()
    monkeypatch.setattr(module, "pypandoc", None)
    output = tmp_path / "note.docx"
    markdown = (
        "# Regole sui pagamenti\n\n## Ambito\n\n"
        "Testo **da verificare** con `riferimento` e "
        "[fonte UE](https://example.eu/procedure?id=(2023)).\n\n"
        "[altro](http://example.eu/guide) e "
        "[collegamento non attivo](javascript:alert(1))\n"
    )

    assert module.try_write_docx(markdown, output)

    document = Document(output)
    assert document.paragraphs[0].style.name == "Title"
    assert document.paragraphs[1].style.name == "Heading 1"
    assert document.paragraphs[0].text == "Regole sui pagamenti"
    assert document.element.xpath(".//w:hyperlink/w:r/w:t/text()") == [
        "fonte UE",
        "altro",
    ]
    targets = {
        relation.target_ref
        for relation in document.part.rels.values()
        if relation.reltype == RELATIONSHIP_TYPE.HYPERLINK
    }
    assert targets == {
        "https://example.eu/procedure?id=(2023)",
        "http://example.eu/guide",
    }
    assert any(
        run.bold and run.text == "da verificare" for run in document.paragraphs[2].runs
    )
    assert any(
        run.font.name == "Consolas" and run.text == "riferimento"
        for run in document.paragraphs[2].runs
    )
    assert "collegamento non attivo" in document.paragraphs[3].text
    assert "**" not in document.paragraphs[2].text
    assert not document.styles["Title"].element.xpath(".//w:pBdr")


def test_native_fallback_preserves_comparison_as_word_table(tmp_path: Path) -> None:
    module = _package_module()
    output = tmp_path / "comparison.docx"
    markdown = (
        "# Confronto\n\n"
        "| Tema | Regola |\n"
        "| :--- | ---: |\n"
        "| Termine | **60 giorni**, salvo deroga |\n"
        "| Fonte | [Direttiva](https://example.eu/directive) |\n\n"
        "Una precisazione | che rimane testo.\n"
    )

    assert module._write_docx_fallback(markdown, output)

    document = Document(output)
    assert len(document.tables) == 1
    table = document.tables[0]
    assert len(table.rows) == 3
    assert [cell.text for cell in table.rows[0].cells] == ["Tema", "Regola"]
    assert table.cell(1, 1).text == "60 giorni, salvo deroga"
    assert table.cell(1, 1).paragraphs[0].runs[1].bold
    assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    assert document.element.xpath(".//w:hyperlink/w:r/w:t/text()") == ["Direttiva"]
    assert document.paragraphs[-1].text == "Una precisazione | che rimane testo."
    assert document.paragraphs[-1].paragraph_format.space_before.pt >= 6
