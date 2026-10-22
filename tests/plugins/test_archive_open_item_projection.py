"""Exercise a real three-invoice Archive run through the public review transport."""

from __future__ import annotations

import copy
import json
import os
import shutil
import subprocess
import sys
from contextlib import contextmanager
from pathlib import Path

import pytest

from tests.open_item_reconciliation.test_reviewed_tabular_adapter import spec
from tests.plugins.test_archive_automatic_reporting import (
    archive_core,
    archive_script_imports,
    client_case,
    reporting_runtime,
)

__all__: list[str] = []
PLUGIN = Path(__file__).resolve().parents[2] / "plugins/open-item-reconciliation"
AUTHORED_TITLE = (
    "Verifica partite aperte al 30 settembre 2026 — incassi fino al 20 ottobre 2026"
)
AUTHORED_NARRATIVE = (
    "Bozza da dati fittizi. Popolazione al 30 settembre 2026; verifica degli incassi successivi fino al 20 ottobre 2026. "
    "Revisione dell'agente agent.codex; nessuna conferma umana o approvazione professionale."
)


@pytest.fixture
def reconciled_case(archive_core, client_case, tmp_path, request, monkeypatch):
    sources = {
        "population.csv": (
            "InvoiceX;DateX;ValueX;DueX;PartyX;PrivateX\n"
            "F001;01/09/2026;1200.00;30/09/2026;Alfa;LOCAL_ONLY_1\n"
            "F002;05/09/2026;800.00;30/09/2026;Beta;LOCAL_ONLY_2\n"
            "F003;15/09/2026;500.00;15/10/2026;Gamma;LOCAL_ONLY_3\n"
        ),
        "bank.csv": (
            "InvoiceX;PostedX;ValueX;PartyX;PrivateX\n"
            "F001;05/10/2026;1200.00;Alfa;LOCAL_ONLY_4\n"
            "F002;10/10/2026;300.00;Beta;LOCAL_ONLY_5\n"
        ),
    }
    mapped_kind = getattr(request, "param", None) == "mapped_evidence_type"
    if mapped_kind:
        header, *rows = sources["bank.csv"].splitlines()
        sources["bank.csv"] = (
            header + ";KindX\n" + "".join(row + ";transfer\n" for row in rows)
        )
    imported = []
    for name, text in sources.items():
        path = tmp_path / name
        path.write_text(text)
        imported.append(
            archive_core.import_studio_client_document(
                client_case.client_id,
                path,
                "source",
                engagement_id=client_case.engagement_id,
                state_dir=client_case.state_dir,
            )
        )
    prepared = archive_core.prepare_studio_client_workflow(
        client_case.engagement_id,
        "open-item-reconciliation",
        input_ids=[entry["input_id"] for entry in imported],
        purpose="PRIVATE_ENGAGEMENT_PURPOSE",
        state_dir=client_case.state_dir,
    )
    run_id = prepared["run"]["run_id"]
    archive_core.start_studio_client_workflow(
        client_case.client_id,
        client_case.engagement_id,
        run_id,
        state_dir=client_case.state_dir,
    )
    loaded = archive_core.ledger.load_run(
        client_case.client_root, client_case.engagement_id, run_id
    )
    decisions = {}
    for binding in loaded["context"]["input_bindings"]:
        path = Path(binding["path"])
        population = path.name == "population.csv"
        decision = spec(path, "open_items" if population else "bank_statement")
        decision["perimeter"]["party_ref"] = None
        decision["tabular"]["columns"]["party_ref"] = "PartyX"
        if population:
            decision["tabular"]["last_data_row"] = 4
            decision["tabular"]["columns"]["due_date"] = "DueX"
        else:
            del decision["tabular"]["columns"]["document_date"]
            if mapped_kind:
                decision["tabular"]["columns"]["evidence_type"] = "KindX"
                decision["tabular"]["default_evidence_type"] = None
                decision["tabular"]["evidence_type_map"] = {"transfer": "external_bank"}
        decisions[str(path)] = decision
    assumptions = Path(loaded["output_dir"]) / "assumptions.json"
    assumptions.write_text(
        json.dumps(
            {
                "assurance_run_date": "2026-10-20",
                "scope_year": 2026,
                "cutoff_date": "2026-09-30",
                "amount_tolerance": "0.00",
                "post_cutoff_events_excluded": False,
                "reviewed_source_decisions": decisions,
            }
        )
    )
    helper_arguments = [
        "--client-engagement",
        str(Path(loaded["run_root"]) / "context.json"),
        "--assumptions-json",
        str(assumptions),
        "--output-subdirectory",
        "reconciliation",
    ]
    variant = getattr(request, "param", "")
    if variant in {"model_authored", "model_authored_changed", "authored_cli"}:
        helper_arguments.extend(
            ["--title", AUTHORED_TITLE, "--narrative", AUTHORED_NARRATIVE]
        )
        if variant != "authored_cli":
            from tests.plugins.test_native_execution_reporting import (
                submit_native_arguments,
            )

            submit_native_arguments(
                tmp_path,
                monkeypatch,
                "open-item-reconciliation",
                helper_arguments,
                script="scripts/raw_input_runner.py",
            )
        if variant == "model_authored_changed":
            helper_arguments[
                helper_arguments.index("--title") + 1
            ] += " changed after the request"
    completed = subprocess.run(
        [
            sys.executable,
            "-I",
            "-B",
            str(PLUGIN / "scripts/raw_input_runner.py"),
            *helper_arguments,
        ],
        capture_output=True,
        text=True,
        timeout=90,
    )
    assert completed.returncode == 0, completed.stderr
    assert json.loads(completed.stdout)["status"] == "ready_for_review"
    assert "PRIVATE_ENGAGEMENT_PURPOSE" not in completed.stdout
    loaded["calculation_dir"] = str(Path(loaded["output_dir"]) / "reconciliation")
    return loaded


@contextmanager
def review_transport(tmp_path):
    """Use the same owned emission bridge as installed Vera."""
    node = shutil.which("node")
    if node is None:
        pytest.skip("Node.js is required for the real MCP transport.")
    vera = PLUGIN.parent / "vera"
    launcher = tmp_path / "bridge.cjs"
    launcher.write_text(
        "const {spawn}=require('node:child_process');\n"
        f"const {{bridge}}=require({json.dumps(str(vera / 'scripts/archive_emission_mcp.cjs'))});\n"
        f"const child=spawn(process.execPath,[{json.dumps(str(PLUGIN / 'mcp/server.cjs'))}],{{stdio:['pipe','pipe','inherit']}});\n"
        f"bridge(child,'open-item-reconciliation',{json.dumps(str(vera))});\n"
    )
    with subprocess.Popen(
        [node, str(launcher)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        env={
            **os.environ,
            "VERA_REPORTING_PYTHON": sys.executable,
            "OPEN_ITEM_RECONCILIATION_PYTHON": sys.executable,
        },
    ) as process:

        def call(name, arguments):
            process.stdin.write(
                json.dumps(
                    {
                        "jsonrpc": "2.0",
                        "id": 1,
                        "method": "tools/call",
                        "params": {"name": name, "arguments": arguments},
                    }
                )
                + "\n"
            )
            process.stdin.flush()
            response = json.loads(process.stdout.readline())
            assert not response["result"].get("isError"), response
            return response["result"]

        try:
            yield call
        finally:
            process.stdin.close()
            process.wait(timeout=15)


def report(loaded):
    from archive_reporting import snapshot_run

    return json.loads(Path(snapshot_run(loaded)["report_path"]).read_text())


def worksheet_selection(path, name, column=None):
    """Select an actual worksheet in Archive's flattened physical-row order."""
    from openpyxl import load_workbook

    if name is None:
        return {}
    workbook = load_workbook(path, read_only=True)
    offset = 0
    try:
        for sheet in workbook:
            rows = [
                row
                for row in sheet.iter_rows(values_only=True)
                if any(value is not None for value in row)
            ]
            if sheet.title == name:
                options = {"offset": offset, "limit": len(rows)}
                if column is not None:
                    options["columns"] = [rows[0].index(column)]
                return options
            offset += len(rows)
    finally:
        workbook.close()
    raise AssertionError(f"Expected generated worksheet {name}")


def word_paragraphs(path):
    """Locate assertions in the actual document body, including table cells."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    document = Document(path)
    paragraphs = [
        Paragraph(element, document)
        for element in document.element.body.iter(qn("w:p"))
        if Paragraph(element, document).text
    ]
    return document, paragraphs


@pytest.mark.parametrize(
    "reconciled_case,bank_cells",
    [(None, 8), ("mapped_evidence_type", 10)],
    indirect=["reconciled_case"],
)
@pytest.mark.parametrize("artifact_folder", ["", "assurance_final_outputs"])
def test_canonical_json_retains_exact_row_and_allocation_source_cells(
    reconciled_case, artifact_folder, bank_cells
):
    from archive_reporting import read_artifact

    path = (
        Path(reconciled_case["calculation_dir"])
        / artifact_folder
        / "reconciliation_results.json"
    )

    packet = read_artifact(reconciled_case, path)

    assert "LOCAL_ONLY" not in json.dumps(packet)
    measured = report(reconciled_case)
    assert measured["collection_status"] == "complete", measured["errors"]
    inputs = {
        row["label"]: row for row in measured["sources"] if row["role"] == "input"
    }
    assert inputs["population.csv"]["sent_cells"] == 15
    assert inputs["bank.csv"]["sent_cells"] == bank_cells


@pytest.mark.parametrize("artifact_folder", ["", "assurance_final_outputs"])
def test_changed_canonical_json_cannot_reuse_producer_provenance(
    reconciled_case, artifact_folder
):
    from archive_reporting import read_artifact

    path = (
        Path(reconciled_case["calculation_dir"])
        / artifact_folder
        / "reconciliation_results.json"
    )
    value = json.loads(path.read_bytes())
    value["review_rows"][0]["review_notes"] = "UNTRACED_REPLACEMENT"
    path.write_text(json.dumps(value))

    read_artifact(reconciled_case, path)

    measured = report(reconciled_case)
    assert measured["collection_status"] == "incomplete"
    assert any(
        error.startswith("derived_output_lineage_unavailable:")
        for error in measured["errors"]
    )
    assert (
        sum(row["sent_cells"] for row in measured["sources"] if row["role"] == "input")
        == 0
    )


@pytest.mark.parametrize(
    "reconciled_case", [None, "mapped_evidence_type"], indirect=True
)
def test_assurance_receipt_retains_only_echoed_amount_and_allocation_cells(
    reconciled_case,
):
    from archive_reporting import read_artifact

    path = Path(reconciled_case["calculation_dir"]) / "assurance_receipts.json"

    packet = read_artifact(reconciled_case, path)

    assert "LOCAL_ONLY" not in json.dumps(packet)
    measured = report(reconciled_case)
    assert measured["collection_status"] == "complete", measured["errors"]
    inputs = {
        row["label"]: row for row in measured["sources"] if row["role"] == "input"
    }
    # Three invoice amounts and the two allocated counterparties; two bank
    # amounts and their counterparties. Invoice dates and unused text stay local.
    assert inputs["population.csv"]["sent_cells"] == 5
    assert inputs["bank.csv"]["sent_cells"] == 4


def test_changed_assurance_receipt_cannot_reuse_producer_provenance(reconciled_case):
    from archive_reporting import read_artifact

    path = Path(reconciled_case["calculation_dir"]) / "assurance_receipts.json"
    value = json.loads(path.read_bytes())
    value["rendered_value_addresses"][0]["value"] = "987654.00"
    path.write_text(json.dumps(value))

    read_artifact(reconciled_case, path)

    measured = report(reconciled_case)
    assert measured["collection_status"] == "incomplete"
    assert any(
        error.startswith("derived_output_lineage_unavailable:")
        for error in measured["errors"]
    )
    assert (
        sum(row["sent_cells"] for row in measured["sources"] if row["role"] == "input")
        == 0
    )


@pytest.mark.parametrize(
    "reconciled_case,expected_status",
    [
        ("model_authored", "complete"),
        ("authored_cli", "incomplete"),
        ("model_authored_changed", "incomplete"),
    ],
    indirect=["reconciled_case"],
)
@pytest.mark.parametrize("artifact_folder", ["", "assurance_final_outputs"])
def test_word_authored_text_uses_exact_native_request_provenance(
    reconciled_case, expected_status, artifact_folder
):
    from archive_reporting import read_artifact

    path = (
        Path(reconciled_case["calculation_dir"])
        / artifact_folder
        / "relazione_riconciliazione_audit.docx"
    )

    packet = read_artifact(reconciled_case, path)

    assert AUTHORED_TITLE in json.dumps(packet, ensure_ascii=False)
    assert AUTHORED_NARRATIVE in json.dumps(packet, ensure_ascii=False)
    measured = report(reconciled_case)
    assert measured["collection_status"] == expected_status, measured["errors"]
    inputs = {
        row["label"]: row for row in measured["sources"] if row["role"] == "input"
    }
    assert inputs["population.csv"]["sent_cells"] == 9
    assert inputs["bank.csv"]["sent_cells"] == 6


@pytest.mark.parametrize("artifact_folder", ["", "assurance_final_outputs"])
@pytest.mark.parametrize(
    "selection,population_cells,bank_cells",
    [("title", 0, 0), ("partial_amounts", 1, 1), ("whole_document", 9, 6)],
)
def test_word_reads_retain_exact_source_cells(
    reconciled_case, artifact_folder, selection, population_cells, bank_cells
):
    from archive_reporting import read_artifact

    path = (
        Path(reconciled_case["calculation_dir"])
        / artifact_folder
        / "relazione_riconciliazione_audit.docx"
    )
    options = {}
    if selection == "title":
        options = {"offset": 0, "limit": 1}
    elif selection == "partial_amounts":
        document, paragraphs = word_paragraphs(path)
        table = next(
            table
            for table in document.tables
            if "Residuo aperto" in [cell.text for cell in table.rows[0].cells]
        )
        amount_paragraph = table.rows[1].cells[1].paragraphs[0]._p
        offset = [paragraph._p for paragraph in paragraphs].index(amount_paragraph)
        options = {"offset": offset, "limit": 3}

    packet = read_artifact(reconciled_case, path, **options)

    assert packet["data"]["rows"]
    measured = report(reconciled_case)
    assert measured["collection_status"] == "complete", measured["errors"]
    inputs = {
        row["label"]: row for row in measured["sources"] if row["role"] == "input"
    }
    assert inputs["population.csv"]["sent_cells"] == population_cells
    assert inputs["bank.csv"]["sent_cells"] == bank_cells


@pytest.mark.parametrize("reconciled_case", ["mapped_evidence_type"], indirect=True)
@pytest.mark.parametrize("artifact_folder", ["", "assurance_final_outputs"])
def test_word_evidence_labels_include_mapped_source_type_cells(
    reconciled_case, artifact_folder
):
    from archive_reporting import read_artifact

    path = (
        Path(reconciled_case["calculation_dir"])
        / artifact_folder
        / "relazione_riconciliazione_audit.docx"
    )

    read_artifact(reconciled_case, path)

    measured = report(reconciled_case)
    assert measured["collection_status"] == "complete", measured["errors"]
    inputs = {
        row["label"]: row for row in measured["sources"] if row["role"] == "input"
    }
    assert inputs["population.csv"]["sent_cells"] == 9
    assert inputs["bank.csv"]["sent_cells"] == 8


@pytest.mark.parametrize(
    "custom_field", ["title", "narrative", "next_steps", "metadata"]
)
@pytest.mark.parametrize("read_custom", [False, True])
def test_custom_word_text_stays_unknown_without_tainting_unrelated_amounts(
    reconciled_case, custom_field, read_custom
):
    from archive_projection import register_table_projection
    from archive_reporting import read_artifact

    from tests.open_item_reconciliation.test_workpaper_outputs import load_outputs

    output = Path(reconciled_case["calculation_dir"])
    source = json.loads((output / "normalized_records.json").read_bytes())[0]
    path = output / "custom-report.docx"
    options = {
        "title": "Test report",
        "metadata": {},
        "next_steps": [],
        "narrative": "",
    }
    custom_value = "UNTRACED_CUSTOM_TEXT"
    options[custom_field] = (
        [custom_value]
        if custom_field == "next_steps"
        else (
            {"Custom context": custom_value}
            if custom_field == "metadata"
            else custom_value
        )
    )
    cells = []
    load_outputs().write_word_report(
        path,
        **options,
        summary_rows=[
            {
                "Record refs": source["record_id"],
                "reconciliation_status": "open_supported",
                "rows": 1,
                "amount": "1200.00",
            }
        ],
        assumptions={},
        language="en",
        projection_cells=cells,
    )
    register_table_projection("open-item-reconciliation", path, cells)
    _, paragraphs = word_paragraphs(path)
    target = custom_value if read_custom else "1,200.00"
    offset = next(
        index for index, paragraph in enumerate(paragraphs) if target in paragraph.text
    )

    packet = read_artifact(reconciled_case, path, offset=offset, limit=1)

    assert target in packet["data"]["rows"][0][0]
    measured = report(reconciled_case)
    assert measured["collection_status"] == (
        "incomplete" if read_custom else "complete"
    )
    inputs = {
        row["label"]: row for row in measured["sources"] if row["role"] == "input"
    }
    assert inputs["population.csv"]["sent_cells"] == (0 if read_custom else 1)
    assert inputs["bank.csv"]["sent_cells"] == 0


def test_changed_word_cannot_reuse_original_source_provenance(reconciled_case):
    from archive_reporting import read_artifact
    from docx import Document

    path = (
        Path(reconciled_case["calculation_dir"])
        / "relazione_riconciliazione_audit.docx"
    )
    document = Document(path)
    document.paragraphs[0].text = "UNREGISTERED_REPLACEMENT"
    document.save(path)

    read_artifact(reconciled_case, path, offset=0, limit=1)

    measured = report(reconciled_case)
    assert measured["collection_status"] == "incomplete"
    assert any(
        error.startswith("derived_output_lineage_unavailable:")
        for error in measured["errors"]
    )
    assert (
        sum(row["sent_cells"] for row in measured["sources"] if row["role"] == "input")
        == 0
    )


@pytest.mark.parametrize(
    "relative",
    [
        "relazione_riconciliazione_audit.docx",
        "riconciliazione_audit.xlsx",
        "richieste_mirate_evidenze.xlsx",
        "reconciliation_results.json",
        "assurance_final_outputs/relazione_riconciliazione_audit.docx",
        "assurance_final_outputs/riconciliazione_audit.xlsx",
        "assurance_final_outputs/richieste_mirate_evidenze.xlsx",
        "assurance_final_outputs/reconciliation_results.json",
        "assurance_receipts.json",
    ],
)
def test_generated_deliverable_read_retains_its_source_provenance(
    archive_core, client_case, reconciled_case, relative
):
    loaded = reconciled_case
    path = Path(loaded["calculation_dir"]) / relative

    packet = archive_core.read_studio_run_artifact(
        client_case.client_id,
        client_case.engagement_id,
        loaded["run"]["run_id"],
        path,
        state_dir=client_case.state_dir,
    )

    assert packet["data"]["rows"]
    assert "LOCAL_ONLY_" not in json.dumps(packet)
    measured = report(loaded)
    assert measured["collection_status"] == "complete", measured["errors"]
    inputs = {
        row["label"]: row for row in measured["sources"] if row["role"] == "input"
    }
    assert inputs["population.csv"]["not_sent_cells"] >= 3
    assert inputs["bank.csv"]["not_sent_cells"] >= 2


@pytest.mark.parametrize("artifact_folder", ["", "assurance_final_outputs"])
@pytest.mark.parametrize(
    "sheet,column,population_cells,bank_cells",
    [
        (None, None, 9, 2),
        ("sintesi", "importo_totale", 3, 1),
        ("evidenze_da_integrare", "importo", 1, 1),
        ("evidenze_da_integrare", "evidenza_gia_disponibile", 0, 2),
        ("evidenze_da_integrare", "riferimento_esistente", 0, 1),
        ("evidenze_da_integrare", "dato_mancante_mirato", 0, 0),
    ],
)
def test_targeted_requests_retain_embedded_amount_and_date_sources(
    reconciled_case, artifact_folder, sheet, column, population_cells, bank_cells
):
    from archive_reporting import read_artifact

    path = (
        Path(reconciled_case["calculation_dir"])
        / artifact_folder
        / "richieste_mirate_evidenze.xlsx"
    )
    options = worksheet_selection(path, sheet, column)

    read_artifact(reconciled_case, path, **options)

    measured = report(reconciled_case)
    assert measured["collection_status"] == "complete", measured["errors"]
    inputs = {
        row["label"]: row for row in measured["sources"] if row["role"] == "input"
    }
    assert inputs["population.csv"]["sent_cells"] == population_cells
    assert inputs["bank.csv"]["sent_cells"] == bank_cells


@pytest.mark.parametrize(
    "column,population_cells,bank_cells,status",
    [
        ("dato_mancante_mirato", 1, 1, "complete"),
        ("evidenza_gia_disponibile", 0, 1, "complete"),
        ("destinatario_operativo", 0, 0, "incomplete"),
    ],
)
def test_request_without_later_candidate_traces_residual_prose_and_unknown_owner(
    reconciled_case, column, population_cells, bank_cells, status
):
    from archive_projection import register_table_projection
    from archive_reporting import read_artifact

    from tests.open_item_reconciliation.test_build_missing_evidence_requests import (
        load_missing_requests,
    )
    from tests.open_item_reconciliation.test_workpaper_outputs import load_outputs

    output = Path(reconciled_case["calculation_dir"])
    rows = json.loads((output / "reconciliation_results.json").read_bytes())[
        "reconciliation_rows"
    ]
    normalized = json.loads((output / "normalized_records.json").read_bytes())
    field_maps = load_outputs().workpaper_row_origins(
        "Reconciliation detail", rows, {"Normalized records": normalized}
    )
    origins = {
        "Reconciliation detail": {
            row["record_id"]: fields
            for row, fields in zip(rows, field_maps, strict=True)
        }
    }
    requests = load_missing_requests()
    projection_rows = {}
    pack = requests.build_missing_evidence_request_pack(
        rows,
        normalized_records=normalized,
        source_origins=origins,
        projection_rows=projection_rows,
        entity_name="UNTRACED_OWNER",
    )
    cells = []
    path = output / "requests-without-later-candidates.xlsx"
    requests.write_missing_evidence_workbook(
        path, pack, projection_rows=projection_rows, projection_cells=cells
    )
    register_table_projection("open-item-reconciliation", path, cells)
    selection = worksheet_selection(path, "evidenze_da_integrare", column)

    packet = read_artifact(reconciled_case, path, **selection)

    assert packet["data"]["rows"]
    measured = report(reconciled_case)
    assert measured["collection_status"] == status, measured["errors"]
    inputs = {
        row["label"]: row for row in measured["sources"] if row["role"] == "input"
    }
    assert inputs["population.csv"]["sent_cells"] == population_cells
    assert inputs["bank.csv"]["sent_cells"] == bank_cells


@pytest.mark.parametrize("artifact_folder", ["", "assurance_final_outputs"])
@pytest.mark.parametrize(
    "sheet_name,selected_column,expected_population,expected_bank",
    [
        ("Normalized records", None, 15, 8),
        ("Righe normalizzate", None, 15, 8),
        ("Summary", "amount", 3, 0),
        ("Reconciliation detail", "allocated_amount", 0, 1),
        ("Reconciliation detail", "residual_amount", 1, 1),
        ("Review", "residual_amount", 1, 1),
        (None, None, 15, 8),
        ("External evidence aggregate", "cash_flow_signed_total", 0, 0),
        ("External evidence aggregate", "amount_total", 0, 2),
        ("Open item aging", "partially_paid_amount", 1, 0),
        ("Open item aging", "min_age_days", 3, 0),
        ("Evidence concentration", "share_of_abs_amount_percent", 3, 0),
        ("Document source map", "document_dates", 3, 2),
        ("Post-cutoff candidates", "evidence_date", 0, 2),
        ("Cutoff window movements", "days_from_cutoff", 3, 2),
    ],
)
def test_audit_workbook_selected_rows_retain_exact_source_cells(
    reconciled_case,
    artifact_folder,
    sheet_name,
    selected_column,
    expected_population,
    expected_bank,
):
    from archive_reporting import read_artifact

    loaded = reconciled_case
    path = (
        Path(loaded["calculation_dir"]) / artifact_folder / "riconciliazione_audit.xlsx"
    )
    options = worksheet_selection(path, sheet_name, selected_column)

    packet = read_artifact(loaded, path, **options)

    assert "LOCAL_ONLY_" not in json.dumps(packet)
    measured = report(loaded)
    assert measured["collection_status"] == "complete", measured["errors"]
    inputs = {
        row["label"]: row for row in measured["sources"] if row["role"] == "input"
    }
    assert inputs["population.csv"]["sent_cells"] == expected_population
    assert inputs["bank.csv"]["sent_cells"] == expected_bank


@pytest.mark.parametrize(
    "column,expected_status,expected_value",
    [(0, "complete", "1200"), (1, "incomplete", "UNTRACED_NOTE")],
)
def test_audit_writer_keeps_supplied_notes_unknown_without_tainting_amount_reads(
    reconciled_case,
    column,
    expected_status,
    expected_value,
):
    from archive_projection import register_table_projection
    from archive_reporting import read_artifact

    from tests.open_item_reconciliation.test_workpaper_outputs import load_outputs

    loaded = reconciled_case
    output = Path(loaded["calculation_dir"])
    source_row = json.loads((output / "normalized_records.json").read_bytes())[0]
    path = output / "selected-review.xlsx"
    cells = []
    load_outputs().write_excel_workpaper(
        path,
        {
            "Normalized records": [source_row],
            "Reconciliation detail": [{**source_row, "review_notes": "UNTRACED_NOTE"}],
        },
        preferred_headers={"Reconciliation detail": ["amount", "review_notes"]},
        projection_cells=cells,
    )
    register_table_projection("open-item-reconciliation", path, cells)

    packet = read_artifact(loaded, path, offset=3, limit=1, columns=[column])

    assert packet["data"]["rows"] == [[expected_value]]
    measured = report(loaded)
    assert measured["collection_status"] == expected_status


def test_supplied_check_note_is_not_registered_as_generated_metadata(reconciled_case):
    from archive_projection import register_table_projection
    from archive_reporting import read_artifact

    from tests.open_item_reconciliation.test_workpaper_outputs import load_outputs

    path = Path(reconciled_case["calculation_dir"]) / "supplied-checks.xlsx"
    cells = []
    load_outputs().write_excel_workpaper(
        path,
        {"Checks": [{"note": "UNTRACED_CHECK_NOTE"}]},
        projection_cells=cells,
    )
    register_table_projection("open-item-reconciliation", path, cells)

    packet = read_artifact(reconciled_case, path, offset=1, limit=1)

    assert packet["data"]["rows"] == [["UNTRACED_CHECK_NOTE"]]
    assert report(reconciled_case)["collection_status"] == "incomplete"


@pytest.mark.parametrize("reconciled_case", ["mapped_evidence_type"], indirect=True)
@pytest.mark.parametrize(
    "sheet,column",
    [
        ("External evidence detail", "evidence_type"),
        ("Post-cutoff candidates", "evidence_type"),
        ("Document source map", "evidence_types_present"),
        ("Reconciliation detail", "matched_evidence_type"),
    ],
)
def test_audit_evidence_labels_retain_mapped_source_type_cells(
    reconciled_case, sheet, column
):
    from archive_reporting import read_artifact

    path = Path(reconciled_case["calculation_dir"]) / "riconciliazione_audit.xlsx"
    selection = worksheet_selection(path, sheet, column)

    read_artifact(reconciled_case, path, **selection)

    measured = report(reconciled_case)
    assert measured["collection_status"] == "complete", measured["errors"]
    inputs = {
        row["label"]: row for row in measured["sources"] if row["role"] == "input"
    }
    assert inputs["population.csv"]["sent_cells"] == 0
    assert inputs["bank.csv"]["sent_cells"] == 2


@pytest.mark.parametrize("artifact_folder", ["", "assurance_final_outputs"])
@pytest.mark.parametrize(
    "selection,population_cells,bank_cells",
    [("legend", 0, 0), ("partial_amounts", 1, 1), ("whole_workbook", 9, 6)],
)
def test_accountant_workbook_reads_preserve_exact_original_cell_provenance(
    reconciled_case, artifact_folder, selection, population_cells, bank_cells
):
    from archive_reporting import read_artifact
    from openpyxl import load_workbook

    loaded = reconciled_case
    path = (
        Path(loaded["calculation_dir"])
        / artifact_folder
        / "scheda_operativa_commercialista.xlsx"
    )
    read_options = {}
    if selection == "legend":
        read_options = {"offset": 0, "limit": 4}
    elif selection == "partial_amounts":
        with path.open("rb") as handle:
            book = load_workbook(handle, read_only=True)
            legend_rows = sum(
                any(value is not None for value in row)
                for row in book["Legenda"].iter_rows(values_only=True)
            )
            book.close()
        read_options = {"offset": legend_rows + 2, "limit": 1, "columns": [4, 6, 7]}

    read_artifact(loaded, path, **read_options)

    measured = report(loaded)
    assert measured["collection_status"] == "complete", measured["errors"]
    inputs = {
        row["label"]: row for row in measured["sources"] if row["role"] == "input"
    }
    assert inputs["population.csv"]["sent_cells"] == population_cells
    assert inputs["bank.csv"]["sent_cells"] == bank_cells


@pytest.mark.parametrize(
    "name,sheet",
    [
        ("scheda_operativa_commercialista.xlsx", "Legenda"),
        ("riconciliazione_audit.xlsx", "Indice"),
        ("richieste_mirate_evidenze.xlsx", "istruzioni"),
    ],
)
def test_changed_workbook_cannot_reuse_original_source_provenance(
    reconciled_case,
    name,
    sheet,
):
    from archive_reporting import read_artifact
    from openpyxl import load_workbook

    loaded = reconciled_case
    path = Path(loaded["calculation_dir"]) / name
    workbook = load_workbook(path)
    workbook[sheet]["B2"] = "UNREGISTERED_REPLACEMENT"
    workbook.save(path)
    workbook.close()

    read_artifact(loaded, path, offset=0, limit=4)

    measured = report(loaded)
    assert measured["collection_status"] == "incomplete"
    assert any(
        error.startswith("derived_output_lineage_unavailable:")
        for error in measured["errors"]
    )
    assert (
        sum(row["sent_cells"] for row in measured["sources"] if row["role"] == "input")
        == 0
    )


@pytest.mark.parametrize(
    "name,read_options,source_cells",
    [
        ("artifact_card.md", {}, 0),
        ("artifact_card.md", {"offset": 0, "limit": 4}, 0),
        ("final_artifacts.json", {}, 1),
        ("ui_decisions.json", {}, 0),
    ],
)
def test_initial_workflow_status_files_have_producer_provenance(
    reconciled_case, name, read_options, source_cells
):
    from archive_reporting import read_artifact

    loaded = reconciled_case
    path = Path(loaded["calculation_dir"]) / name

    read_artifact(loaded, path, **read_options)

    measured = report(loaded)
    assert measured["collection_status"] == "complete", measured["errors"]
    assert (
        sum(
            source["sent_cells"]
            for source in measured["sources"]
            if source["role"] == "input"
        )
        == source_cells
    )


@pytest.mark.parametrize("route", ["browser", "mcp"])
@pytest.mark.parametrize("operation", ["save", "apply"])
@pytest.mark.parametrize("with_note", [False, True])
def test_saved_decision_titles_keep_sources_and_arbitrary_notes_remain_unmeasured(
    reconciled_case, tmp_path, route, operation, with_note
):
    from archive_reporting import read_artifact

    loaded = reconciled_case
    output = Path(loaded["calculation_dir"])
    review = json.loads((output / "review_payload.json").read_text())
    args = {
        "client_engagement": str(Path(loaded["run_root"]) / "context.json"),
        **{
            key + "_path": str(output / (key + ".json"))
            for key in ("run_intake", "review_payload", "final_artifacts")
        },
        "decision_source": "codex_agent",
        "reviewer": "agent.test",
        "decisions": [
            {
                "item_id": item["id"],
                "action": "accept",
                **({"note": "UNREGISTERED_REVIEW_TEXT"} if with_note else {}),
            }
            for item in review["items"]
        ],
        "expected_predecessor_checkpoint": json.loads(
            (output / "assurance_receipts.json").read_text()
        )["content_sha256"],
    }

    if route == "mcp":
        with review_transport(tmp_path) as call:
            call(f"{operation}_open_item_reconciliation_decisions", args)
    else:
        completed = subprocess.run(
            [
                sys.executable,
                "-I",
                "-B",
                "-c",
                "import json, runpy, sys; "
                "api = runpy.run_path(sys.argv[1]); "
                "api[sys.argv[2] + '_decisions'](sys.argv[3], json.loads(sys.stdin.read()))",
                str(PLUGIN / "scripts/review_server.py"),
                operation,
                str(output),
            ],
            input=json.dumps(args),
            capture_output=True,
            text=True,
            timeout=90,
        )
        assert completed.returncode == 0, completed.stderr
    read_artifact(loaded, output / "ui_decisions.json")

    measured = report(loaded)
    inputs = {
        row["label"]: row for row in measured["sources"] if row["role"] == "input"
    }
    if with_note:
        assert measured["collection_status"] == "incomplete"
        assert any(
            error.startswith("derived_output_lineage_unavailable:")
            for error in measured["errors"]
        )
        assert inputs["population.csv"]["sent_cells"] == 0
    else:
        assert measured["collection_status"] == "complete", measured["errors"]
        assert inputs["population.csv"]["sent_cells"] == 6
    assert inputs["bank.csv"]["sent_cells"] == 0


def test_edited_status_file_cannot_inherit_producer_metadata(reconciled_case):
    from archive_reporting import read_artifact

    loaded = reconciled_case
    path = Path(loaded["calculation_dir"]) / "artifact_card.md"
    path.write_text(path.read_text() + "UNREGISTERED_REVIEW_TEXT\n")

    read_artifact(loaded, path)

    measured = report(loaded)
    assert measured["collection_status"] == "incomplete"
    assert any(
        error.startswith("derived_output_lineage_unavailable:")
        for error in measured["errors"]
    )


@pytest.mark.parametrize("route", ["browser", "mcp"])
def test_applied_status_keeps_qa_source_provenance_without_private_review_notes(
    reconciled_case, tmp_path, route
):
    from archive_reporting import read_artifact

    loaded = reconciled_case
    output = Path(loaded["calculation_dir"])
    review = json.loads((output / "review_payload.json").read_text())
    args = {
        "client_engagement": str(Path(loaded["run_root"]) / "context.json"),
        **{
            key + "_path": str(output / (key + ".json"))
            for key in ("run_intake", "review_payload", "final_artifacts")
        },
        "decision_source": "codex_agent",
        "reviewer": "agent.test",
        "decisions": [
            {"item_id": item["id"], "action": "accept", "note": "PRIVATE_REVIEW_NOTE"}
            for item in review["items"]
        ],
        "expected_predecessor_checkpoint": json.loads(
            (output / "assurance_receipts.json").read_text()
        )["content_sha256"],
    }

    if route == "mcp":
        with review_transport(tmp_path) as call:
            call("apply_open_item_reconciliation_decisions", args)
    else:
        completed = subprocess.run(
            [
                sys.executable,
                "-I",
                "-B",
                "-c",
                "import json, runpy, sys; "
                "api = runpy.run_path(sys.argv[1]); "
                "api['apply_decisions'](sys.argv[2], json.loads(sys.stdin.read()))",
                str(PLUGIN / "scripts/review_server.py"),
                str(output),
            ],
            input=json.dumps(args),
            capture_output=True,
            text=True,
            timeout=90,
        )
        assert completed.returncode == 0, completed.stderr
    packet = read_artifact(loaded, output / "final_artifacts.json")

    assert "PRIVATE_REVIEW_NOTE" not in json.dumps(packet)
    measured = report(loaded)
    assert measured["collection_status"] == "complete", measured["errors"]
    inputs = {
        row["label"]: row for row in measured["sources"] if row["role"] == "input"
    }
    assert inputs["population.csv"]["sent_cells"] == 1
    assert inputs["bank.csv"]["sent_cells"] == 0


def test_pending_review_counts_do_not_claim_source_cell_exposure(
    reconciled_case, tmp_path
):
    loaded = reconciled_case
    output = Path(loaded["calculation_dir"])
    with review_transport(tmp_path) as call:
        index = call(
            "validate_open_item_reconciliation_review",
            {
                "review_payload_path": str(output / "review_payload.json"),
                "client_engagement": str(Path(loaded["run_root"]) / "context.json"),
            },
        )["structuredContent"]
        rendered = call(
            "render_open_item_reconciliation_review",
            {
                "persistence_token": index["review_reference"]["persistence_token"],
            },
        )
        assert "review_payload" not in rendered["structuredContent"]
        assert (
            rendered["_meta"]["private_review_payload"]["review_payload"]["item_count"]
            == 8
        )
        handle = next(
            item["case_handle"]
            for item in index["model_context_index"]["cases"]
            if item["item_type"] == "check_exception"
        )
        selected = call(
            "get_open_item_reconciliation_case_context",
            {
                "persistence_token": index["review_reference"]["persistence_token"],
                "case_handles": [handle],
            },
        )["structuredContent"]

    # The unique exact payment is already verified; two items still need review.
    assert selected["cases"][0]["semantic_fields"]["actual"] == 2
    measured = report(loaded)
    assert measured["collection_status"] == "complete", measured["errors"]
    assert (
        sum(
            entry["sent_cells"]
            for entry in measured["sources"]
            if entry["role"] == "input"
        )
        == 0
    )


@pytest.mark.parametrize("exact_identifiers,expected_cells", [(False, 6), (True, 7)])
def test_selected_partial_payment_has_exact_cells_and_retains_private_columns(
    reconciled_case, tmp_path, exact_identifiers, expected_cells
):
    loaded = reconciled_case
    output = Path(loaded["calculation_dir"])
    with review_transport(tmp_path) as call:
        index = call(
            "validate_open_item_reconciliation_review",
            {
                "review_payload_path": str(output / "review_payload.json"),
                "client_engagement": str(Path(loaded["run_root"]) / "context.json"),
            },
        )["structuredContent"]
        initial = report(loaded)
        assert initial["collection_status"] == "complete", initial["errors"]
        assert (
            sum(s["sent_cells"] for s in initial["sources"] if s["role"] == "input")
            == 0
        )
        partial = next(
            item["case_handle"]
            for item in index["model_context_index"]["cases"]
            if item["item_type"] == "manual_review"
        )
        selected = call(
            "get_open_item_reconciliation_case_context",
            {
                "persistence_token": index["review_reference"]["persistence_token"],
                "case_handles": [partial],
                "include_exact_identifiers": exact_identifiers,
            },
        )["structuredContent"]

    fields = selected["cases"][0]["semantic_fields"]
    assert fields["party_ref"] == "Beta"
    assert fields["allocated_amount"] == "300.00"
    assert fields["residual_amount"] == "500.00"
    assert fields["supporting_bank_date"] == "2026-10-10"
    assert "LOCAL_ONLY" not in json.dumps(selected)
    measured = report(loaded)
    assert measured["collection_status"] == "complete", measured["errors"]
    inputs = [entry for entry in measured["sources"] if entry["role"] == "input"]
    assert sum(entry["sent_cells"] for entry in inputs) == expected_cells
    assert sum(entry["not_sent_cells"] for entry in inputs) == 28 - expected_cells


@pytest.mark.parametrize(
    "reconciled_case,regeneration_route,exact_identifiers",
    [
        (None, "native", False),
        (None, "native", True),
        (None, "cli", False),
        (None, "cli", True),
        (None, "cli-retained", False),
        (None, "cli-retained", True),
        ("model_authored", "native", False),
    ],
    indirect=["reconciled_case"],
)
def test_public_review_save_apply_keeps_notes_private_and_report_complete(
    reconciled_case, tmp_path, exact_identifiers, regeneration_route
):
    from archive_reporting import read_artifact

    loaded = reconciled_case
    output = Path(loaded["calculation_dir"])
    review = json.loads((output / "review_payload.json").read_text())
    initial = json.loads((output / "final_artifacts.json").read_text())
    assert initial["status"] == "written_pending_review"
    assert initial["assurance"]["report_ready"] is False
    checkpoint = json.loads((output / "assurance_receipts.json").read_text())[
        "content_sha256"
    ]
    args = {
        "client_engagement": str(Path(loaded["run_root"]) / "context.json"),
        **{
            key + "_path": str(output / (key + ".json"))
            for key in ("run_intake", "review_payload", "final_artifacts")
        },
        "decision_source": "codex_agent",
        "reviewer": "agent.test",
        "decisions": [
            {"item_id": item["id"], "action": "accept", "note": "PRIVATE_REVIEW_NOTE"}
            for item in review["items"]
            if item["id"].startswith("review:")
        ],
    }
    with review_transport(tmp_path) as call:
        index = call(
            "validate_open_item_reconciliation_review",
            {
                "review_payload_path": str(output / "review_payload.json"),
                "client_engagement": str(Path(loaded["run_root"]) / "context.json"),
            },
        )["structuredContent"]
        selected = call(
            "get_open_item_reconciliation_case_context",
            {
                "persistence_token": index["review_reference"]["persistence_token"],
                "case_handles": [
                    item["case_handle"]
                    for item in index["model_context_index"]["cases"]
                ],
                "include_exact_identifiers": exact_identifiers,
            },
        )["structuredContent"]
        assert selected["case_count"] == 4
        saved = call("save_open_item_reconciliation_decisions", args)
        applied = call(
            "apply_open_item_reconciliation_decisions",
            {
                **args,
                "expected_predecessor_checkpoint": checkpoint,
            },
        )

    assert saved["structuredContent"]["decision_count"] == 3
    assert applied["structuredContent"]["decision_count"] == 3
    pending = json.loads((output / "final_artifacts.json").read_text())
    assert applied["structuredContent"]["application_status"] == "awaiting_regeneration"
    assert applied["structuredContent"]["blocker_count"] == 0
    assert pending["status"] == "awaiting_regeneration"
    assert pending["blockers"] == []
    assert (
        json.loads((output / "assurance_gates.json").read_text())["report_ready"]
        is False
    )
    assert pending["pending_steps"][0]["kind"] == "assurance_replay_required"
    assert pending["review_handoff"]["required_before_final_delivery"] is False
    assert "PRIVATE_REVIEW_NOTE" not in json.dumps(saved["structuredContent"])
    assert "PRIVATE_REVIEW_NOTE" not in json.dumps(applied["structuredContent"])
    assert "PRIVATE_REVIEW_NOTE" in json.dumps(saved["_meta"]["private_review_payload"])
    assert "PRIVATE_REVIEW_NOTE" in json.dumps(
        applied["_meta"]["private_review_payload"]
    )
    measured = report(loaded)
    assert measured["collection_status"] == "complete", measured["errors"]

    # Complete the real producer -> public save/apply -> native regeneration
    # integration, retaining the original sources and predecessor snapshot.
    sources_before = {
        entry["path"]: Path(entry["path"]).read_bytes()
        for entry in loaded["context"]["input_bindings"]
    }
    history = output / "assurance_transition_history" / checkpoint
    retained_before = {
        str(path.relative_to(history)): path.read_bytes()
        for path in history.rglob("*")
        if path.is_file()
    }
    command = [
        sys.executable,
        "-I",
        "-B",
        str(PLUGIN.parent / "vera/scripts/native_open_item_bridge.py"),
        str(PLUGIN),
    ]
    request = json.dumps(
        {
            "operation": "regenerate",
            "output_dir": str(output),
            "context": str(Path(loaded["run_root"]) / "context.json"),
            "arguments": {"expected_predecessor_checkpoint": checkpoint},
        }
    )
    if regeneration_route.startswith("cli"):
        command = [
            sys.executable,
            "-I",
            "-B",
            str(PLUGIN / "scripts/raw_input_runner.py"),
            "--client-engagement",
            str(Path(loaded["run_root"]) / "context.json"),
            "--output-subdirectory",
            "reconciliation",
            "--expected-predecessor-checkpoint",
            checkpoint,
        ]
        if regeneration_route == "cli":
            command.extend(
                [
                    "--assumptions-json",
                    str(Path(loaded["output_dir"]) / "assumptions.json"),
                ]
            )
        request = None
    regenerated = subprocess.run(
        command,
        input=request,
        capture_output=True,
        text=True,
        timeout=90,
    )
    assert regenerated.returncode == 0, regenerated.stderr
    summary = json.loads(regenerated.stdout)
    assert summary["regenerated"] is True
    assert summary["report_ready"] is True
    assert summary.get("status", "final_ready") == "final_ready"
    final = json.loads((output / "final_artifacts.json").read_text())
    assert final["status"] == "final_ready"
    assert final["review_handoff"]["status"] == "review_completed"
    assert final["review_handoff"]["required_before_final_delivery"] is False
    assert final["review_handoff"]["server"]["required"] is False
    assert "ui_decisions.json is pending" not in " ".join(final["caveats"])
    assert "Open the browser review server" not in " ".join(final["next_actions"])
    assert sources_before == {path: Path(path).read_bytes() for path in sources_before}
    assert retained_before == {
        str(path.relative_to(history)): path.read_bytes()
        for path in history.rglob("*")
        if path.is_file()
    }
    # Promotion changes paths and workbook bytes. Both current copies must carry
    # their own provenance, without borrowing the predecessor's declaration.
    read_artifact(loaded, output / "scheda_operativa_commercialista.xlsx")
    read_artifact(
        loaded,
        output / "assurance_final_outputs/scheda_operativa_commercialista.xlsx",
    )
    audit_path = output / "riconciliazione_audit.xlsx"
    selection = worksheet_selection(
        audit_path, "Reconciliation detail", "residual_amount"
    )
    read_artifact(loaded, audit_path, **selection)
    read_artifact(
        loaded, output / "assurance_final_outputs" / audit_path.name, **selection
    )
    read_artifact(loaded, output / "relazione_riconciliazione_audit.docx")
    read_artifact(
        loaded, output / "assurance_final_outputs/relazione_riconciliazione_audit.docx"
    )
    read_artifact(loaded, output / "richieste_mirate_evidenze.xlsx")
    read_artifact(
        loaded, output / "assurance_final_outputs/richieste_mirate_evidenze.xlsx"
    )
    read_artifact(loaded, output / "reconciliation_results.json")
    read_artifact(
        loaded, output / "assurance_final_outputs/reconciliation_results.json"
    )
    read_artifact(loaded, output / "assurance_receipts.json")
    read_artifact(loaded, output / "artifact_card.md")
    read_artifact(loaded, output / "final_artifacts.json")
    read_artifact(loaded, output / "ui_decisions.json")
    with review_transport(tmp_path) as call:
        call(
            "validate_open_item_reconciliation_review",
            {
                "review_payload_path": str(output / "review_payload.json"),
                "client_engagement": str(Path(loaded["run_root"]) / "context.json"),
            },
        )
    measured = report(loaded)
    assert measured["collection_status"] == "complete", measured["errors"]


def test_public_assurance_replay_reports_status_without_reemitting_source_rows(
    reconciled_case,
):
    loaded = reconciled_case
    output = Path(loaded["calculation_dir"])
    completed = subprocess.run(
        [
            sys.executable,
            "-I",
            "-B",
            str(PLUGIN / "scripts/audit_assurance.py"),
            "--client-engagement",
            str(Path(loaded["run_root"]) / "context.json"),
            "validate-run-json",
            str(output),
        ],
        capture_output=True,
        text=True,
        timeout=90,
    )

    assert completed.returncode == 0, completed.stderr
    summary = json.loads(completed.stdout)
    assert summary["ok"] is True
    assert summary["reconciliation_row_count"] == 3
    assert summary["assurance"]["report_ready"] is False
    assert "reconciliation_rows" not in completed.stdout
    assert "Alfa" not in completed.stdout
    assert "PRIVATE" not in completed.stdout
    measured = report(loaded)
    assert measured["collection_status"] == "complete", measured["errors"]
    assert measured["emissions"][-1]["channel"] == "cli:open-item-reconciliation"
    assert measured["emissions"][-1]["unmapped"] is False
    assert sum(source["sent_cells"] for source in measured["sources"]) == 0


@pytest.mark.parametrize("changed_field", ["amount", "review_notes"])
@pytest.mark.parametrize("item_type", ["manual_review", "closure_evidence_review"])
def test_edited_review_cannot_claim_the_producers_source_coordinates(
    reconciled_case, tmp_path, changed_field, item_type
):
    loaded = reconciled_case
    output = Path(loaded["calculation_dir"])
    path = output / "review_payload.json"
    review = json.loads(path.read_text())
    item = next(item for item in review["items"] if item["item_type"] == item_type)
    item["data"][changed_field] = "UNKNOWN_CASE_DATA"
    path.write_text(json.dumps(review))

    with review_transport(tmp_path) as call:
        index = call(
            "validate_open_item_reconciliation_review",
            {
                "review_payload_path": str(path),
                "client_engagement": str(Path(loaded["run_root"]) / "context.json"),
            },
        )["structuredContent"]
        handle = next(
            item["case_handle"]
            for item in index["model_context_index"]["cases"]
            if item["item_type"] == item_type
        )
        selected = call(
            "get_open_item_reconciliation_case_context",
            {
                "persistence_token": index["review_reference"]["persistence_token"],
                "case_handles": [handle],
            },
        )["structuredContent"]

    assert selected["cases"][0]["semantic_fields"][changed_field] == "UNKNOWN_CASE_DATA"
    measured = report(loaded)
    assert measured["collection_status"] == "incomplete"
    assert "sent_data_has_no_source_coordinates" in measured["errors"]


@pytest.mark.parametrize("note_location", ["data", "evidence"])
def test_unknown_producer_check_note_is_not_classified_as_locator_metadata(
    reconciled_case, tmp_path, note_location
):
    from archive_projection import register_private_value

    loaded = reconciled_case
    output = Path(loaded["calculation_dir"])
    path = output / "review_payload.json"
    review = json.loads(path.read_text())
    check = next(
        item for item in review["items"] if item["item_type"] == "check_exception"
    )
    containers = {"data": check["data"], "evidence": check["evidence"][0]}
    containers[note_location]["note"] = "UNMAPPED_SOURCE_VALUE"
    path.write_text(json.dumps(review))
    # Even a producer-bound value needs the declared field transformation;
    # the private digest alone does not make arbitrary diagnostic text metadata.
    register_private_value("open-item-reconciliation", output, review)

    with review_transport(tmp_path) as call:
        index = call(
            "validate_open_item_reconciliation_review",
            {
                "review_payload_path": str(path),
                "client_engagement": str(Path(loaded["run_root"]) / "context.json"),
            },
        )["structuredContent"]
        handle = next(
            item["case_handle"]
            for item in index["model_context_index"]["cases"]
            if item["item_type"] == "check_exception"
        )
        selected = call(
            "get_open_item_reconciliation_case_context",
            {
                "persistence_token": index["review_reference"]["persistence_token"],
                "case_handles": [handle],
                "include_exact_identifiers": True,
            },
        )["structuredContent"]

    assert "UNMAPPED_SOURCE_VALUE" in json.dumps(selected)
    measured = report(loaded)
    assert measured["collection_status"] == "incomplete"
    assert "sent_data_has_no_source_coordinates" in measured["errors"]


def test_private_replay_requested_directly_still_records_its_unmapped_rows(
    reconciled_case,
):
    loaded = reconciled_case
    completed = subprocess.run(
        [
            sys.executable,
            "-I",
            "-B",
            str(PLUGIN / "scripts/audit_assurance.py"),
            "--client-engagement",
            str(Path(loaded["run_root"]) / "context.json"),
            "validate-run-json",
            loaded["calculation_dir"],
            "--private-review-bridge",
        ],
        capture_output=True,
        text=True,
        timeout=90,
    )

    assert completed.returncode == 0, completed.stderr
    response = json.loads(completed.stdout)
    assert len(response["result"]["reconciliation_rows"]) == 3
    assert "Alfa" in completed.stdout
    measured = report(loaded)
    assert measured["collection_status"] == "incomplete"
    assert measured["emissions"][-1]["channel"] == "cli:open-item-reconciliation"
    assert measured["emissions"][-1]["unmapped"] is True


@pytest.mark.parametrize("claim", ["no_residual", "fully_closed"])
def test_partial_payment_cannot_erase_its_balance_or_claim_full_closure(
    reconciled_case, monkeypatch, claim
):
    from tests.plugins.test_open_item_reconciliation_plugin import load_raw_input_runner

    loaded = reconciled_case
    runner = load_raw_input_runner()
    original = runner.finalize_assurance_run

    def altered_claim(**kwargs):
        rows = copy.deepcopy(kwargs["reconciliation_rows"])
        partial = next(
            row for row in rows if row["reconciliation_status"] == "partially_paid"
        )
        if claim == "no_residual":
            partial["residual_amount"] = "0.00"
        else:
            partial["reconciliation_status"] = "closed"
        return original(**{**kwargs, "reconciliation_rows": rows})

    monkeypatch.setattr(runner, "finalize_assurance_run", altered_claim)
    assumptions = json.loads(
        (Path(loaded["output_dir"]) / "assumptions.json").read_text()
    )

    with pytest.raises(
        ValueError,
        match=(
            "exact evidence allocation and retained residual"
            if claim == "no_residual"
            else "balanced evidence allocation relationship"
        ),
    ):
        runner.run_raw_input_reconciliation(
            input_dir=loaded["context"]["input_dir"],
            prepared_client_engagement=loaded["context"],
            assumptions=assumptions,
            output_subdirectory="reconciliation",
        )
