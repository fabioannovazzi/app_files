from __future__ import annotations

import hashlib
import json
import socket
import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from lxml import etree as ET

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "plugins/lucia/scripts"))
from legal_documents import compare, prepare
from legal_docx import W, inspect_docx, paragraph_text, parse_xml
from legal_proofreading import main as proofreading_main
from legal_proofreading import render_audit, scaffold, validate_audit
from legal_word_bridge import main as word_main
from legal_word_bridge import verify_word_result


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def refuse(*args, **kwargs):
        raise AssertionError("Lucia document helpers must not connect to a server")

    monkeypatch.setattr(socket, "create_connection", refuse)
    monkeypatch.setattr(socket.socket, "connect", refuse)


def rewrite(source, target, transform):
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    transform(parts)
    with ZipFile(target, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return target


def document(tmp_path, text="Pagamento entro 30 giorni."):
    path = tmp_path / "original.docx"
    doc = Document()
    doc.add_paragraph(text)
    doc.add_paragraph("Foro di Milano.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Allegato A"
    doc.save(path)
    return path


def revised(
    source, target, *, comment=False, tracked=True, replacement="60", existing=False
):
    def change(parts):
        root = ET.fromstring(parts["word/document.xml"])
        paragraph = next(root.iter(W + "p"))
        run = paragraph.find(W + "r")
        paragraph.remove(run)
        for value, kind, cid in [
            ("Pagamento entro ", None, None),
            ("30", "del", "11"),
            (replacement, "ins", "12"),
            (" giorni.", None, None),
        ]:
            if not tracked and kind == "del":
                continue
            owner = (
                ET.SubElement(
                    paragraph, W + kind, {W + "id": cid, W + "author": "Lucia"}
                )
                if kind and tracked
                else paragraph
            )
            r = ET.SubElement(owner, W + "r")
            ET.SubElement(
                r,
                W + ("delText" if kind == "del" and tracked else "t"),
                {"{http://www.w3.org/XML/1998/namespace}space": "preserve"},
            ).text = value
        if existing:
            second = list(root.iter(W + "p"))[1]
            second.clear()
            insertion = ET.SubElement(
                second, W + "ins", {W + "id": "7", W + "author": "Previous reviewer"}
            )
            ET.SubElement(ET.SubElement(insertion, W + "r"), W + "t").text = (
                "Foro di Milano."
            )
        if comment:
            ET.SubElement(paragraph, W + "commentRangeStart", {W + "id": "0"})
            ET.SubElement(paragraph, W + "commentRangeEnd", {W + "id": "0"})
            ET.SubElement(
                ET.SubElement(paragraph, W + "r"),
                W + "commentReference",
                {W + "id": "0"},
            )
            comments = ET.Element(W + "comments")
            c = ET.SubElement(
                comments, W + "comment", {W + "id": "0", W + "author": "Lucia"}
            )
            ET.SubElement(
                ET.SubElement(ET.SubElement(c, W + "p"), W + "r"), W + "t"
            ).text = "Confermare il termine."
            parts["word/comments.xml"] = ET.tostring(comments)
            rels = ET.fromstring(parts["word/_rels/document.xml.rels"])
            ET.SubElement(
                rels,
                "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                Id="rId999",
                Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
                Target="comments.xml",
            )
            parts["word/_rels/document.xml.rels"] = ET.tostring(rels)
            content_types = ET.fromstring(parts["[Content_Types].xml"])
            ET.SubElement(
                content_types,
                "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                PartName="/word/comments.xml",
                ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
            )
            parts["[Content_Types].xml"] = ET.tostring(content_types)
        parts["word/document.xml"] = ET.tostring(root)

    return rewrite(source, target, change)


def word_request(source, mode="tracked"):
    return {
        "mode": mode,
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "edits": [
            {
                "id": "E1",
                "kind": "replace",
                "anchor": "word/document.xml#p1",
                "old": "30",
                "new": "60",
                "basis": "Termine concordato.",
            }
        ],
    }


def completed_audit(tmp_path):
    source = tmp_path / "sample.txt"
    source.write_text(
        "Alfa S.r.l.\nArticolo 1. Termine: 30 giorni.\n", encoding="utf-8"
    )
    run = tmp_path / "review"
    pack = prepare(run, [source], "controllo-documento", ["Proofreading"])
    audit = json.loads(scaffold(run).read_text())
    audit["context"].update(
        {
            key: "Non fornito; sola revisione testuale"
            for key in audit["context"]
            if key != "assumptions"
        }
    )
    audit["context"]["language"] = "it"
    record = audit["documents"]["D001"]
    record["reviewed_anchors"] = [u["anchor"] for u in pack["sources"][0]["units"]]
    for mapping in record["maps"].values():
        mapping["note"] = "Testo breve letto; nessuna definizione formale ulteriore."
    for review_pass in record["passes"].values():
        review_pass.update(
            status="complete", note="Controllo eseguito sul testo disponibile."
        )
    record["visual_review"].update(
        status="complete", evidence="File di testo letto integralmente."
    )
    return run, audit


def test_docx_inspection_separates_revision_views_and_table_paragraphs(tmp_path):
    source = document(tmp_path)
    result = revised(source, tmp_path / "tracked.docx")

    inspected = inspect_docx(result)

    assert inspected["paragraphs"][0]["text"] == "Pagamento entro 60 giorni."
    assert inspected["paragraphs"][0]["original_text"] == "Pagamento entro 30 giorni."
    assert inspected["paragraphs"][2]["in_table"] is True
    assert inspected["paragraphs"][0]["revisions"][0]["text"] == "30"


def test_docx_inspection_exposes_hidden_hyperlink_fields_and_metadata(tmp_path):
    source = document(tmp_path)

    def add_metadata(parts):
        root = ET.fromstring(parts["word/document.xml"])
        p = next(root.iter(W + "p"))
        link = ET.SubElement(
            p,
            W + "hyperlink",
            {
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id": "rId999"
            },
        )
        ET.SubElement(ET.SubElement(link, W + "r"), W + "t").text = "studio@example.it"
        r = ET.SubElement(p, W + "r")
        ET.SubElement(r, W + "instrText").text = (
            ' HYPERLINK "https://example.invalid/hidden" '
        )
        props = ET.SubElement(r, W + "rPr")
        ET.SubElement(props, W + "vanish")
        parts["word/document.xml"] = ET.tostring(root)
        rels = ET.fromstring(parts["word/_rels/document.xml.rels"])
        ET.SubElement(
            rels,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            Id="rId999",
            Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            Target="mailto:wrong@example.it",
            TargetMode="External",
        )
        parts["word/_rels/document.xml.rels"] = ET.tostring(rels)

    artifact = rewrite(source, tmp_path / "links.docx", add_metadata)

    inspection = inspect_docx(artifact)

    paragraph = inspection["paragraphs"][0]
    assert (
        paragraph["hyperlinks"][0]["relationship"]["Target"]
        == "mailto:wrong@example.it"
    )
    assert (
        paragraph["fields"][0]["instruction"]
        == ' HYPERLINK "https://example.invalid/hidden" '
    )
    assert paragraph["runs"][-1]["properties"][0]["property"] == "vanish"
    assert "docProps/core.xml" in inspection["properties"]
    assert inspection["style_defaults"]


@pytest.mark.parametrize(
    "view, expected", [("final", "60"), ("original", "30"), ("all", "3060")]
)
def test_revision_text_views_are_explicit(view, expected):
    root = parse_xml(
        f'<w:p xmlns:w="{W[1:-1]}"><w:del><w:r><w:delText>30</w:delText></w:r></w:del><w:ins><w:r><w:t>60</w:t></w:r></w:ins></w:p>'.encode()
    )
    assert paragraph_text(root, view) == expected


def test_xml_entity_declaration_is_rejected():
    with pytest.raises(ValueError, match="DTD"):
        parse_xml(
            b'<!DOCTYPE document [<!ENTITY x SYSTEM "file:///etc/passwd">]><document>&x;</document>'
        )


def test_utf16_xml_entity_declaration_is_rejected():
    xml = '<!DOCTYPE document [<!ENTITY x "hidden">]><document>&x;</document>'
    with pytest.raises(ValueError, match="DTD"):
        parse_xml(xml.encode("utf-16"))


def test_word_result_rejects_spaces_that_word_would_drop(tmp_path):
    source = document(tmp_path)
    tracked = revised(source, tmp_path / "tracked.docx")

    def remove_space_preservation(parts):
        root = ET.fromstring(parts["word/document.xml"])
        for node in root.iter(W + "t"):
            node.attrib.pop("{http://www.w3.org/XML/1998/namespace}space", None)
        parts["word/document.xml"] = ET.tostring(root)

    broken = rewrite(tracked, tmp_path / "broken.docx", remove_space_preservation)

    result = verify_word_result(source, broken, word_request(source))

    assert result["valid"] is False
    assert any("Word may join words" in error for error in result["errors"])


def test_proofreading_whitespace_does_not_count_as_source_evidence(tmp_path):
    run, audit = completed_audit(tmp_path)
    audit["documents"]["D001"]["maps"]["entities"]["entries"] = [
        {
            "label": "Alfa",
            "note": "Parte",
            "citations": [{"source_id": "D001", "anchor": "line:1", "quote": " \t "}],
        }
    ]

    result = validate_audit(run, audit)

    assert result["valid"] is False
    assert any("quotation not found" in error for error in result["errors"])


def test_literal_comparison_uses_final_text_without_revision_or_metadata_noise(
    tmp_path,
):
    source = document(tmp_path)
    result = revised(source, tmp_path / "tracked.docx")
    run = tmp_path / "comparison"
    prepare(run, [source, result], "confronto-documenti", ["Termine"])

    diff = compare(run, "D001", "D002").read_text()

    assert "+Pagamento entro 60 giorni." in diff
    assert "3060" not in diff
    assert "hyperlinks" not in diff


def test_original_revision_comparison_retains_prior_text(tmp_path):
    source = document(tmp_path)
    result = revised(source, tmp_path / "tracked.docx")
    run = tmp_path / "comparison"
    prepare(run, [source, result], "confronto-documenti", ["Termine"])

    diff = compare(run, "D001", "D002", view="original").read_text()

    assert diff.startswith("No differences")


def test_host_word_result_verifies_native_replacement_and_original_roundtrip(tmp_path):
    source = document(tmp_path)
    result = revised(source, tmp_path / "tracked.docx")

    verification = verify_word_result(source, result, word_request(source))

    assert verification["valid"] is True
    assert verification["checks"] == [{"id": "E1", "native_revision": True}]
    assert verification["requires_visual_review"] is True


def test_host_word_result_rejects_clean_replacement_when_tracked_requested(tmp_path):
    source = document(tmp_path)
    result = revised(source, tmp_path / "clean.docx", tracked=False)

    verification = verify_word_result(source, result, word_request(source))

    assert verification["valid"] is False
    assert any("native insertion/deletion" in e for e in verification["errors"])


def test_host_word_result_detects_silently_missed_edit(tmp_path):
    source = document(tmp_path)
    result = tmp_path / "unchanged.docx"
    result.write_bytes(source.read_bytes())

    verification = verify_word_result(source, result, word_request(source))

    assert verification["valid"] is False
    assert any("proposed text differs" in e for e in verification["errors"])


def test_host_word_result_validates_comment_at_requested_anchor(tmp_path):
    source = document(tmp_path)
    result = revised(source, tmp_path / "commented.docx", comment=True)
    request = word_request(source)
    request["edits"].append(
        {
            "id": "C1",
            "kind": "comment",
            "anchor": "word/document.xml#p1",
            "old": "30",
            "comment": "Confermare il termine.",
            "basis": "Istruzione dello studio.",
        }
    )

    verification = verify_word_result(source, result, request)

    assert verification["valid"] is True
    assert {"id": "C1", "comment_anchored": True} in verification["checks"]


def test_host_word_result_refuses_source_hash_drift(tmp_path):
    source = document(tmp_path)
    result = revised(source, tmp_path / "tracked.docx")
    request = word_request(source)
    request["source_sha256"] = "different"

    with pytest.raises(ValueError, match="source hash"):
        verify_word_result(source, result, request)


def test_proofreading_cannot_be_complete_with_an_unfinished_pass(tmp_path):
    run, audit = completed_audit(tmp_path)
    audit["documents"]["D001"]["passes"]["formatting"]["status"] = "not-reviewed"

    result = validate_audit(run, audit)

    assert result["valid"] is True
    assert result["complete"] is False


def test_proofreading_rejects_a_made_up_source_quotation(tmp_path):
    run, audit = completed_audit(tmp_path)
    audit["findings"] = [
        {
            "id": "F1",
            "severity": "high",
            "category": "entities",
            "kind": "error",
            "issue": "Nome diverso",
            "reason": "Identità incerta",
            "fix": "Chiedere conferma",
            "citations": [
                {"source_id": "D001", "anchor": "line:1", "quote": "Alfa S.p.A."}
            ],
        }
    ]

    result = validate_audit(run, audit)

    assert result["valid"] is False
    assert any("quotation not found" in error for error in result["errors"])


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_proofreading_renders_local_report_with_coverage_and_escaped_content(
    tmp_path, language
):
    run, audit = completed_audit(tmp_path)
    audit["context"]["language"] = language
    audit["takeaways"] = ["<script>alert('client data')</script>"]
    path = run / "proofreading.json"
    path.write_text(json.dumps(audit))

    report = render_audit(run, path).read_text()

    assert f'<html lang="{language}">' in report
    assert "&lt;script&gt;" in report
    assert "<script>" not in report
    assert "Alfa" not in report  # No invented finding in a clean control.


def test_failed_proofreading_rerender_retires_previous_output(tmp_path):
    run, audit = completed_audit(tmp_path)
    path = run / "proofreading.json"
    path.write_text(json.dumps(audit))
    render_audit(run, path)
    audit["documents"] = {}
    path.write_text(json.dumps(audit))

    with pytest.raises(ValueError, match="Retain every"):
        render_audit(run, path)

    assert not (run / "proofreading.html").exists()
    assert (run / "previous-1-proofreading.html").exists()


def test_word_bridge_rejects_lost_preexisting_revision(tmp_path):
    base = document(tmp_path)
    source = revised(base, tmp_path / "existing.docx", replacement="30", existing=True)
    request = word_request(source)
    # A host editor silently flattened an earlier reviewer's edit.
    result = revised(base, tmp_path / "flattened.docx")

    verification = verify_word_result(source, result, request)

    assert verification["valid"] is False
    assert any("pre-existing revision" in e for e in verification["errors"])


def test_word_bridge_rejects_unwired_comment_even_if_text_exists(tmp_path):
    source = document(tmp_path)
    commented = revised(source, tmp_path / "commented.docx", comment=True)
    result = rewrite(
        commented,
        tmp_path / "unwired.docx",
        lambda parts: parts.pop("word/_rels/document.xml.rels"),
    )

    verification = verify_word_result(source, result, word_request(source))

    assert verification["valid"] is False
    assert any("relationship or content type" in e for e in verification["errors"])


def test_word_bridge_rejects_unrelated_package_changes(tmp_path):
    source = document(tmp_path)
    edited = revised(source, tmp_path / "edited.docx")
    result = rewrite(
        edited,
        tmp_path / "lost-styles.docx",
        lambda parts: parts.pop("word/styles.xml"),
    )

    verification = verify_word_result(source, result, word_request(source))

    assert verification["valid"] is False
    assert any("word/styles.xml" in e for e in verification["errors"])


def test_word_bridge_rejects_ambiguous_same_word_without_offset(tmp_path):
    source = document(tmp_path, text="30 oppure 30 giorni.")
    result = tmp_path / "copy.docx"
    result.write_bytes(source.read_bytes())

    with pytest.raises(ValueError, match="ambiguous quote"):
        verify_word_result(source, result, word_request(source))


def test_word_bridge_rejects_overlapping_approved_edits(tmp_path):
    source = document(tmp_path)
    result = revised(source, tmp_path / "edited.docx")
    request = word_request(source)
    request["edits"].append({**request["edits"][0], "id": "E2"})

    with pytest.raises(ValueError, match="overlap"):
        verify_word_result(source, result, request)


@pytest.mark.parametrize("enforcement", ["1", "true", "on"])
def test_word_bridge_rejects_enforced_protection_in_original(tmp_path, enforcement):
    source = document(tmp_path)

    def protect(parts):
        root = ET.fromstring(parts["word/settings.xml"])
        ET.SubElement(
            root,
            W + "documentProtection",
            {W + "enforcement": enforcement, W + "edit": "readOnly"},
        )
        parts["word/settings.xml"] = ET.tostring(root)

    protected = rewrite(source, tmp_path / "protected.docx", protect)
    result = revised(source, tmp_path / "unprotected-result.docx")

    with pytest.raises(ValueError, match="enforced document protection"):
        verify_word_result(protected, result, word_request(protected))


def test_document_inspection_does_not_treat_disabled_protection_as_enforced(tmp_path):
    source = document(tmp_path)

    def protect(parts):
        root = ET.fromstring(parts["word/settings.xml"])
        ET.SubElement(root, W + "documentProtection", {W + "enforcement": "false"})
        parts["word/settings.xml"] = ET.tostring(root)

    result = rewrite(source, tmp_path / "not-protected.docx", protect)

    assert inspect_docx(result)["protection_enforced"] is False


def test_word_bridge_cli_saves_failure_report_for_missed_edits(tmp_path):

    source = document(tmp_path)
    result = tmp_path / "copy.docx"
    result.write_bytes(source.read_bytes())
    request = tmp_path / "request.json"
    request.write_text(json.dumps(word_request(source)))
    report = tmp_path / "verification.json"

    status = word_main(
        [
            "--original",
            str(source),
            "--result",
            str(result),
            "--request",
            str(request),
            "--report",
            str(report),
        ]
    )

    assert status == 2
    assert json.loads(report.read_text())["valid"] is False


def test_proofreading_cli_writes_supported_finding_and_reference_map(tmp_path):

    run, audit = completed_audit(tmp_path)
    quote = {"source_id": "D001", "anchor": "line:1", "quote": "Alfa S.r.l."}
    audit["documents"]["D001"]["maps"]["entities"]["entries"] = [
        {"label": "Alfa", "note": "Denominazione nelle premesse", "citations": [quote]}
    ]
    audit["findings"] = [
        {
            "id": "F1",
            "severity": "low",
            "category": "entities",
            "kind": "uncertain",
            "issue": "Poteri non documentati",
            "reason": "Il nome non prova i poteri",
            "fix": "Richiedere la prova dei poteri",
            "citations": [quote],
        }
    ]
    (run / "proofreading.json").write_text(json.dumps(audit))

    status = proofreading_main(["render", "--run-dir", str(run)])

    assert status == 0
    assert "Alfa S.r.l." in (run / "proofreading.html").read_text()
    assert "Richiedere la prova" in (run / "proofreading.html").read_text()


def test_proofreading_rejects_empty_completion_note_and_missing_visual_evidence(
    tmp_path,
):
    run, audit = completed_audit(tmp_path)
    audit["documents"]["D001"]["passes"]["numbering"]["note"] = ""
    audit["documents"]["D001"]["visual_review"]["evidence"] = ""

    verification = validate_audit(run, audit)

    assert verification["valid"] is False
    assert len(verification["errors"]) == 2
