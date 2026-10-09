"""Mechanism regressions; these tests do not establish native-host acceptance."""

from __future__ import annotations

import importlib.util
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests._plugin_cli import workflow_cli
from tests.plugins import test_archive_organization_plugin as archive
from tests.plugins import test_check_entries_plugin as vouching
from tests.plugins import test_concordato_plan_review_plugin as concordato
from tests.plugins import test_concordato_preventivo_semantics as concordato_semantics
from tests.plugins import test_journal_sampling_plugin as sampling
from tests.plugins import test_open_item_reconciliation_plugin as open_items
from tests.plugins import test_report_builder_plugin as reports
from tests.plugins import test_vera_new_client_component as new_client
from tests.plugins.test_bilancio_native_workspace import NODE
from tests.plugins.test_journal_bank_reconciliation_plugin import (
    _prepare_sealed_mcp_review_run,
)
from tests.plugins.test_treasury_delivery import SCRIPTS, archived_case

ROOT = Path(__file__).resolve().parents[2]
SERVER = ROOT / "plugins/vera/mcp/workspace.cjs"


def workspace_module():
    # Mirror Python's sibling-import path when the maintained adapter runs as a CLI.
    scripts = str(ROOT / "plugins/vera/scripts")
    if scripts not in sys.path:
        sys.path.insert(0, scripts)
    spec = importlib.util.spec_from_file_location(
        "tested_vera_native_workspace",
        ROOT / "plugins/vera/scripts/native_workspace.py",
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def configure(monkeypatch, base: Path, bindings: list[dict]) -> dict[str, str]:
    path = base / "workspace-bindings.json"
    path.write_text(
        json.dumps(
            {
                "tenant_id": "fictional-studio",
                "actor_id": "fictional-reviewer",
                "bindings": bindings,
            }
        )
    )
    env = {
        "VERA_WORKSPACE_BINDINGS": str(path),
        "VERA_WORKSPACE_TENANT_ID": "fictional-studio",
        "VERA_WORKSPACE_ACTOR_ID": "fictional-reviewer",
        "VERA_WORKSPACE_ROLES": "REVIEWER",
        "VERA_WORKSPACE_PYTHON": sys.executable,
    }
    for key, value in env.items():
        monkeypatch.setenv(key, value)
    return env


def archived_review_binding(output: Path, component: str) -> tuple[dict, Path]:
    candidate = output
    while not (candidate / "context.json").is_file():
        candidate = candidate.parent
    context_path = candidate / "context.json"
    context = json.loads(context_path.read_bytes())
    return {
        "work_ref": "fictional-" + component,
        "client_root": str(context_path.parents[5]),
        "client_id": context["client_id"],
        "engagement_id": context["engagement_id"],
        "run_id": context["run_id"],
        "workflow_id": component,
    }, candidate / "outputs"


@pytest.fixture
def sampling_workspace(tmp_path, monkeypatch):
    output, _ = sampling._real_journal_review_case(sampling.load_core(), tmp_path)
    binding, root_output = archived_review_binding(output, "journal-sampling")
    configure(monkeypatch, tmp_path, [binding])
    return workspace_module(), binding, root_output, output


@pytest.fixture
def vouching_workspace(tmp_path, monkeypatch):
    _, output, _, _ = vouching._supported_assurance_run(monkeypatch, tmp_path)
    binding, root_output = archived_review_binding(output, "check-entries")
    configure(monkeypatch, tmp_path, [binding])
    return workspace_module(), binding, root_output, output


@pytest.fixture
def new_client_workspace(tmp_path, monkeypatch, vera_workflow_workspace):
    # Match the existing component tests' fixed fixture clock in the Node child.
    # This does not attest currentness of the bundled legal source registry.
    clock = tmp_path / "new-client-clock.cjs"
    clock.write_text(
        "const RealDate = Date; global.Date = class extends RealDate {"
        "constructor(...args) { super(...(args.length ? args : ['2026-07-20T12:00:00Z'])); }"
        "static now() { return new RealDate('2026-07-20T12:00:00Z').getTime(); } };"
    )
    monkeypatch.setenv("NODE_OPTIONS", "--require " + str(clock))
    output = new_client._generate_package(vera_workflow_workspace)
    binding, _ = archived_review_binding(output, "new-client")
    configure(monkeypatch, tmp_path, [binding])
    return workspace_module(), binding, output


@pytest.fixture
def archive_workspace(tmp_path, monkeypatch):
    context_path, snapshot, original_files = archive.scenario._prepared_run(tmp_path)
    result = archive.core.build_review_package(
        context_path, archive.scenario._proposals(tmp_path, snapshot), language="it"
    )
    output = Path(result["output_dir"])
    binding, _ = archived_review_binding(output, "archive-organization")
    configure(monkeypatch, tmp_path, [binding])
    return workspace_module(), binding, output, original_files


@pytest.fixture(params=[False, True], ids=["root", "intake"])
def file_preparation_workspace(request, tmp_path, monkeypatch, vera_workflow_workspace):
    content = (
        "Certificazione Unica 2025. Codice fiscale TSTUSR80A01H501U. "
        "Sostituto d'imposta Fornitore Test SRL. Redditi lavoro dipendente 24.000,00."
    )
    workspace = vera_workflow_workspace(
        "client-file-preparation",
        input_files={
            "CU_2025.txt": content,
            "second_source.txt": content + "\nFonte distinta da rivedere.",
        },
    )
    output = (
        workspace["output_dir"] / "intake" if request.param else workspace["output_dir"]
    )
    subprocess.run(
        workflow_cli(
            ROOT
            / "plugins/client-file-preparation/scripts/build_file_preparation_outputs.py"
        )
        + [
            str(workspace["input_dir"]),
            "--client-engagement",
            str(workspace["context_path"]),
            "--year",
            "2025",
            "--language",
            "it",
            "--out",
            str(output),
        ],
        capture_output=True,
        text=True,
        check=True,
    )
    binding, root_output = archived_review_binding(output, "client-file-preparation")
    configure(monkeypatch, tmp_path, [binding])
    return workspace_module(), binding, root_output, output


@pytest.fixture(params=[False, True], ids=["root", "report"])
def report_workspace(request, tmp_path, monkeypatch, vera_workflow_workspace):
    workspace = vera_workflow_workspace(
        "report-builder", input_files={"budget.csv": "line,amount\nA,10\nB,20\n"}
    )
    output = (
        workspace["output_dir"] / "report" if request.param else workspace["output_dir"]
    )
    source = Path(workspace["context"]["input_bindings"][0]["path"])
    reports.load_core().build_report(
        source,
        output,
        report_type="management_report",
        language="it",
        document_language="auto",
        run_id=workspace["run_id"],
        client_engagement=workspace["context"],
    )
    binding, root_output = archived_review_binding(output, "report-builder")
    configure(monkeypatch, tmp_path, [binding])
    return workspace_module(), binding, root_output, output, source


@pytest.fixture(params=[False, True], ids=["root", "reviewed-output"])
def concordato_workspace(request, tmp_path, monkeypatch):
    core = concordato.load_core()
    semantic = concordato_semantics._load_semantic()
    received = tmp_path / "received"
    received.mkdir()
    for name, rows in {
        "case_material_a.xlsx": [
            ["Plan component", "Amount"],
            ["Cash generation", 500],
        ],
        "case_material_b.xlsx": [
            ["Attestation section", "Reference"],
            ["Feasibility", "Section four"],
        ],
        "case_material_c.xlsx": [["Creditor", "Claim"], ["Secured Bank", 600]],
    }.items():
        concordato._save_workbook(received / name, rows)
    inputs, root_output, run_id, _ = concordato._start_managed_concordato_run(
        tmp_path, received
    )
    inspection = core.run_concordato_review(
        inputs,
        tmp_path / "fixture-inspection",
        reference_date="2026-03-31",
        language="it",
        tolerance="0.01",
    )
    # Reuse the existing fictional professional model, then restore the exact
    # imported paths before the maintained producer binds its source receipts.
    model = concordato_semantics._reviewed_case_model(
        semantic,
        [
            {**row, "relative_path": Path(row["relative_path"]).name}
            for row in inspection.inventory
        ],
        missing_attestation=True,
    )
    paths = {
        row["source_artifact_ref"]: row["relative_path"] for row in inspection.inventory
    }
    for document in model["document_perimeter"]["documents"]:
        document["relative_path"] = paths[document["source_artifact_ref"]]
    recipe = semantic.review_concordato_case_model(
        inspection.inventory,
        model,
        reviewer_ref="fictional-qualified-reviewer",
        reviewed_on="2026-07-26",
        reference_date="2026-03-31",
    )
    output = root_output / "reviewed-output" if request.param else root_output
    core.run_concordato_review(
        inputs,
        output,
        reference_date="2026-03-31",
        language="it",
        document_language="it",
        tolerance="0.01",
        semantic_recipe=recipe,
        run_id=run_id,
        input_path_ref="inputs",
        output_path_ref=output.relative_to(root_output.parent).as_posix(),
    )
    binding, _ = archived_review_binding(output, "concordato-plan-review")
    configure(monkeypatch, tmp_path, [binding])
    return workspace_module(), binding, root_output, output, inputs


def test_concordato_native_selected_context_keeps_procedure_and_omits_technical_metadata(
    concordato_workspace,
):
    service, binding, _, output, _ = concordato_workspace
    payload = json.loads((output / "review_payload.json").read_bytes())
    selected = next(
        row for row in payload["items"] if row["item_type"] == "procedure_identity"
    )

    context = service.dispatch(
        "vera_workspace_explain",
        {"work_ref": binding["work_ref"], "item_id": selected["id"]},
    )["untrusted_evidence"]

    assert context["returned_item_count"] == 1
    assert context["limit"] == 1
    assert context["technical_metadata_removed"] is True
    assert context["items"][0]["id"] == selected["id"]
    assert (
        context["items"][0]["data"]["debtor_name"] == "Northwind Restructuring S.p.A."
    )
    assert "source_artifact_ref" not in json.dumps(context)
    assert "output_path" not in json.dumps(context)
    assert str(output) not in json.dumps(context)


def test_concordato_native_source_discussion_uses_engine_aliases(concordato_workspace):
    service, binding, _, output, _ = concordato_workspace
    payload = json.loads((output / "review_payload.json").read_bytes())
    selected = next(
        row for row in payload["items"] if row["item_type"] == "source_inventory"
    )

    context = service.dispatch(
        "vera_workspace_explain",
        {"work_ref": binding["work_ref"], "item_id": selected["id"]},
    )["untrusted_evidence"]

    serialized = json.dumps(context)
    assert context["source_labels_replaced_with_stable_aliases"] is True
    assert context["returned_item_count"] == 1
    assert "[source-" in serialized
    assert "case_material_" not in serialized
    assert "sha256" not in serialized
    assert "source_artifact_ref" not in serialized


def test_concordato_native_save_reopens_proposal_without_mutating_confirmed_case(
    concordato_workspace,
):
    service, binding, _, output, _ = concordato_workspace
    before = (output / "concordato_case_model.json").read_bytes()
    selected = next(
        row
        for row in json.loads((output / "review_payload.json").read_bytes())["items"]
        if row["item_type"] == "procedure_identity"
    )
    current = service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "item_id": selected["id"]},
    )
    note = "Confrontare il riferimento con il documento depositato."

    result = service.dispatch(
        "vera_workspace_save",
        {
            "work_ref": binding["work_ref"],
            "item_id": selected["id"],
            "revision": current["revision"],
            "human_reviewed": True,
            "idempotency_key": "fictional-concordato-save",
            "decisions": [
                {
                    "item_id": selected["id"],
                    "action": "mark_unclear",
                    "reviewer_note": note,
                }
            ],
        },
    )

    assert result["saved"] is True
    assert (output / "concordato_case_model.json").read_bytes() == before
    reopened = service.dispatch(
        "vera_workspace_view", {"work_ref": binding["work_ref"]}
    )
    assert reopened["data"]["ui_decisions"]["decisions"][0]["reviewer_note"] == note
    assert reopened["data"]["final_artifacts"]["final_ready"] is False


def test_concordato_native_memo_apply_refuses_before_unreplayable_word_write(
    concordato_workspace,
):
    service, binding, root_output, output, inputs = concordato_workspace
    case_bytes = (output / "concordato_case_model.json").read_bytes()
    sources = {
        path.relative_to(inputs).as_posix(): path.read_bytes()
        for path in inputs.rglob("*")
        if path.is_file()
    }
    selected = next(
        row
        for row in json.loads((output / "review_payload.json").read_bytes())["items"]
        if row["item_type"] == "codex_review_memo"
    )
    current = service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "item_id": selected["id"]},
    )
    memo = "# Memo di revisione\n\nL'attestazione resta da acquisire e verificare."

    before = {
        path.relative_to(output).as_posix(): path.read_bytes()
        for path in output.rglob("*")
        if path.is_file()
    }

    with pytest.raises(ValueError, match="non supera il replay del motore"):
        service.dispatch(
            "vera_workspace_apply",
            {
                "work_ref": binding["work_ref"],
                "item_id": selected["id"],
                "revision": current["revision"],
                "human_reviewed": True,
                "idempotency_key": "fictional-concordato-memo",
                "decisions": [
                    {"item_id": selected["id"], "action": "edit", "edit_value": memo}
                ],
            },
        )

    assert {
        path.relative_to(output).as_posix(): path.read_bytes()
        for path in output.rglob("*")
        if path.is_file()
    } == before
    assert (output / "concordato_case_model.json").read_bytes() == case_bytes
    assert {
        path.relative_to(inputs).as_posix(): path.read_bytes()
        for path in inputs.rglob("*")
        if path.is_file()
    } == sources
    reopened = service.dispatch(
        "vera_workspace_view", {"work_ref": binding["work_ref"]}
    )
    assert reopened["data"]["final_artifacts"]["final_ready"] is False
    assert any(row["item_type"] == "semantic_issue" for row in reopened["items"])
    assert service.review_directory(root_output, "concordato-plan-review") == output


def test_concordato_native_acceptance_does_not_resolve_confirmed_semantic_gap(
    concordato_workspace,
):
    service, binding, _, output, _ = concordato_workspace
    before = (output / "concordato_case_model.json").read_bytes()
    selected = next(
        row
        for row in json.loads((output / "review_payload.json").read_bytes())["items"]
        if row["item_type"] == "semantic_issue"
    )
    current = service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "item_id": selected["id"]},
    )

    service.dispatch(
        "vera_workspace_apply",
        {
            "work_ref": binding["work_ref"],
            "item_id": selected["id"],
            "revision": current["revision"],
            "human_reviewed": True,
            "idempotency_key": "fictional-concordato-issue",
            "decisions": [
                {
                    "item_id": selected["id"],
                    "action": "accept",
                    "reviewer_note": "Rilievo preso in carico; evidenza ancora mancante.",
                }
            ],
        },
    )

    assert (output / "concordato_case_model.json").read_bytes() == before
    final = json.loads((output / "final_artifacts.json").read_bytes())
    assert final["final_ready"] is False
    assert json.loads(before)["case_model"]["issues"][0]["status"] == "open"
    reopened = service.dispatch(
        "vera_workspace_view", {"work_ref": binding["work_ref"]}
    )
    assert reopened["data"]["final_artifacts"]["final_ready"] is False


def test_concordato_native_report_edit_keeps_confirmed_report_and_replays_revision(
    concordato_workspace,
):
    service, binding, _, output, _ = concordato_workspace
    original = (output / "concordato_semantic_review.md").read_bytes()
    current = service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "item_id": "semantic-review"},
    )
    revision = "Revisione proposta: attestazione ancora mancante."

    service.dispatch(
        "vera_workspace_apply",
        {
            "work_ref": binding["work_ref"],
            "item_id": "semantic-review",
            "revision": current["revision"],
            "human_reviewed": True,
            "idempotency_key": "fictional-concordato-report",
            "decisions": [
                {"item_id": "semantic-review", "action": "edit", "edit_value": revision}
            ],
        },
    )

    assert (output / "concordato_semantic_review.md").read_bytes() == original
    applied = json.loads((output / "applied_decisions.json").read_bytes())
    assert applied["effects"][0]["artifact_update"] == "revision_artifact_written"
    assert (output / applied["effects"][0]["revision_artifact"]).read_text() == revision
    reopened = service.dispatch(
        "vera_workspace_view", {"work_ref": binding["work_ref"]}
    )
    assert reopened["data"]["final_artifacts"]["final_ready"] is False


def test_concordato_native_saved_memo_edit_blocks_merged_apply_without_output_change(
    concordato_workspace,
):
    service, binding, _, output, _ = concordato_workspace
    current = service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "item_id": "codex-review-memo"},
    )
    service.dispatch(
        "vera_workspace_save",
        {
            "work_ref": binding["work_ref"],
            "item_id": "codex-review-memo",
            "revision": current["revision"],
            "human_reviewed": True,
            "idempotency_key": "fictional-concordato-save-memo",
            "decisions": [
                {
                    "item_id": "codex-review-memo",
                    "action": "edit",
                    "edit_value": "Proposta di memo.",
                }
            ],
        },
    )
    current = service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "item_id": "semantic-issue-1"},
    )
    before = {
        path.relative_to(output).as_posix(): path.read_bytes()
        for path in output.rglob("*")
        if path.is_file()
    }

    with pytest.raises(ValueError, match="non supera il replay del motore"):
        service.dispatch(
            "vera_workspace_apply",
            {
                "work_ref": binding["work_ref"],
                "item_id": "semantic-issue-1",
                "revision": current["revision"],
                "human_reviewed": True,
                "idempotency_key": "fictional-concordato-merged",
                "decisions": [{"item_id": "semantic-issue-1", "action": "accept"}],
            },
        )

    assert {
        path.relative_to(output).as_posix(): path.read_bytes()
        for path in output.rglob("*")
        if path.is_file()
    } == before


def test_concordato_native_changed_source_refuses_before_review_write(
    concordato_workspace,
):
    service, binding, _, output, inputs = concordato_workspace
    service.load_binding(binding)
    before = (output / "ui_decisions.json").read_bytes()
    source = next(inputs.rglob("case_material_a.xlsx"))
    source.write_bytes(source.read_bytes() + b"changed source bytes")

    with pytest.raises(sys.modules["client_ledger"].LedgerError, match="receipt"):
        service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})

    assert (output / "ui_decisions.json").read_bytes() == before


def test_report_native_selected_section_excludes_private_control_and_other_sections(
    report_workspace,
):
    service, binding, root_output, output, source = report_workspace
    review = json.loads((output / "review_payload.json").read_bytes())
    selected = next(
        item
        for item in review["items"]
        if item["item_type"] == "report_section"
        and item["data"]["status"] == "assigned"
    )

    context = service.dispatch(
        "vera_workspace_explain",
        {"work_ref": binding["work_ref"], "item_id": selected["id"]},
    )

    assert context["untrusted_evidence"]["id"] == selected["id"]
    serialized = json.dumps(context)
    assert "inspection_control.json" not in serialized
    assert "source_index.json" not in serialized
    assert "review_integrity.json" not in serialized
    assert str(source) not in serialized
    assert all(
        item["id"] not in serialized
        for item in review["items"]
        if item["id"] != selected["id"]
    )
    assert service.review_directory(root_output, "report-builder") == output


def test_report_native_apply_regenerates_narrative_and_preserves_pending_numeric_review(
    report_workspace,
):
    service, binding, _, output, source = report_workspace
    review = json.loads((output / "review_payload.json").read_bytes())
    selected = next(
        item
        for item in review["items"]
        if item["item_type"] == "report_section"
        and item["data"]["status"] == "assigned"
    )
    before = source.read_bytes()
    current = service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "item_id": selected["id"]},
    )
    replacement = "La tabella richiede una revisione professionale delle misure."

    result = service.dispatch(
        "vera_workspace_apply",
        {
            "work_ref": binding["work_ref"],
            "item_id": selected["id"],
            "revision": current["revision"],
            "human_reviewed": True,
            "idempotency_key": "fictional-report-edit",
            "decisions": [
                {"item_id": selected["id"], "action": "edit", "edit_value": replacement}
            ],
        },
    )

    assert result["saved"] is True
    assert result["revision"] != current["revision"]
    assert source.read_bytes() == before
    assert replacement in (output / "report_draft.md").read_text()
    assert replacement in "\n".join(
        paragraph.text
        for paragraph in reports.Document(output / "report.docx").paragraphs
    )
    final = json.loads((output / "final_artifacts.json").read_bytes())
    assert final["status"] != "final_ready"
    assert not (output / "numeric_evidence_ledger.json").exists()
    reopened = service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "item_id": selected["id"]},
    )
    saved = reopened["data"]["ui_decisions"]["decisions"]
    assert saved[0]["edit_value"] == replacement
    assert reopened["selection"]["data"]["codex_comment"] == replacement


def test_report_native_rejects_changed_imported_source_before_review(report_workspace):
    service, binding, _, output, source = report_workspace
    service.load_binding(binding)
    before = (output / "ui_decisions.json").read_bytes()
    source.write_text("line,amount\nA,999\nB,20\n")

    with pytest.raises(sys.modules["client_ledger"].LedgerError, match="receipt"):
        service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})

    assert (output / "ui_decisions.json").read_bytes() == before


@pytest.fixture
def reviewed_report_workspace(report_workspace):
    service, binding, root_output, output, source = report_workspace
    selected = next(
        item
        for item in json.loads((output / "review_payload.json").read_bytes())["items"]
        if item["item_type"] == "report_section"
        and item["data"]["status"] == "assigned"
    )
    view = service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "item_id": selected["id"]},
    )
    service.dispatch(
        "vera_workspace_apply",
        {
            "work_ref": binding["work_ref"],
            "item_id": selected["id"],
            "revision": view["revision"],
            "human_reviewed": True,
            "idempotency_key": "fictional-report-predecessor",
            "decisions": [
                {
                    "item_id": selected["id"],
                    "action": "edit",
                    "edit_value": "Prima narrativa revisionata.",
                }
            ],
        },
    )
    current = service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "item_id": selected["id"]},
    )
    return service, binding, root_output, output, source, selected, current


def test_report_native_successor_replays_retained_checkpoint_and_current_narrative(
    reviewed_report_workspace,
):
    service, binding, root_output, output, source, selected, current = (
        reviewed_report_workspace
    )
    checkpoint_path = service.ui_state_directory(root_output) / "report-checkpoint.json"
    previous = json.loads(checkpoint_path.read_bytes())["current_checkpoint"]
    source_bytes = source.read_bytes()
    replacement = "La revisione successiva mantiene i controlli numerici pendenti."

    result = service.dispatch(
        "vera_workspace_apply",
        {
            "work_ref": binding["work_ref"],
            "item_id": selected["id"],
            "revision": current["revision"],
            "human_reviewed": True,
            "idempotency_key": "fictional-report-successor",
            "decisions": [
                {"item_id": selected["id"], "action": "edit", "edit_value": replacement}
            ],
        },
    )

    assert result["saved"] is True
    stored = json.loads(checkpoint_path.read_bytes())
    assert stored["predecessor_checkpoint"] == previous
    assert stored["current_checkpoint"] != previous
    applied = json.loads((output / "applied_decisions.json").read_bytes())
    assert applied["predecessor_checkpoint"] == previous
    assert len(applied["review_history_paths"]) == 1
    assert source.read_bytes() == source_bytes
    reopened = service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "item_id": selected["id"]},
    )
    assert reopened["selection"]["data"]["codex_comment"] == replacement
    explanation = service.dispatch(
        "vera_workspace_explain",
        {
            "work_ref": binding["work_ref"],
            "item_id": selected["id"],
            "revision": reopened["revision"],
        },
    )
    assert explanation["untrusted_evidence"]["data"]["codex_comment"] == replacement
    assert reopened["data"]["final_artifacts"]["status"] != "final_ready"


def test_report_native_save_preserves_rendered_report_and_retains_new_authority(
    report_workspace,
):
    service, binding, root_output, output, _ = report_workspace
    selected = next(
        item
        for item in json.loads((output / "review_payload.json").read_bytes())["items"]
        if item["item_type"] == "report_section"
        and item["data"]["status"] == "assigned"
    )
    current = service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "item_id": selected["id"]},
    )
    before = (output / "report_draft.md").read_bytes()
    checkpoint_path = service.ui_state_directory(root_output) / "report-checkpoint.json"
    proposal = "Proposta conservata senza rigenerare il report."

    result = service.dispatch(
        "vera_workspace_save",
        {
            "work_ref": binding["work_ref"],
            "item_id": selected["id"],
            "revision": current["revision"],
            "human_reviewed": True,
            "idempotency_key": "fictional-report-save",
            "decisions": [
                {"item_id": selected["id"], "action": "edit", "edit_value": proposal}
            ],
        },
    )

    assert result["saved"] is True
    assert (output / "report_draft.md").read_bytes() == before
    stored = json.loads(checkpoint_path.read_bytes())
    assert stored["run_id"] == binding["run_id"]
    assert len(stored["current_checkpoint"]) == 64
    reopened = service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "item_id": selected["id"]},
    )
    assert reopened["data"]["ui_decisions"]["decisions"][0]["edit_value"] == proposal
    assert reopened["selection"]["data"]["codex_comment"] == ""


def test_report_native_successor_draft_preserves_complete_official_artifact_tree(
    reviewed_report_workspace,
):
    service, binding, _, output, _, selected, current = reviewed_report_workspace
    before = {
        path.relative_to(output).as_posix(): path.read_bytes()
        for path in output.rglob("*")
        if path.is_file()
    }
    fields = {"action": "edit", "edit_value": "Proposta della revisione successiva."}

    result = service.dispatch(
        "vera_workspace_draft_save",
        {
            "work_ref": binding["work_ref"],
            "item_id": selected["id"],
            "revision": current["revision"],
            "fields": fields,
        },
    )

    assert result["draft_saved"] is True
    assert {
        path.relative_to(output).as_posix(): path.read_bytes()
        for path in output.rglob("*")
        if path.is_file()
    } == before
    saved = service.dispatch(
        "vera_workspace_draft_read", {"work_ref": binding["work_ref"]}
    )["draft"]
    assert saved["fields"] == fields
    assert saved["revision"] == current["revision"]


@pytest.mark.parametrize(
    "guard,message",
    [
        ("missing", "Retained report checkpoint is required"),
        ("wrong", "Report Builder persisted review authorization failed"),
    ],
)
def test_report_native_successor_refuses_lost_or_changed_external_checkpoint(
    reviewed_report_workspace, guard, message
):
    service, binding, root_output, output, _, selected, current = (
        reviewed_report_workspace
    )
    before = {
        path.relative_to(output).as_posix(): path.read_bytes()
        for path in output.rglob("*")
        if path.is_file()
    }
    path = service.ui_state_directory(root_output) / "report-checkpoint.json"
    if guard == "missing":
        path.unlink()
    else:
        stored = json.loads(path.read_bytes())
        stored["current_checkpoint"] = "0" * 64
        path.write_text(json.dumps(stored))

    with pytest.raises(ValueError, match=message):
        service.dispatch(
            "vera_workspace_apply",
            {
                "work_ref": binding["work_ref"],
                "item_id": selected["id"],
                "revision": current["revision"],
                "human_reviewed": True,
                "idempotency_key": "fictional-report-rejected",
                "decisions": [
                    {
                        "item_id": selected["id"],
                        "action": "edit",
                        "edit_value": "Questa proposta deve essere rifiutata.",
                    }
                ],
            },
        )

    assert {
        path.relative_to(output).as_posix(): path.read_bytes()
        for path in output.rglob("*")
        if path.is_file()
    } == before


def test_file_preparation_native_explanation_uses_exact_verified_model_handoff(
    file_preparation_workspace,
):
    service, binding, root_output, output = file_preparation_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    item = next(
        row for row in view["items"] if row["item_type"] == "document_inventory"
    )
    context, records, _ = service.file_preparation_handoff(output)

    explained = service.dispatch(
        "vera_workspace_explain",
        {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "item_id": item["id"],
        },
    )

    projection = explained["untrusted_evidence"]
    assert projection["selection_mode"] == "maintained_model_handoff"
    assert projection["items"] == [
        row for row in records if row.get("document_ref") == item["id"]
    ]
    assert projection["content_policy"] == context["content_policy"]
    assert "draft-client-email" not in json.dumps(explained)
    assert str(output) not in json.dumps(explained)
    assert service.review_directory(root_output, "client-file-preparation") == output


def test_file_preparation_native_edit_apply_reseals_exact_draft_and_reopens(
    file_preparation_workspace,
):
    service, binding, root_output, output = file_preparation_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    item = next(
        row for row in view["items"] if row["item_type"] == "draft_client_email"
    )
    original = (output / "04_bozza_email_cliente.md").read_bytes()
    replacement = "Oggetto: Documenti aggiornati\n\nGentile cliente, confermi i documenti ancora mancanti."
    request = {
        "work_ref": binding["work_ref"],
        "revision": view["revision"],
        "item_id": item["id"],
        "human_reviewed": True,
        "idempotency_key": "fictional-file-preparation-save",
        "decisions": [
            {
                "item_id": item["id"],
                "action": "edit",
                "edit_value": replacement,
                "reviewer_note": "Fictional exact draft explicitly reviewed",
            }
        ],
    }
    saved = service.dispatch("vera_workspace_save", request)
    assert (output / "04_bozza_email_cliente.md").read_bytes() == original
    assert saved["revision"] != view["revision"]
    reopened = service.dispatch(
        "vera_workspace_view", {"work_ref": binding["work_ref"]}
    )

    applied = service.dispatch(
        "vera_workspace_apply",
        {
            **request,
            "revision": reopened["revision"],
            "idempotency_key": "fictional-file-preparation-apply",
        },
    )

    assert (output / "04_bozza_email_cliente.md").read_text() == replacement
    fresh = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    assert fresh["revision"] == applied["revision"]
    assert fresh["data"]["ui_decisions"]["decisions"][0]["edit_value"] == replacement
    manifest = json.loads((output / "final_artifacts.json").read_bytes())
    email = next(
        row for row in manifest["outputs"] if row["path"] == "04_bozza_email_cliente.md"
    )
    assert email["required_text"] == [replacement]
    assert manifest["review_application"]["target_update_count"] >= 1
    paths = service.dispatch(
        "vera_workspace_outputs", {"work_ref": binding["work_ref"]}
    )
    assert any(
        row["path"] == str(output / "04_bozza_email_cliente.md")
        for row in paths["outputs"]
    )
    assert (root_output.parent / ".native-workspace").is_dir()


def test_file_preparation_native_rejects_changed_model_page_before_discussion(
    file_preparation_workspace,
):
    service, binding, _, output = file_preparation_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    context = json.loads((output / "model_handoff.json").read_bytes())
    page = output / context["pagination"]["pages"][0]["path"]
    page.write_bytes(page.read_bytes() + b"\n")

    with pytest.raises(
        ValueError, match="Model handoff page identity or content changed"
    ):
        service.dispatch(
            "vera_workspace_explain",
            {
                "work_ref": binding["work_ref"],
                "revision": view["revision"],
                "item_id": view["items"][0]["id"],
            },
        )


def test_file_preparation_native_request_apply_updates_only_reviewed_email_handoff(
    file_preparation_workspace,
):
    service, binding, _, output = file_preparation_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    item = next(
        row for row in view["items"] if row["item_type"] == "missing_document_request"
    )
    reviewed_request = (
        "Fictional specific document request reviewed by the professional"
    )

    service.dispatch(
        "vera_workspace_apply",
        {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "item_id": item["id"],
            "human_reviewed": True,
            "idempotency_key": "fictional-specific-request",
            "decisions": [
                {
                    "item_id": item["id"],
                    "action": "request_more_documents",
                    "requested_documents": [reviewed_request],
                }
            ],
        },
    )

    context, records, _ = service.file_preparation_handoff(output)
    email_requests = [row for row in records if row["kind"] == "email_request"]
    assert len(email_requests) == 1
    assert email_requests[0]["review_item_ref"] == item["id"]
    assert email_requests[0]["request_text"] == reviewed_request
    assert email_requests[0]["client_reference"] == "CLIENT-001"
    assert context["phase_access"]["email_drafting"] == ["email_request"]
    assert "email_request" not in context["phase_access"]["file_preparation_review"]
    fresh = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    assert fresh["data"]["final_artifacts"]["status"] != "final_ready"


def test_new_client_native_review_reopens_without_mutating_domain_or_activating_relationship(
    new_client_workspace,
):
    service, binding, output = new_client_workspace
    before = (output / "case_facts_validated.json").read_bytes()
    original = json.loads((output / "final_artifacts.json").read_bytes())
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    item = view["items"][0]
    request = {
        "work_ref": binding["work_ref"],
        "revision": view["revision"],
        "item_id": item["id"],
        "human_reviewed": True,
        "idempotency_key": "fictional-new-client-save",
        "decisions": [
            {
                "item_id": item["id"],
                "action": "accept",
                "reviewer_note": "Fictional proposal explicitly reviewed; remaining gaps retained",
            }
        ],
    }
    saved = service.dispatch("vera_workspace_save", request)
    reopened = service.dispatch(
        "vera_workspace_view", {"work_ref": binding["work_ref"]}
    )
    applied = service.dispatch(
        "vera_workspace_apply",
        {
            **request,
            "revision": reopened["revision"],
            "idempotency_key": "fictional-new-client-apply",
        },
    )

    assert saved["revision"] != view["revision"]
    assert (
        service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})[
            "revision"
        ]
        == applied["revision"]
    )
    final = json.loads((output / "final_artifacts.json").read_bytes())
    assert (
        final["export_gate"]["domain_blockers"]
        == original["export_gate"]["domain_blockers"]
    )
    assert (
        final["export_gate"]["artifact_blockers"]
        == original["export_gate"]["artifact_blockers"]
    )
    assert final["relationship_activation_performed"] is False
    assert final["signature_performed"] is False
    assert final["client_communication_sent"] is False
    assert (output / "case_facts_validated.json").read_bytes() == before
    assert list(output.glob("applied_decisions.history.*.json"))
    assert (
        json.loads((output / "ui_decisions.json").read_bytes())["decision_revision"]
        == 2
    )


def approve_native_archive(service, binding):
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    for index, item in enumerate(view["items"]):
        service.dispatch(
            "vera_workspace_save",
            {
                "work_ref": binding["work_ref"],
                "revision": view["revision"],
                "item_id": item["id"],
                "human_reviewed": True,
                "idempotency_key": "fictional-archive-save-" + str(index),
                "decisions": [
                    {
                        "item_id": item["id"],
                        "action": "accept",
                        "reviewer_note": "Fictional reviewed archive destination",
                    }
                ],
            },
        )
        view = service.dispatch(
            "vera_workspace_view", {"work_ref": binding["work_ref"]}
        )
    item = view["items"][0]
    service.dispatch(
        "vera_workspace_apply",
        {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "item_id": item["id"],
            "human_reviewed": True,
            "idempotency_key": "fictional-archive-compile",
            "decisions": [
                {
                    "item_id": item["id"],
                    "action": "accept",
                    "reviewer_note": "Fictional reviewed archive destination",
                }
            ],
        },
    )
    return service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "view": "ARCHIVE_EXECUTION"},
    )


def test_new_client_native_apply_rejects_expired_package_without_changing_dossier(
    new_client_workspace,
    monkeypatch,
    tmp_path,
):
    service, binding, output = new_client_workspace
    before = (output / "final_artifacts.json").read_bytes()
    clock = tmp_path / "expired-clock.cjs"
    clock.write_text(
        "const RealDate = Date; global.Date = class extends RealDate {"
        "constructor(...args) { super(...(args.length ? args : ['2099-01-01T12:00:00Z'])); }"
        "static now() { return new RealDate('2099-01-01T12:00:00Z').getTime(); } };"
    )
    monkeypatch.setenv("NODE_OPTIONS", "--require " + str(clock))
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    item = view["items"][0]

    with pytest.raises(ValueError, match="temporal validity expired"):
        service.dispatch(
            "vera_workspace_apply",
            {
                "work_ref": binding["work_ref"],
                "revision": view["revision"],
                "item_id": item["id"],
                "human_reviewed": True,
                "idempotency_key": "fictional-expired-new-client",
                "decisions": [{"item_id": item["id"], "action": "accept"}],
            },
        )

    assert (output / "final_artifacts.json").read_bytes() == before
    assert not list(output.glob("applied_decisions.history.*.json"))


@pytest.mark.parametrize("guard", ["stale", "viewer", "not_reviewed", "selected_item"])
def test_archive_native_execution_rejects_invalid_authority_without_moving_files(
    archive_workspace,
    monkeypatch,
    guard,
):
    service, binding, output, original_files = archive_workspace
    view = approve_native_archive(service, binding)
    request = {
        "work_ref": binding["work_ref"],
        "revision": view["revision"],
        "human_reviewed": True,
        "execution_approved": True,
        "operation": "apply",
        "idempotency_key": "fictional-invalid-archive",
    }
    expected_error = ValueError
    if guard == "stale":
        request["revision"] = "stale"
    elif guard == "viewer":
        monkeypatch.setenv("VERA_WORKSPACE_ROLES", "VIEWER")
        expected_error = PermissionError
    elif guard == "not_reviewed":
        request["human_reviewed"] = False
    else:
        request["item_id"] = view["items"][0]["id"]

    with pytest.raises(expected_error):
        service.dispatch("vera_workspace_archive_execute", request)

    assert not (output / "apply_journal.json").exists()
    assert all(
        (Path(binding["client_root"]) / name).read_bytes() == content
        for name, content in original_files.items()
    )


def test_archive_native_mcp_execute_rejects_forged_ticket_before_domain_mutation(
    archive_workspace,
):
    service, binding, output, original_files = archive_workspace
    view = approve_native_archive(service, binding)
    request = {
        "jsonrpc": "2.0",
        "id": 1,
        "method": "tools/call",
        "params": {
            "name": "vera_workspace_archive_execute",
            "arguments": {
                "work_ref": binding["work_ref"],
                "revision": view["revision"],
                "human_reviewed": True,
                "execution_approved": True,
                "operation": "apply",
                "idempotency_key": "fictional-forged-archive",
                "review_ticket": "forged.invalid",
            },
        },
    }

    completed = subprocess.run(
        [NODE, str(SERVER)],
        input=json.dumps(request) + "\n",
        text=True,
        capture_output=True,
        check=True,
    )

    assert json.loads(completed.stdout)["result"]["isError"] is True
    assert "Invalid review ticket" in completed.stdout
    assert not (output / "apply_journal.json").exists()
    assert all(
        (Path(binding["client_root"]) / name).read_bytes() == content
        for name, content in original_files.items()
    )


def test_archive_native_compile_does_not_move_files_and_separate_execution_can_rollback(
    archive_workspace,
):
    service, binding, output, original_files = archive_workspace
    client = Path(binding["client_root"])
    view = approve_native_archive(service, binding)
    assert all(
        (client / name).read_bytes() == content
        for name, content in original_files.items()
    )
    projection = view["data"]["archive_execution"]
    assert projection["approved_plan"]["change_count"] == 2
    assert projection["execution_requires_separate_explicit_approval"] is True
    assert "source_sha256" not in json.dumps(projection)
    assert "drive_root_folder_id" not in json.dumps(projection)
    request = {
        "work_ref": binding["work_ref"],
        "revision": view["revision"],
        "human_reviewed": True,
        "idempotency_key": "fictional-archive-execute",
        "operation": "apply",
    }
    with pytest.raises(ValueError, match="Separate explicit approval"):
        service.dispatch("vera_workspace_archive_execute", request)
    assert all(
        (client / name).read_bytes() == content
        for name, content in original_files.items()
    )

    applied = service.dispatch(
        "vera_workspace_archive_execute", {**request, "execution_approved": True}
    )

    assert applied["status"] == "applied"
    assert not (client / "Comunicazione_36bis_2024.pdf").exists()
    assert (
        client / "AdE/2024/36-bis/2024-06-14_comunicazione-36-bis_Example-Srl.pdf"
    ).read_bytes() == original_files["Comunicazione_36bis_2024.pdf"]
    assert (
        service.dispatch(
            "vera_workspace_archive_execute", {**request, "execution_approved": True}
        )
        == applied
    )
    reopened = service.dispatch(
        "vera_workspace_view",
        {"work_ref": binding["work_ref"], "view": "ARCHIVE_EXECUTION"},
    )
    assert reopened["data"]["archive_execution"]["journal"]["status"] == "applied"
    rolled_back = service.dispatch(
        "vera_workspace_archive_execute",
        {
            **request,
            "execution_approved": True,
            "revision": reopened["revision"],
            "idempotency_key": "fictional-archive-rollback",
            "operation": "rollback",
        },
    )
    assert rolled_back["status"] == "rolled_back"
    assert all(
        (client / name).read_bytes() == content
        for name, content in original_files.items()
    )
    assert (
        json.loads((output / "apply_journal.json").read_bytes())["status"]
        == "rolled_back"
    )


def test_archive_native_execution_rejects_changed_source_without_moving_others(
    archive_workspace,
):
    service, binding, _, original_files = archive_workspace
    view = approve_native_archive(service, binding)
    client = Path(binding["client_root"])
    (client / "Comunicazione_36bis_2024.pdf").write_bytes(
        b"fictional changed source after approval"
    )

    with pytest.raises(ValueError, match="changed|Changed"):
        service.dispatch(
            "vera_workspace_archive_execute",
            {
                "work_ref": binding["work_ref"],
                "revision": view["revision"],
                "human_reviewed": True,
                "execution_approved": True,
                "operation": "apply",
                "idempotency_key": "fictional-changed-source",
            },
        )

    assert (
        client / "inbox/Comunicazione_36bis_copia.pdf"
    ).read_bytes() == original_files["inbox/Comunicazione_36bis_copia.pdf"]
    assert (client / "notes.txt").read_bytes() == original_files["notes.txt"]


def test_sampling_nested_review_explanation_uses_only_selected_model_projection(
    sampling_workspace,
):
    service, binding, root_output, output = sampling_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    item = next(row for row in view["items"] if row["item_type"] == "sampled_entry")
    context = json.loads((output / "model_review_context.json").read_bytes())

    result = service.dispatch(
        "vera_workspace_explain",
        {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "item_id": item["id"],
        },
    )

    assert result["untrusted_evidence"]["item"] == next(
        row for row in context["review"]["items"] if row["id"] == item["id"]
    )
    assert "items" not in result["untrusted_evidence"]
    assert "sample-journal.xlsx" not in json.dumps(result)
    paths = service.dispatch(
        "vera_workspace_outputs", {"work_ref": binding["work_ref"]}
    )
    assert any(
        row["path"] == str(output / "journal_sample.xlsx") for row in paths["outputs"]
    )
    assert service.review_directory(root_output, "journal-sampling") == output


def test_sampling_save_apply_reopens_replayed_successor_and_keeps_draft_outside_outputs(
    sampling_workspace,
):
    service, binding, root_output, output = sampling_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    item = next(row for row in view["items"] if row["item_type"] == "sampled_entry")
    request = {
        "work_ref": binding["work_ref"],
        "revision": view["revision"],
        "item_id": item["id"],
        "human_reviewed": True,
        "idempotency_key": "fictional-sampling-save",
        "decisions": [
            {
                "item_id": item["id"],
                "action": "accept",
                "reviewer_note": "Fictional selected-entry review",
            }
        ],
    }

    saved = service.dispatch("vera_workspace_save", request)
    reopened = service.dispatch(
        "vera_workspace_view", {"work_ref": binding["work_ref"]}
    )
    applied = service.dispatch(
        "vera_workspace_apply",
        {
            **request,
            "revision": reopened["revision"],
            "idempotency_key": "fictional-sampling-apply",
        },
    )

    assert saved["revision"] != view["revision"]
    assert applied["revision"] != reopened["revision"]
    assert (
        service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})[
            "revision"
        ]
        == applied["revision"]
    )
    applied_record = json.loads((output / "applied_decisions.json").read_bytes())
    assert applied_record["decision_count"] == 1
    assert applied_record["effects"][0]["item_id"] == item["id"]
    assert applied_record["application_status"] == "partial_review_applied"
    assert list((output / "assurance_history").iterdir())
    assert (root_output.parent / ".native-workspace").is_dir()
    assert not (root_output / ".native-workspace").exists()
    with pytest.raises(ValueError, match="Stale"):
        service.dispatch(
            "vera_workspace_save",
            {**request, "idempotency_key": "fictional-stale-sampling"},
        )


def test_vouching_nested_selected_context_uses_engine_minimization(vouching_workspace):
    service, binding, _, output = vouching_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    item = next(row for row in view["items"] if row["item_type"] == "supported_entry")

    result = service.dispatch(
        "vera_workspace_explain",
        {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "item_id": item["id"],
        },
    )

    evidence = result["untrusted_evidence"]
    assert evidence["case_count"] == 1
    assert evidence["include_exact_identifiers"] is False
    assert evidence["cases"][0]["item_type"] == "supported_entry"
    assert (
        "prepared entry and support artifact IDs" in evidence["minimization"]["omitted"]
    )
    assert str(output) not in json.dumps(result)
    assert "M-1001" not in json.dumps(result)


def test_vouching_native_note_edit_regenerates_workbook_through_engine(
    vouching_workspace,
):
    service, binding, _, output = vouching_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    item = next(row for row in view["items"] if row["item_type"] == "supported_entry")
    note = "Fictional note from native qualification"

    result = service.dispatch(
        "vera_workspace_apply",
        {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "item_id": item["id"],
            "human_reviewed": True,
            "idempotency_key": "fictional-vouching-edit",
            "decisions": [
                {"item_id": item["id"], "action": "edit", "edit_value": note}
            ],
        },
    )

    assert result["status"] == "applied"
    workbook = load_workbook(output / "check_results.xlsx", read_only=True)
    assert any(note in row for sheet in workbook for row in sheet.values)
    assert (
        service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})[
            "revision"
        ]
        == result["revision"]
    )


def test_open_items_exact_context_uses_engine_projection(tmp_path, monkeypatch):
    output, _ = open_items._portable_audit_transaction_case(tmp_path)
    binding, _ = archived_review_binding(output, "open-item-reconciliation")
    configure(monkeypatch, tmp_path, [binding])
    service = workspace_module()
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})

    result = service.dispatch(
        "vera_workspace_explain",
        {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "item_id": view["items"][0]["id"],
        },
    )

    assert result["untrusted_evidence"]["case_count"] == 1
    assert result["untrusted_evidence"]["include_exact_identifiers"] is False
    assert str(output) not in json.dumps(result)


@pytest.mark.parametrize("nested", [False, True])
def test_open_items_assured_apply_requires_separate_checkpoint_and_reopens_partial_review(
    tmp_path,
    monkeypatch,
    nested,
):
    root_output = open_items._running_audit_output(tmp_path)
    output = root_output / "reconciliation" if nested else root_output
    context_path = root_output.parent / "context.json"
    context = open_items.load_assurance().load_client_engagement_context_file(
        context_path,
        expected_workflow_id="open-item-reconciliation",
    )
    workflow = open_items.load_reconciliation_workflow()
    workflow.build_reconciliation_artifacts(
        output_dir=output,
        run_id=context["run_id"],
        client_engagement=context,
        open_items=[
            {
                "record_id": "fictional-open-1",
                "document_key": "FICT-1|2026",
                "document_no": "FICT-1",
                "document_date": "2026-01-01",
                "amount": "100.00",
                "currency": "EUR",
            }
        ],
        evidence_rows=[],
        assumptions={
            "scope_year": "2026",
            "amount_tolerance": "0",
            "assurance_run_date": "2026-07-25",
        },
        require_completed_review=False,
        fail_on_check_errors=False,
        language="it",
    )
    # This synthetic test retains the original engine checkpoint in test memory,
    # independently of the candidate directory submitted for the subsequent apply.
    checkpoint = json.loads((output / "assurance_receipts.json").read_bytes())[
        "content_sha256"
    ]
    binding, _ = archived_review_binding(output, "open-item-reconciliation")
    configure(monkeypatch, tmp_path, [binding])
    service = workspace_module()
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    item = next(row for row in view["items"] if "accept" in row["allowed_actions"])
    request = {
        "work_ref": binding["work_ref"],
        "revision": view["revision"],
        "item_id": item["id"],
        "human_reviewed": True,
        "idempotency_key": "fictional-open-items-apply",
        "decisions": [
            {
                "item_id": item["id"],
                "action": "accept",
                "reviewer_note": "Fictional review; source qualification remains pending",
            }
        ],
    }
    before = open_items._audit_tree_image(output)

    with pytest.raises(ValueError, match="predecessor checkpoint is required"):
        service.dispatch("vera_workspace_apply", request)
    assert open_items._audit_tree_image(output) == before
    with pytest.raises(ValueError, match="predecessor checkpoint does not match"):
        service.dispatch(
            "vera_workspace_apply",
            {**request, "expected_predecessor_checkpoint": "0" * 64},
        )
    assert open_items._audit_tree_image(output) == before
    result = service.dispatch(
        "vera_workspace_apply",
        {**request, "expected_predecessor_checkpoint": checkpoint},
    )

    assert result["status"] == "applied"
    assert result["revision"] != view["revision"]
    assert (
        service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})[
            "revision"
        ]
        == result["revision"]
    )
    assert (output / "riconciliazione_audit.xlsx").is_file()
    assert (output / "relazione_riconciliazione_audit.docx").is_file()
    gates = json.loads((output / "assurance_gates.json").read_bytes())
    assert gates["report_ready"] is False
    applied_record = json.loads((output / "applied_decisions.json").read_bytes())
    assert (
        applied_record["professional_review"]["successor_assurance_replayed"] is False
    )


def test_nested_review_rejects_ambiguous_and_linked_locations(tmp_path):
    service = workspace_module()
    sample = tmp_path / "sample"
    sample.mkdir()
    (tmp_path / "review_payload.json").write_text("{}")
    (sample / "review_payload.json").write_text("{}")
    with pytest.raises(ValueError, match="Ambiguous"):
        service.review_directory(tmp_path, "journal-sampling")
    (sample / "review_payload.json").unlink()
    sample.rmdir()
    sample.symlink_to(tmp_path, target_is_directory=True)
    with pytest.raises(ValueError, match="symbolic links"):
        service.review_directory(tmp_path, "journal-sampling")


def treasury_binding(base: Path, ref: str = "treasury-fictional") -> tuple[dict, Path]:
    run = archived_case(base)
    context = run["context"]
    manifest = next(
        Path(row["path"])
        for row in context["input_bindings"]
        if row["path"].endswith(".json")
    )
    prepared = subprocess.run(
        [
            sys.executable,
            str(SCRIPTS / "run_treasury.py"),
            "prepare",
            "--client-engagement",
            str(run["context_path"]),
            "--manifest",
            str(manifest),
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert prepared.returncode == 0, prepared.stderr
    return (
        {
            "work_ref": ref,
            "client_root": str(run["client"]),
            "client_id": context["client_id"],
            "engagement_id": context["engagement_id"],
            "run_id": context["run_id"],
            "workflow_id": "treasury-forecast",
        },
        Path(context["output_dir"]),
    )


@pytest.fixture
def treasury_workspace(tmp_path, monkeypatch):
    binding, output = treasury_binding(tmp_path)
    env = configure(monkeypatch, tmp_path, [binding])
    return workspace_module(), binding, output, env


def test_catalogue_replays_only_authorised_archive_runs(treasury_workspace):
    service, binding, _, _ = treasury_workspace
    result = service.dispatch("vera_workspace_open", {})
    assert [row["work_ref"] for row in result["works"]] == [binding["work_ref"]]
    assert "client_root" not in result["works"][0]


def test_exact_explanation_never_includes_other_events(treasury_workspace):
    service, binding, _, _ = treasury_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    selected = view["items"][0]
    result = service.dispatch(
        "vera_workspace_explain",
        {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "item_id": selected["event_id"],
        },
    )
    assert result["untrusted_evidence"] == selected
    assert "events" not in result


def test_treasury_save_recalculates_reopens_and_preserves_immutable_predecessor(
    treasury_workspace,
):
    service, binding, output, _ = treasury_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    selected = view["items"][0]
    request = {
        "work_ref": binding["work_ref"],
        "revision": view["revision"],
        "human_reviewed": True,
        "idempotency_key": "fictional-date-1",
        "item_id": selected["event_id"],
        "decisions": {
            selected["event_id"]: {
                "expected_date": "2026-10-01",
                "basis": "Fictional date explicitly reviewed for mechanism qualification",
            }
        },
    }

    result = service.dispatch("vera_workspace_save", request)

    reopened = service.dispatch(
        "vera_workspace_view", {"work_ref": binding["work_ref"]}
    )
    assert reopened["revision"] == result["revision"]
    assert reopened["revision"] != view["revision"]
    assert (output / "versions" / view["revision"] / "forecast.json").is_file()
    record = json.loads(
        (output / "versions" / reopened["revision"] / "forecast.json").read_bytes()
    )
    assert record["decisions"][selected["event_id"]]["expected_date"] == "2026-10-01"
    workbook = load_workbook(
        output / "versions" / reopened["revision"] / "tesoreria.xlsx",
        read_only=True,
        data_only=True,
    )
    assert any(
        "2026-10-01" in [str(value) for value in row]
        for sheet in workbook
        for row in sheet.values
    )
    workbook.close()


def test_duplicate_submission_replays_receipt_without_second_version(
    treasury_workspace,
):
    service, binding, output, _ = treasury_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    selected = view["items"][0]
    request = {
        "work_ref": binding["work_ref"],
        "revision": view["revision"],
        "human_reviewed": True,
        "idempotency_key": "duplicate-1",
        "item_id": selected["event_id"],
        "decisions": {
            selected["event_id"]: {
                "expected_date": "2026-10-01",
                "basis": "Fictional reviewed date",
            }
        },
    }
    saved = service.dispatch("vera_workspace_save", request)
    history = (output / "treasury_session.json").read_bytes()

    repeated = service.dispatch("vera_workspace_save", request)

    assert repeated == saved
    assert (output / "treasury_session.json").read_bytes() == history


def test_stale_explanation_is_rejected_after_review(treasury_workspace):
    service, binding, _, _ = treasury_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    selected = view["items"][0]
    service.dispatch(
        "vera_workspace_save",
        {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "human_reviewed": True,
            "idempotency_key": "stale-save",
            "item_id": selected["event_id"],
            "decisions": {
                selected["event_id"]: {
                    "expected_date": "2026-10-01",
                    "basis": "Fictional reviewed date",
                }
            },
        },
    )

    with pytest.raises(ValueError, match="Stale state"):
        service.dispatch(
            "vera_workspace_explain",
            {
                "work_ref": binding["work_ref"],
                "revision": view["revision"],
                "item_id": selected["event_id"],
            },
        )


@pytest.mark.parametrize(
    "field,value",
    [
        ("VERA_WORKSPACE_ACTOR_ID", "other-reviewer"),
        ("VERA_WORKSPACE_TENANT_ID", "other-studio"),
    ],
)
def test_configured_actor_and_tenant_cannot_be_substituted(
    treasury_workspace, monkeypatch, field, value
):
    service, _, _, _ = treasury_workspace
    monkeypatch.setenv(field, value)
    with pytest.raises(PermissionError, match="configured host actor"):
        service.dispatch("vera_workspace_open", {})


def test_viewer_cannot_persist_a_draft(treasury_workspace, monkeypatch):
    service, binding, _, _ = treasury_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    monkeypatch.setenv("VERA_WORKSPACE_ROLES", "VIEWER")
    with pytest.raises(PermissionError, match="Reviewer authority"):
        service.dispatch(
            "vera_workspace_draft_save",
            {
                "work_ref": binding["work_ref"],
                "revision": view["revision"],
                "fields": {"basis": "unsaved"},
            },
        )


def test_draft_recovery_does_not_change_forecast(treasury_workspace):
    service, binding, output, _ = treasury_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    before = (output / "treasury_session.json").read_bytes()
    fields = {"expected_date": "2026-10-01", "basis": "Unsaved fictional proposal"}

    service.dispatch(
        "vera_workspace_draft_save",
        {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "fields": fields,
        },
    )

    assert (
        service.dispatch(
            "vera_workspace_draft_read", {"work_ref": binding["work_ref"]}
        )["draft"]["fields"]
        == fields
    )
    assert (output / "treasury_session.json").read_bytes() == before


def test_other_run_reference_is_rejected(treasury_workspace):
    service, _, _, _ = treasury_workspace
    with pytest.raises(PermissionError, match="not authorised"):
        service.dispatch("vera_workspace_view", {"work_ref": "another-client-run"})


def test_draft_does_not_expand_the_official_artifact_tree(treasury_workspace):
    from tests.model_data_helpers import write_no_model_report

    service, binding, output, _ = treasury_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    write_no_model_report(output, "treasury-forecast", binding["run_id"])
    declarations = [
        {
            "artifact_id": f"treasury_{index}",
            "path": path.relative_to(output).as_posix(),
            "purpose": "Fictional mechanism qualification",
            "audience": "internal",
            "media_type": "application/octet-stream",
        }
        for index, path in enumerate(
            sorted(path for path in output.rglob("*") if path.is_file())
        )
    ]
    service.dispatch(
        "vera_workspace_draft_save",
        {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "fields": {"basis": "Unapplied fictional draft"},
        },
    )
    from client_ledger import finalize_run, validate_run_artifacts

    finalized = finalize_run(
        Path(binding["client_root"]),
        binding["engagement_id"],
        binding["run_id"],
        declarations,
    )

    assert finalized["run"]["status"] == "ready_for_review"
    assert validate_run_artifacts(
        Path(binding["client_root"]), binding["engagement_id"], binding["run_id"]
    )["artifacts"]
    assert (
        service.dispatch(
            "vera_workspace_draft_read", {"work_ref": binding["work_ref"]}
        )["draft"]["fields"]["basis"]
        == "Unapplied fictional draft"
    )


def test_treasury_finalized_run_reopens_read_only(treasury_workspace, monkeypatch):
    # The preceding finalization test covers sealing; finalized inspection must
    # use the same public ledger and read-only assurance contract.
    from tests.model_data_helpers import write_no_model_report

    service, binding, output, _ = treasury_workspace
    write_no_model_report(output, "treasury-forecast", binding["run_id"])
    monkeypatch.syspath_prepend(str(ROOT / "plugins/studio-archive/scripts"))
    from client_ledger import finalize_run

    finalize_run(
        Path(binding["client_root"]),
        binding["engagement_id"],
        binding["run_id"],
        [
            {
                "artifact_id": f"artifact_{index}",
                "path": path.relative_to(output).as_posix(),
                "purpose": "Fictional read-only qualification",
                "audience": "internal",
                "media_type": "application/octet-stream",
            }
            for index, path in enumerate(
                sorted(p for p in output.rglob("*") if p.is_file())
            )
        ],
    )
    reopened = service.dispatch(
        "vera_workspace_view", {"work_ref": binding["work_ref"]}
    )
    assert reopened["run_status"] == "ready_for_review"
    with pytest.raises(ValueError, match="Resume the Studio Archive run"):
        service.dispatch(
            "vera_workspace_draft_save",
            {
                "work_ref": binding["work_ref"],
                "revision": reopened["revision"],
                "fields": {},
            },
        )


def test_treasury_decision_cannot_target_another_opened_event(treasury_workspace):
    service, binding, output, _ = treasury_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    before = (output / "treasury_session.json").read_bytes()
    with pytest.raises(ValueError, match="selected event"):
        service.dispatch(
            "vera_workspace_save",
            {
                "work_ref": binding["work_ref"],
                "revision": view["revision"],
                "item_id": view["items"][0]["event_id"],
                "human_reviewed": True,
                "idempotency_key": "wrong-event",
                "decisions": {
                    view["items"][1]["event_id"]: {
                        "expected_date": "2026-10-01",
                        "basis": "Fictional unrelated event",
                    }
                },
            },
        )
    assert (output / "treasury_session.json").read_bytes() == before


@pytest.fixture
def bank_workspace(tmp_path, monkeypatch):
    output, _, review, _, _ = _prepare_sealed_mcp_review_run(
        tmp_path, language="it", portable=True
    )
    context = json.loads((output.parent / "context.json").read_bytes())
    binding = {
        "work_ref": "bank-fictional",
        "client_root": str(tmp_path / "Managed Customer"),
        "client_id": context["client_id"],
        "engagement_id": context["engagement_id"],
        "run_id": context["run_id"],
        "workflow_id": "journal-bank-reconciliation",
    }
    configure(monkeypatch, tmp_path, [binding])
    service = workspace_module()
    return service, binding, output, review


def test_bank_explanation_uses_existing_selected_case_projection(bank_workspace):
    service, binding, output, _ = bank_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    item = view["items"][0]

    result = service.dispatch(
        "vera_workspace_explain",
        {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "item_id": item["id"],
        },
    )

    assert result["untrusted_evidence"]["case_count"] == 1
    assert result["untrusted_evidence"]["include_exact_identifiers"] is False
    assert str(output) not in json.dumps(result)
    assert "source_path" not in json.dumps(result["untrusted_evidence"]["cases"])


def test_bank_edit_applies_to_native_workbook_and_saved_state(bank_workspace):
    service, binding, output, review = bank_workspace
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    selected = next(
        row for row in review["items"] if row["item_type"] == "matched_pair"
    )
    note = "Reviewed fictional bank movement from native workspace"

    result = service.dispatch(
        "vera_workspace_apply",
        {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "human_reviewed": True,
            "idempotency_key": "bank-apply",
            "item_id": selected["id"],
            "decisions": [
                {
                    "item_id": selected["id"],
                    "action": "edit",
                    "edit_value": note,
                    "reviewer_note": note,
                }
            ],
        },
    )

    reopened = service.dispatch(
        "vera_workspace_view", {"work_ref": binding["work_ref"]}
    )
    assert result["status"] == "applied"
    assert reopened["data"]["ui_decisions"]["decisions"][0]["edit_value"] == note
    workbook = load_workbook(
        output / "journal_bank_reconciliation.xlsx", read_only=True, data_only=True
    )
    assert any(note in row for sheet in workbook for row in sheet.values)
    workbook.close()


def test_bank_edit_preserves_another_saved_row_decision(bank_workspace):
    service, binding, _, review = bank_workspace
    first, second = review["items"][:2]
    view = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})
    service.dispatch(
        "vera_workspace_save",
        {
            "work_ref": binding["work_ref"],
            "revision": view["revision"],
            "item_id": first["id"],
            "human_reviewed": True,
            "idempotency_key": "row-one",
            "decisions": [
                {
                    "item_id": first["id"],
                    "action": "accept",
                    "reviewer_note": "First fictional review",
                }
            ],
        },
    )
    current = service.dispatch("vera_workspace_view", {"work_ref": binding["work_ref"]})

    service.dispatch(
        "vera_workspace_save",
        {
            "work_ref": binding["work_ref"],
            "revision": current["revision"],
            "item_id": second["id"],
            "human_reviewed": True,
            "idempotency_key": "row-two",
            "decisions": [
                {
                    "item_id": second["id"],
                    "action": "mark_unclear",
                    "reviewer_note": "Second fictional review",
                }
            ],
        },
    )

    reopened = service.dispatch(
        "vera_workspace_view", {"work_ref": binding["work_ref"]}
    )
    decisions = reopened["data"]["ui_decisions"]["decisions"]
    assert len(decisions) == 2
    assert decisions[0]["item_id"] == first["id"]
    assert decisions[0]["action"] == "accept"
    assert decisions[0]["reviewer_note"] == "First fictional review"
    assert decisions[1]["item_id"] == second["id"]
    assert decisions[1]["action"] == "mark_unclear"
    assert decisions[1]["reviewer_note"] == "Second fictional review"


def test_signed_ticket_rejects_forgery_before_domain_save(treasury_workspace):
    _, binding, output, env = treasury_workspace
    before = (output / "treasury_session.json").read_bytes()
    request = {
        "jsonrpc": "2.0",
        "id": 1,
        "method": "tools/call",
        "params": {
            "name": "vera_workspace_save",
            "arguments": {
                "work_ref": binding["work_ref"],
                "revision": "forged",
                "human_reviewed": True,
                "idempotency_key": "forged",
                "review_ticket": "forged.invalid",
                "decisions": {},
            },
        },
    }

    result = subprocess.run(
        [NODE, str(SERVER)],
        input=json.dumps(request) + "\n",
        env={**os.environ, **env},
        text=True,
        capture_output=True,
        check=True,
    )

    assert json.loads(result.stdout)["result"]["isError"] is True
    assert "Invalid review ticket" in result.stdout
    assert (output / "treasury_session.json").read_bytes() == before
