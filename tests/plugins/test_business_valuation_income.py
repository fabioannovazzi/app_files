"""Finite income, independently supplied residual and evidence-bound review."""

from __future__ import annotations

import hashlib
import shutil
from copy import deepcopy
from decimal import Decimal
from pathlib import Path

import pytest
from test_business_valuation import (
    FIXTURE,
    build_valuation,
    case_data,
    compile_html,
    normalized_case,
    planning_case,
    prepare_archive_run,
    read_json,
    restore_imports,
    review,
    run_valuation,
    write_package,
)


def finite_case(base: dict | None = None) -> dict:
    """Keep other fixture methods independent; append one evidenced income model."""
    case = deepcopy(base) if base is not None else case_data()
    case["sources"].append(
        {
            **case["sources"][0],
            "id": "income-basis",
            "path": "finite-income.txt",
            "sha256": hashlib.sha256(
                (FIXTURE / "finite-income.txt").read_bytes()
            ).hexdigest(),
            "description": "Fictional income horizon and independent residual assumptions",
        }
    )
    template = case["inputs"][0]
    case["inputs"].extend(
        [
            {
                **template,
                "id": identifier,
                "value": value,
                "unit": unit,
                "source_ids": ["income-basis"],
                "description": f"Fictional {identifier}",
                "locator": "finite-income.txt: declared amounts and assumptions",
            }
            for identifier, value, unit in (
                ("residual", "500", "EUR"),
                ("second-rate", "0.2", "ratio"),
                ("horizon-rate", "0.3", "ratio"),
            )
        ]
    )
    case["mandate"].update(valuation_date="2024-12-31", information_cutoff="2024-12-31")
    case["methods"].append(
        {
            "id": "finite",
            "kind": "INCOME_EQUITY_FINITE",
            "selected": True,
            "rationale": "Synthetic finite economic life of two years; supplied residual requires review.",
            "inputs": {"incomes": ["income", "income"], "residual_value": "residual"},
            "limitations": [
                "Income availability and capital coherence are not proven by arithmetic."
            ],
            "timing": {
                "valuation_date": "2024-12-31",
                "period_end_dates": ["2025-12-31", "2026-12-31"],
                "cash_flow_timing": "end_period",
                "day_count": "ACT/365F",
                "rate_compounding": "effective_annual",
                "rate_model": "flat",
                "rate_ids": ["rate"],
                "rationale": "Explicit full-year fictional incomes; residual at final period end.",
            },
            "income_basis": {
                "capital_maintenance": "Synthetic income includes expenses assumed to maintain capacity.",
                "reinvestment": "No incremental reinvestment assumed; substantive review still required.",
                "distributions": "Synthetic availability assumption; no legal distributability assessment.",
                "residual_basis": "Independent 500 EUR equity residual, excluding supplied period incomes.",
                "source_ids": ["income-basis"],
                "locator": "finite-income.txt: all assumptions",
                "status": "confirmed",
            },
        }
    )
    return case


@pytest.mark.parametrize(
    ("timing", "expected"),
    [
        ({}, 586.7768595041322),
        ({"cash_flow_timing": "mid_period"}, 595.2478166245717),
        ({"rate_compounding": "continuous"}, 581.7221936503851),
        (
            {
                "rate_model": "spot_curve",
                "rate_ids": ["rate", "second-rate"],
                "terminal_discount_rate": "horizon-rate",
            },
            456.2115235192158,
        ),
        (
            {"rate_model": "forward_curve", "rate_ids": ["rate", "second-rate"]},
            545.4545454545454,
        ),
        ({"period_end_dates": ["2025-06-30", "2025-12-31"]}, 640.8381574573095),
    ],
)
def test_finite_income_has_independent_numeric_results(timing, expected):
    case = finite_case()
    case["methods"][-1]["timing"].update(timing)
    result = build_valuation(case, FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert float(values["finite/value"]) == pytest.approx(expected, rel=1e-12)
    assert result["methods"][-1]["equity_id"] == "finite/value"
    assert (
        result["methods"][-1]["timing"]["terminal_value_convention"]
        == "explicit_equity_residual_at_horizon"
    )


def test_midperiod_income_does_not_move_residual_to_midperiod():
    case = finite_case()
    case["methods"][-1]["timing"]["cash_flow_timing"] = "mid_period"
    result = build_valuation(case, FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert Decimal(values["finite/time/2/cash"]) == Decimal("1.5")
    assert Decimal(values["finite/discount/terminal-horizon"]) == Decimal("1.21")
    assert Decimal(values["finite/pv_residual"]) == Decimal(
        "413.2231404958677685950413223140495867769"
    )


@pytest.mark.parametrize(
    ("identifier", "value", "expected"),
    [
        ("income", "-100", 239.6694214876033),
        ("residual", "-500", -239.66942148760327),
        ("residual", "0", 173.55371900826447),
        ("rate", "0", 700.0),
    ],
)
def test_finite_signed_amounts_and_zero_are_not_replaced_by_perpetuity(
    identifier, value, expected
):
    case = finite_case()
    next(item for item in case["inputs"] if item["id"] == identifier)["value"] = value
    result = build_valuation(case, FIXTURE)
    values = {row["id"]: row["value"] for row in result["calculations"]}
    assert float(values["finite/value"]) == pytest.approx(expected, rel=1e-12)
    assert "finite/terminal_value" not in values


@pytest.mark.parametrize("field", ["timing", "income_basis"])
def test_finite_method_requires_its_explicit_contract(field):
    case = finite_case()
    del case["methods"][-1][field]
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"
    assert result["methods"][0]["status"] == "ready_for_professional_review"


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("residual_value", "income"),
        ("residual_value", "missing"),
        ("residual_value", "rate"),
        ("incomes", []),
        ("incomes", ["rate"]),
        ("terminal_growth", "growth"),
    ],
)
def test_wrong_income_residual_or_implicit_growth_blocks_only_finite_method(
    field, value
):
    case = finite_case()
    case["methods"][-1]["inputs"][field] = value
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"
    assert result["methods"][0]["status"] == "ready_for_professional_review"


def test_missing_residual_remains_blocked_in_report(tmp_path):
    case = finite_case()
    next(item for item in case["inputs"] if item["id"] == "residual")["value"] = None
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"
    assert "Missing input: residual" in compile_html(result)


def test_finite_equity_rejects_a_second_debt_deduction():
    case = finite_case()
    case["methods"][-1]["bridge"] = case["methods"][0]["bridge"]
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("capital_maintenance", ""),
        ("residual_basis", None),
        ("source_ids", ["absent"]),
        ("source_ids", []),
        ("status", "accepted"),
    ],
)
def test_unusable_basis_never_approves_the_method(field, value):
    case = finite_case()
    case["methods"][-1]["income_basis"][field] = value
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "blocked"


def test_proposed_basis_keeps_existing_arithmetic_partial():
    case = finite_case()
    case["methods"][-1]["income_basis"]["status"] = "proposed"
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "partial"
    assert result["methods"][-1]["value_id"] == "finite/value"


def test_unreviewed_basis_source_keeps_method_partial():
    case = finite_case()
    case["sources"][-1]["status"] = "unverified"
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["status"] == "partial"
    assert "income-basis" in result["methods"][-1]["source_ids"]


@pytest.mark.parametrize(
    "field",
    [
        "capital_maintenance",
        "reinvestment",
        "distributions",
        "residual_basis",
        "locator",
    ],
)
def test_changed_income_basis_invalidates_only_dependent_review(field):
    case = finite_case()
    before = build_valuation(case, FIXTURE)
    case["methods"][-1]["review"] = review(before["methods"][-1]["dependency_sha256"])
    case["methods"][0]["review"] = review(before["methods"][0]["dependency_sha256"])
    case["methods"][-1]["income_basis"][field] += " Revised assumption."
    result = build_valuation(case, FIXTURE)
    assert result["methods"][-1]["stale_review"] is True
    assert result["methods"][0]["status"] == "accepted_workpaper"


def test_income_keeps_normalization_formula_and_pending_review():
    result = build_valuation(finite_case(normalized_case()), FIXTURE)
    values = {row["id"]: row for row in result["calculations"]}
    assert result["methods"][-1]["normalization_ids"] == ["income-2026"]
    assert result["methods"][-1]["status"] == "partial"
    assert values["finite/input/income"]["arguments"] == [
        "normalization/income-2026/adjusted"
    ]


def test_plan_fcff_cannot_be_relabelled_as_finite_income(tmp_path):
    case = finite_case(planning_case(tmp_path))
    shutil.copyfile(FIXTURE / "finite-income.txt", tmp_path / "finite-income.txt")
    case["mandate"].update(valuation_date="2026-12-31", information_cutoff="2026-12-31")
    case["methods"][-1]["timing"].update(
        valuation_date="2026-12-31", period_end_dates=["2027-12-31", "2028-12-31"]
    )
    case["methods"][-1]["inputs"]["incomes"] = ["flow", "flow"]
    result = build_valuation(case, tmp_path, replay_parent=tmp_path / "replay")
    assert result["methods"][-1]["status"] == "blocked"
    assert "cannot be relabelled" in result["methods"][-1]["reason"]


def test_finite_export_keeps_literal_basis_and_equity_residual_formula(tmp_path):
    from openpyxl import load_workbook

    case = finite_case()
    basis = case["methods"][-1]["income_basis"]
    basis["capital_maintenance"] = '=HYPERLINK("https://example.invalid")'
    basis["distributions"] = '<script>alert("source")</script>'
    result = build_valuation(case, FIXTURE)
    output = tmp_path / "report"
    write_package(result, FIXTURE, output)
    workbook = load_workbook(output / "valuation_workbook.xlsx", data_only=False)
    sheet = workbook["Base reddituale"]
    assert sheet["C2"].data_type == "s"
    assert sheet["C2"].value == basis["capital_maintenance"]
    html = (output / "valuation_report.html").read_text()
    assert "<script>alert" not in html
    assert "&lt;script&gt;alert" in html
    assert "Non è derivato" in html
    assert read_json(output / "valuation.json") == result
    calculations = workbook["Calcoli"]
    rows = {r[0].value: r[0].row for r in calculations.iter_rows(min_row=2)}
    assert (
        calculations.cell(rows["finite/pv_residual"], 2).value
        == f"=B{rows['finite/input/residual']}/B{rows['finite/discount/terminal-horizon']}"
    )


def test_finite_method_runs_with_exact_archive_receipts(tmp_path, record_property):
    case_path, context = prepare_archive_run(
        tmp_path,
        supplied_case=finite_case(),
        source_files=[FIXTURE / "evidence.txt", FIXTURE / "finite-income.txt"],
    )
    output = run_valuation.run_case(case_path, context)
    report = read_json(Path(output["output_dir"]) / "valuation.json")
    assert report["methods"][-1]["status"] == "ready_for_professional_review"
    assert report["methods"][-1]["income_basis"]["status"] == "confirmed"
    record_property("finite_income_output", output["output_dir"])
