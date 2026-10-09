"""Public Concordato memo application must remain independently replayable."""

from __future__ import annotations

import importlib.util
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any

import pytest
from docx import Document

__all__: list[str] = []

ROOT = Path(__file__).resolve().parents[2]
PLUGIN = Path(
    os.environ.get(
        "VERA_CONCORDATO_TEST_PLUGIN_ROOT", ROOT / "plugins/concordato-plan-review"
    )
)
TEST_HOST = os.environ.get("VERA_CONCORDATO_TEST_HOST", "codex")
MEMO_FILE, MEMO_HEADING = {
    "codex": ("codex_run_review.md", "Memo revisore Codex"),
    "cowork": ("run_review.md", "Memo revisore Claude"),
}[TEST_HOST]


@pytest.fixture(scope="module")
def concordato_tools() -> Any:
    """Reuse source fixtures while exercising the selected real implementation."""

    path = ROOT / "tests/plugins/test_concordato_plan_review_plugin.py"
    spec = importlib.util.spec_from_file_location("concordato_memo_test_tools", path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    module.SCRIPT_DIR = PLUGIN / "scripts"
    module.CORE_PATH = module.SCRIPT_DIR / "concordato_plan_core.py"
    return module


def _prepare_run(tools: Any, tmp_path: Path, nested: bool) -> tuple[Path, dict, Path]:
    core = tools.load_core()
    received = tmp_path / "received"
    received.mkdir()
    tools._save_workbook(
        received / "plan.xlsx", [["Voce", "Importo"], ["Debiti tributari", 100]]
    )
    inputs, root_output, run_id, context = tools._start_managed_concordato_run(
        tmp_path, received
    )
    inspection = core.run_concordato_review(
        inputs, tmp_path / "inspection", tolerance="0"
    )
    recipe = tools._reviewed_source_recipe(
        core,
        inspection,
        {
            str(item["relative_path"]): "concordato_plan"
            for item in inspection.inventory
            if item.get("supported")
        },
    )
    output = root_output / "reviewed-output" if nested else root_output
    core.run_concordato_review(
        inputs,
        output,
        tolerance="0",
        recipe=recipe,
        run_id=run_id,
        input_path_ref="inputs",
        output_path_ref="outputs/reviewed-output" if nested else "outputs",
    )
    return output, tools._concordato_reference_args(output, context), inputs


def _call(tools: Any, name: str, arguments: dict) -> dict:
    responses = tools._call_mcp_server(
        [
            {
                "jsonrpc": "2.0",
                "id": 1,
                "method": "tools/call",
                "params": {"name": name, "arguments": arguments},
            }
        ],
        server_path=PLUGIN / "mcp/server.cjs",
        timeout_seconds=30,
    )
    return tools._private_tool_payload(responses[0])


def _memo_decision(output: Path) -> dict:
    review = json.loads((output / "review_payload.json").read_bytes())
    item = next(
        row for row in review["items"] if row["item_type"] == "codex_review_memo"
    )
    return {
        "item_id": item["id"],
        "action": "edit",
        "edit_value": "# Memo professionale\n\n- Attestazione da acquisire.\nConclusioni sospese.",
    }


def _files(root: Path) -> dict[str, bytes]:
    return {
        path.relative_to(root).as_posix(): path.read_bytes()
        for path in root.rglob("*")
        if path.is_file()
    }


@pytest.mark.parametrize("nested", [False, True], ids=["root", "reviewed-output"])
@pytest.mark.parametrize("save_first", [False, True], ids=["direct", "saved"])
def test_public_memo_application_replays_without_changing_facts(
    concordato_tools: Any, tmp_path: Path, nested: bool, save_first: bool
) -> None:
    output, reference, inputs = _prepare_run(concordato_tools, tmp_path, nested)
    facts = (output / "concordato_case_model.json").read_bytes()
    sources = _files(inputs)
    numeric = (output / "concordato_review_summary.docx").read_bytes()
    predecessor = json.loads((output / "workflow_output_closure.json").read_bytes())
    decision = _memo_decision(output)
    arguments = {**reference, "decisions": [decision]}
    if save_first:
        saved = _call(concordato_tools, "save_concordato_plan_decisions", arguments)
        assert saved["ok"] is True, saved
        arguments = {
            **reference,
            "decisions": json.loads((output / "ui_decisions.json").read_bytes())[
                "decisions"
            ],
        }
        predecessor = json.loads((output / "workflow_output_closure.json").read_bytes())

    applied = _call(concordato_tools, "apply_concordato_plan_decisions", arguments)
    validated = _call(concordato_tools, "validate_concordato_plan_review", reference)

    assert applied["ok"] is True, applied
    assert validated["ok"] is True, validated
    assert (output / MEMO_FILE).read_text() == decision["edit_value"]
    assert (output / "concordato_case_model.json").read_bytes() == facts
    assert _files(inputs) == sources
    assert (output / "concordato_review_summary.docx").read_bytes() == numeric
    summary = Document(output / "concordato_preventivo_review_summary.docx")
    assert "Attestazione da acquisire." in [
        paragraph.text for paragraph in summary.paragraphs
    ]
    closure = json.loads((output / "workflow_output_closure.json").read_bytes())
    assert closure["previous_closure_content_sha256"] == predecessor["content_sha256"]
    assert closure["phase"] == "review_apply_finalization"
    assert (
        json.loads((output / "final_artifacts.json").read_bytes())["final_ready"]
        is False
    )


@pytest.mark.parametrize(
    "artifact",
    [
        "concordato_preventivo_review_summary.docx",
        "concordato_review_summary.docx",
        "memo",
    ],
)
def test_public_validate_rejects_altered_memo_artifacts_without_writes(
    concordato_tools: Any, tmp_path: Path, artifact: str
) -> None:
    output, reference, _ = _prepare_run(concordato_tools, tmp_path, False)
    applied = _call(
        concordato_tools,
        "apply_concordato_plan_decisions",
        {**reference, "decisions": [_memo_decision(output)]},
    )
    assert applied["ok"] is True, applied
    target = output / (MEMO_FILE if artifact == "memo" else artifact)
    if target.suffix == ".docx":
        document = Document(target)
        document.add_paragraph("Unapproved additional conclusion.")
        document.save(target)
    else:
        target.write_text("Unapproved replacement memo.\n")
    before = _files(output)

    result = _call(concordato_tools, "validate_concordato_plan_review", reference)

    assert result["ok"] is False, result
    assert _files(output) == before


def test_public_validate_without_memo_preserves_the_complete_output(
    concordato_tools: Any, tmp_path: Path
) -> None:
    output, reference, _ = _prepare_run(concordato_tools, tmp_path, False)
    before = _files(output)

    result = _call(concordato_tools, "validate_concordato_plan_review", reference)

    assert result["ok"] is True, result
    assert _files(output) == before


def _apply_memo(tools: Any, output: Path, reference: dict, text: str) -> dict:
    decision = {**_memo_decision(output), "edit_value": text}
    result = _call(
        tools, "apply_concordato_plan_decisions", {**reference, "decisions": [decision]}
    )
    assert result["ok"] is True, result
    return decision


@pytest.mark.parametrize("nested", [False, True], ids=["root", "reviewed-output"])
@pytest.mark.parametrize("save_first", [False, True], ids=["direct", "saved"])
def test_public_memo_replacement_retains_prior_authority_and_replays(
    concordato_tools: Any, tmp_path: Path, nested: bool, save_first: bool
) -> None:
    output, reference, inputs = _prepare_run(concordato_tools, tmp_path, nested)
    first_text = "# Prima revisione\n\nDocumentazione ancora mancante."
    first = _apply_memo(concordato_tools, output, reference, first_text)
    authority_before = _files(output / "revisions/authority")
    facts = (output / "concordato_case_model.json").read_bytes()
    numeric = (output / "concordato_review_summary.docx").read_bytes()
    sources = _files(inputs)
    replacement = {
        **first,
        "edit_value": "# Seconda revisione\n\nDocumentazione ricevuta; conclusioni sospese.",
    }
    arguments = {**reference, "decisions": [replacement]}
    if save_first:
        saved = _call(concordato_tools, "save_concordato_plan_decisions", arguments)
        assert saved["ok"] is True, saved

    result = _call(concordato_tools, "apply_concordato_plan_decisions", arguments)

    assert result["ok"] is True, result
    validated = _call(concordato_tools, "validate_concordato_plan_review", reference)
    assert validated["ok"] is True, validated
    after = _files(output / "revisions/authority")
    assert {path: after[path] for path in authority_before} == authority_before
    applied = json.loads((output / "applied_decisions.json").read_bytes())
    assert len(applied["review_history"]) == 1
    original = json.loads(
        (
            output / applied["review_history"][0]["applied_decisions"]["path"]
        ).read_bytes()
    )
    assert original["decisions"][0]["edit_value"] == first_text
    summary = Document(output / "concordato_preventivo_review_summary.docx")
    paragraphs = [paragraph.text for paragraph in summary.paragraphs]
    assert paragraphs.count(MEMO_HEADING) == 2
    assert "Documentazione ancora mancante." in paragraphs
    assert "Documentazione ricevuta; conclusioni sospese." in paragraphs
    assert (output / MEMO_FILE).read_text() == replacement["edit_value"]
    assert (output / "concordato_case_model.json").read_bytes() == facts
    assert (output / "concordato_review_summary.docx").read_bytes() == numeric
    assert _files(inputs) == sources


@pytest.mark.parametrize(
    "repeat_literal", [False, True], ids=["third-edit", "literal-retry"]
)
def test_public_later_memo_apply_preserves_history_and_original_backups(
    concordato_tools: Any, tmp_path: Path, repeat_literal: bool
) -> None:
    output, reference, _ = _prepare_run(concordato_tools, tmp_path, False)
    first = _apply_memo(
        concordato_tools, output, reference, "# Prima\n\nNota iniziale."
    )
    second = _apply_memo(
        concordato_tools, output, reference, "# Seconda\n\nNota successiva."
    )
    retained = _files(output / "revisions/authority")
    original_backups = _files(output / "revisions/originals")
    docx_before = (output / "concordato_preventivo_review_summary.docx").read_bytes()
    text = second["edit_value"] if repeat_literal else "# Terza\n\nEsame ancora aperto."
    decision = {**first, "edit_value": text}

    result = _call(
        concordato_tools,
        "apply_concordato_plan_decisions",
        {**reference, "decisions": [decision]},
    )

    assert result["ok"] is True, result
    assert _files(output / "revisions/originals") == original_backups
    after = _files(output / "revisions/authority")
    assert {path: after[path] for path in retained} == retained
    applied = json.loads((output / "applied_decisions.json").read_bytes())
    assert len(applied["review_history"]) == 2
    summary = Document(output / "concordato_preventivo_review_summary.docx")
    expected_sections = 2 if repeat_literal else 3
    assert [paragraph.text for paragraph in summary.paragraphs].count(
        MEMO_HEADING
    ) == expected_sections
    if repeat_literal:
        assert (
            output / "concordato_preventivo_review_summary.docx"
        ).read_bytes() == docx_before
    validated = _call(concordato_tools, "validate_concordato_plan_review", reference)
    assert validated["ok"] is True, validated


def test_public_save_after_apply_preserves_applied_state_and_replay(
    concordato_tools: Any, tmp_path: Path
) -> None:
    output, reference, _ = _prepare_run(concordato_tools, tmp_path, False)
    first = _apply_memo(
        concordato_tools, output, reference, "# Applicato\n\nGiudizio sospeso."
    )
    applied_before = (output / "applied_decisions.json").read_bytes()
    summary_before = (output / "concordato_preventivo_review_summary.docx").read_bytes()
    draft = {**first, "edit_value": "# Bozza\n\nUlteriore verifica richiesta."}

    result = _call(
        concordato_tools,
        "save_concordato_plan_decisions",
        {**reference, "decisions": [draft]},
    )

    assert result["ok"] is True, result
    assert (output / "applied_decisions.json").read_bytes() == applied_before
    assert (
        output / "concordato_preventivo_review_summary.docx"
    ).read_bytes() == summary_before
    assert (output / MEMO_FILE).read_text() == first["edit_value"]
    validated = _call(concordato_tools, "validate_concordato_plan_review", reference)
    assert validated["ok"] is True, validated


@pytest.mark.parametrize("receipt_kind", ["current_ui", "previous_apply"])
def test_public_replay_cli_rejects_changed_review_authority_before_closure_check(
    concordato_tools: Any, tmp_path: Path, receipt_kind: str
) -> None:
    output, reference, _ = _prepare_run(concordato_tools, tmp_path, False)
    _apply_memo(
        concordato_tools, output, reference, "# Prima\n\nEvidenze da verificare."
    )
    _apply_memo(
        concordato_tools, output, reference, "# Seconda\n\nConclusioni sospese."
    )
    applied = json.loads((output / "applied_decisions.json").read_bytes())
    receipt = (
        applied["ui_decisions_receipt"]
        if receipt_kind == "current_ui"
        else applied["review_history"][0]["applied_decisions"]
    )
    target = output / receipt["path"]
    changed = json.loads(target.read_bytes())
    changed["reviewer"] = "Unapproved replacement attribution"
    target.write_text(json.dumps(changed, ensure_ascii=False, indent=2) + "\n")
    before = _files(output)

    result = subprocess.run(
        [
            sys.executable,
            str(PLUGIN / "scripts/replay_assurance.py"),
            "--output-dir",
            str(output),
            "--client-engagement",
            reference["client_engagement"],
        ],
        capture_output=True,
        text=True,
        check=False,
        timeout=30,
    )

    assert result.returncode == 1
    assert json.loads(result.stdout) == {
        "ok": False,
        "error": "Concordato review authority receipt changed",
    }
    assert _files(output) == before


def test_public_apply_keeps_complete_review_history_out_of_model_summary(
    concordato_tools: Any, tmp_path: Path
) -> None:
    output, reference, _ = _prepare_run(concordato_tools, tmp_path, False)
    first_text = "# Prima revisione\n\nNota privata del revisore iniziale."
    first = _apply_memo(concordato_tools, output, reference, first_text)
    second_text = "# Seconda revisione\n\nNota privata del revisore successivo."
    arguments = {**reference, "decisions": [{**first, "edit_value": second_text}]}

    responses = concordato_tools._call_mcp_server(
        [
            {
                "jsonrpc": "2.0",
                "id": 1,
                "method": "tools/call",
                "params": {
                    "name": "apply_concordato_plan_decisions",
                    "arguments": arguments,
                },
            }
        ],
        server_path=PLUGIN / "mcp/server.cjs",
        timeout_seconds=30,
    )

    result = responses[0]["result"]
    model = result["structuredContent"]
    assert model["ok"] is True, model
    model_text = json.dumps(model, ensure_ascii=False)
    assert "review_history" not in model_text
    assert "revisions/authority" not in model_text
    assert "Nota privata del revisore iniziale." not in model_text
    assert "Nota privata del revisore successivo." not in model_text
    private = result["_meta"]["widget_payload"]["applied_decisions"]
    assert len(private["review_history"]) == 1
    assert private["decisions"][0]["edit_value"] == second_text


@pytest.mark.parametrize("linked_parent", ["revisions", "revisions/authority"])
def test_public_replay_refuses_linked_receipt_parent_before_external_read(
    concordato_tools: Any, tmp_path: Path, linked_parent: str
) -> None:
    output, reference, _ = _prepare_run(concordato_tools, tmp_path, False)
    _apply_memo(
        concordato_tools, output, reference, "# Revisione\n\nConclusioni sospese."
    )
    receipt_parent = output / linked_parent
    external = tmp_path / "outside_customer_receipts"
    shutil.move(receipt_parent, external)
    receipt_parent.symlink_to(external, target_is_directory=True)
    before = _files(output)
    external_before = _files(external)
    script = (
        "import pathlib, runpy, sys\n"
        "script, linked, external, *arguments = sys.argv[1:]\n"
        "linked = pathlib.Path(linked)\n"
        "external = pathlib.Path(external).resolve()\n"
        "def refuse_external_open(event, args):\n"
        "    if event == 'open' and isinstance(args[0], (str, bytes)):\n"
        "        path = pathlib.Path(args[0])\n"
        "        if path.is_relative_to(linked) or path.resolve().is_relative_to(external):\n"
        "            sys.stderr.write('OUTSIDE_CASE_READ_ATTEMPT\\n')\n"
        "            raise RuntimeError('outside customer receipt read attempted')\n"
        "sys.addaudithook(refuse_external_open)\n"
        "sys.path.insert(0, str(pathlib.Path(script).parent))\n"
        "sys.argv = [script, *arguments]\n"
        "runpy.run_path(script, run_name='__main__')\n"
    )

    result = subprocess.run(
        [
            sys.executable,
            "-c",
            script,
            str(PLUGIN / "scripts/replay_assurance.py"),
            str(receipt_parent),
            str(external),
            "--output-dir",
            str(output),
            "--client-engagement",
            reference["client_engagement"],
        ],
        capture_output=True,
        text=True,
        check=False,
        timeout=30,
    )

    assert "OUTSIDE_CASE_READ_ATTEMPT" not in result.stderr
    assert result.returncode == 1
    assert json.loads(result.stdout) == {
        "ok": False,
        "error": "Concordato review authority directories must be real directories",
    }
    assert _files(output) == before
    assert _files(external) == external_before
