"""File intake, persistence, artifact readback and the real archive boundary."""

from __future__ import annotations

import csv
import http.client
import io
import json
import subprocess
import sys
import threading
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(
    0, str(Path(__file__).resolve().parents[2] / "plugins/treasury-forecast/scripts")
)

from treasury_core import TreasuryError
from treasury_inputs import HEADERS, load_inputs, read_json, safe_path, write_templates
from treasury_report import write_artifacts
from treasury_server import make_server
from treasury_session import (
    create_session,
    current_record,
    review_session,
    save_scenario,
)

from tests.plugins.test_treasury_forecast import SCRIPTS, accept, first, second
from tests.plugins.test_vera_client_workflow_filesystem import _load_customer_ledger


def csv_sources(data: dict) -> dict[str, str]:
    result = {}
    for name, header in HEADERS.items():
        buffer = io.StringIO(newline="")
        writer = csv.DictWriter(buffer, fieldnames=header)
        writer.writeheader()
        writer.writerows(data[name])
        result[name + ".csv"] = buffer.getvalue()
    return result


def manifest(data: dict) -> dict:
    return {
        "schema_version": "vera.treasury_manifest.v1",
        **{
            key: data[key]
            for key in (
                "company_id",
                "company_name",
                "currency",
                "as_of",
                "horizon_end",
                "coverage",
            )
        },
        "tables": {name: {"path": name + ".csv"} for name in HEADERS},
        "invoice_files": [],
        "previous": None,
    }


def intake(tmp_path: Path) -> tuple[dict, Path]:
    root = tmp_path / "sources"
    root.mkdir()
    for name, content in csv_sources(first()).items():
        (root / name).write_text(content)
    return manifest(first()), root


def session(tmp_path: Path) -> tuple[Path, dict, Path]:
    source = tmp_path / "supplied.json"
    source.write_text(json.dumps(first()))
    output = tmp_path / "outputs"
    record = create_session(
        output,
        first(),
        None,
        source_paths=[source],
        context_path=tmp_path / "context.json",
    )
    return output, record, source


def test_exact_csv_intake_preserves_outstanding_amounts_and_source_hashes(tmp_path):
    supplied, root = intake(tmp_path)
    data, previous = load_inputs(
        supplied, root, client_id="client-a", engagement_id="engagement-a"
    )
    assert data["open_items"][0]["amount"] == "50000.00"
    assert len(data["sources"]) == 6
    assert previous is None


def test_xlsx_intake_reads_exact_sheets_and_rejects_formulas(tmp_path):
    supplied = manifest(first())
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, header in HEADERS.items():
        sheet = workbook.create_sheet(name)
        sheet.append(header)
        for row in first()[name]:
            sheet.append([row[key] for key in header])
        supplied["tables"][name] = {"path": "input.xlsx", "sheet": name}
    workbook["accounts"]["B2"] = 40000
    workbook.save(tmp_path / "input.xlsx")
    data, _ = load_inputs(
        supplied, tmp_path, client_id="client-a", engagement_id="engagement-a"
    )
    assert data["accounts"][0]["balance"] == "40000.00"
    workbook["accounts"]["B2"] = "=40000"
    workbook.save(tmp_path / "input.xlsx")
    with pytest.raises(TreasuryError, match="formulas"):
        load_inputs(
            supplied, tmp_path, client_id="client-a", engagement_id="engagement-a"
        )
    workbook.close()


@pytest.mark.parametrize(
    "contents",
    [
        "account_id,total\na,20\n",
        "account_id,balance\na,20,extra\n",
        "account_id,balance\na,1.000,00\n",
    ],
)
def test_malformed_csv_is_rejected(tmp_path, contents):
    supplied, root = intake(tmp_path)
    (root / "accounts.csv").write_text(contents)
    with pytest.raises(TreasuryError):
        load_inputs(supplied, root, client_id="client-a", engagement_id="engagement-a")


def test_malformed_xml_is_a_readable_intake_failure(tmp_path):
    supplied, root = intake(tmp_path)
    (root / "bad.xml").write_text("<Fattura><")
    supplied["invoice_files"] = ["bad.xml"]
    with pytest.raises(TreasuryError, match="Unreadable FatturaPA"):
        load_inputs(supplied, root, client_id="client-a", engagement_id="engagement-a")


def test_fatturapa_evidence_keeps_payment_terms_separate_from_outstanding(tmp_path):
    supplied, root = intake(tmp_path)
    invoice = """<FatturaElettronica><FatturaElettronicaHeader>
<CedentePrestatore><DatiAnagrafici><IdFiscaleIVA><IdPaese>IT</IdPaese><IdCodice>01234567890</IdCodice></IdFiscaleIVA><Anagrafica><Denominazione>Synthetic supplier</Denominazione></Anagrafica></DatiAnagrafici></CedentePrestatore>
<CessionarioCommittente><DatiAnagrafici><CodiceFiscale>99999999999</CodiceFiscale><Anagrafica><Denominazione>Synthetic customer</Denominazione></Anagrafica></DatiAnagrafici></CessionarioCommittente>
</FatturaElettronicaHeader><FatturaElettronicaBody><DatiGenerali><DatiGeneraliDocumento><TipoDocumento>TD01</TipoDocumento><Divisa>EUR</Divisa><Data>2026-09-01</Data><Numero>XML-1</Numero><ImportoTotaleDocumento>122.00</ImportoTotaleDocumento></DatiGeneraliDocumento></DatiGenerali>
<DatiPagamento><CondizioniPagamento>TP02</CondizioniPagamento><DettaglioPagamento><ModalitaPagamento>MP05</ModalitaPagamento><DataScadenzaPagamento>2026-09-30</DataScadenzaPagamento><ImportoPagamento>122.00</ImportoPagamento></DettaglioPagamento></DatiPagamento>
</FatturaElettronicaBody></FatturaElettronica>"""
    (root / "invoice.xml").write_text(invoice)
    (root / "copy.xml").write_text(invoice)
    supplied["invoice_files"] = ["invoice.xml", "copy.xml"]
    data, _ = load_inputs(
        supplied, root, client_id="client-a", engagement_id="engagement-a"
    )
    assert len(data["invoice_evidence"]) == 1
    assert data["invoice_evidence"][0]["payments"][0]["due_date"] == "2026-09-30"
    assert data["open_items"][0]["amount"] == "50000.00"
    (root / "copy.xml").write_text(invoice.replace("122.00", "244.00"))
    with pytest.raises(TreasuryError, match="same invoice identity"):
        load_inputs(supplied, root, client_id="client-a", engagement_id="engagement-a")


def test_intake_rejects_path_escape_and_symlink(tmp_path):
    source = tmp_path / "outside.csv"
    source.write_text("x")
    root = tmp_path / "inputs"
    root.mkdir()
    (root / "link.csv").symlink_to(source)
    with pytest.raises(TreasuryError, match="relative"):
        safe_path(root, "../outside.csv")
    with pytest.raises(TreasuryError, match="Linked"):
        safe_path(root, "link.csv")


def test_duplicate_json_fields_are_rejected(tmp_path):
    path = tmp_path / "input.json"
    path.write_text('{"amount":"10","amount":"20"}')
    with pytest.raises(TreasuryError, match="Duplicate"):
        read_json(path)


def test_templates_do_not_overwrite_existing_sources(tmp_path):
    write_templates(tmp_path)
    with pytest.raises(TreasuryError, match="no overwrite"):
        write_templates(tmp_path)
    assert (tmp_path / "accounts.csv").read_text().strip() == "account_id,balance"


def test_saved_review_changes_forecast_and_survives_reopen(tmp_path):
    output, original, source = session(tmp_path)
    changed = review_session(
        output,
        expected_record_sha256=original["record_sha256"],
        decisions={
            "item:C101": {
                "expected_date": "2026-10-02",
                "basis": "Confirmed customer forecast date",
            }
        },
    )
    state, reopened = current_record(output)
    workbook = load_workbook(
        output / "versions" / changed["record_sha256"] / "tesoreria.xlsx",
        data_only=True,
    )
    assert reopened["minimum_daily_cash"] == "-35000.00"
    assert workbook["Sintesi"]["B6"].value == -35000
    assert len(state["history"]) == 2
    assert (output / "versions" / original["record_sha256"] / "forecast.json").exists()
    workbook.close()
    resumed = create_session(
        output,
        first(),
        None,
        source_paths=[source],
        context_path=tmp_path / "context.json",
    )
    assert resumed["record_sha256"] == reopened["record_sha256"]


def test_acceptance_is_bound_and_accepted_session_cannot_be_edited(tmp_path):
    output, record, _ = session(tmp_path)
    accepted = review_session(
        output,
        expected_record_sha256=record["record_sha256"],
        review={
            "proposal_sha256": record["proposal_sha256"],
            "reviewer_ref": "Synthetic reviewer",
            "reviewed_at": "2026-09-10",
            "conclusion": "Assumptions reviewed.",
        },
    )
    assert current_record(output)[1]["status"] == "accepted"
    with pytest.raises(TreasuryError, match="immutable"):
        review_session(
            output, expected_record_sha256=accepted["record_sha256"], decisions={}
        )


@pytest.mark.parametrize(
    "changed_file", ["source", "prepared", "workbook", "inventory"]
)
def test_changed_bound_content_blocks_review(tmp_path, changed_file):
    output, record, source = session(tmp_path)
    targets = {
        "source": source,
        "prepared": output / "prepared_inputs.json",
        "workbook": output / "versions" / record["record_sha256"] / "tesoreria.xlsx",
        "inventory": output
        / "versions"
        / record["record_sha256"]
        / "artifact_manifest.json",
    }
    targets[changed_file].write_text("{}")
    with pytest.raises(TreasuryError):
        review_session(
            output, expected_record_sha256=record["record_sha256"], decisions={}
        )


def test_stale_and_concurrent_reviews_preserve_current_forecast(tmp_path):
    output, record, _ = session(tmp_path)
    with pytest.raises(TreasuryError, match="Stale"):
        review_session(output, expected_record_sha256="0" * 64)
    (output / ".treasury-review.lock").write_text("another writer")
    with pytest.raises(TreasuryError, match="in progress"):
        review_session(output, expected_record_sha256=record["record_sha256"])
    assert current_record(output)[1] == record


def test_failed_render_keeps_previous_version_current(tmp_path, monkeypatch):
    output, record, _ = session(tmp_path)

    def fail(*args):
        raise OSError("Synthetic disk failure")

    monkeypatch.setattr("treasury_session.write_artifacts", fail)
    with pytest.raises(OSError, match="disk failure"):
        review_session(
            output,
            expected_record_sha256=record["record_sha256"],
            decisions={
                "item:C101": {"expected_date": "2026-10-02", "basis": "Reviewed"}
            },
        )
    assert current_record(output)[1] == record
    assert not list(output.glob(".treasury-stage-*"))


def test_alternative_is_saved_without_mutating_baseline(tmp_path):
    output, record, _ = session(tmp_path)
    scenario = save_scenario(
        output,
        expected_record_sha256=record["record_sha256"],
        dates={"item:C101": "2026-10-02"},
    )
    assert scenario["minimum_daily_cash"] == "-35000.00"
    assert current_record(output)[1] == record
    assert len(list(output.glob("scenario-*.json"))) == 1


def test_report_escapes_supplied_markup_and_excel_formula_text(tmp_path):
    data = first()
    data["company_name"] = "<script>alert(1)</script>"
    data["planned_flows"][0]["description"] = '=HYPERLINK("https://invalid.test")'
    record = accept(data)
    write_artifacts(tmp_path / "version", record)
    workbook = load_workbook(tmp_path / "version/tesoreria.xlsx", data_only=False)
    assert "<script>" not in (tmp_path / "version/report.html").read_text()
    assert workbook["Flussi"]["B5"].data_type == "s"
    assert workbook["Flussi"]["B5"].value.startswith("=HYPERLINK")
    workbook.close()


def test_http_review_requires_token_origin_and_recalculates(tmp_path):
    output, record, _ = session(tmp_path)
    server, token = make_server(output, lambda: None)
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    connection = http.client.HTTPConnection("127.0.0.1", server.server_port, timeout=5)
    try:
        connection.request("GET", "/api/state")
        response = connection.getresponse()
        assert response.status == 403
        response.read()
        body = json.dumps(
            {
                "record_sha256": record["record_sha256"],
                "decisions": {
                    "item:C101": {
                        "expected_date": "2026-10-02",
                        "basis": "Reviewed customer date",
                    }
                },
            }
        )
        headers = {"X-Treasury-Token": token, "Content-Type": "application/json"}
        connection.request("POST", "/api/review", body, headers)
        response = connection.getresponse()
        assert response.status == 403
        response.read()
        headers["Origin"] = f"http://127.0.0.1:{server.server_port}"
        connection.request("POST", "/api/review", body, headers)
        response = connection.getresponse()
        assert response.status == 200
        response.read()
        assert current_record(output)[1]["minimum_daily_cash"] == "-35000.00"
        connection.request("POST", "/api/review", "[]", headers)
        response = connection.getresponse()
        assert response.status == 409
        response.read()
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        worker.join(timeout=5)


def archived_case(tmp_path: Path) -> dict:
    ledger = _load_customer_ledger()
    client = tmp_path / "Studio client"
    client.mkdir()
    client_id = "client_111111111111111111111111"
    ledger.create_client_manifest(client, client_id)
    engagement_id = ledger.create_engagement(client, client_id, "Synthetic treasury")[
        "engagement_id"
    ]
    supplied = manifest(first())
    input_ids = []
    staging = tmp_path / "staging"
    staging.mkdir()
    for name, content in csv_sources(first()).items():
        path = staging / name
        path.write_text(content)
        receipt = ledger.import_document(
            client, client_id, engagement_id, path, "source"
        )["receipt"]
        input_ids.append(receipt["input_id"])
        supplied["tables"][path.stem][
            "path"
        ] = f"imports/{receipt['input_id']}/{Path(receipt['relative_path']).name}"
    path = staging / "manifest.json"
    path.write_text(json.dumps(supplied))
    input_ids.append(
        ledger.import_document(client, client_id, engagement_id, path, "source")[
            "receipt"
        ]["input_id"]
    )
    prepared = ledger.prepare_run(
        client,
        client_id,
        engagement_id,
        "treasury-forecast",
        "0.1.0",
        input_ids=input_ids,
    )
    running = ledger.start_run(client, engagement_id, prepared["run"]["run_id"])
    return {
        "ledger": ledger,
        "client": client,
        "engagement_id": engagement_id,
        **running,
    }


def test_cli_prepares_and_finalizes_actual_archive_run(tmp_path):
    from tests.model_data_helpers import write_no_model_report

    workspace = archived_case(tmp_path)
    context = workspace["context"]
    manifest_path = next(
        Path(row["path"])
        for row in context["input_bindings"]
        if row["path"].endswith(".json")
    )
    command = [
        sys.executable,
        str(SCRIPTS / "run_treasury.py"),
        "prepare",
        "--client-engagement",
        str(workspace["context_path"]),
        "--manifest",
        str(manifest_path),
    ]
    result = subprocess.run(command, capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stderr
    output = Path(workspace["output_dir"])
    assert current_record(output)[1]["opening_cash"] == "40000.00"
    write_no_model_report(output, "treasury-forecast", context["run_id"])
    declarations = [
        {
            "artifact_id": f"treasury_{index}",
            "path": path.relative_to(output).as_posix(),
            "purpose": "Synthetic workflow validation",
            "audience": "internal",
            "media_type": "application/octet-stream",
        }
        for index, path in enumerate(
            sorted(path for path in output.rglob("*") if path.is_file())
        )
    ]
    finalized = workspace["ledger"].finalize_run(
        workspace["client"], workspace["engagement_id"], context["run_id"], declarations
    )
    assert finalized["run"]["status"] == "ready_for_review"
    rejected = subprocess.run(command, capture_output=True, text=True, timeout=30)
    assert rejected.returncode != 0


def test_reused_bank_evidence_id_is_rejected_across_updates():
    previous = accept(
        second(),
        previous=accept(first()),
        decisions={
            "item:C101": {
                "expected_date": "2026-10-02",
                "basis": "Customer expected date",
            }
        },
    )
    data = second()
    data["as_of"] = "2026-09-19"
    data["bank_movements"][0]["date"] = "2026-09-19"
    with pytest.raises(TreasuryError, match="Previously consumed"):
        from treasury_core import build_forecast

        build_forecast(data, previous=previous)


def test_cli_review_context_scenario_and_templates_use_bound_outputs(tmp_path):
    import run_treasury

    workspace = archived_case(tmp_path)
    context = workspace["context"]
    bound = ["--client-engagement", str(workspace["context_path"])]
    manifest_path = next(
        row["path"]
        for row in context["input_bindings"]
        if row["path"].endswith(".json")
    )
    assert run_treasury.main(["prepare", *bound, "--manifest", manifest_path]) == 0
    output = Path(workspace["output_dir"])
    record = current_record(output)[1]
    review = output / "date-review.json"
    review.write_text(
        json.dumps(
            {
                "record_sha256": record["record_sha256"],
                "decisions": {
                    "item:C101": {
                        "expected_date": "2026-10-02",
                        "basis": "Reviewed expected date",
                    }
                },
            }
        )
    )
    assert run_treasury.main(["review", *bound, "--request", str(review)]) == 0
    assert run_treasury.main(["context", *bound, "--event-id", "item:C101"]) == 0
    alternative = output / "alternative-request.json"
    alternative.write_text(json.dumps({"item:C101": "2026-09-15"}))
    assert run_treasury.main(["scenario", *bound, "--request", str(alternative)]) == 0
    assert len(list(output.glob("scenario-*.json"))) == 1
    assert run_treasury.main(["context", *bound, "--event-id", "item:UNKNOWN"]) == 2
    assert (output / "treasury_blocked.json").is_file()
    templates = tmp_path / "templates"
    assert run_treasury.main(["templates", "--output", str(templates)]) == 0
    assert (templates / "accounts.csv").is_file()


def test_dependency_check_reports_supported_and_incompatible_versions(monkeypatch):
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "treasury_dependency_test", SCRIPTS / "check_dependencies.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    monkeypatch.setattr(sys, "argv", ["check_dependencies.py"])
    assert module.main() == 0
    monkeypatch.setattr(module.importlib.metadata, "version", lambda name: "2.6.4")
    with pytest.raises(SystemExit, match="shared environment"):
        module.main()


def test_report_writes_explained_update_and_incomplete_status(tmp_path):
    from treasury_core import build_forecast

    updated = accept(
        second(),
        previous=accept(first()),
        decisions={
            "item:C101": {"expected_date": "2026-10-02", "basis": "Customer date"}
        },
    )
    write_artifacts(tmp_path / "updated", updated)
    assert "Variazioni rispetto" in (tmp_path / "updated/report.md").read_text()
    workbook = load_workbook(tmp_path / "updated/tesoreria.xlsx", data_only=True)
    assert workbook["Sintesi"]["B6"].value == -16000
    workbook.close()
    missing = first()
    missing["open_items"][0]["due_date"] = ""
    write_artifacts(tmp_path / "missing", build_forecast(missing))
    workbook = load_workbook(tmp_path / "missing/tesoreria.xlsx", data_only=True)
    assert workbook["Sintesi"]["B7"].value == "Previsione incompleta"
    assert "Previsione incompleta" in (tmp_path / "missing/report.html").read_text()
    workbook.close()


def test_http_assets_state_download_and_saved_alternative(tmp_path):
    output, record, _ = session(tmp_path)
    server, token = make_server(output, lambda: None)
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    connection = http.client.HTTPConnection("127.0.0.1", server.server_port, timeout=5)
    headers = {
        "X-Treasury-Token": token,
        "Origin": f"http://127.0.0.1:{server.server_port}",
    }
    try:
        for route in (
            "/",
            "/review.js",
            "/review.css",
            "/api/state",
            "/api/workbook",
            "/api/report",
        ):
            connection.request("GET", route, headers=headers)
            response = connection.getresponse()
            assert response.status == 200
            assert response.read()
        connection.request("GET", "/api/state?offset=-1", headers=headers)
        response = connection.getresponse()
        assert response.status == 409
        response.read()
        connection.request(
            "POST",
            "/api/scenario",
            json.dumps(
                {
                    "record_sha256": record["record_sha256"],
                    "dates": {"item:C101": "2026-10-02"},
                }
            ),
            headers,
        )
        response = connection.getresponse()
        assert response.status == 200
        response.read()
        assert len(list(output.glob("scenario-*.json"))) == 1
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        worker.join(timeout=5)
