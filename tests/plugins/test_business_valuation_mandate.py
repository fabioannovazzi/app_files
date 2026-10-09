"""Engagement evidence, explicit standards and invalidation without qualification."""

from __future__ import annotations

import hashlib
import shutil
from copy import deepcopy
from decimal import Decimal
from pathlib import Path

import pytest

from tests.plugins.test_business_valuation import (
    FIXTURE,
    ValuationError,
    build_valuation,
    case_data,
    claimed_case,
    read_json,
    restore_imports,
    review,
    reviewed_normalized_case,
    write_package,
)


def reviewed_engagement_case() -> dict:
    """Prepare synthetic attestations only for exercising dependency changes."""
    case = reviewed_normalized_case()
    case["claims"] = claimed_case()["claims"]
    prepared = build_valuation(case, FIXTURE)
    case["mandate_details"]["review"] = review(
        prepared["mandate_assessment"]["dependency_sha256"]
    )
    case["claims"][0]["review"] = review(prepared["claims"][0]["dependency_sha256"])
    case["conclusion"] = {
        "text": "Invented conclusion for dependency testing only",
        "method_ids": ["fcff"],
        "review": None,
    }
    prepared = build_valuation(case, FIXTURE)
    case["conclusion"]["review"] = review(prepared["conclusion"]["dependency_sha256"])
    return case


@pytest.mark.parametrize(
    "field",
    [
        "expert_identity",
        "written_mandate",
        "remuneration",
        "delivery_terms",
        "amendments",
    ],
)
@pytest.mark.parametrize(
    ("missing", "value"),
    [
        ("record", None),
        ("value", None),
        ("status", "proposed"),
        ("source_ids", []),
        ("locator", None),
    ],
)
def test_missing_engagement_evidence_keeps_values_but_prevents_completion(
    field: str, missing: str, value: object
) -> None:
    case = case_data()
    if missing == "record":
        del case["mandate_details"][field]
    else:
        case["mandate_details"][field][missing] = value

    result = build_valuation(case, FIXTURE)

    assert result["mandate_assessment"]["status"] == "partial"
    assert f"Mandate {field}: evidence or confirmation pending" in result["issues"]
    assert result["status"] == "partial"
    assert Decimal(
        next(
            row["value"] for row in result["calculations"] if row["id"] == "fcff/equity"
        )
    ) == Decimal("750")


@pytest.mark.parametrize("absent", [True, False])
def test_missing_standard_selection_stays_unknown_without_date_inference(
    absent: bool,
) -> None:
    case = case_data()
    if absent:
        del case["mandate_details"]["standards"]
    else:
        case["mandate_details"]["standards"] = []

    result = build_valuation(case, FIXTURE)

    assert result["status"] == "partial"
    assert result["mandate_assessment"]["details"].get("standards") in (None, [])
    assert (
        "Mandate standards: explicit selection and evidence pending" in result["issues"]
    )
    assert result["piv_conformity"] == "not_assessed"


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("name", None),
        ("edition", None),
        ("adoption_reason", None),
        ("departures", None),
        ("status", "proposed"),
        ("source_ids", []),
        ("locator", None),
    ],
)
def test_unconfirmed_standard_choice_is_incomplete(field: str, value: object) -> None:
    case = case_data()
    case["mandate_details"]["standards"][0][field] = value

    result = build_valuation(case, FIXTURE)

    assert result["mandate_assessment"]["status"] == "partial"
    assert (
        "Standard synthetic-protocol: selection or evidence pending" in result["issues"]
    )
    assert result["methods"][0]["status"] == "ready_for_professional_review"


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("standards", "text"),
        ("standards", [{}]),
        ("expert_identity", "text"),
        (
            "standards",
            [
                {
                    "id": "invented",
                    "name": "invented",
                    "edition": "invented",
                    "adoption_reason": "invented",
                    "departures": "invented",
                    "status": "approved",
                    "source_ids": [],
                    "locator": None,
                }
            ],
        ),
    ],
)
def test_malformed_engagement_schema_rejects_before_source_reads(
    field: str, value: object, tmp_path: Path
) -> None:
    case = case_data()
    case["mandate_details"][field] = value

    with pytest.raises(ValuationError, match="schema"):
        build_valuation(case, tmp_path / "no-source-access")


def test_duplicate_standard_identifier_is_rejected() -> None:
    case = case_data()
    case["mandate_details"]["standards"].append(
        deepcopy(case["mandate_details"]["standards"][0])
    )

    with pytest.raises(ValuationError, match="Duplicate mandate standard"):
        build_valuation(case, FIXTURE)


def test_unresolved_standard_evidence_is_rejected() -> None:
    case = case_data()
    case["mandate_details"]["standards"][0]["source_ids"] = ["absent"]

    with pytest.raises(ValuationError, match="Unresolved mandate standard evidence"):
        build_valuation(case, FIXTURE)


@pytest.mark.parametrize("report_date", ["2026-12-31", "2027-01-15"])
def test_explicit_standard_is_preserved_without_edition_or_conformity_inference(
    report_date: str,
) -> None:
    case = case_data()
    case["mandate_details"]["report_date"]["value"] = report_date

    result = build_valuation(case, FIXTURE)

    assert result["mandate_assessment"]["status"] == "ready_for_professional_review"
    assert (
        result["mandate_assessment"]["details"]["standards"]
        == case["mandate_details"]["standards"]
    )
    assert result["piv_conformity"] == "not_assessed"
    assert result["review_identity"] == "local_attestation_not_authenticated"
    assert (
        result["legal_purpose_qualification"] == "requires_separate_professional_review"
    )


@pytest.mark.parametrize(
    "path",
    [
        ("expert_identity", "value"),
        ("written_mandate", "value"),
        ("remuneration", "value"),
        ("delivery_terms", "value"),
        ("amendments", "value"),
        ("standards", 0, "name"),
        ("standards", 0, "edition"),
        ("standards", 0, "adoption_reason"),
        ("standards", 0, "departures"),
        ("standards", 0, "locator"),
    ],
)
def test_changed_engagement_choice_expires_all_dependent_reviews(path: tuple) -> None:
    case = reviewed_engagement_case()
    parent = case["mandate_details"]
    for key in path[:-1]:
        parent = parent[key]
    parent[path[-1]] = "Changed fictional engagement evidence"

    result = build_valuation(case, FIXTURE)

    assert result["mandate_assessment"]["stale_review"] is True
    assert result["methods"][0]["stale_review"] is True
    assert result["normalizations"][0]["adjustments"][0]["stale_review"] is True
    assert result["claims"][0]["stale_review"] is True
    assert result["conclusion"]["status"] == "draft"
    assert Decimal(
        next(
            row["value"] for row in result["calculations"] if row["id"] == "fcff/equity"
        )
    ) == Decimal("750")


def standard_source_case(tmp_path: Path) -> dict:
    """Bind the selected protocol to its own synthetic source, separate from numbers."""
    case = case_data()
    shutil.copy(FIXTURE / "evidence.txt", tmp_path / "evidence.txt")
    source = tmp_path / "selection.txt"
    source.write_text(
        "Fictional selection of protocol fixture-v1 for software testing only.\n"
    )
    case["sources"].append(
        {
            **case["sources"][0],
            "id": "standard-choice",
            "path": source.name,
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "description": "Fictional standard selection",
        }
    )
    case["mandate_details"]["standards"][0].update(
        source_ids=["standard-choice"], locator="selection.txt: whole note"
    )
    return case


def test_unreviewed_standard_source_prevents_acceptance(tmp_path: Path) -> None:
    case = standard_source_case(tmp_path)
    case["sources"][-1]["status"] = "unverified"

    result = build_valuation(case, tmp_path)

    assert "Mandate source review pending" in result["issues"]
    assert result["mandate_assessment"]["status"] == "partial"
    assert result["methods"][0]["status"] == "partial"


def test_replaced_standard_source_expires_mandate_and_method_reviews(
    tmp_path: Path,
) -> None:
    case = standard_source_case(tmp_path)
    initial = build_valuation(case, tmp_path)
    case["mandate_details"]["review"] = review(
        initial["mandate_assessment"]["dependency_sha256"]
    )
    case["methods"][0]["review"] = review(initial["methods"][0]["dependency_sha256"])
    source = tmp_path / "selection.txt"
    source.write_text(
        "Revised fictional selection evidence; numerical inputs unchanged.\n"
    )
    case["sources"][-1]["sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()

    result = build_valuation(case, tmp_path)

    assert result["mandate_assessment"]["stale_review"] is True
    assert result["methods"][0]["stale_review"] is True
    assert result["calculations"] == initial["calculations"]


def test_standards_export_to_every_format_as_literal_evidence(tmp_path: Path) -> None:
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader

    case = case_data()
    untrusted = '=1+1 <script>alert("x")</script>'
    case["mandate_details"]["standards"][0]["name"] = untrusted
    report = build_valuation(case, FIXTURE)
    output = tmp_path / "exports"

    write_package(report, FIXTURE, output)

    workbook = load_workbook(output / "valuation_workbook.xlsx")
    cells = [
        cell for row in workbook["Incarico"] for cell in row if cell.value == untrusted
    ]
    assert len(cells) == 1
    assert cells[0].data_type == "s"
    assert (
        "&lt;script&gt;alert(&quot;x&quot;)&lt;/script&gt;"
        in (output / "valuation_report.html").read_text()
    )
    assert "fixture-v1" in (output / "valuation_report.md").read_text()
    assert untrusted in "\n".join(
        p.text for p in Document(output / "valuation_report.docx").paragraphs
    )
    assert untrusted in "\n".join(
        p.extract_text() for p in PdfReader(output / "valuation_report.pdf").pages
    )
    assert (
        read_json(output / "mandate.json")["data"]["assessment"]
        == report["mandate_assessment"]
    )
    assert (
        read_json(output / "professional_review.json")["data"]["mandate"]
        == report["mandate_assessment"]
    )


def test_missing_new_mandate_fields_remain_visible_in_exports(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    case = case_data()
    del case["mandate_details"]["remuneration"]
    del case["mandate_details"]["standards"]
    report = build_valuation(case, FIXTURE)
    output = tmp_path / "exports"

    write_package(report, FIXTURE, output)

    rows = list(load_workbook(output / "valuation_workbook.xlsx")["Incarico"].values)
    assert ("Compenso e condizioni", "Da acquisire", "Incompleto", None, None) in rows
    assert ("Standard ed edizione", "Da acquisire", "Incompleto", None, None) in rows
    text = (output / "valuation_report.md").read_text()
    assert "Compenso e condizioni: Da acquisire" in text
    assert "Standard ed edizione: Da acquisire" in text
