from __future__ import annotations

import json
import os
import shlex
import subprocess
import sys
from pathlib import Path

import pytest

from tests.plugins.test_vera_customer_folder_ledger import (
    archive_core,
    archive_script_imports,
    client_case,
)
from tests.plugins.test_vera_studio_archive_component import isolated_mcp_launcher

PIPELINES = (
    "journal-bank-reconciliation",
    "open-item-reconciliation",
    "journal-sampling",
    "vouching",
    "fatture-xml-check",
    "purchase-invoice-review",
    "centrale-rischi-review",
    "financial-report-builder",
    "report-enti-locali",
    "management-control-pack",
    "variance-analysis",
    "agenzia-acquisition",
    "archive-organization",
)
USER_ID = "397d2de9-0ec5-4db8-a01a-b91ae41d13b0"


@pytest.mark.skipif(os.name != "posix", reason="POSIX permission-change behavior")
def test_private_report_database_can_record_without_changing_permissions(
    tmp_path, monkeypatch
):
    from archive_report_store import RunStore

    store = RunStore(tmp_path / "run")
    store.save({"status": "running"})

    def denied(*args, **kwargs):
        raise PermissionError("Host denies permission changes")

    monkeypatch.setattr(Path, "chmod", denied)

    store.event("first-send", {"characters": 12})

    [event] = store.events()
    assert event["event_id"] == "first-send"
    assert event["characters"] == 12
    assert store.state() == {"status": "running"}


@pytest.mark.skipif(os.name != "posix", reason="POSIX permission-change behavior")
def test_nonprivate_report_database_still_requires_permission_repair(
    tmp_path, monkeypatch
):
    from archive_report_store import RunStore

    store = RunStore(tmp_path / "run")
    store.save({"status": "running"})
    store.path.chmod(0o644)
    before = store.path.read_bytes()

    def denied(*args, **kwargs):
        raise PermissionError("Host denies permission changes")

    monkeypatch.setattr(Path, "chmod", denied)

    with pytest.raises(PermissionError, match="permission changes"):
        store.save({"status": "complete"})

    assert store.path.read_bytes() == before


@pytest.mark.parametrize(
    "operation",
    [
        "vera_workspace_transformation_author_context",
        "vera_workspace_transformation_author_stage",
        "vera_workspace_practice_context",
    ],
)
def test_non_archive_model_response_does_not_require_an_archive_run(
    monkeypatch, operation
):
    import native_workspace

    def no_archive():
        raise PermissionError("No Archive workspace is connected")

    monkeypatch.setattr(native_workspace, "configuration", no_archive)
    args = {"operation": operation, "work_ref": "separately-authorized-work"}

    binding = native_workspace.bind_model_output(args)

    assert binding is None


def test_reportable_model_response_still_requires_its_authorized_archive(
    monkeypatch,
):
    import native_workspace

    def no_archive():
        raise PermissionError("No Archive workspace is connected")

    monkeypatch.setattr(native_workspace, "configuration", no_archive)
    args = {
        "operation": "vera_workspace_bank_explain",
        "work_ref": "unconnected-bank-run",
    }

    with pytest.raises(PermissionError, match="No Archive workspace"):
        native_workspace.bind_model_output(args)


@pytest.fixture(autouse=True)
def reporting_runtime(tmp_path, monkeypatch, archive_script_imports):
    import archive_report_receipt
    import archive_reporting

    original_stamp = archive_report_receipt.stamp_report

    monkeypatch.setenv("VERA_REPORTING_STATE_DIR", str(tmp_path / "reporting-state"))
    monkeypatch.setenv("CODEX_THREAD_ID", "test-session")
    monkeypatch.setenv(
        "VERA_STUDIO_ARCHIVE_SESSION_ID", str(tmp_path / "private-state")
    )
    monkeypatch.setattr(archive_reporting, "_identity", lambda: USER_ID)
    monkeypatch.setattr(
        archive_report_receipt,
        "stamp_report",
        lambda path: {"status": "pending", "reason": "offline_test"},
    )
    yield original_stamp


def start(core, case, workflow="variance-analysis"):
    workflow = {
        "vouching": "check-entries",
        "purchase-invoice-review": "passive-invoice-audit",
        "financial-report-builder": "report-builder",
    }.get(workflow, workflow)
    prepared = core.prepare_studio_client_workflow(
        case.engagement_id,
        workflow,
        input_ids=[case.imported["input_id"]],
        state_dir=case.state_dir,
    )
    run_id = prepared["run"]["run_id"]
    core.start_studio_client_workflow(
        case.client_id, case.engagement_id, run_id, state_dir=case.state_dir
    )
    return core.ledger.load_run(case.client_root, case.engagement_id, run_id)


def read_sent(loaded, **selection):
    from archive_reporting import read_artifact

    path = Path(loaded["context"]["input_bindings"][0]["path"])
    return read_artifact(loaded, path, **selection)


def finish(core, case, loaded):
    output = Path(loaded["output_dir"])
    (output / "result.txt").write_text("Local professional result\n")
    core.finalize_studio_client_workflow(
        case.client_id,
        case.engagement_id,
        loaded["run"]["run_id"],
        [
            {
                "artifact_id": "result",
                "path": "result.txt",
                "purpose": "Professional result",
                "audience": "review",
                "media_type": "text/plain",
            }
        ],
        state_dir=case.state_dir,
    )
    core.complete_studio_client_workflow(
        case.client_id,
        case.engagement_id,
        loaded["run"]["run_id"],
        state_dir=case.state_dir,
    )
    return json.loads((output / "model_data_report.json").read_text(encoding="utf-8"))


@pytest.mark.parametrize("workflow", PIPELINES)
def test_selected_archive_run_always_creates_report_at_start_and_terminal_write(
    archive_core, client_case, workflow
):
    loaded = start(archive_core, client_case, workflow)
    output = Path(loaded["output_dir"])
    initial = json.loads(
        (output / "model_data_report.json").read_text(encoding="utf-8")
    )
    assert initial["outcome"] == "running"
    assert initial["user_id"] == USER_ID
    assert not (output / "model_data_report_input.json").exists()

    report = finish(archive_core, client_case, loaded)

    assert report["outcome"] == "completed"
    assert report["run_id"] == loaded["run"]["run_id"]
    assert report["collection_status"] == "complete"
    assert report["sources"][0]["not_sent_quantity"] == 1
    archive_core.ledger.validate_run_artifacts(
        client_case.client_root, client_case.engagement_id, report["run_id"]
    )


def test_full_exposure_still_creates_report_and_zero_local_remainder(
    archive_core, client_case
):
    loaded = start(archive_core, client_case)
    read_sent(loaded)

    report = finish(archive_core, client_case, loaded)

    assert report["collection_status"] == "complete"
    assert report["sources"][0]["sent_quantity"] == 1
    assert report["sources"][0]["not_sent_quantity"] == 0


@pytest.mark.parametrize("terminal", ["failed", "cancelled"])
def test_failure_and_cancellation_retain_an_incomplete_report(
    archive_core, client_case, terminal
):
    loaded = start(archive_core, client_case)
    action = (
        archive_core.ledger.fail_run
        if terminal == "failed"
        else archive_core.ledger.cancel_run
    )
    args = [client_case.client_root, client_case.engagement_id, loaded["run"]["run_id"]]
    if terminal == "failed":
        args.append("Synthetic interruption")

    action(*args)

    report = json.loads(
        (Path(loaded["output_dir"]) / "model_data_report.json").read_text(
            encoding="utf-8"
        )
    )
    assert report["outcome"] == terminal
    assert report["finished_at"] is not None


@pytest.mark.parametrize(
    "workflow", ["financial-analysis", "new-client", "contenzioso-tributario"]
)
def test_unselected_archive_runs_do_not_create_reports(
    archive_core, client_case, workflow
):
    loaded = start(archive_core, client_case, workflow)

    assert not (Path(loaded["output_dir"]) / "model_data_report.json").exists()


def csv_case(core, case, tmp_path, line_ending="\n"):
    source = tmp_path / "population.csv"
    text = "id,amount\n" + "".join(f"{index},{index * 2}\n" for index in range(10000))
    source.write_bytes(text.replace("\n", line_ending).encode("utf-8"))
    imported = core.import_studio_client_document(
        case.client_id,
        source,
        "journal",
        engagement_id=case.engagement_id,
        state_dir=case.state_dir,
    )
    case.imported = imported
    return start(core, case)


def test_terminal_report_retries_real_authenticated_proof_without_changing_work(
    archive_core, client_case, tmp_path, monkeypatch, reporting_runtime
):
    import archive_report_receipt
    import archive_reporting
    import operational_incidents
    from archive_report_store import Registry

    from tests.modules.run_receipts.test_archive_binding import archive_client

    client, owner, _, receipts, signer = archive_client(tmp_path, monkeypatch)
    monkeypatch.setattr(
        archive_reporting, "_identity", lambda: str(owner.installation_id)
    )
    monkeypatch.setattr(archive_report_receipt, "stamp_report", reporting_runtime)
    monkeypatch.setattr(
        operational_incidents,
        "reporting_credentials",
        lambda: owner.model_dump(mode="json"),
    )
    requests = []

    def transport(operation, envelope):
        assert operation == "archive-stamp"
        requests.append(envelope["receipt_request"])
        response = client.post(
            "/api/operational-incidents/archive-stamp", json=envelope
        )
        response.raise_for_status()
        return response.json()

    def offline(operation, envelope):
        raise OSError("Synthetic offline connection")

    monkeypatch.setattr(operational_incidents, "workflow_transport", offline)
    loaded = start(archive_core, client_case)
    read_sent(loaded)
    report = finish(archive_core, client_case, loaded)
    root = Path(loaded["run_root"])
    report_path = root / "outputs/model_data_report.json"
    before = report_path.read_bytes()
    professional_before = (root / "outputs/result.txt").read_bytes()
    assert (
        json.loads((root / "model_data_reporting.json").read_text(encoding="utf-8"))[
            "server_receipt"
        ]["status"]
        == "pending"
    )
    monkeypatch.setattr(operational_incidents, "workflow_transport", transport)

    # A normal Archive operation retries automatically; no Stop hook or agent
    # calling the receipt transport is required.
    archive_core._selected_ledger_root(
        client_case.client_id,
        client_case.engagement_id,
        state_dir=client_case.state_dir,
    )

    receipt = json.loads(
        (root / "outputs/model_data_receipt.json").read_text(encoding="utf-8")
    )
    assert receipt["user_id"] == report["user_id"]
    assert receipt["archive_run_id"] == report["run_id"]
    assert report_path.read_bytes() == before
    assert (root / "outputs/result.txt").read_bytes() == professional_before
    assert set(requests[0]) == {
        "schema_version",
        "product",
        "receipt_id",
        "plugin_version",
        "report_sha256",
        "user_id",
        "archive_run_id",
        "report_revision",
    }
    assert Registry().pending_receipts() == []
    record = receipts.get(receipt["receipt_id"])
    assert signer.verify(record.signed_payload, record.signature)
    archive_core.ledger.validate_run_artifacts(
        client_case.client_root, client_case.engagement_id, report["run_id"]
    )

    import io
    from types import SimpleNamespace
    from urllib.parse import urlsplit

    import notarized_run_receipt

    standalone = tmp_path / "copied-report.json"
    standalone.write_bytes(before)

    def open_verified(url, timeout):
        parsed = urlsplit(url.full_url)
        response = client.get(parsed.path + "?" + parsed.query)
        response.raise_for_status()
        return io.BytesIO(response.content)

    monkeypatch.setattr(
        archive_report_receipt.urllib.request,
        "build_opener",
        lambda *args: SimpleNamespace(open=open_verified),
    )
    monkeypatch.setattr(
        notarized_run_receipt,
        "_plugin_version",
        lambda root: pytest.fail(
            "Verification must not depend on the current plugin version"
        ),
    )

    verified = archive_report_receipt.verify_report(
        standalone, user_id=report["user_id"]
    )
    assert verified["report_sha256"] == receipt["report_sha256"]
    with pytest.raises(ValueError, match="another Vera userID"):
        archive_report_receipt.verify_report(standalone, user_id=USER_ID)
    standalone.write_bytes(before + b"\n")
    with pytest.raises(ValueError, match="canonical"):
        archive_report_receipt.verify_report(standalone, user_id=report["user_id"])


@pytest.fixture
def signed_archive_report(
    archive_core, client_case, tmp_path, monkeypatch, reporting_runtime
):
    import io
    from types import SimpleNamespace
    from urllib.parse import urlsplit

    import archive_report_receipt
    import archive_reporting
    import operational_incidents

    from tests.modules.run_receipts.test_archive_binding import archive_client

    client, owner, _, _, _ = archive_client(tmp_path, monkeypatch)
    monkeypatch.setattr(
        archive_reporting, "_identity", lambda: str(owner.installation_id)
    )
    monkeypatch.setattr(
        operational_incidents,
        "reporting_credentials",
        lambda: owner.model_dump(mode="json"),
    )

    def stamp(operation, envelope):
        assert operation == "archive-stamp"
        response = client.post(
            "/api/operational-incidents/archive-stamp", json=envelope
        )
        response.raise_for_status()
        return response.json()

    monkeypatch.setattr(operational_incidents, "workflow_transport", stamp)
    monkeypatch.setattr(archive_report_receipt, "stamp_report", reporting_runtime)
    loaded = start(archive_core, client_case)
    finish(archive_core, client_case, loaded)
    output = Path(loaded["output_dir"])
    report = output / "model_data_report.json"
    receipt = output / "model_data_receipt.json"
    requests = []

    def open_verified(request, *, timeout):
        requests.append(request)
        parsed = urlsplit(request.full_url)
        response = client.get(parsed.path + "?" + parsed.query)
        response.raise_for_status()
        return io.BytesIO(response.content)

    return SimpleNamespace(
        report=report,
        receipt=receipt,
        owner=str(owner.installation_id),
        opener=open_verified,
        requests=requests,
        original={
            path: path.read_bytes()
            for path in Path(loaded["run_root"]).rglob("*")
            if path.is_file()
        },
    )


def test_general_verifier_accepts_archive_receipt_without_rewriting_run(
    signed_archive_report,
):
    import notarized_run_receipt

    case = signed_archive_report
    result = notarized_run_receipt.verify_model_data_receipt(
        case.receipt, report_path=case.report, opener=case.opener
    )

    assert result["status"] == "valid"
    assert result["user_id"] == case.owner
    assert result["archive_run_id"] == json.loads(case.report.read_bytes())["run_id"]
    assert result["report_revision"] == 1
    assert len(case.requests) == 1
    assert {path: path.read_bytes() for path in case.original} == case.original


@pytest.mark.parametrize("explicit_owner", [False, True])
def test_general_verify_cli_accepts_archive_receipt(
    signed_archive_report, monkeypatch, capsys, explicit_owner
):
    from types import SimpleNamespace

    import archive_report_receipt
    import notarized_run_receipt

    case = signed_archive_report
    monkeypatch.setattr(
        archive_report_receipt.urllib.request,
        "build_opener",
        lambda *args: SimpleNamespace(open=case.opener),
    )

    status = notarized_run_receipt.main(
        [
            "verify",
            "--report",
            str(case.report),
            "--receipt",
            str(case.receipt),
        ]
        + (["--user-id", case.owner] if explicit_owner else [])
    )

    assert status == 0
    assert json.loads(capsys.readouterr().out)["user_id"] == case.owner


def test_general_archive_verifier_uses_current_user_before_network(
    signed_archive_report, monkeypatch
):
    import notarized_run_receipt
    import operational_incidents

    case = signed_archive_report
    monkeypatch.setattr(
        operational_incidents,
        "reporting_credentials",
        lambda: {"installation_id": USER_ID},
    )

    with pytest.raises(
        notarized_run_receipt.NotarizedRunReceiptError, match="another Vera userID"
    ):
        notarized_run_receipt.verify_model_data_receipt(
            case.receipt, report_path=case.report, opener=case.opener
        )

    assert case.requests == []


@pytest.mark.parametrize("product", ["clara", "lucia"])
def test_other_product_verifier_cannot_claim_vera_archive_receipt(
    signed_archive_report, monkeypatch, product
):
    import notarized_run_receipt

    case = signed_archive_report
    monkeypatch.setattr(notarized_run_receipt, "PRODUCT", product)

    with pytest.raises(notarized_run_receipt.NotarizedRunReceiptError):
        notarized_run_receipt.verify_model_data_receipt(
            case.receipt, report_path=case.report, opener=case.opener
        )

    assert case.requests == []


@pytest.mark.parametrize(
    "field,value",
    [
        ("user_id", USER_ID),
        ("archive_run_id", "run_" + "e" * 24),
        ("report_revision", 2),
        ("report_sha256", "f" * 64),
        ("signature", "alteredSignature"),
    ],
)
def test_general_archive_verifier_rejects_changed_local_proof(
    signed_archive_report, field, value
):
    import notarized_run_receipt

    case = signed_archive_report
    altered = json.loads(case.receipt.read_bytes())
    altered[field] = value
    case.receipt.write_text(json.dumps(altered))

    with pytest.raises(
        notarized_run_receipt.NotarizedRunReceiptError, match="does not match"
    ):
        notarized_run_receipt.verify_model_data_receipt(
            case.receipt, report_path=case.report, opener=case.opener
        )

    assert json.loads(case.receipt.read_bytes()) == altered
    assert case.report.read_bytes() == case.original[case.report]


def test_late_receipt_cannot_reseal_modified_professional_output(
    archive_core, client_case
):
    from archive_report_receipt import retry_pending
    from archive_report_store import Registry

    loaded = start(archive_core, client_case)
    finish(archive_core, client_case, loaded)
    (Path(loaded["output_dir"]) / "result.txt").write_text(
        "Altered professional result"
    )

    with pytest.raises(ValueError, match="changed|match"):
        retry_pending(Registry())


def test_retry_does_not_recreate_a_removed_run(archive_core, client_case):
    import shutil

    from archive_report_receipt import retry_pending
    from archive_report_store import Registry

    loaded = start(archive_core, client_case)
    finish(archive_core, client_case, loaded)
    root = Path(loaded["run_root"])
    shutil.rmtree(root)

    retry_pending(Registry())

    assert not root.exists()
    assert Registry().pending_receipts() == []


def test_invalid_pending_proof_does_not_block_other_archive_work(
    archive_core, client_case
):
    loaded = start(archive_core, client_case)
    finish(archive_core, client_case, loaded)
    output = Path(loaded["output_dir"])
    manifest = Path(loaded["run_root"]) / "artifact_manifest.json"
    sealed = manifest.read_bytes()
    report = (output / "model_data_report.json").read_bytes()
    (output / "result.txt").write_text("Changed after sealing")

    result = start(archive_core, client_case, "report-builder")

    assert result["run"]["workflow_id"] == "report-builder"
    assert result["run"]["status"] == "running"
    assert manifest.read_bytes() == sealed
    assert (output / "model_data_report.json").read_bytes() == report
    assert (output / "result.txt").read_text() == "Changed after sealing"


def test_reporting_sources_and_pending_proof_move_with_the_customer_folder(
    archive_core, client_case, monkeypatch
):
    from archive_report_store import RunStore, registry_for_run

    monkeypatch.delenv("VERA_REPORTING_STATE_DIR")
    loaded = start(archive_core, client_case)
    original = client_case.client_root
    renamed = original.with_name(original.name + " Renamed")
    original.rename(renamed)
    moved = archive_core.ledger.load_run(
        renamed, client_case.engagement_id, loaded["run"]["run_id"]
    )

    packet = read_sent(moved)
    output = Path(moved["output_dir"])
    (output / "result.txt").write_text("Local professional result\n")
    archive_core.ledger.finalize_run(
        renamed,
        client_case.engagement_id,
        moved["run"]["run_id"],
        [
            {
                "artifact_id": "result",
                "path": "result.txt",
                "purpose": "Result",
                "audience": "review",
                "media_type": "text/plain",
            }
        ],
    )
    archive_core.ledger.complete_run(
        renamed, client_case.engagement_id, moved["run"]["run_id"]
    )

    report = json.loads((output / "model_data_report.json").read_text())
    assert packet["data"]["rows"]
    assert report["collection_status"] == "complete", report["errors"]
    assert report["sources"][0]["not_sent_quantity"] == 0
    assert registry_for_run(Path(moved["run_root"])).pending_receipts() == [
        Path(moved["run_root"])
    ]
    assert Path(RunStore(Path(moved["run_root"])).sources()[0]["path"]).is_file()
    assert not original.exists()


def test_read_rejects_empty_column_selection(archive_core, client_case):
    from archive_reporting import read_artifact

    loaded = start(archive_core, client_case)

    with pytest.raises(ValueError, match="at least one"):
        read_artifact(
            loaded, Path(loaded["context"]["input_bindings"][0]["path"]), columns=[]
        )


@pytest.mark.parametrize("relative", [True, False])
def test_run_artifact_read_uses_selected_run_independent_of_process_directory(
    archive_core, client_case, tmp_path, monkeypatch, relative
):
    loaded = start(archive_core, client_case)
    artifact = Path(loaded["output_dir"]) / "inspection.json"
    artifact.write_text('{"candidate_rows": 3}')
    monkeypatch.chdir(tmp_path)
    selected_path = "outputs/inspection.json" if relative else str(artifact)

    packet = archive_core.read_studio_run_artifact(
        client_case.client_id,
        client_case.engagement_id,
        loaded["run"]["run_id"],
        selected_path,
        state_dir=client_case.state_dir,
    )

    assert "candidate_rows" in json.dumps(packet["data"])
    assert packet["vera_reporting"]["run_id"] == loaded["run"]["run_id"]


def test_relative_run_artifact_read_cannot_escape_to_another_run(
    archive_core, client_case
):
    first = start(archive_core, client_case)
    other = start(archive_core, client_case, "report-builder")
    artifact = Path(other["output_dir"]) / "inspection.json"
    artifact.write_text('{"other_case": true}')

    with pytest.raises(archive_core.ArchiveError, match="bound input"):
        archive_core.read_studio_run_artifact(
            client_case.client_id,
            client_case.engagement_id,
            first["run"]["run_id"],
            f"../{other['run']['run_id']}/outputs/inspection.json",
            state_dir=client_case.state_dir,
        )


def test_local_only_report_makes_no_transport_call(
    archive_core, client_case, monkeypatch, reporting_runtime
):
    import archive_report_receipt
    import operational_incidents

    monkeypatch.setattr(archive_report_receipt, "stamp_report", reporting_runtime)
    monkeypatch.setattr(
        operational_incidents,
        "workflow_transport",
        lambda *args: pytest.fail("Local-only run attempted transport"),
    )
    loaded = start(archive_core, client_case)
    (Path(loaded["output_dir"]) / "model_data_report_options.json").write_text(
        '{"local_only":true}'
    )

    report = finish(archive_core, client_case, loaded)

    status = json.loads(
        (Path(loaded["run_root"]) / "model_data_reporting.json").read_text(
            encoding="utf-8"
        )
    )
    assert report["collection_closed"] is True
    assert status["server_receipt"] == {
        "status": "not_requested",
        "reason": "local_only",
    }


@pytest.mark.parametrize("delimiter", [",", ";", "\t", "|"])
def test_csv_selection_exposes_the_delimiter_used_for_exact_source_rows(
    tmp_path, delimiter
):
    import csv

    from archive_report_measurements import measure_source, read_selection

    path = tmp_path / "source.csv"
    with path.open("w", newline="") as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        writer.writerows(
            [
                ["ID", "Description", "Amount"],
                ["F001", "first, quoted; description", "1200.00"],
                ["F002", "second, quoted; description", "800.00"],
            ]
        )
    source = measure_source(path, source_id="invoices")

    data, evidence = read_selection(source, offset=1, limit=1, columns=[0, 2])

    assert data["source_format"] == {
        "kind": "csv",
        "delimiter": delimiter,
        "delimiter_basis": "sniffed",
    }
    assert data["columns"] == ["ID", "Amount"]
    assert data["rows"] == [["F002", "800.00"]]
    assert evidence["start"] == 1
    assert evidence["stop"] == 2
    assert evidence["columns"] == [0, 2]


def test_single_column_preview_identifies_default_delimiter_as_unconfirmed(tmp_path):
    from archive_report_measurements import measure_source, read_selection

    path = tmp_path / "source.csv"
    path.write_text("ID\nF001\nF002\n")
    source = measure_source(path, source_id="ids")

    data, _ = read_selection(source)

    assert data["source_format"] == {
        "kind": "csv",
        "delimiter": ",",
        "delimiter_basis": "default",
    }
    assert data["rows"] == [["F001"], ["F002"]]


def test_xlsx_selection_preserves_physical_rows_columns_and_formula(tmp_path):
    from archive_report_measurements import measure_source, read_selection
    from openpyxl import Workbook

    path = tmp_path / "source.xlsx"
    book = Workbook()
    book.active.append(["Account", "Value"])
    book.active.append(["A", "=2+3"])
    book.create_sheet("Second").append(["B", 7])
    book.save(path)
    source = measure_source(path, source_id="book")

    data, evidence = read_selection(source, offset=1, limit=1, columns=[1])

    assert source["quantity"] == 3
    assert source["cells"] == 6
    assert data["rows"] == [["=2+3"]]
    assert evidence["start"] == 1
    assert evidence["stop"] == 2
    assert evidence["columns"] == [1]


def test_scanned_or_empty_pdf_cannot_claim_measured_exposure(tmp_path):
    from archive_report_measurements import measure_source, read_selection
    from pypdf import PdfWriter

    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(path)
    source = measure_source(path, source_id="scan")

    with pytest.raises(ValueError, match="no complete measurement"):
        read_selection(source)

    assert source["measurement_error"] is not None


def test_identity_failure_keeps_local_report(archive_core, client_case, monkeypatch):
    import archive_reporting

    def unavailable():
        raise OSError("No identity storage")

    monkeypatch.setattr(archive_reporting, "_identity", unavailable)
    loaded = start(archive_core, client_case)

    report = finish(archive_core, client_case, loaded)

    assert report["user_id"] is None
    assert "vera_machine_identity_unavailable" in report["errors"]
    assert report["collection_status"] == "incomplete"


@pytest.mark.parametrize("changed", [False, True])
def test_acquisition_registers_only_verified_originals_in_its_archive_run(
    archive_core, client_case, tmp_path, monkeypatch, changed
):
    from archive_report_store import RunStore

    monkeypatch.syspath_prepend(
        str(Path(__file__).resolve().parents[2] / "plugins/browser-automation/scripts")
    )
    from ade_acquisition import archive_runs, contracts, storage

    monkeypatch.setattr(archive_runs, "_archive", lambda: archive_core)
    loaded = start(archive_core, client_case, "agenzia-acquisition")
    worker = tmp_path / "acquisition-worker"
    worker.mkdir()
    (worker / "archive-runs.json").write_text(
        json.dumps(
            {
                "01234567890": {
                    "context_path": str(Path(loaded["run_root"]) / "context.json")
                }
            }
        )
    )
    retained = storage.Store(tmp_path / "retained-originals")
    client = contracts.Client.parse(
        {
            "name": "Synthetic client",
            "tax_code": "01234567890",
            "vat_number": "01234567890",
        }
    )
    record = retained.retain(
        client,
        "ricevute",
        "synthetic-download",
        2026,
        1,
        "original.xml",
        b"<invoice>private original</invoice>",
        {},
    )
    if changed:
        retained.safe_path(record["artifacts"][0]["path"]).write_bytes(b"changed")

    archive_runs.record_acquired_event(
        worker, retained, {"client": client.vat_number, "record": record}
    )

    sources = RunStore(Path(loaded["run_root"])).sources()
    assert len(sources) == (1 if changed else 2)
    assert (
        any(item["source_id"].startswith("acquired_") for item in sources)
        is not changed
    )


def test_mcp_entrypoint_records_its_exact_return_without_a_host_observer(
    archive_core, client_case, tmp_path, isolated_mcp_launcher
):

    from tests.plugins.test_vera_studio_archive_component import _mcp_tool

    loaded = csv_case(archive_core, client_case, tmp_path)
    arguments = {
        "client_id": client_case.client_id,
        "engagement_id": client_case.engagement_id,
        "run_id": loaded["run"]["run_id"],
        "path": loaded["context"]["input_bindings"][0]["path"],
        "offset": 10,
        "limit": 2,
        "columns": [1],
    }

    archive_core._release_state_leases()
    response = _mcp_tool(
        "read_studio_run_artifact", arguments, state_dir=client_case.state_dir
    )

    assert response["isError"] is False, response
    packet = json.loads(response["content"][0]["text"])
    assert packet["data"]["rows"] == [["20"], ["22"]]
    report = finish(archive_core, client_case, loaded)
    assert report["sources"][0]["sent_quantity"] == 2
    assert report["sources"][0]["sent_cells"] == 2


def test_vera_return_records_exact_rows_without_host_hooks(
    archive_core, client_case, tmp_path
):
    from archive_reporting import read_artifact

    loaded = csv_case(archive_core, client_case, tmp_path)
    path = Path(loaded["context"]["input_bindings"][0]["path"])
    packet = read_artifact(loaded, path, offset=10, limit=2, columns=[1])

    report = finish(archive_core, client_case, loaded)

    assert packet["data"]["rows"] == [["20"], ["22"]]
    assert report["measurement_scope"] == "vera_sent_data"
    assert report["collection_status"] == "complete"
    assert report["sources"][0]["sent_quantity"] == 2
    assert report["sources"][0]["sent_cells"] == 2
    assert report["sources"][0]["not_sent_quantity"] == 9998
    assert report["collection_closed"] is True


@pytest.mark.parametrize("line_ending", ["\n", "\r\n"])
def test_repeated_and_full_vera_sends_count_distinct_source_rows(
    archive_core, client_case, tmp_path, line_ending
):
    loaded = csv_case(archive_core, client_case, tmp_path, line_ending)
    read_sent(loaded, limit=100)
    read_sent(loaded, limit=100)
    read_sent(loaded)

    report = finish(archive_core, client_case, loaded)

    assert report["sources"][0]["sent_quantity"] == 10000
    assert report["sources"][0]["sent_cells"] == 20000
    assert report["sources"][0]["not_sent_quantity"] == 0
    assert report["collection_status"] == "complete"


def test_unmapped_vera_send_preserves_payload_but_cannot_claim_source_remainder(
    archive_core, client_case
):
    from archive_reporting import record_send

    loaded = start(archive_core, client_case)

    record_send(
        Path(loaded["run_root"]), "Some extra case-derived content", channel="worker"
    )
    report = finish(archive_core, client_case, loaded)

    assert report["collection_status"] == "incomplete"
    assert "sent_data_has_no_source_coordinates" in report["errors"]
    assert report["sources"][0]["not_sent_quantity"] is None
    assert report["sent_characters"] > 0


@pytest.mark.parametrize("extra_field", [False, True])
def test_complete_producer_envelope_cannot_cover_added_content(
    archive_core, client_case, extra_field
):
    from archive_projection import register_value
    from archive_reporting import record_send, snapshot_run

    loaded = start(archive_core, client_case)
    payload = {
        "content": [{"type": "text", "text": "Prepared 2 rows"}],
        "structuredContent": {"count": 2},
        "isError": False,
    }
    register_value("variance-analysis", Path(loaded["output_dir"]), payload, [])
    if extra_field:
        payload["structuredContent"]["private_value"] = "Unregistered source value"
    record_send(Path(loaded["run_root"]), payload, channel="test_mcp")
    report = json.loads(Path(snapshot_run(loaded)["report_path"]).read_text())

    assert report["collection_status"] == ("incomplete" if extra_field else "complete")


@pytest.mark.parametrize("unregistered_line", ["", "Private case value: 123\n"])
def test_cli_combined_status_preserves_unknown_lines_as_unmapped(
    archive_core, client_case, unregistered_line
):
    from archive_projection import register_metadata
    from archive_report_store import RunStore, digest
    from archive_reporting import record_send, snapshot_run

    loaded = start(archive_core, client_case, "journal-bank-reconciliation")
    output = Path(loaded["output_dir"])
    first = "INFO: bank_rows=3\n"
    second = "INFO: journal_rows=2\n"
    register_metadata("journal-bank-reconciliation", output, first)
    register_metadata("journal-bank-reconciliation", output, second)
    payload = {"stdout": "", "stderr": first + second + unregistered_line}

    record_send(
        Path(loaded["run_root"]), payload, channel="cli:journal-bank-reconciliation"
    )

    report = json.loads(Path(snapshot_run(loaded)["report_path"]).read_text())
    sent = [
        event
        for event in RunStore(Path(loaded["run_root"])).events()
        if event["kind"] == "vera_send"
    ][-1]
    assert sent["response_sha256"] == digest(payload)
    assert sent["selections"] == []
    assert sent["unmapped"] is bool(unregistered_line)
    assert report["collection_status"] == (
        "incomplete" if unregistered_line else "complete"
    )


@pytest.mark.parametrize("limit", [None, 1])
@pytest.mark.parametrize("source_bearing", [False, True])
def test_partial_text_is_measured_only_for_source_free_producer_metadata(
    table_projection_case, limit, source_bearing
):
    from archive_projection import register_row_projection
    from archive_report_store import file_digest
    from archive_reporting import read_artifact, snapshot_run

    loaded = table_projection_case
    artifact = Path(loaded["output_dir"]) / "review_notes.md"
    content = "# Review\nTwo matches; one residual.\n"
    artifact.write_text(content, encoding="utf-8")
    register_row_projection(
        "journal-bank-reconciliation",
        artifact,
        content,
        [{"record_id": "first", "fields": ["amount"]}] if source_bearing else [],
        artifact_sha256=file_digest(artifact),
    )

    read_artifact(loaded, artifact, limit=limit)

    report = json.loads(Path(snapshot_run(loaded)["report_path"]).read_text())
    complete = not source_bearing or limit is None
    assert report["collection_status"] == ("complete" if complete else "incomplete")
    assert next(source for source in report["sources"] if source["role"] == "input")[
        "sent_cells"
    ] == int(source_bearing and complete)


@pytest.mark.parametrize(
    "changed", [None, "file_bytes", "producer_review", "expected_json"]
)
def test_saved_artifact_projection_requires_exact_writer_and_review_digests(
    table_projection_case, changed
):
    from archive_projection import register_artifact_projection, register_private_value
    from archive_report_store import digest, file_digest
    from archive_reporting import read_artifact, snapshot_run

    loaded = table_projection_case
    output = Path(loaded["output_dir"])
    artifact = output / "decisions.json"
    review = {"items": [{"id": "first", "title": "7"}]}
    value = {"title": "7"}
    artifact.write_text(json.dumps(value))
    artifact_sha256 = file_digest(artifact)
    register_private_value("journal-bank-reconciliation", output, review)
    if changed == "file_bytes":
        artifact.write_text(json.dumps({"title": "Unregistered replacement"}))
    elif changed == "producer_review":
        review["items"][0]["title"] = "Edited review"
    elif changed == "expected_json":
        value["title"] = "Different writer output"

    register_artifact_projection(
        "journal-bank-reconciliation",
        artifact,
        data_sha256=digest(value),
        artifact_sha256=artifact_sha256,
        records=[{"record_id": "first", "fields": ["amount"]}],
        private_sha256=digest(review),
    )
    read_artifact(loaded, artifact)

    report = json.loads(Path(snapshot_run(loaded)["report_path"]).read_text())
    assert report["collection_status"] == (
        "complete" if changed is None else "incomplete"
    )
    assert next(row for row in report["sources"] if row["role"] == "input")[
        "sent_cells"
    ] == (1 if changed is None else 0)


@pytest.fixture
def table_projection_case(archive_core, client_case, tmp_path, request):
    from archive_projection import register_rows
    from archive_report_store import RunStore

    source = tmp_path / "table.csv"
    source.write_text("amount,private_note\n7,PRIVATE_ONE\n7,PRIVATE_TWO\n")
    sheet_name = None
    if getattr(request, "param", None) == "xlsx":
        from openpyxl import Workbook

        source = source.with_suffix(".xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet_name = sheet.title
        sheet.append(["amount", "private_note"])
        sheet.append([7, "PRIVATE_ONE"])
        sheet.append([7, "PRIVATE_TWO"])
        workbook.save(source)
        workbook.close()
    client_case.imported = archive_core.import_studio_client_document(
        client_case.client_id,
        source,
        "journal",
        engagement_id=client_case.engagement_id,
        state_dir=client_case.state_dir,
    )
    loaded = start(archive_core, client_case, "journal-bank-reconciliation")
    [measured] = RunStore(Path(loaded["run_root"])).sources()
    register_rows(
        "journal-bank-reconciliation",
        Path(loaded["output_dir"]),
        {
            record_id: {
                "source": {
                    "source_path": measured["path"],
                    "source_sha256": measured["sha256"],
                    "source_column_count": 2,
                    "sheet_name": sheet_name,
                },
                "physical_row": physical_row,
                "fields": {"amount": [0]},
            }
            for record_id, physical_row in (("first", 2), ("second", 3))
        },
    )
    return loaded


@pytest.mark.parametrize("source_bearing", [False, True])
@pytest.mark.parametrize("invalidated", [None, "parent", "revision"])
def test_metadata_edit_preserves_only_proven_parent_source_operands(
    table_projection_case, source_bearing, invalidated
):
    from archive_projection import register_metadata_revision, register_row_projection
    from archive_report_store import file_digest
    from archive_reporting import read_artifact, snapshot_run

    loaded = table_projection_case
    artifact = Path(loaded["output_dir"]) / "status.json"
    value = {"sample": 7} if source_bearing else {"count": 2}
    artifact.write_text(json.dumps(value))
    register_row_projection(
        "journal-bank-reconciliation",
        artifact,
        value,
        [{"record_id": "first", "fields": ["amount"]}] if source_bearing else [],
        artifact_sha256=file_digest(artifact),
    )
    if invalidated == "parent":
        value["unknown_note"] = "Unregistered source content"
        artifact.write_text(json.dumps(value))
    previous_sha256 = file_digest(artifact)
    value["status"] = "ready"
    artifact.write_text(json.dumps(value))

    register_metadata_revision(
        "journal-bank-reconciliation",
        artifact,
        previous_sha256,
        expected_sha256=(
            "0" * 64 if invalidated == "revision" else file_digest(artifact)
        ),
    )
    read_artifact(loaded, artifact)

    report = json.loads(Path(snapshot_run(loaded)["report_path"]).read_text())
    assert report["collection_status"] == ("incomplete" if invalidated else "complete")
    assert next(row for row in report["sources"] if row["role"] == "input")[
        "sent_cells"
    ] == int(source_bearing and not invalidated)


@pytest.mark.parametrize(
    "columns,complete,sent_cells", [([0], True, 1), ([1], False, 0), ([0, 1], False, 1)]
)
def test_partial_table_provenance_keeps_unknown_cells_explicit_when_selected(
    table_projection_case, columns, complete, sent_cells
):
    from archive_projection import register_table_projection
    from archive_reporting import read_artifact, snapshot_run

    loaded = table_projection_case
    artifact = Path(loaded["output_dir"]) / "partial.csv"
    artifact.write_text("amount,unmapped_note\n7,UNMAPPED_VALUE\n")
    register_table_projection(
        "journal-bank-reconciliation",
        artifact,
        [[[{"record_id": "first", "fields": ["amount"]}], None]],
    )

    read_artifact(loaded, artifact, columns=columns)

    measured = json.loads(Path(snapshot_run(loaded)["report_path"]).read_text())
    assert measured["collection_status"] == ("complete" if complete else "incomplete")
    [source] = [row for row in measured["sources"] if row["role"] == "input"]
    assert source["sent_cells"] == sent_cells


@pytest.mark.parametrize("first_to_read", [True, False])
def test_identical_table_bytes_keep_each_producers_source_row(
    table_projection_case, first_to_read
):
    from archive_projection import register_table_projection
    from archive_report_store import RunStore
    from archive_reporting import read_artifact, snapshot_run

    loaded = table_projection_case
    first = Path(loaded["output_dir"]) / "first.csv"
    second = first.with_name("second.csv")
    first.write_text("amount\n7\n")
    second.write_bytes(first.read_bytes())
    register_table_projection(
        "journal-bank-reconciliation",
        first,
        [[[{"record_id": "first", "fields": ["amount"]}]]],
    )
    register_table_projection(
        "journal-bank-reconciliation",
        second,
        [[[{"record_id": "second", "fields": ["amount"]}]]],
    )

    read_artifact(loaded, first if first_to_read else second)

    store = RunStore(Path(loaded["run_root"]))
    sent = [event for event in store.events() if event["kind"] == "vera_send"][-1]
    [origin] = [
        selection
        for selection in sent["selections"]
        if selection["basis"] == "vera_producer_projection"
    ]
    assert origin["start"] == (0 if first_to_read else 1)
    assert origin["columns"] == [0]
    report = json.loads(Path(snapshot_run(loaded)["report_path"]).read_text())
    assert report["collection_status"] == "complete", report["errors"]
    [source] = [item for item in report["sources"] if item["role"] == "input"]
    assert source["sent_cells"] == 1
    assert source["not_sent_cells"] == 3


@pytest.mark.parametrize(
    "record",
    [
        {"record_id": "unregistered", "fields": ["amount"]},
        {"record_id": "first", "fields": ["unregistered_field"]},
    ],
)
def test_table_with_missing_source_evidence_remains_unaccounted(
    table_projection_case, record
):
    from archive_projection import register_table_projection
    from archive_reporting import read_artifact, snapshot_run

    loaded = table_projection_case
    path = Path(loaded["output_dir"]) / "derived.csv"
    path.write_text("amount\n7\n")
    register_table_projection("journal-bank-reconciliation", path, [[[record]]])

    read_artifact(loaded, path)

    report = json.loads(Path(snapshot_run(loaded)["report_path"]).read_text())
    assert report["collection_status"] == "incomplete"
    assert any(
        error.startswith("derived_output_lineage_unavailable:")
        for error in report["errors"]
    )


@pytest.mark.parametrize("cells", [[], [[[]], [[]]], [[[], []]]])
def test_table_registration_rejects_incomplete_output_cell_maps(
    table_projection_case, cells
):
    from archive_projection import register_table_projection

    path = Path(table_projection_case["output_dir"]) / "derived.csv"
    path.write_text("amount\n7\n")

    with pytest.raises(ValueError, match="exact shape"):
        register_table_projection("journal-bank-reconciliation", path, cells)


@pytest.mark.parametrize("table_projection_case", ["xlsx"], indirect=True)
def test_table_projection_scans_each_workbook_once(table_projection_case, monkeypatch):
    import openpyxl
    from archive_projection import register_table_projection
    from archive_report_store import RunStore
    from archive_reporting import read_artifact, snapshot_run

    loaded = table_projection_case
    path = Path(loaded["output_dir"]) / "derived.csv"
    path.write_text("amount,copy\n7,7\n7,7\n")
    original = openpyxl.load_workbook
    opened = []

    def counted(*args, **kwargs):
        opened.append(args[0])
        return original(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", counted)
    first = [{"record_id": "first", "fields": ["amount"]}]
    second = [{"record_id": "second", "fields": ["amount"]}]

    register_table_projection(
        "journal-bank-reconciliation", path, [[first, first], [second, second]]
    )

    assert len(opened) == 1
    read_artifact(loaded, path, offset=1, limit=1, columns=[0])
    store = RunStore(Path(loaded["run_root"]))
    sent = [event for event in store.events() if event["kind"] == "vera_send"][-1]
    [origin] = [
        selection
        for selection in sent["selections"]
        if selection["basis"] == "vera_producer_projection"
    ]
    assert origin["start"] == 2  # Worksheet headers occupy a measured row.
    report = json.loads(Path(snapshot_run(loaded)["report_path"]).read_text())
    assert report["collection_status"] == "complete", report["errors"]
    [source] = [item for item in report["sources"] if item["role"] == "input"]
    assert source["sent_cells"] == 1
    assert source["not_sent_cells"] == 5
