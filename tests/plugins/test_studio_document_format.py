from __future__ import annotations

import copy
import importlib.util
import json
import subprocess
import sys
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
COMM = ROOT / "plugins/comunicazione-professionale/scripts"
REPORT = ROOT / "plugins/report-builder"


def _helpers(name: str) -> Any:
    spec = importlib.util.spec_from_file_location(
        name, ROOT / f"tests/plugins/{name}.py"
    )
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _command(
    script: str, *args: str, check: bool = True
) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(COMM / script), *args],
        text=True,
        capture_output=True,
        check=check,
    )


def _json(path: Path, payload: Any) -> Path:
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def _studio(
    tmp_path: Path, *, name: str = "Studio Alpha", serif: bool = False
) -> tuple[Path, dict[str, Any]]:
    helper = _helpers("test_comunicazione_professionale")
    workspace = tmp_path / name.replace(" ", "-")
    _command(
        "initialize_workspace.py",
        "--workspace",
        str(workspace),
        "--workspace-id",
        workspace.name,
        "--owner",
        "Synthetic reviewer",
        "--retention-owner",
        "Synthetic reviewer",
        "--confirmed-by-user",
    )
    logo = tmp_path / f"{workspace.name}.png"
    Image.new("RGB", (240, 80), "#152f52" if not serif else "#5c302b").save(logo)
    brand = helper._brand()
    brand.update(studio_name=name, logo_path=str(logo))
    base = _json(
        tmp_path / f"{workspace.name}-base.json",
        {"brand_profile": brand, "profile": helper._new_studio_profile()},
    )
    settings = {
        "font_family": "Times New Roman" if serif else "Arial",
        "body_font_size_pt": 12 if serif else 10,
        "line_spacing_pt": 16 if serif else 13,
        "paragraph_after_pt": 10 if serif else 5,
        "heading_1_size_pt": 22 if serif else 16,
        "heading_2_size_pt": 15 if serif else 12,
        "heading_color": "5C302B" if serif else "152F52",
        "header_text": name,
        "footer_text": f"{name} - document review",
        "page_numbers": True,
        "use_logo": True,
        "table_header_fill": "5C302B" if serif else "152F52",
        "table_header_color": "FFFFFF",
        "table_alternate_fill": "F5EEEA" if serif else "EFF4FA",
        "number_format": "decimal_point" if serif else "decimal_comma",
        "date_format": "mdy" if serif else "dmy",
        "signature_lines": [name, "Accountancy team"],
    }
    preference = _json(tmp_path / f"{workspace.name}-settings.json", settings)
    sample = tmp_path / f"{workspace.name}-sample.docx"
    document = Document()
    document.add_heading(name, level=1)
    document.add_paragraph("Representative style example with synthetic text.")
    document.save(sample)
    _command(
        "review_document_format.py",
        "prepare",
        "--workspace",
        str(workspace),
        "--review-id",
        "format-review-001",
        "--settings",
        str(preference),
        "--sample",
        str(sample),
        "--base-profile",
        str(base),
    )
    review = workspace / "runs/format-review-001"
    payload = json.loads((review / "format_review.json").read_text())
    denied = _command(
        "review_document_format.py",
        "approve",
        "--review-dir",
        str(review),
        "--review-digest",
        payload["review_digest"],
        "--reviewer",
        "Synthetic reviewer",
        check=False,
    )
    assert denied.returncode != 0
    assert not (workspace / "studio_profile.json").exists()
    _command(
        "review_document_format.py",
        "approve",
        "--review-dir",
        str(review),
        "--review-digest",
        payload["review_digest"],
        "--reviewer",
        "Synthetic reviewer",
        "--confirmed-by-user",
    )
    core = _helpers("test_report_builder_plugin").load_core()
    return workspace, core.load_studio_format(
        workspace, studio_id=workspace.name, studio_name=name
    )


def test_two_studio_profiles_reuse_approved_assets_without_mixing(
    tmp_path: Path,
) -> None:
    alpha, alpha_format = _studio(tmp_path)
    beta, beta_format = _studio(tmp_path, name="Studio Beta", serif=True)
    core = _helpers("test_report_builder_plugin").load_core()
    assert (
        core.load_studio_format(
            alpha, studio_id="Studio-Alpha", studio_name="Studio Alpha"
        )
        == alpha_format
    )
    assert alpha_format["settings"]["font_family"] == "Arial"
    assert beta_format["settings"]["font_family"] == "Times New Roman"
    assert alpha_format["logo_base64"] != beta_format["logo_base64"]
    assert "samples" not in alpha_format and "profile" not in alpha_format
    with pytest.raises(ValueError, match="identity"):
        core.load_studio_format(
            beta, studio_id="Studio-Alpha", studio_name="Studio Alpha"
        )


@pytest.mark.parametrize("target", ["profile", "logo"])
def test_studio_profile_rejects_changed_approved_content(
    tmp_path: Path, target: str
) -> None:
    workspace, _ = _studio(tmp_path)
    path = workspace / "studio_profile.json"
    payload = json.loads(path.read_text())
    if target == "profile":
        payload["profile"]["document"]["docx"]["header_text"] = "Another studio"
        _json(path, payload)
    else:
        (
            workspace / payload["brand_assets"]["logo"]["workspace_relative_path"]
        ).write_bytes(b"changed")
    core = _helpers("test_report_builder_plugin").load_core()
    with pytest.raises(ValueError, match="digest|changed"):
        core.load_studio_format(
            workspace, studio_id="Studio-Alpha", studio_name="Studio Alpha"
        )


def test_profile_revision_archives_previous_standard_and_preserves_communications(
    tmp_path: Path,
) -> None:
    workspace, _ = _studio(tmp_path)
    before = json.loads((workspace / "studio_profile.json").read_text())
    settings = _json(
        tmp_path / "revision.json", {"font_family": "Arial", "paragraph_after_pt": 12}
    )
    _command(
        "review_document_format.py",
        "prepare",
        "--workspace",
        str(workspace),
        "--review-id",
        "format-review-002",
        "--settings",
        str(settings),
    )
    review = workspace / "runs/format-review-002"
    digest = json.loads((review / "format_review.json").read_text())["review_digest"]
    _command(
        "review_document_format.py",
        "approve",
        "--review-dir",
        str(review),
        "--review-digest",
        digest,
        "--reviewer",
        "Synthetic",
        "--confirmed-by-user",
    )
    after = json.loads((workspace / "studio_profile.json").read_text())
    assert after["version"] == 2
    assert (
        json.loads((workspace / "profiles/studio_profile-v001.json").read_text())
        == before
    )
    assert after["profile"]["email"] == before["profile"]["email"]
    assert (
        after["profile"]["document"]["layout"]
        == before["profile"]["document"]["layout"]
    )


@pytest.mark.parametrize(
    "serif, expected_date", [(False, "01/10/2026"), (True, "10/01/2026")]
)
def test_docx_format_changes_presentation_and_preserves_exact_financial_closure(
    tmp_path: Path, serif: bool, expected_date: str
) -> None:
    workspace, profile = _studio(
        tmp_path, name="Studio Beta" if serif else "Studio Alpha", serif=serif
    )
    helper = _helpers("test_report_builder_plugin")
    core, output = helper._build_reviewed_detail_report(tmp_path)
    analysis = core.read_json(output / "report_analysis.json")
    before_ledger = core.read_json(output / "numeric_evidence_ledger.json")
    recipe = core.read_json(output / "used_recipe.json")
    recipe["studio_format"] = profile
    recipe["report_date"] = "2026-10-01"
    recipe["executive_summary"] = (
        "A longer reviewed narrative exercises wrapping without changing financial evidence. "
        * 40
    )
    core.write_json(output / "used_recipe.json", recipe)
    core.write_report_docx(
        recipe,
        analysis,
        core.read_json(output / "report_audit.json"),
        output / "report.docx",
    )
    ledger = core.write_numeric_evidence_ledger(output, analysis)
    document = Document(output / "report.docx")
    assert ledger["entries"] == before_ledger["entries"]
    assert core.read_json(output / "report_analysis.json") == analysis
    assert document.styles["Normal"].font.name == profile["settings"]["font_family"]
    assert (
        document.sections[0].header.paragraphs[0].text.endswith(profile["studio_name"])
    )
    assert "PAGE" in document.sections[0].footer._element.xml
    assert len(document.sections) == 2
    assert document.sections[1].header.is_linked_to_previous
    assert document.tables[2].rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    assert document.tables[2].rows[1]._tr.xpath("./w:trPr/w:cantSplit")
    assert not document.tables[2].rows[1]._tr.xpath("./w:trPr/w:trHeight")
    assert expected_date in [p.text for p in document.paragraphs]
    assert any(
        p.text == "\n".join(profile["settings"]["signature_lines"])
        for p in document.paragraphs
    )
    # Regeneration uses the frozen settings, even if the studio's current profile changes.
    (workspace / "studio_profile.json").unlink()
    replay = tmp_path / "regenerated.docx"
    core.write_report_docx(
        recipe, analysis, core.read_json(output / "report_audit.json"), replay
    )
    assert replay.read_bytes() == (output / "report.docx").read_bytes()


def test_shared_writer_rejects_omitted_word_standard_without_mutation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workspace, _ = _studio(tmp_path)
    profile_path = workspace / "studio_profile.json"
    before = profile_path.read_bytes()
    stored = json.loads(before)
    proposal = copy.deepcopy(stored["profile"])
    del proposal["document"]["docx"]
    monkeypatch.syspath_prepend(str(COMM))
    spec = importlib.util.spec_from_file_location(
        "studio_format_profile_writer", COMM / "promote_studio_profile.py"
    )
    assert spec and spec.loader
    writer = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(writer)

    with writer.workflow_lock(workspace):
        with pytest.raises(ValueError, match="DOCX preferences are missing"):
            writer.persist_studio_profile(
                workspace,
                profile=proposal,
                brand_profile=stored["brand_profile"],
                workspace_id=workspace.name,
                logo=None,
                approved_from=stored["approved_from"],
            )

    assert profile_path.read_bytes() == before
    assert not (workspace / "profiles").exists()


def test_word_format_distinguishes_totals_from_five_column_text_preview(
    tmp_path: Path,
) -> None:
    _, profile = _studio(tmp_path)
    core = _helpers("test_report_builder_plugin").load_core()
    document = Document()
    core.add_numeric_totals_table(
        document,
        [
            {
                "column": "Revenue",
                "sum": "1234.50",
                "currency": "EUR",
                "unit": "currency",
                "scale": "1",
            }
        ],
    )
    preview = document.add_table(rows=2, cols=5)
    preview.cell(1, 1).text = "A source description, not a reviewed total"

    core.apply_studio_format(document, profile)

    assert (
        document.tables[0].cell(1, 1).paragraphs[0].alignment
        == WD_ALIGN_PARAGRAPH.RIGHT
    )
    assert preview.cell(1, 1).paragraphs[0].alignment != WD_ALIGN_PARAGRAPH.RIGHT


def test_docx_format_schema_matches_the_shared_profile_extension() -> None:
    communication = json.loads(
        (COMM.parent / "schemas/model_contribution.schema.json").read_text()
    )
    report = json.loads((REPORT / "assets/studio-docx-format.schema.json").read_text())
    assert communication["$defs"]["documentProfile"]["properties"]["docx"] == report


def test_format_adoption_rejects_changed_selected_example(tmp_path: Path) -> None:
    workspace, _ = _studio(tmp_path)
    settings = _json(
        tmp_path / "changed-evidence-settings.json", {"paragraph_after_pt": 8}
    )
    sample = tmp_path / "selected-example.txt"
    sample.write_text("Selected formatting evidence.", encoding="utf-8")
    _command(
        "review_document_format.py",
        "prepare",
        "--workspace",
        str(workspace),
        "--review-id",
        "format-review-002",
        "--settings",
        str(settings),
        "--sample",
        str(sample),
    )
    review = workspace / "runs/format-review-002"
    payload = json.loads((review / "format_review.json").read_text())
    Path(payload["samples"][0]["snapshot_path"]).write_text(
        "Changed evidence.", encoding="utf-8"
    )
    before = (workspace / "studio_profile.json").read_bytes()

    result = _command(
        "review_document_format.py",
        "approve",
        "--review-dir",
        str(review),
        "--review-digest",
        payload["review_digest"],
        "--reviewer",
        "Synthetic",
        "--confirmed-by-user",
        check=False,
    )

    assert result.returncode != 0
    assert "evidence changed" in result.stderr
    assert (workspace / "studio_profile.json").read_bytes() == before


def test_selected_format_survives_assured_native_review_regeneration(
    tmp_path: Path,
) -> None:
    workspace, profile = _studio(tmp_path)
    helper = _helpers("test_report_builder_plugin")
    source = tmp_path / "income_statement.csv"
    source.write_text("line,amount\nRevenue,1234.50\nCost,-234.50\n", encoding="utf-8")
    managed = helper._managed_report_run(tmp_path, source)
    core = managed["core"]
    output = managed["output_dir"]
    core.build_report(
        managed["source_path"],
        output,
        run_id=managed["run_id"],
        client_engagement=managed["context"],
        studio_workspace=workspace,
        studio_id=workspace.name,
        studio_name="Studio Alpha",
    )
    saved = core.read_json(output / "used_recipe.json")["studio_format"]
    assert saved == profile
    review = core.read_json(output / "review_payload.json")
    item = next(
        item
        for item in review["items"]
        if item["item_type"] == "report_section"
        and item["data"]["section"] == "income_statement"
    )
    response = helper._call_mcp_server(
        "tools/call",
        {
            "name": "apply_report_builder_decisions",
            "arguments": {
                "client_engagement": str(managed["context_path"]),
                "run_intake": core.read_json(output / "run_intake.json"),
                "review_payload": review,
                "final_artifacts": core.read_json(output / "final_artifacts.json"),
                "decisions": [
                    {
                        "item_id": item["id"],
                        "action": "edit",
                        "edit_value": "The reviewer confirmed this longer report narrative after inspecting the mapped source evidence.",
                        "reviewer_note": "Reviewed narrative",
                    }
                ],
                "decision_source": "pytest_studio_format",
                "reviewer": "Synthetic reviewer",
            },
        },
    )
    assert response["structuredContent"]["ok"] is True, response
    assert core.read_json(output / "used_recipe.json")["studio_format"] == saved
    assert Document(output / "report.docx").styles["Normal"].font.name == "Arial"
    assert (
        Document(output / "report.docx")
        .sections[0]
        .header.paragraphs[0]
        .text.endswith("Studio Alpha")
    )
