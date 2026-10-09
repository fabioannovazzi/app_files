"""Real public CR producers through native signed scopes on fictional evidence."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.plugins.test_centrale_rischi_review import (
    _reviewed_recipe,
    _write_native_pdf_fixture,
    _write_source,
)
from tests.plugins.test_vera_client_workflow_filesystem import _load_customer_ledger
from tests.plugins.test_vera_native_archive_lifecycle import rpc_program
from tests.plugins.test_vera_native_workspace import ROOT, configure, workspace_module

LONG_SOURCE_CELL = "Valore completo Ω — non dedurre il significato. " * 8


def write_extended_source(path: Path) -> None:
    """Fictional rows beyond both public preview and native page bounds, plus empty tables."""
    _write_source(path)
    workbook = load_workbook(path)
    sheet = workbook.active
    for index in range(31):
        sheet.append(
            [
                "2025-02",
                LONG_SOURCE_CELL,
                "Categoria ignota",
                f"  Termine ignoto {index:02d} Ω  ",
                "Termine residuo ignoto",
                1,
                1,
                1,
                "",
                0,
                "",
            ]
        )
    auxiliary = workbook.create_sheet("Popolazione separata")
    auxiliary.append(["Codice", "Importo"])
    auxiliary.append(["Ausiliario fittizio", 901])
    for index in range(32):
        workbook.create_sheet(f"Vuota {index:02d}").append(["Campo"])
    workbook.save(path)
    workbook.close()


@pytest.fixture
def cr_run(tmp_path, monkeypatch, request):
    """Register fictional CR source receipts before starting its public Archive run."""
    ledger = _load_customer_ledger()
    monkeypatch.syspath_prepend(str(ROOT / "plugins/vera/scripts"))
    folder = tmp_path / "Fictional CR customer"
    folder.mkdir()
    client = "client_111111111111111111111111"
    ledger.create_client_manifest(folder, client)
    engagement = ledger.create_engagement(folder, client, "Fictional CR")
    if getattr(request, "param", "table") == "pdf":
        source = tmp_path / "fictional-cr.pdf"
        _write_native_pdf_fixture(source)
    else:
        source = tmp_path / "fictional-cr.xlsx"
        if getattr(request, "param", "table") == "extended":
            write_extended_source(source)
        else:
            _write_source(source)
    imported = ledger.import_document(
        folder, client, engagement["engagement_id"], source, "source"
    )["receipt"]
    prepared = ledger.prepare_run(
        folder,
        client,
        engagement["engagement_id"],
        "centrale-rischi-review",
        "test-version",
        input_ids=[imported["input_id"]],
    )
    loaded = ledger.start_run(
        folder, engagement["engagement_id"], prepared["run"]["run_id"]
    )
    binding = {
        "work_ref": "fictional-cr",
        "client_root": str(folder),
        "client_id": client,
        "engagement_id": engagement["engagement_id"],
        "run_id": prepared["run"]["run_id"],
        "workflow_id": "centrale-rischi-review",
    }
    env = configure(monkeypatch, tmp_path, [binding])
    return env, Path(loaded["output_dir"]), binding


INSPECT = """
const initial=payload(call('vera_workspace_cr_setup',{work_ref:'fictional-cr'}));
const inspectArgs={work_ref:'fictional-cr',revision:initial.revision,review_ticket:initial.review_ticket,input_ids:[initial.items[0].id],human_reviewed:true,confirmed:true,idempotency_key:'fictional-inspect'};
const inspected=payload(call('vera_workspace_cr_inspect',inspectArgs));
const selected=payload(call('vera_workspace_cr_view',{work_ref:'fictional-cr',source_ref:inspected.source_ref}));
const grantArgs={work_ref:'fictional-cr',source_ref:inspected.source_ref,revision:selected.revision,review_ticket:selected.review_ticket,question:'Spiega le esposizioni fittizie.  ',human_reviewed:true,confirmed:true,idempotency_key:'fictional-grant'};
const grant=payload(call('vera_workspace_cr_grant',grantArgs));
const model=(name,args)=>{const r=call(name,args);if(r.isError)throw new Error(r.content[0].text);return r.structuredContent;};
const exact={work_ref:'fictional-cr',grant_ref:grant.grant_ref};
const context=model('vera_workspace_cr_model_context',exact);
"""


def proposal_script(changes: str = "") -> str:
    recipe = _reviewed_recipe({"inventory_sha256": "", "tables": [{"table_id": ""}]})
    return (
        INSPECT
        + f"""
const proposal={{...context.template,...{json.dumps(recipe)},inventory_sha256:context.template.inventory_sha256,table_id:context.context.tables[0].table_id,mapping_review:context.template.mapping_review}};
{changes}
const stageArgs={{...exact,expected_proposal_revision:context.proposal_revision,proposal,idempotency_key:'fictional-stage'}};
const staged=model('vera_workspace_cr_stage',stageArgs);
const candidate=payload(call('vera_workspace_cr_proposal_read',exact));
const calculateArgs={{...exact,source_ref:inspected.source_ref,revision:candidate.revision,review_ticket:candidate.review_ticket,proposal_sha256:candidate.proposal_sha256,reviewer:'Fictional professional',reviewed_at:'2026-10-07T15:00:00+02:00',human_reviewed:true,confirmed:true,idempotency_key:'fictional-calculate'}};
"""
    )


@pytest.mark.parametrize(
    ("preview_key", "phase", "filename", "format_name"),
    [
        ("html", "report", "centrale_rischi_dashboard_reviewed.html", "html"),
        ("markdown", "report", "centrale_rischi_report.md", "markdown"),
        ("facts", "calculated", "centrale_rischi_facts.md", "markdown"),
    ],
)
def test_cr_full_public_pipeline_preserves_outputs_and_draft_boundary(
    cr_run,
    preview_key,
    phase,
    filename,
    format_name,
):
    env, output, _ = cr_run
    result = rpc_program(
        env,
        proposal_script()
        + """
const calculated=payload(call('vera_workspace_cr_calculate',calculateArgs));
const retry=payload(call('vera_workspace_cr_calculate',calculateArgs));
const analysis=payload(call('vera_workspace_cr_view',{work_ref:'fictional-cr',source_ref:calculated.source_ref}));
const commentGrant=payload(call('vera_workspace_cr_grant',{work_ref:'fictional-cr',source_ref:calculated.source_ref,revision:analysis.revision,review_ticket:analysis.review_ticket,question:'Commenta i dati fittizi senza rating.',human_reviewed:true,confirmed:true,idempotency_key:'fictional-comment-grant'}));
const commentExact={work_ref:'fictional-cr',grant_ref:commentGrant.grant_ref};
const commentContext=model('vera_workspace_cr_model_context',commentExact);
const comment={...commentContext.template,observations:[{text:'Esposizione utilizzata del mese fittizio.',evidence_refs:['metric:'+commentContext.context.metrics[0].metric_id]}],hypotheses:[],questions:['Quali documenti spiegano la variazione?'],limitations:['Le fonti sono fittizie.']};
model('vera_workspace_cr_stage',{...commentExact,expected_proposal_revision:commentContext.proposal_revision,proposal:comment,idempotency_key:'fictional-comment-stage'});
const commentCandidate=payload(call('vera_workspace_cr_proposal_read',commentExact));
const reportArgs={...commentExact,source_ref:calculated.source_ref,revision:commentCandidate.revision,review_ticket:commentCandidate.review_ticket,proposal_sha256:commentCandidate.proposal_sha256,human_reviewed:true,confirmed:true,idempotency_key:'fictional-finalize'};
const report=payload(call('vera_workspace_cr_finalize',reportArgs));
const reopen=payload(call('vera_workspace_cr_view',{work_ref:'fictional-cr',source_ref:report.source_ref}));
const preview=(source_ref,item_id)=>payload(call('vera_workspace_cr_view',{work_ref:'fictional-cr',source_ref,item_id}));
const html=preview(report.source_ref,'centrale_rischi_dashboard_reviewed.html');
const markdown=preview(report.source_ref,'centrale_rischi_report.md');
const facts=preview(calculated.source_ref,'centrale_rischi_facts.md');
const childRefusal=call('vera_workspace_cr_view',{work_ref:'fictional-cr',source_ref:report.source_ref,item_id:'centrale_rischi_report.md',member_ref:'json:'+JSON.stringify({path:[],offset:0})});
const traversal=call('vera_workspace_cr_view',{work_ref:'fictional-cr',source_ref:report.source_ref,item_id:'../centrale_rischi_report.md'});
const result={initial,inspected,context,calculated,retry,analysis,report,reopen,html,markdown,facts,childRefusal,traversal,summary:call('vera_workspace_cr_setup',{work_ref:'fictional-cr'}).structuredContent};
""",
    )
    analysis = json.loads(
        (
            output
            / result["calculated"]["source_ref"]
            / "centrale_rischi_analysis.json"
        ).read_bytes()
    )
    assert result["retry"] == result["calculated"]
    assert result["context"]["question"].endswith("  ")
    assert analysis["assurance_levels"]["professional"] == "pending"
    assert result["report"]["status"] == "draft_pending_professional_review"
    assert result["report"]["run_completed"] is False
    assert result["report"]["professional_approval"] is False
    assert len(list(output.glob("cr-*"))) == 3
    assert (
        output / result["calculated"]["source_ref"] / "centrale_rischi_analysis.xlsx"
    ).is_file()
    assert (
        output / result["report"]["source_ref"] / "centrale_rischi_report.md"
    ).is_file()
    assert str(output) not in json.dumps(result["context"])
    assert "inputs" not in result["summary"]
    assert result["childRefusal"]["isError"] is True
    assert result["traversal"]["isError"] is True
    assert "centrale_rischi_report.md" in {
        row["id"] for row in result["reopen"]["items"]
    }
    from native_bank_preparation import file_hash

    artifact = output / result[phase]["source_ref"] / filename
    preview = result[preview_key]["selection"]["artifact_preview"]
    assert preview == {
        "format": format_name,
        "content": artifact.read_bytes().decode("utf-8"),
        "sha256": file_hash(artifact),
        "complete": True,
    }
    assert result[preview_key]["data"]["professional_approval"] is False


@pytest.mark.parametrize(
    "change",
    [
        "proposal.mapping_review={status:'reviewed',reviewer:'Model',reviewed_at:'now'};",
        "proposal.inventory_sha256='f'.repeat(64);",
        "delete proposal.value_mappings;",
    ],
)
def test_cr_model_cannot_manufacture_mapping_approval_or_change_provenance(
    cr_run, change
):
    env, output, _ = cr_run
    script = proposal_script(change).split("const staged=")[0]
    result = rpc_program(
        env, script + "const result=call('vera_workspace_cr_stage',stageArgs);"
    )
    assert result["isError"] is True
    assert len(list(output.glob("cr-*"))) == 1


@pytest.mark.parametrize(
    "change",
    [
        "calculateArgs.review_ticket='forged';",
        "calculateArgs.proposal_sha256='f'.repeat(64);",
        "calculateArgs.reviewed_at='2026-10-07';",
        "calculateArgs.reviewer=' ';",
    ],
)
def test_cr_rejects_forged_or_unreviewed_calculation_without_public_writes(
    cr_run, change
):
    env, output, _ = cr_run
    result = rpc_program(
        env,
        proposal_script()
        + change
        + "const result=call('vera_workspace_cr_calculate',calculateArgs);",
    )
    assert result["isError"] is True
    assert len(list(output.glob("cr-*"))) == 1


def test_cr_reconciled_mode_stays_unsupported_without_uncertain_output(cr_run):
    env, output, _ = cr_run
    result = rpc_program(
        env,
        proposal_script("proposal.analysis_mode='reconciled';")
        + "const result=call('vera_workspace_cr_calculate',calculateArgs);",
    )
    assert result["isError"] is True
    assert "Reconciled mode requires" in result["content"][0]["text"]
    assert len(list(output.glob("cr-*"))) == 1


def test_cr_blocked_controls_are_persisted_and_cannot_grant_commentary(cr_run):
    env, output, _ = cr_run
    result = rpc_program(
        env,
        proposal_script("proposal.control_totals.used='1499';")
        + """
const calculated=payload(call('vera_workspace_cr_calculate',calculateArgs));
const view=payload(call('vera_workspace_cr_view',{work_ref:'fictional-cr',source_ref:calculated.source_ref}));
const refusal=call('vera_workspace_cr_grant',{work_ref:'fictional-cr',source_ref:calculated.source_ref,revision:view.revision,review_ticket:view.review_ticket,question:'Commenta.',confirmed:true,human_reviewed:true,idempotency_key:'blocked-grant'});
const result={calculated,refusal};
""",
    )
    assert result["calculated"]["status"] == "blocked"
    assert result["calculated"]["exit_status"] == 2
    assert result["refusal"]["isError"] is True
    assert (
        output / result["calculated"]["source_ref"] / "execution_receipt.json"
    ).is_file()


def test_cr_changed_artifacts_and_interrupted_writes_block_native_closure(
    cr_run, monkeypatch
):
    env, output, binding = cr_run
    rpc_program(env, INSPECT + "const result=inspected;")
    api = workspace_module()
    from native_centrale_rischi import audit_run

    private = api.ui_state_directory(output)
    api.atomic_json(
        private / "cr-request-interrupted.json", {"request_sha256": "f" * 64}
    )
    assert audit_run(output, api)["recovery_required"] is True
    setup = api.dispatch("vera_workspace_cr_setup", {"work_ref": binding["work_ref"]})
    assert setup["can_write"] is False
    first = next(output.glob("cr-*"))
    (first / "unknown.txt").write_text("Changed artifact population")
    with pytest.raises(ValueError, match="artifacts changed"):
        audit_run(output, api)


def test_cr_unknown_input_id_cannot_read_or_create_another_run(cr_run):
    env, output, _ = cr_run
    result = rpc_program(
        env,
        INSPECT.split("const inspected=")[0]
        + "inspectArgs.input_ids=['foreign-input'];const result=call('vera_workspace_cr_inspect',inspectArgs);",
    )
    assert result["isError"] is True
    assert not list(output.glob("cr-*"))


@pytest.mark.parametrize("cr_run", ["pdf"], indirect=True)
def test_cr_digital_pdf_retains_original_normalization_and_separate_populations(cr_run):
    env, output, _ = cr_run
    columns = {
        "reference_month": "reference_month",
        "intermediary": "intermediary",
        "risk_category": "category",
        "original_duration": "original_duration",
        "residual_duration": "residual_duration",
        "granted": "granted",
        "operational_granted": "operational_granted",
        "used": "used",
        "guarantee_type": "guarantee_type",
        "guaranteed_amount": "guaranteed_amount",
        "record_status": "record_status",
        "valid_from": "valid_from",
        "valid_to": "valid_to",
        "source_page": "source_page",
        "source_region": "source_region",
        "source_row_locator": "source_row_locator",
        "extraction_confidence": "extraction_confidence",
    }
    mappings = {
        "original_term": {"Oltre cinque anni": "long"},
        "residual_term": {"Oltre 1 anno": "over_one_year"},
        "exposure_family": {"RISCHI A SCADENZA": "performing"},
    }
    script = (
        INSPECT
        + f"""
const proposal={{...context.template,entity:'Fictional PDF company',analysis_objective:'Inspect fictional digital PDF',table_id:context.context.tables[0].table_id,columns:{json.dumps(columns)},value_mappings:{json.dumps(mappings)},control_totals:{{used:'45000'}}}};
model('vera_workspace_cr_stage',{{...exact,expected_proposal_revision:context.proposal_revision,proposal,idempotency_key:'pdf-stage'}});
const candidate=payload(call('vera_workspace_cr_proposal_read',exact));
const calculated=payload(call('vera_workspace_cr_calculate',{{...exact,source_ref:inspected.source_ref,revision:candidate.revision,review_ticket:candidate.review_ticket,proposal_sha256:candidate.proposal_sha256,reviewer:'Fictional PDF reviewer',reviewed_at:'2026-10-07T15:00:00+02:00',human_reviewed:true,confirmed:true,idempotency_key:'pdf-calculate'}}));
const result={{inspected,calculated}};
"""
    )
    result = rpc_program(env, script)
    inspection = output / result["inspected"]["source_ref"]
    calculation = output / result["calculated"]["source_ref"]
    normalized_receipt = json.loads(
        (inspection / "pdf_normalization_receipt.json").read_bytes()
    )
    analysis = json.loads((calculation / "centrale_rischi_analysis.json").read_bytes())
    assert (inspection / "centrale_rischi_normalized.xlsx").read_bytes() == (
        calculation / "centrale_rischi_normalized.xlsx"
    ).read_bytes()
    assert normalized_receipt["status"] == "pending_professional_review"
    assert analysis["source"]["current_row_count"] == 1
    assert analysis["source"]["previous_row_count"] == 1
    assert analysis["coverage"]["guarantees_received"] == "available"
    assert analysis["coverage"]["information_requests"] == "available"
    assert analysis["assurance_levels"]["professional"] == "pending"


def test_cr_proposal_cas_preserves_history_and_requires_exact_full_selection(cr_run):
    env, output, _ = cr_run
    result = rpc_program(
        env,
        proposal_script()
        + """
const stale=call('vera_workspace_cr_stage',{...stageArgs,idempotency_key:'stale-stage'});
const refreshed=model('vera_workspace_cr_model_context',exact);
const changed={...proposal,analysis_objective:'A second explicitly fictional objective.'};
model('vera_workspace_cr_stage',{...exact,expected_proposal_revision:refreshed.proposal_revision,proposal:changed,idempotency_key:'second-stage'});
const refused=call('vera_workspace_cr_calculate',calculateArgs);
const retained=payload(call('vera_workspace_cr_grants',{work_ref:'fictional-cr',source_ref:inspected.source_ref}));
const result={stale,refused,retained};
""",
    )
    assert result["stale"]["isError"] is True
    assert result["refused"]["isError"] is True
    assert (
        len(
            list(
                (output.parent / ".native-workspace").glob("cr-grant-*-proposal-*.json")
            )
        )
        == 2
    )
    assert result["retained"]["grants"][0]["question"].endswith("  ")
    assert result["retained"]["grants"][0]["proposal_available"] is True
    assert len(list(output.glob("cr-*"))) == 1


@pytest.mark.parametrize(
    "offset, expected_count, has_more", [(0, 30, True), (30, 1, False)]
)
def test_cr_grants_page_all_questions_and_distinguish_pending_proposals(
    cr_run, offset, expected_count, has_more
):
    env, _, _ = cr_run
    result = rpc_program(
        env,
        proposal_script()
        + f"""
for(let index=0;index<30;index++)payload(call('vera_workspace_cr_grant',{{...grantArgs,question:'Fictional pending question '+index,idempotency_key:'pending-grant-'+index}}));
const result=payload(call('vera_workspace_cr_grants',{{work_ref:'fictional-cr',source_ref:inspected.source_ref,offset:{offset}}}));
""",
    )
    assert result["total"] == 31
    assert len(result["grants"]) == expected_count
    assert result["offset"] == offset
    assert result["has_more"] is has_more
    assert {row["proposal_available"] for row in result["grants"]} <= {True, False}


def test_cr_json_member_navigation_reaches_rows_without_latest_version_inference(
    cr_run,
):
    env, output, _ = cr_run
    result = rpc_program(
        env,
        proposal_script()
        + """
const calculated=payload(call('vera_workspace_cr_calculate',calculateArgs));
const rows=payload(call('vera_workspace_cr_view',{work_ref:'fictional-cr',source_ref:calculated.source_ref,item_id:'centrale_rischi_analysis.json',member_ref:'json:'+JSON.stringify({path:['exposures'],offset:0})}));
const denied=call('vera_workspace_cr_view',{work_ref:'fictional-cr',source_ref:calculated.source_ref,item_id:'../inspection.json'});
const missing=call('vera_workspace_cr_view',{work_ref:'fictional-cr',source_ref:'cr-'+'f'.repeat(64)});
const result={rows,denied,missing};
""",
    )
    assert result["rows"]["selection"]["prepared_context"]["total"] > 0
    assert result["denied"]["isError"] is True
    assert result["missing"]["isError"] is True


@pytest.mark.parametrize("cr_run", ["extended"], indirect=True)
def test_cr_complete_source_rows_preserve_long_cells_and_all_tables(cr_run):
    env, _, _ = cr_run
    result = rpc_program(
        env,
        INSPECT
        + """
const base={work_ref:'fictional-cr',source_ref:inspected.source_ref};
const read=source_selector=>payload(call('vera_workspace_cr_source',{...base,source_selector}));
const tables=read({table_id:'',kind:'rows',column:'',offset:0});
const tail=read({table_id:'',kind:'rows',column:'',offset:30});
const selectedRows=read({table_id:tables.source.page.entries[0].table_id,kind:'rows',column:'',offset:30});
const auxiliary=read({table_id:tables.source.page.entries[1].table_id,kind:'rows',column:'',offset:0});
const result={tables,tail,selected:selectedRows,auxiliary};
""",
    )
    assert result["tables"]["source"]["page"]["total"] == 34
    assert len(result["tables"]["source"]["page"]["entries"]) == 30
    assert len(result["tail"]["source"]["page"]["entries"]) == 4
    page = result["selected"]["source"]["page"]
    assert page["total"] == 36
    assert len(page["entries"]) == 6
    assert page["entries"][0]["source_row"] == 32
    assert page["entries"][0]["values"]["Intermediario"] == LONG_SOURCE_CELL
    assert (
        page["entries"][0]["values"]["Durata originaria"] == "  Termine ignoto 25 Ω  "
    )
    assert result["selected"]["model_grant_issued"] is False
    assert result["selected"]["professional_approval"] is False
    assert result["auxiliary"]["source"]["page"]["entries"] == [
        {"source_row": 2, "values": {"Codice": "Ausiliario fittizio", "Importo": 901}}
    ]


@pytest.mark.parametrize("cr_run", ["extended"], indirect=True)
def test_cr_observed_values_use_public_mapping_keys_without_classification(cr_run):
    env, _, _ = cr_run
    result = rpc_program(
        env,
        INSPECT
        + """
const source_selector={table_id:context.context.tables[0].table_id,kind:'values',column:'Durata originaria',offset:30};
const result=payload(call('vera_workspace_cr_source',{work_ref:'fictional-cr',source_ref:inspected.source_ref,source_selector}));
""",
    )
    page = result["source"]["page"]
    assert page["total"] == 34
    assert len(page["entries"]) == 4
    assert page["entries"][0] == {
        "value": "Termine ignoto 27 Ω",
        "count": 1,
        "first_source_row": 34,
    }
    assert page["has_more"] is False
    assert "class" not in page["entries"][0]
    assert page["missing_cell_count"] == 0


@pytest.mark.parametrize("cr_run", ["extended"], indirect=True)
def test_cr_exact_signed_source_page_grant_exposes_only_selected_page(cr_run):
    env, output, _ = cr_run
    result = rpc_program(
        env,
        INSPECT
        + """
const source_selector={table_id:context.context.tables[0].table_id,kind:'rows',column:'Intermediario',offset:30};
const exactSource={work_ref:'fictional-cr',source_ref:inspected.source_ref,source_selector};
const page=payload(call('vera_workspace_cr_source',exactSource));
const grantSource=payload(call('vera_workspace_cr_grant',{...exactSource,item_id:page.selection.id,revision:page.revision,review_ticket:page.review_ticket,question:'Leggi questa pagina esatta.  ',confirmed:true,human_reviewed:true,idempotency_key:'exact-source-grant'}));
const actual=model('vera_workspace_cr_model_context',{work_ref:'fictional-cr',grant_ref:grantSource.grant_ref});
const result={page,actual};
""",
    )
    assert result["actual"]["selected_source_page"] == result["page"]["source"]
    assert result["actual"]["question"] == "Leggi questa pagina esatta.  "
    assert len(result["actual"]["selected_source_page"]["page"]["entries"]) == 6
    assert set(
        result["actual"]["selected_source_page"]["page"]["entries"][0]["values"]
    ) == {"Intermediario"}
    assert str(output) not in json.dumps(result["actual"])
    assert result["actual"]["template"]["mapping_review"]["status"] == "pending"
    assert result["actual"]["professional_approval"] is False


@pytest.mark.parametrize(
    "change",
    [
        "sourceArgs.source_selector.offset=30;",
        "sourceArgs.item_id='source-page-'+'f'.repeat(64);",
        "delete sourceArgs.source_selector;",
        "sourceArgs.review_ticket=selected.review_ticket;",
    ],
)
def test_cr_refuses_changed_or_unselected_source_page_grants(cr_run, change):
    env, output, _ = cr_run
    result = rpc_program(
        env,
        INSPECT
        + """
const source_selector={table_id:context.context.tables[0].table_id,kind:'rows',column:'Intermediario',offset:0};
const page=payload(call('vera_workspace_cr_source',{work_ref:'fictional-cr',source_ref:inspected.source_ref,source_selector}));
const sourceArgs={...grantArgs,source_selector,item_id:page.selection.id,revision:page.revision,review_ticket:page.review_ticket,idempotency_key:'invalid-source-grant'};
"""
        + change
        + "const result=call('vera_workspace_cr_grant',sourceArgs);",
    )
    assert result["isError"] is True
    assert len(list(output.glob("cr-*"))) == 1


@pytest.mark.parametrize(
    "change",
    [
        "selector.table_id='../outside';",
        "selector.column='Unregistered column';",
        "selector.offset=-1;",
        "selector.kind='semantic';",
        "selector.path='/private/tmp/outside';",
    ],
)
def test_cr_source_navigation_rejects_foreign_coordinates_and_paths(cr_run, change):
    env, _, _ = cr_run
    result = rpc_program(
        env,
        INSPECT
        + """
const selector={table_id:context.context.tables[0].table_id,kind:'rows',column:'',offset:0};
"""
        + change
        + "const result=call('vera_workspace_cr_source',{work_ref:'fictional-cr',source_ref:inspected.source_ref,source_selector:selector});",
    )
    assert result["isError"] is True


@pytest.mark.parametrize("cr_run", ["pdf"], indirect=True)
def test_cr_source_pages_reuse_receipted_pdf_normalization_and_keep_provenance(cr_run):
    env, _, _ = cr_run
    result = rpc_program(
        env,
        INSPECT
        + """
const tables=payload(call('vera_workspace_cr_source',{work_ref:'fictional-cr',source_ref:inspected.source_ref,source_selector:{table_id:'',kind:'rows',column:'',offset:0}}));
const page=payload(call('vera_workspace_cr_source',{work_ref:'fictional-cr',source_ref:inspected.source_ref,source_selector:{table_id:tables.source.page.entries[0].table_id,kind:'rows',column:'',offset:0}}));
const result={tables,page};
""",
    )
    assert result["tables"]["source"]["source_kind"] == "native_pdf_extraction"
    rows = result["page"]["source"]["page"]["entries"]
    assert rows[0]["values"]["source_page"] == 1
    assert rows[0]["values"]["source_row_locator"]
    assert rows[0]["values"]["extraction_confidence"]
    assert result["page"]["professional_approval"] is False
