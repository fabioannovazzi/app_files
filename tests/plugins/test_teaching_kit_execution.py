"""Execute the packaged teaching inputs through current, unmocked pipelines.

These tests establish input compatibility and actual artifact generation. They
do not simulate a learner, certify professional judgement or test native voice.
"""

from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import mimetypes
import subprocess
import sys
from decimal import Decimal
from pathlib import Path

import pytest

from tests.plugins._teaching_release import prepared_kit, record_native_check

ROOT = Path(__file__).resolve().parents[2]


def _read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def _write(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def _run(program, *args):
    result = subprocess.run(
        [sys.executable, str(ROOT / program), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=60,
        check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr


def _bound_case(
    tmp_path,
    monkeypatch,
    workflow,
    ledger_workflow,
    phase,
    product="vera",
    language="it",
):
    monkeypatch.syspath_prepend(str(ROOT / "plugins/_shared/vendor/modules"))
    from courseware.library import CourseLibrary

    kit = CourseLibrary(ROOT / "plugins" / product, {workflow}).render(
        workflow, language, tmp_path / "kit"
    )
    sources = kit["source_files" if phase == "demo" else "practice_files"]
    spec = importlib.util.spec_from_file_location(
        "kit_execution_ledger", ROOT / "plugins/studio-archive/scripts/client_ledger.py"
    )
    ledger = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(ledger)
    (tmp_path / ".vera-onboarding-local-only").write_text(
        "Mechanical teaching-kit verification; no user profile or telemetry.\n",
        encoding="utf-8",
    )
    case = tmp_path / "case"
    case.mkdir()
    client = "client_0123456789abcdef01234567"
    ledger.create_client_manifest(case, client)
    engagement = ledger.create_engagement(case, client, "Fictional kit verification")[
        "engagement_id"
    ]
    imported = [
        ledger.import_document(
            case,
            client,
            engagement,
            Path(source),
            (
                "journal"
                if ledger_workflow == "journal-sampling"
                and Path(source).suffix.lower() in {".csv", ".xlsx", ".xls"}
                else "source"
            ),
        )
        for source in sources
    ]
    if ledger_workflow == "treasury-forecast":
        tables = {
            Path(source).stem: {
                "path": "imports/"
                + item["receipt"]["input_id"]
                + "/"
                + Path(item["receipt"]["path"]).name
            }
            for source, item in zip(sources, imported, strict=True)
            if Path(source).suffix == ".csv"
        }
        manifest_path = case / "manifest.json"
        _write(
            manifest_path,
            {
                "schema_version": "vera.treasury_manifest.v1",
                "company_id": "officina-arco-fictional",
                "company_name": "Officina Arco — caso fittizio",
                "currency": "EUR",
                "as_of": "2026-03-31",
                "horizon_end": "2026-05-31",
                "coverage": "Un conto bancario e le sole partite residue e uscite previste fornite; nessuna vendita futura non fatturata e nessun calcolo di imposte.",
                "tables": tables,
                "invoice_files": [],
                "previous": None,
            },
        )
        imported.append(
            ledger.import_document(case, client, engagement, manifest_path, "source")
        )
    version = _read(ROOT / "plugins" / product / ".codex-plugin/plugin.json")["version"]
    run = ledger.prepare_run(
        case,
        client,
        engagement,
        ledger_workflow,
        version,
        input_ids=[item["receipt"]["input_id"] for item in imported],
        purpose="Local fictional kit verification; no professional approval",
    )
    return ledger.start_run(case, engagement, run["run"]["run_id"])


def _complete_teaching_case(run, client_root, *, artifact_ids=None):
    """Seal actual local test artifacts so a later run can import the result."""
    from tests.model_data_helpers import write_no_model_report

    spec = importlib.util.spec_from_file_location(
        "teaching_completion_ledger",
        ROOT / "plugins/studio-archive/scripts/client_ledger.py",
    )
    ledger = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(ledger)
    output = Path(run["output_dir"])
    context = run["context"]
    artifact_ids = artifact_ids or {}
    declarations = [
        {
            "artifact_id": artifact_ids.get(
                path.relative_to(output).as_posix(), f"internal.teaching.{index}"
            ),
            "path": path.relative_to(output).as_posix(),
            "purpose": "Preserve actual fictional teaching verification output",
            "audience": "review",
            "media_type": mimetypes.guess_type(path.name)[0]
            or "application/octet-stream",
        }
        for index, path in enumerate(
            sorted(p for p in output.rglob("*") if p.is_file())
        )
    ]
    declarations.extend(
        write_no_model_report(output, context["workflow_id"], context["run_id"])
    )
    ledger.finalize_run(
        client_root, context["engagement_id"], context["run_id"], declarations
    )
    completed = ledger.complete_run(
        client_root, context["engagement_id"], context["run_id"]
    )
    assert completed["run"]["status"] == "completed"
    return ledger


def _save_review_delivery(output, memo, record, prose, workflow, follow_up):
    """Replay the reviewed fixture's ordinary source challenge and delivery step."""
    note = prose["review_note"]
    if workflow == "aml" or follow_up:
        note += (
            "\n\n"
            + prose[workflow + ("_demo_correction" if not follow_up else "_correction")]
        )
    (output / "codex_run_review.md").write_text(note + "\n", encoding="utf-8")
    (output / "artifact_card.md").write_text(
        f"[Memo]({memo})\n\n[{prose['record']}]({record})\n\n"
        f"[{prose['review_link']}]({output / 'codex_run_review.md'})\n",
        encoding="utf-8",
    )


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
@pytest.mark.parametrize("lesson_phase", ["demo", "practice"])
def test_assetti_kit_runs_assessment_and_same_case_follow_up(
    tmp_path, monkeypatch, record_property, language, lesson_phase
):
    run = _bound_case(
        tmp_path,
        monkeypatch,
        "adeguati-assetti",
        "adeguati-assetti",
        "demo",
        language=language,
    )
    prose = _read(ROOT / "tests/fixtures/teaching_reviews/narratives.json")[language]
    client_root = tmp_path / "case"
    words = {
        "it": [
            "Responsabilità, chiusura mensile e uso delle informazioni per gli incassi; altri processi non valutati.",
            "Officina Arco: una sede e dodici addetti, con contabilità esterna e responsabilità interne descritte.",
            "Valutare il processo esistente e una sostituzione praticabile, senza imporre una struttura più complessa del necessario.",
            "La procedura definisce invio, consegna e riesame; non nomina un sostituto di Sara.",
            "Nel ciclo di febbraio sono documentati invio, consegna e decisione; la promessa del cliente non prova il pagamento.",
            "Nel ciclo di marzo invio e consegna sono tardivi dopo l’assenza di Sara; il riesame non è ancora annotato.",
            "Il primo ciclo dimostra uso delle informazioni in quel periodo; la continuità e la sostituzione rimangono da verificare.",
            "La nuova evidenza mostra un ritardo effettivo nel ciclo esaminato. Il precedente ciclo riuscito resta valido, ma non risolve la debolezza emersa durante l’assenza.",
            "Concordare la sostituzione per invio documenti e accesso allo scadenziario e verificarla nel prossimo ciclo.",
            "Proposta per Elena, da concordare con Sara e lo studio; nessun incarico assegnato automaticamente.",
            "Da concordare prima della prossima assenza e verificare al successivo riepilogo.",
            "Conferma delle responsabilità, accesso e prova datata di invio, consegna e riesame.",
            "Chi può sostituire Sara e con quali accessi durante un’assenza?",
            "Gli altri processi e la continuità fuori dai cicli forniti non sono valutati. Nessuna certificazione o decisione professionale è registrata.",
        ],
        "en": [
            "Responsibilities, monthly closing and collection information; other processes are not assessed.",
            "Officina Arco: one site and twelve staff, with external accounting and stated internal responsibilities.",
            "Assess the existing process and a workable substitute without imposing unnecessary complexity.",
            "The procedure defines dispatch, delivery and review but names no substitute for Sara.",
            "The February cycle records dispatch, delivery and a decision; a customer promise does not prove payment.",
            "March-cycle dispatch and delivery are late after Sara’s absence; the review is not yet recorded.",
            "The first cycle demonstrates use of information in that period; continuity and substitution still need evidence.",
            "New evidence shows an actual delay in the reviewed cycle. The earlier successful cycle remains valid but does not resolve the weakness during absence.",
            "Agree cover for document dispatch and access to due dates, then check it in the next cycle.",
            "Proposed for Elena, to agree with Sara and the practice; no task is assigned automatically.",
            "To agree before the next absence and check at the next monthly summary.",
            "Confirmed responsibilities, access and dated dispatch, delivery and review evidence.",
            "Who can cover Sara’s tasks, and with which access, during an absence?",
            "Other processes and continuity outside supplied cycles are not assessed. No certification or professional decision is recorded.",
        ],
        "fr": [
            "Responsabilités, clôture mensuelle et informations de recouvrement ; autres processus non évalués.",
            "Officina Arco : un site, douze collaborateurs, comptabilité externe et responsabilités internes décrites.",
            "Évaluer le processus existant et un remplacement praticable sans complexité inutile.",
            "La procédure définit envoi, livraison et revue sans nommer de remplaçant de Sara.",
            "Le cycle de février documente envoi, livraison et décision ; une promesse ne prouve pas le paiement.",
            "Dans le cycle de mars, envoi et livraison sont tardifs après l’absence de Sara ; la revue n’est pas encore notée.",
            "Le premier cycle démontre l’utilisation des informations sur cette période ; continuité et remplacement restent à vérifier.",
            "Les nouvelles preuves montrent un retard réel sur le cycle examiné. Le succès précédent reste valable sans résoudre la faiblesse pendant l’absence.",
            "Convenir du remplacement pour l’envoi et l’accès à l’échéancier puis le vérifier au prochain cycle.",
            "Proposition pour Elena, à convenir avec Sara et le cabinet ; aucune tâche attribuée automatiquement.",
            "À convenir avant la prochaine absence et vérifier à la prochaine synthèse.",
            "Responsabilités confirmées, accès et preuves datées d’envoi, livraison et revue.",
            "Qui peut remplacer Sara et avec quels accès pendant son absence ?",
            "Autres processus et continuité hors cycles fournis non évalués. Aucune certification ni décision professionnelle enregistrée.",
        ],
        "de": [
            "Zuständigkeiten, Monatsabschluss und Forderungsinformationen; andere Prozesse nicht beurteilt.",
            "Officina Arco: ein Standort, zwölf Beschäftigte, externe Buchhaltung und beschriebene interne Zuständigkeiten.",
            "Vorhandenen Prozess und praktikable Vertretung ohne unnötige Komplexität beurteilen.",
            "Das Verfahren regelt Versand, Lieferung und Prüfung, nennt aber keine Vertretung für Sara.",
            "Der Februarzyklus dokumentiert Versand, Lieferung und Entscheidung; die Kundenzusage belegt keine Zahlung.",
            "Im Märzzyklus erfolgen Versand und Lieferung nach Saras Abwesenheit verspätet; die Prüfung ist noch nicht dokumentiert.",
            "Der erste Zyklus belegt Informationsnutzung in diesem Zeitraum; Kontinuität und Vertretung bleiben zu prüfen.",
            "Neue Nachweise zeigen eine tatsächliche Verzögerung im geprüften Zyklus. Der frühere Erfolg bleibt gültig, löst aber die Schwäche während der Abwesenheit nicht.",
            "Vertretung für Unterlagenversand und Fälligkeitszugriff vereinbaren und im nächsten Zyklus prüfen.",
            "Vorschlag für Elena, mit Sara und der Kanzlei abzustimmen; keine automatische Aufgabenzuweisung.",
            "Vor der nächsten Abwesenheit vereinbaren und bei der nächsten Monatsübersicht prüfen.",
            "Bestätigte Zuständigkeiten, Zugriff und datierte Versand-, Liefer- und Prüfnachweise.",
            "Wer kann Sara während einer Abwesenheit mit welchen Zugriffsrechten vertreten?",
            "Andere Prozesse und Kontinuität außerhalb gelieferter Zyklen nicht beurteilt. Keine Zertifizierung oder fachliche Entscheidung erfasst.",
        ],
        "es": [
            "Responsabilidades, cierre mensual e información de cobros; otros procesos no evaluados.",
            "Officina Arco: una sede, doce personas, contabilidad externa y responsabilidades internas descritas.",
            "Evaluar proceso existente y sustitución viable sin complejidad innecesaria.",
            "El procedimiento define envío, entrega y revisión sin nombrar sustituto de Sara.",
            "El ciclo de febrero documenta envío, entrega y decisión; una promesa no prueba el pago.",
            "En el ciclo de marzo, envío y entrega se retrasan tras la ausencia de Sara; la revisión aún no consta.",
            "El primer ciclo demuestra uso de información en ese período; continuidad y sustitución requieren pruebas.",
            "Las nuevas pruebas muestran un retraso real en el ciclo revisado. El éxito anterior sigue válido, pero no resuelve la debilidad durante la ausencia.",
            "Acordar sustitución para enviar documentos y acceder a vencimientos y comprobarla en el próximo ciclo.",
            "Propuesta para Elena, a acordar con Sara y el despacho; no se asignan tareas automáticamente.",
            "Acordar antes de la próxima ausencia y verificar con el siguiente resumen.",
            "Responsabilidades confirmadas, acceso y pruebas fechadas de envío, entrega y revisión.",
            "¿Quién puede sustituir a Sara y con qué accesos durante una ausencia?",
            "Otros procesos y continuidad fuera de los ciclos aportados no evaluados. No se registra certificación ni decisión profesional.",
        ],
    }[language]
    previous = None
    for phase in ("demo", "practice"):
        output = Path(run["output_dir"])
        input_root = Path(run["context"]["run_root"]) / "inputs"
        paths = [Path(item["path"]) for item in run["context"]["input_bindings"]]
        sources = [
            {
                "id": f"S{index}",
                "path": path.relative_to(input_root).as_posix(),
                "title": path.name,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            }
            for index, path in enumerate(paths)
        ]
        by_name = {
            path.name: source["id"] for path, source in zip(paths, sources, strict=True)
        }
        policy_citation = {
            "source_id": by_name[f"procedure-{language}.md"],
            "locator": prose["paragraph"],
        }
        operation_citation = {
            "source_id": by_name[
                f"{'update' if previous else 'operation'}-{language}.md"
            ],
            "locator": prose["paragraph"],
        }
        company_citation = {
            "source_id": by_name[f"company-{language}.md"],
            "locator": prose["paragraph"],
        }
        finding_citations = [policy_citation, operation_citation, company_citation]
        operation = words[5 if previous else 4]
        assessment = words[7 if previous else 6]
        review = {
            "schema_version": 1,
            "jurisdiction": "IT",
            "language": language,
            "as_of": "2026-04-30" if previous else "2026-03-31",
            "scope": words[0],
            "company_context": words[1],
            "proportionality_basis": words[2],
            "sources": sources,
            "legal_basis": [
                {
                    "title": "CNDCEC — Assetti organizzativi, amministrativi e contabili: check-list operative",
                    "url": "https://commercialisti.it/documenti-studio/assetti-organizzativi-amministrativi-e-contabili-check-list-operative/",
                    "locator": prose["assetti_locator"],
                    "checked_at": "2026-09-14",
                    "applicability": prose["source_basis"],
                }
            ],
            "observations": [
                {
                    "id": "O1",
                    "area": prose["monthly"],
                    "description": words[3],
                    "proportionality": words[2],
                    "assessment": words[3],
                    "evidence_state": "documented",
                    "citations": [policy_citation, company_citation],
                },
                {
                    "id": "O2",
                    "area": prose["cycle"],
                    "description": operation,
                    "proportionality": words[2],
                    "assessment": assessment,
                    "evidence_state": "operating_evidence",
                    "citations": [operation_citation],
                },
            ],
            "findings": [
                {
                    "id": "F1",
                    "observation_ids": ["O1", "O2"],
                    "observation": words[3],
                    "interpretation": assessment,
                    "alternatives": prose[
                        "practice_alternative" if previous else "demo_alternative"
                    ],
                    "follow_up": words[12],
                    "citations": finding_citations,
                }
            ],
            "actions": [
                {
                    "id": "A1",
                    "finding_ids": ["F1"],
                    "proposal": words[8],
                    "owner": words[9],
                    "timing": words[10],
                    "priority_reason": prose["risk"],
                    "completion_evidence_needed": words[11],
                    "status": "proposed",
                }
            ],
            "assessment": assessment,
            "assessment_citations": finding_citations,
            "limitations": words[13],
            "intelligent_review": {
                "version": 1,
                "coverage": [
                    {
                        "id": "C1",
                        "area": prose["coverage"],
                        "status": "assessed",
                        "reason": words[0],
                        "observation_ids": ["O1", "O2"],
                    },
                    {
                        "id": "C2",
                        "area": prose["excluded"],
                        "status": "unresolved",
                        "reason": words[13],
                        "observation_ids": [],
                    },
                ],
                "processes": [
                    {
                        "id": "P1",
                        "process": words[0],
                        "risk": prose["risk"],
                        "responsibility": prose["responsibility"],
                        "control": words[3],
                        "information_flow": prose["flow"],
                        "operation": operation,
                        "gap": words[12],
                        "observation_ids": ["O1", "O2"],
                    }
                ],
                "questions": [
                    {
                        "id": "Q1",
                        "question": words[12],
                        "why_it_matters": prose["risk"],
                        "evidence_needed": words[11],
                        "status": prose[
                            "practice_answer" if previous else "unanswered"
                        ],
                        "observation_ids": ["O1", "O2"],
                    }
                ],
                "chronology": [
                    {
                        "id": "T1",
                        "event_date": (
                            "2026-04-28" if previous else "2026-03-24 / 2026-03-25"
                        ),
                        "known_at": (
                            prose["received"] if previous else "2026-03-24 / 2026-03-25"
                        ),
                        "recipient": (
                            prose["recipient"] if previous else "Elena Bianchi"
                        ),
                        "event": operation,
                        "response": prose[
                            "practice_response" if previous else "demo_response"
                        ],
                        "uncertainty": (
                            prose["temporal_limit"] if previous else words[4]
                        ),
                        "observation_ids": ["O2"],
                    }
                ],
                "decision_brief": assessment + " " + words[8],
                "action_ids": ["A1"],
                "next_review": words[10] + " " + words[11],
            },
        }
        if previous:
            review["previous"] = {
                "source_id": by_name[previous[0].name],
                "record_sha256": previous[2],
            }
            review["changes_since_previous"] = words[7]
            review["prior_action_review"] = {
                "A1": {
                    "status": "open",
                    "assessment": words[7],
                    "citations": [operation_citation],
                    "current_action_ids": ["A1"],
                }
            }
        request = output / "review_input.json"
        _write(request, review)
        _run(
            "plugins/adeguati-assetti/scripts/assetti_review.py",
            "--client-engagement",
            run["context_path"],
            "--review",
            request,
        )
        path = next(output.glob("adeguati-assetti-*.json"))
        record = _read(path)
        assert record["status"] == "draft_for_review"
        assert next(output.glob("adeguati-assetti-*.md")).stat().st_size > 2000
        memo_path = next(output.glob("adeguati-assetti-*.md"))
        _save_review_delivery(
            output, memo_path, path, prose, "assetti", previous is not None
        )
        memo_text = memo_path.read_text()
        assert "../inputs/" in memo_text
        assert "**observation:**" not in memo_text
        assert "draft_for_review" not in memo_text
        assert "Test-only" not in memo_text
        assert record["review"]["language"] == language
        ledger = _complete_teaching_case(run, client_root)
        if previous:
            assert record["previous_record_sha256"] == previous[2]
            assert previous[0].read_bytes() == previous[1]
            assert record["review"]["prior_action_review"]["A1"]["status"] == "open"
        elif lesson_phase == "practice":
            previous = (path, path.read_bytes(), record["record_sha256"])
            client, engagement = (
                run["context"]["client_id"],
                run["context"]["engagement_id"],
            )
            practice = list((tmp_path / "kit/files/practice").glob("*.md"))
            imports = [
                ledger.import_document(client_root, client, engagement, item, "source")
                for item in [*practice, path]
            ]
            prepared = ledger.prepare_run(
                client_root,
                client,
                engagement,
                "adeguati-assetti",
                run["context"]["workflow_version"],
                input_ids=[item["receipt"]["input_id"] for item in imports],
            )
            run = ledger.start_run(client_root, engagement, prepared["run"]["run_id"])
        if phase == lesson_phase:
            break
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="adeguati-assetti",
        language=language,
        phase=lesson_phase,
    )


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
@pytest.mark.parametrize("lesson_phase", ["demo", "practice"])
def test_aml_kit_saves_initial_and_linked_current_review(
    tmp_path, monkeypatch, record_property, language, lesson_phase
):
    """Create both real records; fixture prose never becomes learner approval."""
    run = _bound_case(
        tmp_path, monkeypatch, "aml-review", "aml-review", "demo", language=language
    )
    prose = _read(ROOT / "tests/fixtures/teaching_reviews/narratives.json")[language]
    client_root = tmp_path / "case"
    texts = {
        "it": [
            "Revisione del dossier fittizio di Officina Arco e del finanziamento soci.",
            "Le quote e i poteri sono dichiarati; mancano riscontri indipendenti.",
            "Il dossier non consente ancora una conclusione sull’identificazione e il controllo.",
            "Le informazioni potrebbero essere corrette; l’assenza di documenti non prova irregolarità.",
            "Verificare identità, poteri e catena di controllo con evidenze pertinenti.",
            "Contratto e pagamento del finanziamento coincidono per parti e importo. Le persone e il controllo richiedono verifica; nessun giudizio di sospetto o accettazione è espresso.",
            "La nuova dichiarazione include Marta Riva e cambia la quota di Paolo. Il finanziamento rimane documentato dalle stesse fonti; servono riscontri sulla cessione e sulla nuova socia.",
        ],
        "en": [
            "Review of Officina Arco’s fictional file and shareholder loan.",
            "Ownership and authority are declared; independent corroboration is missing.",
            "The file does not yet support a conclusion on identification and control.",
            "The statements may be correct; missing documents do not prove wrongdoing.",
            "Verify identity, authority and control links with relevant evidence.",
            "Loan agreement and payment agree on parties and amount. People and control need verification; no suspicion or acceptance decision is expressed.",
            "The new declaration adds Marta Riva and changes Paolo’s share. The loan remains supported by the same sources; transfer and new-owner evidence is needed.",
        ],
        "fr": [
            "Revue du dossier fictif d’Officina Arco et du prêt d’associé.",
            "Participations et pouvoirs sont déclarés ; les corroborations indépendantes manquent.",
            "Le dossier ne permet pas encore de conclure sur l’identification et le contrôle.",
            "Les déclarations peuvent être exactes ; les pièces manquantes ne prouvent pas une irrégularité.",
            "Vérifier identité, pouvoirs et liens de contrôle avec les preuves pertinentes.",
            "Contrat et paiement concordent sur parties et montant. Personnes et contrôle restent à vérifier ; aucune décision de soupçon ou d’acceptation n’est exprimée.",
            "La nouvelle déclaration ajoute Marta Riva et modifie la participation de Paolo. Le prêt reste étayé par les mêmes sources ; il faut vérifier cession et nouvelle associée.",
        ],
        "de": [
            "Prüfung der fiktiven Akte von Officina Arco und des Gesellschafterdarlehens.",
            "Beteiligungen und Befugnisse sind angegeben; unabhängige Bestätigungen fehlen.",
            "Die Akte erlaubt noch keine Schlussfolgerung zu Identifikation und Kontrolle.",
            "Die Angaben können richtig sein; fehlende Unterlagen beweisen keinen Verstoß.",
            "Identität, Befugnisse und Kontrollverbindungen anhand geeigneter Nachweise prüfen.",
            "Darlehensvertrag und Zahlung stimmen bei Parteien und Betrag überein. Personen und Kontrolle sind zu prüfen; keine Verdachts- oder Annahmeentscheidung wird getroffen.",
            "Die neue Erklärung ergänzt Marta Riva und ändert Paolos Anteil. Das Darlehen bleibt durch dieselben Quellen belegt; Übertragung und neue Gesellschafterin benötigen Nachweise.",
        ],
        "es": [
            "Revisión del expediente ficticio de Officina Arco y del préstamo de socio.",
            "Participaciones y poderes están declarados; faltan corroboraciones independientes.",
            "El expediente aún no permite concluir sobre identificación y control.",
            "Las declaraciones podrían ser correctas; la ausencia de documentos no prueba irregularidad.",
            "Verificar identidad, poderes y vínculos de control con pruebas pertinentes.",
            "Contrato y pago coinciden en partes e importe. Personas y control requieren verificación; no se expresa decisión de sospecha ni aceptación.",
            "La nueva declaración añade a Marta Riva y cambia la cuota de Paolo. El préstamo sigue respaldado por las mismas fuentes; se requieren pruebas sobre transmisión y nueva socia.",
        ],
    }[language]
    previous = None
    for phase in ("demo", "practice"):
        output = Path(run["output_dir"])
        input_root = Path(run["context"]["run_root"]) / "inputs"
        paths = [Path(item["path"]) for item in run["context"]["input_bindings"]]
        sources = [
            {
                "id": f"S{index}",
                "path": path.relative_to(input_root).as_posix(),
                "title": path.name,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            }
            for index, path in enumerate(paths)
        ]
        by_name = {
            path.name: source["id"] for path, source in zip(paths, sources, strict=True)
        }
        ownership = by_name[
            f"ownership{'-update' if phase == 'practice' else ''}-{language}.md"
        ]
        loan = by_name[f"loan-{language}.md"]
        client_source = by_name[f"client-{language}.md"]
        review = {
            "schema_version": 1,
            "jurisdiction": "IT",
            "language": language,
            "as_of": "2026-08-31" if phase == "demo" else "2026-09-10",
            "scope": texts[0],
            "sources": sources,
            "legal_basis": [
                {
                    "title": "CNDCEC — Antiriciclaggio",
                    "url": "https://commercialisti.it/norme-per-la-professione/norme-tecniche/antiriciclaggio/",
                    "locator": prose["aml_locator"],
                    "checked_at": "2026-09-14",
                    "applicability": prose["source_basis"],
                }
            ],
            "findings": [
                {
                    "id": "F1",
                    "observation": prose[f"aml_{phase}_ownership"],
                    "interpretation": texts[2],
                    "alternatives": texts[3],
                    "follow_up": prose[f"aml_{phase}_followup"],
                    "citations": [
                        {"source_id": ownership, "locator": prose["paragraph"]}
                    ],
                }
            ],
            "assessment": texts[5] + (" " + texts[6] if previous else ""),
            "assessment_citations": [
                {"source_id": client_source, "locator": prose["paragraph"]},
                {"source_id": loan, "locator": prose["paragraph"]},
                {"source_id": ownership, "locator": prose["paragraph"]},
            ],
            "limitations": prose["aml_limit"],
        }
        review["findings"].append(
            {
                "id": "F2",
                "observation": prose["loan_observation"],
                "interpretation": prose["loan_interpretation"],
                "alternatives": prose["loan_alternative"],
                "follow_up": prose["loan_followup"],
                "citations": [{"source_id": loan, "locator": prose["paragraph"]}],
            }
        )
        if previous:
            review["previous"] = {
                "source_id": by_name[previous[0].name],
                "record_sha256": previous[2],
            }
            review["changes_since_previous"] = texts[6]
        request = output / "review_input.json"
        _write(request, review)
        _run(
            "plugins/aml-review/scripts/aml_review.py",
            "--client-engagement",
            run["context_path"],
            "--review",
            request,
        )
        record_path = next(output.glob("aml-review-*.json"))
        record = _read(record_path)
        assert record["status"] == "draft_for_review"
        assert record["calculation"] is None
        assert next(output.glob("aml-review-*.md")).stat().st_size > 1000
        memo_path = next(output.glob("aml-review-*.md"))
        _save_review_delivery(
            output, memo_path, record_path, prose, "aml", previous is not None
        )
        memo_text = memo_path.read_text()
        assert "../inputs/" in memo_text
        assert "**observation:**" not in memo_text
        assert "draft_for_review" not in memo_text
        assert "Test-only" not in memo_text
        assert record["review"]["language"] == language
        ledger = _complete_teaching_case(run, client_root)
        if previous:
            assert record["previous_record_sha256"] == previous[2]
            assert previous[0].read_bytes() == previous[1]
            assert "Marta Riva" in next(output.glob("aml-review-*.md")).read_text()
        elif lesson_phase == "practice":
            previous = (record_path, record_path.read_bytes(), record["record_sha256"])
            client = run["context"]["client_id"]
            engagement = run["context"]["engagement_id"]
            practice = list((tmp_path / "kit/files/practice").glob("*.md"))
            imports = [
                ledger.import_document(client_root, client, engagement, path, "source")
                for path in [*practice, record_path]
            ]
            prepared = ledger.prepare_run(
                client_root,
                client,
                engagement,
                "aml-review",
                run["context"]["workflow_version"],
                input_ids=[item["receipt"]["input_id"] for item in imports],
            )
            run = ledger.start_run(client_root, engagement, prepared["run"]["run_id"])
        if phase == lesson_phase:
            break
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="aml-review",
        language=language,
        phase=lesson_phase,
    )


REPORT_NARRATIVES = {
    "it": [
        "La relazione riassume i tre prospetti forniti per Officina Arco. Il risultato operativo e il flusso netto di cassa sono positivi; la situazione patrimoniale quadra. Il perimetro è gestionale: cause commerciali, imposte e altre sezioni non sono documentate.",
        "I ricavi superano il costo del venduto e le spese operative. Il risultato riportato deriva dalle righe del prospetto, senza sommare nuovamente il subtotale già presente. Le cause commerciali richiedono informazioni ulteriori.",
        "Le attività sono pareggiate da passività e patrimonio netto. Il controllo di quadratura non è una posta aggiuntiva e non prova da solo la completezza dei saldi.",
        "Il flusso operativo supera le uscite nette per investimenti e finanziamenti. La variazione netta è coerente con la cassa iniziale indicata nel contesto e la cassa finale del prospetto.",
    ],
    "en": [
        "This report summarises the three statements supplied for Officina Arco. Operating profit and net cash flow are positive, and the balance sheet balances. This is a management report: commercial causes, taxes and other sections are not documented.",
        "Revenue exceeds cost of sales and operating expenses. The reported result comes from the statement rows, without adding the existing subtotal again. Explaining the commercial causes requires further information.",
        "Assets balance against liabilities and equity. The balance check is not an additional item and does not by itself establish that all balances are complete.",
        "Operating cash flow exceeds the net outflows for investment and financing. The net movement reconciles the opening cash supplied in the context to the closing cash in the statement.",
    ],
    "fr": [
        "Ce rapport résume les trois états fournis pour Officina Arco. Le résultat opérationnel et le flux net de trésorerie sont positifs, et le bilan est équilibré. Le périmètre est celui d’un rapport de gestion : les causes commerciales, les impôts et les autres rubriques ne sont pas documentés.",
        "Le chiffre d’affaires dépasse le coût des ventes et les charges opérationnelles. Le résultat découle des lignes de l’état, sans additionner à nouveau le sous-total existant. Les causes commerciales nécessitent des informations complémentaires.",
        "Les actifs sont équilibrés par les passifs et les capitaux propres. Le contrôle d’équilibre n’est pas un poste supplémentaire et ne prouve pas à lui seul l’exhaustivité des soldes.",
        "Le flux opérationnel dépasse les sorties nettes liées aux investissements et au financement. La variation nette rapproche la trésorerie initiale indiquée dans le contexte de la trésorerie finale de l’état.",
    ],
    "de": [
        "Dieser Bericht fasst die drei bereitgestellten Aufstellungen für Officina Arco zusammen. Das operative Ergebnis und der Nettozahlungsstrom sind positiv; die Bilanz ist ausgeglichen. Der Bericht dient der internen Steuerung. Geschäftliche Ursachen, Steuern und weitere Bereiche sind nicht dokumentiert.",
        "Die Umsatzerlöse übersteigen die Umsatzkosten und betrieblichen Aufwendungen. Das ausgewiesene Ergebnis ergibt sich aus den Positionen der Aufstellung, ohne die vorhandene Zwischensumme nochmals zu addieren. Zur Erklärung der geschäftlichen Ursachen sind weitere Informationen erforderlich.",
        "Die Vermögenswerte entsprechen den Verbindlichkeiten und dem Eigenkapital. Die Kontrollsumme ist keine zusätzliche Position und belegt allein nicht die Vollständigkeit der Salden.",
        "Der operative Zahlungsstrom übersteigt die Nettoabflüsse für Investitionen und Finanzierung. Die Nettoveränderung stimmt den im Kontext angegebenen Anfangsbestand mit dem Endbestand der liquiden Mittel ab.",
    ],
    "es": [
        "Este informe resume los tres estados facilitados para Officina Arco. El resultado operativo y el flujo neto de caja son positivos, y el balance cuadra. El alcance es de gestión: no se documentan las causas comerciales, los impuestos ni las demás secciones.",
        "Los ingresos superan el coste de ventas y los gastos operativos. El resultado procede de las líneas del estado, sin volver a sumar el subtotal existente. Explicar las causas comerciales requiere información adicional.",
        "Los activos se equilibran con los pasivos y el patrimonio neto. El control de cuadre no es una partida adicional ni demuestra por sí solo que los saldos estén completos.",
        "El flujo operativo supera las salidas netas por inversión y financiación. La variación neta concilia la caja inicial indicada en el contexto con la caja final del estado.",
    ],
}


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
@pytest.mark.parametrize(
    "phase,end,profit,cash_movement",
    [
        ("demo", "2026-02-28", "53000", "7000"),
        ("practice", "2026-03-31", "85000", "15000"),
    ],
)
def test_financial_report_kit_builds_current_source_bound_document(
    tmp_path, monkeypatch, language, phase, end, profit, cash_movement, record_property
):
    """Review this known fixture in the test, then build the normal report."""
    monkeypatch.syspath_prepend(str(ROOT / "tests/plugins"))
    from test_report_builder_plugin import _numeric_review_args, load_core

    run = _bound_case(
        tmp_path,
        monkeypatch,
        "financial-report-builder",
        "report-builder",
        phase,
        language=language,
    )
    source = next(
        Path(item["path"])
        for item in run["context"]["input_bindings"]
        if Path(item["path"]).suffix == ".xlsx"
    )
    original = source.read_bytes()
    work = Path(run["output_dir"])
    output = work / "report"
    core = load_core()
    inspection = core.inspect_inputs(
        source,
        work / "inspection",
        language=language,
        document_language=language,
        report_type="management_report",
    )
    recipe = inspection.suggested_recipe
    recipe["entity"] = "Officina Arco"
    recipe["period"] = f"2026-01-01 to {end}"
    # Model-authored readings of these exact fictional statements, replayed
    # only as test inputs. The distributed kit contains no finished narrative.
    recipe["executive_summary"] = REPORT_NARRATIVES[language][0]
    for index, section in enumerate(
        ("income_statement", "balance_sheet", "cash_flow"), start=1
    ):
        recipe["sections"][section]["codex_comment"] = REPORT_NARRATIVES[language][
            index
        ]
    # These are explicit interpretations of the authored fixture, not a
    # generic classifier, a saved learner decision or a shipped approval.
    assignments = {
        "Income Statement": ("income_statement", 5),
        "Balance Sheet": ("balance_sheet", 9),
        "Cash Flow": ("cash_flow", 5),
    }
    for section in recipe["sections"].values():
        section["assigned_table"] = ""
    for table in inspection.inspection["tables"]:
        section, subtotal_row = assignments[table["sheet_name"]]
        recipe["sections"][section]["assigned_table"] = table["table_id"]
        recipe = core.review_numeric_measure_columns(
            inspection.inspection,
            recipe,
            section_key=section,
            **_numeric_review_args(
                inspection.inspection,
                table["table_id"],
                ["Amount EUR"],
                excluded_cell_rows={"Amount EUR": {subtotal_row}},
            ),
            reviewer_ref="reviewer.teaching_fixture",
            reviewed_on="2026-09-14",
            numeric_locale="en",
            currency="EUR",
            unit="currency",
            scale="1",
            parse_policy="strict_all_nonblank_v1",
        )
    recipe_path = work / "reviewed_recipe.json"
    _write(recipe_path, recipe)
    core.build_report(
        source,
        output,
        recipe_path=recipe_path,
        language=language,
        document_language=language,
        report_type="management_report",
        run_id=run["context"]["run_id"],
        client_engagement=run["context"],
    )
    evidence = _read(output / "numeric_evidence_ledger.json")
    totals = [
        entry for entry in evidence["entries"] if entry["evidence_id"].endswith(".sum")
    ]
    details = [
        entry for entry in evidence["entries"] if ".row_" in entry["evidence_id"]
    ]
    assert {entry["value"] for entry in totals} == {
        profit,
        "0",
        cash_movement,
    }
    assert len(totals) == 3
    assert len(details) == 13
    assert all(
        entry["source"]["value"] == entry["value"] for entry in evidence["entries"]
    )
    assert {
        item["artifact_ref"]
        for entry in evidence["entries"]
        for item in entry["outputs"]
    } == {
        "output.report_docx",
        "output.report_draft",
        "output.report_tables",
    }
    assert (output / "report.docx").stat().st_size > 10000
    assert (output / "report_tables.xlsx").is_file()
    assert (output / "review_payload.json").is_file()
    assert (output / "final_artifacts.json").is_file()
    assert source.read_bytes() == original
    markdown = (output / "report_draft.md").read_text(encoding="utf-8")
    assert REPORT_NARRATIVES[language][0] in markdown
    assert "[numeric source value withheld]" not in markdown
    assert "[excluded from calculation]" not in markdown
    assert {
        "it": "Escluso dal totale",
        "en": "Excluded from total",
        "fr": "Exclu du total",
        "de": "Von der Summe ausgeschlossen",
        "es": "Excluido del total",
    }[language] in markdown
    _complete_teaching_case(run, tmp_path / "case")
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="financial-report-builder",
        language=language,
        phase=phase,
    )


@pytest.mark.parametrize("product", ["vera", "clara"])
@pytest.mark.parametrize("language", ["it", "en"])
@pytest.mark.parametrize("phase", ["demo", "practice"])
@prepared_kit("vera/business-planning", "clara/business-planning")
def test_business_plan_kit_runs_and_revises_the_owning_product_case(
    tmp_path, monkeypatch, record_property, product, language, phase
):
    from tests.plugins._business_teaching import planning_case, write_delivery

    monkeypatch.syspath_prepend(str(ROOT / "plugins/_shared/vendor/modules"))
    from courseware.library import CourseLibrary

    previous = None
    earlier_outputs = {}
    if product == "vera":
        run = _bound_case(
            tmp_path,
            monkeypatch,
            "business-planning",
            "business-planning",
            "demo",
            language=language,
        )
    else:
        monkeypatch.syspath_prepend(str(ROOT / "plugins/clara/scripts"))
        import advisor_case_core as clara

        kit = CourseLibrary(ROOT / "plugins/clara", {"business-planning"}).render(
            "business-planning", language, tmp_path / "kit"
        )
        workspace = tmp_path / "clara-case"
        clara.initialize_case(
            workspace,
            client="Ciclo Arco",
            project="Business planning",
            objective="Assess the fictional pilot",
            audience="internal",
            output_language=language,
        )
        request = workspace / "business_plan_case.json"
        base_output = workspace / "business-plan"
        (workspace / "inputs").mkdir()
        (workspace / ".clara-onboarding-local-only").write_text(
            "Fictional local lesson\n"
        )
    for current_phase in (("demo",) if phase == "demo" else ("demo", "practice")):
        if product == "vera":
            output = Path(run["output_dir"])
            source_root = Path(run["context"]["run_root"]) / "inputs"
            paths = [Path(item["path"]) for item in run["context"]["input_bindings"]]
            request = output / "business_plan_case.json"
            report = output / "plan"
            args = ["--client-engagement", run["context_path"]]
            entry = "run_business_plan.py"
        else:
            source_root = workspace
            paths = []
            for source in kit[
                "source_files" if current_phase == "demo" else "practice_files"
            ]:
                path = workspace / "inputs" / Path(source).name
                if path.exists():
                    assert path.read_bytes() == Path(source).read_bytes()
                path.write_bytes(Path(source).read_bytes())
                paths.append(path)
            if previous:
                parent = workspace / "inputs/prior-plan.json"
                parent.write_bytes(previous[1])
                paths.append(parent)
            report = base_output / current_phase
            output = report
            args = ["--case-workspace", workspace]
            entry = "run_strategic_plan.py"
        case = planning_case(
            paths, source_root, language, previous[2] if previous else None
        )
        _write(request, case)
        if product == "clara":
            for source in case["sources"]:
                clara.register_material(
                    workspace,
                    source_root / source["path"],
                    summary="Fictional local teaching input; not professionally reviewed.",
                    source_metadata={
                        "sha256": source["sha256"],
                        "role": source["role"],
                    },
                )
        result = subprocess.run(
            [
                sys.executable,
                str(ROOT / "plugins/business-planning/scripts" / entry),
                "--case",
                str(request),
                "--source-root",
                str(source_root),
                "--output-dir",
                str(report),
                *map(str, args),
            ],
            capture_output=True,
            text=True,
            check=False,
            timeout=60,
        )
        assert result.returncode == 2, result.stdout + result.stderr
        plan_path = report / "business_plan.json"
        assert plan_path.is_file(), result.stdout + result.stderr
        plan = _read(plan_path)
        assert plan["status"] == "partial"
        assert len(plan["accepted_narrative"]) == 10
        assert all(n["provisional"] for n in plan["accepted_narrative"])
        assert plan["case"]["review"] == {"status": "pending"}
        assert (
            plan["calculations"]["pilot/2027-01/commercial_operating_result"]["value"]
            == "-260"
        )
        assert (
            plan["calculations"]["pilot/2027-03/commercial_operating_result"]["value"]
            == "280"
        )
        assert plan["case"]["assessment"]["decision"] == (
            "redesign" if previous else "test"
        )
        html = (report / "business_plan_review.html").read_text(encoding="utf-8")
        assert (
            "Ipotesi economiche del test"
            if language == "it"
            else "Pilot economics assumptions"
        ) in html
        assert "-260" in html and "280" in html
        assert _read(report / "validation.json")["canonical_replay"] == "passed"
        receipt = _read(report / "execution_receipt.json")
        assert receipt["status"] == "partial" and receipt["pdf_error"] is None
        for artifact in receipt["outputs"]:
            assert (
                hashlib.sha256((report / artifact["path"]).read_bytes()).hexdigest()
                == artifact["sha256"]
            )
        write_delivery(
            output, report, plan, language, previous[0] if previous else None
        )
        if product == "vera":
            ledger = _complete_teaching_case(
                run,
                tmp_path / "case",
                artifact_ids={
                    "plan/business_plan_review.html": "business-planning.report"
                },
            )
        else:
            for path in (report / "business_plan_review.html", plan_path):
                clara.register_material(
                    workspace,
                    path,
                    summary="Actual compiled fictional plan; professional review pending.",
                    source_metadata={
                        "sha256": hashlib.sha256(path.read_bytes()).hexdigest()
                    },
                )
            assert clara.validate_case_workspace(workspace) == []
        for path, content in earlier_outputs.items():
            assert path.read_bytes() == content
        if previous:
            assert previous[0].read_bytes() == previous[1]
            assert (
                plan["planning_cycle"]["parent_content_sha256"]
                == previous[2]["content_sha256"]
            )
            assert plan["planning_cycle"]["withheld_narrative_ids"] == []
            assert ("sabato" if language == "it" else "Saturday") in html
            assert {
                key: (row["value"], row["unit"])
                for key, row in plan["calculations"].items()
            } == {
                key: (row["value"], row["unit"])
                for key, row in previous[2]["calculations"].items()
            }
            assert plan["case"]["case_id"] == previous[2]["case"]["case_id"]
            assert plan["case"]["cycle"]["id"] != previous[2]["case"]["cycle"]["id"]
            assert set(plan["case"]["cycle"]["reassessed_ids"]) == {
                n["id"] for n in plan["case"]["narrative"]
            }
        else:
            previous = (plan_path, plan_path.read_bytes(), plan)
            earlier_outputs = {
                p: p.read_bytes() for p in output.rglob("*") if p.is_file()
            }
            if product == "vera" and phase == "practice":
                client = run["context"]["client_id"]
                engagement = run["context"]["engagement_id"]
                parent = tmp_path / "prior-plan.json"
                parent.write_bytes(previous[1])
                imported = [
                    ledger.import_document(
                        tmp_path / "case", client, engagement, p, "source"
                    )
                    for p in [
                        *sorted((tmp_path / "kit/files/practice").iterdir()),
                        parent,
                    ]
                ]
                prepared = ledger.prepare_run(
                    tmp_path / "case",
                    client,
                    engagement,
                    "business-planning",
                    _read(ROOT / "plugins/vera/.codex-plugin/plugin.json")["version"],
                    input_ids=[i["receipt"]["input_id"] for i in imported],
                    purpose="Revise the same fictional planning case from new supplied evidence",
                )
                run = ledger.start_run(
                    tmp_path / "case", engagement, prepared["run"]["run_id"]
                )
    record_native_check(
        record_property,
        root=ROOT,
        product=product,
        workflow="business-planning",
        language=language,
        phase=phase,
    )


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
@pytest.mark.parametrize("phase,expected", [("demo", "38000"), ("practice", "31000")])
def test_financial_analysis_kit_runs_current_net_debt_pack(
    tmp_path, monkeypatch, record_property, language, phase, expected
):
    from tests.plugins._financial_teaching import write_financial_review

    monkeypatch.syspath_prepend(str(ROOT / "plugins/_shared/vendor/modules"))
    from vera_financial_analysis import (
        build_data_package_manifest,
        build_dataset_contract,
        build_fdd_case,
    )

    run = _bound_case(
        tmp_path,
        monkeypatch,
        "financial-analysis",
        "financial-analysis",
        phase,
        language=language,
    )
    source = next(
        Path(item["path"])
        for item in run["context"]["input_bindings"]
        if Path(item["path"]).suffix == ".csv"
    )
    original = source.read_bytes()
    with source.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 4
    dates = {row["as_of_date"] for row in rows}
    assert len(dates) == 1
    as_of = dates.pop()
    output = Path(run["output_dir"])
    prepared_source = output / source.name
    prepared_source.write_bytes(original)
    assert prepared_source.read_bytes() == original
    package = build_data_package_manifest(
        package_id="package.teaching.v1",
        snapshot_id="snapshot.teaching.v1",
        reporting_perimeter={
            "entity_refs": ["entity.arco"],
            "period_start": as_of,
            "period_end": as_of,
            "currency_refs": ["EUR"],
        },
        sensitivity="confidential",
        sources=[
            {
                "source_id": "source.balances",
                "artifact_ref": "artifact.balances",
                "file_name": source.name,
                "locator": source.name,
                "byte_count": len(original),
                "sha256": hashlib.sha256(original).hexdigest(),
                "snapshot_id": "snapshot.teaching.v1",
                "dataset_contract_ref": "dataset.balances.v1",
            }
        ],
    )
    dataset = build_dataset_contract(
        dataset_contract_id="dataset.balances.v1",
        dataset_id="balances",
        version="v1",
        grain="one identified balance at the stated date",
        keys=["row_id"],
        fields=[
            {
                "name": "row_id",
                "concept_id": "row.identity",
                "data_type": "text",
                "nullable": False,
                "unit": "identifier",
                "currency": None,
                "aggregation": "none",
                "period_role": "none",
            },
            {
                "name": "amount",
                "concept_id": "financial.amount",
                "data_type": "decimal",
                "nullable": False,
                "unit": "EUR_units",
                "currency": "EUR",
                "aggregation": "sum",
                "period_role": "period_end",
            },
        ],
        period={
            "calendar": "gregorian",
            "grain": "month",
            "start": as_of,
            "end": as_of,
        },
        source_artifact_refs=["artifact.balances"],
    )
    # Explicit interpretation of the supplied fictional request in this test.
    # Native learners review their choices; this receipt is never in a kit.
    classification = {
        "loan": "debt",
        "cash": "cash",
        "deposit": "excluded",
        "payables": "excluded",
    }
    review = {
        "status": "reviewed",
        "reviewed_on": "2026-09-14",
        "reviewer_ref": "reviewer.teaching_fixture",
        "basis": "Test-only interpretation of the authored loan and available-cash definition.",
    }
    case = build_fdd_case(
        case_id="case.teaching",
        scope_id="scope.arco",
        entity_refs=["entity.arco"],
        pack_id="net_debt",
        currency="EUR",
        unit="EUR_units",
        reporting_period={"start": as_of, "end": as_of},
        package=package,
        datasets=[dataset],
        request_id="request.net_debt.v1",
        review=review,
        reviewed_decisions=[{"decision_ref": "decision.definition", **review}],
        inputs={
            "as_of_date": as_of,
            "items": [
                {
                    "item_id": "item." + row["row_id"],
                    "economic_effect_id": "effect." + row["row_id"],
                    "description": row["description"],
                    "as_of_date": as_of,
                    "classification": classification[row["row_id"]],
                    "amount": row["amount"],
                    "included": classification[row["row_id"]] != "excluded",
                    "decision_ref": "decision.definition",
                    "evidence_refs": ["artifact.balances"],
                }
                for row in rows
            ],
        },
    )
    stack = case["contract_stack"]
    bundle = {
        "schema_version": "vera.fdd_execution_bundle.v2",
        "fdd_case": case,
        **{
            key: stack[key]
            for key in ("package", "datasets", "relationships", "crosswalks", "request")
        },
    }
    # Use the current canonical serializer, not a separately invented seal.
    from tests.plugins._financial_analysis_test_loader import (
        load_financial_analysis_scripts,
    )

    modules = load_financial_analysis_scripts(
        ROOT / "plugins/financial-analysis/scripts"
    )
    bundle["content_sha256"] = modules.kernel.canonical_json_sha256(bundle)
    case_path = output / "case.json"
    _write(case_path, bundle)
    prepared = output / "prepared"
    _run(
        "plugins/financial-analysis/scripts/run_pack.py",
        "--pack",
        "net_debt",
        "--case",
        case_path,
        "--output-dir",
        prepared,
        "--client-engagement",
        run["context_path"],
    )
    receipt = _read(prepared / "pack_execution_receipt.json")
    result = _read(prepared / "fdd_result.json")
    assert receipt["status"] == "passed"
    assert receipt["report_ready"] is False
    assert {item["metric_id"]: item["value"] for item in result["metrics"]}[
        "net_debt"
    ] == expected
    assert result["source_tie_out"]["status"] == "not_assessed"
    assert (prepared / "financial_analysis_contract_audit.json").is_file()
    assert (prepared / "model_use_manifest.json").is_file()
    assert source.read_bytes() == original
    note = write_financial_review(output, language)
    assert f"**{expected}**" in note.read_text(encoding="utf-8")
    _complete_teaching_case(run, tmp_path / "case")
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="financial-analysis",
        language=language,
        phase=phase,
    )


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
@pytest.mark.parametrize(
    "phase,baseline,actual", [("demo", 47000, 53000), ("practice", 77000, 85000)]
)
def test_variance_kit_runs_current_full_amount_comparison(
    tmp_path, monkeypatch, record_property, language, phase, baseline, actual
):
    from tests.plugins._variance_teaching import review_context, write_variance_review

    copy = review_context(language)
    run = _bound_case(
        tmp_path,
        monkeypatch,
        "variance-analysis",
        "variance-analysis",
        phase,
        language=language,
    )
    source = next(
        item["path"]
        for item in run["context"]["input_bindings"]
        if item["path"].endswith(".csv")
    )
    output = Path(run["output_dir"])
    scripts = "plugins/variance-analysis/scripts/"
    _run(
        scripts + "inspect_inputs.py",
        source,
        "--output-dir",
        output / "inspection",
        "--language",
        language,
        "--client-engagement",
        run["context_path"],
    )
    recipe = _read(output / "inspection/suggested_recipe.json")
    assert recipe["mappings"]["baseline_period"] == "PL"
    assert recipe["mappings"]["comparison_period"] == "AC"
    # Fixed interpretation of this authored fixture, not a production mapping
    # heuristic or a review approval distributed with the lesson.
    recipe["mappings"].update(
        amount_column="Amount", units_column=None, dimensions=["Category", "Month"]
    )
    recipe["accounting_review"] = {
        "perimeter": {
            "status": "established",
            "description": copy["perimeter"],
        },
        "source_tie_out": {
            "baseline_source_total": baseline,
            "comparison_source_total": actual,
            "tolerance": 0.01,
        },
        "favorable_adverse_convention": {
            "status": "established",
            "description": copy["signs"],
        },
        "materiality": {"status": "not_applied", "basis": copy["materiality"]},
    }
    reviewed_recipe = output / "reviewed-recipe.json"
    _write(reviewed_recipe, recipe)
    _run(
        scripts + "run_variance.py",
        source,
        "--output-dir",
        output / "variance",
        "--recipe",
        reviewed_recipe,
        "--currency",
        "EUR",
        "--language",
        language,
        "--client-engagement",
        run["context_path"],
    )
    for filename in (
        "model_use_manifest.json",
        "variance_audit.json",
        "waterfall.png",
        "root_cause_sweep_summary.json",
        "root_cause_client_report.docx",
        "root_cause_client_report.md",
        "review_payload.json",
        "final_artifacts.json",
    ):
        artifact = output / "variance" / filename
        assert artifact.is_file() and artifact.stat().st_size > 0
    audit = _read(output / "variance/variance_audit.json")
    assert "approved_for_client_use" not in json.dumps(audit)
    with (output / "variance/root_cause_total_bridge.csv").open() as stream:
        bridge = list(csv.DictReader(stream))
    assert sum(Decimal(row["variance_amount"]) for row in bridge) == actual - baseline
    assert any(Decimal(row["amount_baseline"]) < 0 for row in bridge)
    summary = _read(output / "variance/root_cause_sweep_summary.json")
    for alternative in summary["alternatives"]:
        selected = sum(
            Decimal(value.strip())
            for value in alternative["selected_amounts"].split("|")
        )
        assert (
            selected + Decimal(str(alternative["other_residual"])) == actual - baseline
        )

    standard = _read(output / "variance/standard_variance_context.json")
    assert standard["pvm_available"] is False
    assert standard["dominant_component"] == {
        "variance_type": "Total variance",
        "variance_amount": actual - baseline,
    }
    assert all(row["variance_type"] == "Total variance" for row in bridge)
    assert not (output / "variance/pvm_decomposition_ladder.png").exists()
    with (output / "variance/total_by_dimension_bridge.csv").open() as stream:
        category_rows = list(csv.DictReader(stream))
    costs = [row for row in category_rows if Decimal(row["amount_baseline"]) < 0]
    assert len(costs) == 2
    assert all(row["percent_delta"] == "" for row in costs)
    write_variance_review(output, language, phase)
    _complete_teaching_case(run, tmp_path / "case")
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="variance-analysis",
        language=language,
        phase=phase,
    )


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
@pytest.mark.parametrize("phase,population", [("demo", 13), ("practice", 15)])
def test_journal_sampling_kit_normalizes_and_samples_current_bound_source(
    tmp_path, monkeypatch, record_property, language, phase, population
):
    run = _bound_case(
        tmp_path,
        monkeypatch,
        "journal-sampling",
        "journal-sampling",
        phase,
        language=language,
    )
    _sample_teaching_journal(run, monkeypatch, language, population=population)
    _complete_teaching_case(
        run,
        tmp_path / "case",
        artifact_ids={
            "normalization/normalized_journal.csv": "prepared.normalized_journal",
            "normalization/normalization_diagnostics.json": "internal.normalization_diagnostics",
            "sample/journal_sample.csv": "prepared.journal_sample_csv",
        },
    )
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="journal-sampling",
        language=language,
        phase=phase,
    )


def _sample_teaching_journal(
    run, monkeypatch, language, *, population, size=5, include_accounts=None
):
    source = Path(
        next(
            i["path"]
            for i in run["context"]["input_bindings"]
            if i["role"] == "journal"
        )
    )
    original = source.read_bytes()
    output = Path(run["output_dir"])
    normalization = output / "normalization"
    scripts = "plugins/journal-sampling/scripts/"
    common = (
        "--client-engagement",
        run["context_path"],
        "--language",
        language,
        "--document-language",
        "en",
    )
    _run(scripts + "inspect_journal.py", source, "--output-dir", normalization, *common)
    recipe_path = normalization / "suggested_recipe.json"
    recipe = _read(recipe_path)
    entry = recipe["files"][source.name]
    # These source meanings are authored in context-<language>.md. The fixed
    # regression review is not shipped as an approved recipe for the learner.
    entry["mapping"].update(movement_number="Entry", line_number="Line")
    entry.update(
        posting_identity="movement_number_and_line_number",
        decimal_separator=".",
        thousands_separator=None,
        excluded_monetary_columns=[],
        carry_forward_fields=[],
    )
    proposed = output / "mapped-recipe.json"
    _write(proposed, recipe)
    _run(
        scripts + "inspect_journal.py",
        source,
        "--output-dir",
        normalization,
        "--recipe",
        proposed,
        *common,
    )
    # Use the current receipt constructor on the exact newly inspected mapping.
    monkeypatch.syspath_prepend(str(ROOT / "tests/plugins"))
    from test_journal_sampling_plugin import _approve_suggested_recipe

    _approve_suggested_recipe(recipe_path)
    _run(
        scripts + "normalize_journal.py",
        source,
        "--output-dir",
        normalization,
        "--recipe",
        recipe_path,
        *common,
    )
    _run(
        scripts + "run_sample.py",
        normalization / "normalized_journal.csv",
        "--output-dir",
        output / "sample",
        "--method",
        "random",
        "--size",
        str(size),
        *(["--include-accounts", include_accounts] if include_accounts else []),
        *common,
    )
    with (normalization / "normalized_journal.csv").open(
        encoding="utf-8-sig", newline=""
    ) as stream:
        normalized = list(csv.DictReader(stream))
    with (output / "sample/journal_sample.csv").open(
        encoding="utf-8-sig", newline=""
    ) as stream:
        sampled = list(csv.DictReader(stream))
    assert len(normalized) == population
    assert len(sampled) == size
    assert all(row in normalized for row in sampled)
    assert source.read_bytes() == original
    review = _read(output / "sample/model_review_context.json")["review"]
    assert review["summary"]["seed"] == 42
    assert review["summary"]["method"] == "random"
    assert review["summary"]["population_size_before_filters"] == population
    assert all(item["status"] == "needs_review" for item in review["items"])
    for filename in (
        "journal_sample.xlsx",
        "sampling_audit.json",
        "sample_reproducibility.json",
        "sample_material_value_ledger.json",
        "sample_assurance_gates.json",
        "sample_assurance_envelope.json",
        "sample_output_receipts.json",
        "model_review_context.json",
        "final_artifacts.json",
    ):
        path = output / "sample" / filename
        assert path.is_file() and path.stat().st_size > 0


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
@pytest.mark.parametrize("phase,count", [("demo", 3), ("practice", 4)])
@prepared_kit("vera/purchase-invoice-review")
def test_purchase_invoice_kit_inputs_match_actual_bookings(
    tmp_path, monkeypatch, language, phase, count
):
    """Check source compatibility only; the native reviewer is a separate check."""
    monkeypatch.syspath_prepend(str(ROOT / "plugins/_shared/vendor/modules"))
    from courseware.library import CourseLibrary

    kit = CourseLibrary(ROOT / "plugins/vera", {"purchase-invoice-review"}).render(
        "purchase-invoice-review", language, tmp_path / "kit"
    )
    sources = kit["source_files" if phase == "demo" else "practice_files"]
    ledger = next(Path(p) for p in sources if p.endswith(".csv"))
    invoices = next(Path(p).parent for p in sources if p.endswith(".xml"))
    with ledger.open(encoding="utf-8", newline="") as stream:
        headers = csv.DictReader(stream).fieldnames
    mapping = tmp_path / "mapping.json"
    # The kit declares these canonical source fields. No reviewed mapping is
    # precomputed for the native teacher or distributed as an approved input.
    _write(mapping, {**{key: key for key in headers}, "number_format": "canonical"})
    spec = importlib.util.spec_from_file_location(
        "kit_passive_audit_core",
        ROOT / "plugins/passive-invoice-audit/scripts/audit_core.py",
    )
    core = importlib.util.module_from_spec(spec)
    monkeypatch.setitem(sys.modules, spec.name, core)
    spec.loader.exec_module(core)
    loaded = core.load_ledger(ledger, mapping)
    parsed = core.parse_invoice_population(invoices, tmp_path / "parsed")
    items, orphans = core.match_population(parsed, loaded, core.CENT)
    assert len(loaded) == count * 3
    assert len(parsed) == count
    assert len(items) == count
    assert orphans == []
    assert all(item["match_state"] == "matched" for item in items)
    assert all("invoice_number_exact" in item["match_evidence"] for item in items)


def _run_centrale_rischi_teaching_case(run, phase, total, row_count):
    """Replay the authored interpretation of the supplied fictional export."""
    source = next(
        item["path"]
        for item in run["context"]["input_bindings"]
        if item["path"].endswith(".csv")
    )
    output = Path(run["output_dir"])
    context = run["context_path"]
    scripts = "plugins/centrale-rischi-review/scripts/"
    _run(
        scripts + "inspect_inputs.py",
        "--input",
        source,
        "--output-dir",
        output / "inspection",
        "--client-engagement",
        context,
    )
    inventory = _read(output / "inspection/inspection.json")
    assert inventory["tables"][0]["row_count"] == row_count
    # The kit declares these meanings explicitly in caso.md. This fixed review
    # belongs only to this fictional regression case; it is never a classifier
    # or a professional approval copied into a user's lesson.
    recipe = {
        "schema_version": "vera.centrale_rischi_recipe.v2",
        "workflow_id": "centrale-rischi-review",
        "inventory_sha256": inventory["inventory_sha256"],
        "entity": "Officina Arco Srl — caso fittizio",
        "currency": "EUR",
        "analysis_mode": "trend",
        "analysis_objective": "Preparare l’incontro con il cliente: affidamenti, utilizzi e andamento mensile.",
        "audience": "professional",
        "source_kind": "tabular_export",
        "source_document_sha256": "",
        "table_id": inventory["tables"][0]["table_id"],
        "columns": {
            "reference_month": "Mese",
            "intermediary": "Intermediario",
            "risk_category": "Categoria",
            "original_duration": "Durata originaria",
            "residual_duration": "Durata residua",
            "granted": "Accordato",
            "operational_granted": "Operativo",
            "used": "Utilizzato",
            "guarantee_type": "Garanzia",
            "guaranteed_amount": "Garantito",
        },
        "value_mappings": {
            "original_term": {"Fino a 1 anno": "short", "Oltre 5 anni": "long"},
            "residual_term": {
                "Fino a 1 anno": "within_one_year",
                "Oltre 1 anno": "over_one_year",
            },
            "exposure_family": {
                "A revoca": "performing",
                "A scadenza": "performing",
                "Autoliquidante": "performing",
            },
        },
        "control_totals": {"used": total},
        "control_tolerance": "0.01",
        "mapping_review": {
            "status": "reviewed",
            "reviewer": "Significati dichiarati nel caso didattico",
            "reviewed_at": "2026-09-14T07:00:00+02:00",
        },
    }
    recipe_path = output / "inspection/reviewed_recipe.json"
    _write(recipe_path, recipe)
    _run(
        scripts + "run_analysis.py",
        "--input",
        source,
        "--recipe",
        recipe_path,
        "--output-dir",
        output / "analysis",
        "--client-engagement",
        context,
    )
    analysis_path = output / "analysis/centrale_rischi_analysis.json"
    analysis = _read(analysis_path)
    assert analysis["status"] == "complete"
    metrics = {item["metric_id"]: item for item in analysis["metrics"]}
    assert metrics["cr.total_used"]["value"] == total
    assert metrics["financial.dscr"]["availability"] == "unavailable"
    assert analysis["controls"][0]["status"] == "passed"
    commentary = _read(output / "analysis/commentary_template.json")
    commentary["observations"] = [
        {
            "text": (
                "A marzo l’utilizzato è 119.000 EUR su 170.000 EUR di accordato operativo. Il margine calcolato è 51.000 EUR; il totale coincide con il riscontro del caso."
                if phase == "demo"
                else "Ad aprile l’utilizzato è 117.000 EUR su 170.000 EUR di accordato operativo. Il margine calcolato sale a 53.000 EUR; il totale coincide con il riscontro del caso."
            ),
            "evidence_refs": [
                "metric:cr.total_used",
                "metric:cr.total_operational_granted",
                "metric:cr.available_resources",
            ],
        },
        {
            "text": (
                "La categoria A scadenza presso Banca Levante rappresenta 82.000 EUR, il 68,91% dell’utilizzato; i due rapporti presso Banca Aurora sommano 37.000 EUR. Questa distribuzione indica dove concentrare il confronto sugli affidamenti."
                if phase == "demo"
                else "La categoria A scadenza presso Banca Levante rappresenta 80.000 EUR, il 68,38% dell’utilizzato; i due rapporti presso Banca Aurora sommano ancora 37.000 EUR. La concentrazione cambia poco rispetto a marzo."
            ),
            "evidence_refs": [
                "metric:cr.top_intermediary_share_pct",
                "metric:cr.category.902be1e530.used",
                "metric:cr.original_term.short_share_pct",
            ],
        },
        {
            "text": (
                "L’utilizzato resta a 128.000 EUR tra gennaio e febbraio e scende di 9.000 EUR a marzo. Rispetto a febbraio: A revoca diminuisce di 6.000 EUR, A scadenza di 2.000 EUR e Autoliquidante di 1.000 EUR. L’export descrive i saldi, non le cause dei movimenti."
                if phase == "demo"
                else "Rispetto a marzo l’utilizzato scende di 2.000 EUR: A revoca e A scadenza diminuiscono ciascuna di 2.000 EUR, mentre Autoliquidante cresce di 2.000 EUR. L’aumento su quest’ultima categoria compensa in parte le altre riduzioni; la causa resta da verificare."
            ),
            "evidence_refs": [
                "metric:cr.used_mom_change",
                "metric:cr.category.565c6a7b77.used_change",
                "metric:cr.category.902be1e530.used_change",
                "metric:cr.category.c01763df9e.used_change",
            ],
        },
        {
            "text": (
                "L’utilizzo per categoria è 56% per A revoca, 91,11% per A scadenza e 30% per Autoliquidante. La lettura va mantenuta distinta per categoria. Le durate originarie distinguono 37.000 EUR a breve e 82.000 EUR a lungo; la durata residua indica invece entro oppure oltre un anno."
                if phase == "demo"
                else "L’utilizzo per categoria è 52% per A revoca, 88,89% per A scadenza e 36,67% per Autoliquidante. Le durate originarie distinguono 37.000 EUR a breve e 80.000 EUR a lungo; la durata residua resta una lettura separata."
            ),
            "evidence_refs": [
                "metric:cr.category.565c6a7b77.utilization_pct",
                "metric:cr.category.902be1e530.utilization_pct",
                "metric:cr.category.c01763df9e.utilization_pct",
                "metric:cr.original_term.short_share_pct",
                "metric:cr.original_term.long_share_pct",
            ],
        },
    ]
    commentary["questions"] = [
        "Quali incassi, rimborsi o nuove operazioni spiegano le variazioni dei saldi? Per rispondere servono movimenti o documenti bancari, non soltanto i saldi mensili.",
        "Gli affidamenti per categoria e le scadenze rispondono al fabbisogno previsto? Portare all’incontro condizioni contrattuali e previsioni di cassa, se disponibili.",
    ]
    commentary["limitations"] = [
        "L’export didattico copre soltanto i rapporti e i mesi forniti; non è un report ufficiale della Banca d’Italia né una riconciliazione con documenti esterni.",
        "Non sono state fornite informazioni pregiudizievoli: i campi vuoti non attestano assenza di eventi. Il margine calcolato non dimostra disponibilità immediata di nuova liquidità.",
        "Mancano bilancio, EBITDA e flussi di cassa: PFN/EBITDA, Debt/Equity e DSCR restano non disponibili. Il commento è una bozza da rivedere con il professionista.",
    ]
    commentary_path = output / "analysis/commentary.json"
    _write(commentary_path, commentary)
    _run(
        scripts + "finalize_analysis.py",
        "--analysis",
        analysis_path,
        "--commentary",
        commentary_path,
        "--output-dir",
        output / "review",
        "--client-engagement",
        context,
    )
    assert (output / "analysis/centrale_rischi_analysis.xlsx").stat().st_size > 0
    assert (output / "review/centrale_rischi_dashboard_reviewed.html").is_file()
    assert (
        _read(output / "review/commentary_receipt.json")["status"]
        == "draft_pending_professional_review"
    )
    assert analysis["coverage"]["pregiudizievoli"] == "unavailable"
    assert analysis["latest_reference_month"] == (
        "2026-03" if phase == "demo" else "2026-04"
    )
    (output / "artifact_card.md").write_text(
        f"# Analisi Centrale Rischi — {analysis['latest_reference_month']}\n\n"
        f"[Cruscotto con commento]({output / 'review/centrale_rischi_dashboard_reviewed.html'})\n\n"
        f"[Prospetto Excel]({output / 'analysis/centrale_rischi_analysis.xlsx'})\n\n"
        f"[Sintesi e questioni da approfondire]({output / 'review/centrale_rischi_report.md'})\n\n"
        "Calcolo completato e totale di controllo riscontrato; bozza in attesa di revisione professionale. "
        "Apri il cruscotto sul mese corrente e usa Esposizioni, Serie mensile e Controlli per risalire ai dati.\n",
        encoding="utf-8",
    )
    (output / "codex_run_review.md").write_text(
        "# Riesame del commento\n\n"
        "Confrontati export, caso didattico, indicatori, categorie e serie mensile. "
        "I saldi e le variazioni sono separati dalle cause da verificare. "
        "La copertura pregiudizievole è non disponibile perché il caso non fornisce tali informazioni. "
        "Le durate originaria e residua rimangono distinte; nessun rating o indice basato su dati assenti è proposto. "
        "Il commento resta una proposta, senza approvazione professionale.\n",
        encoding="utf-8",
    )


@pytest.mark.parametrize("phase", ["demo", "practice"])
def test_centrale_rischi_kit_runs_actual_inspection_analysis_and_delivery(
    tmp_path, monkeypatch, record_property, phase
):
    run = _bound_case(
        tmp_path,
        monkeypatch,
        "centrale-rischi-review",
        "centrale-rischi-review",
        "demo",
    )
    first_outputs = {}
    for current_phase, total, row_count in [
        ("demo", "119000", 9),
        ("practice", "117000", 12),
    ]:
        _run_centrale_rischi_teaching_case(run, current_phase, total, row_count)
        ledger = _complete_teaching_case(run, tmp_path / "case")
        if current_phase == phase:
            break
        output = Path(run["output_dir"])
        first_outputs = {p: p.read_bytes() for p in output.rglob("*") if p.is_file()}
        context = run["context"]
        imports = [
            ledger.import_document(
                tmp_path / "case",
                context["client_id"],
                context["engagement_id"],
                p,
                "source",
            )
            for p in (tmp_path / "kit/files/practice").iterdir()
            if p.is_file()
        ]
        prepared = ledger.prepare_run(
            tmp_path / "case",
            context["client_id"],
            context["engagement_id"],
            "centrale-rischi-review",
            context["workflow_version"],
            input_ids=[i["receipt"]["input_id"] for i in imports],
        )
        run = ledger.start_run(
            tmp_path / "case", context["engagement_id"], prepared["run"]["run_id"]
        )
    for p, original in first_outputs.items():
        assert p.read_bytes() == original
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="centrale-rischi-review",
        language="it",
        phase=phase,
    )
    assert _read(tmp_path / "kit/course-provenance.json")["execution_receipt"] is False


@pytest.mark.parametrize(
    "phase,numbers",
    [
        ("demo", ["DEMO-001", "DEMO-002", "DEMO-003"]),
        ("practice", ["DEMO-001", "DEMO-002", "DEMO-003", "DEMO-004"]),
    ],
)
@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_xml_kit_runs_actual_managed_parser_with_demo_and_practice_files(
    tmp_path, monkeypatch, phase, numbers, language, record_property
):
    run = _bound_case(
        tmp_path,
        monkeypatch,
        "fatture-xml-check",
        "client-file-preparation",
        phase,
        language=language,
    )
    output = Path(run["output_dir"]) / "invoices"
    _run(
        "plugins/client-file-preparation/scripts/parse_fatturapa_xml.py",
        run["context"]["input_dir"],
        "--year",
        "2026",
        "--out",
        output,
        "--client-engagement",
        run["context_path"],
    )
    with (output / "fatture_summary.csv").open(encoding="utf-8", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert sorted(row["invoice_number"] for row in rows) == numbers
    assert {row["currency"] for row in rows} == {"EUR"}
    assert sum(Decimal(row["total_amount"]) for row in rows) == Decimal(
        "2440.00" if phase == "demo" else "3416.00"
    )
    assert {row["invoice_date"][:7] for row in rows} == (
        {"2026-03"} if phase == "demo" else {"2026-03", "2026-04"}
    )
    assert (output / "formal_anomalies.md").is_file()
    _complete_teaching_case(run, tmp_path / "case")
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="fatture-xml-check",
        language=language,
        phase=phase,
    )


@pytest.mark.parametrize(
    "product,language",
    [
        ("vera", "it"),
        ("vera", "en"),
        ("clara", "it"),
        ("clara", "en"),
        ("clara", "fr"),
        ("clara", "de"),
        ("clara", "es"),
    ],
)
@pytest.mark.parametrize(
    "phase,end,actual,budget,variance",
    [
        ("demo", "2026-02-28", "53000", "47000", "6000"),
        ("practice", "2026-03-31", "85000", "77000", "8000"),
    ],
)
@prepared_kit("vera/management-control-pack", "clara/reporting-engine")
def test_budget_kits_run_the_owning_product_report_pipeline(
    tmp_path,
    monkeypatch,
    product,
    language,
    phase,
    end,
    actual,
    budget,
    variance,
    record_property,
):
    if product == "vera":
        run = _bound_case(
            tmp_path,
            monkeypatch,
            "management-control-pack",
            "management-control-pack",
            phase,
            language=language,
        )
        source = next(
            item["path"]
            for item in run["context"]["input_bindings"]
            if item["path"].endswith(".xlsx")
        )
        output = Path(run["output_dir"])
        _run(
            "plugins/management-control-pack/scripts/inspect_inputs.py",
            "--input",
            source,
            "--output-dir",
            output / "inspection",
            "--client-engagement",
            run["context_path"],
        )
    else:
        monkeypatch.syspath_prepend(str(ROOT / "plugins/_shared/vendor/modules"))
        from courseware.library import CourseLibrary

        kit = CourseLibrary(ROOT / "plugins/clara", {"reporting-engine"}).render(
            "reporting-engine", language, tmp_path / "kit"
        )
        source = next(
            p
            for p in kit["source_files" if phase == "demo" else "practice_files"]
            if p.endswith(".xlsx")
        )
        output = tmp_path / "clara-report"
        spec = importlib.util.spec_from_file_location(
            "kit_budget_report",
            ROOT / "plugins/clara/modules/reporting-engine/scripts/budget_report.py",
        )
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        assert (
            module.main(
                [
                    "inspect",
                    "--input",
                    source,
                    "--output-dir",
                    str(output / "inspection"),
                ]
            )
            == 0
        )
    inventory = _read(output / "inspection/inspection.json")
    tables = {t["table_label"]: t["table_id"] for t in inventory["tables"]}
    recipe = {
        "schema_version": "vera.management_control_recipe.v1",
        "workflow_id": "management-control-pack",
        "inventory_sha256": inventory["inventory_sha256"],
        "language": language,
        "entity": "Arco — fictional teaching data",
        "reporting_period": {"start": "2026-01-01", "end": end, "cutoff": end},
        "currency": "EUR",
        "fiscal_year_start_month": 1,
        "number_format": "dot_decimal",
        "date_format": "%Y-%m-%d",
        "tables": {
            "general_ledger": {
                "table_id": tables["GL"],
                "columns": {
                    "date": "Date",
                    "account_code": "Account",
                    "category": "Category",
                    "amount": "Amount",
                },
            },
            "budget": {
                "table_id": tables["Budget"],
                "columns": {"date": "Date", "category": "Category", "amount": "Amount"},
            },
        },
        "category_roles": {
            "Revenue": "revenue",
            "COGS": "cogs",
            "Operating expenses": "operating_expense",
        },
        "category_multipliers": {},
        "aging_buckets": [30, 60, 90],
        "top_customers": 10,
        "control_totals": {"general_ledger": actual, "budget": budget},
        "control_tolerance": "0.01",
        "mapping_review": {
            "status": "reviewed",
            "reviewer": "Declared fictional kit meanings only",
            "reviewed_at": "2026-09-14T07:00:00+02:00",
        },
        "audience": "public_demo",
    }
    recipe_path = output / "reviewed_recipe.json"
    _write(recipe_path, recipe)
    args = [
        "--input",
        source,
        "--recipe",
        str(recipe_path),
        "--output-dir",
        str(output / "report"),
    ]
    if product == "vera":
        _run(
            "plugins/management-control-pack/scripts/run_pack.py",
            *args,
            "--client-engagement",
            run["context_path"],
        )
    else:
        assert module.main(["run", *args]) == 0
    context = _read(output / "report/model_context.json")
    assert context["status"] != "blocked"
    pack = _read(output / "report/management_control_pack.json")
    assert pack["language"] == language
    metrics = {item["metric_id"]: item["value"] for item in context["metrics"]}
    assert metrics["budget.total.ebitda_variance"] == variance
    assert metrics["pnl.total.ebitda"] == actual
    from openpyxl import load_workbook

    summary, report_title, comparison = {
        "it": ("Sintesi", "Controllo di gestione", "Consuntivo e budget"),
        "en": ("Summary", "Management Control Pack", "Actual and budget"),
        "fr": ("Synthèse", "Rapport de gestion", "Réalisé et budget"),
        "de": ("Zusammenfassung", "Controllingbericht", "Ist und Budget"),
        "es": ("Resumen", "Informe de gestión", "Real y presupuesto"),
    }[language]
    workbook = load_workbook(output / "report/management_control_pack.xlsx")
    assert workbook.sheetnames[0] == summary
    workbook.close()
    markdown = (output / "report/management_control_facts.md").read_text()
    assert markdown.startswith(f"# {report_title}")
    dashboard = (output / "report/management_control_dashboard.html").read_text()
    assert f'<html lang="{language}">' in dashboard
    assert comparison in dashboard
    record_native_check(
        record_property,
        product=product,
        workflow="management-control-pack" if product == "vera" else "reporting-engine",
        language=language,
        phase=phase,
        root=ROOT,
    )


@pytest.mark.parametrize("phase", ["demo", "practice"])
def test_treasury_kit_inputs_produce_real_reviewable_forecast_without_acceptance(
    tmp_path, monkeypatch, record_property, phase
):
    run = _bound_case(
        tmp_path, monkeypatch, "treasury-forecast", "treasury-forecast", phase
    )
    manifest = next(
        item["path"]
        for item in run["context"]["input_bindings"]
        if item["path"].endswith("manifest.json")
    )
    _run(
        "plugins/treasury-forecast/scripts/run_treasury.py",
        "prepare",
        "--client-engagement",
        run["context_path"],
        "--manifest",
        manifest,
    )
    output = Path(run["output_dir"])
    artifacts = _read(output / "final_artifacts.json")
    forecast = _read(output / artifacts["forecast"])
    assert forecast["company_name"] == "Officina Arco — caso fittizio"
    assert len(forecast["events"]) == 5
    assert not forecast["review"]
    first_version = output / artifacts["forecast"]
    first_bytes = first_version.read_bytes()
    # These assumptions are stated in the fictional case, not learner decisions
    # or a professional acceptance. Exercise the native versioned date review.
    request = output / "date-review-request.json"
    _write(
        request,
        {
            "record_sha256": forecast["record_sha256"],
            "decisions": {
                "item:ar-01": {
                    "expected_date": "2026-04-10",
                    "basis": "caso.md: incasso di Aurora confermato nel caso fittizio.",
                },
                "item:ap-01": {
                    "expected_date": "2026-04-15",
                    "basis": "caso.md: pagamento a Levante previsto nel caso fittizio.",
                },
                "item:ar-02": {
                    "expected_date": "2026-04-24",
                    "basis": "caso.md: incasso di Borgo confermato nel caso fittizio.",
                },
            },
        },
    )
    _run(
        "plugins/treasury-forecast/scripts/run_treasury.py",
        "review",
        "--client-engagement",
        run["context_path"],
        "--request",
        request,
    )
    artifacts = _read(output / "final_artifacts.json")
    forecast = _read(output / artifacts["forecast"])
    assert forecast["status"] == "draft_for_review"
    assert forecast["calculation_complete"] is True
    assert forecast["review"] is None
    assert forecast["opening_cash"] == "25000.00"
    assert forecast["minimum_daily_cash"] == "25000.00"
    assert forecast["daily"][-1]["closing_cash"] == "29500.00"
    assert first_version.read_bytes() == first_bytes
    assert (output / artifacts["report"]).is_file()
    assert (output / artifacts["workbook"]).stat().st_size > 0
    note = [
        "# Previsione di cassa di Officina Arco",
        "",
        "Il prospetto parte dal saldo effettivo del 31 marzo e mostra i cinque "
        "flussi forniti fino al 31 maggio. Le date delle tre partite sono state "
        "registrate con la motivazione indicata in caso.md; i due altri pagamenti "
        "derivano dal prospetto della direzione fittizia.",
        "",
        "| Prospetto iniziale | EUR |",
        "| --- | ---: |",
        f"| Cassa iniziale | {forecast['opening_cash']} |",
        f"| Minimo giornaliero | {forecast['minimum_daily_cash']} |",
        f"| Cassa al 31 maggio | {forecast['daily'][-1]['closing_cash']} |",
        "",
    ]
    if phase == "practice":
        baseline_bytes = (output / artifacts["forecast"]).read_bytes()
        scenario_request = output / "hypothetical-dates-request.json"
        _write(scenario_request, {"item:ar-01": "2026-04-20"})
        _run(
            "plugins/treasury-forecast/scripts/run_treasury.py",
            "scenario",
            "--client-engagement",
            run["context_path"],
            "--request",
            scenario_request,
        )
        scenario_path = next(output.glob("scenario-*.json"))
        scenario = _read(scenario_path)
        assert scenario["status"] == "hypothetical"
        assert scenario["baseline_record_sha256"] == forecast["record_sha256"]
        assert scenario["minimum_daily_cash"] == "16000.00"
        assert scenario["daily"][-1]["closing_cash"] == "29500.00"
        baseline_days = {row["date"]: row for row in forecast["daily"]}
        scenario_days = {row["date"]: row for row in scenario["daily"]}
        assert baseline_days["2026-04-15"]["closing_cash"] == "34000.00"
        assert scenario_days["2026-04-15"]["closing_cash"] == "16000.00"
        assert scenario_days["2026-04-20"]["closing_cash"] == "34000.00"
        assert (output / artifacts["forecast"]).read_bytes() == baseline_bytes
        assert _read(output / "final_artifacts.json") == artifacts
        note.extend(
            [
                "## Ritardo dell’incasso di Aurora",
                "",
                "L’alternativa sposta solo i 18.000 EUR dal 10 al 20 aprile. "
                "La cassa è più bassa dal 10 al 19 aprile; dal 20 aprile le "
                "disponibilità tornano a coincidere. Il saldo finale uguale "
                "non elimina la necessità di controllare quelle settimane.",
                "",
                "| Data | Previsione iniziale EUR | Alternativa EUR |",
                "| --- | ---: | ---: |",
            ]
        )
        for date in ("2026-04-10", "2026-04-15", "2026-04-20", "2026-05-31"):
            note.append(
                f"| {date} | {baseline_days[date]['closing_cash']} | "
                f"{scenario_days[date]['closing_cash']} |"
            )
        note.extend(
            [
                "",
                f"Minimo nell’alternativa: {scenario['minimum_daily_cash']} EUR. "
                f"[Calcolo nativo dell’alternativa]({scenario_path.name}).",
                "",
                "La funzione conserva l’alternativa in un JSON separato; questa "
                "tabella ne legge i risultati. Il report e il workbook originali "
                "restano quelli della previsione iniziale.",
                "",
            ]
        )
    note.extend(
        [
            "## Cosa controllare e come ripetere",
            "",
            "Apri il report, segui i saldi settimanali e ritrova ogni evento "
            "nel workbook. La copertura include soltanto i flussi dichiarati; "
            "date e incassi restano ipotesi da rivedere. La bozza non è stata "
            "professionalmente accettata e non esegue pagamenti.",
            "",
            "Per ripetere, fornisci saldi effettivi, partite residue e flussi "
            "futuri con le relative date. Per un aggiornamento successivo servono "
            "anche movimenti, regolamenti e la precedente versione accettata. "
            "L’alternativa ipotetica non può sostituire quel precedente.",
            "",
            f"[Report]({artifacts['report']}) · "
            f"[Workbook]({artifacts['workbook']}) · "
            f"[Record corrente]({artifacts['forecast']})",
            "",
        ]
    )
    (output / "codex_run_review.md").write_text("\n".join(note), encoding="utf-8")
    (output / "artifact_card.md").write_text(
        "# Previsione di cassa di Officina Arco\n\n"
        "Bozza da rivedere al 31 marzo 2026, con orizzonte al 31 maggio. "
        "Il calcolo è completo per i cinque flussi forniti; restano da "
        "valutare copertura, ipotesi sulle date e accettazione professionale.\n\n"
        f"[Report corrente]({output / artifacts['report']})\n\n"
        f"[Workbook corrente]({output / artifacts['workbook']})\n\n"
        f"[Lettura dei risultati ed esercizio]({output / 'codex_run_review.md'})\n",
        encoding="utf-8",
    )
    _complete_teaching_case(run, tmp_path / "case")
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="treasury-forecast",
        language="it",
        phase=phase,
    )


@pytest.mark.parametrize("phase", ["demo", "practice"])
@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_lucia_kit_inputs_become_bound_evidence_in_the_actual_matter_package(
    tmp_path, monkeypatch, phase, language, record_property
):
    monkeypatch.syspath_prepend(str(ROOT / "plugins/apertura-pratica/scripts"))
    from tests.plugins._matter_teaching import prepare_intake, review_and_deliver

    run = _bound_case(
        tmp_path,
        monkeypatch,
        "apertura-pratica",
        "apertura-pratica",
        "demo",
        product="lucia",
        language=language,
    )
    scripts = "plugins/apertura-pratica/scripts/"
    first_context = run["context"]
    first_outputs = {}
    previous_intake = None
    previous_output = None
    client_root = tmp_path / "case"
    for current_phase in ("demo", "practice"):
        output = Path(run["output_dir"])
        _run(
            scripts + "initialize_workspace.py",
            output,
            "--opening-mode",
            "new_client_new_matter",
            "--client-reference",
            "beta-fictional",
            "--matter-reference",
            "delivery-fictional",
            "--language",
            language,
        )
        for item in run["context"]["input_bindings"]:
            name = Path(item["path"]).name
            role = (
                "firm_record"
                if name == "register.csv"
                else (
                    "correspondence"
                    if name.startswith("update-")
                    else "client_supplied"
                )
            )
            _run(scripts + "add_evidence.py", output, item["path"], "--role", role)
        intake = prepare_intake(output, language, previous_intake)
        _run(scripts + "prepare_review.py", output)
        evidence = intake["evidence_register"]
        assert len(evidence) == (3 if current_phase == "demo" else 4)
        assert all(
            hashlib.sha256((output / item["stored_path"]).read_bytes()).hexdigest()
            == item["sha256"]
            for item in evidence
        )
        assert [party["display_name"] for party in intake["parties"]] == [
            "Beta Laboratorio Srl",
            "Gamma Forniture Srl",
            "Elena Esempio",
        ]
        assert intake["conflict_check"]["register_scope"] == "partial"
        assert intake["conflict_check"]["professional_decision"]["status"] == "pending"
        assert intake["deadline_review"]["status"] == "pending"
        assert intake["confirmed_facts"] == []
        review_and_deliver(output, language, _run, previous_output)
        assert (
            intake["matter"]["summary"]
            in (output / "matter_opening_memo.md").read_text()
        )
        final = _read(output / "final_artifacts.json")
        assert final["status"] == "blocked"
        assert final["conflict_cleared_by_software"] is False
        assert final["engagement_accepted_by_software"] is False
        assert final["source_files_modified"] is False
        for artifact in _read(output / "artifact_manifest.json")["artifacts"]:
            assert (
                hashlib.sha256((output / artifact["path"]).read_bytes()).hexdigest()
                == artifact["sha256"]
            )
        ledger = _complete_teaching_case(run, client_root)
        if current_phase == "practice":
            assert intake["matter"]["summary"].startswith(
                previous_intake["matter"]["summary"]
            )
            assert (
                intake["matter"]["objective"] != previous_intake["matter"]["objective"]
            )

            def source_bound_value(record, field):
                value = json.dumps(record[field], sort_keys=True)
                for source in record["evidence_register"]:
                    value = value.replace(
                        json.dumps(source["evidence_id"]), json.dumps(source["sha256"])
                    )
                return value

            assert source_bound_value(intake, "engagement") == source_bound_value(
                previous_intake, "engagement"
            )
            assert source_bound_value(intake, "parties") == source_bound_value(
                previous_intake, "parties"
            )
            assert run["context"]["engagement_id"] == first_context["engagement_id"]
            assert run["context"]["run_id"] != first_context["run_id"]
            assert all(
                path.read_bytes() == content for path, content in first_outputs.items()
            )
            assert (
                _read(output / "applied_decisions.json")["intake_sha256"]
                != _read(previous_output / "applied_decisions.json")["intake_sha256"]
            )
        if phase == current_phase:
            break
        previous_intake = intake
        previous_output = output
        first_outputs = {
            path: path.read_bytes() for path in output.rglob("*") if path.is_file()
        }
        sources = sorted((tmp_path / "kit/files/practice").glob("*"))
        assert {path.name for path in sources} == {
            f"request-{language}.md",
            f"supply-{language}.md",
            f"update-{language}.md",
            "register.csv",
        }
        imports = [
            ledger.import_document(
                client_root,
                first_context["client_id"],
                first_context["engagement_id"],
                source,
                "source",
            )
            for source in sources
        ]
        prepared = ledger.prepare_run(
            client_root,
            first_context["client_id"],
            first_context["engagement_id"],
            "apertura-pratica",
            first_context["workflow_version"],
            input_ids=[item["receipt"]["input_id"] for item in imports],
            purpose="Fictional delivery proposal; preserve initial dossier and unresolved professional review",
        )
        run = ledger.start_run(
            client_root, first_context["engagement_id"], prepared["run"]["run_id"]
        )
    record_native_check(
        record_property,
        root=ROOT,
        product="lucia",
        workflow="apertura-pratica",
        language=language,
        phase=phase,
    )


@pytest.mark.parametrize("phase,expected", [("demo", 2), ("practice", 3)])
@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_bank_reconciliation_kit_runs_current_comparison(
    tmp_path, monkeypatch, phase, expected, language, record_property
):
    run = _bound_case(
        tmp_path,
        monkeypatch,
        "journal-bank-reconciliation",
        "journal-bank-reconciliation",
        phase,
        language=language,
    )
    files = {
        Path(item["path"]).name: Path(item["path"])
        for item in run["context"]["input_bindings"]
    }
    month = "march" if phase == "demo" else "april"
    bank = files[f"bank-{month}.csv"]
    journal = files[f"bank-ledger-{month}.csv"]
    output = Path(run["output_dir"])
    scripts = "plugins/journal-bank-reconciliation/scripts/"
    inspection = output / "inspection"
    _run(
        scripts + "inspect_inputs.py",
        bank,
        journal,
        "--client-engagement",
        run["context_path"],
        "--output-dir",
        inspection,
        "--language",
        language,
    )
    monkeypatch.syspath_prepend(str(ROOT / scripts))
    spec = importlib.util.spec_from_file_location(
        "journal_bank_core", ROOT / scripts / "journal_bank_core.py"
    )
    core = importlib.util.module_from_spec(spec)
    monkeypatch.setitem(sys.modules, spec.name, core)
    spec.loader.exec_module(core)
    recipe_path = inspection / "suggested_recipe.json"
    recipe = _read(recipe_path)
    # Fixed fictional case authority only: exact individual outgoing payments,
    # no tolerance, no reused evidence. Never shipped as a learner's approval.
    policy = {
        **recipe["relationship"]["policy"],
        "amount_tolerance": "0",
        "date_window_days": 0,
    }
    recipe["matching"]["amount_tolerance"] = "0"
    recipe["matching"]["date_window_days"] = 0
    refs = [
        row["artifact_id"]
        for row in _read(inspection / "input_receipts.json")["receipts"]
        if row["artifact_id"].startswith(("source.bank.", "source.journal."))
    ]
    recipe["relationship"] = {
        "policy": policy,
        "review_content_sha256": core.canonical_json_sha256({"policy": policy}),
        "decision": core.build_relationship_review_receipt(
            decision_id="decision.relationship",
            reviewer_ref="reviewer.teaching_fixture",
            reviewed_on="2026-09-14",
            source_artifact_refs=refs,
            policy=policy,
        ),
    }
    _write(recipe_path, recipe)
    result = output / "reconciliation"
    _run(
        scripts + "run_reconciliation.py",
        bank,
        journal,
        "--client-engagement",
        run["context_path"],
        "--output-dir",
        result,
        "--recipe",
        recipe_path,
        "--language",
        language,
        "--tolerance",
        "0",
        "--date-window-days",
        "0",
    )
    audit = _read(result / "reconciliation_audit.json")
    assert audit["matched_count"] == expected
    assert audit["unmatched_bank_count"] == 0
    assert audit["unmatched_journal_count"] == 0
    assert list(result.glob("*.xlsx"))
    assert (result / "review_notes.md").is_file()
    assert (
        _read(result / "assurance_gates.json")["gates"]["source"]["status"] == "passed"
    )
    _complete_teaching_case(run, tmp_path / "case")
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="journal-bank-reconciliation",
        language=language,
        phase=phase,
    )


@pytest.mark.parametrize("phase,closed", [("demo", 2), ("practice", 3)])
@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_open_item_kit_runs_current_pdf_ingestion_and_workpapers(
    tmp_path, monkeypatch, phase, closed, language, record_property
):
    run = _bound_case(
        tmp_path,
        monkeypatch,
        "open-item-reconciliation",
        "open-item-reconciliation",
        phase,
        language=language,
    )
    output = Path(run["output_dir"])
    input_root = Path(run["context"]["input_dir"])
    decisions = {}
    for item in run["context"]["input_bindings"]:
        source = Path(item["path"])
        is_population = source.name == f"open-items-{language}.pdf"
        decisions[source.relative_to(input_root).as_posix()] = {
            "role": "open_items" if is_population else "bank_statement",
            "adapter_family": (
                "open_items_text_v1" if is_population else "bank_statement_text_v1"
            ),
            "reviewer_ref": "reviewer.teaching_fixture",
            "reviewed_on": "2026-09-14",
            "perimeter": {
                "entity_ref": "entity.arco",
                "party_ref": "party.servizi_esempio",
                "currency": "EUR",
                "unit": "currency_amount",
                "direction_policy": "supplier",
                "allocation_policy": "one_to_one",
            },
            "money": {
                "decimal_separator": ",",
                "thousands_separator": ".",
                "reported_unit": "EUR",
                "reported_increment": "0.01",
            },
            "date": {"order": "day_first"},
        }
    assumptions = {
        "scope_year": "2026",
        "cutoff_date": "2026-03-31" if phase == "demo" else "2026-04-30",
        "assurance_run_date": "2026-09-14",
        "post_cutoff_events_excluded": True,
        "reviewed_source_decisions": decisions,
        "counterparty_keywords": ["servizi esempio"],
        "ocr_scanned": False,
    }
    # Use the workflow's supported Python entrypoint. Its validated output tree
    # is closed; a test parameter file is not a workflow deliverable.
    # A fresh interpreter also prevents pytest module cleanup from reloading
    # PyMuPDF's native extension between language cases.
    execution = subprocess.run(
        [
            sys.executable,
            "-I",
            "-B",
            "-c",
            """
import importlib.util, json, sys
spec = importlib.util.spec_from_file_location('kit_open_item_runner', sys.argv[1])
runner = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = runner
spec.loader.exec_module(runner)
context = runner.load_client_engagement_context_file(
    sys.argv[2], expected_workflow_id='open-item-reconciliation')
runner.run_raw_input_reconciliation(
    input_dir=context['input_dir'], prepared_client_engagement=context,
    assumptions=json.loads(sys.argv[3]), language=sys.argv[4],
    title=json.loads(sys.argv[5])[sys.argv[4]])
""",
            str(ROOT / "plugins/open-item-reconciliation/scripts/raw_input_runner.py"),
            run["context_path"],
            json.dumps(assumptions),
            language,
            json.dumps(
                {
                    "it": "Officina Arco — esercitazione di riconciliazione",
                    "en": "Officina Arco — reconciliation exercise",
                    "fr": "Officina Arco — exercice de rapprochement",
                    "de": "Officina Arco — Abstimmungsübung",
                    "es": "Officina Arco — ejercicio de conciliación",
                }
            ),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=60,
        check=False,
    )
    assert execution.returncode == 0, execution.stdout + execution.stderr
    assert (output / "riconciliazione_audit.xlsx").is_file()
    assert (output / "relazione_riconciliazione_audit.docx").is_file()
    assert (output / "review_ui.html").is_file()
    canonical = _read(output / "assurance_final_outputs/reconciliation_results.json")
    assert not canonical["source_processing"]["extraction_errors"]
    rows = canonical["reconciliation_rows"]
    assert len(rows) == 3
    assert sum(row["reconciliation_status"] == "closed" for row in rows) == closed
    assert (
        sum(row["reconciliation_status"] == "unresolved" for row in rows) == 3 - closed
    )
    if phase == "practice":
        third = next(row for row in rows if row["document_key"] == "3FF|2026")
        assert third["supporting_bank_date"] == "2026-04-04"
    assert _read(output / "final_artifacts.json")["status"] != "final_ready"

    # The first table is the learner's actual result, in the chosen language.
    from docx import Document

    report = Document(output / "relazione_riconciliazione_audit.docx")
    amount_headers = {
        "it": "Importo",
        "en": "Amount",
        "fr": "Montant",
        "de": "Betrag",
        "es": "Importe",
    }
    expected_amounts = {
        "it": ("1.952,00", "2.440,00"),
        "en": ("1,952.00", "2,440.00"),
        "fr": ("1952,00", "2440,00"),
        "de": ("1.952,00", "2.440,00"),
        "es": ("1.952,00", "2.440,00"),
    }
    summary = report.tables[0]
    assert summary.rows[0].cells[2].text == amount_headers[language]
    assert summary.rows[1].cells[1].text == str(closed)
    assert (
        summary.rows[1].cells[2].text == expected_amounts[language][phase == "practice"]
    )
    assert len(summary.rows) == (3 if phase == "demo" else 2)

    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="open-item-reconciliation",
        language=language,
        phase=phase,
    )


@pytest.mark.parametrize("phase", ["demo", "practice"])
@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_new_client_kit_runs_actual_dossier_from_reported_facts(
    tmp_path, monkeypatch, phase, language, record_property
):
    from tests.plugins._new_client_teaching import (
        case_input,
        complete_case,
        review_and_deliver,
        words,
    )

    run = _bound_case(
        tmp_path, monkeypatch, "new-client", "new-client", "demo", language=language
    )
    client_root = tmp_path / "case"
    scripts = "plugins/new-client/scripts/"
    previous_intake = None
    first_outputs = {}
    first_context = run["context"]
    for current_phase in ("demo", "practice"):
        output = Path(run["output_dir"])
        _run(
            scripts + "initialize_case.py",
            "--case-dir",
            output,
            "--client-engagement",
            run["context_path"],
            "--client-reference",
            "ARCO-DEMO",
            "--assessment-date",
            "2026-09-14",
            "--language",
            language,
        )
        path = output / "new_client_input.json"
        intake = case_input(_read(path), run, language, previous_intake)
        _write(path, intake)
        _run(
            scripts + "package_new_client.py",
            "--input",
            path,
            "--client-engagement",
            run["context_path"],
            "--output-dir",
            output,
        )
        memo = (output / "studio_new_client_memo.md").read_text()
        assert memo.index("Officina Arco Srl") < memo.index("RE =")
        assert words(language)["service"] in memo
        assert "`None`" not in memo
        aml_audit = _read(output / "aml_calculation_audit.json")
        assert aml_audit["status"] == "blocked_incomplete_scores"
        assert aml_audit["calculated_band"] is None
        assert aml_audit["effective_risk"] is None
        assert len(aml_audit["missing_score_ids"]) == 11
        assert _read(output / "monitoring_plan.json")["next_review_date"] is None
        assert intake["engagement"]["start_date"] is None
        saved = _read(output / "case_facts_validated.json")
        assert "Officina Arco Srl" in json.dumps(saved)
        assert ("Sara Campione" in json.dumps(saved)) == (current_phase == "practice")
        review_and_deliver(run, language, follow_up=current_phase == "practice")
        final = _read(output / "final_artifacts.json")
        assert final["professional_review_required"] is True
        assert final["relationship_activation_performed"] is False
        assert final["signature_performed"] is False
        assert final["client_communication_sent"] is False
        assert final["export_gate"]["relationship_ready"] is False
        questions = (output / "client_questions.md").read_text()
        assert "Elena Esempio" in questions and "Paolo Prova" in questions
        assert words(language)["questions"][3] in questions
        ledger = complete_case(run, client_root, _run, language)
        if current_phase == "practice":
            assert run["context"]["client_id"] == first_context["client_id"]
            assert run["context"]["engagement_id"] == first_context["engagement_id"]
            assert run["context"]["run_id"] != first_context["run_id"]
            assert intake["party_facts"][:-1] == previous_intake["party_facts"]
            assert intake["engagement"] == previous_intake["engagement"]
            contact = intake["party_facts"][-1]
            assert contact["evidence_ids"] == ["contact-update"]
            assert words(language)["contact_role"] in contact["value"]
            assert all(
                path.read_bytes() == content for path, content in first_outputs.items()
            )
        if current_phase == phase:
            break
        previous_intake = intake
        first_outputs = {
            path: path.read_bytes() for path in output.rglob("*") if path.is_file()
        }
        sources = sorted((tmp_path / "kit/files/practice").glob("*.md"))
        assert {p.name for p in sources} == {
            f"profile-{language}.md",
            f"contact-update-{language}.md",
        }
        imports = [
            ledger.import_document(
                client_root,
                first_context["client_id"],
                first_context["engagement_id"],
                source,
                "source",
            )
            for source in sources
        ]
        prepared = ledger.prepare_run(
            client_root,
            first_context["client_id"],
            first_context["engagement_id"],
            "new-client",
            first_context["workflow_version"],
            input_ids=[item["receipt"]["input_id"] for item in imports],
            purpose="Fictional contact update in the same engagement; no professional approval",
        )
        run = ledger.start_run(
            client_root, first_context["engagement_id"], prepared["run"]["run_id"]
        )
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="new-client",
        language=language,
        phase=phase,
    )


@pytest.mark.parametrize("phase,expected_documents", [("demo", 2), ("practice", 3)])
@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_fiscal_data_kit_runs_actual_preparation_and_extraction(
    tmp_path, monkeypatch, phase, expected_documents, language, record_property
):
    run = _bound_case(
        tmp_path,
        monkeypatch,
        "dati-fiscali-strutturati",
        "client-file-preparation",
        phase,
        language=language,
    )
    output = Path(run["output_dir"])
    _run(
        "plugins/client-file-preparation/scripts/build_file_preparation_outputs.py",
        run["context"]["input_dir"],
        "--client-engagement",
        run["context_path"],
        "--year",
        "2026",
        "--jurisdiction",
        "italy",
        "--language",
        language,
        "--no-ocr",
    )
    with (output / "extracted/structured_fiscal_fields.csv").open(
        encoding="utf-8-sig", newline=""
    ) as stream:
        fields = list(csv.DictReader(stream))
    assert len({row["relative_path"] for row in fields}) == expected_documents
    codes = {row["field_code"]: row["normalized_value"] for row in fields}
    assert codes["redditi_lavoro_dipendente"] == "28000.00"
    assert codes["ritenute_irpef"] == "5000.00"
    debit_values = {
        row["normalized_value"]
        for row in fields
        if row["field_code"] == "importo_debito"
    }
    assert debit_values == ({"1200.00"} if phase == "demo" else {"1200.00", "900.00"})
    assert (output / "08_dati_fiscali_strutturati.md").is_file()
    assert (output / "extracted/extraction_report.md").is_file()
    handoff = _read(output / "model_handoff.json")
    assert handoff
    _complete_teaching_case(run, tmp_path / "case")
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="dati-fiscali-strutturati",
        language=language,
        phase=phase,
    )


@pytest.mark.parametrize("phase,documents", [("demo", 1), ("practice", 2)])
@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_notice_kit_runs_actual_intake_and_keeps_source_dates(
    tmp_path, monkeypatch, phase, documents, language, record_property
):
    from tests.plugins._notice_teaching import write_notice_review

    run = _bound_case(
        tmp_path,
        monkeypatch,
        "avviso-intake",
        "client-file-preparation",
        phase,
        language=language,
    )
    output = Path(run["output_dir"])
    _run(
        "plugins/client-file-preparation/scripts/build_file_preparation_outputs.py",
        run["context"]["input_dir"],
        "--client-engagement",
        run["context_path"],
        "--year",
        "2026",
        "--jurisdiction",
        "italy",
        "--language",
        language,
        "--no-ocr",
    )
    extracted = [
        json.loads(line)
        for line in (output / "extracted/documents.jsonl")
        .read_text(encoding="utf-8")
        .splitlines()
        if line
    ]
    assert len(extracted) == documents
    table = (output / "avviso/deadlines_and_amounts.csv").read_text(encoding="utf-8")
    assert "02/04/2026" in table
    assert "30/04/2026" in table
    assert "DEMO-NOTICE-01" in table
    assert (output / "avviso/avviso_intake_memo.md").is_file()
    assert (output / "model_handoff.json").is_file()
    assert _read(output / "final_artifacts.json")["status"] != "final_ready"
    sources = [Path(item["path"]) for item in run["context"]["input_bindings"]]
    # The native table is an extraction aid. The skill also calls for a Codex
    # review note grounded in the actual readable sources. These fixed model
    # interpretations are test-only, never a prepared learner answer.
    notice = next(path for path in sources if path.name == "avviso.md")
    assert "Non quantifica un debito" in notice.read_text(encoding="utf-8")
    if phase == "practice":
        reply = next(path for path in sources if path.name == "client-update.md")
        assert "non una ricevuta di notifica" in reply.read_text(encoding="utf-8")
        assert "non è allegata" in reply.read_text(encoding="utf-8")
    note = write_notice_review(output, sources, language, phase)
    reviewed = note.read_text(encoding="utf-8")
    assert all(
        value in reviewed
        for value in ("DEMO-NOTICE-01", "02/04/2026", "30/04/2026", "2025", "F24")
    )
    assert ("06/04/2026" in reviewed) == (phase == "practice")
    assert all(str(path) in reviewed for path in sources)
    _complete_teaching_case(run, tmp_path / "case")
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="avviso-intake",
        language=language,
        phase=phase,
    )


EMAIL_REQUESTS = {
    "it": (
        "Conferma se hai altre CU per la campagna 2026.",
        "Conferma se ci hai consegnato tutti gli F24 della campagna 2026.",
    ),
    "en": (
        "Please confirm whether you have other CUs for the 2026 campaign.",
        "Please confirm whether you have supplied all F24s for the 2026 campaign.",
    ),
    "fr": (
        "Merci de confirmer si vous avez d’autres CU pour la campagne 2026.",
        "Merci de confirmer si tous les F24 de la campagne 2026 ont été transmis.",
    ),
    "de": (
        "Bitte bestätigen Sie, ob weitere CU für die Kampagne 2026 vorliegen.",
        "Bitte bestätigen Sie, ob sämtliche F24 für die Kampagne 2026 übermittelt wurden.",
    ),
    "es": (
        "Confirma si tienes otras CU para la campaña de 2026.",
        "Confirma si has entregado todos los F24 de la campaña de 2026.",
    ),
}
EMAIL_FRAME = {
    "it": (
        "Bozza email cliente",
        "Oggetto: Conferme sui documenti ricevuti",
        "Buongiorno CLIENT-001,",
        "Grazie,\nLo studio",
    ),
    "en": (
        "Client email draft",
        "Subject: Confirmation of documents received",
        "Hello CLIENT-001,",
        "Thank you,\nThe practice",
    ),
    "fr": (
        "Projet de courriel au client",
        "Objet : Confirmation des documents reçus",
        "Bonjour CLIENT-001,",
        "Merci,\nLe cabinet",
    ),
    "de": (
        "Entwurf der Mandanten-E-Mail",
        "Betreff: Bestätigung der erhaltenen Unterlagen",
        "Guten Tag CLIENT-001,",
        "Vielen Dank,\nDie Kanzlei",
    ),
    "es": (
        "Borrador de correo al cliente",
        "Asunto: Confirmación de documentos recibidos",
        "Hola CLIENT-001:",
        "Gracias,\nEl despacho",
    ),
}


def _client_review_tool(output, context, name, decisions):
    import shutil

    node = shutil.which("node")
    assert node, "The declared Node runtime is required for persistent review."
    request = {
        "jsonrpc": "2.0",
        "id": 1,
        "method": "tools/call",
        "params": {
            "name": name,
            "arguments": {
                "client_engagement": str(context),
                "run_intake": _read(output / "run_intake.json"),
                "review_payload": _read(output / "review_payload.json"),
                "final_artifacts": _read(output / "final_artifacts.json"),
                "ui_decisions": _read(output / "ui_decisions.json"),
                "decisions": decisions,
                "reviewer": "fictional-kit-regression-reviewer",
                "decision_source": "synthetic-regression-fixture",
            },
        },
    }
    process = subprocess.run(
        [node, str(ROOT / "plugins/client-file-preparation/mcp/server.cjs"), "--stdio"],
        input=json.dumps(request) + "\n",
        capture_output=True,
        text=True,
        timeout=60,
        cwd=ROOT,
        check=False,
    )
    assert process.returncode == 0, process.stderr
    response = next(
        json.loads(line)
        for line in process.stdout.splitlines()
        if json.loads(line).get("id") == 1
    )
    result = response["result"]["structuredContent"]
    assert result["ok"] is True, result
    return result


def _reviewed_email_requests(output):
    handoff = _read(output / "model_handoff.json")
    items = []
    for page in handoff["pagination"]["pages"]:
        path = output / page["path"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == page["sha256"]
        items.extend(_read(path)["items"])
    return [item for item in items if item["kind"] == "email_request"]


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
@pytest.mark.parametrize("phase", ["demo", "practice"])
def test_client_email_kit_runs_intake_review_and_sealed_draft_replacement(
    tmp_path, monkeypatch, language, phase, record_property
):
    run = _bound_case(
        tmp_path,
        monkeypatch,
        "email-cliente",
        "client-file-preparation",
        phase,
        language=language,
    )
    output = Path(run["output_dir"])
    _run(
        "plugins/client-file-preparation/scripts/build_file_preparation_outputs.py",
        run["context"]["input_dir"],
        "--client-engagement",
        run["context_path"],
        "--year",
        "2026",
        "--jurisdiction",
        "italy",
        "--language",
        language,
        "--no-ocr",
    )
    assert _reviewed_email_requests(output) == []
    payload = _read(output / "review_payload.json")
    missing = [
        item
        for item in payload["items"]
        if item["item_type"] == "missing_document_request"
    ]
    cu = next(item for item in missing if "CU" in item["data"]["request_text"])
    f24 = next(item for item in missing if "F24" in item["data"]["request_text"])
    decisions = []
    # Model-authored review for these exact fictional sources. The practice
    # reply closes the F24 question; the CU question is still open. These
    # fixture decisions are not bundled or copied into a learner's session.
    for item in payload["items"]:
        decision = {"item_id": item["id"], "action": "accept"}
        if item["item_type"] == "missing_document_request":
            decision["action"] = "reject"
        if item["id"] == cu["id"] or (item["id"] == f24["id"] and phase == "demo"):
            decision.update(
                action="request_more_documents",
                requested_documents=[
                    EMAIL_REQUESTS[language][0 if item["id"] == cu["id"] else 1]
                ],
            )
        if item["id"] == "draft-client-email":
            decision["action"] = "skip"
        decisions.append(decision)
    for action in ("save", "apply"):
        _client_review_tool(
            output,
            run["context_path"],
            f"{action}_client_file_preparation_decisions",
            decisions,
        )
    requests = _reviewed_email_requests(output)
    assert len(requests) == (2 if phase == "demo" else 1)
    assert all(item["client_reference"] == "CLIENT-001" for item in requests)
    title, subject, greeting, closing = EMAIL_FRAME[language]
    # This test draft consumes only the actual reviewed request page items.
    replacement = (
        f"# {title}\n\n{subject}\n\n{greeting}\n\n"
        + "\n".join("- " + item["request_text"] for item in requests)
        + f"\n\n{closing}"
    )
    before = (output / "04_bozza_email_cliente.md").read_bytes()
    for decision in decisions:
        if decision["item_id"] == "draft-client-email":
            decision.update(action="edit", edit_value=replacement)
    for action in ("save", "apply"):
        _client_review_tool(
            output,
            run["context_path"],
            f"{action}_client_file_preparation_decisions",
            decisions,
        )
    assert (output / "04_bozza_email_cliente.md").read_text() == replacement
    final = _read(output / "final_artifacts.json")
    record = next(
        item for item in final["outputs"] if item["path"] == "04_bozza_email_cliente.md"
    )
    assert record["sha256"] == hashlib.sha256(replacement.encode()).hexdigest()
    assert any(
        path.read_bytes() == before
        for path in output.rglob("*.md")
        if path.name != "04_bozza_email_cliente.md"
    )
    assert (EMAIL_REQUESTS[language][1] in replacement) == (phase == "demo")
    _complete_teaching_case(run, tmp_path / "case")
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="email-cliente",
        language=language,
        phase=phase,
    )


def _execute_teaching_sales_plan(
    run, client_root, growth, net_sales, margin, language, prior_plan=None
):
    """Execute the current scenario engine inside one explicitly prepared run."""
    from tests.plugins._sales_plan_teaching import (
        assumption_readback,
        deliver_plan,
        validate_links,
    )

    output = Path(run["output_dir"])
    source = next(
        Path(item["path"])
        for item in run["context"]["input_bindings"]
        if Path(item["path"]).suffix == ".csv"
    )
    original = source.read_bytes()
    (output / source.name).write_bytes(original)
    assumption_source = next(
        Path(item["path"])
        for item in run["context"]["input_bindings"]
        if Path(item["path"]).suffix == ".md"
    )
    readback = assumption_readback(output, source, assumption_source, language, growth)
    case = _read(ROOT / "plugins/sales-plan/evals/synthetic/case.json")
    case["case_id"] = "ciclo-arco-teaching"
    case["purpose"] = (
        "Execute the fictional sales-plan request; test-only interpretation."
    )
    case["preparation_recipe"]["dimension_columns"] = ["product"]
    case["preparation_recipe"]["period_mapping"] = [
        {"source_period": "2026-01", "target_period": "2027-01"},
        {"source_period": "2026-02", "target_period": "2027-02"},
    ]
    case["files"]["actual_sales"] = {
        "path": source.name,
        "sha256": hashlib.sha256(original).hexdigest(),
    }
    case["reviewed_assumptions"] = {
        "status": "reviewed",
        "reviewed_by": "reviewer.teaching_fixture",
        "reviewed_at": "2026-09-14",
        "review_basis": (
            "Test-only confirmation fixture for the units-growth request; no learner approval. "
            "Read-back SHA-256: " + hashlib.sha256(readback.read_bytes()).hexdigest()
        ),
        "assumptions": [
            {
                "assumption_id": "units-growth",
                "driver": "units_pct",
                "change_pct": growth,
                "scope": {"product": ["Urban bicycle", "Trekking bicycle"]},
                "effective_periods": ["2027-01", "2027-02"],
                "priority": 100,
                "rationale": "Requested scenario: unchanged unit prices; discounts and COGS proportional to sales.",
            }
        ],
    }
    request = output / "case.json"
    _write(request, case)
    plan = output / "plan"
    _run(
        "plugins/sales-plan/scripts/run_plan.py",
        "--case",
        request,
        "--client-engagement",
        run["context_path"],
        "--output-dir",
        plan,
    )
    receipt = _read(plan / "plan_execution_receipt.json")
    assert receipt["status"] == "passed"
    assert receipt["report_ready"] is False
    with (plan / "scenario_summary.csv").open(encoding="utf-8", newline="") as stream:
        totals = {
            row["metric"]: row
            for row in csv.DictReader(stream)
            if row["summary_level"] == "total"
        }
    assert totals["net_sales_reporting"]["actual"] == "86450"
    assert totals["net_sales_reporting"]["plan"] == net_sales
    assert totals["gross_margin_reporting"]["plan"] == margin
    assert totals["net_sales_reporting"]["delta_pct_rounded_4dp"] == growth
    with (plan / "assumption_application_ledger.csv").open(
        encoding="utf-8", newline=""
    ) as stream:
        applied = list(csv.DictReader(stream))
    assert len(applied) == 4
    assert {row["assumption_id"] for row in applied} == {"units-growth"}
    assert {row["change_pct"] for row in applied} == {growth}
    assert {row["status"] for row in applied} == {"applied"}
    assert (plan / "model_use_manifest.json").is_file()
    _run(
        "plugins/sales-plan/scripts/model_use.py",
        "--manifest",
        plan / "model_use_manifest.json",
        "--client-engagement",
        run["context_path"],
        "--reason",
        "Trace the January urban-bicycle Plan back to its Actual row and units assumption.",
        "--source-row-id",
        "urban-jan",
        "--column",
        "period",
        "--column",
        "product",
        "--column",
        "units",
        "--column",
        "gross_sales_reporting",
        "--column",
        "discount_reporting",
        "--column",
        "net_sales_reporting",
        "--column",
        "cogs_reporting",
        "--column",
        "gross_margin_reporting",
    )
    drilldowns = list((plan / "model_drilldowns").glob("scenario_rows_*.json"))
    assert len(drilldowns) == 1
    detail = _read(drilldowns[0])
    assert detail["matched_row_count"] == 2
    assert detail["full_population_rows_scanned_locally"] == 8
    traced = {row["scenario"]: row for row in detail["rows"]}
    assert set(traced) == {"AC", "PL"}
    assert traced["AC"]["period"] == "2026-01"
    assert traced["PL"]["period"] == "2027-01"
    assert {row["source_row_id"] for row in detail["rows"]} == {"urban-jan"}
    assert {row["product"] for row in detail["rows"]} == {"Urban bicycle"}
    assert Decimal(traced["AC"]["gross_sales_reporting"]) / Decimal(
        traced["AC"]["units"]
    ) == Decimal(traced["PL"]["gross_sales_reporting"]) / Decimal(traced["PL"]["units"])
    deliver_plan(plan, language, growth, drilldowns[0], prior_plan)
    ledger = _complete_teaching_case(run, client_root)
    validate_links(output)
    assert source.read_bytes() == original
    return ledger, plan


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
@pytest.mark.parametrize("phase", ["demo", "practice"])
def test_sales_plan_kit_runs_current_reviewed_assumptions(
    tmp_path, monkeypatch, record_property, language, phase
):
    initial = _bound_case(
        tmp_path, monkeypatch, "sales-plan", "sales-plan", "demo", language=language
    )
    client_root = tmp_path / "case"
    ledger, first_plan = _execute_teaching_sales_plan(
        initial, client_root, "10", "95095", "35035", language
    )
    record_property("teaching_initial_output", str(first_plan))
    preserved = {
        path: path.read_bytes()
        for path in Path(initial["output_dir"]).rglob("*")
        if path.is_file()
    }
    if phase == "demo":
        record_property("teaching_output", str(first_plan))
        record_native_check(
            record_property,
            root=ROOT,
            product="vera",
            workflow="sales-plan",
            language=language,
            phase=phase,
        )
        return
    context = initial["context"]
    client, engagement = context["client_id"], context["engagement_id"]
    # Practice imports its own changed request but uses the same Actual bytes.
    practice_sources = sorted((tmp_path / "kit/files/practice").iterdir())
    assert {path.name for path in practice_sources} == {
        "actual-sales.csv",
        f"assumptions-{language}.md",
    }
    imports = [
        ledger.import_document(client_root, client, engagement, path, "source")
        for path in practice_sources
    ]
    prepared = ledger.prepare_run(
        client_root,
        client,
        engagement,
        "sales-plan",
        context["workflow_version"],
        input_ids=[item["receipt"]["input_id"] for item in imports],
        new_run=True,
        purpose="Fictional alternative scenario requested by the practice exercise",
    )
    alternative = ledger.start_run(client_root, engagement, prepared["run"]["run_id"])
    _, alternative_plan = _execute_teaching_sales_plan(
        alternative, client_root, "5", "90772.5", "33442.5", language, first_plan
    )
    assert alternative["context"]["engagement_id"] == engagement
    assert alternative["context"]["run_id"] != context["run_id"]
    assert all(path.read_bytes() == data for path, data in preserved.items())
    assert _read(first_plan / "plan_execution_receipt.json")["report_ready"] is False
    assert alternative_plan != first_plan
    record_property("teaching_output", str(alternative_plan))
    record_native_check(
        record_property,
        root=ROOT,
        product="vera",
        workflow="sales-plan",
        language=language,
        phase=phase,
    )
