from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path
from typing import Any

import polars as pl
import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
PLUGIN_ROOT = ROOT / "plugins" / "clara" / "modules" / "reporting-engine"
FIXTURE_ROOT = PLUGIN_ROOT / "fixtures" / "semantic_layer"


def _load_module(name: str, path: Path) -> Any:
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _intake_module() -> Any:
    return _load_module(
        "reporting_engine_dataset_intake_test",
        PLUGIN_ROOT / "scripts" / "dataset_intake.py",
    )


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


@pytest.mark.parametrize(
    ("suffix", "expected_format"),
    (
        (".csv", "csv"),
        (".xlsx", "xlsx"),
        (".parquet", "parquet"),
    ),
)
def test_first_upload_profiles_each_supported_dataset_format(
    tmp_path: Path,
    suffix: str,
    expected_format: str,
) -> None:
    intake = _intake_module()
    dataset = tmp_path / f"business_metrics{suffix}"
    rows = {
        "Month": ["2026-01-01", "2026-02-01"],
        "Sales": [100, 120],
        "Discount": [5, 6],
        "COGS": [60, 72],
    }
    if suffix == ".csv":
        pl.DataFrame(rows).write_csv(dataset)
    elif suffix == ".xlsx":
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(list(rows))
        for values in zip(*rows.values(), strict=True):
            worksheet.append(list(values))
        workbook.save(dataset)
    else:
        pl.DataFrame(rows).write_parquet(dataset)

    receipt = intake.run_dataset_intake(
        dataset,
        dataset_contract_id=f"business_metrics_{expected_format}",
        output_dir=tmp_path / f"intake_{expected_format}",
    )

    profile = _read_json(
        tmp_path / f"intake_{expected_format}" / "dataset_profile.json"
    )
    assert receipt["status"] == "review_required"
    assert profile["source"]["format"] == expected_format


def test_first_upload_never_auto_maps_literal_business_metric_headers(
    tmp_path: Path,
) -> None:
    intake = _intake_module()
    dataset = tmp_path / "literal_headers.csv"
    dataset.write_text(
        "Month,Sales,Discount,COGS\n" "2026-01-01,100,5,60\n" "2026-02-01,120,6,72\n",
        encoding="utf-8",
    )
    output_dir = tmp_path / "intake"

    receipt = intake.run_dataset_intake(
        dataset,
        dataset_contract_id="literal_business_metrics",
        output_dir=output_dir,
    )

    assert receipt["status"] == "review_required"
    assert {path.name for path in output_dir.iterdir() if path.is_file()} >= {
        "dataset_profile.json",
        "semantic_layer.draft.json",
        "semantic_authoring_context.json",
        "dataset_intake.json",
    }
    layer = _read_json(output_dir / "semantic_layer.draft.json")
    assert set(layer["business_metric_mappings"]) == {"sales", "discount", "cogs"}
    assert all(
        mapping["state"] == "unknown"
        and mapping["metric_id"] is None
        and mapping["candidate_metric_ids"] == []
        for mapping in layer["business_metric_mappings"].values()
    )
    assert {
        metric["binding"]["column"]
        for metric in layer["metrics"]
        if metric["binding"]["binding_type"] == "column"
    } >= {"Sales", "Discount", "COGS"}
    assert _read_json(output_dir / "dataset_intake.json") == receipt


def test_first_upload_surfaces_ambiguous_metric_candidates_without_choosing(
    tmp_path: Path,
) -> None:
    intake = _intake_module()
    dataset = tmp_path / "ambiguous_sales.csv"
    dataset.write_text(
        "Month,Gross Sales,Net Sales,Discount Amount,Cost of Sales\n"
        "2026-01-01,120,100,20,60\n"
        "2026-02-01,150,125,25,75\n",
        encoding="utf-8",
    )
    output_dir = tmp_path / "intake"

    receipt = intake.run_dataset_intake(
        dataset,
        dataset_contract_id="ambiguous_business_metrics",
        output_dir=output_dir,
    )

    context = _read_json(output_dir / "semantic_authoring_context.json")
    review = context["business_metric_mapping_review"]
    candidate_ids = {
        candidate["metric_id"] for candidate in review["metric_candidates"]
    }
    assert receipt["status"] == "review_required"
    assert set(review["role_guidance"]) == {"sales", "discount", "cogs"}
    assert {"metric.gross_sales", "metric.net_sales"} <= candidate_ids
    assert all(
        mapping["state"] == "unknown" for mapping in review["current_mappings"].values()
    )


def test_reviewed_business_metric_mapping_is_reused_on_later_snapshot(
    tmp_path: Path,
) -> None:
    intake = _intake_module()
    reviewed_layer_path = FIXTURE_ROOT / "retail_monthly.semantic.json"
    original_layer_bytes = reviewed_layer_path.read_bytes()
    reviewed_layer = _read_json(reviewed_layer_path)
    output_dir = tmp_path / "refresh"

    receipt = intake.run_dataset_intake(
        FIXTURE_ROOT / "retail_monthly_refresh.csv",
        dataset_contract_id="retail_monthly",
        output_dir=output_dir,
        semantic_layer_path=reviewed_layer_path,
    )

    attachment = _read_json(output_dir / "snapshot_attachment.json")
    context = _read_json(output_dir / "semantic_authoring_context.json")
    assert receipt["status"] == "mapping_reused"
    assert attachment["attachment_status"] == "attached"
    assert attachment["semantic_version"] == reviewed_layer["semantic_version"] == 1
    assert (
        context["semantic_layer_draft"]["business_metric_mappings"]
        == reviewed_layer["business_metric_mappings"]
    )
    assert context["business_metric_mapping_review"]["current_mappings"] == (
        reviewed_layer["business_metric_mappings"]
    )
    assert reviewed_layer_path.read_bytes() == original_layer_bytes


def test_cli_rejects_semantic_layer_from_another_dataset_contract(
    tmp_path: Path,
) -> None:
    intake = _intake_module()
    output_dir = tmp_path / "mismatch"

    return_code = intake.main(
        [
            str(FIXTURE_ROOT / "retail_monthly_refresh.csv"),
            "--dataset-contract-id",
            "another_retail_asset",
            "--output-dir",
            str(output_dir),
            "--semantic-layer",
            str(FIXTURE_ROOT / "retail_monthly.semantic.json"),
        ]
    )

    receipt = _read_json(output_dir / "dataset_intake.json")
    assert return_code == 1
    assert receipt["status"] == "rejected"
    assert receipt["dataset_contract_id"] == "another_retail_asset"
    assert receipt["semantic_layer_dataset_contract_id"] == "retail_monthly"


def test_intake_refuses_to_overwrite_existing_run_artifacts(tmp_path: Path) -> None:
    intake = _intake_module()
    output_dir = tmp_path / "existing"
    output_dir.mkdir()
    receipt_path = output_dir / "dataset_intake.json"
    original_content = '{"status":"keep"}\n'
    receipt_path.write_text(original_content, encoding="utf-8")

    with pytest.raises(FileExistsError, match="will not overwrite"):
        intake.run_dataset_intake(
            FIXTURE_ROOT / "retail_monthly.csv",
            dataset_contract_id="retail_monthly",
            output_dir=output_dir,
        )

    assert receipt_path.read_text(encoding="utf-8") == original_content


@pytest.mark.parametrize(
    "cell_value, error",
    [("=1+2", "Uncached Excel formula"), ("#DIV/0!", "Excel cell error")],
)
def test_intake_rejects_unresolved_excel_values(
    tmp_path: Path, cell_value: str, error: str
) -> None:
    intake = _intake_module()
    workbook = Workbook()
    workbook.active.append(["Customer", "Revenue"])
    workbook.active.append(["Synthetic", cell_value])
    dataset = tmp_path / "unresolved.xlsx"
    workbook.save(dataset)
    output = tmp_path / "intake"

    with pytest.raises(ValueError, match=error):
        intake.run_dataset_intake(
            dataset, dataset_contract_id="synthetic", output_dir=output
        )

    assert not (output / "dataset_profile.json").exists()


def test_csv_intake_retains_explicit_delimiter_and_decimal_contract(
    tmp_path: Path,
) -> None:
    dataset = tmp_path / "regional.csv"
    dataset.write_text("Month;Revenue\n2026-01-01;12,50\n", encoding="utf-8")
    intake = _intake_module()
    output = tmp_path / "intake"

    intake.run_dataset_intake(
        dataset,
        dataset_contract_id="regional",
        output_dir=output,
        csv_options={"separator": ";", "decimal_comma": True},
    )

    profile = _read_json(output / "dataset_profile.json")
    assert profile["source"]["parser_options"]["separator"] == ";"
    assert profile["source"]["parser_options"]["decimal_comma"] is True
    assert profile["columns"]["Revenue"]["sample_values"] == [12.5]


def test_excel_blank_is_retained_and_duplicate_headers_have_explicit_names(
    tmp_path: Path,
) -> None:
    profile = _load_module(
        "clara_profile_blank_test", PLUGIN_ROOT / "scripts/profile_dataset.py"
    )
    workbook = Workbook()
    workbook.active.append(["Revenue", "Revenue", "Label"])
    workbook.active.append([None, 12, "Synthetic"])
    dataset = tmp_path / "blank.xlsx"
    workbook.save(dataset)

    frame, source = profile.load_dataset_frame(dataset, sheet_name=None)

    assert frame.to_dicts() == [
        {"Revenue": None, "Revenue_2": 12, "Label": "Synthetic"}
    ]
    assert source["formula_count"] == 0
    assert source["formula_cache"] == "no_formulas"


def test_empty_excel_sheet_rejects_and_closes_both_workbooks(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    profile = _load_module(
        "clara_profile_close_test", PLUGIN_ROOT / "scripts/profile_dataset.py"
    )
    dataset = tmp_path / "empty.xlsx"
    Workbook().save(dataset)
    original = profile.load_workbook
    closed = []

    def tracked_load(*args, **kwargs):
        workbook = original(*args, **kwargs)
        close = workbook.close

        def tracked_close():
            closed.append(kwargs["data_only"])
            close()

        workbook.close = tracked_close
        return workbook

    monkeypatch.setattr(profile, "load_workbook", tracked_load)

    with pytest.raises(ValueError, match="no header row"):
        profile.load_dataset_frame(dataset, sheet_name=None)

    assert closed == [True, False]


def test_late_csv_type_conflict_fails_without_silent_null_substitution(
    tmp_path: Path,
) -> None:
    profile = _load_module(
        "clara_profile_late_csv_test", PLUGIN_ROOT / "scripts/profile_dataset.py"
    )
    dataset = tmp_path / "late.csv"
    dataset.write_text("Revenue\n" + "12\n" * 10000 + "not-a-number\n")

    with pytest.raises(pl.exceptions.ComputeError, match="not-a-number"):
        profile.load_dataset_frame(dataset, sheet_name=None)


def test_excel_cached_formula_is_reported_without_claiming_recalculation(
    tmp_path: Path,
) -> None:
    from zipfile import ZipFile

    profile = _load_module(
        "clara_profile_cached_formula_test", PLUGIN_ROOT / "scripts/profile_dataset.py"
    )
    workbook = Workbook()
    workbook.active.append(["Revenue"])
    workbook.active.append(["=1+2"])
    original = tmp_path / "original.xlsx"
    workbook.save(original)
    dataset = tmp_path / "cached.xlsx"
    with ZipFile(original) as source, ZipFile(dataset, "w") as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"<f>1+2</f><v></v>", b"<f>1+2</f><v>3</v>")
            target.writestr(info, content)

    frame, source = profile.load_dataset_frame(dataset, sheet_name=None)

    assert frame.to_dicts() == [{"Revenue": 3}]
    assert source["formula_count"] == 1
    assert source["formula_cache"] == "present_not_recalculated"
