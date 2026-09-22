from __future__ import annotations

import hashlib
import importlib.util
import json
import socket
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfWriter

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "plugins/lucia/scripts/legal_documents.py"
SPEC = importlib.util.spec_from_file_location("lucia_legal_documents", SCRIPT)
legal = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = legal
SPEC.loader.exec_module(legal)


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def refuse(*args, **kwargs):
        raise AssertionError("Document workflows must not use network access")

    monkeypatch.setattr(socket, "create_connection", refuse)
    monkeypatch.setattr(socket.socket, "connect", refuse)


def text_pack(
    tmp_path,
    *,
    workflow="revisione-contratti",
    text="Payment in 30 days.\nLiability is capped at EUR 1000.\n",
):
    source = tmp_path / "contract.txt"
    source.write_text(text)
    run = tmp_path / "run"
    files = [source]
    if workflow == "confronto-documenti":
        second = tmp_path / "second.txt"
        second.write_text(text)
        files.append(second)
    pack = legal.prepare(run, files, workflow, ["Payment"])
    review = json.loads((run / "review.json").read_text(encoding="utf-8"))
    review["context"].update(
        represented_party="Customer",
        jurisdiction="Unknown; no enforceability opinion",
        instructions="Review payment terms",
        relationship="Customer role known; purchase purpose not supplied",
        formation="Negotiation history not supplied",
        forum="Not established",
        legal_basis="No legal sources supplied; textual review only",
        firm_instructions="No firm playbook supplied",
        language="en",
    )
    for source, item in zip(pack["sources"], review["items"]):
        review["coverage"][source["id"]]["reviewed_anchors"] = [
            unit["anchor"] for unit in source["units"]
        ]
        item.update(
            status="supported",
            summary="Payment in 30 days",
            reasoning="The customer has 30 days to pay.",
            proposal="Confirm commercial acceptance.",
            citations=[
                {
                    "source_id": source["id"],
                    "anchor": "line:1",
                    "quote": "Payment in 30 days.",
                }
            ],
        )
    return run, review


def write_review(run, review):
    path = run / "review.json"
    path.write_text(json.dumps(review))
    return path


def test_prepare_keeps_long_document_tail_and_original(tmp_path):
    content = "Ordinary provision.\n" * 5000 + "Liability is uncapped.\n"
    run, _ = text_pack(tmp_path, text=content)

    pack = legal.read_pack(run)

    assert pack["sources"][0]["units"][-1] == {
        "anchor": "line:5001",
        "text": "Liability is uncapped.",
    }
    assert (run / "originals/D001.txt").read_text(encoding="utf-8") == content


@pytest.mark.parametrize("filename", ["evidence.json", "originals/D001.txt"])
def test_changed_evidence_or_original_rejects_use(tmp_path, filename):
    run, _ = text_pack(tmp_path)
    (run / filename).write_text("changed")

    with pytest.raises(ValueError, match="changed"):
        legal.read_pack(run)


def test_existing_run_is_not_overwritten(tmp_path):
    run, _ = text_pack(tmp_path)

    with pytest.raises(ValueError, match="already exists"):
        legal.prepare(
            run, [tmp_path / "contract.txt"], "revisione-contratti", ["Payment"]
        )


def test_unreadable_and_unsupported_sources_remain_in_pack(tmp_path):
    bad = tmp_path / "corrupt.pdf"
    bad.write_bytes(b"Not a PDF")
    unknown = tmp_path / "old.doc"
    unknown.write_bytes(b"legacy")

    pack = legal.prepare(
        tmp_path / "run", [bad, unknown], "revisione-documentale", ["Liability"]
    )

    assert [source["id"] for source in pack["sources"]] == ["D001", "D002"]
    assert "PDF could not be read" in pack["sources"][0]["error"]
    assert "Unsupported format .doc" in pack["sources"][1]["error"]


def test_blank_pdf_page_cannot_establish_absence(tmp_path):
    source = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    writer.write(source)
    run = tmp_path / "run"
    legal.prepare(run, [source], "revisione-documentale", ["Term"])
    review = json.loads((run / "review.json").read_text(encoding="utf-8"))
    review["context"].update(
        represented_party="Buyer",
        jurisdiction="Unknown",
        instructions="Review",
        relationship="Unknown",
        formation="Unknown",
        forum="Unknown",
        legal_basis="No sources supplied",
        firm_instructions="None supplied",
    )
    review["coverage"]["D001"]["reviewed_anchors"] = ["page:1"]
    review["items"][0]["status"] = "not-stated"

    result = legal.validate_review(run, review)

    assert not result["valid"]
    assert any("absence" in error for error in result["errors"])


@pytest.mark.parametrize(
    ("quote", "valid"),
    [
        ("Payment  in\n30 days.", True),
        ("Payment in 15 days.", False),
        ("payment in 30 days.", False),
        ("", False),
    ],
)
def test_quote_verifies_literal_source_not_similar_claim(tmp_path, quote, valid):
    run, review = text_pack(tmp_path)
    review["items"][0]["citations"][0]["quote"] = quote

    result = legal.validate_review(run, review)

    assert result["valid"] is valid


def test_quote_must_use_the_correct_anchor(tmp_path):
    run, review = text_pack(tmp_path)
    review["items"][0]["citations"][0]["anchor"] = "line:2"

    result = legal.validate_review(run, review)

    assert not result["valid"]
    assert result["quote_checks"][0]["matched"] is False


@pytest.mark.parametrize(
    "mutation", ["missing", "duplicate", "unknown_anchor", "wrong_shape"]
)
def test_invalid_cells_or_coverage_are_rejected(tmp_path, mutation):
    run, review = text_pack(tmp_path)
    if mutation == "missing":
        review["items"] = []
    elif mutation == "duplicate":
        review["items"].append(review["items"][0].copy())
    elif mutation == "unknown_anchor":
        review["coverage"]["D001"]["reviewed_anchors"].append("line:999")
    else:
        review["items"] = {"not": "a list"}

    result = legal.validate_review(run, review)

    assert not result["valid"]
    assert result["errors"]


def test_partial_review_preserves_limits_context_and_sources_in_all_outputs(tmp_path):
    run, review = text_pack(tmp_path)
    limitation = "The referenced schedule and signing authority were not supplied."
    review["coverage"]["D001"]["limitations"] = [limitation]
    review["takeaways"] = ["Obtain the missing schedule."]

    target = legal.render(run, write_review(run, review))

    report = target.read_text(encoding="utf-8")
    assert limitation in report
    assert "Partial review" in report
    assert "Customer" in report
    workbook = load_workbook(run / "review.xlsx")
    context = str(list(workbook["Context"].values))
    assert limitation in context
    assert "Customer" in context
    assert "quotation occurrence only" in context
    assert "Obtain the missing schedule" in context
    assert "Original SHA-256" in str(list(workbook["Sources"].values))
    assert limitation in (run / "review.csv").read_text(encoding="utf-8")


def test_invalid_rerender_retires_previous_outputs(tmp_path):
    run, review = text_pack(tmp_path)
    legal.render(run, write_review(run, review))
    review["items"][0]["citations"][0]["quote"] = "invented"

    with pytest.raises(ValueError, match="validation failed"):
        legal.render(run, write_review(run, review))

    assert not (run / "review.html").exists()
    assert not (run / "delivery.json").exists()
    assert list(run.glob("previous-*-review.html"))


def test_render_escapes_html_and_disables_spreadsheet_formulas(tmp_path):
    run, review = text_pack(tmp_path)
    review["items"][0]["proposal"] = '=HYPERLINK("https://example.invalid","open")'
    review["takeaways"] = ["<script>alert('client')</script>"]

    target = legal.render(run, write_review(run, review))

    assert "<script>" not in target.read_text(encoding="utf-8")
    assert "&lt;script&gt;" in target.read_text(encoding="utf-8")
    workbook = load_workbook(run / "review.xlsx")
    assert workbook["Findings"]["F2"].data_type == "s"
    assert workbook["Findings"]["F2"].value.startswith("'=")


def test_comparison_delivers_topic_by_document_with_explicit_difference(tmp_path):
    run, review = text_pack(tmp_path, workflow="confronto-documenti")
    review["differences"] = {"Payment": "Both versions retain the same payment term."}

    target = legal.render(run, write_review(run, review))

    assert "Difference and implication" in target.read_text(encoding="utf-8")
    workbook = load_workbook(run / "review.xlsx")
    assert list(workbook["Comparison"].values)[0] == (
        "Topic",
        "D001 · contract.txt",
        "D002 · second.txt",
        "Difference and implication",
    )
    assert workbook["Comparison"]["D2"].value == review["differences"]["Payment"]
    assert review["differences"]["Payment"] in (run / "review.csv").read_text(
        encoding="utf-8"
    )


def test_comparison_requires_difference_explanations(tmp_path):
    run, review = text_pack(tmp_path, workflow="confronto-documenti")

    result = legal.validate_review(run, review)

    assert not result["valid"]


def test_literal_comparison_finds_late_document_change(tmp_path):
    before = tmp_path / "v1.txt"
    after = tmp_path / "v2.txt"
    before.write_text("Unchanged\n" * 8000 + "Notice: 30 days\n")
    after.write_text("Unchanged\n" * 8000 + "Notice: 5 days\n")
    run = tmp_path / "run"
    legal.prepare(run, [before, after], "confronto-documenti", ["Notice"])

    result = legal.compare(run, "D001", "D002")

    assert "-Notice: 30 days" in result.read_text(encoding="utf-8")
    assert "+Notice: 5 days" in result.read_text(encoding="utf-8")


def test_docx_draft_preserves_original_unmodified_parts_and_run_styles(tmp_path):
    source = tmp_path / "template.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Client: ").italic = True
    paragraph.add_run("[CLI").bold = True
    paragraph.add_run("ENT]")
    paragraph.add_run(". Payment unchanged.").underline = True
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Confidential schedule"
    document.sections[0].header.paragraphs[0].text = "Private draft"
    document.save(source)
    original_bytes = source.read_bytes()
    run = tmp_path / "run"
    pack = legal.prepare(run, [source], "redazione-da-modello", ["Client"])
    changes = run / "changes.json"
    changes.write_text(
        json.dumps(
            [
                {
                    "anchor": "word/document.xml#p1",
                    "old": "[CLIENT]",
                    "new": "Alfa S.r.l.",
                    "basis": "Supplied client identity",
                }
            ]
        )
    )

    target = legal.draft(run, "D001", changes)

    assert source.read_bytes() == original_bytes
    changed = Document(target)
    assert changed.paragraphs[0].text == "Client: Alfa S.r.l.. Payment unchanged."
    assert changed.paragraphs[0].runs[0].italic
    assert changed.paragraphs[0].runs[1].bold
    assert changed.paragraphs[0].runs[-1].underline
    assert any(
        unit["text"] == "Confidential schedule" for unit in pack["sources"][0]["units"]
    )
    assert any(unit["text"] == "Private draft" for unit in pack["sources"][0]["units"])
    with ZipFile(source) as original, ZipFile(target) as edited:
        assert {
            name: original.read(name)
            for name in original.namelist()
            if name != "word/document.xml"
        } == {
            name: edited.read(name)
            for name in edited.namelist()
            if name != "word/document.xml"
        }


def test_draft_rejects_ambiguous_replacement_without_output(tmp_path):
    run, _ = text_pack(
        tmp_path, text="[CLIENT] and [CLIENT]\n", workflow="redazione-da-modello"
    )
    changes = run / "changes.json"
    changes.write_text(
        json.dumps(
            [
                {
                    "anchor": "line:1",
                    "old": "[CLIENT]",
                    "new": "Alfa",
                    "basis": "Client instruction",
                }
            ]
        )
    )

    with pytest.raises(ValueError, match="ambiguous"):
        legal.draft(run, "D001", changes)

    assert not (run / "draft.txt").exists()


@pytest.mark.parametrize("workflow", legal.WORKFLOWS)
def test_upstream_snapshots_retain_pinned_bytes_and_mit_notice(workflow):
    directory = ROOT / "plugins/lucia/skills" / workflow / "references/upstream"
    provenance = json.loads((directory / "PROVENANCE.json").read_text(encoding="utf-8"))

    actual = {
        entry["bundled_file"]: hashlib.sha256(
            (directory / entry["bundled_file"]).read_bytes()
        ).hexdigest()
        for entry in provenance["files"]
    }

    assert actual == {
        entry["bundled_file"]: entry["sha256"] for entry in provenance["files"]
    }
    assert provenance["commit"] == "ce62e6a2d3f47e1d3567a4f2edc61898cfe9e78a"
    assert "Copyright (c) 2026 Mike" in (directory / "LICENSE").read_text(
        encoding="utf-8"
    )


@pytest.mark.parametrize(
    "archive_name",
    ["lucia-plugin.zip", "lucia-chatgpt-upload.zip", "lucia-claude-plugin.zip"],
)
def test_shipped_helper_runs_with_its_schema_and_labels(tmp_path, archive_name):
    archive_path = ROOT / "plugin_packages/lucia" / archive_name
    script_names = (
        "legal_documents.py",
        "legal_documents_review.schema.json",
        "legal_documents_labels.json",
    )
    with ZipFile(archive_path) as archive:
        for filename in script_names:
            member = next(
                name
                for name in archive.namelist()
                if name.endswith("scripts/" + filename)
            )
            (tmp_path / filename).write_bytes(archive.read(member))
    run, review = text_pack(tmp_path)
    write_review(run, review)

    result = subprocess.run(
        [
            sys.executable,
            str(tmp_path / "legal_documents.py"),
            "render",
            "--run-dir",
            str(run),
            "--review",
            str(run / "review.json"),
        ],
        capture_output=True,
        text=True,
        check=False,
        timeout=30,
    )

    assert result.returncode == 0, result.stderr
    assert (run / "delivery.json").is_file()
    assert "Customer" in (run / "review.html").read_text(encoding="utf-8")


@pytest.mark.parametrize("language", ["it", "en", "fr", "de", "es"])
def test_italian_review_basis_survives_every_export_in_any_language(tmp_path, language):
    run, review = text_pack(tmp_path)
    review["coverage"]["D001"]["limitations"] = ["Factual questions remain open."]
    review["context"].update(
        language=language,
        jurisdiction="English law in clause 8",
        forum="Milan courts in clause 9; distinct from governing law",
        relationship="Alfa S.r.l. receives confidential information for its business",
        formation="Negotiation emails not provided; standard-form status unresolved",
        legal_basis="2026-09-22: no authority supplied; enforceability remains open",
        firm_instructions="Studio NDA v3: EUR 200000 cap is a negotiation preference",
    )

    target = legal.render(run, write_review(run, review))

    labels = json.loads(
        SCRIPT.with_name("legal_documents_labels.json").read_text(encoding="utf-8")
    )[language]
    workbook = load_workbook(run / "review.xlsx")
    exported_context = dict(list(workbook[labels["context"]].values)[1:])
    assert (
        exported_context[labels["context_labels"]["jurisdiction"]]
        == "English law in clause 8"
    )
    assert (
        exported_context[labels["context_labels"]["forum"]]
        == review["context"]["forum"]
    )
    assert review["context"]["legal_basis"] in target.read_text(encoding="utf-8")
    assert labels["coverage_incomplete"] in target.read_text(encoding="utf-8")
    assert review["context"]["firm_instructions"] in (run / "review.csv").read_text(
        encoding="utf-8"
    )
    assert (
        exported_context[labels["context_labels"]["formation"]]
        == review["context"]["formation"]
    )


def test_missing_legal_basis_cannot_disappear_from_report_context(tmp_path):
    run, review = text_pack(tmp_path)
    del review["context"]["legal_basis"]

    result = legal.validate_review(run, review)

    assert not result["valid"]
    assert "legal_basis" in str(result["errors"])


def test_unknown_law_and_relationship_allow_explicitly_limited_review(tmp_path):
    run, review = text_pack(tmp_path)
    review["context"].update(jurisdiction="Unknown", relationship="Unknown")

    result = legal.validate_review(run, review)

    assert result["valid"]
    assert review["context"]["jurisdiction"] == "Unknown"
    assert (
        review["context"]["legal_basis"]
        == "No legal sources supplied; textual review only"
    )
