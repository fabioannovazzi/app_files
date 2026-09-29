"""Render the same validated valuation register into local reviewable workpapers."""

from __future__ import annotations

import csv
import hashlib
import html
import json
from pathlib import Path

from valuation_case import build_valuation, require
from valuation_engine import decimal

__all__ = ["write_package", "report_sections", "write_workbook", "compile_html"]

NOTICE = "Carte di lavoro per revisione professionale. Non costituiscono perizia firmata, attestazione di conformità PIV o qualificazione della finalità legale."
STATUS = {
    "partial": "Incompleto",
    "blocked": "Bloccato",
    "ready_for_professional_review": "Da rivedere",
    "accepted_workpaper": "Revisione registrata",
    "excluded": "Escluso",
    "draft": "Bozza",
}


def display(value: str) -> str:
    """Round only the reader-facing display; retain original decimal strings."""
    return f"{decimal(value):,.2f}"


def report_sections(report: dict) -> list[tuple[str, list[str]]]:
    """Produce one reader-facing outline shared by every document format."""
    case = report["case"]
    mandate = case["mandate"]
    amounts = {row["id"]: row for row in report["calculations"]}
    sections = [
        (
            "Incarico e perimetro",
            [
                f"{case['entity_name']} · {mandate['subject']}",
                f"Finalità: {mandate['purpose']}",
                f"Data valutativa: {mandate['valuation_date']} · Limite informativo: {mandate['information_cutoff']}",
                f"Configurazione: {mandate['basis_of_value']} · Premessa: {mandate['premise']}",
                f"Diritti e perimetro della quota: {mandate['rights']}",
                f"Destinatari: {case['audience']} · Valuta: {case['currency']}",
                f"Stato: {STATUS[report['status']]}",
                (
                    "Dati sintetici di prova."
                    if case["synthetic"]
                    else "Dati del caso selezionato."
                ),
                NOTICE,
            ],
        )
    ]
    for method in report["methods"]:
        rows = [
            f"Stato: {STATUS[method['status']]}",
            method.get("rationale", method.get("reason", "")),
        ]
        if "value_id" in method:
            value = amounts[method["value_id"]]
            basis = (
                "Valore operativo"
                if method["value_type"] == "operating_enterprise"
                else "Valore del capitale proprio"
            )
            rows.append(
                f"{basis}: {display(value['value'])} {case['currency']} [{value['id']}]"
            )
            if method["equity_id"] and method["equity_id"] != method["value_id"]:
                equity = amounts[method["equity_id"]]
                rows.append(
                    f"Capitale proprio dopo raccordo: {display(equity['value'])} {case['currency']} [{equity['id']}]"
                )
            elif method["equity_id"] is None:
                rows.append(
                    "Raccordo al capitale proprio non fornito; il valore operativo non è il valore della quota."
                )
            for label, ref in method["detail_ids"].items():
                rows.append(
                    f"{label}: {display(amounts[ref]['value'])} {amounts[ref]['unit']} [{ref}]"
                )
            rows.extend(method["limitations"])
            if method["stale_review"]:
                rows.append(
                    "La revisione precedente non è valida per queste dipendenze."
                )
        sections.append((f"Metodo {method['kind']} {method['method_id']}", rows))
    if report["sensitivity"]:
        rows = []
        for scenario in report["sensitivity"]:
            if scenario["status"] == "blocked":
                rows.append(f"{scenario['id']}: non calcolabile — {scenario['reason']}")
            else:
                row = amounts[scenario["equity_id"] or scenario["value_id"]]
                rows.append(
                    f"{scenario['id']}: {display(row['value'])} {case['currency']} [{row['id']}]"
                )
        sections.append(
            (
                "Sensitività delle ipotesi",
                [
                    "Risultati condizionati alle ipotesi dichiarate; non sono intervalli statistici.",
                    *rows,
                ],
            )
        )
    conclusion = report["conclusion"]
    sections.append(
        (
            "Confronto e conclusione",
            [
                (
                    f"{STATUS[conclusion['status']]}: {conclusion['text']}"
                    if conclusion
                    else "Conclusione professionale non ancora registrata. I risultati dei metodi rimangono separati; non è applicata una media automatica."
                )
            ],
        )
    )
    sections.append(
        (
            "Limiti e verifiche aperte",
            [
                mandate["professional_limitations"],
                *case["limitations"],
                *report["issues"],
                "DCF e metodo misto: periodi annuali interi a fine anno. Nessuna conversione automatica di periodi frazionari o mensili.",
                "La registrazione del revisore è una dichiarazione locale, non autenticazione dell'identità o firma professionale.",
            ],
        )
    )
    sections.append(
        (
            "Fonti e ipotesi",
            [
                f"{row['id']} · {row['description']} · {row['path']} · SHA-256 {row['sha256']}"
                for row in case["sources"]
            ]
            + [
                f"{row['id']}: {row['description']} · {row['value'] if row['value'] is not None else 'mancante'} {row['unit']} · {row['kind']} / {row['status']} · fonte {', '.join(row['source_ids'])} · {row['locator']}"
                for row in case["inputs"]
            ],
        )
    )
    sections.append(
        (
            "Quali dati arrivano al modello",
            [
                "Il modello può leggere i documenti selezionati, gli estratti di bilancio e piano, le ipotesi, le fonti benchmark, le rettifiche e le decisioni del professionista. Questi materiali possono contenere dati reali di clienti e persone; non vengono anonimizzati automaticamente.",
                "Il motore locale calcola i metodi e produce il registro e le esportazioni senza chiamare servizi esterni. La scelta e l'interpretazione dei dati restano nel contesto del modello del servizio utilizzato. Le fonti pubbliche sono ricercate dal modello usando solo quesiti pubblici. Il resoconto dei dati effettivamente letti è consegnato separatamente nel report privacy del run.",
            ],
        )
    )
    return sections


def compile_html(report: dict) -> str:
    """Escape every authored field; the report loads no remote assets or code."""
    escape = html.escape
    parts = [
        "<!doctype html><html lang='it'><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>Valutazione d’impresa</title><style>body{font:17px/1.6 Arial,sans-serif;color:#15253a;margin:0;background:#fff}main{max-width:920px;margin:60px auto;padding:0 28px}h1{font-size:40px;line-height:1.15}h2{font-size:24px;color:#002060;margin-top:40px}p{overflow-wrap:anywhere}section{border-top:1px solid #dce2e8;margin-top:32px}small{color:#526171}summary{cursor:pointer}table{border-collapse:collapse;width:100%;font-size:13px}td,th{padding:9px;text-align:left;border-bottom:1px solid #ddd;overflow-wrap:anywhere}code{font-size:12px}@media print{main{margin:0}section{break-inside:auto}h2{break-after:avoid}}</style><main>",
        f"<small>VERA · {escape(STATUS[report['status']])}</small><h1>Valutazione d’impresa</h1>",
    ]
    for heading, paragraphs in report_sections(report):
        parts.append(f"<section><h2>{escape(heading)}</h2>")
        parts.extend(f"<p>{escape(paragraph)}</p>" for paragraph in paragraphs)
        parts.append("</section>")
    parts.append("</main></html>")
    return "".join(parts)


def _literal(cell: object, value: str) -> None:
    cell.value = value
    cell.data_type = "s"


def write_workbook(path: Path, report: dict) -> None:
    """Link input cells and formula nodes; prevent formula injection from labels."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Sintesi"
    inputs = workbook.create_sheet("Dati")
    calculations = workbook.create_sheet("Calcoli")
    sources = workbook.create_sheet("Fonti")
    summary.append(["Valutazione d'impresa", report["case"]["entity_name"]])
    summary.append(["Stato", STATUS[report["status"]]])
    summary.append(["Uso", NOTICE])
    summary.append(["Metodo", "Base", "Valore", "Capitale proprio", "Stato"])
    inputs.append(
        [
            "ID",
            "Descrizione",
            "Valore",
            "Unità",
            "Tipo",
            "Revisione",
            "Fonti",
            "Posizione",
            "Valore esatto",
        ]
    )
    input_rows = {}
    for index, item in enumerate(report["case"]["inputs"], 2):
        input_rows[item["id"]] = index
        inputs.append(
            [
                item["id"],
                item["description"],
                float(decimal(item["value"])) if item["value"] is not None else None,
                item["unit"],
                item["kind"],
                item["status"],
                ", ".join(item["source_ids"]),
                item["locator"],
                item["value"],
            ]
        )
        for column in (1, 2, 4, 5, 6, 7, 8, 9):
            _literal(
                inputs.cell(index, column), str(inputs.cell(index, column).value or "")
            )
    calculations.append(
        [
            "Calculation ID",
            "Formula",
            "Unità",
            "Valore esatto motore",
            "Input ID",
            "Source ID",
        ]
    )
    row_ids = {}
    for index, row in enumerate(report["calculations"], 2):
        args = [f"B{row_ids[ref]}" for ref in row["arguments"]]
        if row["op"] == "input":
            formula = f"='Dati'!C{input_rows[row['input_ids'][0]]}"
        elif row["op"] == "constant":
            formula = float(decimal(row["value"]))
        elif row["op"] == "sum":
            formula = f"=SUM({','.join(args)})" if args else "=0"
        else:
            symbol = {"subtract": "-", "multiply": "*", "divide": "/", "power": "^"}[
                row["op"]
            ]
            formula = f"={args[0]}{symbol}{args[1]}"
        calculations.append(
            [
                row["id"],
                formula,
                row["unit"],
                row["value"],
                ", ".join(row["input_ids"]),
                ", ".join(row["source_ids"]),
            ]
        )
        row_ids[row["id"]] = index
        for column in (1, 3, 4, 5, 6):
            _literal(
                calculations.cell(index, column),
                str(calculations.cell(index, column).value or ""),
            )
    for method in report["methods"]:
        value = (
            f"='Calcoli'!B{row_ids[method['value_id']]}"
            if "value_id" in method
            else None
        )
        equity = (
            f"='Calcoli'!B{row_ids[method['equity_id']]}"
            if method.get("equity_id")
            else None
        )
        summary.append(
            [
                method["kind"],
                method.get("value_type", ""),
                value,
                equity,
                STATUS[method["status"]],
            ]
        )
    sources.append(["ID", "Descrizione", "Percorso", "SHA-256", "Revisione"])
    for index, source in enumerate(report["case"]["sources"], 2):
        for column, key in enumerate(
            ("id", "description", "path", "sha256", "status"), 1
        ):
            _literal(sources.cell(index, column), source[key])
    _literal(summary["B1"], report["case"]["entity_name"])
    for sheet in workbook:
        sheet.freeze_panes = "C5" if sheet == summary else "C2"
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = (
            sheet.dimensions if sheet != summary else f"A4:E{sheet.max_row}"
        )
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.font = Font(
                    name="Arial",
                    size=10,
                    color="002060" if cell.data_type == "f" else "172737",
                )
                if cell.data_type == "n" or cell.data_type == "f":
                    cell.number_format = '#,##0.00;[Red](#,##0.00);"–"'
        for cell in sheet[4 if sheet == summary else 1]:
            cell.fill = PatternFill("solid", fgColor="002060")
            cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        for column in "ABCDEFGHI":
            sheet.column_dimensions[column].width = 24 if column not in "AB" else 40
        sheet.print_options.horizontalCentered = True
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    workbook.save(path)


def _write_documents(directory: Path, report: dict) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    document = Document()
    document.sections[0].top_margin = Inches(0.75)
    document.sections[0].bottom_margin = Inches(0.75)
    for name in ("Normal", "Title", "Heading 1"):
        document.styles[name].font.name = "Arial"
        document.styles[name].font.color.rgb = RGBColor.from_string(
            "000000"
            if name == "Title"
            else "002060" if name == "Heading 1" else "172737"
        )
    document.styles["Normal"].font.size = Pt(10)
    document.add_paragraph("Valutazione d’impresa", "Title")
    styles = getSampleStyleSheet()
    styles["Normal"].fontSize = 10
    styles["Normal"].leading = 14
    styles["Normal"].alignment = TA_LEFT
    content = [Paragraph("Valutazione d’impresa", styles["Title"]), Spacer(1, 12)]
    for heading, paragraphs in report_sections(report):
        document.add_heading(heading, level=1)
        content.append(Paragraph(html.escape(heading), styles["Heading2"]))
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
            content.extend(
                [Paragraph(html.escape(paragraph), styles["Normal"]), Spacer(1, 7)]
            )
    document.save(directory / "valuation_report.docx")

    def footer(canvas: object, doc: object) -> None:
        canvas.setFont("Helvetica", 8)
        canvas.drawString(
            42, 24, f"Vera · {STATUS[report['status']]} · Carte di lavoro"
        )
        canvas.drawRightString(A4[0] - 42, 24, str(doc.page))

    SimpleDocTemplate(
        str(directory / "valuation_report.pdf"),
        pagesize=A4,
        rightMargin=42,
        leftMargin=42,
        topMargin=42,
        bottomMargin=42,
    ).build(content, onFirstPage=footer, onLaterPages=footer)


def write_package(
    report: dict, source_root: Path, output: Path, *, replay_parent: Path | None = None
) -> list[dict]:
    """Replay before exporting, never overwrite a previous valuation revision."""
    require(
        report
        == build_valuation(report["case"], source_root, replay_parent=replay_parent),
        "Valuation differs from canonical replay",
    )
    require(
        not output.exists(),
        "Use a new revision directory; previous outputs are immutable",
    )
    output.mkdir(parents=True)
    (output / "valuation.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (output / "valuation_report.html").write_text(
        compile_html(report), encoding="utf-8"
    )
    (output / "valuation_report.md").write_text(
        "# Valutazione d’impresa\n\n"
        + "\n\n".join(
            f"## {heading}\n\n" + "\n\n".join(rows)
            for heading, rows in report_sections(report)
        ),
        encoding="utf-8",
    )
    with (output / "calculations.csv").open(
        "w", newline="", encoding="utf-8"
    ) as handle:
        writer = csv.writer(handle)
        writer.writerow(
            ["id", "value", "unit", "op", "arguments", "input_ids", "source_ids"]
        )
        for row in report["calculations"]:
            writer.writerow(
                [
                    row["id"],
                    row["value"],
                    row["unit"],
                    row["op"],
                    " | ".join(row["arguments"]),
                    " | ".join(row["input_ids"]),
                    " | ".join(row["source_ids"]),
                ]
            )
    write_workbook(output / "valuation_workbook.xlsx", report)
    _write_documents(output, report)
    artifacts = [
        {
            "path": path.name,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "bytes": path.stat().st_size,
        }
        for path in sorted(output.iterdir())
        if path.is_file()
    ]
    (output / "artifacts.json").write_text(
        json.dumps(artifacts, indent=2) + "\n", encoding="utf-8"
    )
    return artifacts
