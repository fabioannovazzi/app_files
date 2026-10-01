"""Consume the approved communications profile without importing client history."""

from __future__ import annotations

import base64
import hashlib
import json
import re
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from jsonschema import Draft202012Validator

__all__ = [
    "load_studio_format",
    "apply_studio_format",
    "display_number",
    "display_date",
]


def _digest(value: Any) -> str:
    return hashlib.sha256(
        json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            allow_nan=False,
        ).encode()
    ).hexdigest()


def load_studio_format(
    workspace: Path, *, studio_id: str, studio_name: str
) -> dict[str, Any]:
    """Read one explicitly selected, identity-bound approved studio profile."""
    root = workspace.expanduser().resolve()
    manifest = json.loads((root / "workspace.json").read_text(encoding="utf-8"))
    payload = json.loads((root / "studio_profile.json").read_text(encoding="utf-8"))
    if (
        Path(manifest["bound_path"]).resolve() != root
        or manifest["workspace_id"] != studio_id
        or payload["workspace_id"] != studio_id
        or payload["studio_name"] != studio_name
    ):
        raise ValueError(
            "Studio formatting identity or workspace binding does not match"
        )
    if payload.get("accepted_as_studio_standard") is not True or not payload.get(
        "approved_from"
    ):
        raise ValueError("Studio formatting requires an approved studio standard")
    fields = {
        key: payload[key]
        for key in ("studio_name", "brand_profile", "brand_assets", "profile")
    }
    if _digest(fields) != payload["format_digest"]:
        raise ValueError("Studio profile digest does not match its approved content")
    document = payload["profile"]["document"]
    layout = document["layout"]
    settings = {
        "font_family": {"Times": "Times New Roman", "Helvetica": "Arial"}.get(
            document["font_family"], document["font_family"]
        ),
        "body_font_size_pt": layout["body_font_size_pt"],
        "line_spacing_pt": layout["body_leading_pt"],
        "heading_1_size_pt": layout["heading_font_size_pt"],
        "heading_2_size_pt": layout["heading_font_size_pt"],
        "paragraph_before_pt": 0,
        "paragraph_after_pt": 6,
        "heading_color": "000000",
        "header_text": "",
        "footer_text": "",
        "page_numbers": False,
        "use_logo": False,
        "logo_width_mm": 30,
        "header_distance_mm": 8,
        "footer_distance_mm": 8,
        "table_font_size_pt": 9,
        "table_header_fill": "E8EEF5",
        "table_header_color": "000000",
        "table_alternate_fill": "FFFFFF",
        "number_format": "canonical",
        "date_format": "iso",
        "signature_lines": [],
        **document.get("docx", {}),
    }
    schema = json.loads(
        (
            Path(__file__).resolve().parents[1]
            / "assets"
            / "studio-docx-format.schema.json"
        ).read_text(encoding="utf-8")
    )
    Draft202012Validator(schema).validate(settings)
    if settings["line_spacing_pt"] < settings["body_font_size_pt"]:
        raise ValueError("Document line spacing is below the body font size")
    margins = {
        side: layout[f"{side}_margin_mm"] for side in ("top", "bottom", "left", "right")
    }
    if (
        any(
            not isinstance(n, (int, float)) or not 8 <= n <= 70
            for n in margins.values()
        )
        or margins["left"] + margins["right"] > 100
    ):
        raise ValueError("Studio report margins leave insufficient A4 content width")
    if (
        margins["top"] < settings["header_distance_mm"] + 8
        or margins["bottom"] < settings["footer_distance_mm"] + 5
        or (
            settings["use_logo"]
            and margins["top"] <= settings["header_distance_mm"] + 8
        )
    ):
        raise ValueError("Studio report margins leave insufficient header/footer room")
    logo = None
    if settings["use_logo"]:
        asset = payload["brand_assets"].get("logo")
        if not isinstance(asset, dict):
            raise ValueError("Approved DOCX logo was requested but is unavailable")
        path = (root / asset["workspace_relative_path"]).resolve()
        if not path.is_relative_to(
            root / "studio_assets"
        ) or path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
            raise ValueError("DOCX logo must be a PNG/JPEG within studio_assets")
        raw = path.read_bytes()
        if len(raw) > 5_000_000 or hashlib.sha256(raw).hexdigest() != asset["sha256"]:
            raise ValueError("Studio logo is too large or has changed since approval")
        logo = base64.b64encode(raw).decode("ascii")
    # Only presentation and attribution enter a client run: no samples, voice,
    # source register, history, or other clients' content.
    return {
        "studio_id": studio_id,
        "studio_name": studio_name,
        "version": payload["version"],
        "format_digest": payload["format_digest"],
        "settings": settings,
        "margins_mm": margins,
        "logo_base64": logo,
        "font_status": "requested_only_not_embedded; renderer availability must be checked",
    }


def display_number(value: str, settings: dict[str, Any]) -> str:
    """Group canonical decimals without rounding or changing their precision."""
    mode = settings.get("number_format", "canonical")
    if mode == "canonical":
        return value
    if not re.fullmatch(r"-?\d+(?:\.\d+)?", value):
        raise ValueError("Only canonical reviewed numeric totals may be formatted")
    sign = "-" if value.startswith("-") else ""
    integer, dot, fraction = value.lstrip("-").partition(".")
    groups = [integer[max(0, i - 3) : i] for i in range(len(integer), 0, -3)]
    thousands, decimal = (".", ",") if mode == "decimal_comma" else (",", ".")
    return sign + thousands.join(reversed(groups)) + (decimal + fraction if dot else "")


def display_date(value: str, settings: dict[str, Any]) -> str:
    """Format only an explicit ISO report date; leave period/source text alone."""
    parsed = date.fromisoformat(value)
    return parsed.strftime(
        {"iso": "%Y-%m-%d", "dmy": "%d/%m/%Y", "mdy": "%m/%d/%Y"}[
            settings["date_format"]
        ]
    )


def apply_studio_format(document: Any, profile: dict[str, Any]) -> None:
    """Apply approved presentation after report content has been constructed."""
    settings = profile["settings"]
    for section in document.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        for side, value in profile["margins_mm"].items():
            setattr(section, f"{side}_margin", Mm(value))
        section.header_distance = Mm(settings["header_distance_mm"])
        section.footer_distance = Mm(settings["footer_distance_mm"])
    for name in (
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "List Bullet",
        "Header",
        "Footer",
    ):
        style = document.styles[name]
        style.font.name = settings["font_family"]
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            if "theme" in attr.lower():
                del fonts.attrib[attr]
        fonts.set(qn("w:eastAsia"), settings["font_family"])
        fonts.set(qn("w:cs"), settings["font_family"])
        style.font.size = Pt(settings["body_font_size_pt"])
        style.paragraph_format.space_before = Pt(settings["paragraph_before_pt"])
        style.paragraph_format.space_after = Pt(settings["paragraph_after_pt"])
        style.paragraph_format.line_spacing = Pt(settings["line_spacing_pt"])
        style.paragraph_format.widow_control = True
        if name in {"Title", "Heading 1", "Heading 2"}:
            style.font.size = Pt(
                settings["heading_2_size_pt"]
                if name == "Heading 2"
                else settings["heading_1_size_pt"]
            )
            style.font.color.rgb = RGBColor.from_string(settings["heading_color"])
            style.paragraph_format.line_spacing = None
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(
                max(10, settings["paragraph_before_pt"])
            )
    header = document.sections[0].header.paragraphs[0]
    if settings["use_logo"]:
        run = header.add_run()
        shape = run.add_picture(
            BytesIO(base64.b64decode(profile["logo_base64"], validate=True)),
            width=Mm(settings["logo_width_mm"]),
        )
        # Scale a tall logo to leave room for the header text and body.
        max_height = Mm(
            profile["margins_mm"]["top"] - settings["header_distance_mm"] - 8
        )
        if shape.height > max_height:
            shape.width = int(shape.width * max_height / shape.height)
            shape.height = max_height
        if settings["header_text"]:
            header.add_run("\n")
    header.add_run(settings["header_text"])
    header.paragraph_format.line_spacing = 1
    header.paragraph_format.space_after = Pt(0)
    footer = document.sections[0].footer.paragraphs[0]
    footer.add_run(settings["footer_text"])
    if settings["page_numbers"]:
        footer.add_run("  ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
    footer.paragraph_format.line_spacing = 1
    footer.paragraph_format.space_after = Pt(0)
    for table in document.tables:
        numeric = bool(
            table._tbl.xpath('./w:tblPr/w:tblCaption[@w:val="Reviewed numeric totals"]')
        )
        metadata = bool(
            table._tbl.xpath('./w:tblPr/w:tblCaption[@w:val="Report metadata"]')
        )
        for row_index, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            header_row = row_index == 0 and not metadata
            if header_row:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for cell_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shaded = header_row or (metadata and cell_index == 0)
                for node in cell._tc.get_or_add_tcPr().findall(qn("w:shd")):
                    node.getparent().remove(node)
                fill = (
                    settings["table_header_fill"]
                    if shaded
                    else (
                        settings["table_alternate_fill"]
                        if row_index % 2 == 0
                        else "FFFFFF"
                    )
                )
                node = OxmlElement("w:shd")
                node.set(qn("w:fill"), fill)
                cell._tc.get_or_add_tcPr().append(node)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.1
                    if numeric and cell_index == 1:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    for run in paragraph.runs:
                        run.font.name = settings["font_family"]
                        run.font.size = Pt(settings["table_font_size_pt"])
                        run.font.color.rgb = RGBColor.from_string(
                            settings["table_header_color"] if shaded else "334155"
                        )
