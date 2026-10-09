"""Exact maintained passive-invoice job contracts for the native workspace.

This foundation is not yet an MCP entry point. Callers must supply the reviewed
source roles and controls; no source, accounting treatment or worker is inferred.
"""

from __future__ import annotations

import ast
import json
import sqlite3
import tempfile
from collections import Counter
from contextlib import closing
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from urllib.parse import quote
from zipfile import ZipFile

from native_bank_preparation import file_hash, tree_hash
from native_financial_analysis import json_page

__all__ = ["build_plan", "inspect_job", "run_job"]

WORKFLOW = "passive-invoice-audit"
ROLES = {"invoices", "ledger", "ledger_mapping", "history", "chart", "worker_selection"}
CONTROLS = {
    "ledger_sheet",
    "chunk_size",
    "concurrency",
    "max_retries",
    "reasoning_effort",
    "amount_tolerance",
}
MAX_FILE_BYTES = 256 * 1024 * 1024
MAX_ROWS = 100_000


def job_population(output: Path) -> dict:
    """Exclude only SQLite's transient shared-memory read-lock bookkeeping.

    Every physical file still passes the ordinary-file guard. The main database
    and complete WAL bytes remain checkpointed; no business output is excluded.
    """
    return {
        key: value
        for key, value in tree_hash(output).items()
        if key != "audit.sqlite3-shm"
    }


def implementation(producer: Any) -> dict:
    """Checkpoint actual producer/adapter code, not a claimed historical attestation."""
    scripts = Path(producer.__file__).resolve().parent
    files = [path for path in scripts.iterdir() if path.suffix in {".py", ".json"}]
    files.extend(
        Path(__file__).with_name(name)
        for name in (
            "native_passive_audit.py",
            "native_passive_audit_bridge.py",
            "native_bank_preparation.py",
            "native_financial_analysis.py",
        )
    )
    for base in (scripts.parent.parent, scripts.parent.parent.parent / "modules"):
        for relative in (
            "client-file-preparation/scripts/parse_fatturapa_xml.py",
            "journal-bank-reconciliation/scripts/semantic_review.py",
        ):
            path = base / relative
            if path.is_file():
                files.append(path)
    return {str(path): file_hash(path) for path in sorted(set(files))}


def verify_sources(plan: dict) -> None:
    """Recheck all authoritative receipts before/after execution or inspection."""
    for binding in plan["context"]["input_bindings"]:
        ordinary(Path(binding["path"]))
        if file_hash(Path(binding["path"])) != binding["sha256"]:
            raise ValueError(
                "Passive-invoice source changed during inspection or execution"
            )


def ordinary(path: Path, limit: int = MAX_FILE_BYTES) -> bytes:
    """Exact regular single-link bytes enforce the physical evidence boundary."""
    if any(parent.is_symlink() for parent in (path, *path.parents)):
        raise ValueError("Passive-invoice evidence cannot use linked paths")
    if not path.is_file() or path.stat().st_nlink != 1 or path.stat().st_size > limit:
        raise ValueError("Missing, linked or oversized passive-invoice evidence")
    return path.read_bytes()


def object_file(path: Path, limit: int = 5 * 1024 * 1024) -> dict:
    value = json.loads(ordinary(path, limit))
    if not isinstance(value, dict):
        raise ValueError("Expected a passive-invoice JSON object")
    return value


def jsonl(path: Path) -> list[dict]:
    rows = [json.loads(line) for line in ordinary(path).splitlines() if line.strip()]
    if len(rows) > MAX_ROWS or any(not isinstance(row, dict) for row in rows):
        raise ValueError("Passive-invoice population exceeds native inspection limit")
    return rows


def build_plan(producer: Any, context: dict, recipe: dict, workers: Any) -> dict:
    """Bind explicit roles to receipts and reuse the producer's fingerprint contract."""
    if context["workflow_id"] != WORKFLOW:
        raise PermissionError("Passive-invoice plan belongs to another workflow")
    if not isinstance(recipe, dict) or set(recipe) != {"inputs", "controls"}:
        raise ValueError("Expected exact passive-invoice inputs and controls")
    roles, controls = recipe["inputs"], recipe["controls"]
    if (
        not isinstance(roles, dict)
        or not {"invoices", "ledger", "ledger_mapping"} <= set(roles)
        or set(roles) - ROLES
        or not isinstance(controls, dict)
        or set(controls) != CONTROLS
    ):
        raise ValueError("Unknown or incomplete passive-invoice recipe")
    for key in ("chunk_size", "concurrency", "max_retries"):
        if type(controls[key]) is not int:
            raise ValueError("Passive-invoice execution counts must be integers")
    if controls["ledger_sheet"] is not None and (
        not isinstance(controls["ledger_sheet"], str)
        or not controls["ledger_sheet"].strip()
        or len(controls["ledger_sheet"]) > 200
    ):
        raise ValueError("Invalid explicit ledger sheet")
    if not isinstance(controls["reasoning_effort"], str):
        raise ValueError("Invalid explicit reasoning effort")
    tolerance = controls["amount_tolerance"]
    if not isinstance(tolerance, str) or len(tolerance) > 40:
        raise ValueError("Expected an explicit finite decimal tolerance")
    try:
        amount = Decimal(tolerance)
    except InvalidOperation as exc:
        raise ValueError("Expected an explicit finite decimal tolerance") from exc
    if not amount.is_finite():
        raise ValueError("Expected an explicit finite decimal tolerance")
    bindings = {row["binding_id"]: row for row in context["input_bindings"]}
    if len(bindings) != len(context["input_bindings"]):
        raise ValueError("Duplicate source binding")

    def resolve(identity: Any) -> Path:
        if not isinstance(identity, str) or identity not in bindings:
            raise PermissionError("Passive-invoice source is not an exact run receipt")
        row = bindings[identity]
        path = Path(row["path"])
        ordinary(path)
        if file_hash(path) != row["sha256"]:
            raise ValueError("Passive-invoice source receipt changed")
        return path

    invoice_ids = roles["invoices"]
    if (
        not isinstance(invoice_ids, list)
        or not invoice_ids
        or len(invoice_ids) > 10_000
        or any(not isinstance(item, str) for item in invoice_ids)
        or len(set(invoice_ids)) != len(invoice_ids)
    ):
        raise ValueError("Choose an exact nonduplicate invoice population")
    paths = {key: resolve(value) for key, value in roles.items() if key != "invoices"}
    invoices = [resolve(identity) for identity in invoice_ids]
    if len(invoices) == 1:
        invoice_source = invoices[0]
        if invoice_source.suffix.lower() not in {".xml", ".zip", ".json"}:
            raise ValueError("Unsupported invoice population format")
    else:
        invoice_source = Path(context["input_dir"])
        if any(path.suffix.lower() != ".xml" for path in invoices):
            raise ValueError("Multiple invoice sources must all be XML")
        physical = {
            path
            for path in invoice_source.rglob("*")
            if path.suffix.lower() == ".xml" and not path.is_dir()
        }
        if physical != set(invoices):
            raise ValueError(
                "Directory invoice population differs from explicit receipts"
            )
        for path in physical:
            ordinary(path)
    if invoice_source.suffix.lower() == ".json" and invoice_source.is_file():
        # The maintained CH-GE extractor verifies source locators and its review.
        # Additionally require every underlying document to be registered in this run.
        reviewed = object_file(invoice_source)
        for row in reviewed["invoices"]:
            relative = Path(row["source_path"])
            if relative.is_absolute() or ".." in relative.parts:
                raise PermissionError("Reviewed invoice source leaves its population")
            source = invoice_source.parent / relative
            source_bindings = [
                b for b in bindings.values() if Path(b["path"]) == source
            ]
            if len(source_bindings) != 1:
                raise PermissionError("Reviewed invoice document is not registered")
            resolve(source_bindings[0]["binding_id"])
    runtime = workers.configured_runtime()
    if runtime == "cowork-haiku":
        if "worker_selection" in paths or controls["reasoning_effort"] != "low":
            raise ValueError("Cowork Haiku rejects Codex model or effort overrides")
        model, effort, selection = "haiku", "low", None
    elif runtime == "codex-luna":
        model, effort, selection = workers.load_worker_selection(
            paths.get("worker_selection"),
            workflow_id=WORKFLOW,
            reasoning_effort=controls["reasoning_effort"],
        )
    else:
        raise ValueError("Unsupported packaged semantic runtime")
    config = producer.AuditConfig(
        chunk_size=controls["chunk_size"],
        concurrency=controls["concurrency"],
        max_retries=controls["max_retries"],
        amount_tolerance=amount,
        reasoning_effort=effort,
        worker_runtime="cowork" if runtime == "cowork-haiku" else "codex-native",
        worker_model=model,
        worker_selection=selection,
    )
    config.validate()
    history = jsonl(paths["history"]) if "history" in paths else []
    chart = object_file(paths["chart"]) if "chart" in paths else {}
    if any(
        not isinstance(key, str) or not isinstance(value, str)
        for key, value in chart.items()
    ):
        raise ValueError("Chart must be the maintained string-to-string object")
    fingerprint = producer._input_fingerprint(
        invoice_source,
        paths["ledger"],
        paths["ledger_mapping"],
        config,
        ledger_sheet=controls["ledger_sheet"],
        history=history,
        chart_of_accounts=chart,
    )
    return {
        "invoice_source": invoice_source,
        "paths": paths,
        "config": config,
        "history": history,
        "chart": chart,
        "fingerprint": fingerprint,
        "context": context,
        "recipe": recipe,
    }


def run_job(producer: Any, plan: dict, runner: Any) -> dict:
    """Use the unchanged public job lifecycle; never delete or replace its database.

    The future native launcher must authorize the run, retain durable intent and
    supervise long execution. This synchronous foundation supplies no shortcut
    around worker qualification or an incomplete Cowork handoff.
    """
    context = plan["context"]
    verify_sources(plan)
    if object_file(Path(context["run_manifest_path"]))["status"] != "running":
        raise ValueError("Resume the Archive run before executing an invoice job")
    result = producer.run_audit(
        invoice_source=plan["invoice_source"],
        ledger_path=plan["paths"]["ledger"],
        mapping_path=plan["paths"]["ledger_mapping"],
        output_dir=Path(context["output_dir"]),
        runner=runner,
        config=plan["config"],
        ledger_sheet=plan["recipe"]["controls"]["ledger_sheet"],
        history=plan["history"],
        chart_of_accounts=plan["chart"],
        client_run_id=context["run_id"],
        client_run_root=Path(context["run_root"]),
    )
    verify_sources(plan)
    return result


def inspect_job(producer: Any, plan: dict, args: dict) -> dict:
    """Replay mechanical matching and packet identities without invoking a model.

    The SQLite snapshot, complete JSONL population and public chunk checkpoints
    must agree. Consistency establishes provenance, not a professional decision
    or authenticated provider/model attestation.
    """
    if set(args) - {"view", "offset", "item_id", "source_ref", "revision"}:
        raise ValueError("Unknown passive-invoice inspection argument")
    offset = args.get("offset", 0)
    view = args.get("view", "exceptions")
    if (
        type(offset) is not int
        or offset < 0
        or view not in {"exceptions", "population", "chunks", "orphans"}
    ):
        raise ValueError("Unknown passive-invoice page or view")
    context, config = plan["context"], plan["config"]
    output = Path(context["output_dir"])
    population = job_population(output)
    if len(population) > 10_000:
        raise ValueError("Passive-invoice artifact count exceeds native limit")
    source_code = implementation(producer)
    checkpoint = producer._sha256_json(
        [context, plan["recipe"], population, source_code]
    )
    if args.get("revision") and args["revision"] != checkpoint:
        raise ValueError("Passive-invoice job changed; reopen the exact population")
    database = output / "audit.sqlite3"
    if not database.exists():
        if population:
            raise ValueError(
                "Unrecognized invoice outputs require specialist inspection"
            )
        return {
            "revision": checkpoint,
            "status": "not_started",
            "items": [],
            "total": 0,
            "selection": None,
        }
    ordinary(database)
    if (output / "audit.sqlite3-wal").exists() and not (
        output / "audit.sqlite3-shm"
    ).exists():
        raise ValueError("Live database is unavailable for read-only inspection")
    uri = "file:" + quote(str(database), safe="/") + "?mode=ro"
    if not (output / "audit.sqlite3-wal").exists():
        # A closed producer has no WAL; immutable reads create no lock sidecars.
        # The before/after full population detects a concurrent writer or new WAL.
        uri += "&immutable=1"
    with closing(sqlite3.connect(uri, uri=True, timeout=1)) as connection:
        connection.execute("PRAGMA query_only=ON")
        connection.execute("BEGIN")
        metadata = dict(connection.execute("SELECT key,value FROM metadata"))
        if (
            metadata.get("run_fingerprint") != plan["fingerprint"]
            or metadata.get("schema_version") != producer.SCHEMA_VERSION
        ):
            raise ValueError("SQLite job belongs to different inputs or controls")
        if (
            connection.execute("SELECT COUNT(*) FROM audit_items").fetchone()[0]
            > MAX_ROWS
        ):
            raise ValueError("Passive-invoice job exceeds native population limit")
        stored = connection.execute(
            "SELECT invoice_id,item_json,packet_json,semantic_json,semantic_model,semantic_effort,chunk_id,final_state,packet_sha256 FROM audit_items ORDER BY invoice_id"
        ).fetchall()
        chunk_rows = connection.execute(
            "SELECT chunk_id,packet_sha256,invoice_ids_json,status,attempt_count,error FROM chunks ORDER BY chunk_id"
        ).fetchall()
    chunks = [
        dict(
            zip(
                (
                    "chunk_id",
                    "packet_sha256",
                    "invoice_ids_json",
                    "status",
                    "attempt_count",
                    "error",
                ),
                row,
            )
        )
        for row in chunk_rows
    ]
    if len(chunks) > MAX_ROWS or len({row["chunk_id"] for row in chunks}) != len(
        chunks
    ):
        raise ValueError("Invalid chunk population")
    if any(
        row["status"]
        not in {"pending", "running", "completed", "failed", "awaiting_semantic_review"}
        for row in chunks
    ):
        raise ValueError("Unknown producer chunk state")
    if any(row["status"] in {"pending", "running"} for row in chunks):
        if job_population(output) != population:
            raise ValueError("Passive-invoice job changed during inspection")
        return {
            "revision": checkpoint,
            "status": "in_progress",
            "items": [],
            "total": 0,
            "selection": None,
            "chunk_counts": dict(Counter(row["status"] for row in chunks)),
            "published_summary_current": False,
        }
    summary = object_file(output / "run_summary.json")
    if (
        summary["run_fingerprint"] != plan["fingerprint"]
        or summary["client_run_id"] != context["run_id"]
        or summary["schema_version"] != producer.SCHEMA_VERSION
    ):
        raise ValueError("Invoice summary belongs to another exact run")
    rows = jsonl(output / "full_population.jsonl")
    by_id = {row["invoice"]["invoice_id"]: row for row in rows}
    if len(by_id) != len(rows) or set(by_id) != {row[0] for row in stored}:
        raise ValueError("SQLite and JSONL invoice populations differ")
    # Public XML/ledger parsing and arithmetic are mechanically replayable. The
    # temporary ZIP staging cannot alter the authoritative job or model evidence.
    with tempfile.TemporaryDirectory(prefix="vera-passive-replay-") as temporary:
        invoices = producer.parse_invoice_population(
            plan["invoice_source"], Path(temporary).resolve()
        )
        ledger = producer.load_ledger(
            plan["paths"]["ledger"],
            plan["paths"]["ledger_mapping"],
            sheet=plan["recipe"]["controls"]["ledger_sheet"],
        )
        items, orphans = producer.match_population(
            invoices, ledger, config.amount_tolerance
        )
    replayed = {row["invoice"]["invoice_id"]: row for row in items}
    if set(replayed) != set(by_id):
        raise ValueError("Invoice JSONL differs from the registered source population")
    matched_packets = {
        identity: producer.build_packet(row, plan["history"], plan["chart"])
        for identity, row in replayed.items()
        if row["match_state"] == "matched"
    }
    results = {}
    chunk_members = set()
    for chunk in chunks:
        directory = output / "luna_chunks" / chunk["chunk_id"]
        packets = json.loads(ordinary(directory / "audit_packets.json"))
        identities = json.loads(chunk["invoice_ids_json"])
        if (
            not isinstance(packets, list)
            or len(packets) > 50
            or [p["invoice_id"] for p in packets] != identities
            or producer._sha256_json(packets) != chunk["packet_sha256"]
            or chunk["chunk_id"] != "chunk-" + chunk["packet_sha256"][:16]
        ):
            raise ValueError("Chunk packet population differs from SQLite")
        if (
            any(
                packet != matched_packets.get(packet["invoice_id"])
                for packet in packets
            )
            or len(set(identities)) != len(identities)
            or chunk_members.intersection(identities)
        ):
            raise ValueError(
                "Chunk evidence differs from actual invoice/ledger packets"
            )
        chunk_members.update(identities)
        if ordinary(directory / "luna_prompt.md") != producer.build_luna_prompt(
            packets
        ).encode() or object_file(
            directory / "luna_output_schema.json"
        ) != producer.luna_output_schema(
            identities
        ):
            raise ValueError("Chunk prompt or schema changed")
        if chunk["status"] == "completed":
            recovered = producer._recover_checkpoint(
                directory / producer.CHUNK_RESULT_NAME,
                chunk_id=chunk["chunk_id"],
                packet_sha256=chunk["packet_sha256"],
                invoice_ids=identities,
                reasoning_effort=config.reasoning_effort,
                worker_runtime=config.worker_runtime,
                worker_model=config.worker_model,
                worker_selection=config.worker_selection,
            )
            for result in recovered["results"].values():
                if result["invoice_id"] in results:
                    raise ValueError("Invoice semantic result reused across chunks")
                results[result["invoice_id"]] = (chunk["chunk_id"], result)
    if chunk_members != set(matched_packets):
        raise ValueError("Semantic chunks do not cover the exact matched population")
    for (
        identity,
        item_json,
        packet_json,
        semantic_json,
        model,
        effort,
        chunk_id,
        final_state,
        packet_sha256,
    ) in stored:
        source = replayed[identity]
        item, row = json.loads(item_json), by_id[identity]
        if item != source or any(
            row.get(key) != value for key, value in source.items()
        ):
            raise ValueError("Stored invoice or matching evidence changed")
        packet = (
            producer.build_packet(source, plan["history"], plan["chart"])
            if source["match_state"] == "matched"
            else None
        )
        semantic = json.loads(semantic_json) if semantic_json else None
        if (
            (json.loads(packet_json) if packet_json else None) != packet
            or row["semantic_packet"] != packet
            or packet_sha256 != (producer._sha256_json(packet) if packet else None)
        ):
            raise ValueError(
                "Stored semantic packet differs from actual booked evidence"
            )
        expected = results.get(identity)
        if semantic is not None and (
            expected != (chunk_id, semantic)
            or model != config.worker_model
            or effort
            != (
                "host_default"
                if config.worker_runtime == "cowork"
                else config.reasoning_effort
            )
        ):
            raise ValueError(
                "Stored semantic result differs from its checkpoint or worker"
            )
        if semantic is None and (
            expected is not None
            or model is not None
            or effort is not None
            or chunk_id is not None
        ):
            raise ValueError("Incomplete semantic row has inconsistent worker evidence")
        actual_state, reasons = producer._finalize_item(source, semantic)
        if (
            row["semantic_result"],
            row["semantic_model"],
            row["semantic_reasoning_effort"],
            row["semantic_chunk_id"],
            row["final_state"],
            row["final_exception_reasons"],
            row["client_run_id"],
            row["workflow_version"],
            row["schema_version"],
        ) != (
            semantic,
            model,
            effort,
            chunk_id,
            actual_state,
            reasons,
            context["run_id"],
            producer.WORKFLOW_VERSION,
            producer.SCHEMA_VERSION,
        ) or final_state != actual_state:
            raise ValueError("Final invoice screening state changed")
        if (
            semantic
            and row["professional_should_inspect"]
            != semantic["professional_should_inspect"]
        ):
            raise ValueError(
                "Professional inspection wording differs from worker evidence"
            )
    if jsonl(output / "ledger_entries_without_invoice.jsonl") != orphans:
        raise ValueError("Ledger-orphan population differs from source evidence")
    expected_status = (
        "awaiting_semantic_review"
        if any(row["status"] == "awaiting_semantic_review" for row in chunks)
        else (
            "failed"
            if any(row["status"] == "failed" for row in chunks)
            else "completed"
        )
    )
    expected_counts = {
        "population": len(rows),
        "invoices_requiring_professional_attention": sum(
            row["final_state"] == "professional_review_required" for row in rows
        ),
        "ledger_entry_without_invoice": len(orphans),
        "luna_chunks_total": len(chunks),
        "luna_chunks_completed": sum(row["status"] == "completed" for row in chunks),
        "luna_chunks_failed": sum(row["status"] == "failed" for row in chunks),
        **{
            key: sum(row["match_state"] == key for row in rows)
            for key in (
                "matched",
                "ambiguous_match",
                "invoice_not_found_in_ledger",
                "duplicate_candidate",
            )
        },
        **{
            "luna_"
            + key: sum(
                (row["semantic_result"] or {}).get("status") == key for row in rows
            )
            for key in ("no_issue_detected", "review_required", "insufficient_evidence")
        },
        "luna_not_run_or_failed": sum(row["semantic_result"] is None for row in rows),
    }
    if (
        summary["status"] != expected_status
        or summary["semantic_worker_requested"] != config.worker_model
        or summary["semantic_runtime"]
        != ("cowork_subagent" if config.worker_runtime == "cowork" else "codex_native")
        or any(summary[key] != value for key, value in expected_counts.items())
    ):
        raise ValueError("Invoice summary counters or worker state changed")
    header = "# Vera — Intelligent Passive-Invoice Audit\n\n"
    footer = "\n\n`no_issue_detected` is a screening result, not proof that accounting treatment is correct.\n"
    markdown = ordinary(output / "run_summary.md").decode()
    if not markdown.startswith(header) or not markdown.endswith(footer):
        raise ValueError("Published summary text structure changed")
    lines = markdown[len(header) : -len(footer)].splitlines()
    # The producer saves sorted JSON but renders insertion order. Compare the
    # complete exact line population, including the sole dictionary value.
    expected_lines = [
        f"- {key}: {value}"
        for key, value in summary.items()
        if key not in {"limitations", "luna_usage", "luna_recovery_sources"}
    ]
    recovery = [line for line in lines if line.startswith("- luna_recovery_sources: ")]
    if (
        len(recovery) != 1
        or ast.literal_eval(recovery[0].partition(": ")[2])
        != summary["luna_recovery_sources"]
        or sorted(line for line in lines if line not in recovery)
        != sorted(expected_lines)
    ):
        raise ValueError("Published summary text differs from its JSON evidence")
    # Compare public workbook values, not ZIP timestamps or semantic judgments.
    from openpyxl import load_workbook

    def workbook_values(path: Path) -> list:
        ordinary(path, 50 * 1024 * 1024)
        with ZipFile(path) as archive:
            members = archive.infolist()
            if (
                len(members) > 1000
                or sum(row.file_size for row in members) > 200 * 1024 * 1024
            ):
                raise ValueError("Exception workbook exceeds expanded inspection limit")
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            if any(sheet.max_row * sheet.max_column > 2_000_000 for sheet in workbook):
                raise ValueError("Exception workbook exceeds native inspection limit")
            sheets = []
            for sheet in workbook:
                values = list(sheet.values)
                if sheet.title == "Summary":
                    # Complete public metric rows survive sorted JSON persistence;
                    # only their order and Python container rendering can change.
                    metrics = []
                    for row in values[3:]:
                        if len(row) != 2 or not isinstance(row[0], str):
                            raise ValueError("Unexpected summary workbook structure")
                        value = row[1]
                        if row[0] in {"luna_usage", "luna_recovery_sources"}:
                            value = repr(
                                json.loads(
                                    json.dumps(ast.literal_eval(value), sort_keys=True)
                                )
                            )
                        metrics.append((row[0], value))
                    values = [*values[:3], *sorted(metrics)]
                sheets.append((sheet.title, values))
            return sheets
        finally:
            workbook.close()

    with tempfile.TemporaryDirectory(prefix="vera-passive-workpaper-") as temporary:
        replay = Path(temporary).resolve() / "workpaper.xlsx"
        producer._write_exception_workpaper(
            replay,
            summary,
            [
                row
                for row in rows
                if row["final_state"] == "professional_review_required"
            ],
            orphans,
        )
        if workbook_values(output / "exception_workpaper.xlsx") != workbook_values(
            replay
        ):
            raise ValueError("Exception workbook differs from the public producer")
    verify_sources(plan)
    if implementation(producer) != source_code:
        raise ValueError("Passive-invoice implementation changed during inspection")
    if job_population(output) != population:
        raise ValueError("Passive-invoice job changed during inspection")
    entries = {
        "exceptions": [
            ("invoice:" + row["invoice"]["invoice_id"], row)
            for row in rows
            if row["final_state"] == "professional_review_required"
        ],
        "population": [
            ("invoice:" + row["invoice"]["invoice_id"], row) for row in rows
        ],
        "chunks": [("chunk:" + row["chunk_id"], row) for row in chunks],
        "orphans": [
            ("orphan:" + producer._sha256_json(row)[:32], row) for row in orphans
        ],
    }[view]
    selected = next(
        (row for identity, row in entries if identity == args.get("item_id")), None
    )
    if args.get("item_id") and selected is None:
        raise ValueError(
            "Selected invoice evidence belongs to another view or population"
        )
    if args.get("source_ref") and selected is None:
        raise ValueError("Choose an exact item before navigating its JSON evidence")
    selection = (
        json_page(selected, args.get("source_ref")) if selected is not None else None
    )
    result = {
        "revision": checkpoint,
        "status": expected_status,
        "summary": summary,
        "items": [
            {
                "id": identity,
                "group": view,
                "title": row.get("invoice", {}).get("invoice_number", identity),
            }
            for identity, row in entries[offset : offset + 30]
        ],
        "total": len(entries),
        "selection": selection,
        "artifacts": sorted(population),
        "provider_attestation": False,
        "professional_approval": False,
        "archive_completed": False,
    }
    if len(json.dumps(result).encode()) > 2_000_000:
        raise ValueError("Passive-invoice projection exceeds native limit")
    return result
