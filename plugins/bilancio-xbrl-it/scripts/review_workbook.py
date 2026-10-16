"""Export a revision-bound review document, without importing spreadsheet edits."""

from __future__ import annotations

import hashlib
import json
from datetime import datetime
from io import BytesIO
from typing import Any, Mapping
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from area_review import build_area_review

__all__ = ["build_review_workbook"]


def build_review_workbook(case: Mapping[str, Any]) -> bytes:
    """Use the same living projection as the dashboard; retain exact money text."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    review = build_area_review(case)
    book = Workbook()
    book.remove(book.active)

    def sheet(name: str, headers: list[str], rows: list[list[Any]]) -> None:
        ws = book.create_sheet(name)
        ws.append(headers)
        for row in rows:
            if any(
                len(
                    json.dumps(v, ensure_ascii=False)
                    if isinstance(v, (dict, list))
                    else str(v or "")
                )
                > 32000
                for v in row
            ):
                raise ValueError(
                    "Review workbook cell exceeds 32,000 characters; use the detailed source review"
                )
            ws.append(
                [
                    (
                        json.dumps(v, ensure_ascii=False, sort_keys=True)
                        if isinstance(v, (dict, list))
                        else str(v if v is not None else "")
                    )
                    for v in row
                ]
            )
        for cell in ws[1]:
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="163B65")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                # Untrusted source text must never become a formula or hyperlink.
                if len(str(cell.value or "")) > 32000:
                    raise ValueError(
                        "A review workbook cell exceeds the explicit 32,000 character bound"
                    )
                cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        from openpyxl.utils import get_column_letter

        for column in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(column)].width = (
                28 if column < 3 else 45
            )
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False

    sheet(
        "Riepilogo",
        ["Voce", "Valore"],
        [
            ["Società", case["entity"]["legal_name"]],
            ["Periodo", case["period"]],
            ["Revisione", case["revision_id"]],
            ["Stato revisione", review["status"]],
            ["Percorso", review["profile"]],
            ["Pacchetto OIC", review["oic_rule_pack"]],
            ["Impronta pacchetto OIC", review["oic_rule_pack_checksum"]],
            [
                "Catalogo e impronta",
                [review["catalogue_id"], review["catalogue_sha256"]],
            ],
            [
                "Limiti",
                "Documento di revisione. Saldi esatti EUR conservati come testo. Nessuna approvazione, importazione delle modifiche Excel o deposito. Copertura OIC parziale; giudizi e fonti restano da verificare.",
            ],
        ],
    )
    sheet(
        "Aree",
        ["Area", "Titolo", "Stato", "Motivazione", "Impronta dipendenze"],
        [
            [a["area_id"], a["title"], a["status"], a["reason"], a["basis_sha256"]]
            for a in review["workpapers"]
        ],
    )
    from closing_workflow import closing_status

    closing = closing_status(case) if case.get("closing_workflow") else {}
    workflow = case.get("closing_workflow") or {}
    sheet(
        "Fascicolo",
        [
            "Requisito",
            "Descrizione",
            "Stato",
            "Motivazione",
            "Fonti",
            "Richiesta in bozza",
        ],
        [
            [
                r["requirement_id"],
                r["description"],
                r["status"],
                r["reason"],
                r.get("evidence_refs", []),
                r.get("request_from_client", False),
            ]
            for r in closing.get("requirements", [])
        ],
    )
    sheet(
        "Screening",
        ["Versione", "Documento", "Versione selezionata", "Calibrazione", "Raccordo"],
        [
            [
                v["version_id"],
                v["document_id"],
                v["version_id"] == closing.get("selected_version_id"),
                v["trial_balance"]["calibration"],
                v["reconciliation"],
            ]
            for v in workflow.get("versions", [])
        ],
    )
    sheet(
        "Rettifiche",
        [
            "Proposta",
            "Oggetto",
            "Conto",
            "Dare EUR esatto",
            "Avere EUR esatto",
            "Motivazione",
            "Fonti",
            "Decisioni",
        ],
        [
            [
                p["proposal_id"],
                p["subject_key"],
                line["account_code"],
                line["debit"],
                line["credit"],
                p["reason"],
                p["evidence"],
                [
                    {
                        "record_id": e["record_id"],
                        "outcome": e["decision"]["outcome"],
                        "current_state": closing["decision_states"][e["record_id"]],
                    }
                    for e in workflow.get("events", [])
                    if e["kind"] == "ADJUSTMENT"
                    and e["decision"]["proposal_id"] == p["proposal_id"]
                ],
            ]
            for p in workflow.get("proposals", [])
            for line in p["lines"]
        ],
    )
    sheet(
        "Conti",
        [
            "Area",
            "Conto",
            "Descrizione",
            "Saldo EUR esatto",
            "Saldo precedente EUR esatto",
            "Fonti",
        ],
        [
            [
                a["area_id"],
                e["account_code"],
                e["account_description"],
                e["closing_signed"],
                e.get("prior_closing_signed"),
                e.get("source_refs", []),
            ]
            for a in review["workpapers"]
            for e in a["accounts"]
        ],
    )
    reconciliation_index = {
        r["fact_id"]: r for a in review["workpapers"] for r in a["reconciliations"]
    }
    sheet(
        "Raccordi",
        ["Area", "Conto", "Mapping revisionato", "Fatti e prospetti"],
        [
            [
                a["area_id"],
                link["account_id"],
                link,
                [reconciliation_index[ref] for ref in link["fact_ids"]],
            ]
            for a in review["workpapers"]
            for link in a["links"]
        ],
    )
    sheet(
        "Catalogo OIC",
        ["Principio", "Titolo", "Applicabilità", "Copertura", "Versione fonte", "URL"],
        [
            [
                s["id"],
                s["title"],
                s["status"],
                s["coverage_level"],
                s["source_version"],
                s["source_url"],
            ]
            for s in review["standards"]
        ],
    )
    sheet(
        "Controlli",
        ["Area", "Controllo", "Obiettivo", "Stato", "Decisione"],
        [
            [c["area_id"], c["id"], c["objective"], c["status"], c["review"]]
            for c in review["controls"]
        ],
    )
    sheet(
        "Documenti qualitativi",
        ["Proposta", "Documento", "Periodo", "Brano", "Aree", "Stato"],
        [
            [
                q["id"],
                q["document_id"],
                q["period"],
                q["quote"],
                q["area_ids"],
                q["status"],
            ]
            for q in review["qualitative"]
        ],
    )
    sheet(
        "Prospetti richiesti",
        ["Area", "Richiesti", "Mancanti", "Prospetti acquisiti"],
        [
            [
                a["area_id"],
                a["required_schedule_types"],
                a["missing_schedule_types"],
                [
                    {
                        "schedule_id": s.get("schedule_id"),
                        "schedule_type": s["schedule_type"],
                        "issues": s.get("issues", []),
                    }
                    for s in a["schedules"]
                ],
            ]
            for a in review["workpapers"]
            if a["status"] != "EXCLUDED"
        ],
    )
    sheet(
        "Questioni aperte",
        ["Area", "Domande", "Questioni aperte"],
        [
            [a["area_id"], a["questions"], a["open_issues"]]
            for a in review["workpapers"]
            if a["questions"] or a["open_issues"]
        ],
    )
    sheet(
        "Nota integrativa",
        ["Area", "Sezioni", "Testi registrati"],
        [
            [a["area_id"], a["note_sections"], a["notes"]]
            for a in review["workpapers"]
            if a["status"] != "EXCLUDED"
        ],
    )
    sheet(
        "Decisioni",
        ["Tipo", "Autore", "Revisione", "Decisione", "Impronta"],
        [
            [
                e["kind"],
                e["actor"],
                e["revision_id"],
                {
                    "reason": e.get("reason", e.get("payload", {}).get("reason")),
                    "evidence": e.get(
                        "evidence", e.get("payload", {}).get("evidence_refs", [])
                    ),
                    "status": e.get("payload", {}).get("status"),
                    "proposal_id": e.get("payload", {}).get("proposal_id"),
                },
                e["sha256"],
            ]
            for e in [
                *(case.get("closing_workflow") or {}).get("events", []),
                *review["events"],
            ]
        ],
    )
    sheet(
        "Fonti",
        ["Documento", "Nome", "SHA-256", "Scopo"],
        [
            [d["document_id"], d["file_name"], d["sha256"], d["purpose"]]
            for d in case.get("source_documents", [])
        ],
    )
    book.properties.title = "Bilancio OIC · Documento di revisione"
    book.properties.description = (
        "Revisione "
        + case["revision_id"]
        + "; contenuto "
        + hashlib.sha256(json.dumps(case, sort_keys=True).encode()).hexdigest()
    )
    buffer = BytesIO()
    frozen = datetime.fromisoformat(case["updated_at"].replace("Z", "+00:00")).replace(
        tzinfo=None
    )
    book.properties.created = frozen
    book.properties.modified = frozen
    book.save(buffer)
    canonical = BytesIO()
    from lxml import etree

    with (
        ZipFile(buffer) as source,
        ZipFile(canonical, "w", compression=ZIP_DEFLATED) as destination,
    ):
        for name in sorted(source.namelist()):
            content = source.read(name)
            if name == "docProps/core.xml":
                # This XML is generated locally, never parsed from user input.
                root = etree.fromstring(
                    content,
                    parser=etree.XMLParser(resolve_entities=False, no_network=True),
                )
                for child in root:
                    if child.tag.endswith("}modified"):
                        child.text = frozen.isoformat() + "Z"
                content = etree.tostring(root)
            item = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            item.compress_type = ZIP_DEFLATED
            destination.writestr(item, content)
    return canonical.getvalue()
