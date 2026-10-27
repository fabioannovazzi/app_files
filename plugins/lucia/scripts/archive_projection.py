"""Producer-owned source coordinates for bounded packets, before any delivery."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from archive_emit import bound_run
from archive_report_store import RunStore, digest, file_digest, run_lock
from archive_reporting_failure import isolated_projection

__all__ = [
    "is_recording",
    "model_argument_origin",
    "register_model_arguments",
    "register_metadata",
    "register_metadata_revision",
    "register_artifact_projection",
    "register_private_value",
    "register_projection",
    "register_rows",
    "register_row_projection",
    "register_table_projection",
    "register_value",
    "run_cli",
]


def run_cli(main: Any, component_file: Path) -> int | None:
    """Delegate direct component CLI delivery to Vera's owned recorder."""
    from archive_cli import run_entrypoint

    return run_entrypoint(main, component_file)


@isolated_projection
def is_recording(component: str, output: Path) -> bool:
    """Identify an explicitly bound selected run through Archive's stored state."""
    return bound_run(component, {"output": str(output)}) is not None


def _positions(
    source: dict[str, Any],
    entry: dict[str, Any],
    workbooks: dict[str, dict[str, dict[int, int]]],
) -> dict[int, int] | None:
    """Convert producer physical rows to the same units used at Archive start."""
    if source["unit"] == "rows" and Path(source["path"]).suffix.lower() in {
        ".csv",
        ".tsv",
    }:
        return {
            row["physical_row"]: row["physical_row"] - 2
            for row in entry["rows"]
            if row["physical_row"] > 1
        }
    if source["unit"] != "worksheet_rows":
        return None
    from openpyxl import load_workbook

    if source["sha256"] not in workbooks:
        sheets = {}
        index = 0
        workbook = load_workbook(source["path"], read_only=True, data_only=False)
        try:
            for sheet in workbook:
                positions = {}
                for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                    if any(value is not None for value in values):
                        positions[number] = index
                        index += 1
                sheets[sheet.title] = positions
        finally:
            workbook.close()
        workbooks[source["sha256"]] = sheets
    return workbooks[source["sha256"]].get(entry["sheet_name"], {})


@isolated_projection
def register_metadata(component: str, output: Path, payload: Any) -> None:
    """Declare an exact producer status message, never a semantic classification."""
    root = bound_run(component, {"output": str(output)})
    if root is None:
        return
    with run_lock(root):
        RunStore(root).event(
            "metadata:" + digest(payload),
            {
                "kind": "prepared_metadata",
                "data_sha256": digest(payload),
            },
        )


@isolated_projection
def register_model_arguments(
    component: str, output: Path, arguments: list[str]
) -> None:
    """Remember exact text received from the model, without recording a send.

    The native tool calls this before execution. Hashes are scoped to this run
    and component; reading a file or running a plain CLI cannot create this
    evidence. Returning this same text adds no newly exposed input-file cells.
    """
    root = bound_run(component, {"output": str(output)})
    if root is None:
        return
    hashes = sorted({digest(value) for value in arguments})
    with run_lock(root):
        store = RunStore(root)
        state = store.state()
        if state is None or state["collection_closed"]:
            return
        store.event(
            "model_arguments:" + digest([component, hashes]),
            {
                "kind": "model_supplied_arguments",
                "component": component,
                "value_sha256": hashes,
            },
        )


@isolated_projection
def model_argument_origin(
    component: str, output: Path, value: str
) -> list[dict[str, Any]] | None:
    """Resolve exact model-supplied text; unknown or changed text stays unknown."""
    root = bound_run(component, {"output": str(output)})
    if root is not None:
        with run_lock(root):
            if any(
                event["kind"] == "model_supplied_arguments"
                and event["component"] == component
                and digest(value) in event["value_sha256"]
                for event in RunStore(root).events()
            ):
                return []
    return None


@isolated_projection
def register_rows(
    component: str, output: Path, records: dict[str, dict[str, Any]]
) -> None:
    """Retain the normalizer's field origins without counting preparation as a send."""
    root = bound_run(component, {"output": str(output)})
    if root is None:
        return
    with run_lock(root):
        RunStore(root).event(
            "rows:" + digest(records),
            {
                "kind": "prepared_source_rows",
                "component": component,
                "records": records,
            },
        )


@isolated_projection
def register_private_value(component: str, output: Path, value: Any) -> None:
    """Bind a locally produced review packet; this does not declare a model send."""
    root = bound_run(component, {"output": str(output)})
    if root is not None:
        with run_lock(root):
            RunStore(root).event(
                "private:" + digest(value),
                {
                    "kind": "prepared_private_value",
                    "component": component,
                    "data_sha256": digest(value),
                },
            )


@isolated_projection
def register_metadata_revision(
    component: str,
    path: Path,
    previous_sha256: str,
    *,
    expected_sha256: str | None = None,
) -> None:
    """Preserve known source operands through a producer's metadata-only edit."""
    root = bound_run(component, {"output": str(path)})
    if root is None or path.is_symlink():
        return
    with run_lock(root):
        store = RunStore(root)
        parent = next(
            (
                event
                for event in store.events()
                if event["kind"] == "prepared_projection"
                and event["artifact_sha256"] == previous_sha256
            ),
            None,
        )
        if parent is None:
            return
        sources = store.sources()
        referenced_paths = set()
        for selection in parent["selections"]:
            source = next(
                (
                    item
                    for item in sources
                    if item["role"] == "input"
                    and item["source_id"] == selection["source_id"]
                    and item["sha256"] == selection["source_sha256"]
                ),
                None,
            )
            if source is None:
                return
            referenced_paths.add(Path(source["path"]))
        content = path.read_bytes()
        artifact_sha256 = hashlib.sha256(content).hexdigest()
        if expected_sha256 is not None and artifact_sha256 != expected_sha256:
            return
        value = json.loads(content)
        _verify_source_snapshots(sources, referenced_paths)
        if file_digest(path) != artifact_sha256:
            raise ValueError("Producer metadata edit changed during registration")
        store.event(
            "projection:" + artifact_sha256,
            {
                "kind": "prepared_projection",
                "artifact_sha256": artifact_sha256,
                "data_sha256": digest(value),
                "selections": parent["selections"],
            },
        )


@isolated_projection
def register_row_projection(
    component: str,
    output: Path,
    value: Any,
    records: list[dict[str, Any]],
    *,
    private_sha256: str | None = None,
    artifact_sha256: str | None = None,
) -> None:
    """Resolve only fields selected by the producer, never infer origin from values."""
    root = bound_run(component, {"output": str(output)})
    if root is None:
        return
    with run_lock(root):
        events = RunStore(root).events()
        if private_sha256 is not None and not any(
            event["kind"] == "prepared_private_value"
            and event["component"] == component
            and event["data_sha256"] == private_sha256
            for event in events
        ):
            return  # An edited or supplied review cannot claim producer provenance.
        entries = _row_entries(_row_origins(events, component), records)
    if entries is not None:
        register_value(
            component, output, value, entries, artifact_sha256=artifact_sha256
        )


@isolated_projection
def register_artifact_projection(
    component: str,
    path: Path,
    *,
    data_sha256: str,
    artifact_sha256: str,
    records: list[dict[str, Any]],
    private_sha256: str,
    previous_sha256: str | None = None,
) -> None:
    """Bind a writer's exact JSON file to its producer-owned review origins."""
    root = bound_run(component, {"output": str(path)})
    if (
        root is None
        or path.is_symlink()
        or not path.resolve().is_relative_to(root / "outputs")
        or file_digest(path) != artifact_sha256
    ):
        return
    value = json.loads(path.read_bytes())
    if digest(value) != data_sha256 or file_digest(path) != artifact_sha256:
        return
    if previous_sha256 is not None:
        with run_lock(root):
            if not any(
                event["kind"] == "prepared_private_value"
                and event["component"] == component
                and event["data_sha256"] == private_sha256
                for event in RunStore(root).events()
            ):
                return
        register_metadata_revision(
            component, path, previous_sha256, expected_sha256=artifact_sha256
        )
        return
    register_row_projection(
        component,
        path,
        value,
        records,
        private_sha256=private_sha256,
        artifact_sha256=artifact_sha256,
    )


def _row_origins(events: list[dict[str, Any]], component: str) -> dict[str, Any]:
    origins = {}
    for event in events:
        if event["kind"] == "prepared_source_rows" and event["component"] == component:
            origins.update(event["records"])
    return origins


def _row_entries(
    origins: dict[str, Any], records: list[dict[str, Any]]
) -> list[dict[str, Any]] | None:
    entries = []
    for record in records:
        if (
            not isinstance(record, dict)
            or not isinstance(record.get("record_id"), str)
            or not isinstance(record.get("fields"), list)
            or any(not isinstance(field, str) for field in record["fields"])
        ):
            return None
        origin = origins.get(record["record_id"])
        if origin is None or any(
            field not in origin["fields"] for field in record["fields"]
        ):
            return None  # Missing producer evidence remains an incomplete report.
        columns = sorted(
            {column for field in record["fields"] for column in origin["fields"][field]}
        )
        entries.append(
            {
                **origin["source"],
                "rows": [{"physical_row": origin["physical_row"], "columns": columns}],
            }
        )
    return entries


@isolated_projection
def register_table_projection(
    component: str, path: Path, cells: list[list[list[dict[str, Any]] | None]]
) -> None:
    """Bind each produced table cell to its declared record fields and exact bytes.

    Empty references mean a producer-created constant or locator. None or missing
    record evidence leaves that cell unknown; no column-name or value matching
    is used. An unread unknown cell cannot invalidate another cell's provenance.
    """
    from archive_report_measurements import measure_source

    root = bound_run(component, {"output": str(path)})
    if root is None:
        return
    artifact = measure_source(path, source_id="producer_table", role="output")
    if artifact["measurement_error"] or artifact["unit"] not in {
        "rows",
        "worksheet_rows",
        "body_paragraphs",
    }:
        return
    if len(cells) != artifact["quantity"] or any(
        len(row) != artifact["columns"] for row in cells
    ):
        raise ValueError("Producer table provenance must cover its exact shape")
    with run_lock(root):
        store = RunStore(root)
        state = store.state()
        if state is None or state["collection_closed"]:
            raise ValueError("Producer projection requires an active Archive run")
        origins = _row_origins(store.events(), component)
        sources = store.sources()
        checked_sources: set[Path] = set()
        workbooks: dict[str, dict[str, dict[int, int]]] = {}
        resolved = []
        for row in cells:
            resolved_row: list[list[dict[str, Any]] | None] = []
            for records in row:
                entries = None if records is None else _row_entries(origins, records)
                if entries is None:
                    resolved_row.append(None)
                    continue
                selections = _source_selections(
                    sources, entries, checked_sources, workbooks
                )
                resolved_row.append(selections)
            resolved.append(resolved_row)
        if file_digest(path) != artifact["sha256"]:
            raise ValueError("Producer table changed while it was registered")
        _verify_source_snapshots(sources, checked_sources)
        store.event(
            "table_projection:" + digest([artifact["path"], artifact["sha256"]]),
            {
                "kind": "prepared_table_projection",
                "artifact_path": artifact["path"],
                "artifact_sha256": artifact["sha256"],
                "unit": artifact["unit"],
                "columns": artifact["columns"],
                "cells": resolved,
            },
        )


@isolated_projection
def register_projection(
    component: str, path: Path, entries: list[dict[str, Any]]
) -> None:
    """Bind a generated JSON packet to exact cells in the started run's inputs.

    Only pipeline code calls this while constructing its bounded output. This is
    preparation evidence; no send is counted until Vera returns the packet.
    """
    if not is_recording(component, path):
        return
    if path.is_symlink():
        raise ValueError("A producer projection must be an ordinary file")
    packet_hash = file_digest(path)
    value = json.loads(path.read_text(encoding="utf-8"))
    if file_digest(path) != packet_hash:
        raise ValueError("Producer projection changed while it was registered")
    register_value(component, path, value, entries, artifact_sha256=packet_hash)


@isolated_projection
def register_value(
    component: str,
    output: Path,
    value: Any,
    entries: list[dict[str, Any]],
    *,
    artifact_sha256: str | None = None,
) -> None:
    """Retain producer provenance for this exact value, without counting a send."""
    root = bound_run(component, {"output": str(output)})
    if root is None:
        return
    with run_lock(root):
        store = RunStore(root)
        state = store.state()
        if state is None or state["collection_closed"]:
            raise ValueError("Producer projection requires an active Archive run")
        sources = store.sources()
        checked_sources: set[Path] = set()
        selections = _source_selections(sources, entries, checked_sources, {})
        if selections is None:
            return
        _verify_source_snapshots(sources, checked_sources)
        store.event(
            "projection:" + (artifact_sha256 or digest(value)),
            {
                "kind": "prepared_projection",
                "artifact_sha256": artifact_sha256,
                "data_sha256": digest(value),
                "selections": selections,
            },
        )


def _source_selections(
    sources: list[dict[str, Any]],
    entries: list[dict[str, Any]],
    checked_sources: set[Path],
    workbooks: dict[str, dict[str, dict[int, int]]],
) -> list[dict[str, Any]] | None:
    """Resolve producer coordinates against the run's unchanged input snapshots."""
    selections = []
    for entry in entries:
        source = next(
            (
                item
                for item in sources
                if item["role"] == "input"
                and item["sha256"] == entry["source_sha256"]
                and str(Path(entry["source_path"]).resolve())
                in item.get("aliases", [item["path"]])
            ),
            None,
        )
        if source is None or source["measurement_error"]:
            return None  # Unsupported source provenance must remain explicit.
        if entry.get("source_column_count", source["columns"]) != source["columns"]:
            return None  # The source reader and population measure must agree.
        source_path = Path(source["path"])
        if source_path not in checked_sources and (
            source_path.is_symlink() or file_digest(source_path) != source["sha256"]
        ):
            raise ValueError("Producer projection source changed after Archive start")
        positions = _positions(source, entry, workbooks)
        if (
            source_path not in checked_sources
            and file_digest(source_path) != source["sha256"]
        ):
            raise ValueError("Producer projection source changed during measurement")
        checked_sources.add(source_path)
        if positions is None:
            return None
        for row in entry["rows"]:
            index, columns = positions.get(row["physical_row"]), row["columns"]
            if index is None:
                continue  # CSV column names are metadata, not data rows.
            if (
                type(index) is not int
                or not 0 <= index < source["quantity"]
                or any(
                    type(column) is not int or not 0 <= column < source["columns"]
                    for column in columns
                )
            ):
                raise ValueError("Producer projection lies outside its source")
            if columns:
                selections.append(
                    {
                        "source_id": source["source_id"],
                        "source_sha256": source["sha256"],
                        "unit": source["unit"],
                        "start": index,
                        "stop": index + 1,
                        "columns": columns,
                        "basis": "vera_producer_projection",
                    }
                )
    return selections


def _verify_source_snapshots(
    sources: list[dict[str, Any]], checked_sources: set[Path]
) -> None:
    """Detect input changes across a registration without hashing per output cell."""
    for source in sources:
        path = Path(source["path"])
        if path in checked_sources and (
            path.is_symlink() or file_digest(path) != source["sha256"]
        ):
            raise ValueError("Producer projection source changed during registration")
