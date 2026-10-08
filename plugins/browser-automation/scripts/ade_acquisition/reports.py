"""Local, reviewable acquisition outputs; missing amounts are never filled with zero."""

from __future__ import annotations

import csv
import html
import json
import os
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from .contracts import AcquisitionError, Plan, money
from .engine import load_events
from .storage import atomic_json

__all__ = ["write_reports", "cash_records", "prepare_f24"]
AMOUNTS = {"amount": "ammontare", "taxable": "imponibile", "vat": "imposta"}


def _text_cell(value: Any) -> Any:
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _save_bytes(path: Path, value: bytes) -> None:
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(descriptor, "wb") as stream:
        stream.write(value)
        stream.flush()
        os.fsync(stream.fileno())


def cash_records(
    events: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Normalize fields by observed column names, preserving failed rows for review."""
    records, issues = [], []
    ids: set[tuple[str, str]] = set()
    for event in events:
        if event["kind"] != "cash":
            continue
        for index, raw in enumerate(event["rows"], 1):
            row: dict[str, Any] = {
                "client": event["client"],
                "device_type": event["device"],
                "source_year": event["year"],
                "source_quarter": event["quarter"],
                "raw": raw,
                "issues": [],
            }

            def field(prefix: str) -> str | None:
                found = [
                    value
                    for key, value in raw.items()
                    if key.lower().startswith(prefix)
                ]
                return found[0] if len(found) == 1 else None

            row["id"] = field("id invio") or ""
            row["device"] = field("matricola") or ""
            raw_date = field("data e ora rilevazione")
            row["detected_at"] = raw_date or ""
            row["sent_at"] = field("data e ora trasmissione") or ""
            row["day"] = ""
            try:
                detected = datetime.strptime(raw_date or "", "%d/%m/%Y %H:%M:%S")
                row["day"] = detected.date().isoformat()
                if detected.hour < 6:
                    row["issues"].append("early-hour-check-business-date")
            except ValueError:
                row["issues"].append("detection-date-unrecognized")
            for key, prefix in AMOUNTS.items():
                try:
                    parsed = money(field(prefix))
                    row[key] = str(parsed) if parsed is not None else None
                    if parsed is None:
                        row["issues"].append(f"{key}-missing")
                except AcquisitionError:
                    row[key] = None
                    row["issues"].append(f"{key}-invalid")
            if not row["id"]:
                row["issues"].append("submission-id-missing")
            elif (event["client"], row["id"]) in ids:
                row["issues"].append("submission-id-repeated")
            ids.add((event["client"], row["id"]))
            if not row["device"]:
                row["issues"].append("device-id-missing")
            if any(
                value.strip()
                for key, value in raw.items()
                if key.lower().startswith("periodo di inattivit")
            ):
                row["issues"].append("inactivity-period-check")
            if row["issues"]:
                issues.append(
                    {
                        "client": event["client"],
                        "device": event["device"],
                        "row": index,
                        "issues": row["issues"],
                    }
                )
            records.append(row)
    return records, issues


def _aggregate(records: list[dict[str, Any]], monthly: bool) -> list[list[Any]]:
    groups: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        period = record["day"][:7] if monthly else record["day"]
        if period:
            groups[
                (record["client"], period, "" if monthly else record["device"])
            ].append(record)
    result = []
    for (client, period, device), group in sorted(groups.items()):
        blocked = any(
            "submission-id-repeated" in row["issues"]
            or "submission-id-missing" in row["issues"]
            for row in group
        )
        totals = [
            (
                None
                if blocked or any(row[key] is None for row in group)
                else sum((Decimal(row[key]) for row in group), Decimal("0.00"))
            )
            for key in AMOUNTS
        ]
        notes = sorted({note for row in group for note in row["issues"]})
        result.append([client, period, device, len(group), *totals, "; ".join(notes)])
    return result


def write_reports(run: Path, plan: Plan, state: dict[str, Any]) -> dict[str, Any]:
    """Write JSON/XLSX/CSV/HTML with source counts, incomplete scopes and local file links."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    state = {key: value for key, value in state.items() if key != "reports_ready"}
    events = load_events(run)
    cash, issues = cash_records(events)
    clients = {client.vat_number: client for client in plan.clients}
    invoice_events = [event for event in events if event["kind"] == "invoice"]
    stamp_events = [event for event in events if event["kind"] == "stamp"]
    report = {
        "schema_version": "vera-agenzia-report/v1",
        "status": state,
        "cash_review_issues": issues,
        "cash_records": cash,
        "invoice_results": invoice_events,
        "stamp_duty": stamp_events,
        "f24": "Review amounts, payment status and due dates before preparing a draft; nothing is submitted.",
    }
    atomic_json(run / "report.json", report)
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(name: str, headers: list[str], rows: list[list[Any]]) -> None:
        ws = workbook.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([_text_cell(value) for value in row])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="153E67")
        for index, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(index)].width = min(
                42, max(16, len(header) + 3)
            )
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    sheet(
        "Esecuzione",
        ["Campo", "Valore"],
        [[key, value] for key, value in state.items() if key != "scopes"],
    )
    sheet(
        "Ambiti",
        [
            "P.IVA",
            "Operazione",
            "Anno",
            "Trimestre",
            "Esito",
            "Righe portale",
            "Errore",
        ],
        [
            [
                scope["client"],
                scope["operation"],
                scope["year"],
                scope["quarter"],
                scope["status"],
                scope.get("portal_count"),
                scope.get("error", ""),
            ]
            for scope in state.get("scopes", [])
        ],
    )
    invoice_rows = [
        [
            event["client"],
            event["operation"],
            event["year"],
            event["quarter"],
            event["status"],
            event.get("error", ""),
            event.get("record", {}).get("sha256", ""),
            "; ".join(
                item["path"] for item in event.get("record", {}).get("artifacts", [])
            ),
        ]
        for event in invoice_events
    ]
    invoice_headers = [
        "P.IVA",
        "Categoria",
        "Anno",
        "Trimestre",
        "Esito",
        "Errore",
        "SHA256 originale",
        "File locali",
    ]
    sheet("Fatture", invoice_headers, invoice_rows)
    raw_rows = [
        [
            row["client"],
            row["id"],
            row["device"],
            row["device_type"],
            row["detected_at"],
            row["sent_at"],
            *(Decimal(row[key]) if row[key] is not None else None for key in AMOUNTS),
            "; ".join(row["issues"]),
            json.dumps(row["raw"], ensure_ascii=False),
        ]
        for row in cash
    ]
    sheet(
        "Invii corrispettivi",
        [
            "P.IVA",
            "Id invio",
            "Matricola",
            "Tipo",
            "Rilevazione",
            "Trasmissione",
            "Ammontare",
            "Imponibile",
            "Imposta",
            "Da verificare",
            "Riga fonte",
        ],
        raw_rows,
    )
    aggregate_headers = [
        "P.IVA",
        "Periodo rilevazione",
        "Matricola",
        "Invii",
        "Ammontare",
        "Imponibile",
        "Imposta",
        "Da verificare",
    ]
    sheet("Corrispettivi giornalieri", aggregate_headers, _aggregate(cash, False))
    sheet("Corrispettivi mensili", aggregate_headers, _aggregate(cash, True))
    stamp_rows = []
    for event in stamp_events:
        evidence = event["evidence"]
        stamp_rows.append(
            [
                event["client"],
                event["year"],
                event["quarter"],
                str(2520 + event["quarter"]),
                evidence["status"],
                Decimal(evidence["amount"]) if evidence["amount"] is not None else None,
                evidence.get("list_a_count"),
                evidence.get("list_b_count"),
                evidence.get("document_count"),
                evidence["payment_status"],
                evidence.get("attestations", ""),
                "Importo portale dell'intero trimestre; verificare il residuo da versare e la scadenza.",
            ]
        )
    sheet(
        "Bolli e prospetto F24",
        [
            "P.IVA",
            "Anno",
            "Trimestre",
            "Codice tributo",
            "Esito",
            "Bollo portale",
            "Elenco A",
            "Elenco B",
            "Documenti",
            "Stato pagamento",
            "Attestazioni",
            "Verifica",
        ],
        stamp_rows,
    )
    buffer = BytesIO()
    workbook.save(buffer)
    _save_bytes(run / "riepilogo.xlsx", buffer.getvalue())
    descriptor = os.open(
        run / "fatture.csv", os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600
    )
    with os.fdopen(descriptor, "w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, delimiter=";")
        writer.writerow(invoice_headers)
        writer.writerows([_text_cell(value) for value in row] for row in invoice_rows)
    rows_html = "".join(
        "<tr>"
        + "".join(
            f"<td>{html.escape(str(value if value is not None else ''))}</td>"
            for value in row
        )
        + "</tr>"
        for row in invoice_rows
    )
    scopes_html = "".join(
        f"<li>{html.escape(clients[scope['client']].name)} · {html.escape(scope['operation'])} · {scope['year']} T{scope['quarter']}: {html.escape(scope['status'])} {html.escape(scope.get('error', ''))}</li>"
        for scope in state.get("scopes", [])
    )
    mode_notice = (
        "<p><strong>ESEMPIO DIDATTICO: sorgente locale fittizia, nessun accesso al portale.</strong></p>"
        if state.get("execution_mode") == "synthetic_local_source_no_portal"
        else ""
    )
    document = f"""<!doctype html><html lang="it"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>Vera · acquisizione Agenzia</title><style>body{{font:16px system-ui;color:#142b44;margin:48px auto;padding:0 24px;max-width:1100px;line-height:1.5}}h1{{font-size:32px}}table{{border-collapse:collapse;display:block;overflow:auto}}td,th{{text-align:left;border-bottom:1px solid #ccd5df;padding:10px;max-width:350px;overflow-wrap:anywhere}}a{{color:#135b9d}}</style><h1>Acquisizione Agenzia delle Entrate</h1>{mode_notice}<p>Esito: <strong>{html.escape(state['state'])}</strong>. Nuovi originali: {state.get('downloaded', 0)}. Già presenti e verificati: {state.get('verified_existing', 0)}. Errori documenti: {state.get('failed_documents', 0)}.</p><p>Ambiti completi: {state.get('scopes_complete', 0)} / {state.get('scopes_expected', '?')}. Segnalazioni corrispettivi: {len(issues)}.</p><p><a href="riepilogo.xlsx">Riepilogo Excel</a> · <a href="fatture.csv">Registro fatture CSV</a> · <a href="report.json">Evidenze JSON</a> · <a href="model_data_report.md">Dati e modello</a></p><ul>{scopes_html}</ul><p>I bolli si riferiscono all'intero trimestre. Il prospetto non determina il residuo dovuto o la scadenza e non dispone pagamenti. Le firme P7M non vengono validate.</p><table><thead><tr>{''.join('<th>'+html.escape(h)+'</th>' for h in invoice_headers)}</tr></thead><tbody>{rows_html}</tbody></table></html>"""
    _save_bytes(run / "report.html", document.encode())
    from .disclosure import write_disclosure

    write_disclosure(
        run, state.get("downloaded", 0) + state.get("verified_existing", 0)
    )
    return {
        "reports_ready": True,
        "cash_review_issue_count": len(issues),
        "files": [
            "report.html",
            "riepilogo.xlsx",
            "fatture.csv",
            "report.json",
            "model_data_report.md",
        ],
    }


def prepare_f24(run: Path, plan: Plan, reviewed: dict[str, Any]) -> Path:
    """Render a non-submittable F24 working draft from explicitly reviewed residual amounts."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    if (
        not isinstance(reviewed, dict)
        or set(reviewed) != {"schema_version", "reviewed", "items"}
        or reviewed["schema_version"] != "vera-agenzia-f24/v1"
        or reviewed["reviewed"] is not True
        or not isinstance(reviewed["items"], list)
        or not reviewed["items"]
    ):
        raise AcquisitionError("f24-explicit-review-required")
    sources = {
        (event["client"], event["year"], event["quarter"]): event["evidence"]
        for event in load_events(run)
        if event["kind"] == "stamp"
    }
    clients = {c.vat_number: c for c in plan.clients}
    validated, seen = [], set()
    for item in reviewed["items"]:
        if not isinstance(item, dict) or set(item) != {
            "client",
            "year",
            "quarter",
            "amount",
            "due_date",
            "note",
        }:
            raise AcquisitionError("f24-item-invalid")
        key = (item["client"], item["year"], item["quarter"])
        if key in seen or key not in sources or sources[key]["status"] != "acquired":
            raise AcquisitionError("f24-source-unavailable-or-duplicate")
        seen.add(key)
        try:
            amount = Decimal(item["amount"])
            due = date.fromisoformat(item["due_date"])
        except (InvalidOperation, ValueError, TypeError) as exc:
            raise AcquisitionError("f24-amount-or-date-invalid") from exc
        if (
            not isinstance(item["amount"], str)
            or not amount.is_finite()
            or amount <= 0
            or amount != amount.quantize(Decimal("0.01"))
            or amount > Decimal("999999999.99")
            or not isinstance(item["note"], str)
            or not item["note"].strip()
        ):
            raise AcquisitionError("f24-review-incomplete")
        validated.append(
            {
                **item,
                "amount": str(amount),
                "due_date": due.isoformat(),
                "source_amount": sources[key]["amount"],
                "payment_status": sources[key]["payment_status"],
            }
        )
    from uuid import uuid4

    basename = f"f24-prospetto-{uuid4().hex}"
    buffer = BytesIO()
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Prospetto F24 — BOZZA DA VERIFICARE", styles["Title"]),
        Paragraph(
            "Non è un modello telematico e non dispone alcun pagamento. Importi residui e scadenze sono quelli approvati dal professionista.",
            styles["Normal"],
        ),
        Spacer(1, 16),
    ]
    for item in validated:
        client = clients[item["client"]]
        story.extend(
            [
                Paragraph(
                    html.escape(
                        f"{client.name} — CF {client.tax_code} — P.IVA {client.vat_number}"
                    ),
                    styles["Heading2"],
                ),
                Paragraph(
                    html.escape(
                        f"Scadenza indicata: {item['due_date']}. {item['note']}"
                    ),
                    styles["Normal"],
                ),
            ]
        )
        table = Table(
            [
                ["Sezione", "Codice", "Anno", "Trimestre", "Debito EUR"],
                [
                    "ERARIO",
                    str(2520 + item["quarter"]),
                    str(item["year"]),
                    str(item["quarter"]),
                    item["amount"],
                ],
            ],
            repeatRows=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#153e67")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
                ]
            )
        )
        story.extend([Spacer(1, 10), table, Spacer(1, 16)])
    SimpleDocTemplate(buffer, pagesize=A4).build(story)
    path = run / f"{basename}.pdf"
    _save_bytes(path, buffer.getvalue())
    atomic_json(
        run / f"{basename}.json",
        {"schema_version": "vera-agenzia-f24-draft/v1", "items": validated},
    )
    return path
