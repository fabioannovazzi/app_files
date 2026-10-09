"""Explicit acquisition inputs and exact format/amount handling.

These checks enforce mechanically verifiable identities and representations;
they never decide accounting treatment or interpret an ambiguous portal state.
"""

from __future__ import annotations

import calendar
import csv
import hashlib
import json
import re
from dataclasses import asdict, dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

__all__ = [
    "AcquisitionError",
    "Client",
    "Plan",
    "CATEGORIES",
    "money",
    "quarters",
    "read_clients",
    "canonical_hash",
    "safe_name",
]
CATEGORIES = {
    "emesse": "fatture/emesse",
    "estere": "transfrontaliere/emesse",
    "ricevute": "fatture/ricevute",
    "mancata-consegna": "fatture/mc",
}
OPERATIONS = frozenset((*CATEGORIES, "corrispettivi", "bolli"))


class AcquisitionError(ValueError):
    """A bounded failure code suitable for status output, without page contents."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


def canonical_hash(value: Any) -> str:
    """Identify exact structured evidence, not its professional meaning."""
    return hashlib.sha256(
        json.dumps(
            value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        ).encode()
    ).hexdigest()


def safe_name(value: str) -> str:
    """Produce a bounded filename segment; client identity remains a separate ID."""
    value = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", value).strip(" .")[:80]
    return value or "documento"


def _identifier(value: Any, *, vat: bool = False) -> str:
    if isinstance(value, bool) or value is None:
        raise AcquisitionError("client-identifier-missing")
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            raise AcquisitionError("client-identifier-invalid")
        result = str(int(value)).zfill(11)
    else:
        result = str(value).strip().upper()
    pattern = r"[0-9]{11}" if vat else r"(?:[0-9]{11}|[A-Z0-9]{16})"
    if not re.fullmatch(pattern, result):
        raise AcquisitionError("client-identifier-invalid")
    return result


@dataclass(frozen=True)
class Client:
    """Exact taxpayer identity; same names never merge separate clients."""

    name: str
    tax_code: str
    vat_number: str
    archive_client_id: str = ""
    archive_engagement_id: str = ""

    @classmethod
    def parse(cls, data: dict[str, Any]) -> Client:
        allowed = set(cls.__dataclass_fields__)
        if not isinstance(data, dict) or set(data) - allowed:
            raise AcquisitionError("client-fields-invalid")
        name = data.get("name")
        if not isinstance(name, str) or not name.strip() or len(name) > 200:
            raise AcquisitionError("client-name-invalid")
        archive = [
            data.get(k, "") for k in ("archive_client_id", "archive_engagement_id")
        ]
        if any(not isinstance(v, str) for v in archive) or bool(archive[0]) != bool(
            archive[1]
        ):
            raise AcquisitionError("archive-binding-incomplete")
        return cls(
            name.strip(),
            _identifier(data.get("tax_code")),
            _identifier(data.get("vat_number"), vat=True),
            *archive,
        )


@dataclass(frozen=True)
class Plan:
    """One immutable requested population, re-enumerated on every run."""

    clients: tuple[Client, ...]
    date_from: date
    date_to: date
    operations: tuple[str, ...]
    access_mode: str = "delega_diretta"
    work_identity: str = ""
    max_documents: int = 20000

    @classmethod
    def parse(cls, data: dict[str, Any]) -> Plan:
        required = {"schema_version", "clients", "date_from", "date_to", "operations"}
        allowed = required | {"access_mode", "work_identity", "max_documents"}
        if (
            not isinstance(data, dict)
            or not required <= set(data)
            or set(data) - allowed
        ):
            raise AcquisitionError("plan-fields-invalid-no-credentials-allowed")
        if data["schema_version"] != "vera-agenzia-plan/v1":
            raise AcquisitionError("plan-version-invalid")
        try:
            start, end = date.fromisoformat(data["date_from"]), date.fromisoformat(
                data["date_to"]
            )
        except (ValueError, TypeError) as exc:
            raise AcquisitionError("plan-dates-invalid") from exc
        if start > end or start.year < 2000 or (end - start).days > 3660:
            raise AcquisitionError("plan-range-invalid")
        if (
            not isinstance(data["clients"], list)
            or not 1 <= len(data["clients"]) <= 1000
        ):
            raise AcquisitionError("plan-clients-required")
        clients = tuple(Client.parse(c) for c in data["clients"])
        if len({c.vat_number for c in clients}) != len(clients):
            raise AcquisitionError("duplicate-client-identity")
        operations = data["operations"]
        if (
            not isinstance(operations, list)
            or not operations
            or any(not isinstance(v, str) or v not in OPERATIONS for v in operations)
            or len(set(operations)) != len(operations)
        ):
            raise AcquisitionError("plan-operations-invalid")
        mode = data.get("access_mode", "delega_diretta")
        work = data.get("work_identity", "")
        if mode not in {"delega_diretta", "incaricato"}:
            raise AcquisitionError("access-mode-invalid")
        if mode == "incaricato":
            work = _identifier(work)
        elif work:
            raise AcquisitionError("work-identity-unexpected")
        limit = data.get("max_documents", 20000)
        if (
            isinstance(limit, bool)
            or not isinstance(limit, int)
            or not 1 <= limit <= 200000
        ):
            raise AcquisitionError("document-limit-invalid")
        return cls(clients, start, end, tuple(operations), mode, work, limit)

    def payload(self) -> dict[str, Any]:
        result = asdict(self)
        result.update(
            schema_version="vera-agenzia-plan/v1",
            date_from=self.date_from.isoformat(),
            date_to=self.date_to.isoformat(),
        )
        result["clients"] = [asdict(c) for c in self.clients]
        result["operations"] = list(self.operations)
        return result


def quarters(start: date, end: date) -> list[tuple[int, int, date, date]]:
    """Partition the exact inclusive period without gaps or duplicated dates."""
    result = []
    for year in range(start.year, end.year + 1):
        for quarter in range(1, 5):
            month = quarter * 3
            first = date(year, month - 2, 1)
            last = date(year, month, calendar.monthrange(year, month)[1])
            if first <= end and last >= start:
                result.append((year, quarter, max(start, first), min(end, last)))
    return result


def money(value: str | None) -> Decimal | None:
    """Parse explicit Italian monetary text; absent/invalid values never become zero."""
    if value is None or not str(value).strip() or str(value).strip() in {"-", "—"}:
        return None
    text = str(value).strip().replace("€", "").replace("EUR", "").strip()
    text = text.replace("\u00a0", "")
    if not re.fullmatch(r"-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d{1,2})?", text):
        raise AcquisitionError("monetary-value-invalid")
    try:
        return Decimal(text.replace(".", "").replace(",", ".")).quantize(
            Decimal("0.01")
        )
    except InvalidOperation as exc:
        raise AcquisitionError("monetary-value-invalid") from exc


def read_clients(path: Path, sheet: str | None = None) -> list[dict[str, Any]]:
    """Read declared CSV/XLSX client columns; reject incomplete rows with row numbers."""
    aliases = {
        "name": {"NOMINATIVO CLIENTE", "DENOMINAZIONE", "CLIENTE", "NAME"},
        "tax_code": {"CODICE FISCALE", "TAX_CODE"},
        "vat_number": {"PARTITA IVA", "P.IVA", "VAT_NUMBER"},
    }
    if path.suffix.lower() == ".csv":
        text = path.read_text(encoding="utf-8-sig")
        delimiter = ";" if text.count(";") >= text.count(",") else ","
        tables = [list(csv.reader(text.splitlines(), delimiter=delimiter))]
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        with path.open("rb") as stream:
            workbook = load_workbook(stream, data_only=True, read_only=True)
            try:
                if sheet and sheet not in workbook.sheetnames:
                    raise AcquisitionError("client-sheet-not-found")
                tables = [
                    list(ws.values)
                    for ws in workbook.worksheets
                    if not sheet or ws.title == sheet
                ]
            finally:
                workbook.close()
    else:
        raise AcquisitionError("client-file-format-unsupported")
    candidates = []
    for rows in tables:
        for index, row in enumerate(rows[:40]):
            headers = [re.sub(r"\s+", " ", str(v or "")).strip().upper() for v in row]
            matches = {
                key: [i for i, h in enumerate(headers) if h in names]
                for key, names in aliases.items()
            }
            if all(len(indices) == 1 for indices in matches.values()):
                clients = []
                for number, values in enumerate(rows[index + 1 :], index + 2):
                    if not any(v is not None and str(v).strip() for v in values):
                        continue
                    selected = {
                        key: values[indices[0]] if indices[0] < len(values) else None
                        for key, indices in matches.items()
                    }
                    try:
                        clients.append(asdict(Client.parse(selected)))
                    except AcquisitionError as exc:
                        raise AcquisitionError(
                            f"client-row-{number}-{exc.code}"
                        ) from exc
                candidates.append(clients)
                break
    if len(candidates) != 1 or not candidates[0]:
        raise AcquisitionError("client-table-missing-or-ambiguous")
    if len({c["vat_number"] for c in candidates[0]}) != len(candidates[0]):
        raise AcquisitionError("duplicate-client-identity")
    return candidates[0]
