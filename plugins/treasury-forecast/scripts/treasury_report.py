"""Render the same treasury record to reviewable native and text artifacts."""

from __future__ import annotations

import csv
import html
import json
from pathlib import Path
from typing import Any

from treasury_core import money, validate_record

__all__ = ["render_html", "render_markdown", "write_artifacts"]

STYLE = """
:root { color-scheme: light; font-family: 'Instrument Sans', Arial, sans-serif; color:#122b45; background:#fff; }
body { max-width:1120px; margin:48px auto; padding:0 24px 64px; line-height:1.5; }
h1 { font-size:34px; letter-spacing:-.025em; margin-bottom:8px; } h2 { margin-top:32px; font-size:21px; }
.label { color:#42617d; } .summary { display:flex; gap:40px; flex-wrap:wrap; margin:28px 0; }
.summary strong { display:block; font-size:26px; } table { border-collapse:collapse; width:100%; font-size:14px; }
th,td { padding:10px 12px; border-bottom:1px solid #dce4ed; text-align:left; vertical-align:top; }
th { color:#42617d; font-weight:600; } .scroll { overflow:auto; } .negative { color:#b3261e; }
.notice { border-left:3px solid #008cba; padding-left:16px; margin:24px 0; }
button { background:#002060; color:white; border:0; padding:10px 16px; border-radius:4px; cursor:pointer; }
button.secondary { background:#eef3f8; color:#002060; } input,textarea { font:inherit; padding:8px; border:1px solid #a7b8c9; border-radius:3px; }
input[type=date] { width:145px; } textarea { width:240px; min-height:42px; } .actions { display:flex; gap:12px; flex-wrap:wrap; margin:24px 0; }
small { color:#42617d; } @media(max-width:600px) { body { margin-top:24px; padding:0 16px 40px; } h1 { font-size:28px; } }
"""


def _display(amount: str) -> str:
    return (
        f"€ {money(amount):,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    )


def render_markdown(record: dict[str, Any]) -> str:
    """State calculation status prominently; do not turn missing dates into safety."""
    validate_record(record)
    lines = [
        f"# Budget di tesoreria — {record['company_name']}",
        "",
        f"Situazione al {record['as_of']} · Orizzonte al {record['horizon_end']}",
        "",
        f"Stato: **{record['status']}**",
        "",
        record["coverage"],
        "",
    ]
    if record["calculation_complete"]:
        lines.extend(
            [
                f"Disponibilità iniziale: {_display(record['opening_cash'])}.",
                f"Minimo saldo giornaliero previsto: {_display(record['minimum_daily_cash'])}.",
                f"Prima chiusura giornaliera negativa: {record['first_negative_day'] or 'nessuna nel periodo indicato'}.",
            ]
        )
    else:
        lines.append(
            "**Previsione incompleta: mancano date attese. I prospetti parziali non attestano la copertura dei pagamenti.**"
        )
    comparison = record["comparison"]
    if comparison:
        lines.extend(
            [
                "",
                "## Variazioni rispetto alla previsione precedente",
                "",
                f"Confronto fino al {comparison['through']}: {_display(comparison['closing_variance'])}.",
                f"Differenza tra cassa effettiva iniziale e precedente previsione: {_display(comparison['opening_variance'])}.",
            ]
        )
        for change in comparison["changes"]:
            before, after = change["previous"], change["current"]
            old = (
                f"{_display(before['cash_amount'])} il {before['expected_date']}"
                if before
                else "assente"
            )
            new = (
                f"{_display(after['cash_amount'])} il {after['expected_date']}"
                if after
                else "non più previsto"
            )
            lines.append(f"- {change['event_id']}: {old} → {new}.")
        if comparison["horizon_extended"]:
            lines.append(
                "Il periodo aggiunto dopo l'orizzonte precedente è esposto separatamente dal confronto."
            )
    lines.extend(["", "## Incassi e pagamenti attesi", ""])
    for event in record["events"]:
        lines.append(
            f"- {event['event_id']} · {event['description']}: {_display(event['cash_amount'])} · {event['expected_date'] or 'data da definire'}. Base: {event['basis']}"
        )
    if record["issues"] or record["evidence_notes"]:
        lines.extend(["", "## Questioni e movimenti da esaminare", ""])
        for row in [*record["issues"], *record["evidence_notes"]]:
            lines.append(f"- {row.get('event_id', row.get('id', ''))}: {row['detail']}")
    if record["review"]:
        lines.extend(
            [
                "",
                "## Decisione professionale",
                "",
                record["review"]["conclusion"],
                f"Revisore dichiarato: {record['review']['reviewer_ref']} · {record['review']['reviewed_at']}",
            ]
        )
    lines.extend(
        [
            "",
            "I saldi sono chiusure giornaliere. Gli incassi attesi restano ipotesi; il prospetto non garantisce liquidità infragiornaliera o completezza delle informazioni fornite.",
            "",
            f"Versione: {record['record_sha256']}",
            "",
        ]
    )
    return "\n".join(lines)


def render_html(record: dict[str, Any]) -> str:
    """Render a portable report; the separate live review saves actual decisions."""
    validate_record(record)
    escape = html.escape
    if record["calculation_complete"]:
        summary = f"<div class='summary'><div>Cassa iniziale<strong>{_display(record['opening_cash'])}</strong></div><div>Minimo giornaliero<strong>{_display(record['minimum_daily_cash'])}</strong></div><div>Prima data negativa<strong>{record['first_negative_day'] or 'Nessuna'}</strong></div></div>"
    else:
        summary = "<p class='notice'>Previsione incompleta: definire le date attese prima di utilizzare i saldi.</p>"
    rows = "".join(
        f"<tr><td>{escape(row['description'])}<br><small>{escape(row['event_id'])}</small></td><td>{escape(row['expected_date'] or 'Da definire')}</td><td>{_display(row['cash_amount'])}</td><td>{escape(row['basis'])}</td></tr>"
        for row in record["events"]
    )
    return f"<!doctype html><html lang='it'><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>Budget di tesoreria</title><style>{STYLE}</style><body><p class='label'>Vera · Budget di tesoreria</p><h1>{escape(record['company_name'])}</h1><p>{escape(record['as_of'])} — {escape(record['horizon_end'])} · {escape(record['status'])}</p><p>{escape(record['coverage'])}</p>{summary}<h2>Incassi e pagamenti attesi</h2><div class='scroll'><table><thead><tr><th>Voce</th><th>Data attesa</th><th>Flusso</th><th>Base della previsione</th></tr></thead><tbody>{rows}</tbody></table></div><h2>Prospetto e variazioni</h2><pre style='white-space:pre-wrap;font:inherit'>{escape(render_markdown(record))}</pre></body></html>"


def _csv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    with path.open("x", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(columns)
        for row in rows:
            values = []
            for key in columns:
                value = row.get(key, "")
                text = "" if value is None else str(value)
                # Text formula prefixes are escaped; monetary columns stay numeric.
                if key not in {
                    "amount",
                    "cash_amount",
                    "net_cash",
                    "closing_cash",
                    "minimum_daily_cash",
                    "cash_change_in_common_period",
                } and text.startswith(("=", "+", "-", "@", "\t", "\r")):
                    text = "'" + text
                values.append(text)
            writer.writerow(values)


def write_artifacts(directory: Path, record: dict[str, Any]) -> list[str]:
    """Write an immutable version package using the canonical calculated values."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    validate_record(record)
    directory.mkdir(parents=True, exist_ok=False)
    (directory / "forecast.json").write_text(
        json.dumps(record, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    (directory / "report.md").write_text(render_markdown(record), encoding="utf-8")
    (directory / "report.html").write_text(render_html(record), encoding="utf-8")
    tables = {
        "Giorni": (record["daily"], ["date", "net_cash", "closing_cash"]),
        "Settimane": (
            record["weekly"],
            ["week_start", "net_cash", "closing_cash", "minimum_daily_cash"],
        ),
        "Flussi": (
            record["events"],
            [
                "event_id",
                "description",
                "expected_date",
                "amount",
                "cash_amount",
                "basis",
                "decision_origin",
            ],
        ),
        "Variazioni": (
            (record["comparison"] or {}).get("changes", []),
            ["event_id", "cash_change_in_common_period"],
        ),
        "Questioni": (
            [*record["issues"], *record["evidence_notes"]],
            ["kind", "event_id", "id", "amount", "detail"],
        ),
    }
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Sintesi"
    summary.append(["Budget di tesoreria", record["company_name"]])
    summary.append(["Stato", record["status"]])
    summary.append(["Data situazione", record["as_of"]])
    summary.append(["Orizzonte", record["horizon_end"]])
    summary.append(["Cassa iniziale EUR", money(record["opening_cash"])])
    summary.append(
        [
            "Minimo giornaliero EUR",
            (
                money(record["minimum_daily_cash"])
                if record["calculation_complete"]
                else "Previsione incompleta"
            ),
        ]
    )
    summary.append(
        [
            "Prima data negativa",
            (
                (record["first_negative_day"] or "Nessuna")
                if record["calculation_complete"]
                else "Previsione incompleta"
            ),
        ]
    )
    summary.append(["Copertura dichiarata", record["coverage"]])
    summary.append(
        [
            "Decisione",
            record["review"]["conclusion"] if record["review"] else "Da registrare",
        ]
    )
    numeric = {
        "amount",
        "cash_amount",
        "net_cash",
        "closing_cash",
        "minimum_daily_cash",
        "cash_change_in_common_period",
    }
    for name, (rows, columns) in tables.items():
        _csv(directory / f"{name.lower()}.csv", rows, columns)
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for row in rows:
            sheet.append(
                [
                    (
                        money(row[key])
                        if key in numeric and row.get(key) not in {None, ""}
                        else row.get(key, "")
                    )
                    for key in columns
                ]
            )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                elif cell.value is not None:
                    cell.number_format = "#,##0.00;[Red](#,##0.00)"
                cell.font = Font(name="Instrument Sans", size=11, color="122B45")
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="002060")
            cell.font = Font(name="Instrument Sans", bold=True, color="FFFFFF")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 26
    workbook.save(directory / "tesoreria.xlsx")
    workbook.close()
    return sorted(path.name for path in directory.iterdir())
