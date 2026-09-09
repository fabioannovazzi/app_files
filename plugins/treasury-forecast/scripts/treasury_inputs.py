"""Read the published treasury tables and optional linked FatturaPA evidence."""

from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

# Exception type only; parsing delegates to the size-limited, DTD-rejecting parser.
from xml.etree.ElementTree import ParseError  # nosec B405
from zipfile import BadZipFile, ZipFile

from treasury_core import TreasuryError, money, validate_record

__all__ = [
    "HEADERS",
    "load_inputs",
    "manifest_paths",
    "read_json",
    "safe_path",
    "write_templates",
]

ROOT = Path(__file__).resolve().parents[1]
HEADERS = {
    "accounts": ["account_id", "balance"],
    "bank_movements": ["movement_id", "account_id", "date", "amount", "description"],
    "open_items": [
        "item_id",
        "party_id",
        "party_name",
        "side",
        "document_type",
        "document_number",
        "document_date",
        "installment",
        "due_date",
        "amount",
        "replaces_flow_id",
    ],
    "planned_flows": [
        "flow_id",
        "side",
        "amount",
        "expected_date",
        "description",
        "basis",
    ],
    "allocations": [
        "allocation_id",
        "movement_id",
        "target_type",
        "target_id",
        "amount",
    ],
    "adjustments": [
        "adjustment_id",
        "item_id",
        "date",
        "amount",
        "kind",
        "evidence_ref",
    ],
}
MONEY_COLUMNS = {"balance", "amount"}
MAX_SOURCE_BYTES = 50 * 1024 * 1024


def safe_path(root: Path, relative: str) -> Path:
    """Permit regular input files within the selected run, with no symlinks."""
    path = Path(relative)
    if path.is_absolute() or ".." in path.parts or not path.parts:
        raise TreasuryError("Source paths must be relative to the run inputs")
    cursor = root
    for part in path.parts:
        cursor = cursor / part
        if cursor.is_symlink():
            raise TreasuryError("Linked input paths are unsupported")
    resolved = cursor.resolve(strict=True)
    if not resolved.is_relative_to(root.resolve()) or not resolved.is_file():
        raise TreasuryError("Source escapes the input boundary")
    if resolved.stat().st_size > MAX_SOURCE_BYTES:
        raise TreasuryError("Input file exceeds the supported size")
    return resolved


def read_json(path: Path) -> dict[str, Any]:
    """Read a bounded object; duplicate fields cannot silently replace evidence."""

    def unique(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
        result: dict[str, Any] = {}
        for key, value in pairs:
            if key in result:
                raise TreasuryError(f"Duplicate JSON field: {key}")
            result[key] = value
        return result

    if path.stat().st_size > MAX_SOURCE_BYTES:
        raise TreasuryError("JSON input is too large")
    value = json.loads(path.read_text(encoding="utf-8"), object_pairs_hook=unique)
    if not isinstance(value, dict):
        raise TreasuryError("Expected a JSON object")
    return value


def manifest_paths(manifest: dict[str, Any], root: Path) -> list[Path]:
    """Resolve every source before a caller validates archive receipts."""
    if manifest.get("schema_version") != "vera.treasury_manifest.v1":
        raise TreasuryError("Unsupported input manifest")
    if set(manifest["tables"]) != set(HEADERS):
        raise TreasuryError("All six named tables are required, including empty tables")
    paths = [safe_path(root, item["path"]) for item in manifest["tables"].values()]
    paths.extend(safe_path(root, path) for path in manifest.get("invoice_files", []))
    if manifest.get("previous"):
        paths.append(safe_path(root, manifest["previous"]["path"]))
    return sorted(set(paths))


def _cell(value: Any, column: str) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        if column not in {"date", "document_date", "due_date", "expected_date"}:
            raise TreasuryError("Date cell appears in a non-date column")
        return (
            value.date().isoformat()
            if isinstance(value, datetime)
            else value.isoformat()
        )
    if column in MONEY_COLUMNS:
        if isinstance(value, bool):
            raise TreasuryError("Boolean is not a monetary value")
        return f"{money(str(value)):.2f}"
    if not isinstance(value, str):
        raise TreasuryError(
            f"{column} must be a text cell; preserve stable IDs explicitly"
        )
    return value.strip()


def _table(path: Path, name: str, sheet: str | None) -> list[dict[str, str]]:
    """Parse exact supported headers; there is no heuristic layout selection."""
    header = HEADERS[name]
    rows: list[dict[str, str]] = []

    def consume(values: list[Any]) -> None:
        if len(values) != len(header):
            raise TreasuryError(f"Ragged row in {name}")
        if len(rows) >= 100000:
            raise TreasuryError("Table exceeds the supported row count")
        rows.append({key: _cell(value, key) for key, value in zip(header, values)})

    if path.suffix.lower() == ".csv":
        if sheet:
            raise TreasuryError("CSV inputs do not have sheets")
        with path.open(newline="", encoding="utf-8-sig") as stream:
            reader = csv.reader(stream, strict=True)
            if next(reader, None) != header:
                raise TreasuryError(
                    f"Unsupported {name} CSV headers; use the published template"
                )
            for row in reader:
                consume(row)
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        try:
            with ZipFile(path) as archive:
                if (
                    sum(info.file_size for info in archive.infolist())
                    > 250 * 1024 * 1024
                ):
                    raise TreasuryError("Expanded workbook is too large")
        except BadZipFile as exc:
            raise TreasuryError("Invalid XLSX package") from exc
        workbook = load_workbook(
            path, read_only=True, data_only=False, keep_links=False
        )
        try:
            if sheet not in workbook.sheetnames:
                raise TreasuryError(f"Missing exact sheet for {name}")
            iterator = workbook[sheet].iter_rows()
            first = next(iterator, ())
            if [cell.value for cell in first] != header:
                raise TreasuryError(f"Unsupported {name} XLSX headers")
            for cells in iterator:
                if any(cell.data_type == "f" for cell in cells):
                    raise TreasuryError(
                        "Input formulas are unsupported; supply exported values"
                    )
                values = [cell.value for cell in cells]
                if any(value is not None for value in values):
                    consume(values)
        finally:
            workbook.close()
    else:
        raise TreasuryError("Only the specified CSV and XLSX tables are supported")
    return rows


def _invoices(manifest: dict[str, Any], root: Path) -> list[dict[str, Any]]:
    if not manifest.get("invoice_files"):
        return []
    parser_path = ROOT.parent / "client-file-preparation/scripts/parse_fatturapa_xml.py"
    spec = importlib.util.spec_from_file_location("treasury_fatturapa", parser_path)
    if spec is None or spec.loader is None:
        raise TreasuryError("The bundled FatturaPA parser is unavailable")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    result = []
    identities: dict[tuple[str, ...], str] = {}
    seen_bytes: set[str] = set()
    for relative in manifest["invoice_files"]:
        path = safe_path(root, relative)
        if path.suffix.lower() != ".xml":
            raise TreasuryError(
                "Supply extracted FatturaPA XML; ZIP/P7M are not input formats"
            )
        file_digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if file_digest in seen_bytes:
            continue
        seen_bytes.add(file_digest)
        try:
            invoices = module.parse_fatturapa_audit_file(path, root)
        except (ParseError, ValueError) as exc:
            raise TreasuryError(f"Unreadable FatturaPA evidence: {relative}") from exc
        for invoice in invoices:
            identity = tuple(
                str(invoice[key])
                for key in (
                    "supplier_vat",
                    "customer_tax_id",
                    "invoice_number",
                    "invoice_date",
                    "document_type",
                )
            )
            if not all(identity) or invoice["currency"] != "EUR":
                raise TreasuryError("Invoice identity or currency is unsupported")
            if identity in identities:
                raise TreasuryError(
                    "Different XML bytes represent the same invoice identity; review the sources"
                )
            identities[identity] = file_digest
            result.append({**invoice, "source_sha256": file_digest})
    return result


def load_inputs(
    manifest: dict[str, Any],
    root: Path,
    *,
    client_id: str,
    engagement_id: str,
) -> tuple[dict[str, Any], dict[str, Any] | None]:
    """Load qualified tables; XML supplies evidence, never unpaid balances."""
    paths = manifest_paths(manifest, root)
    data = {
        "schema_version": "vera.treasury_inputs.v1",
        "client_id": client_id,
        "engagement_id": engagement_id,
        **{
            key: manifest[key]
            for key in (
                "company_id",
                "company_name",
                "currency",
                "as_of",
                "horizon_end",
                "coverage",
            )
        },
        "sources": [
            {
                "path": path.relative_to(root.resolve()).as_posix(),
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            }
            for path in paths
        ],
    }
    for name, source in manifest["tables"].items():
        data[name] = _table(safe_path(root, source["path"]), name, source.get("sheet"))
    data["invoice_evidence"] = _invoices(manifest, root)
    previous = None
    if manifest.get("previous"):
        previous = read_json(safe_path(root, manifest["previous"]["path"]))
        validate_record(previous, accepted=True)
        if previous["record_sha256"] != manifest["previous"]["record_sha256"]:
            raise TreasuryError("Wrong predecessor forecast")
    return data, previous


def write_templates(output: Path) -> None:
    """Create exact empty input templates without replacing existing files."""
    output.mkdir(parents=True, exist_ok=True)
    paths = [output / f"{name}.csv" for name in HEADERS]
    if any(path.exists() for path in paths):
        raise TreasuryError("Input templates already exist; no overwrite performed")
    for name, header in HEADERS.items():
        with (output / f"{name}.csv").open("x", newline="", encoding="utf-8") as stream:
            csv.writer(stream).writerow(header)
