"""Render focused costing within the existing pack delivery and review flow."""

from __future__ import annotations

import html
from collections.abc import Mapping
from decimal import Decimal
from pathlib import Path
from typing import Any

from management_report_copy import localize_number
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

__all__ = ["render_costing_html", "render_costing_markdown", "write_costing_excel"]

LABELS = {
    "title": ("Costi e margini", "Costs and margins"),
    "draft": (
        "Bozza da rivedere con il professionista",
        "Draft for professional review",
    ),
    "synthetic": ("Esempio con dati sintetici", "Synthetic demonstration"),
    "question": ("Decisione da supportare", "Decision to support"),
    "methods": ("Perimetro e riconciliazione", "Scope and reconciliation"),
    "decisions": ("Alternative decisionali", "Decision alternatives"),
    "controls": ("Controlli", "Controls"),
    "limits": ("Limiti e dati mancanti", "Limitations and missing evidence"),
    "observations": ("Osservazioni", "Observations"),
    "hypotheses": ("Ipotesi", "Hypotheses"),
    "questions": ("Domande", "Questions"),
    "limitations": ("Limiti del commento", "Commentary limitations"),
    "object": ("Oggetto", "Object"),
    "method": ("Metodo", "Method"),
    "revenue": ("Ricavi", "Revenue"),
    "variable_cost": ("Costi variabili", "Variable costs"),
    "specific_fixed_cost": ("Fissi specifici", "Specific fixed costs"),
    "allocated_fixed_cost": ("Fissi allocati", "Allocated fixed costs"),
    "margin": ("Margine", "Margin"),
    "status": ("Stato", "Status"),
    "entity_profit": ("Risultato operativo", "Operating profit"),
    "retained_cost": ("Costi trattenuti", "Retained costs"),
    "unused_capacity": ("Capacità inutilizzata", "Unused capacity"),
    "reason": ("Motivo", "Reason"),
    "decision": ("Alternativa", "Alternative"),
    "delta_profit": ("Variazione risultato", "Profit change"),
    "stranded_fixed_cost": ("Fissi che restano", "Stranded fixed costs"),
    "basis": ("Ipotesi rivedute", "Reviewed assumptions"),
    "role": ("Controllo", "Control"),
    "actual": ("Calcolato", "Calculated"),
    "expected": ("Fonte riveduta", "Reviewed source"),
    "difference": ("Differenza", "Difference"),
    "available": ("Disponibile", "Available"),
    "unavailable": ("Non disponibile", "Unavailable"),
    "ready_for_review": ("Calcolato; da rivedere", "Calculated; review pending"),
    "partial": ("Parziale", "Partial"),
    "blocked": ("Bloccato", "Blocked"),
    "passed": ("Quadrato", "Passed"),
    "failed": ("Non quadrato", "Failed"),
    "direct_costing": (
        "Direct costing · margine I",
        "Direct costing · contribution margin",
    ),
    "direct_costing_evoluto": (
        "Direct evoluto · margine II",
        "Direct costing · second margin",
    ),
    "full_costing": ("Full costing", "Full costing"),
    "abc": ("Activity-based costing", "Activity-based costing"),
    "reviewed_source_revenue": ("Ricavi / totale fonte", "Revenue / source total"),
    "reviewed_source_costs": ("Costi / totale fonte", "Costs / source total"),
    "accounting_profit_bridge": (
        "Ponte risultato contabile",
        "Accounting profit bridge",
    ),
}
LIMITS = {
    "Source totals are reviewed input assertions; arithmetic does not prove source completeness.": "I totali fonte sono dati riveduti in ingresso: la quadratura aritmetica non prova la completezza delle fonti.",
    "Consumed operating costs only: no automatic statutory or tax inventory valuation.": "Il perimetro riguarda costi operativi consumati nel periodo; non valuta automaticamente le rimanenze civilistiche o fiscali.",
    "Allocations do not establish avoidability, savings or a decision to discontinue an activity.": "Un riparto non dimostra costi evitabili, risparmi o la convenienza a chiudere un’attività.",
    "Each view contains non-overlapping objects of one dimension; different views must not be added together.": "Ogni vista contiene oggetti non sovrapposti di una sola dimensione. Non sommare viste per prodotto, cliente e canale.",
    "ERP connectors, food/recipe costing and production optimization are outside this implemented path.": "Questo percorso non comprende connettori gestionali, food cost e ricette o ottimizzazione della produzione.",
    "No accounting-profit bridge supplied: reconciliation is limited to the reviewed costing source totals.": "Manca il ponte con il risultato contabile: la riconciliazione riguarda soltanto i totali fonte riveduti del costing.",
    "Synthetic demonstration: these results are not client or market evidence.": "Dati sintetici: questi risultati non costituiscono evidenza di un cliente o del mercato.",
}
NUMBERS = {
    "revenue",
    "variable_cost",
    "specific_fixed_cost",
    "allocated_fixed_cost",
    "margin",
    "entity_profit",
    "retained_cost",
    "unused_capacity",
    "delta_profit",
    "stranded_fixed_cost",
    "actual",
    "expected",
    "difference",
}


def label(key: str, language: str) -> str:
    return LABELS[key][0 if language == "it" else 1] if key in LABELS else key


def display(value: Any, key: str, language: str) -> str:
    if value is None:
        return "—"
    if key in NUMBERS:
        return localize_number(f"{Decimal(str(value)):,.2f}", language)
    return label(str(value), language)


def tables(
    pack: Mapping[str, Any],
) -> list[tuple[str, list[str], list[dict[str, Any]]]]:
    result = []
    for method in pack["sections"]["costing"]["requested_methods"]:
        rows = [
            row
            for row in pack["sections"]["costing"]["rows"]
            if row["method"] == method
        ]
        if rows:
            columns = ["object", "revenue", "variable_cost"]
            if method != "direct_costing":
                columns.append("specific_fixed_cost")
            if method in {"full_costing", "abc"}:
                columns.append("allocated_fixed_cost")
            result.append((method, [*columns, "margin"], rows))
    result.append(
        (
            "methods",
            [
                "method",
                "status",
                "entity_profit",
                "retained_cost",
                "unused_capacity",
                "reason",
            ],
            pack["sections"]["costing_methods"]["rows"],
        )
    )
    if pack["sections"]["costing_decisions"]["rows"]:
        result.append(
            (
                "decisions",
                ["decision", "delta_profit", "stranded_fixed_cost", "basis"],
                pack["sections"]["costing_decisions"]["rows"],
            )
        )
    result.append(
        (
            "controls",
            ["role", "status", "actual", "expected", "difference"],
            pack["controls"],
        )
    )
    return result


def render_costing_html(
    pack: Mapping[str, Any], commentary: Mapping[str, Any] | None = None
) -> str:
    """Render all selected methods and caveats without external assets or scripts."""
    language = pack["language"]
    esc = html.escape
    blocks = []
    for name, columns, rows in tables(pack):
        headings = "".join(
            f'<th scope="col">{esc(label(key, language))}</th>' for key in columns
        )
        body = "".join(
            "<tr>"
            + "".join(
                f'<td class="{"number" if key in NUMBERS else "text"}">{esc(display(row.get(key), key, language))}</td>'
                for key in columns
            )
            + "</tr>"
            for row in rows
        )
        blocks.append(
            f'<section><h2>{esc(label(name, language))}</h2><div class="table-scroll" tabindex="0" role="region" aria-label="{esc(label(name, language))}"><table><thead><tr>{headings}</tr></thead><tbody>{body}</tbody></table></div></section>'
        )
    if commentary:
        for key in ("observations", "hypotheses", "questions", "limitations"):
            items = "".join(f"<li>{esc(item['text'])}</li>" for item in commentary[key])
            if items:
                blocks.append(
                    f"<section><h2>{esc(label(key, language))}</h2><ul>{items}</ul></section>"
                )
    limits = "".join(
        f"<li>{esc(LIMITS.get(item, item) if language == 'it' else item)}</li>"
        for item in pack["limitations"]
    )
    synthetic = (
        f"<p>{esc(label('synthetic', language))}</p>" if pack["demonstration"] else ""
    )
    period = pack["reporting_period"]
    return f"""<!doctype html><html lang="{language}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{esc(label('title', language))} · {esc(pack['entity'])}</title>
<style>*{{box-sizing:border-box}}body{{margin:0;background:#fff;color:#182537;font:16px/1.6 "Instrument Sans",sans-serif}}main{{width:min(1120px,calc(100% - 40px));margin:auto;padding:40px 0 70px}}header{{border-top:5px solid #002060;padding:28px 0}}h1{{font-size:clamp(32px,5vw,54px);line-height:1.1;letter-spacing:-.035em;margin:10px 0 20px}}h2{{font-size:25px;color:#002060;line-height:1.25}}section{{padding:24px 0;border-top:1px solid #d7e1eb}}.meta{{color:#4f6074}}.question{{max-width:72ch}}table{{border-collapse:collapse;width:100%;font-size:14px}}th,td{{padding:12px 15px;border-bottom:1px solid #d7e1eb;vertical-align:top;text-align:left}}th{{background:#f4f7fb;color:#002060}}.number{{text-align:right;white-space:nowrap;font-variant-numeric:tabular-nums}}.text{{min-width:140px;overflow-wrap:anywhere}}.table-scroll{{overflow:auto}}.table-scroll:focus-visible{{outline:2px solid #0070c0;outline-offset:4px}}li{{margin-bottom:10px}}@media(max-width:640px){{main{{width:calc(100% - 28px);padding-top:20px}}h2{{font-size:22px}}}}</style></head><body><main><header><p>Vera · {esc(label('title', language))}</p><h1>{esc(pack['entity'])}</h1><p class="meta">{period['start']} / {period['end']} · EUR · {esc(label(pack['status'], language))}</p><p>{esc(label('draft', language))}</p>{synthetic}<p class="question"><strong>{esc(label('question', language))}:</strong> {esc(pack['sections']['costing']['question'])}</p></header>{''.join(blocks)}<section><h2>{esc(label('limits', language))}</h2><ul>{limits}</ul></section></main></body></html>"""


def render_costing_markdown(
    pack: Mapping[str, Any], commentary: Mapping[str, Any] | None = None
) -> str:
    """Write the same calculated schedules and distinct commentary to Markdown."""
    language = pack["language"]

    def safe(value: Any) -> str:
        return str(value).replace("|", "/").replace("\n", " ")

    lines = [
        f"# {label('title', language)} — {safe(pack['entity'])}",
        "",
        label("draft", language),
        "",
        pack["sections"]["costing"]["question"],
        "",
    ]
    if pack["demonstration"]:
        lines.extend([label("synthetic", language), ""])
    for name, columns, rows in tables(pack):
        lines.extend(
            [
                f"## {label(name, language)}",
                "",
                "| " + " | ".join(label(key, language) for key in columns) + " |",
                "| " + " | ".join("---" for _ in columns) + " |",
            ]
        )
        lines.extend(
            "| "
            + " | ".join(safe(display(row.get(key), key, language)) for key in columns)
            + " |"
            for row in rows
        )
        lines.append("")
    if commentary:
        for key in ("observations", "hypotheses", "questions", "limitations"):
            if commentary[key]:
                lines.extend(
                    [
                        f"## {label(key, language)}",
                        "",
                        *[f"- {item['text']}" for item in commentary[key]],
                        "",
                    ]
                )
    lines.extend(
        [
            f"## {label('limits', language)}",
            "",
            *[
                f"- {LIMITS.get(item, item) if language == 'it' else item}"
                for item in pack["limitations"]
            ],
        ]
    )
    return "\n".join(lines) + "\n"


def write_costing_excel(path: Path, pack: Mapping[str, Any]) -> None:
    """Persist calculated schedules and reviewed inputs; JSON holds exact cents."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Sintesi" if pack["language"] == "it" else "Summary"

    def safe(value: Any) -> Any:
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    for row in (
        ("Vera", label("title", pack["language"])),
        ("Entity", pack["entity"]),
        ("Question", pack["sections"]["costing"]["question"]),
        ("Status", pack["status"]),
        ("Review", label("draft", pack["language"])),
        ("Synthetic", str(pack["demonstration"])),
    ):
        summary.append(tuple(safe(value) for value in row))
    for name, columns, rows in tables(pack):
        sheet = workbook.create_sheet(name[:31])
        sheet.append([label(key, pack["language"]) for key in columns])
        for row in rows:
            sheet.append(
                [
                    (
                        float(Decimal(row[key]))
                        if key in NUMBERS and row.get(key) is not None
                        else safe(display(row.get(key), key, pack["language"]))
                    )
                    for key in columns
                ]
            )
        for index, key in enumerate(columns, start=1):
            if key in NUMBERS:
                for column in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
                    for cell in column:
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = "#,##0.00;[Red](#,##0.00)"
    workpapers = pack["costing_workpapers"]["case"]
    costs = workbook.create_sheet("Costi" if pack["language"] == "it" else "Costs")
    keys = (
        "id",
        "label",
        "amountCents",
        "behavior",
        "traceability",
        "objectId",
        "poolId",
        "evidenceIds",
    )
    costs.append(keys)
    for row in workpapers["costs"]:
        costs.append(
            [
                safe(
                    ", ".join(row[key])
                    if isinstance(row.get(key), list)
                    else row.get(key)
                )
                for key in keys
            ]
        )
    limits = workbook.create_sheet(
        "Limiti" if pack["language"] == "it" else "Limitations"
    )
    limits.append([label("limits", pack["language"])])
    for item in pack["limitations"]:
        limits.append(
            [safe(LIMITS.get(item, item) if pack["language"] == "it" else item)]
        )
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name="Instrument Sans", color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="002060")
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = (
                25 if column > 1 else 32
            )
    workbook.save(path)
    workbook.close()
