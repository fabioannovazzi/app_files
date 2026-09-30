"""Compose draft dossiers from reviewed records without certifying legal adequacy.

One plain-text document model drives both exports, preserving exact source
references and unresolved states. Rendering performs no semantic inference.
"""

from __future__ import annotations

import io
from html import escape
from pathlib import Path
from typing import Any

from .casebook import check_casebook
from .contracts import ContractError, canonical_hash, indexed
from .coordination import reconcile_declarations

__all__ = ["compose_dossier", "render_docx", "render_pdf"]

STATUS = {
    "EVIDENCED": "Fatto documentato nel registro",
    "PROPOSED": "Proposta da verificare",
    "UNRESOLVED": "Questione irrisolta",
}


def compose_dossier(
    proposal: dict[str, Any], result: dict[str, Any], decision: dict[str, Any]
) -> dict[str, Any]:
    """Build a source-located draft; missing template paragraphs stay explicit."""
    if (
        decision["proposal_digest"] != canonical_hash(proposal)
        or result["case_id"] != proposal["case"]["case_id"]
    ):
        raise ContractError("Dossier inputs belong to different reviewed versions")
    blocks: list[dict[str, Any]] = []

    def paragraph(text: str) -> None:
        blocks.append({"kind": "paragraph", "text": text})

    def heading(text: str, level: int = 1) -> None:
        blocks.append({"kind": "heading", "text": text, "level": level})

    def table(headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
        blocks.append(
            {"kind": "table", "headers": headers, "rows": rows, "widths": widths}
        )

    case, rules = proposal["case"], proposal["rules"]
    book = proposal.get("casebook")
    assessment = (
        check_casebook(book, case, rules, {r["key"] for r in proposal["controls"]})
        if book
        else None
    )
    paragraph(
        "BOZZA SINTETICA DA RIVEDERE"
        if case["demo"]
        else "BOZZA DA RIVEDERE DAL PROFESSIONISTA"
    )
    paragraph(
        "Il fascicolo raccoglie la posizione proposta, i calcoli e i riferimenti alla versione riesaminata. Le componenti escluse o sospese restano visibili. Il documento non è firmato e non attesta il riconoscimento della spettanza o dell’idoneità da parte dell’Amministrazione."
    )
    paragraph(
        f"Pratica {case['case_id']}. Periodo di fruizione {case['claim_period_id']}."
    )
    if book:
        paragraph(
            f"Contribuente: {book['taxpayer']['name']}. Obiettivo: {book['objective']}"
        )
    heading("Sintesi degli importi")
    table(
        ["Stato", "Redditi EUR", "IRAP EUR"],
        [[s, v["income"], v["irap"]] for s, v in result["bases"].items()],
        [40, 30, 30],
    )
    paragraph(
        f"Deduzione aggiuntiva redditi EUR {result['additional_deduction']['income']}; IRAP EUR {result['additional_deduction']['irap']}. Si riferisce alle sole componenti incluse. Non è un credito d’imposta né un risparmio fiscale già realizzato."
    )
    paragraph(
        f"Verifica documentale separata: {result['penalty_protection']['status']}. Firma, marca temporale, poteri e conservazione richiedono esiti distinti; la generazione del file non li verifica."
    )
    template = book["template"] if book else None
    heading("Indice documentale e fonte")
    if template:
        paragraph(
            f"Indice proposto per la pratica: {template['version']}; fonte {template['source_id']}; prova {template['evidence_id']}. La corrispondenza all’indice vigente e l’applicabilità restano oggetto di riesame professionale."
        )
        for simplification in template["simplifications"]:
            paragraph(
                f"Semplificazione {simplification['issue']} - {STATUS[simplification['status']]}: {simplification['conclusion']}. Prove: {', '.join(simplification['evidence_ids'])}; fonti: {', '.join(simplification['source_ids'])}."
            )
    else:
        paragraph(
            "DA COMPLETARE: indice vigente e applicabile non acquisito. Le sezioni seguenti costituiscono soltanto una struttura di lavoro."
        )
    facts = indexed(book["facts"], "fact_id") if book else {}
    paragraphs = indexed(book["paragraphs"], "paragraph_id") if book else {}
    for section in ("A", "B"):
        heading(f"Sezione {section}")
        required = (
            [p for p in template["required_paragraphs"] if p["section"] == section]
            if template
            else []
        )
        ordered = required + [
            p
            for p in paragraphs.values()
            if p["section"] == section
            and p["paragraph_id"] not in {r["paragraph_id"] for r in required}
        ]
        if not ordered:
            paragraph(
                "DA COMPLETARE: paragrafi, fatti e prove di questa sezione non sono registrati."
            )
        for item in ordered:
            heading(item["title"], 2)
            row = paragraphs.get(item["paragraph_id"])
            if row is None:
                paragraph(
                    f"DA COMPLETARE: paragrafo {item['paragraph_id']} richiesto dall’indice della pratica."
                )
                continue
            paragraph(f"{row['paragraph_id']} - {STATUS[row['status']]}.")
            paragraph(row["text"])
            for fid in row["fact_ids"]:
                fact = facts[fid]
                paragraph(
                    f"Fatto {fid} - {STATUS[fact['status']]}: {fact['statement']}. Prove: {', '.join(fact['evidence_ids'])}; {fact['locator']}."
                )
            paragraph("Fonti: " + (", ".join(row["source_ids"]) or "DA COMPLETARE"))
        for row in proposal["narratives"]:
            if row["section"] == section:
                heading("Testo istruttorio proposto", 2)
                paragraph(
                    "Proposta da verificare e raccordare ai paragrafi dell’indice."
                )
                paragraph(row["text"])
                paragraph(
                    "Prove: " + ", ".join(row["evidence_ids"]) + "; " + row["locator"]
                )
    heading("Raccordo tra costi attività progetti e beni")
    costs = indexed(case["costs"], "cost_id")
    allocations = indexed(case["allocations"], "allocation_id")
    for line in result["lines"]:
        allocation, cost = allocations[line["allocation_id"]], costs[line["cost_id"]]
        heading(f"Allocazione {line['allocation_id']}", 2)
        paragraph(
            f"{line['status']}; redditi EUR {line['income_amount']}; IRAP EUR {line['irap_amount']}. Motivi: {', '.join(line['reasons']) or 'Controlli della versione riesaminata superati'}."
        )
        paragraph(
            f"Costo {cost['cost_id']}; periodo {cost['period_id']}; riga {cost['ledger_row_key']}; prova {cost['evidence_id']}; bene {allocation['ip_id']}; progetto {allocation['project_id']}; attività {allocation['activity_id']}. Criterio: {allocation['allocation_method']}."
        )
    if "normalization_record" in proposal:
        heading("Normalizzazione della popolazione contabile")
        normal = proposal["normalization_record"]["result"]
        for total in normal["reviewed_control_totals"]:
            paragraph(
                f"Totale originale {total['value']} {total['currency']}; {total['evidence_id']}, {total['locator']}."
            )
        for item in normal["trace"]:
            paragraph(
                f"Costo {item['cost_id']}: EUR {item['book_amount']}; delta di arrotondamento {item['rounding_delta']}. {item['rationale']}"
            )
            for row in item["components"]:
                paragraph(
                    f"{row['evidence_id']}, {row['locator']}, riga {row['original_ledger_key']}: {row['source_amount']} {row['source_currency']}; EUR/unità {row['eur_per_unit']}; riferimento cambio {row['fx_rate_id'] or 'non necessario'}."
                )
        for rate in normal["fx_rates"]:
            paragraph(
                f"Cambio {rate['rate_id']}: {rate['evidence_id']}, {rate['locator']}; data {rate['rate_date']}. {rate['rationale']}"
            )
        for excluded in normal["excluded_rows"] + normal["non_data_rows"]:
            paragraph(f"Riga non inclusa {excluded['row_ref']}: {excluded['reason']}")
        for duplicate in normal["duplicate_findings"]:
            paragraph(
                f"Duplicato candidato {duplicate['finding_id']}: {duplicate['status']}; costi {', '.join(duplicate['cost_ids'])}. Nessuna eliminazione automatica."
            )
    if book and assessment:
        heading("Coordinamento con altri incentivi")
        for incentive in assessment["incentives"]:
            paragraph(
                f"{incentive['incentive_id']}: {incentive['assessment']}; {incentive['status']}."
            )
            for name, value in incentive["outputs"].items():
                paragraph(f"{name}: {value['value']} {value['unit']}.")
        for incentive in book["incentives"]:
            paragraph(
                incentive["conclusion"]
                + " Fonti: "
                + ", ".join(incentive["source_ids"])
            )
            paragraph(
                f"Azione: {incentive['action']}. Scadenza proposta: {incentive['due_date'] or 'non definita'}. {incentive['deadline_reason']}"
            )
        heading("Raccordo dichiarativo annuale")
        bridge = reconcile_declarations(
            book["declarations"], result, assessment["incentives"]
        )
        for model in book["declarations"]:
            paragraph(
                f"Modello {model['model_year']}; periodo {model['period_id']}; istruzioni {model['instructions_version']}; prova {model['model_evidence_id']}; fonte {model['instructions_source_id']}."
            )
        for row in bridge["rows"]:
            paragraph(
                f"{row['form']} / {row['field']} / {row['role']}: atteso {row['expected']}; riportato {row['reported'] if row['reported'] is not None else 'non riportato'}; {row['status']}; differenza {row['difference'] if row['difference'] is not None else 'non applicabile'}."
            )
        heading("Riesame critico")
        for issue in book["adversarial_review"]:
            paragraph(f"{issue['issue_id']}: {issue['argument']}")
            paragraph(
                f"Risposta proposta: {issue['response']}. Decisione proposta: {issue['decision']}. Redditi a rischio EUR {issue['income_at_risk']}; IRAP EUR {issue['irap_at_risk']}."
            )
            paragraph(
                "Prove: "
                + ", ".join(issue["evidence_ids"])
                + "; fonti: "
                + ", ".join(issue["source_ids"])
            )
        heading("Documenti mancanti e prossimi passi")
        for request in book["missing_documents"]:
            paragraph(
                f"{request['request_id']} - {request['priority']}: {request['request']}. Motivo: {request['reason']}. Referente: {request['owner']}. Conseguenza: {request['consequence']}. Prossimo passo: {request['next_step']}."
            )
        for gap in assessment["gaps"]:
            paragraph(gap["control_key"] + ": " + gap["reason"])
        heading("Richieste dell Ufficio e consegne")
        if not book["office_requests"]:
            paragraph(
                "Nessuna richiesta dell’Ufficio è registrata in questa versione della pratica."
            )
        for request in book["office_requests"]:
            paragraph(
                f"{request['request_id']}: ricevuta il {request['received_on']}; prova {request['evidence_id']}; consegna entro {request['delivery_due_on']}; fonte {request['deadline_source_id']}. Versione consegnata: {request['delivered_version'] or 'non registrata'}; ricevuta: {request['delivery_receipt_evidence_id'] or 'non registrata'}."
            )
    heading("Matrice dei controlli")
    for control in proposal["controls"]:
        heading(control["key"], 2)
        paragraph(f"{control['status']}: {control['conclusion']}")
        paragraph(
            "Prove: "
            + (", ".join(control["evidence_ids"]) or "mancanti")
            + "; fonti: "
            + (", ".join(control["source_ids"]) or "mancanti")
        )
    heading("Registro delle prove e delle fonti")
    for evidence in case["evidence"]:
        paragraph(
            f"{evidence['evidence_id']}: {evidence['description']}; SHA256 {evidence['sha256']}."
        )
    for source in rules["sources"]:
        paragraph(
            f"{source['source_id']}: prova {source['snapshot_evidence_id']}; SHA256 {source['snapshot_sha256']}."
        )
    paragraph(
        f"Regole {rules['ruleset_id']} / {rules['version']}; stato {rules['status']}; riesame {rules['reviewed_on']}; controllo fonti {rules['sources_checked_on']}."
    )
    heading("Versione e decisione")
    paragraph(
        f"Revisore dichiarato: {decision['reviewer']}; riferimento decisione {decision['confirmation_ref']}; data {decision['reviewed_on']}; livello di identità {decision['identity_assurance']}."
    )
    paragraph(
        f"Proposta SHA256 {decision['proposal_digest']}. Input SHA256 {result['input_hash']}. Regole SHA256 {result['rules_hash']}. Risultato SHA256 {result['result_hash']}."
    )
    paragraph(
        "I documenti restano bozze modificabili. Una modifica ai file esportati richiede una nuova verifica e una nuova versione prima di qualsiasi approvazione. Gli hash identificano i byte; non costituiscono firma elettronica o marca temporale."
    )
    document = {
        "schema_version": "1.0",
        "title": "Fascicolo Patent Box",
        "proposal_digest": decision["proposal_digest"],
        "result_hash": result["result_hash"],
        "status": "DRAFT_UNSIGNED",
        "blocks": blocks,
    }
    document["document_hash"] = canonical_hash(document)
    return document


def render_docx(document: dict[str, Any]) -> bytes:
    """Create editable Word paragraphs and tables with repeating header rows."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Header", "Footer"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(
            10
            if name in ("Normal", "Header", "Footer")
            else 24 if name == "Title" else 15 if name == "Heading 1" else 11
        )
        style.paragraph_format.space_after = Pt(6)
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.12
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(6)
    section.header.paragraphs[0].text = "Patent Box | Bozza non firmata"
    footer = section.footer.paragraphs[0]
    footer.text = "Pagina "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = document["title"]
    doc.core_properties.author = "Vera"
    doc.core_properties.comments = (
        "Bozza non firmata; riferimenti nella versione riesaminata."
    )
    doc.add_paragraph(document["title"], "Title")
    for block in document["blocks"]:
        if block["kind"] == "heading":
            doc.add_heading(block["text"], level=block["level"])
        elif block["kind"] == "paragraph":
            doc.add_paragraph(block["text"])
        else:
            table = doc.add_table(rows=1, cols=len(block["headers"]))
            table.autofit = False
            properties = table._tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement("w:" + edge)
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn("w:" + key), value)
                borders.append(border)
            properties.append(borders)
            header = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(header)
            for index, text in enumerate(block["headers"]):
                table.rows[0].cells[index].text = text
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EDEDED")
                table.rows[0].cells[index]._tc.get_or_add_tcPr().append(shading)
            for values in block["rows"]:
                for cell, text in zip(table.add_row().cells, values):
                    cell.text = text
            for row in table.rows:
                for index, cell in enumerate(row.cells):
                    cell.width = Cm(16.6 * block["widths"][index] / 100)
                    cell.vertical_alignment = 1
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_before = Pt(4)
                        paragraph.paragraph_format.space_after = Pt(4)
                        paragraph.alignment = 0 if index == 0 else 2
            doc.add_paragraph()
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render_pdf(document: dict[str, Any]) -> bytes:
    """Render with embedded bundled fonts; escape all source text as plain content."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        LongTable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        TableStyle,
    )

    fonts = Path(__file__).resolve().parents[1] / "assets/fonts"
    for name, filename in (
        ("PBRegular", "InstrumentSans-Regular.ttf"),
        ("PBBold", "InstrumentSans-SemiBold.ttf"),
    ):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(fonts / filename)))
    supported = pdfmetrics.getFont("PBRegular").face.charToGlyph
    texts = [document["title"]]
    for block in document["blocks"]:
        texts += (
            [block["text"]]
            if block["kind"] != "table"
            else block["headers"] + [cell for row in block["rows"] for cell in row]
        )
    if any(
        ord(char) not in supported
        for text in texts
        for char in text
        if not char.isspace()
    ):
        raise ContractError(
            "Dossier contains characters unsupported by the bundled PDF font; select a reviewed font before exporting"
        )
    body = ParagraphStyle(
        "Body",
        fontName="PBRegular",
        fontSize=9.5,
        leading=13.5,
        spaceAfter=5,
        splitLongWords=True,
    )
    title = ParagraphStyle(
        "Title", parent=body, fontName="PBBold", fontSize=23, leading=29, spaceAfter=16
    )
    h1 = ParagraphStyle(
        "H1",
        parent=body,
        fontName="PBBold",
        fontSize=14,
        leading=18,
        spaceBefore=12,
        spaceAfter=7,
        keepWithNext=True,
    )
    h2 = ParagraphStyle(
        "H2",
        parent=body,
        fontName="PBBold",
        fontSize=10.5,
        leading=14,
        spaceBefore=6,
        spaceAfter=5,
        keepWithNext=True,
    )

    def para(text: str, style: Any = body) -> Any:
        return Paragraph(escape(text, quote=False).replace("\n", "<br/>"), style)

    story = [para(document["title"], title)]
    for block in document["blocks"]:
        if block["kind"] == "heading":
            story.append(para(block["text"], h1 if block["level"] == 1 else h2))
        elif block["kind"] == "paragraph":
            story.append(para(block["text"]))
        else:
            data = [
                [para(c) for c in row] for row in [block["headers"], *block["rows"]]
            ]
            table = LongTable(
                data,
                colWidths=[(A4[0] - 124) * w / 100 for w in block["widths"]],
                repeatRows=1,
                hAlign="LEFT",
            )
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9D9D9")),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EDEDED")),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("TOPPADDING", (0, 0), (-1, -1), 7),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                    ]
                )
            )
            story += [table, Spacer(1, 8)]

    def page(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("PBRegular", 8)
        canvas.drawString(62, A4[1] - 32, "Patent Box | Bozza non firmata")
        canvas.drawRightString(A4[0] - 62, 30, f"Pagina {doc.page}")
        canvas.restoreState()

    buffer = io.BytesIO()
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=62,
        leftMargin=62,
        topMargin=55,
        bottomMargin=50,
        title=document["title"],
        author="Vera",
    )
    pdf.build(story, onFirstPage=page, onLaterPages=page)
    return buffer.getvalue()
