"""Render the same annual evidence into reviewable Word, Markdown and Excel."""

from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

from annual_statements import CHECKS, FIELDS, FORMULAS, RATIOS

__all__ = ["render_outputs"]


def _display(value: str | None) -> str:
    return "n.d." if value is None else value


def _tables(result: dict[str, Any]) -> list[tuple[str, list[list[str]]]]:
    years = result["years"]
    tables = []
    for title, population, keys, value_key in (
        ("Dati dichiarati nelle fonti", result["facts"], FIELDS, "value"),
        (
            "Indicatori e definizioni",
            result["metrics"],
            {m["key"]: m["label"] for m in result["metrics"]},
            "value",
        ),
        (
            "Riconciliazione delle fonti",
            result["checks"],
            {k: label for k, (label, _terms) in CHECKS.items()},
            "difference",
        ),
    ):
        by_key = {(item["key"], item["year"]): item for item in population}
        rows = [["Voce", *years]]
        for key, label in keys.items():
            values = []
            for year in years:
                item = by_key[(key, year)]
                text = _display(item[value_key])
                if value_key == "difference":
                    text += " / " + item["status"]
                values.append(text)
            rows.append([label, *values])
        tables.append((title, rows))
    return tables


def _paragraphs(result: dict[str, Any]) -> list[tuple[str, list[str]]]:
    exceptions = [
        f"{c['year']} — {c['label']}: {c['difference']} {result['currency']}"
        for c in result["checks"]
        if c["status"] == "exception"
    ]
    if not exceptions:
        exceptions = [
            "Nessuna differenza oltre la tolleranza nei controlli disponibili; verificare i controlli non disponibili e la copertura delle fonti."
        ]
    observations = [
        f"[{note['kind']}] {note['text']} ("
        + "; ".join(
            f"{ref['source_id']}: {ref['locator']}" for ref in note["references"]
        )
        + ")"
        for note in result["commentary"]
    ]
    if not observations:
        observations = [
            "Commento professionale da completare dopo la lettura dei risultati e delle note."
        ]
    coverage = [
        f"{source_id}: {entry['status']} — {entry['basis']}"
        for source_id, entry in result["coverage"].items()
    ]
    return [
        ("Eccezioni da riesaminare", exceptions),
        ("Analisi e domande", observations),
        ("Definizioni", result["definitions"]),
        ("Copertura documentale dichiarata dal revisore", coverage),
        ("Limiti", result["limitations"]),
    ]


def _word_and_markdown(result: dict[str, Any]) -> tuple[bytes, bytes]:
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.sections[0].left_margin = Inches(0.65)
    doc.sections[0].right_margin = Inches(0.65)
    title = "Analisi dei bilanci — " + result["entity"]
    intro = f"Periodi: {', '.join(result['years'])}. Valuta: {result['currency']}. Bozza per revisione professionale. Stato: {result['status']}."
    doc.add_heading(title, 0)
    doc.add_paragraph(intro)
    markdown = ["# " + title, "", intro, ""]
    paragraphs = _paragraphs(result)
    for heading, texts in paragraphs[:2]:
        doc.add_heading(heading, 1)
        markdown.extend(["## " + heading, ""])
        for text in texts:
            doc.add_paragraph(text)
            markdown.extend([text, ""])
    for heading, rows in _tables(result):
        doc.add_heading(heading, 1)
        table = doc.add_table(rows=1, cols=len(rows[0]))
        table.style = "Light Shading Accent 1"
        for cell, value in zip(table.rows[0].cells, rows[0]):
            cell.text = value
        for row in rows[1:]:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = value
        markdown.extend(
            [
                "## " + heading,
                "",
                "| " + " | ".join(rows[0]) + " |",
                "| " + " | ".join("---" for _ in rows[0]) + " |",
            ]
        )
        markdown.extend(
            "| " + " | ".join(value.replace("|", "\\|") for value in row) + " |"
            for row in rows[1:]
        )
        markdown.append("")
    for heading, texts in paragraphs[2:]:
        doc.add_heading(heading, 1)
        markdown.extend(["## " + heading, ""])
        for text in texts:
            doc.add_paragraph(text)
            markdown.extend([text, ""])
    doc.add_heading("Tracciabilità", 1)
    trace = "Il file Excel contiene celle originali, mapping e formule modificabili. Il JSON conserva gli importi Decimal, i controlli e i riferimenti; la ricevuta lega gli output alle fonti. Le formule Excel si ricalcolano all'apertura e non hanno cache certificata. Modificare il workbook non aggiorna Word, JSON o ricevuta: rigenerare il pacchetto dopo una revisione."
    doc.add_paragraph(trace)
    markdown.extend(["## Tracciabilità", "", trace, ""])
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue(), "\n".join(markdown).encode("utf-8")


def _text_cell(sheet: Any, row: int, column: int, value: Any) -> None:
    """Preserve user/source strings as text, including formula-looking labels."""
    cell = sheet.cell(row, column, str(value))
    cell.data_type = "s"


def _excel(result: dict[str, Any]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties

    book = openpyxl.Workbook()
    book.remove(book.active)
    sources = book.create_sheet("Fonti")
    sources.append(
        [
            "Voce",
            "Anno",
            "Importo",
            "Fonte",
            "Foglio",
            "Cella",
            "Valore originale",
            "Scala",
            "Stato",
        ]
    )
    cells = {}
    for index, fact in enumerate(result["facts"], 2):
        ref = fact["source"] or {}
        values = [
            fact["key"],
            fact["year"],
            None,
            ref.get("source_id", ""),
            ref.get("sheet", ""),
            ref.get("cell", ""),
            fact.get("source_value", ""),
            ref.get("scale", ""),
            fact["status"],
        ]
        for column, value in enumerate(values, 1):
            if value is not None:
                _text_cell(sources, index, column, value)
        if fact["value"] is not None:
            number = Decimal(fact["value"])
            # Excel guarantees only 15 significant digits. Do not silently lose them.
            if len(number.normalize().as_tuple().digits) > 15:
                raise ValueError(
                    "Amount exceeds Excel precision; an exact values export is required"
                )
            sources.cell(index, 3, number)
        cells[(fact["key"], fact["year"])] = f"'Fonti'!C{index}"
    calculations = book.create_sheet("Calcoli")
    calculations.append(["Voce", *result["years"]])
    rows = {}
    for index, (key, label) in enumerate(
        [
            *FIELDS.items(),
            *((k, v[0]) for k, v in FORMULAS.items()),
            *((k, v[0]) for k, v in RATIOS.items()),
            (
                "revenue_change_percent",
                "Variazione ricavi rispetto al periodo precedente disponibile (%)",
            ),
        ],
        2,
    ):
        rows[key] = index
        _text_cell(calculations, index, 1, label)
        for column, year in enumerate(result["years"], 2):
            letter = get_column_letter(column)
            if key in FIELDS:
                source = cells[(key, year)]
                formula = f'=IF(ISNUMBER({source}),{source},"")'
            elif key in FORMULAS:
                terms = FORMULAS[key][1]
                references = [f"{letter}{rows[k]}" for k, _ in terms]
                expression = "+".join(
                    f"({ref}*{sign})" for ref, (_k, sign) in zip(references, terms)
                )
                formula = f'=IF(COUNT({",".join(references)})={len(references)},{expression},"")'
            elif key in RATIOS:
                _label, n, d, scale = RATIOS[key]
                numerator, denominator = f"{letter}{rows[n]}", f"{letter}{rows[d]}"
                formula = f'=IF(AND(COUNT({numerator},{denominator})=2,{denominator}<>0),ROUND({numerator}/{denominator}*{scale},6),"")'
            elif column == 2:
                formula = '=""'
            else:
                current = f"{letter}{rows['revenue']}"
                previous = f"{get_column_letter(column-1)}{rows['revenue']}"
                formula = f'=IF(AND(COUNT({current},{previous})=2,{previous}<>0),ROUND(({current}/{previous}-1)*100,6),"")'
            calculations.cell(index, column, formula)
    checks = book.create_sheet("Controlli")
    checks.append(
        [
            "Controllo",
            "Anno",
            "Differenza (formula)",
            "Differenza (Decimal)",
            "Stato al calcolo",
            "Tolleranza",
        ]
    )
    for index, check in enumerate(result["checks"], 2):
        terms = CHECKS[check["key"]][1]
        letter = get_column_letter(result["years"].index(check["year"]) + 2)
        refs = [f"'Calcoli'!{letter}{rows[k]}" for k, _ in terms]
        expression = "+".join(f"({ref}*{sign})" for ref, (_k, sign) in zip(refs, terms))
        checks.append(
            [
                check["label"],
                check["year"],
                f'=IF(COUNT({",".join(refs)})={len(refs)},{expression},"")',
                check["difference"],
                check["status"],
                result["tolerance"],
            ]
        )
    values_sheet = book.create_sheet("Risultati Decimal")
    values_sheet.append(["Voce", "Anno", "Valore esatto o indice arrotondato", "Unità"])
    for row, item in enumerate(result["metrics"], 2):
        for col, value in enumerate(
            (item["label"], item["year"], _display(item["value"]), item["unit"]), 1
        ):
            _text_cell(values_sheet, row, col, value)
    notes = book.create_sheet("Metodo e note")
    notes.append(["Tipo", "Contenuto"])
    entries = [
        (
            "Stato",
            "Bozza. Modifiche Excel non aggiornano i risultati Decimal, Word o ricevute. Rigenerare dopo revisione.",
        ),
        ("Valuta", result["currency"]),
        ("Mapping", result["mapping_basis"]),
    ]
    entries.extend(
        (heading, text)
        for heading, paragraphs in _paragraphs(result)
        for text in paragraphs
    )
    entries.extend(
        ("Fonte " + key, f"{source['path']} | SHA256 {source['sha256']}")
        for key, source in result["sources"].items()
    )
    for row, (kind, text) in enumerate(entries, 2):
        _text_cell(notes, row, 1, kind)
        _text_cell(notes, row, 2, text)
    for sheet in book:
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="163A5F")
        sheet.column_dimensions["A"].width = 58
        for column in range(2, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 23
    notes.column_dimensions["B"].width = 115
    book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def render_outputs(result: dict[str, Any]) -> dict[str, bytes]:
    """Produce review artifacts; no model call or financial judgment occurs here."""
    word, markdown = _word_and_markdown(result)
    return {
        "annual_report.docx": word,
        "annual_report.md": markdown,
        "annual_analysis.xlsx": _excel(result),
    }
