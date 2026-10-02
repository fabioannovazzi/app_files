"""Local mechanics for teaching Vera a studio's document format."""

from __future__ import annotations

import argparse
import importlib.util
import json
import logging
import secrets
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

__all__ = [
    "inspect_examples",
    "prepare_review",
    "create_previews",
    "practice_materials",
    "main",
    "prepare_reuse_case",
]
ROOT = Path(__file__).resolve().parents[1]
LOGGER = logging.getLogger(__name__)


def _component(name: str) -> Path:
    """Resolve the packaged component or its repository source."""
    packaged = ROOT / "modules" / name
    return packaged if packaged.is_dir() else ROOT.parent / name


COMM = _component("comunicazione-professionale")
sys.path.insert(0, str(COMM / "scripts"))
from review_document_format import approve_format_review, prepare_format_review
from workflow_core import (
    atomic_write_json,
    canonical_digest,
    file_digest,
    load_json,
    workflow_lock,
)


def _formatting() -> Any:
    spec = importlib.util.spec_from_file_location(
        "studio_formatting",
        _component("report-builder") / "scripts/studio_formatting.py",
    )
    if not spec or not spec.loader:
        raise RuntimeError("The report formatting module is unavailable")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def inspect_examples(samples: list[Path], destination: Path) -> Path:
    """Extract Word formatting evidence without emitting text or author metadata."""
    if destination.exists():
        raise ValueError("Use a fresh evidence path")
    records = []
    for path in samples:
        if path.is_symlink() or not path.is_file() or path.suffix.lower() != ".docx":
            raise ValueError(
                "Structural inspection supports selected regular DOCX files only"
            )
        document = Document(path)
        styles = []
        for style in document.styles:
            if style.type != 1:
                continue
            styles.append(
                {
                    "style_id": style.style_id,
                    "base_style_id": (
                        style.base_style.style_id if style.base_style else None
                    ),
                    "font": style.font.name,
                    "size_pt": style.font.size.pt if style.font.size else None,
                    "color": (
                        str(style.font.color.rgb) if style.font.color.rgb else None
                    ),
                    "before_pt": (
                        style.paragraph_format.space_before.pt
                        if style.paragraph_format.space_before is not None
                        else None
                    ),
                    "after_pt": (
                        style.paragraph_format.space_after.pt
                        if style.paragraph_format.space_after is not None
                        else None
                    ),
                }
            )
        # Count exact formatting occurrences, never classify a dominant style as a studio rule.
        runs: dict[str, int] = {}
        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                signature = json.dumps(
                    {
                        "paragraph_style_id": paragraph.style.style_id,
                        "font": run.font.name,
                        "size_pt": run.font.size.pt if run.font.size else None,
                        "bold": run.bold,
                        "color": (
                            str(run.font.color.rgb) if run.font.color.rgb else None
                        ),
                    },
                    sort_keys=True,
                )
                runs[signature] = runs.get(signature, 0) + 1
        records.append(
            {
                "path": str(path.resolve()),
                "sha256": file_digest(path),
                "sections": [
                    {
                        "page_mm": [s.page_width.mm, s.page_height.mm],
                        "margins_mm": {
                            k: getattr(s, f"{k}_margin").mm
                            for k in ("top", "bottom", "left", "right")
                        },
                        "header_distance_mm": s.header_distance.mm,
                        "footer_distance_mm": s.footer_distance.mm,
                        "has_header": bool(
                            s.header._element.xpath(".//w:t|.//w:drawing")
                        ),
                        "has_footer": bool(
                            s.footer._element.xpath(".//w:t|.//w:fldSimple")
                        ),
                    }
                    for s in document.sections
                ],
                "styles": styles,
                "run_formats": [
                    {**json.loads(k), "occurrences": v} for k, v in runs.items()
                ],
                "table_count": len(document.tables),
                "limitations": "Style and filename identifiers may identify the studio. No document text is emitted. Themes, inherited formatting and font availability need visual/model review; metadata is not a learned standard.",
            }
        )
    atomic_write_json(
        destination, {"kind": "studio_format_evidence", "samples": records}
    )
    return destination


def prepare_review(
    workspace: Path,
    *,
    review_id: str,
    settings: Path,
    samples: list[Path],
    brand: Path | None = None,
    language: str = "it",
) -> Path:
    """Prepare model-selected settings; first-run channel defaults remain labelled."""
    if language not in {"it", "en", "fr", "de", "es"}:
        raise ValueError("Select a supported proposal language")
    base_path = None
    if not (workspace / "studio_profile.json").is_file():
        if brand is None:
            raise ValueError(
                "First setup requires studio identity and selected brand data"
            )
        base_path = workspace / f".{review_id}-base.json"
        if base_path.exists():
            raise ValueError("Temporary base proposal already exists")
        brand_data = load_json(brand)
        baseline_profile = load_json(
            COMM / f"assets/studio-document-format/baseline-profile-{language}.json"
        )

        def expand(value: Any) -> Any:
            if isinstance(value, dict):
                return {k: expand(v) for k, v in value.items()}
            if isinstance(value, list):
                return [expand(v) for v in value]
            if isinstance(value, str):
                return (
                    value.replace("{studio}", brand_data["studio_name"])
                    .replace("{contact}", brand_data["contact_line"])
                    .replace("{location}", "")
                )
            return value

        atomic_write_json(
            base_path,
            {"brand_profile": brand_data, "profile": expand(baseline_profile)},
        )
    try:
        review = prepare_format_review(
            workspace,
            review_id=review_id,
            settings_path=settings,
            samples=samples,
            base_profile_path=base_path,
        )
    finally:
        if base_path is not None:
            base_path.unlink(missing_ok=True)
    create_previews(review)
    return review


def create_previews(review: Path) -> Path:
    """Build short and long synthetic Word previews bound to the exact proposal."""
    with workflow_lock(review.resolve().parent.parent):
        payload = load_json(review / "format_review.json")
        digest = payload["review_digest"]
        if (
            canonical_digest({k: v for k, v in payload.items() if k != "review_digest"})
            != digest
        ):
            raise ValueError("Format proposal changed")
        for record in payload["samples"] + (
            [payload["logo"]] if payload["logo"] else []
        ):
            path = Path(record["snapshot_path"]).resolve()
            if (
                not path.is_relative_to(review.resolve() / "inputs")
                or file_digest(path) != record["sha256"]
            ):
                raise ValueError("Format evidence changed")
        module = _formatting()
        profile = module.resolve_studio_format(
            payload["profile"]["document"],
            studio_id=payload["workspace_id"],
            studio_name=payload["studio_name"],
            version=0,
            format_digest=digest,
            logo_bytes=(
                Path(payload["logo"]["snapshot_path"]).read_bytes()
                if payload["logo"]
                else None
            ),
        )
        records = []
        for name, count in (("short", 4), ("long", 100)):
            document = Document()
            document.add_heading("SYNTHETIC FORMAT PREVIEW", level=1)
            document.add_paragraph(
                "Fictional data. This demonstrates presentation only; it is not client work or an approved report."
            )
            document.add_paragraph(
                module.display_date("2026-09-30", profile["settings"])
            )
            document.add_heading("Document and table review", level=2)
            document.add_paragraph(
                "Check the requested font, spacing, margins, studio header, footer, page numbers and logo. Check repeated table headings and rows on each page of the long example."
            )
            table = document.add_table(rows=1, cols=2)
            caption = OxmlElement("w:tblCaption")
            caption.set(qn("w:val"), "Reviewed numeric totals")
            table._tbl.tblPr.append(caption)
            table.rows[0].cells[0].text = "Fictional item"
            table.rows[0].cells[1].text = "EUR"
            for i in range(count):
                cells = table.add_row().cells
                cells[0].text = (
                    f"Synthetic row {i + 1}: a longer description for checking wrapping without clipped text."
                )
                cells[1].text = module.display_number(
                    "1234.50" if i % 2 else "-234.50", profile["settings"]
                )
            for line in profile["settings"]["signature_lines"]:
                document.add_paragraph(line)
            module.apply_studio_format(document, profile)
            output = review / f"preview-{name}.docx"
            document.save(output)
            records.append({"path": output.name, "sha256": file_digest(output)})
        packet = {
            "kind": "document_format_previews",
            "review_digest": digest,
            "outputs": records,
            "font_status": profile["font_status"],
            "visual_review": "User must open both previews and review actual rendering before approval; byte hashes do not prove visual acceptance.",
        }
        atomic_write_json(review / "preview_manifest.json", packet)
        return review / "preview_manifest.json"


def practice_materials(spec: Path, destination: Path) -> Path:
    """Generate fictional source examples locally from an authored course spec."""
    if destination.exists():
        raise ValueError("Use a fresh course material directory")
    data = load_json(spec)
    destination.mkdir(parents=True, mode=0o700)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = data["font"]
    style.font.size = Pt(data["size_pt"])
    heading = doc.styles["Heading 1"]
    heading.font.name = data["font"]
    heading.font.size = Pt(data["heading_pt"])
    heading.font.color.rgb = RGBColor.from_string(data["color"])
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        for side, value in data["margins_mm"].items():
            setattr(section, f"{side}_margin", Mm(value))
        section.header.paragraphs[0].text = data["studio_name"]
        section.footer.paragraphs[0].text = data["footer"]
    doc.add_heading(data["title"], 1)
    doc.add_paragraph(data["body"])
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = data["item_label"]
    table.rows[0].cells[1].text = "EUR"
    for label, value in data["rows"]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    doc.save(destination / "studio-example.docx")
    (destination / "financial-source.csv").write_text(
        "item,amount\n"
        + "".join(f"{label},{value}\n" for label, value in data["rows"]),
        encoding="utf-8",
    )
    atomic_write_json(
        destination / "source_manifest.json",
        {
            "kind": "synthetic_course_sources",
            "spec_sha256": file_digest(spec),
            "outputs": [
                {"path": p.name, "sha256": file_digest(p)}
                for p in sorted(destination.iterdir())
            ],
        },
    )
    return destination


def prepare_reuse_case(tutorial_case: Path, source: Path, workspace: Path) -> Path:
    """Create an actual synthetic report run inside the bound studio-format lesson."""
    descriptor = load_json(tutorial_case)
    root = tutorial_case.resolve().parent
    if (
        descriptor.get("tutorial") is not True
        or descriptor.get("local_only") is not True
        or descriptor.get("workflow_id") != "studio-document-format"
        or Path(descriptor["directory"]).resolve() != root
        or Path(descriptor["output_dir"]).resolve() != root / "outputs"
    ):
        raise ValueError("Select the bound studio-format tutorial case")
    for path in (source, workspace):
        if path.is_symlink() or not path.resolve().is_relative_to(root / "outputs"):
            raise ValueError(
                "Reuse inputs and studio workspace must stay in this tutorial attempt"
            )
    profile = load_json(workspace / "studio_profile.json")
    _formatting().load_studio_format(
        workspace, studio_id=profile["workspace_id"], studio_name=profile["studio_name"]
    )
    if source.suffix.lower() != ".csv" or not source.is_file():
        raise ValueError("Select the generated fictional financial CSV")
    record_path = root / "report_reuse_case.json"
    if record_path.exists():
        raise ValueError("A reuse report case already exists; resume its actual run")
    spec = importlib.util.spec_from_file_location(
        "studio_format_tutorial_ledger",
        _component("studio-archive") / "scripts/client_ledger.py",
    )
    if not spec or not spec.loader:
        raise RuntimeError("Studio Archive is unavailable")
    ledger = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = ledger
    spec.loader.exec_module(ledger)
    case = root / "outputs/financial-reuse-client"
    case.mkdir(mode=0o700)
    client_id = "client_" + secrets.token_hex(12)
    ledger.create_client_manifest(case, client_id)
    engagement = ledger.create_engagement(
        case, client_id, "Synthetic studio-format report reuse"
    )
    engagement_id = engagement["engagement_id"]
    imported = ledger.import_document(case, client_id, engagement_id, source, "source")
    component = _component("report-builder")
    manifest = component / ".codex-plugin/plugin.json"
    if not manifest.is_file():
        manifest = component / ".claude-plugin/plugin.json"
    version = load_json(manifest)["version"]
    run = ledger.prepare_run(
        case,
        client_id,
        engagement_id,
        "report-builder",
        version,
        input_ids=[imported["receipt"]["input_id"]],
        purpose="Local synthetic studio-format lesson; no hosted transmission",
    )
    started = ledger.start_run(case, engagement_id, run["run"]["run_id"])
    atomic_write_json(
        record_path,
        {
            "tutorial": True,
            "local_only": True,
            "workflow_id": "financial-report-builder",
            "client_root": str(case),
            "studio_workspace": str(workspace.resolve()),
            "studio_id": profile["workspace_id"],
            "studio_name": profile["studio_name"],
            **started,
            "client_engagement_path": started["context"]["context_path"],
            "input_dir": started["context"]["input_dir"],
            "output_dir": started["context"]["output_dir"],
        },
    )
    return record_path


def main(argv: list[str] | None = None) -> int:
    """Expose internal workflow operations; the studio supplies ordinary language."""
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="action", required=True)
    inspect = sub.add_parser("inspect")
    inspect.add_argument("--sample", type=Path, action="append", required=True)
    inspect.add_argument("--output", type=Path, required=True)
    prepare = sub.add_parser("prepare")
    prepare.add_argument("--workspace", type=Path, required=True)
    prepare.add_argument("--review-id", required=True)
    prepare.add_argument("--settings", type=Path, required=True)
    prepare.add_argument("--sample", type=Path, action="append", default=[])
    prepare.add_argument("--brand", type=Path)
    prepare.add_argument(
        "--language", choices=["it", "en", "fr", "de", "es"], default="it"
    )
    preview = sub.add_parser("preview")
    preview.add_argument("--review-dir", type=Path, required=True)
    adopt = sub.add_parser("approve")
    adopt.add_argument("--review-dir", type=Path, required=True)
    adopt.add_argument("--review-digest", required=True)
    adopt.add_argument("--reviewer", required=True)
    adopt.add_argument("--confirmed-by-user", action="store_true")
    practice = sub.add_parser("practice-materials")
    practice.add_argument("--spec", type=Path, required=True)
    practice.add_argument("--output-dir", type=Path, required=True)
    reuse = sub.add_parser("report-case")
    reuse.add_argument("--tutorial-case", type=Path, required=True)
    reuse.add_argument("--source", type=Path, required=True)
    reuse.add_argument("--studio-workspace", type=Path, required=True)
    args = parser.parse_args(argv)
    if args.action == "inspect":
        result = inspect_examples(args.sample, args.output)
    elif args.action == "prepare":
        result = prepare_review(
            args.workspace,
            review_id=args.review_id,
            settings=args.settings,
            samples=args.sample,
            brand=args.brand,
            language=args.language,
        )
    elif args.action == "preview":
        result = create_previews(args.review_dir)
    elif args.action == "approve":
        result = approve_format_review(
            args.review_dir,
            review_digest=args.review_digest,
            reviewer=args.reviewer,
            confirmed_by_user=args.confirmed_by_user,
        )
    elif args.action == "report-case":
        result = prepare_reuse_case(
            args.tutorial_case, args.source, args.studio_workspace
        )
    else:
        result = practice_materials(args.spec, args.output_dir)
    LOGGER.info("Studio format: %s", result)
    return 0


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    raise SystemExit(main())
