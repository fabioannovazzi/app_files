"""Producer-owned source coordinates for bounded packets, before any delivery."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from archive_emit import bound_run
from archive_report_store import RunStore, digest, file_digest, run_lock

__all__ = [
    "is_recording",
    "register_metadata",
    "register_metadata_revision",
    "register_private_value",
    "register_projection",
    "register_rows",
    "register_row_projection",
    "register_value",
    "run_cli",
]


def run_cli(main: Any, component_file: Path) -> int | None:
    """Delegate direct component CLI delivery to Vera's owned recorder."""
    from archive_cli import run_entrypoint

    return run_entrypoint(main, component_file)


def is_recording(component: str, output: Path) -> bool:
    """Identify an explicitly bound selected run through Archive's stored state."""
    return bound_run(component, {"output": str(output)}) is not None


def _positions(source: dict[str, Any], entry: dict[str, Any]) -> dict[int, int] | None:
    """Convert producer physical rows to the same units used at Archive start."""
    if source["unit"] == "rows" and Path(source["path"]).suffix.lower() in {
        ".csv",
        ".tsv",
    }:
        return {
            row["physical_row"]: row["physical_row"] - 2
            for row in entry["rows"]
            if row["physical_row"] > 1
        }
    if source["unit"] != "worksheet_rows":
        return None
    from openpyxl import load_workbook

    wanted = {row["physical_row"] for row in entry["rows"]}
    result = {}
    index = 0
    workbook = load_workbook(source["path"], read_only=True, data_only=False)
    try:
        for sheet in workbook:
            for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                if any(value is not None for value in values):
                    if sheet.title == entry["sheet_name"] and number in wanted:
                        result[number] = index
                    index += 1
    finally:
        workbook.close()
    return result


def register_metadata(component: str, output: Path, payload: Any) -> None:
    """Declare an exact producer status message, never a semantic classification."""
    root = bound_run(component, {"output": str(output)})
    if root is None:
        return
    with run_lock(root):
        RunStore(root).event(
            "metadata:" + digest(payload),
            {
                "kind": "prepared_metadata",
                "data_sha256": digest(payload),
            },
        )


def register_rows(
    component: str, output: Path, records: dict[str, dict[str, Any]]
) -> None:
    """Retain the normalizer's field origins without counting preparation as a send."""
    root = bound_run(component, {"output": str(output)})
    if root is None:
        return
    with run_lock(root):
        RunStore(root).event(
            "rows:" + digest(records),
            {
                "kind": "prepared_source_rows",
                "component": component,
                "records": records,
            },
        )


def register_private_value(component: str, output: Path, value: Any) -> None:
    """Bind a locally produced review packet; this does not declare a model send."""
    root = bound_run(component, {"output": str(output)})
    if root is not None:
        with run_lock(root):
            RunStore(root).event(
                "private:" + digest(value),
                {
                    "kind": "prepared_private_value",
                    "component": component,
                    "data_sha256": digest(value),
                },
            )


def register_metadata_revision(
    component: str, path: Path, previous_sha256: str
) -> None:
    """Carry an exact metadata parent through a producer's metadata-only edit."""
    root = bound_run(component, {"output": str(path)})
    if root is None:
        return
    with run_lock(root):
        known_parent = any(
            event["kind"] == "prepared_projection"
            and event["artifact_sha256"] == previous_sha256
            and not event["selections"]
            for event in RunStore(root).events()
        )
    if known_parent:
        register_projection(component, path, [])


def register_row_projection(
    component: str,
    output: Path,
    value: Any,
    records: list[dict[str, Any]],
    *,
    private_sha256: str | None = None,
) -> None:
    """Resolve only fields selected by the producer, never infer origin from values."""
    root = bound_run(component, {"output": str(output)})
    if root is None:
        return
    with run_lock(root):
        events = RunStore(root).events()
        if private_sha256 is not None and not any(
            event["kind"] == "prepared_private_value"
            and event["component"] == component
            and event["data_sha256"] == private_sha256
            for event in events
        ):
            return  # An edited or supplied review cannot claim producer provenance.
        origins = {}
        for event in events:
            if (
                event["kind"] == "prepared_source_rows"
                and event["component"] == component
            ):
                origins.update(event["records"])
        entries = []
        for record in records:
            if (
                not isinstance(record, dict)
                or not isinstance(record.get("record_id"), str)
                or not isinstance(record.get("fields"), list)
                or any(not isinstance(field, str) for field in record["fields"])
            ):
                return
            origin = origins.get(record["record_id"])
            if origin is None or any(
                field not in origin["fields"] for field in record["fields"]
            ):
                return  # Missing producer evidence remains an incomplete report.
            columns = sorted(
                {
                    column
                    for field in record["fields"]
                    for column in origin["fields"][field]
                }
            )
            entries.append(
                {
                    **origin["source"],
                    "rows": [
                        {"physical_row": origin["physical_row"], "columns": columns}
                    ],
                }
            )
    register_value(component, output, value, entries)


def register_projection(
    component: str, path: Path, entries: list[dict[str, Any]]
) -> None:
    """Bind a generated JSON packet to exact cells in the started run's inputs.

    Only pipeline code calls this while constructing its bounded output. This is
    preparation evidence; no send is counted until Vera returns the packet.
    """
    if not is_recording(component, path):
        return
    if path.is_symlink():
        raise ValueError("A producer projection must be an ordinary file")
    packet_hash = file_digest(path)
    value = json.loads(path.read_text(encoding="utf-8"))
    if file_digest(path) != packet_hash:
        raise ValueError("Producer projection changed while it was registered")
    register_value(component, path, value, entries, artifact_sha256=packet_hash)


def register_value(
    component: str,
    output: Path,
    value: Any,
    entries: list[dict[str, Any]],
    *,
    artifact_sha256: str | None = None,
) -> None:
    """Retain producer provenance for this exact value, without counting a send."""
    root = bound_run(component, {"output": str(output)})
    if root is None:
        return
    with run_lock(root):
        store = RunStore(root)
        state = store.state()
        if state is None or state["collection_closed"]:
            raise ValueError("Producer projection requires an active Archive run")
        sources = store.sources()
        selections = []
        for entry in entries:
            source = next(
                (
                    item
                    for item in sources
                    if item["role"] == "input"
                    and item["sha256"] == entry["source_sha256"]
                    and str(Path(entry["source_path"]).resolve())
                    in item.get("aliases", [item["path"]])
                ),
                None,
            )
            if source is None or source["measurement_error"]:
                return  # Unsupported source provenance must remain explicit.
            if entry.get("source_column_count", source["columns"]) != source["columns"]:
                return  # The source reader and population measure must agree.
            source_path = Path(source["path"])
            if source_path.is_symlink() or file_digest(source_path) != source["sha256"]:
                raise ValueError(
                    "Producer projection source changed after Archive start"
                )
            positions = _positions(source, entry)
            if file_digest(source_path) != source["sha256"]:
                raise ValueError(
                    "Producer projection source changed during measurement"
                )
            if positions is None:
                return
            for row in entry["rows"]:
                index, columns = positions.get(row["physical_row"]), row["columns"]
                if index is None:
                    continue  # CSV column names are metadata, not data rows.
                if (
                    type(index) is not int
                    or not 0 <= index < source["quantity"]
                    or any(
                        type(column) is not int or not 0 <= column < source["columns"]
                        for column in columns
                    )
                ):
                    raise ValueError("Producer projection lies outside its source")
                if columns:
                    selections.append(
                        {
                            "source_id": source["source_id"],
                            "source_sha256": source["sha256"],
                            "unit": source["unit"],
                            "start": index,
                            "stop": index + 1,
                            "columns": columns,
                            "basis": "vera_producer_projection",
                        }
                    )
        store.event(
            "projection:" + (artifact_sha256 or digest(value)),
            {
                "kind": "prepared_projection",
                "artifact_sha256": artifact_sha256,
                "data_sha256": digest(value),
                "selections": selections,
            },
        )
